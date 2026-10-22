"""Populate ten years of tax reconciliation into the Industrial Template Inputs tax table."""

from __future__ import annotations

from pathlib import Path
from typing import Any

from models.annual_update import AnnualTaxReport
from models.new_company import NewCompanyTaxReport, TaxDisclosureStatus, TaxYearResult
from services.annual_tax_research_service import AnnualTaxResearchService
from services.annual_tax_service import AnnualTaxService
from services.new_company_tax_table_service import NewCompanyTaxTableService


class NewCompanyTaxService:
    """Reuse AnnualTaxService mapping per year; do not treat a lone ETR as resolved."""

    def __init__(self) -> None:
        self.annual_tax = AnnualTaxService()
        self.research = AnnualTaxResearchService()
        self.tables = NewCompanyTaxTableService()

    def apply(
        self,
        *,
        analysis_id: str,
        ticker: str,
        workbook_path: Path,
        fiscal_years: list[str],
        company_facts: dict[str, Any] | None = None,
        year_inputs: dict[str, dict[str, Any]] | None = None,
    ) -> NewCompanyTaxReport:
        years: list[TaxYearResult] = []
        warnings: list[str] = []
        for fy in fiscal_years:
            supplied = dict((year_inputs or {}).get(fy) or {})
            researched: dict[str, Any] = {}
            components_present = bool(supplied.get("components"))
            if company_facts and not components_present:
                researched, _questions = self.research.extract_tax_inputs(
                    company_facts=company_facts,
                    fiscal_year=fy,
                    ticker=ticker,
                )
            table_payload: dict[str, Any] = {}
            if supplied.get("candidate_table"):
                table_payload = self.tables.extract(
                    fiscal_year=fy,
                    candidate_table=list(supplied["candidate_table"]),
                    heading=supplied.get("heading"),
                    prior_schema=supplied.get("prior_schema"),
                    accession_number=supplied.get("accession_number"),
                    note_location=supplied.get("note_location") or supplied.get("source"),
                    pretax_income=supplied.get("pretax_income") or researched.get("pretax_income"),
                    income_tax_expense=supplied.get("income_tax_expense")
                    or researched.get("income_tax_expense"),
                    llm_classifier=supplied.get("llm_classifier"),
                )
            payload = {
                **researched,
                **{k: v for k, v in supplied.items() if v is not None and k != "candidate_table"},
            }
            if table_payload:
                payload = {**payload, **{k: v for k, v in table_payload.items() if v is not None}}
            report: AnnualTaxReport = self.annual_tax.apply(
                analysis_id=analysis_id,
                ticker=ticker,
                workbook_path=workbook_path,
                fiscal_year=fy,
                reported_effective_rate=payload.get("reported_effective_rate"),
                filing_components=payload.get("components"),
                source=payload.get("source") or payload.get("etr_source"),
                source_locations=payload.get("locations") or payload.get("source_locations"),
                pretax_income=payload.get("pretax_income"),
                income_tax_expense=payload.get("income_tax_expense"),
            )
            mapped = {c.house_category: c.rate for c in report.mapped_components}
            computed = payload.get("computed_effective_rate")
            if computed is None and report.pretax_income and abs(report.pretax_income) > 1e-9 and report.income_tax_expense is not None:
                computed = report.income_tax_expense / report.pretax_income
            disclosure = payload.get("disclosure_status") or self._disclosure(mapped)
            statutory = mapped.get("statutory_federal")
            category_complete = statutory is not None and report.reported_effective_tax_rate is not None
            year = TaxYearResult(
                fiscal_year=fy,
                reported_effective_rate=report.reported_effective_tax_rate,
                computed_effective_rate=computed,
                statutory_federal=statutory,
                state=mapped.get("state"),
                foreign=mapped.get("foreign"),
                rd_credit=mapped.get("credits"),
                other=mapped.get("other"),
                residual_other=report.residual_other,
                reported_other=payload.get("reported_other"),
                unmapped_lines=list(payload.get("unmapped_lines") or []),
                category_mappings=list(payload.get("category_mappings") or []),
                disclosure_status=disclosure,
                other_calculation=payload.get("other_calculation")
                or (
                    "Other = Effective tax rate − statutory/federal − state − foreign − R&D tax-credit rate"
                    if report.residual_other is not None
                    else None
                ),
                component_sum=sum(
                    (mapped.get(k) or 0.0)
                    for k in ("statutory_federal", "state", "foreign", "credits", "other")
                ),
                reconciliation_delta=report.reconciliation_delta,
                reconciliation_status=payload.get("reconciliation_status") or report.reconciliation_status,
                pretax_income=report.pretax_income,
                income_tax_expense=report.income_tax_expense,
                units="rate_fraction",
                table_units=payload.get("table_units"),
                signs_normalized=bool(payload.get("signs_normalized")),
                cells_written=report.cells_written,
                schedule_populated=report.schedule_populated and category_complete,
                raw_filing_lines=list(payload.get("raw_filing_lines") or payload.get("components") or []),
                normalized_signed_values={
                    "statutory_federal": statutory,
                    "state": mapped.get("state"),
                    "foreign": mapped.get("foreign"),
                    "credits": mapped.get("credits"),
                    "other": mapped.get("other"),
                    "effective_tax_rate": report.reported_effective_tax_rate,
                },
                source=report.annual_report_source,
                source_locations=report.source_locations,
                accession_number=payload.get("accession_number")
                or next(
                    (
                        c.get("accession")
                        for c in (payload.get("components") or [])
                        if isinstance(c, dict) and c.get("accession")
                    ),
                    None,
                ),
                note_or_table_location=payload.get("note_or_table_location") or payload.get("source"),
                confidence=float(payload.get("confidence") or report.confidence or 0.0),
                category_coverage_complete=category_complete,
            )
            if not year.category_coverage_complete:
                warnings.append(
                    f"TEN_YEAR_TAX_COVERAGE_INCOMPLETE: {fy} missing statutory/federal or ETR "
                    "(ETR-only extraction is not coverage)."
                )
            if not year.schedule_populated:
                warnings.append(f"TEN_YEAR_TAX_COVERAGE_INCOMPLETE: {fy} schedule not populated")
            if year.reconciliation_status not in {"ok", ""}:
                warnings.append(f"{year.reconciliation_status}: {fy}")
            if year.reported_effective_rate is not None and year.statutory_federal is None:
                warnings.append(
                    f"TEN_YEAR_TAX_COVERAGE_INCOMPLETE: {fy} ETR found but statutory/federal not mapped"
                )
            years.append(year)

        derived = self._fill_derived_rates(workbook_path, years)
        if derived:
            warnings.append(
                "TAX_RATE_DERIVED_FROM_STATEMENTS: " + ", ".join(derived) + " (no filing reconciliation table: effective rate = 1 - net income / pre-tax income, "
                "shown as the federal rate; confirm)"
            )
        complete = all(y.category_coverage_complete and y.cells_written for y in years) if years else False
        return NewCompanyTaxReport(
            analysis_id=analysis_id,
            ticker=ticker,
            years=years,
            complete=complete,
            warnings=warnings,
            summary=(
                f"Tax: {sum(1 for y in years if y.category_coverage_complete)}/{len(years)} years "
                f"with statutory+ETR coverage; warnings={len(warnings)}."
            ),
        )

    @staticmethod
    def _fill_derived_rates(workbook_path: Path, years: list) -> list[str]:
        """Fill Inputs rows 107-112 for years the filings gave no table for, so NOPAT is taxed instead of treated as untaxed.

        Only blank years are touched. Effective rate = 1 - consolidated net income / pre-tax income (Inputs rows 38 and 37); in percentage mode the rate,
        in USD-millions mode the tax charge, goes to the federal row and the total, and the other rows are 0. Each change is written to the HAP Adjustments tab."""
        from openpyxl import load_workbook

        from services.adjustment_ledger_service import AdjustmentLedger

        wb = load_workbook(workbook_path)
        done: list[str] = []
        try:
            if "Inputs" not in wb.sheetnames:
                return done
            ws = wb["Inputs"]
            cols = {str(ws.cell(1, c).value or "").replace(" ", "").upper(): c for c in range(3, (ws.max_column or 3) + 1)}
            try:                                              # the Inputs headers are links to the statements: read the years the way the mapping does
                from openpyxl.utils import column_index_from_string

                from services.annual_period_service import detect_workbook_years

                for fy, col in detect_workbook_years(workbook_path).items():
                    cols[str(fy).replace(" ", "").upper()] = column_index_from_string(col) if isinstance(col, str) else int(col)
            except Exception:  # noqa: BLE001 - fall back to the header text
                pass
            percent_mode = int(ws["E106"].value or 0) == 1
            ledger = None
            for y in years:
                col = cols.get(str(y.fiscal_year).replace(" ", "").upper())
                if not col or any(ws.cell(r, col).value not in (None, "") for r in range(107, 113)):
                    continue
                pretax, net = ws.cell(37, col).value, ws.cell(38, col).value
                present = lambda v: isinstance(v, (int, float)) or (isinstance(v, str) and v.startswith("="))   # the template links these rows to the statements
                if not present(pretax) or not present(net):
                    continue
                if isinstance(pretax, (int, float)) and pretax <= 0:
                    continue
                rate = (1 - net / pretax) if isinstance(pretax, (int, float)) and isinstance(net, (int, float)) else None
                if rate is not None and not 0 <= rate <= 0.6:
                    continue
                letter = ws.cell(1, col).column_letter
                # percentage mode: the rate itself; USD-millions mode: the tax charge (pre-tax income - net income), the sheet divides it by EBT.
                # Guarded so a loss year or a missing statement cell gives 0 instead of an error.
                formula = (f"=IFERROR(MAX(0,MIN(0.6,1-{letter}38/{letter}37)),0)" if percent_mode
                           else f"=IFERROR(MAX(0,{letter}37-{letter}38),0)")
                ws.cell(112, col).value = formula
                ws.cell(107, col).value = f"={letter}112"
                for r in range(108, 112):
                    ws.cell(r, col).value = 0
                ledger = ledger or AdjustmentLedger(wb)
                ledger.record(
                    sheet="Inputs", cell=f"{letter}112", fiscal_year=y.fiscal_year, category="Missing figure",
                    what="Effective tax rate derived from the statements", original="blank", new=formula + (f" ({rate:.1%})" if rate is not None else ""), amount=rate,
                    reason="No rate-reconciliation table was found for this year. Effective rate = 1 - net income / pre-tax income, shown as federal "
                           "(state, foreign, credits and other = 0). Replace with the 10-K reconciliation when available.",
                    source="workbook statements", method="derived", confidence="medium", mark_cell=False, status="Derived",
                )
                done.append(str(y.fiscal_year))
            if done:
                wb.save(workbook_path)
            return done
        finally:
            wb.close()

    @staticmethod
    def _disclosure(mapped: dict[str, float | None]) -> dict[str, str]:
        out: dict[str, str] = {}
        for house in ("statutory_federal", "state", "foreign", "credits", "other"):
            val = mapped.get(house)
            if val is None:
                out[house] = (
                    TaxDisclosureStatus.INCLUDED_IN_OTHER.value
                    if mapped.get("other") is not None
                    else TaxDisclosureStatus.UNRESOLVED.value
                )
            elif val == 0:
                out[house] = TaxDisclosureStatus.REPORTED_ZERO.value
            else:
                out[house] = TaxDisclosureStatus.REPORTED.value
        return out
