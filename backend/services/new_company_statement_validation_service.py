"""Validate ten years of Bloomberg-prefilled statements against SEC filings."""

from __future__ import annotations

from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

from models.new_company import NewCompanyStatementValidationReport, StatementDiscrepancy
from services.annual_period_service import detect_year_columns
from services.sec_service import SecService

# concept, sheet, label needles, xbrl tags, scale_if_large
_CHECKS: list[tuple[str, str, tuple[str, ...], tuple[str, ...], bool]] = [
    ("revenue", "Income - GAAP", ("revenue", "sales"), ("Revenues", "RevenueFromContractWithCustomerExcludingAssessedTax", "SalesRevenueNet"), True),
    ("gross_profit", "Income - GAAP", ("gross profit",), ("GrossProfit",), True),
    ("operating_income", "Income - GAAP", ("operating income", "operating income (loss)"), ("OperatingIncomeLoss",), True),
    ("pretax_income", "Income - GAAP", ("pretax income", "income before tax", "income before income taxes"), ("IncomeLossFromContinuingOperationsBeforeIncomeTaxesExtraordinaryItemsNoncontrollingInterest", "IncomeLossFromContinuingOperationsBeforeIncomeTaxes"), True),
    ("income_tax_expense", "Income - GAAP", ("income tax expense", "provision for income taxes"), ("IncomeTaxExpenseBenefit",), True),
    ("net_income", "Income - GAAP", ("net income", "net income (loss)"), ("NetIncomeLoss",), True),
    ("diluted_eps", "Income - GAAP", ("diluted eps", "earnings per share diluted", "eps - diluted"), ("EarningsPerShareDiluted",), False),
    ("cfo", "Cash Flow - Standardized", ("cash from operating", "operating activities"), ("NetCashProvidedByUsedInOperatingActivities",), True),
    ("capex", "Cash Flow - Standardized", ("capital expenditure", "capex", "acq of fixed", "purchase of ppe"), ("PaymentsToAcquirePropertyPlantAndEquipment",), True),
    ("cfi", "Cash Flow - Standardized", ("cash from investing", "investing activities"), ("NetCashProvidedByUsedInInvestingActivities",), True),
    ("cff", "Cash Flow - Standardized", ("cash from financing", "financing activities"), ("NetCashProvidedByUsedInFinancingActivities",), True),
    ("cash", "Balance Sheet - Standardized", ("cash and cash equivalents", "cash"), ("CashAndCashEquivalentsAtCarryingValue",), True),
    ("total_assets", "Balance Sheet - Standardized", ("total assets",), ("Assets",), True),
    ("current_liabilities", "Balance Sheet - Standardized", ("total current liabilities", "current liabilities"), ("LiabilitiesCurrent",), True),
    ("total_liabilities", "Balance Sheet - Standardized", ("total liabilities",), ("Liabilities",), True),
    ("equity", "Balance Sheet - Standardized", ("total equity", "total shareholders", "stockholders' equity"), ("StockholdersEquity",), True),
    ("debt", "Balance Sheet - Standardized", ("total debt", "long-term debt", "debt"), ("LongTermDebt", "LongTermDebtNoncurrent"), True),
    ("diluted_shares", "Income - GAAP", ("diluted weighted", "weighted average shares diluted", "diluted shares"), ("WeightedAverageNumberOfDilutedSharesOutstanding",), True),
]

_REL = 0.03
_ABS = 25.0
_MATERIAL_REL = 0.05


def _num(v: Any) -> float | None:
    if v is None or v == "" or isinstance(v, bool):
        return None
    if isinstance(v, (int, float)):
        return float(v)
    try:
        return float(str(v).replace(",", "").replace("(", "-").replace(")", ""))
    except (TypeError, ValueError):
        return None


def _is_formula(v: Any) -> bool:
    return isinstance(v, str) and v.startswith("=")


def _scale_usd(value: float, *, scale: bool) -> float:
    if not scale:
        return value
    if abs(value) >= 10_000:
        return value / 1_000_000.0
    return value


class NewCompanyStatementValidationService:
    """SEC is authoritative. Fill missing values; correct only with filing evidence."""

    def validate(
        self,
        *,
        analysis_id: str,
        ticker: str,
        workbook_path: Path,
        fiscal_years: list[str],
        company_facts: dict[str, Any] | None = None,
        filing_overrides: dict[str, dict[str, float]] | None = None,
    ) -> NewCompanyStatementValidationReport:
        wb = load_workbook(workbook_path, data_only=False)
        sec = SecService()
        items: list[dict[str, Any]] = []
        discrepancies: list[StatementDiscrepancy] = []
        filled: list[StatementDiscrepancy] = []
        corrections: list[StatementDiscrepancy] = []
        unresolved: list[str] = []
        formulas_ok = True
        try:
            for fy in fiscal_years:
                for concept, sheet, needles, tags, scale in _CHECKS:
                    bb = None
                    cell_addr = None
                    formula = False
                    row = col = None
                    if sheet in wb.sheetnames:
                        ws = wb[sheet]
                        cols = detect_year_columns(ws, wb)
                        col = cols.get(fy)
                        row = self._find_row(ws, needles)
                        if col and row:
                            val = ws.cell(row, col).value
                            formula = _is_formula(val)
                            if formula:
                                formulas_ok = formulas_ok and True
                            else:
                                bb = _num(val)
                            cell_addr = f"{sheet}!{get_column_letter(col)}{row}"

                    filing = None
                    source = None
                    if filing_overrides and fy in filing_overrides and concept in filing_overrides[fy]:
                        filing = float(filing_overrides[fy][concept])
                        source = "filing_override"
                    elif company_facts:
                        for tag in tags:
                            fact = sec.find_fact(company_facts, concept, fy, xbrl_tag_hint=tag)
                            if fact is not None and fact.value is not None:
                                filing = _scale_usd(float(fact.value), scale=scale)
                                source = f"SEC {fact.form} {tag} accn={fact.accession_number}"
                                break

                    status = "validated"
                    action = "preserve"
                    if formula:
                        status = "formula_preserved"
                    elif bb is None and filing is not None and row and col and sheet in wb.sheetnames:
                        ws = wb[sheet]
                        cell = ws.cell(row, col)
                        if not _is_formula(cell.value):
                            cell.value = filing
                            action = "fill_missing"
                            status = "filled_from_sec"
                            rec = StatementDiscrepancy(
                                fiscal_year=fy,
                                statement=sheet,
                                concept=concept,
                                sheet=sheet,
                                cell=cell_addr,
                                old_value=None,
                                new_value=filing,
                                source=source,
                                reason="Missing Bloomberg value filled from SEC filing.",
                                impact="Enables downstream formulas that require this concept.",
                                action=action,
                                material=True,
                            )
                            filled.append(rec)
                            bb = filing
                    elif bb is not None and filing is not None:
                        if not self._close(bb, filing):
                            material = abs(bb - filing) / max(abs(filing), 1.0) > _MATERIAL_REL
                            rec = StatementDiscrepancy(
                                fiscal_year=fy,
                                statement=sheet,
                                concept=concept,
                                sheet=sheet,
                                cell=cell_addr,
                                old_value=bb,
                                new_value=filing,
                                source=source,
                                reason="Bloomberg prefilled value differs from SEC-reported amount.",
                                impact="May affect ROIC, tax, and valuation if uncorrected.",
                                action="correct" if material else "preserve",
                                material=material,
                            )
                            discrepancies.append(rec)
                            if material and row and col and sheet in wb.sheetnames:
                                ws = wb[sheet]
                                cell = ws.cell(row, col)
                                if not _is_formula(cell.value):
                                    cell.value = filing
                                    rec.action = "correct"
                                    corrections.append(rec)
                                    bb = filing
                                    status = "corrected_from_sec"
                                else:
                                    unresolved.append(f"{fy}:{concept}")
                                    status = "formula_blocked_correction"
                            else:
                                rec.action = "preserve"
                                status = "immaterial_difference"
                        else:
                            status = "validated"
                    elif bb is None and filing is None:
                        status = "missing_both"
                    elif bb is None:
                        status = "missing_workbook"
                    else:
                        status = "sec_unavailable_preserved"

                    items.append(
                        {
                            "fiscal_year": fy,
                            "concept": concept,
                            "bloomberg_value": bb,
                            "filing_value": filing,
                            "status": status,
                            "cell": cell_addr,
                            "source": source,
                        }
                    )
            wb.save(workbook_path)
        finally:
            wb.close()

        if unresolved:
            status = "blocked"
        elif corrections or discrepancies:
            status = "reviewed"
        else:
            status = "ok"
        return NewCompanyStatementValidationReport(
            analysis_id=analysis_id,
            ticker=ticker,
            items=items,
            discrepancies=discrepancies,
            filled_missing=filled,
            corrections=corrections,
            unresolved_material=unresolved,
            formulas_preserved=formulas_ok,
            status=status,
            summary=(
                f"Statements: {len(items)} checks; filled={len(filled)}; "
                f"corrections={len(corrections)}; discrepancies={len(discrepancies)}; "
                f"unresolved={len(unresolved)}."
            ),
        )

    @staticmethod
    def _find_row(ws, needles: tuple[str, ...]) -> int | None:
        for row in range(1, min(ws.max_row or 1, 160) + 1):
            lab = str(ws.cell(row, 1).value or "").strip().lower()
            if not lab or lab.startswith("+") or lab.startswith("-"):
                continue
            if any(n == lab or n in lab for n in needles):
                return row
        return None

    @staticmethod
    def _close(a: float, b: float) -> bool:
        if abs(a - b) <= _ABS:
            return True
        denom = max(abs(b), abs(a), 1e-9)
        return abs(a - b) / denom <= _REL
