"""Annual workbook lineage and anti-previous-file guard."""

from __future__ import annotations

from pathlib import Path

from openpyxl import load_workbook

from models.annual_update import AnnualBaseWorkbookGuardReport, AnnualWorkbookLineageReport
from services.annual_period_service import AnnualPeriodAlignment, year_columns_for_sheet
from services.annual_update_runner import sha256_file


def build_lineage_report(
    *,
    analysis_id: str,
    ticker: str,
    template_path: Path,
    previous_path: Path,
    working_path: Path,
    final_path: Path,
    alignment: AnnualPeriodAlignment,
    continuity_summary: str = "",
) -> AnnualWorkbookLineageReport:
    working_initial_sha = sha256_file(working_path)
    return AnnualWorkbookLineageReport(
        analysis_id=analysis_id,
        ticker=ticker,
        declared_base="CURRENT_TEMPLATE",
        current_template_filename=template_path.name,
        current_template_sha256=sha256_file(template_path),
        previous_completed_filename=previous_path.name,
        previous_completed_sha256=sha256_file(previous_path),
        working_workbook_initial_sha256=working_initial_sha,
        final_workbook_sha256=sha256_file(final_path),
        template_years=list(alignment.template_years),
        previous_years=list(alignment.previous_years),
        overlap_years=list(alignment.overlap_years),
        dropped_years=list(alignment.dropped_years),
        new_fiscal_year=alignment.new_fiscal_year,
        template_end_year=alignment.template_end_year,
        previous_end_year=alignment.previous_end_year,
        continuity_summary=continuity_summary,
        summary=(
            f"Base=CURRENT_TEMPLATE; new_fy={alignment.new_fiscal_year}; "
            f"template_years={alignment.template_years[0]}..{alignment.template_years[-1]}"
        ),
    )


def assert_base_workbook_guard(
    *,
    analysis_id: str,
    ticker: str,
    template_path: Path,
    previous_path: Path,
    final_path: Path,
    alignment: AnnualPeriodAlignment,
) -> AnnualBaseWorkbookGuardReport:
    template_sha = sha256_file(template_path)
    previous_sha = sha256_file(previous_path)
    final_sha = sha256_file(final_path)
    violations: list[str] = []

    if final_sha == previous_sha:
        violations.append("final workbook SHA-256 identical to previous completed workbook")

    if final_sha == template_sha:
        # Allowed only if continuity made no changes — still suspicious for annual update
        pass

    new_fy = alignment.new_fiscal_year
    tmpl_cols = year_columns_for_sheet(template_path, "Income - GAAP")
    prev_cols = year_columns_for_sheet(previous_path, "Income - GAAP")
    final_cols = year_columns_for_sheet(final_path, "Income - GAAP")
    new_col_t = tmpl_cols.get(new_fy)
    new_col_p = prev_cols.get(new_fy)
    new_col_f = final_cols.get(new_fy)

    template_new_rev = None
    previous_new_rev = None
    final_new_rev = None
    if new_col_t:
        wb = load_workbook(template_path, data_only=True)
        try:
            template_new_rev = wb["Income - GAAP"].cell(11, new_col_t).value
        finally:
            wb.close()
    if new_col_p:
        wb = load_workbook(previous_path, data_only=True)
        try:
            previous_new_rev = wb["Income - GAAP"].cell(11, new_col_p).value
        finally:
            wb.close()
    if new_col_f:
        wb = load_workbook(final_path, data_only=True)
        try:
            final_new_rev = wb["Income - GAAP"].cell(11, new_col_f).value
        finally:
            wb.close()

    if new_col_t is None:
        violations.append(f"new fiscal year {new_fy} missing from template Income statement")
    if template_new_rev is not None and final_new_rev is not None:
        if template_new_rev != final_new_rev:
            violations.append(
                f"new-year Bloomberg revenue changed: template={template_new_rev} final={final_new_rev}"
            )
    if (
        previous_new_rev is not None
        and final_new_rev is not None
        and previous_new_rev == final_new_rev
        and template_new_rev is not None
        and template_new_rev != previous_new_rev
    ):
        violations.append(
            "final new-year revenue matches previous workbook instead of current template Bloomberg data"
        )

    status = "ok" if not violations else "ANNUAL_BASE_WORKBOOK_VIOLATION"
    return AnnualBaseWorkbookGuardReport(
        analysis_id=analysis_id,
        ticker=ticker,
        status=status,
        template_sha256=template_sha,
        previous_sha256=previous_sha,
        final_sha256=final_sha,
        new_fiscal_year=new_fy,
        template_new_year_revenue=template_new_rev,
        previous_new_year_revenue=previous_new_rev,
        final_new_year_revenue=final_new_rev,
        violations=violations,
        summary=status if not violations else "; ".join(violations),
    )
