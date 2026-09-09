"""Quarter-scoped comparison review (not a 10-year analyst review)."""

from __future__ import annotations

from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from models.quarterly_presentation import STATEMENT_SHEETS, QuarterlyStatementKind
from models.quarterly_update import QuarterlyComparisonEntry, QuarterlyReviewReport

LQ_BS = STATEMENT_SHEETS[QuarterlyStatementKind.BALANCE_SHEET]
LQ_IS = STATEMENT_SHEETS[QuarterlyStatementKind.INCOME]
LQ_CF = STATEMENT_SHEETS[QuarterlyStatementKind.CASH_FLOW]

# Industrial Template LQ layout: C = current, D = comparison period (QoQ BS / YoY IS / YTD CF).
# IS also has G/H = YTD vs prior-year YTD.
CURR_COL = 3
PRIOR_COL = 4
YTD_CURR_COL = 7
YTD_PRIOR_COL = 8

_BS_NEEDLES = [
    ("cash", "Cash"),
    ("inventor", "Inventory"),
    ("receiv", "Receivables"),
    ("total current assets", "Current assets"),
    ("property, plant", "PP&E"),
    ("goodwill", "Goodwill"),
    ("total assets", "Total assets"),
    ("accounts payable", "Accounts payable"),
    ("st debt", "Short-term debt"),
    ("lt debt", "Long-term debt"),
    ("total current liabilities", "Current liabilities"),
    ("total liabilities", "Total liabilities"),
    ("total equity", "Equity"),
]
_IS_NEEDLES = [
    ("revenue", "Revenue"),
    ("gross profit", "Gross Profit"),
    ("operating income", "Operating Income"),
    ("net income", "Net Income"),
    ("gross margin", "Gross Margin"),
    ("operating margin", "Operating Margin"),
    ("net margin", "Net Margin"),
]
_CF_NEEDLES = [
    ("cash from operating", "CFO"),
    ("cash from investing", "CFI"),
    ("cash from financing", "CFF"),
    ("acq of fixed", "CapEx"),
]


def _num(v: Any) -> float | None:
    if v is None or isinstance(v, bool):
        return None
    if isinstance(v, (int, float)):
        return float(v)
    try:
        return float(v)
    except (TypeError, ValueError):
        return None


def _chg(a: float | None, b: float | None) -> tuple[float | None, float | None]:
    if a is None or b is None:
        return None, None
    abs_c = b - a
    pct = abs_c / abs(a) if a != 0 else None
    return abs_c, pct


def _norm(label: Any) -> str:
    return " ".join(str(label or "").lower().split())


def _find_rows(ws, needles: list[tuple[str, str]]) -> list[tuple[int, str]]:
    found: list[tuple[int, str]] = []
    for needle, name in needles:
        candidates: list[tuple[int, bool, bool]] = []  # row, unindented, populated
        for r in range(11, 130):
            label = ws.cell(r, 1).value
            if label is None:
                continue
            if needle not in _norm(label):
                continue
            populated = _num(ws.cell(r, CURR_COL).value) is not None
            unindented = not str(label).startswith(" ")
            candidates.append((r, unindented, populated))
        if not candidates:
            continue
        candidates.sort(key=lambda t: (not t[2], not t[1], t[0]))
        found.append((candidates[0][0], name))
    return found


def _detect_fiscal_quarter(ws) -> int | None:
    for r in range(1, 10):
        for c in range(1, 8):
            v = ws.cell(r, c).value
            if isinstance(v, str) and "Q" in v:
                for q in (1, 2, 3, 4):
                    if f"Q{q}" in v.upper().replace(" ", ""):
                        return q
    return None


class QuarterlyReviewService:
    """Build quarter-focused comparisons for BS QoQ, IS YoY+YTD, CF YTD."""

    def review(
        self,
        *,
        analysis_id: str,
        ticker: str,
        workbook_path: Path,
        fiscal_quarter: int | None = None,
    ) -> QuarterlyReviewReport:
        wb = load_workbook(workbook_path, data_only=True)
        comparisons: list[QuarterlyComparisonEntry] = []
        material: list[str] = []

        q = fiscal_quarter
        if q is None and LQ_IS in wb.sheetnames:
            q = _detect_fiscal_quarter(wb[LQ_IS])
        q = q or 3
        ytd_label = {1: "3M", 2: "6M_YTD", 3: "9M_YTD", 4: "12M_YTD"}.get(q, "YTD")

        if LQ_BS in wb.sheetnames:
            ws = wb[LQ_BS]
            for row, name in _find_rows(ws, _BS_NEEDLES):
                prev = _num(ws.cell(row, PRIOR_COL).value)
                curr = _num(ws.cell(row, CURR_COL).value)
                abs_c, pct = _chg(prev, curr)
                comparisons.append(
                    QuarterlyComparisonEntry(
                        statement="balance_sheet",
                        metric=name,
                        baseline_period="previous_quarter",
                        compare_period="latest_quarter",
                        baseline_value=prev,
                        compare_value=curr,
                        absolute_change=abs_c,
                        percentage_change=pct,
                        comparison_type="qoq",
                        note="Latest quarter (col C) vs immediately previous quarter (col D)",
                    )
                )
                if pct is not None and abs(pct) >= 0.10 and abs(abs_c or 0) >= 100:
                    material.append(f"BS QoQ {name}: {pct:.1%}")

        if LQ_IS in wb.sheetnames:
            ws = wb[LQ_IS]
            for row, name in _find_rows(ws, _IS_NEEDLES):
                prior = _num(ws.cell(row, PRIOR_COL).value)
                curr = _num(ws.cell(row, CURR_COL).value)
                abs_c, pct = _chg(prior, curr)
                comparisons.append(
                    QuarterlyComparisonEntry(
                        statement="income_statement",
                        metric=name,
                        baseline_period="prior_year_same_quarter",
                        compare_period="current_quarter",
                        baseline_value=prior,
                        compare_value=curr,
                        absolute_change=abs_c,
                        percentage_change=pct,
                        comparison_type="yoy_quarter",
                        note="Current 3-month period vs same 3-month period one year earlier (C vs D)",
                    )
                )
                if pct is not None and abs(pct) >= 0.10:
                    material.append(f"IS YoY {name}: {pct:.1%}")

                ytd_prior = _num(ws.cell(row, YTD_PRIOR_COL).value)
                ytd_curr = _num(ws.cell(row, YTD_CURR_COL).value)
                if ytd_curr is not None or ytd_prior is not None:
                    abs_y, pct_y = _chg(ytd_prior, ytd_curr)
                    comparisons.append(
                        QuarterlyComparisonEntry(
                            statement="income_statement",
                            metric=name,
                            baseline_period=f"prior_year_{ytd_label}",
                            compare_period=f"current_{ytd_label}",
                            baseline_value=ytd_prior,
                            compare_value=ytd_curr,
                            absolute_change=abs_y,
                            percentage_change=pct_y,
                            comparison_type="ytd",
                            note=f"YTD {ytd_label} vs prior-year equivalent (G vs H) — not vs prior quarter",
                        )
                    )

        if LQ_CF in wb.sheetnames:
            ws = wb[LQ_CF]
            for row, name in _find_rows(ws, _CF_NEEDLES):
                prior = _num(ws.cell(row, PRIOR_COL).value)
                curr = _num(ws.cell(row, CURR_COL).value)
                abs_c, pct = _chg(prior, curr)
                comparisons.append(
                    QuarterlyComparisonEntry(
                        statement="cash_flow",
                        metric=name,
                        baseline_period=f"prior_year_{ytd_label}",
                        compare_period=f"current_{ytd_label}",
                        baseline_value=prior,
                        compare_value=curr,
                        absolute_change=abs_c,
                        percentage_change=pct,
                        comparison_type="ytd",
                        note="Cumulative/YTD like-for-like (C vs D) — standalone quarter CF not invented",
                    )
                )
                if pct is not None and abs(pct) >= 0.15:
                    material.append(f"CF YTD {name}: {pct:.1%}")

        wb.close()
        return QuarterlyReviewReport(
            analysis_id=analysis_id,
            ticker=ticker,
            comparisons=comparisons,
            material_changes=material,
            summary=(
                f"Quarter review: {len(comparisons)} comparisons; "
                f"{len(material)} material movement(s). "
                "BS=QoQ (C vs D); IS=YoY FQ (C vs D) + YTD (G vs H); CF=YTD (C vs D)."
            ),
        )
