"""Lightweight annual restatement detection for quarterly_update.

Does not rewrite history. If material restatement found → RESTATEMENT_REVIEW_REQUIRED
and does NOT trigger full annual validation in the quarter path.
"""

from __future__ import annotations

from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from models.quarterly_update import RestatementCheckReport, RestatementFinding
from workbook_mapping.explicit_mappings import ANNUAL_PERIOD_COLS, ANNUAL_PERIOD_FY_TOKENS

# Spot-check a few material annual lines (USD millions in workbook)
_CHECK_LINES = [
    ("income_statement", "Revenue", "Income - GAAP", 9, "revenue", None),
    ("income_statement", "Net Income", "Income - GAAP", 58, "net income", None),
    ("balance_sheet", "Total Assets", "Balance Sheet - Standardized", 61, "total assets", None),
]

_REL_TOL = 0.02  # 2%
_ABS_TOL_M = 50.0  # USD millions


def _num(v: Any) -> float | None:
    if v is None or v == "":
        return None
    if isinstance(v, bool):
        return None
    if isinstance(v, (int, float)):
        return float(v)
    try:
        return float(v)
    except (TypeError, ValueError):
        return None


class RestatementCheckService:
    """Compare prior workbook annual anchors to SEC for material revisions only."""

    def check(
        self,
        *,
        analysis_id: str,
        ticker: str,
        previous_workbook_path: Path,
        company_facts: dict[str, Any] | None,
        years: int = 3,
    ) -> RestatementCheckReport:
        if company_facts is None:
            return RestatementCheckReport(
                analysis_id=analysis_id,
                ticker=ticker,
                summary="Restatement check skipped — no companyfacts available.",
            )

        from services.sec_service import SecService

        sec = SecService()
        wb = load_workbook(previous_workbook_path, data_only=True)
        findings: list[RestatementFinding] = []

        # Check the most recent N fiscal years present in the template window
        cols = list(ANNUAL_PERIOD_COLS[-years:])
        fys = list(ANNUAL_PERIOD_FY_TOKENS[-years:])

        for statement, metric, sheet, row, concept, tag in _CHECK_LINES:
            if sheet not in wb.sheetnames:
                continue
            ws = wb[sheet]
            for col_letter, fy in zip(cols, fys, strict=True):
                col = ord(col_letter) - ord("A") + 1
                prev_val = _num(ws.cell(row, col).value)
                if prev_val is None:
                    continue
                fact = sec.find_fact(
                    company_facts,
                    concept,
                    fy,
                    xbrl_tag_hint=tag,
                )
                if fact is None or fact.value is None:
                    continue
                sec_m = float(fact.value) / 1_000_000.0
                abs_diff = abs(sec_m - prev_val)
                rel = abs_diff / max(abs(prev_val), 1e-9)
                if abs_diff <= _ABS_TOL_M or rel <= _REL_TOL:
                    continue
                # Ignore EPS-style split multiples elsewhere; these are millions-scale lines
                findings.append(
                    RestatementFinding(
                        statement=statement,
                        metric=metric,
                        period=fy,
                        previous_value=prev_val,
                        sec_value=sec_m,
                        absolute_difference=abs_diff,
                        percentage_difference=rel,
                        decision="RESTATEMENT_REVIEW_REQUIRED",
                        reason=(
                            f"{metric} {fy}: previous workbook {prev_val} vs SEC {sec_m:.1f} "
                            f"(Δ={abs_diff:.1f}, {rel:.1%}) — possible restatement."
                        ),
                    )
                )

        wb.close()
        material = len(findings) > 0
        return RestatementCheckReport(
            analysis_id=analysis_id,
            ticker=ticker,
            material_restatement_detected=material,
            findings=findings,
            annual_validation_triggered=False,  # quarter path never auto-launches full annual
            summary=(
                f"Restatement check: {len(findings)} material finding(s). "
                + (
                    "RESTATEMENT_REVIEW_REQUIRED — historical model left unchanged."
                    if material
                    else "No material restatement detected; annual history left untouched."
                )
            ),
        )
