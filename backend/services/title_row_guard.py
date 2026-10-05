"""Title rows in the Bloomberg statement tabs must stay blank.

A statement tab shows the same line twice: once as a section title (for example "Total Assets" in row 9, "Cash from Operating
Activities" in row 9 of the cash flow tab) and once as the real data row further down (BS_TOT_ASSET, CF_CASH_FROM_OPER, ...).
The real row carries a Bloomberg field code in column B; the title row does not. Numbers belong only on the real row.

- `is_title_row`: a row with a label, no field code in column B, and the same label on another row that does have a field code.
- `real_row_for_label`: the data row to use for a label (the one with a field code).
- `clear_title_rows`: safety net that removes numeric constants a writer put on title rows (formulas are never touched).
"""

from __future__ import annotations

from dataclasses import dataclass, field
from typing import Any

from openpyxl.worksheet.worksheet import Worksheet

STATEMENT_SHEETS = (
    "Balance Sheet - Standardized",
    "Cash Flow - Standardized",
    "Last Quarter BS Standardized",
    "Last Quarter CF Standardized",
)


def _norm(label: Any) -> str:
    return str(label or "").strip().lower().lstrip("+- ").strip()


def _code(ws: Worksheet, row: int) -> str:
    return str(ws.cell(row, 2).value or "").strip()


def title_rows(ws: Worksheet, *, max_row: int = 200) -> set[int]:
    """Rows that repeat a label of a later or earlier row that has a Bloomberg field code in column B."""
    limit = min(ws.max_row or 1, max_row)
    coded: set[str] = set()
    for row in range(1, limit + 1):
        label = _norm(ws.cell(row, 1).value)
        if label and label != "check" and _code(ws, row):
            coded.add(label)
    titles: set[int] = set()
    for row in range(1, limit + 1):
        label = _norm(ws.cell(row, 1).value)
        if label and label != "check" and not _code(ws, row) and label in coded:
            titles.add(row)
    return titles


def is_title_row(ws: Worksheet, row: int) -> bool:
    return row in title_rows(ws)


def real_row_for_label(ws: Worksheet, label: str) -> int | None:
    """The row with this label that carries a field code in column B (None when the tab has no codes)."""
    wanted = _norm(label)
    for row in range(1, min(ws.max_row or 1, 200) + 1):
        if _norm(ws.cell(row, 1).value) == wanted and _code(ws, row):
            return row
    return None


@dataclass
class TitleRowReport:
    cleared: list[str] = field(default_factory=list)

    @property
    def changed(self) -> bool:
        return bool(self.cleared)


def clear_title_rows(wb, sheets: tuple[str, ...] = STATEMENT_SHEETS) -> TitleRowReport:
    """Remove numbers from title rows (columns C onward). Dates and text headers are left alone."""
    report = TitleRowReport()
    for name in sheets:
        if name not in wb.sheetnames:
            continue
        ws = wb[name]
        for row in sorted(title_rows(ws)):
            for col in range(3, (ws.max_column or 3) + 1):
                cell = ws.cell(row, col)
                value = cell.value
                if isinstance(value, (int, float)) and not isinstance(value, bool):
                    cell.value = None
                    report.cleared.append(f"{name}!{cell.coordinate}")
    return report
