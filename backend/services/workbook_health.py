"""Read-only health checks on a completed workbook and its consistency with the SEC buyback evidence.

Never edits the workbook. Findings are warnings for the analyst, not pipeline blockers.
"""

from __future__ import annotations

import json
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

ERROR_VALUES = {"#DIV/0!", "#NAME?", "#REF!", "#VALUE!", "#N/A", "#NUM!", "#NULL!"}
# Sheets whose errors matter to every analysis type. "Last Quarter*" sheets only matter to quarterly updates.
CORE_SHEETS = {
    "Inputs", "All Ratios", "Final Metrics", "Expected Returns & Buybacks", "Enterprise Value",
    "IC & NOPAT & ROIC", "Tax", "Leases", "R&D",
}
BUYBACK_SHEET = "Income - GAAP"
BUYBACK_DOLLAR_LABEL = "$ paid for shares"
MATERIAL_RELATIVE = 0.05
MATERIAL_ABSOLUTE_M = 1.0


def scan_errors(path: Path, *, quarterly: bool = False) -> dict[str, Any]:
    values = load_workbook(path, data_only=True)
    formulas = load_workbook(path, data_only=False)
    try:
        sheets: dict[str, dict[str, Any]] = {}
        for ws in values.worksheets:
            wsf = formulas[ws.title]
            hits: list[tuple[str, str, bool]] = []
            for (row, col), cell in ws._cells.items():
                if isinstance(cell.value, str) and cell.value in ERROR_VALUES:
                    formula = wsf._cells.get((row, col))
                    text = str(formula.value) if formula is not None and formula.value is not None else ""
                    hits.append((cell.coordinate, cell.value, "_xll." in text))
            if hits:
                sheets[ws.title] = {
                    "count": len(hits),
                    "kinds": sorted({kind for _, kind, _ in hits}),
                    "examples": [coord for coord, _, _ in hits[:5]],
                    "bloomberg_addin": sum(1 for _, _, addin in hits if addin),
                }
        findings: list[str] = []
        for name, info in sheets.items():
            if name.lower().startswith("last quarter") and not quarterly:
                continue
            if name in CORE_SHEETS or name.startswith(("Income", "Balance", "Cash Flow")) or quarterly:
                where = ", ".join(info["examples"][:3])
                findings.append(f"{name}: {info['count']} error cell(s) {'/'.join(info['kinds'])} (e.g. {where}).")
        addin_total = sum(info["bloomberg_addin"] for info in sheets.values())
        notes = []
        if addin_total:
            notes.append(
                f"{addin_total} cell(s) call the Bloomberg add-in (_xll.BDP) and show #NAME? without it; "
                "dependent labels may show errors too."
            )
        return {"sheets": sheets, "findings": findings, "notes": notes}
    finally:
        values.close()
        formulas.close()


def buyback_consistency(workbook_path: Path, buyback_report: dict[str, Any]) -> dict[str, Any]:
    """Compare the workbook's '$ Paid for Shares' row with the SEC-derived annual dollars."""
    from services.annual_period_service import detect_year_columns

    wb = load_workbook(workbook_path, data_only=True)
    try:
        if BUYBACK_SHEET not in wb.sheetnames:
            return {"checked": False, "reason": f"No '{BUYBACK_SHEET}' sheet.", "mismatches": []}
        ws = wb[BUYBACK_SHEET]
        row = next(
            (r for r in range(1, min(ws.max_row or 1, 200) + 1)
             if BUYBACK_DOLLAR_LABEL in str(ws.cell(r, 1).value or "").strip().lower()),
            None,
        )
        if row is None:
            return {"checked": False, "reason": "Buyback dollars row not found.", "mismatches": []}
        columns = detect_year_columns(ws, wb)
        mismatches = []
        compared = 0
        for year in buyback_report.get("years", []):
            sec = year.get("dollars")
            col = columns.get(year.get("fiscal_year"))
            if sec is None or not col:
                continue
            value = ws.cell(row, col).value
            book = float(value) if isinstance(value, (int, float)) and not isinstance(value, bool) else 0.0
            compared += 1
            diff = abs(book - float(sec))
            if diff > MATERIAL_ABSOLUTE_M and diff / max(abs(book), abs(float(sec)), 1e-9) > MATERIAL_RELATIVE:
                mismatches.append({"fiscal_year": year["fiscal_year"], "workbook": round(book, 1), "sec": round(float(sec), 1)})
        return {"checked": True, "compared": compared, "mismatches": mismatches}
    finally:
        wb.close()


def health_report(output_dir: Path, *, quarterly: bool = False) -> dict[str, Any]:
    """Findings for the analysis in ``output_dir`` (completed_workbook.xlsx + optional buyback report)."""
    workbook = output_dir / "completed_workbook.xlsx"
    if not workbook.exists():
        return {"available": False, "findings": [], "notes": []}
    report = scan_errors(workbook, quarterly=quarterly)
    findings = list(report["findings"])
    buyback_path = output_dir / "new_company_buyback_report.json"
    consistency = None
    if buyback_path.exists():
        try:
            with buyback_path.open("r", encoding="utf-8") as handle:
                consistency = buyback_consistency(workbook, json.load(handle))
        except (OSError, ValueError):
            consistency = None
        if consistency and consistency["mismatches"]:
            detail = "; ".join(f"{m['fiscal_year']}: workbook {m['workbook']} vs SEC {m['sec']}" for m in consistency["mismatches"][:6])
            findings.append(
                f"Buybacks: the workbook disagrees with SEC data in {len(consistency['mismatches'])} year(s) ({detail}). "
                "The workbook may be from an older run; re-run the analysis."
            )
    return {"available": True, "findings": findings, "notes": report["notes"], "buyback": consistency, "sheets": report["sheets"]}
