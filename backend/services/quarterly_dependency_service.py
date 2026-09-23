"""Reconnect formulas when a quarterly statement layout changes.

Maps semantic financial fields, not raw row numbers. Updates only references
whose field moved. Leaves unrelated formulas untouched.
"""

from __future__ import annotations

import re
from typing import Any

from openpyxl.cell.cell import MergedCell
from openpyxl.workbook.workbook import Workbook

from models.quarterly_presentation import STATEMENT_SHEETS
from services.quarterly_health_service import BODY_END, BODY_START, LABEL_COL, VALUE_COL

# Longer needles first so "cost of revenue" is not classified as revenue.
_FIELD_NEEDLES: tuple[tuple[str, tuple[str, ...]], ...] = (
    ("cost_of_revenue", ("cost of revenue", "cost of sales", "cost of goods")),
    ("gross_profit", ("gross profit",)),
    ("operating_income", ("operating income", "operating profit")),
    ("net_income", ("net income", "net earnings", "net loss")),
    ("basic_eps", ("basic eps", "basic earnings per share", "earnings per share basic")),
    ("diluted_eps", ("diluted eps", "diluted earnings per share", "earnings per share diluted")),
    ("operating_cash_flow", ("cash from operating", "net cash provided by operating", "operating activities")),
    ("investing_cash_flow", ("cash from investing", "investing activities")),
    ("financing_cash_flow", ("cash from financing", "financing activities")),
    ("capex", ("capital expenditure", "purchase of property", "payments to acquire property")),
    ("cash", ("cash and cash equivalents", "cash & cash equivalents", "cash and equivalents")),
    ("revenue", ("total revenue", "net sales", "revenue", "sales")),
)

_QUALIFIED_REF = re.compile(
    r"(?:'(?P<qsheet>[^']+)'|(?P<bsheet>[A-Za-z0-9_]+))!"
    r"(?P<acol>\$?)(?P<col>[A-Z]{1,3})(?P<arow>\$?)(?P<row>\d+)"
)
_STATEMENT_SHEETS = set(STATEMENT_SHEETS.values())


def semantic_field(label: Any) -> str | None:
    text = " ".join(str(label or "").lower().replace("(ytd)", " ").split())
    if not text:
        return None
    if text.startswith("+") or text.startswith("-"):
        return "label:" + text
    adjusted_eps = "adjusted" in text or "non-gaap" in text or "non gaap" in text
    for field, needles in _FIELD_NEEDLES:
        if adjusted_eps and field in {"basic_eps", "diluted_eps"}:
            continue
        if any(needle in text for needle in needles):
            return field
    return "label:" + text


class _SheetMap:
    def __init__(self) -> None:
        self.cell_to_field: dict[str, str] = {}
        self.field_to_cell: dict[str, str] = {}


def _map_sheet(ws) -> _SheetMap:
    mapped = _SheetMap()
    for row in range(BODY_START, BODY_END + 1):
        label = ws.cell(row, LABEL_COL).value
        field = semantic_field(label)
        if not field:
            continue
        value = ws.cell(row, VALUE_COL).value
        if value in (None, "") and not (isinstance(value, str) and str(value).startswith("=")):
            # Keep a labeled row even when the value is still blank so a later
            # fill stays on the original cell.
            if label in (None, ""):
                continue
        addr = ws.cell(row, VALUE_COL).coordinate
        mapped.cell_to_field[addr] = field
        mapped.field_to_cell.setdefault(field, addr)
    return mapped


class QuarterlyDependencyService:
    def snapshot(self, workbook: Workbook) -> dict[str, Any]:
        semantics: dict[str, _SheetMap] = {}
        for sheet in _STATEMENT_SHEETS:
            if sheet in workbook.sheetnames:
                semantics[sheet] = _map_sheet(workbook[sheet])
        formulas: list[dict[str, str]] = []
        for ws in workbook.worksheets:
            max_row = ws.max_row or 1
            max_col = ws.max_column or 1
            for row in ws.iter_rows(min_row=1, max_row=max_row, max_col=max_col):
                for cell in row:
                    if isinstance(cell, MergedCell):
                        continue
                    value = cell.value
                    if not isinstance(value, str) or not value.startswith("="):
                        continue
                    if ws.title in _STATEMENT_SHEETS or any(name in value for name in _STATEMENT_SHEETS):
                        formulas.append({"sheet": ws.title, "cell": cell.coordinate, "formula": value})
        return {"semantics": semantics, "formulas": formulas}

    def reconnect(self, workbook: Workbook, snapshot: dict[str, Any]) -> dict[str, Any]:
        old_maps: dict[str, _SheetMap] = snapshot.get("semantics") or {}
        new_maps: dict[str, _SheetMap] = {}
        for sheet in old_maps:
            if sheet in workbook.sheetnames:
                new_maps[sheet] = _map_sheet(workbook[sheet])
        changes: list[dict[str, Any]] = []
        unresolved: list[dict[str, Any]] = []
        for item in snapshot.get("formulas") or []:
            sheet = item["sheet"]
            if sheet not in workbook.sheetnames:
                unresolved.append(
                    {
                        "sheet": sheet,
                        "cell": item["cell"],
                        "formula": item["formula"],
                        "reason": "Formula sheet removed during quarterly replacement.",
                    }
                )
                continue
            cell = workbook[sheet][item["cell"]]
            if isinstance(cell, MergedCell):
                continue
            current = cell.value
            if current != item["formula"]:
                # Layout replaced this cell. Dependents are updated from their own snapshots.
                continue
            updated, reasons, missing = self._rewrite_formula(
                formula=item["formula"],
                home_sheet=sheet,
                old_maps=old_maps,
                new_maps=new_maps,
            )
            for miss in missing:
                unresolved.append({"sheet": sheet, "cell": item["cell"], "formula": item["formula"], **miss})
            if updated != item["formula"] and reasons:
                cell.value = updated
                for reason in reasons:
                    changes.append(
                        {
                            "sheet": sheet,
                            "cell": item["cell"],
                            "original_formula": item["formula"],
                            "new_formula": updated,
                            "reason": reason["reason"],
                            "field": reason["field"],
                            "period": "latest_reported_quarter",
                            "original_reference": reason["original_reference"],
                            "new_reference": reason["new_reference"],
                        }
                    )
        return {"changed_formulas": changes, "unresolved_dependencies": unresolved}

    def _rewrite_formula(
        self,
        *,
        formula: str,
        home_sheet: str,
        old_maps: dict[str, _SheetMap],
        new_maps: dict[str, _SheetMap],
    ) -> tuple[str, list[dict[str, str]], list[dict[str, str]]]:
        reasons: list[dict[str, str]] = []
        missing: list[dict[str, str]] = []

        def _consider(sheet: str, col: str, row: str, acol: str, arow: str) -> str | None:
            addr = f"{col}{row}"
            old = old_maps.get(sheet)
            new = new_maps.get(sheet)
            if old is None or new is None:
                return None
            field = old.cell_to_field.get(addr)
            # Detail-line labels are not semantic fields. Remap only the
            # canonical metrics whose original template cell actually moved.
            if not field or field.startswith("label:"):
                return None
            new_addr = new.field_to_cell.get(field)
            if new_addr is None:
                missing.append(
                    {
                        "field": field,
                        "original_reference": f"{sheet}!{addr}",
                        "reason": "Semantic field no longer has a value cell after the quarterly layout change.",
                    }
                )
                return None
            if new_addr == addr:
                return None
            new_col, new_row = _split_addr(new_addr)
            reasons.append(
                {
                    "field": field,
                    "original_reference": f"{sheet}!{addr}",
                    "new_reference": f"{sheet}!{new_addr}",
                    "reason": f"{field} moved from {addr} to {new_addr} when the quarterly statement was rewritten.",
                }
            )
            return f"{acol}{new_col}{arow}{new_row}"

        def _qualified(match: re.Match[str]) -> str:
            sheet = match.group("qsheet") or match.group("bsheet")
            if sheet not in _STATEMENT_SHEETS:
                return match.group(0)
            core = _consider(
                sheet,
                match.group("col"),
                match.group("row"),
                match.group("acol"),
                match.group("arow"),
            )
            if core is None:
                return match.group(0)
            prefix = match.group(0).split("!", 1)[0]
            return f"{prefix}!{core}"

        # Same-sheet formulas (growth, subtotals) stay on their own rows.
        # Only qualified cross-sheet references are retargeted when a
        # canonical field moves to a new cell.
        _ = home_sheet
        updated = _QUALIFIED_REF.sub(_qualified, formula)
        return updated, reasons, missing


def _split_addr(addr: str) -> tuple[str, str]:
    match = re.fullmatch(r"([A-Z]{1,3})(\d+)", addr)
    if not match:
        return addr, ""
    return match.group(1), match.group(2)
