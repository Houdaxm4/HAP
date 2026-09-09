"""Completion planner — mode-scoped FILL only for genuinely missing cells."""

from __future__ import annotations

from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from models.completion import CompletionDecision, CompletionEntry, CompletionReport
from services.completion_scope import (
    AnalysisTypeMode,
    WorkbookSection,
    assess_quarterly_bloomberg_health,
    classify_workbook_section,
    infer_target_fiscal_year,
    is_required_for_mode,
    normalize_analysis_type,
    scope_summary,
    source_authority_for,
)
from workbook_mapping.engine import IntentDecision, WriteIntent, WriteIntentReport


def _is_formula(value: Any, data_type: str | None = None) -> bool:
    if data_type == "f":
        return True
    return isinstance(value, str) and value.startswith("=")


def _is_blank_or_unusable(value: Any, data_type: str | None = None) -> bool:
    if _is_formula(value, data_type):
        return False
    if value is None:
        return True
    if isinstance(value, str) and value.strip() == "":
        return True
    return False


_QUARTERLY_SECTIONS = frozenset(
    {
        WorkbookSection.QUARTERLY_INCOME,
        WorkbookSection.QUARTERLY_BALANCE_SHEET,
        WorkbookSection.QUARTERLY_CASH_FLOW,
    }
)


class CompletionService:
    """
    Inspect the uploaded workbook and convert M3 intents into analysis-type-scoped
    completion decisions. Only FILL decisions remain WRITE for the Excel fill engine.
    """

    def plan(
        self,
        *,
        analysis_id: str,
        ticker: str,
        analysis_type: str,
        source_workbook_path: Path,
        intent_report: WriteIntentReport,
        target_fiscal_year: str | None = None,
    ) -> tuple[CompletionReport, WriteIntentReport]:
        mode = normalize_analysis_type(analysis_type)
        workbook = load_workbook(source_workbook_path, data_only=False)
        entries: list[CompletionEntry] = []
        fill_intents: list[WriteIntent] = []

        try:
            target_fy = infer_target_fiscal_year(
                workbook=workbook,
                periods=[i.period for i in intent_report.intents],
                explicit=target_fiscal_year,
            )
            quarterly_health: dict[str, Any] | None = None
            per_q_decisions: dict[str, str] = {}
            if mode == AnalysisTypeMode.QUARTERLY_UPDATE:
                quarterly_health = assess_quarterly_bloomberg_health(workbook)
                per_q_decisions = dict(quarterly_health.get("per_statement_decisions") or {})

            for intent in intent_report.intents:
                section = classify_workbook_section(
                    sheet=intent.sheet,
                    metric=intent.metric,
                    cfm_path=intent.cfm_path,
                    period=intent.period,
                )
                authority = source_authority_for(section)
                required = is_required_for_mode(
                    mode,
                    section,
                    intent.period,
                    target_fiscal_year=target_fy,
                )
                mode_label = mode.value

                # --- Mode scope gate (blank is not automatically a target) ---
                if not required:
                    wb_val, _ = self._read_cell(workbook, intent.sheet, intent.cell)
                    entries.append(
                        self._entry(
                            intent,
                            CompletionDecision.OUT_OF_SCOPE,
                            workbook_value=wb_val,
                            reason=(
                                f"Section '{section.value}' @ {intent.period} is not a "
                                f"completion target for analysis_type={mode_label}"
                            ),
                            analysis_type=mode_label,
                            workbook_section=section.value,
                            required_for_mode=False,
                            source_authority=authority,
                        )
                    )
                    fill_intents.append(
                        intent.model_copy(
                            update={
                                "decision": IntentDecision.SKIP,
                                "write_policy_result": "completion_out_of_scope",
                                "reason": f"OUT_OF_SCOPE for {mode_label}",
                            }
                        )
                    )
                    continue

                # --- Per-statement quarterly presentation gate ---
                if section in _QUARTERLY_SECTIONS and per_q_decisions:
                    stmt_decision = per_q_decisions.get(section.value)
                    sheet_info = (quarterly_health or {}).get("per_sheet", {}).get(
                        intent.sheet, {}
                    )
                    if stmt_decision == "YAHOO_BASIC_TEMPLATE_REQUIRED":
                        wb_val, _ = self._read_cell(workbook, intent.sheet, intent.cell)
                        entries.append(
                            self._entry(
                                intent,
                                CompletionDecision.YAHOO_QUARTERLY_FALLBACK_REQUIRED,
                                workbook_value=wb_val,
                                reason=(
                                    f"YAHOO_BASIC_TEMPLATE_REQUIRED for {section.value}: "
                                    f"{sheet_info.get('decision', stmt_decision)} — "
                                    f"Yahoo basic template handles this statement"
                                ),
                                analysis_type=mode_label,
                                workbook_section=section.value,
                                required_for_mode=True,
                                source_authority="market_internet:yahoo",
                            )
                        )
                        fill_intents.append(
                            intent.model_copy(
                                update={
                                    "decision": IntentDecision.SKIP,
                                    "write_policy_result": "yahoo_basic_template_required",
                                    "reason": "Yahoo basic template replacement handles this statement",
                                }
                            )
                        )
                        continue
                    if stmt_decision == "SEC_10Q_PRESENTATION_REQUIRED":
                        wb_val, _ = self._read_cell(workbook, intent.sheet, intent.cell)
                        entries.append(
                            self._entry(
                                intent,
                                CompletionDecision.SEC_QUARTERLY_FALLBACK_REQUIRED,
                                workbook_value=wb_val,
                                reason=(
                                    f"SEC_10Q_PRESENTATION_REQUIRED for {section.value}: "
                                    f"{sheet_info.get('decision', stmt_decision)} — "
                                    f"no Bloomberg-cell SEC patching"
                                ),
                                analysis_type=mode_label,
                                workbook_section=section.value,
                                required_for_mode=True,
                                source_authority="sec_edgar_10q",
                            )
                        )
                        fill_intents.append(
                            intent.model_copy(
                                update={
                                    "decision": IntentDecision.SKIP,
                                    "write_policy_result": "sec_10q_presentation_required",
                                    "reason": "SEC layout replacement handles this statement",
                                }
                            )
                        )
                        continue
                    if stmt_decision == "BLOOMBERG_PRESERVE":
                        # Do not fill gaps under preserve — validation handles review.
                        wb_val, dtype = self._read_cell(workbook, intent.sheet, intent.cell)
                        if not _is_blank_or_unusable(wb_val, dtype):
                            entries.append(
                                self._entry(
                                    intent,
                                    CompletionDecision.ALREADY_PRESENT,
                                    workbook_value=wb_val,
                                    reason="BLOOMBERG_PRESERVE — prefilled quarterly value retained",
                                    analysis_type=mode_label,
                                    workbook_section=section.value,
                                    required_for_mode=True,
                                    source_authority=authority,
                                )
                            )
                            fill_intents.append(
                                intent.model_copy(
                                    update={
                                        "decision": IntentDecision.SKIP,
                                        "write_policy_result": "bloomberg_preserve",
                                        "reason": "BLOOMBERG_PRESERVE",
                                    }
                                )
                            )
                            continue
                        entries.append(
                            self._entry(
                                intent,
                                CompletionDecision.ALREADY_PRESENT,
                                workbook_value=wb_val,
                                reason=(
                                    "BLOOMBERG_PRESERVE — blank non-critical cell left untouched "
                                    "(isolated fills require BLOOMBERG_FILL_GAPS)"
                                ),
                                analysis_type=mode_label,
                                workbook_section=section.value,
                                required_for_mode=True,
                                source_authority=authority,
                            )
                        )
                        fill_intents.append(
                            intent.model_copy(
                                update={
                                    "decision": IntentDecision.SKIP,
                                    "write_policy_result": "bloomberg_preserve_no_gap_fill",
                                    "reason": "BLOOMBERG_PRESERVE",
                                }
                            )
                        )
                        continue
                    # BLOOMBERG_FILL_GAPS / BLOCKED fall through to normal in-scope logic

                # --- In-scope M3 BLOCK ---
                if intent.decision == IntentDecision.BLOCK:
                    entries.append(
                        self._entry(
                            intent,
                            CompletionDecision.BLOCKED,
                            workbook_value=None,
                            reason=f"M3 BLOCK: {intent.reason}",
                            analysis_type=mode_label,
                            workbook_section=section.value,
                            required_for_mode=True,
                            source_authority=authority,
                        )
                    )
                    fill_intents.append(intent)
                    continue

                # --- In-scope M3 SKIP (no CFM value) ---
                if intent.decision == IntentDecision.SKIP:
                    wb_val, dtype = self._read_cell(workbook, intent.sheet, intent.cell)
                    if wb_val is None and intent.sheet not in workbook.sheetnames:
                        entries.append(
                            self._entry(
                                intent,
                                CompletionDecision.BLOCKED,
                                workbook_value=None,
                                reason=f"Sheet '{intent.sheet}' missing from workbook",
                                analysis_type=mode_label,
                                workbook_section=section.value,
                                required_for_mode=True,
                                source_authority=authority,
                            )
                        )
                        fill_intents.append(
                            intent.model_copy(
                                update={
                                    "decision": IntentDecision.BLOCK,
                                    "write_policy_result": "completion_blocked",
                                    "reason": f"Sheet '{intent.sheet}' missing from workbook",
                                    "value": None,
                                }
                            )
                        )
                        continue
                    if _is_formula(wb_val, dtype):
                        entries.append(
                            self._entry(
                                intent,
                                CompletionDecision.BLOCKED,
                                workbook_value=wb_val,
                                reason="Formula cell; completion will not overwrite",
                                analysis_type=mode_label,
                                workbook_section=section.value,
                                required_for_mode=True,
                                source_authority=authority,
                            )
                        )
                        fill_intents.append(
                            intent.model_copy(
                                update={
                                    "decision": IntentDecision.BLOCK,
                                    "write_policy_result": "completion_blocked_formula",
                                    "reason": "Formula cell; completion will not overwrite",
                                    "value": None,
                                }
                            )
                        )
                        continue
                    if _is_blank_or_unusable(wb_val, dtype):
                        entries.append(
                            self._entry(
                                intent,
                                CompletionDecision.MISSING_SOURCE,
                                workbook_value=wb_val,
                                reason=(
                                    f"In-scope cell blank and CFM has no value for "
                                    f"{intent.cfm_path} @ {intent.period}"
                                ),
                                analysis_type=mode_label,
                                workbook_section=section.value,
                                required_for_mode=True,
                                source_authority=authority,
                            )
                        )
                        fill_intents.append(intent)
                        continue
                    entries.append(
                        self._entry(
                            intent,
                            CompletionDecision.ALREADY_PRESENT,
                            workbook_value=wb_val,
                            reason="Prefilled value present; no CFM candidate to fill",
                            analysis_type=mode_label,
                            workbook_section=section.value,
                            required_for_mode=True,
                            source_authority=authority,
                        )
                    )
                    fill_intents.append(intent)
                    continue

                # --- In-scope WRITE ---
                if intent.sheet not in workbook.sheetnames:
                    entries.append(
                        self._entry(
                            intent,
                            CompletionDecision.BLOCKED,
                            workbook_value=None,
                            reason=f"Sheet '{intent.sheet}' missing from workbook",
                            analysis_type=mode_label,
                            workbook_section=section.value,
                            required_for_mode=True,
                            source_authority=authority,
                        )
                    )
                    fill_intents.append(
                        intent.model_copy(
                            update={
                                "decision": IntentDecision.BLOCK,
                                "write_policy_result": "completion_blocked",
                                "reason": f"Sheet '{intent.sheet}' missing",
                                "value": None,
                            }
                        )
                    )
                    continue

                wb_val, dtype = self._read_cell(workbook, intent.sheet, intent.cell)
                if _is_formula(wb_val, dtype):
                    entries.append(
                        self._entry(
                            intent,
                            CompletionDecision.BLOCKED,
                            workbook_value=wb_val,
                            reason="Formula cell; never overwrite during completion",
                            analysis_type=mode_label,
                            workbook_section=section.value,
                            required_for_mode=True,
                            source_authority=authority,
                        )
                    )
                    fill_intents.append(
                        intent.model_copy(
                            update={
                                "decision": IntentDecision.BLOCK,
                                "write_policy_result": "completion_blocked_formula",
                                "reason": "Formula cell; never overwrite during completion",
                                "value": None,
                            }
                        )
                    )
                    continue

                if intent.value is None:
                    entries.append(
                        self._entry(
                            intent,
                            CompletionDecision.MISSING_SOURCE,
                            workbook_value=wb_val,
                            reason="WRITE intent has null proposed value",
                            analysis_type=mode_label,
                            workbook_section=section.value,
                            required_for_mode=True,
                            source_authority=authority,
                        )
                    )
                    fill_intents.append(
                        intent.model_copy(
                            update={
                                "decision": IntentDecision.SKIP,
                                "write_policy_result": "completion_missing_source",
                                "reason": "WRITE intent has null proposed value",
                            }
                        )
                    )
                    continue

                if not _is_blank_or_unusable(wb_val, dtype):
                    entries.append(
                        self._entry(
                            intent,
                            CompletionDecision.ALREADY_PRESENT,
                            workbook_value=wb_val,
                            reason=(
                                "In-scope cell already populated; deferred to validation "
                                "(SEC will not rewrite merely to match)"
                            ),
                            analysis_type=mode_label,
                            workbook_section=section.value,
                            required_for_mode=True,
                            source_authority=authority,
                        )
                    )
                    fill_intents.append(
                        intent.model_copy(
                            update={
                                "decision": IntentDecision.SKIP,
                                "write_policy_result": "completion_already_present",
                                "reason": (
                                    "ALREADY_PRESENT — preserved prefilled value; "
                                    "no completion write"
                                ),
                            }
                        )
                    )
                    continue

                entries.append(
                    self._entry(
                        intent,
                        CompletionDecision.FILL,
                        workbook_value=wb_val,
                        reason=(
                            f"In-scope required cell blank; filling from {intent.source} "
                            f"({intent.cfm_path} @ {intent.period})"
                        ),
                        analysis_type=mode_label,
                        workbook_section=section.value,
                        required_for_mode=True,
                        source_authority=authority,
                    )
                )
                fill_intents.append(
                    intent.model_copy(
                        update={
                            "decision": IntentDecision.WRITE,
                            "write_policy_result": "completion_fill",
                            "reason": f"FILL — blank in-scope cell from {intent.source}",
                        }
                    )
                )
        finally:
            workbook.close()

        summary = scope_summary(mode)
        report = CompletionReport(
            analysis_id=analysis_id,
            ticker=ticker,
            analysis_type=analysis_type,
            normalized_analysis_type=mode.value,
            target_fiscal_year=target_fy if mode == AnalysisTypeMode.ANNUAL_UPDATE else None,
            entries=entries,
            fill_count=sum(1 for e in entries if e.decision == CompletionDecision.FILL),
            already_present_count=sum(
                1 for e in entries if e.decision == CompletionDecision.ALREADY_PRESENT
            ),
            out_of_scope_count=sum(
                1 for e in entries if e.decision == CompletionDecision.OUT_OF_SCOPE
            ),
            missing_source_count=sum(
                1 for e in entries if e.decision == CompletionDecision.MISSING_SOURCE
            ),
            blocked_count=sum(1 for e in entries if e.decision == CompletionDecision.BLOCKED),
            sec_quarterly_fallback_count=sum(
                1
                for e in entries
                if e.decision == CompletionDecision.SEC_QUARTERLY_FALLBACK_REQUIRED
            ),
            sections_in_scope=summary["sections_in_scope"],
            sections_out_of_scope=summary["sections_out_of_scope"],
            sec_quarterly_fallback_required=bool(
                (quarterly_health or {}).get("sec_quarterly_fallback_required")
            ),
            quarterly_health=quarterly_health,
            assumptions=[
                "Blank cells are filled only when required_for_mode and a valid source exists.",
                "new_company completion scope: tax, PE10, current_data only.",
                "annual_update completion scope: new FY statements + tax/PE10/current for that year.",
                "quarterly_update completion scope: current_data + quarterly statements.",
                "Materially broken Bloomberg quarterly tabs emit YAHOO_QUARTERLY_FALLBACK_REQUIRED "
                "(Yahoo basic template; SEC secondary).",
                "Valid prefilled in-scope values are ALREADY_PRESENT; never silently overwritten.",
            ],
        )
        fill_report = intent_report.model_copy(deep=True)
        fill_report.intents = fill_intents
        fill_report.write_count = sum(
            1 for i in fill_intents if i.decision == IntentDecision.WRITE
        )
        fill_report.skip_count = sum(
            1 for i in fill_intents if i.decision == IntentDecision.SKIP
        )
        fill_report.block_count = sum(
            1 for i in fill_intents if i.decision == IntentDecision.BLOCK
        )
        return report, fill_report

    @staticmethod
    def _read_cell(workbook, sheet: str, cell: str) -> tuple[Any, str | None]:
        if sheet not in workbook.sheetnames:
            return None, None
        c = workbook[sheet][cell]
        return c.value, getattr(c, "data_type", None)

    @staticmethod
    def _entry(
        intent: WriteIntent,
        decision: CompletionDecision,
        *,
        workbook_value: Any,
        reason: str,
        analysis_type: str,
        workbook_section: str,
        required_for_mode: bool,
        source_authority: str,
    ) -> CompletionEntry:
        return CompletionEntry(
            intent_id=intent.intent_id,
            mapping_id=intent.mapping_id,
            sheet=intent.sheet,
            cell=intent.cell,
            cell_ref=intent.cell_ref,
            metric=intent.metric,
            period=intent.period,
            analysis_type=analysis_type,
            workbook_section=workbook_section,
            required_for_mode=required_for_mode,
            source_authority=source_authority,
            decision=decision,
            reason=reason,
            workbook_value=workbook_value,
            proposed_value=intent.value,
            source=intent.source,
        )
