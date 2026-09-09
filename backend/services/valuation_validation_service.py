"""Read-only valuation / 25% MOS / entry-price review for Industrial Template.

House methodology (not a parallel DCF):
  Graham Growth Intrinsic Value = Graham Multiple × EPS (Final Metrics R33)
  Current price = Inputs!B63 (live market; never CRF substitute)
  MOS = (intrinsic − price) / intrinsic
  Required threshold = 25%
  Entry price = intrinsic × (1 − 0.25)
"""

from __future__ import annotations

from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from models.valuation_validation import (
    ValuationAnalystItem,
    ValuationAnalystReview,
    ValuationAttractiveness,
    ValuationDecision,
    ValuationInputStatus,
    ValuationValidationReport,
)

FINAL_METRICS = "Final Metrics"
INPUTS = "Inputs"
INCOME = "Income - GAAP"
ENTERPRISE_VALUE = "Enterprise Value"
HOUSE_MOS_THRESHOLD = 0.25
_IV_ABS_TOL = 1.0  # USD per share
_NEAR_ENTRY_BAND = 0.05  # within 5pp of 25% threshold
# Classic Graham parameters as encoded on Enterprise Value!B51/B52
DEFAULT_PE_MULTIPLIER = 8.5
DEFAULT_GROWTH_MULTIPLIER = 2.0


def _num(value: Any) -> float | None:
    if value is None or value == "":
        return None
    if isinstance(value, bool):
        return None
    if isinstance(value, (int, float)):
        if value != value:
            return None
        return float(value)
    text = str(value).strip()
    if text.startswith("#"):
        return None
    try:
        return float(text)
    except (TypeError, ValueError):
        return None


def margin_of_safety(*, intrinsic: float, price: float) -> float | None:
    if intrinsic <= 0:
        return None
    return (intrinsic - price) / intrinsic


def entry_price_at_mos(*, intrinsic: float, threshold: float = HOUSE_MOS_THRESHOLD) -> float:
    return intrinsic * (1.0 - threshold)


def classify_attractiveness(
    mos: float | None,
    *,
    threshold: float = HOUSE_MOS_THRESHOLD,
) -> ValuationAttractiveness:
    if mos is None:
        return ValuationAttractiveness.INDETERMINATE
    if mos < 0:
        return ValuationAttractiveness.EXPENSIVE
    if mos >= threshold:
        return ValuationAttractiveness.ATTRACTIVE
    if mos >= threshold - _NEAR_ENTRY_BAND:
        return ValuationAttractiveness.NEAR_ENTRY
    return ValuationAttractiveness.WAIT


def graham_intrinsic(*, multiple: float, eps: float) -> float:
    return multiple * eps


def graham_multiple_from_cagr(
    *,
    eps_cagr: float,
    pe_multiplier: float = DEFAULT_PE_MULTIPLIER,
    growth_multiplier: float = DEFAULT_GROWTH_MULTIPLIER,
) -> float:
    """House: PE_mult + (CAGR_as_percent × growth_mult) — Final Metrics!L32."""
    return pe_multiplier + (eps_cagr * 100.0) * growth_multiplier


def eps_cagr_from_series(values: list[float]) -> float | None:
    if len(values) < 2:
        return None
    start, end = values[0], values[-1]
    years = len(values) - 1
    if start <= 0 or end <= 0 or years <= 0:
        return None
    return (end / start) ** (1.0 / years) - 1.0


def sensitivity_table(
    *,
    base_multiple: float,
    eps: float,
    price: float,
    threshold: float = HOUSE_MOS_THRESHOLD,
) -> dict[str, Any]:
    """Small sensitivity set — not a scenario engine."""
    rows = []
    for mult_delta, label in ((-0.15, "multiple_-15%"), (0.0, "base"), (0.15, "multiple_+15%")):
        m = base_multiple * (1.0 + mult_delta)
        iv = graham_intrinsic(multiple=m, eps=eps)
        mos = margin_of_safety(intrinsic=iv, price=price)
        rows.append(
            {
                "case": label,
                "graham_multiple": m,
                "intrinsic": iv,
                "mos": mos,
                "meets_25pct": None if mos is None else mos >= threshold,
                "entry_price": entry_price_at_mos(intrinsic=iv, threshold=threshold),
            }
        )
    for eps_delta, label in ((-0.10, "eps_-10%"), (0.10, "eps_+10%")):
        e = eps * (1.0 + eps_delta)
        iv = graham_intrinsic(multiple=base_multiple, eps=e)
        mos = margin_of_safety(intrinsic=iv, price=price)
        rows.append(
            {
                "case": label,
                "eps": e,
                "intrinsic": iv,
                "mos": mos,
                "meets_25pct": None if mos is None else mos >= threshold,
                "entry_price": entry_price_at_mos(intrinsic=iv, threshold=threshold),
            }
        )
    meets = [r["meets_25pct"] for r in rows if r["meets_25pct"] is not None]
    return {
        "cases": rows,
        "robust_to_adverse": bool(meets) and all(meets),
        "fragile_note": (
            "Conclusion flips under ±15% multiple or ±10% EPS shocks."
            if meets and not all(meets) and any(meets)
            else None
        ),
    }


class ValuationValidationService:
    """Inspect workbook valuation outputs; validate MOS / entry at 25% threshold."""

    def validate(
        self,
        *,
        analysis_id: str,
        ticker: str,
        workbook_path: Path,
        current_price_override: float | None = None,
        current_price_source: str | None = None,
    ) -> ValuationValidationReport:
        path = Path(workbook_path)
        wb_v = load_workbook(path, data_only=True)
        wb_f = load_workbook(path, data_only=False)

        methodology = self._methodology(wb_f)
        inputs, assumptions = self._collect_inputs(
            wb_v,
            current_price_override=current_price_override,
            current_price_source=current_price_source,
        )

        price_status = next(i for i in inputs if i.name == "current_price")
        multiple_status = next(i for i in inputs if i.name == "graham_multiple")
        eps_status = next(i for i in inputs if i.name == "eps")
        iv_wb_status = next(i for i in inputs if i.name == "workbook_intrinsic_graham")

        price = price_status.value
        multiple = multiple_status.value
        eps = eps_status.value
        iv_wb = iv_wb_status.value

        iv_ind = None
        if multiple is not None and eps is not None:
            iv_ind = graham_intrinsic(multiple=multiple, eps=eps)

        iv_diff = None
        if iv_wb is not None and iv_ind is not None:
            iv_diff = iv_wb - iv_ind
        if iv_wb is None and iv_ind is None:
            iv_decision = ValuationDecision.SOURCE_MISSING
        elif iv_wb is not None and iv_ind is not None and abs(iv_diff or 0) <= _IV_ABS_TOL:
            iv_decision = ValuationDecision.VALIDATED
        elif iv_wb is not None and iv_ind is not None:
            iv_decision = ValuationDecision.DISCREPANCY
        else:
            iv_decision = ValuationDecision.REVIEW_REQUIRED

        intrinsic = iv_wb if iv_wb is not None else iv_ind

        mos = None
        entry = None
        gap = None
        pct_decline = None
        meets = None
        if intrinsic is not None and price is not None and price > 0 and intrinsic > 0:
            mos = margin_of_safety(intrinsic=intrinsic, price=price)
            entry = entry_price_at_mos(intrinsic=intrinsic, threshold=HOUSE_MOS_THRESHOLD)
            gap = price - entry
            pct_decline = (price - entry) / price if price else None
            meets = mos >= HOUSE_MOS_THRESHOLD

        attractiveness = classify_attractiveness(mos, threshold=HOUSE_MOS_THRESHOLD)
        if price is None or price <= 0:
            attractiveness = ValuationAttractiveness.INDETERMINATE
            price_status.status = ValuationDecision.SOURCE_MISSING

        sens: dict[str, Any] = {}
        if multiple is not None and eps is not None and price is not None and price > 0:
            sens = sensitivity_table(
                base_multiple=multiple, eps=eps, price=price, threshold=HOUSE_MOS_THRESHOLD
            )

        decision, explanation = self._overall(
            price_decision=price_status.status,
            iv_decision=iv_decision,
            mos=mos,
            attractiveness=attractiveness,
            intrinsic=intrinsic,
            price=price,
            entry=entry,
        )

        max_buy = next((i.value for i in inputs if i.name == "max_price_to_buy"), None)

        summary = (
            f"Valuation review {ticker}: IV={intrinsic}, price={price}, "
            f"MOS={None if mos is None else round(mos, 4)}, "
            f"entry@25%={entry}, attractiveness={attractiveness.value}. {explanation}"
        )

        report = ValuationValidationReport(
            analysis_id=analysis_id,
            ticker=ticker,
            workbook_path=str(path),
            methodology=methodology,
            inputs=inputs,
            assumptions=assumptions,
            workbook_intrinsic_value=iv_wb,
            independent_intrinsic_value=iv_ind,
            intrinsic_difference=iv_diff,
            intrinsic_decision=iv_decision,
            current_price=price,
            current_price_source=price_status.source,
            current_price_decision=price_status.status,
            margin_of_safety=mos,
            required_mos_threshold=HOUSE_MOS_THRESHOLD,
            meets_mos_threshold=meets,
            required_entry_price=entry,
            price_gap_to_entry=gap,
            pct_decline_to_entry=pct_decline,
            workbook_max_price_to_buy=max_buy,
            attractiveness=attractiveness,
            sensitivity=sens,
            decision=decision,
            explanation=explanation,
            summary=summary,
            evidence={
                "graham_iv_cell": "Final Metrics!L33 (latest) / B33 (10y avg)",
                "current_price_cell": "Inputs!B63",
                "max_price_to_buy_cell": "Inputs!B67",
                "house_mos_threshold": HOUSE_MOS_THRESHOLD,
            },
        )
        wb_v.close()
        wb_f.close()
        return report

    def build_analyst_review(self, report: ValuationValidationReport) -> ValuationAnalystReview:
        items: list[ValuationAnalystItem] = []
        items.append(
            ValuationAnalystItem(
                topic="intrinsic_value",
                status=report.intrinsic_decision,
                observation=(
                    f"Workbook Graham IV={report.workbook_intrinsic_value}; "
                    f"independent={report.independent_intrinsic_value}."
                ),
                quantitative_evidence={
                    "workbook": report.workbook_intrinsic_value,
                    "independent": report.independent_intrinsic_value,
                    "diff": report.intrinsic_difference,
                },
                interpretation="House intrinsic is Graham Growth IV (multiple × EPS), not a full DCF.",
            )
        )
        items.append(
            ValuationAnalystItem(
                topic="margin_of_safety",
                status=report.decision,
                observation=report.explanation,
                quantitative_evidence={
                    "mos": report.margin_of_safety,
                    "threshold": report.required_mos_threshold,
                    "meets": report.meets_mos_threshold,
                    "entry": report.required_entry_price,
                    "price": report.current_price,
                    "pct_decline_to_entry": report.pct_decline_to_entry,
                },
                interpretation=(
                    "Undervalued ≠ buyable: require ≥25% MOS. "
                    f"Attractiveness={report.attractiveness.value}."
                ),
            )
        )
        if report.sensitivity:
            items.append(
                ValuationAnalystItem(
                    topic="sensitivity",
                    status=ValuationDecision.WATCH
                    if report.sensitivity.get("fragile_note")
                    else ValuationDecision.VALIDATED,
                    observation=report.sensitivity.get("fragile_note")
                    or "MOS conclusion stable under small multiple/EPS shocks.",
                    quantitative_evidence=report.sensitivity,
                    interpretation="Limited sensitivity — not a full scenario engine.",
                )
            )
        for inp in report.inputs:
            if inp.status in (
                ValuationDecision.SOURCE_MISSING,
                ValuationDecision.DISCREPANCY,
                ValuationDecision.REVIEW_REQUIRED,
            ):
                items.append(
                    ValuationAnalystItem(
                        topic=f"input:{inp.name}",
                        status=inp.status,
                        observation=inp.reason,
                        quantitative_evidence={"value": inp.value, "cell": inp.cell},
                        interpretation="Material valuation input needs attention.",
                    )
                )
        return ValuationAnalystReview(
            analysis_id=report.analysis_id,
            ticker=report.ticker,
            attractiveness=report.attractiveness,
            items=items,
            summary=(
                f"{len(items)} valuation analyst items; "
                f"attractiveness={report.attractiveness.value}; "
                f"decision={report.decision.value}."
            ),
        )

    def _methodology(self, wb_f) -> dict[str, Any]:
        fm = wb_f[FINAL_METRICS] if FINAL_METRICS in wb_f.sheetnames else None
        ev = wb_f[ENTERPRISE_VALUE] if ENTERPRISE_VALUE in wb_f.sheetnames else None
        return {
            "primary_intrinsic": "Graham Growth Intrinsic Value (Final Metrics row 33 / Enterprise Value!B42)",
            "intrinsic_formula": "EPS × (PE_multiplier + CAGR% × growth_multiplier)",
            "graham_multiple_formula": "Enterprise Value!B51 + (EPS_CAGR×100)×Enterprise Value!B52",
            "default_pe_multiplier": DEFAULT_PE_MULTIPLIER,
            "default_growth_multiplier": DEFAULT_GROWTH_MULTIPLIER,
            "sample_L33": fm["L33"].value if fm else None,
            "sample_B42": ev["B42"].value if ev else None,
            "current_price": "Inputs!B63 → Final Metrics!B51 (live market_internet; not CRF)",
            "max_price_to_buy": "Inputs!B67 (Bloomberg CRF entry guide)",
            "max_pe10_to_enter": "Inputs!B66 / Final Metrics!B54",
            "mos_definition": "(intrinsic - price) / intrinsic",
            "house_mos_threshold": HOUSE_MOS_THRESHOLD,
            "entry_price_definition": "intrinsic * (1 - 0.25)",
            "no_parallel_dcf": True,
            "no_workbook_formula_writes": True,
        }

    def _collect_inputs(
        self,
        wb_v,
        *,
        current_price_override: float | None,
        current_price_source: str | None,
    ) -> tuple[list[ValuationInputStatus], dict[str, Any]]:
        fm = wb_v[FINAL_METRICS] if FINAL_METRICS in wb_v.sheetnames else None
        inp = wb_v[INPUTS] if INPUTS in wb_v.sheetnames else None
        income = wb_v[INCOME] if INCOME in wb_v.sheetnames else None
        ev = wb_v[ENTERPRISE_VALUE] if ENTERPRISE_VALUE in wb_v.sheetnames else None

        pe_mult = DEFAULT_PE_MULTIPLIER
        g_mult = DEFAULT_GROWTH_MULTIPLIER
        if ev is not None:
            pe_mult = _num(ev["B51"].value) or pe_mult
            g_mult = _num(ev["B52"].value) or g_mult

        # EPS: prefer Income numeric series (survives openpyxl save); else FM
        eps_series: list[float] = []
        if income is not None:
            for col in range(3, 13):
                v = _num(income.cell(71, col).value)
                if v is not None:
                    eps_series.append(v)
        eps = eps_series[-1] if eps_series else None
        if eps is None and fm is not None:
            eps = _num(fm["L29"].value) or _num(fm["B29"].value)

        cagr = eps_cagr_from_series(eps_series) if len(eps_series) >= 2 else None
        if cagr is None and fm is not None:
            cagr = _num(fm["L31"].value)

        # Workbook IV: cached FM/EV when present
        iv = _num(fm["L33"].value) if fm else None
        iv_src = "Final Metrics!L33"
        if iv is None and fm is not None:
            iv = _num(fm["B33"].value)
            iv_src = "Final Metrics!B33"
        if iv is None and ev is not None:
            iv = _num(ev["B42"].value)
            if iv is not None:
                iv_src = "Enterprise Value!B42"

        multiple = _num(fm["L32"].value) if fm else None
        if multiple is None and fm is not None:
            multiple = _num(fm["B32"].value)
        # Reconstruct house multiple when formula cache missing
        if multiple is None and cagr is not None:
            multiple = graham_multiple_from_cagr(
                eps_cagr=cagr, pe_multiplier=pe_mult, growth_multiplier=g_mult
            )
        # Reconstruct workbook IV if only components available
        if iv is None and multiple is not None and eps is not None:
            iv = graham_intrinsic(multiple=multiple, eps=eps)
            iv_src = "reconstructed_house:EPS×(8.5+2g)"

        price = current_price_override
        price_source = current_price_source
        price_cell = "Inputs!B63"
        if price is None and inp is not None:
            price = _num(inp["B63"].value)
            if price is not None:
                price_source = price_source or "workbook:Inputs!B63"
        if price is None and fm is not None:
            price = _num(fm["B51"].value)
            price_cell = "Final Metrics!B51"
            if price is not None:
                price_source = price_source or "workbook:Final Metrics!B51"

        max_pe = _num(inp["B66"].value) if inp else None
        if max_pe is None and fm is not None:
            max_pe = _num(fm["B54"].value)
        max_buy = _num(inp["B67"].value) if inp else None
        exit_px = _num(inp["B68"].value) if inp else None
        e10 = _num(inp["B64"].value) if inp else None
        pe10 = _num(inp["B65"].value) if inp else None
        wacc = None
        if fm is not None:
            wacc = _num(fm["L7"].value) or _num(fm["B7"].value)

        inputs = [
            ValuationInputStatus(
                name="workbook_intrinsic_graham",
                value=iv,
                cell=iv_src,
                status=ValuationDecision.VALIDATED if iv is not None else ValuationDecision.SOURCE_MISSING,
                reason="Graham Growth Intrinsic Value (house)",
                source="workbook",
            ),
            ValuationInputStatus(
                name="graham_multiple",
                value=multiple,
                cell="Final Metrics!L32 / reconstructed 8.5+2g",
                status=ValuationDecision.VALIDATED
                if multiple is not None
                else ValuationDecision.SOURCE_MISSING,
                reason="Graham Value Multiple",
                source="workbook",
            ),
            ValuationInputStatus(
                name="eps",
                value=eps,
                cell="Income!L71",
                status=ValuationDecision.VALIDATED if eps is not None else ValuationDecision.SOURCE_MISSING,
                reason="Diluted EPS feeding Graham IV",
                source="workbook",
            ),
            ValuationInputStatus(
                name="eps_cagr",
                value=cagr,
                cell="Final Metrics!L31 / EPS series",
                status=ValuationDecision.VALIDATED if cagr is not None else ValuationDecision.SOURCE_MISSING,
                reason="EPS 10y CAGR for Graham multiple",
                source="workbook",
            ),
            ValuationInputStatus(
                name="current_price",
                value=price,
                cell=price_cell,
                status=ValuationDecision.VALIDATED
                if price is not None and price > 0
                else ValuationDecision.SOURCE_MISSING,
                reason=(
                    "Live/current price present"
                    if price is not None and price > 0
                    else "Current price missing — do not invent attractiveness"
                ),
                source=price_source,
            ),
            ValuationInputStatus(
                name="max_pe10_to_enter",
                value=max_pe,
                cell="Inputs!B66",
                status=ValuationDecision.VALIDATED
                if max_pe is not None
                else ValuationDecision.SOURCE_MISSING,
                reason="Entry multiple guide",
                source="workbook/CRF",
            ),
            ValuationInputStatus(
                name="max_price_to_buy",
                value=max_buy,
                cell="Inputs!B67",
                status=ValuationDecision.VALIDATED
                if max_buy is not None
                else ValuationDecision.WATCH,
                reason="CRF max buy price (optional cross-check vs 25% entry)",
                source="bloomberg_custom_run",
            ),
            ValuationInputStatus(
                name="current_e10",
                value=e10,
                cell="Inputs!B64",
                status=ValuationDecision.VALIDATED if e10 is not None else ValuationDecision.WATCH,
                reason="E10 supporting PE10 valuation context",
            ),
            ValuationInputStatus(
                name="current_pe10",
                value=pe10,
                cell="Inputs!B65",
                status=ValuationDecision.VALIDATED if pe10 is not None else ValuationDecision.WATCH,
                reason="Current PE10",
            ),
            ValuationInputStatus(
                name="wacc",
                value=wacc,
                cell="Final Metrics!L7",
                status=ValuationDecision.VALIDATED if wacc is not None else ValuationDecision.WATCH,
                reason="WACC context (not primary Graham driver)",
            ),
            ValuationInputStatus(
                name="exit_price",
                value=exit_px,
                cell="Inputs!B68",
                status=ValuationDecision.WATCH if exit_px is not None else ValuationDecision.SOURCE_MISSING,
                reason="1st exit price (PE10 percentile)",
            ),
        ]
        assumptions = {
            "house_mos_threshold": HOUSE_MOS_THRESHOLD,
            "intrinsic_method": "graham_growth",
            "pe_multiplier": pe_mult,
            "growth_multiplier": g_mult,
            "eps_cagr": cagr,
            "exit_price": exit_px,
        }
        return inputs, assumptions

    def _overall(
        self,
        *,
        price_decision: ValuationDecision,
        iv_decision: ValuationDecision,
        mos: float | None,
        attractiveness: ValuationAttractiveness,
        intrinsic: float | None,
        price: float | None,
        entry: float | None,
    ) -> tuple[ValuationDecision, str]:
        if price_decision == ValuationDecision.SOURCE_MISSING:
            return (
                ValuationDecision.SOURCE_MISSING,
                "Current price missing — valuation attractiveness indeterminate.",
            )
        if iv_decision == ValuationDecision.SOURCE_MISSING:
            return (
                ValuationDecision.SOURCE_MISSING,
                "Intrinsic value unavailable — cannot compute MOS.",
            )
        if iv_decision == ValuationDecision.DISCREPANCY:
            return (
                ValuationDecision.DISCREPANCY,
                "Independent Graham reconstruction differs from workbook IV beyond tolerance.",
            )
        mos_pct = None if mos is None else f"{mos:.1%}"
        return (
            ValuationDecision.VALIDATED
            if iv_decision == ValuationDecision.VALIDATED
            else ValuationDecision.REVIEW_REQUIRED,
            (
                f"Intrinsic={intrinsic}, price={price}, MOS={mos_pct}, "
                f"25% entry={entry}, class={attractiveness.value}. "
                "Below intrinsic alone is not sufficient; need ≥25% MOS to be ATTRACTIVE."
            ),
        )
