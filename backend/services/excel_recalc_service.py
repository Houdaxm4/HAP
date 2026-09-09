"""Excel workbook recalculation via Microsoft Excel COM (Windows).

Preserves formulas and formatting. Writes refreshed cached values into the
workbook so openpyxl data_only=True can read results. Never replaces formulas
with constants.
"""

from __future__ import annotations

import time
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any


EXCEL_COM_METHOD = "excel_com_calculate_full_rebuild"


class ExcelComUnavailable(Exception):
    """Raised when pywin32/win32com cannot be imported."""


class ExcelComFailed(Exception):
    """Raised when Excel COM starts but CalculateFullRebuild fails."""


@dataclass
class ExcelRecalcReport:
    analysis_id: str
    ticker: str
    status: str  # ok | FAILED | UNAVAILABLE
    method: str
    workbook_path: str
    elapsed_ms: float = 0.0
    cells_checked: list[str] = field(default_factory=list)
    missing_cached_values: list[str] = field(default_factory=list)
    formula_errors: list[str] = field(default_factory=list)
    error: str | None = None
    summary: str = ""
    com_invoked: bool = False


def genuine_excel_com_recalc(report: ExcelRecalcReport | None) -> bool:
    """True only after Excel COM CalculateFullRebuild ran and caches verified."""
    return bool(
        report is not None
        and report.status == "ok"
        and report.method == EXCEL_COM_METHOD
        and report.com_invoked
    )


def load_excel_com():
    """Import Windows Excel COM libraries. Isolated so tests mock only this boundary."""
    try:
        import pythoncom  # type: ignore
        import win32com.client as win32com_client  # type: ignore
    except ImportError as exc:
        raise ExcelComUnavailable("pywin32/win32com not installed") from exc
    return win32com_client, pythoncom


# Fixed required post-recalc outputs for Annual Update report authorization.
REQUIRED_ANNUAL_OUTPUTS: tuple[tuple[str, str], ...] = (
    ("Expected Returns & Buybacks", "E14"),
    ("Expected Returns & Buybacks", "F14"),
    ("Enterprise Value", "B20"),
    ("Enterprise Value", "B27"),
    ("Enterprise Value", "B32"),
    ("Enterprise Value", "B42"),
    ("Enterprise Value", "B47"),
    ("Enterprise Value", "B48"),
    ("Enterprise Value", "B6"),
    ("Enterprise Value", "C6"),
)


def _fy_dependent_outputs(path: Path, fiscal_year: str | None) -> list[tuple[str, str]]:
    """Tax / R&D / IC / Final Metrics cells for the new fiscal-year column."""
    if not fiscal_year:
        return []
    try:
        from openpyxl import load_workbook
        from openpyxl.utils import get_column_letter
        from services.annual_period_service import detect_year_columns
    except Exception:  # noqa: BLE001
        return []
    wb = load_workbook(path, data_only=False)
    try:
        if "Inputs" not in wb.sheetnames:
            return []
        cols = detect_year_columns(wb["Inputs"], wb)
        token = fiscal_year if str(fiscal_year).startswith("FY") else f"FY{fiscal_year}"
        col = cols.get(token)
        if not col:
            return []
        letter = get_column_letter(col)
        rd_letter = get_column_letter(col + 2)
        out: list[tuple[str, str]] = [
            ("Tax", f"{letter}15"),
            ("Tax", f"{letter}25"),
            ("R&D", f"{rd_letter}2"),
            ("R&D", f"{rd_letter}3"),
            ("R&D", f"{rd_letter}4"),
            ("Final Metrics", f"{letter}5"),  # ROCE
            ("Final Metrics", f"{letter}8"),  # ROIC - WACC
        ]
        ic_name = next(
            (n for n in wb.sheetnames if "nopat" in n.lower() and "roic" in n.lower()),
            None,
        )
        if ic_name:
            # Rows observed on HAP templates: Invested Capital=7, NOPAT=20, ROIC=23
            out.extend(
                [
                    (ic_name, f"{letter}7"),
                    (ic_name, f"{letter}20"),
                    (ic_name, f"{letter}23"),
                ]
            )
        return out
    finally:
        wb.close()


class ExcelRecalcService:
    """Recalculate an .xlsx in-place using Excel automation when available."""

    def recalculate(
        self,
        *,
        analysis_id: str,
        ticker: str,
        workbook_path: Path,
        fiscal_year: str | None = None,
        required_outputs: tuple[tuple[str, str], ...] | None = None,
    ) -> ExcelRecalcReport:
        path = Path(workbook_path).resolve()
        required_list = list(required_outputs or REQUIRED_ANNUAL_OUTPUTS)
        required_list.extend(_fy_dependent_outputs(path, fiscal_year))
        # de-dupe preserving order
        seen: set[tuple[str, str]] = set()
        required: list[tuple[str, str]] = []
        for item in required_list:
            if item not in seen:
                seen.add(item)
                required.append(item)
        t0 = time.perf_counter()
        formula_snapshot = self._snapshot_formulas(path, tuple(required))

        try:
            self._invoke_calculate_full_rebuild(path)
        except ExcelComUnavailable as exc:
            return ExcelRecalcReport(
                analysis_id=analysis_id,
                ticker=ticker,
                status="UNAVAILABLE",
                method="none",
                workbook_path=str(path),
                error=str(exc),
                summary="WORKBOOK_RECALCULATION_INCOMPLETE: Excel COM unavailable.",
                missing_cached_values=[f"{s}!{c}" for s, c in required],
                com_invoked=False,
            )
        except ExcelComFailed as exc:
            elapsed = (time.perf_counter() - t0) * 1000.0
            return ExcelRecalcReport(
                analysis_id=analysis_id,
                ticker=ticker,
                status="FAILED",
                method=EXCEL_COM_METHOD,
                workbook_path=str(path),
                elapsed_ms=round(elapsed, 1),
                error=str(exc),
                missing_cached_values=[f"{s}!{c}" for s, c in required],
                summary=f"WORKBOOK_RECALCULATION_INCOMPLETE: Excel COM failed: {exc}",
                com_invoked=True,
            )

        elapsed = (time.perf_counter() - t0) * 1000.0
        formula_overwrites = self._verify_formulas_preserved(path, formula_snapshot)
        checked, missing, errors = self._verify_cached_values(path, tuple(required))
        errors = list(errors) + formula_overwrites
        status = "ok" if not missing and not errors else "FAILED"
        return ExcelRecalcReport(
            analysis_id=analysis_id,
            ticker=ticker,
            status=status,
            method=EXCEL_COM_METHOD,
            workbook_path=str(path),
            elapsed_ms=round(elapsed, 1),
            cells_checked=checked,
            missing_cached_values=missing,
            formula_errors=errors,
            com_invoked=True,
            summary=(
                f"Excel recalc {status}: method={EXCEL_COM_METHOD}; "
                f"checked={len(checked)}; missing={len(missing)}; errors={len(errors)}; "
                f"{elapsed/1000:.1f}s."
            ),
        )

    @staticmethod
    def _invoke_calculate_full_rebuild(path: Path) -> None:
        """Windows Excel COM boundary: open, CalculateFullRebuild, save. Do not compute values in Python."""
        win32com_client, pythoncom = load_excel_com()
        excel = None
        wb = None
        try:
            pythoncom.CoInitialize()
            excel = win32com_client.DispatchEx("Excel.Application")
            excel.Visible = False
            excel.DisplayAlerts = False
            excel.AskToUpdateLinks = False
            excel.EnableEvents = False
            try:
                excel.Calculation = -4105  # xlCalculationAutomatic
            except Exception:  # noqa: BLE001
                pass
            wb = excel.Workbooks.Open(str(path), UpdateLinks=0, ReadOnly=False)
            excel.CalculateFullRebuild()
            wb.Save()
            wb.Close(SaveChanges=True)
            wb = None
            # Required save+reopen cycle before cache verification.
            wb = excel.Workbooks.Open(str(path), UpdateLinks=0, ReadOnly=False)
            wb.Close(SaveChanges=False)
            wb = None
        except ExcelComUnavailable:
            raise
        except Exception as exc:  # noqa: BLE001
            raise ExcelComFailed(str(exc)) from exc
        finally:
            try:
                if wb is not None:
                    wb.Close(SaveChanges=False)
            except Exception:  # noqa: BLE001
                pass
            try:
                if excel is not None:
                    excel.Quit()
            except Exception:  # noqa: BLE001
                pass
            try:
                pythoncom.CoUninitialize()
            except Exception:  # noqa: BLE001
                pass

    @staticmethod
    def _verify_cached_values(
        path: Path, required: tuple[tuple[str, str], ...]
    ) -> tuple[list[str], list[str], list[str]]:
        from openpyxl import load_workbook

        checked: list[str] = []
        missing: list[str] = []
        errors: list[str] = []
        error_tokens = ("#REF!", "#DIV/0!", "#VALUE!", "#NAME?", "#N/A", "#NUM!", "#NULL!")
        wb = load_workbook(path, data_only=True)
        try:
            for sheet, addr in required:
                ref = f"{sheet}!{addr}"
                checked.append(ref)
                if sheet not in wb.sheetnames:
                    missing.append(ref)
                    continue
                val = wb[sheet][addr].value
                if val is None or val == "":
                    missing.append(ref)
                elif isinstance(val, str) and any(tok in val for tok in error_tokens):
                    errors.append(f"{ref}={val}")
        finally:
            wb.close()
        return checked, missing, errors

    @staticmethod
    def _snapshot_formulas(
        path: Path, required: tuple[tuple[str, str], ...]
    ) -> dict[tuple[str, str], str]:
        from openpyxl import load_workbook

        snap: dict[tuple[str, str], str] = {}
        wb = load_workbook(path, data_only=False)
        try:
            for sheet, addr in required:
                if sheet not in wb.sheetnames:
                    continue
                val = wb[sheet][addr].value
                if isinstance(val, str) and val.startswith("="):
                    snap[(sheet, addr)] = val
        finally:
            wb.close()
        return snap

    @staticmethod
    def _verify_formulas_preserved(
        path: Path, formula_snapshot: dict[tuple[str, str], str]
    ) -> list[str]:
        """Fail if a cell that was a formula before recalc is no longer a formula."""
        from openpyxl import load_workbook

        errors: list[str] = []
        if not formula_snapshot:
            return errors
        wb = load_workbook(path, data_only=False)
        try:
            for (sheet, addr), before in formula_snapshot.items():
                if sheet not in wb.sheetnames:
                    errors.append(f"FORMULA_OVERWRITTEN: {sheet}!{addr} sheet missing after recalc")
                    continue
                after = wb[sheet][addr].value
                if not (isinstance(after, str) and after.startswith("=")):
                    errors.append(
                        f"FORMULA_OVERWRITTEN: {sheet}!{addr} was formula, now {after!r}"
                    )
                elif after != before:
                    # Recalc should not rewrite formula text; tolerate whitespace-only diffs.
                    if after.replace(" ", "") != before.replace(" ", ""):
                        errors.append(
                            f"FORMULA_CHANGED: {sheet}!{addr} before={before!r} after={after!r}"
                        )
        finally:
            wb.close()
        return errors

    @staticmethod
    def read_cached(path: Path, sheet: str, addr: str) -> Any:
        from openpyxl import load_workbook

        wb = load_workbook(path, data_only=True)
        try:
            if sheet not in wb.sheetnames:
                return None
            return wb[sheet][addr].value
        finally:
            wb.close()
