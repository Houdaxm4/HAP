"""Discover a safe unused adjacent area for HAP_ANALYSIS writes.

Explanatory prose belongs in a bottom-of-sheet notes section. Live
formula blocks stay in their adjacent columns so existing dependencies
are not moved.
"""

from __future__ import annotations

from typing import Any

from openpyxl.cell.cell import MergedCell
from openpyxl.utils import get_column_letter, range_boundaries
from openpyxl.worksheet.worksheet import Worksheet

from services.formula_utils import is_formula
from services.workbook_flag_service import HAP_ANALYSIS_LABEL, style_hap_analysis_cell

NOTES_HEADER = "HAP ANALYSIS — NOTES"


def discover_unused_column(ws: Worksheet, *, start_col: int = 7, end_col: int = 20, scan_rows: int = 40) -> int:
    """Return the first column with no occupied cells in the scan window.

    Does not hard-code a HAP column that might collide across template versions.
    """
    for col in range(start_col, end_col + 1):
        occupied = False
        for row in range(1, scan_rows + 1):
            cell = ws.cell(row, col)
            if isinstance(cell, MergedCell) or cell.value not in (None, ""):
                occupied = True
                break
        if not occupied:
            return col
    return end_col


def cell_is_occupied_or_formula(ws: Worksheet, addr: str) -> bool:
    cell = ws[addr]
    if isinstance(cell, MergedCell):
        return True
    if is_formula(cell.value):
        return True
    return cell.value not in (None, "")


def _range_max_row(ref: str) -> int | None:
    text = str(ref or "").strip()
    if not text:
        return None
    piece = text.split("!")[-1].replace("$", "")
    if not piece:
        return None
    try:
        _min_c, _min_r, _max_c, max_r = range_boundaries(piece)
    except (ValueError, TypeError):
        return None
    return max_r


def discover_occupied_end_row(ws: Worksheet) -> int:
    """Last row that holds model content, including hidden cells and structures.

    Scans values, formulas, merged ranges, tables, print area, and defined
    names. Does not stop at the first visual blank.
    """
    end = 0
    max_row = ws.max_row or 1
    max_col = ws.max_column or 1
    for row in ws.iter_rows(min_row=1, max_row=max_row, max_col=max_col):
        for cell in row:
            if isinstance(cell, MergedCell):
                continue
            if cell.value not in (None, ""):
                end = max(end, cell.row)
    for merged in ws.merged_cells.ranges:
        end = max(end, merged.max_row)
    for table in (getattr(ws, "tables", None) or {}).values():
        found = _range_max_row(getattr(table, "ref", "") or "")
        if found:
            end = max(end, found)
    print_area = getattr(ws, "print_area", None)
    if print_area:
        for piece in str(print_area).split(","):
            found = _range_max_row(piece)
            if found:
                end = max(end, found)
    workbook = ws.parent
    defined = getattr(workbook, "defined_names", None) if workbook is not None else None
    if defined is not None:
        for defn in defined.values():
            attr = str(getattr(defn, "attr_text", "") or "")
            title = ws.title
            if title not in attr and f"'{title}'" not in attr:
                continue
            for piece in attr.split(","):
                if title not in piece:
                    continue
                found = _range_max_row(piece)
                if found:
                    end = max(end, found)
    # Hidden rows are included by the value scan above; openpyxl yields
    # their cells even when row_dimensions.hidden is set.
    return end


class HapAnalysisLayoutService:
    """Write HAP_ANALYSIS blocks into verified unused adjacent columns."""

    def allocate(self, ws: Worksheet, *, start_col: int = 7) -> tuple[int, str]:
        col = discover_unused_column(ws, start_col=start_col)
        letter = get_column_letter(col)
        return col, letter

    def write_block(
        self,
        ws: Worksheet,
        *,
        rows: list[tuple[str, Any]],
        start_col: int | None = None,
        header: str = HAP_ANALYSIS_LABEL,
    ) -> list[str]:
        """Write label/value pairs. Never overwrite occupied/formula/source cells."""
        col = start_col or discover_unused_column(ws)
        written: list[str] = []
        label_cell = ws.cell(1, col)
        value_cell = ws.cell(1, col + 1)
        if cell_is_occupied_or_formula(ws, label_cell.coordinate) or cell_is_occupied_or_formula(
            ws, value_cell.coordinate
        ):
            col = discover_unused_column(ws, start_col=col + 2)
            label_cell = ws.cell(1, col)
            value_cell = ws.cell(1, col + 1)
        if cell_is_occupied_or_formula(ws, label_cell.coordinate) or cell_is_occupied_or_formula(
            ws, value_cell.coordinate
        ):
            return written
        label_cell.value = header
        style_hap_analysis_cell(label_cell)
        written.append(f"{ws.title}!{label_cell.coordinate}")
        row = 2
        for label, value in rows:
            while (
                cell_is_occupied_or_formula(ws, ws.cell(row, col).coordinate)
                or cell_is_occupied_or_formula(ws, ws.cell(row, col + 1).coordinate)
            ):
                row += 1
                if row > 80:
                    return written
            ws.cell(row, col).value = label
            style_hap_analysis_cell(ws.cell(row, col))
            dest = ws.cell(row, col + 1)
            dest.value = value
            style_hap_analysis_cell(dest)
            written.append(f"{ws.title}!{dest.coordinate}")
            row += 1
        return written

    def write_notes_section(
        self,
        ws: Worksheet,
        rows: list[tuple[str, Any]],
        *,
        header: str = NOTES_HEADER,
    ) -> list[str]:
        """Write explanatory notes below the occupied model. Never shifts cells."""
        if not rows:
            return []
        written: list[str] = []
        # Column B is hidden on the Last Quarter tabs: notes go to columns D (label) and E (text) there.
        b_hidden = bool(ws.column_dimensions["B"].hidden) if "B" in ws.column_dimensions else False
        label_col, value_col = (4, 5) if b_hidden else (1, 2)
        header_row = self._notes_header_row(ws, header)
        if header_row is None:
            end = discover_occupied_end_row(ws)
            row = max(end, 0) + 2
            while cell_is_occupied_or_formula(ws, ws.cell(row, 1).coordinate) or cell_is_occupied_or_formula(
                ws, ws.cell(row, value_col).coordinate
            ):
                row += 1
                if row > end + 500:
                    return written
            title = ws.cell(row, 1)
            title.value = header
            style_hap_analysis_cell(title)
            written.append(f"{ws.title}!{title.coordinate}")
            row += 1
        else:
            row = header_row + 1
        limit = row + 500
        for label, value in rows:
            while cell_is_occupied_or_formula(ws, ws.cell(row, label_col).coordinate) or cell_is_occupied_or_formula(
                ws, ws.cell(row, value_col).coordinate
            ):
                row += 1
                if row > limit:
                    return written
            label_cell = ws.cell(row, label_col)
            value_cell = ws.cell(row, value_col)
            label_cell.value = label
            style_hap_analysis_cell(label_cell)
            value_cell.value = value
            style_hap_analysis_cell(value_cell)
            written.append(f"{ws.title}!{value_cell.coordinate}")
            row += 1
        return written

    def place_analysis_notes(
        self,
        ws: Worksheet,
        rows: list[tuple[str, Any]],
        *,
        live_start_col: int | None = None,
    ) -> list[str]:
        """Keep formula rows beside the model; move prose to the bottom notes section."""
        live: list[tuple[str, Any]] = []
        prose: list[tuple[str, Any]] = []
        for label, value in rows:
            if isinstance(value, str) and value.startswith("="):
                live.append((label, value))
            else:
                prose.append((label, value))
        written: list[str] = []
        if live:
            written.extend(
                self.write_block(ws, rows=live, start_col=live_start_col)
                if live_start_col is not None
                else self.write_block(ws, rows=live)
            )
        if prose:
            written.extend(self.write_notes_section(ws, prose))
        return written

    @staticmethod
    def _notes_header_row(ws: Worksheet, header: str) -> int | None:
        max_row = ws.max_row or 1
        for row in range(1, max_row + 1):
            value = ws.cell(row, 1).value
            if isinstance(value, str) and value.strip() == header:
                return row
        return None
