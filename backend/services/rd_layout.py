"""R&D tab layout: year columns and the look-back shift.

The R&D tab shows expense (row 2), capitalized R&D asset (row 3) and amortization (row 4) for each year (row 1). Years before the first
displayed year are the look-back: a useful life of L years needs L - 1 earlier years of expense. When the template has fewer look-back
columns than that, the block B1:N4 is shifted to the right and the new, older years are added on its left (for a 7-year life and a
first displayed year of 2016, columns B to G hold 2010 to 2015). Formulas everywhere in the workbook that point into the moved block are
updated, exactly as Excel does when it inserts cells and shifts them right.
"""

from __future__ import annotations

import re
from copy import copy
from typing import Any

from openpyxl.utils import get_column_letter

from services.formula_utils import parse_cell_refs

RD = "R&D"
BLOCK_ROWS = (1, 2, 3, 4)
_HEADER = re.compile(r"=Inputs!\$?([A-Z]+)\$?1")


def rd_year_columns(ows, inputs_cols: dict[str, int]) -> dict[str, int]:
    """FY token -> column on the R&D tab, read from the header formulas (=Inputs!C1 ...), so a shifted layout works.

    Years that have no header yet (the new year of an annual update) follow the same offset as the others."""
    by_letter = {get_column_letter(c): fy for fy, c in inputs_cols.items() if str(fy).startswith("FY")}
    out: dict[str, int] = {}
    for col in range(2, (ows.max_column or 2) + 1):
        header = str(ows.cell(1, col).value or "").replace(" ", "")
        match = _HEADER.fullmatch(header)
        if match and match.group(1) in by_letter:
            out[by_letter[match.group(1)]] = col
    offset = 2
    if out:
        fy0, col0 = next(iter(out.items()))
        offset = col0 - inputs_cols[fy0]
    for fy, c in inputs_cols.items():
        if str(fy).startswith("FY") and fy not in out:
            out[fy] = c + offset
    return out


def block_width(ows) -> int:
    """Rightmost column that holds anything in rows 1 to 4."""
    last = 1
    for row in BLOCK_ROWS:
        for col in range(1, (ows.max_column or 1) + 1):
            if ows.cell(row, col).value not in (None, ""):
                last = max(last, col)
    return last


def shift_references(formula: str, *, in_rd_sheet: bool, count: int, last_col: int) -> str:
    """Rewrite references into the moved block (columns B..last_col, rows 1 to 4 of the R&D tab) by +count columns."""
    if not isinstance(formula, str) or not formula.startswith("="):
        return formula
    refs = parse_cell_refs(formula)
    pieces: list[str] = []
    last = 0
    previous_target = False
    previous_end = -1
    for ref in refs:
        pieces.append(formula[last: ref["start"]])
        sheet = ref["sheet"]
        range_follow = formula[previous_end:ref["start"]] == ":"        # second half of A1:B2 shares the first half's sheet
        if range_follow:
            target = previous_target
        else:
            target = (sheet is not None and sheet.replace("'", "").upper() == RD) or (sheet is None and in_rd_sheet)
        if target and ref["row"] in BLOCK_ROWS and 2 <= ref["col"] <= last_col:
            prefix = ""
            if ref["sheet"]:
                name = ref["sheet"]
                prefix = f"'{name}'!" if (" " in name or not name.replace("_", "").isalnum()) else f"{name}!"
            col_text = ("$" if ref["col_abs"] else "") + get_column_letter(ref["col"] + count)
            row_text = ("$" if ref["row_abs"] else "") + str(ref["row"])
            pieces.append(f"{prefix}{col_text}{row_text}")
        else:
            pieces.append(ref["text"])
        last = ref["end"]
        previous_target, previous_end = target, ref["end"]
    pieces.append(formula[last:])
    return "".join(pieces)


def insert_lookback_columns(wb, count: int) -> list[str]:
    """Shift B1:N4 of the R&D tab `count` columns to the right and add `count` older year columns on its left.

    Returns the addresses of the new header cells. Does nothing for count <= 0 or when the tab is missing."""
    if count <= 0 or RD not in wb.sheetnames:
        return []
    ows = wb[RD]
    last_col = block_width(ows)
    if last_col < 2:
        return []
    # 1. every formula in the workbook that points into the block follows it
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                value = cell.value
                if isinstance(value, str) and value.startswith("=") and ("R&D" in value.upper() or ws.title == RD):
                    new = shift_references(value, in_rd_sheet=(ws.title == RD), count=count, last_col=last_col)
                    if new != value:
                        cell.value = new
    # 2. move the cells (values, formulas, formats), right to left
    for col in range(last_col, 1, -1):
        for row in BLOCK_ROWS:
            src, dst = ows.cell(row, col), ows.cell(row, col + count)
            dst.value = src.value
            dst._style = copy(src._style)
            dst.number_format = src.number_format
    for col in range(2, 2 + count):
        for row in BLOCK_ROWS:
            cell = ows.cell(row, col)
            cell.value = None
            cell._style = copy(ows.cell(row, 2 + count)._style)
    for width_col in range(last_col + count, 1, -1):
        src_width = ows.column_dimensions[get_column_letter(max(2, width_col - count))].width
        if src_width:
            ows.column_dimensions[get_column_letter(width_col)].width = src_width
    # 3. new, older year headers: each one is the year before its right-hand neighbour
    headers: list[str] = []
    for col in range(2, 2 + count):
        nxt = get_column_letter(col + 1)
        ows.cell(1, col).value = f'="FY "&RIGHT({nxt}1,4)-1'
        headers.append(f"{RD}!{get_column_letter(col)}1")
    return headers
