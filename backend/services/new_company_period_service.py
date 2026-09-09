"""Detect the ten annual fiscal-year columns and the latest quarterly period."""

from __future__ import annotations

from pathlib import Path

from openpyxl import load_workbook

from models.new_company import TenYearPeriodReport
from models.quarterly_presentation import STATEMENT_SHEETS, QuarterlyStatementKind
from services.annual_period_service import (
    detect_workbook_years,
    _resolve_cell_year,
    _column_annual_preference,
    _DATA_START_COL,
    _MAX_COL,
)
from services.quarterly_review_service import _detect_fiscal_quarter

LQ_IS = STATEMENT_SHEETS[QuarterlyStatementKind.INCOME]

_REQUIRED_COUNT = 10
_CONTROL_ROWS = {2, 3}  # template start/end year controls (B3/C3), not FY column headers
_STATEMENT_HEADER_ROWS = (5, 6, 7, 8, 1)


def detect_ten_year_columns(ws, wb=None) -> dict[str, int]:
    """FY column map that ignores start/end-year control cells on rows 2–3."""
    workbook = wb or ws.parent
    mapping: dict[str, int] = {}
    scores: dict[str, int] = {}
    max_col = min(ws.max_column or 1, _MAX_COL)
    for col in range(_DATA_START_COL, max_col + 1):
        for row in _STATEMENT_HEADER_ROWS:
            if row in _CONTROL_ROWS:
                continue
            token = _resolve_cell_year(ws, row, col, workbook)
            if not token:
                continue
            pref = _column_annual_preference(ws, col)
            if token not in mapping or pref > scores.get(token, -999):
                if token in mapping and pref == scores.get(token, -999) and col > mapping[token]:
                    break
                mapping[token] = col
                scores[token] = pref
            break
    return {k: v for k, v in mapping.items() if str(k).startswith("FY")}


def detect_ten_year_workbook_years(path: Path) -> dict[str, int]:
    priority = (
        "Income - GAAP",
        "Balance Sheet - Standardized",
        "Cash Flow - Standardized",
        "Inputs",
    )
    wb = load_workbook(path, data_only=False)
    try:
        for sheet in priority:
            if sheet not in wb.sheetnames:
                continue
            cols = detect_ten_year_columns(wb[sheet], wb)
            if len(cols) >= 8:
                return cols
        cols = detect_workbook_years(path)
        return {k: v for k, v in cols.items() if str(k).startswith("FY")}
    finally:
        wb.close()


def _fy_year(token: str) -> int:
    digits = "".join(ch for ch in str(token) if ch.isdigit())
    return int(digits) if digits else 0


class NewCompanyPeriodService:
    """Template family, ten-year annual range, and latest quarter identification."""

    def detect(self, *, analysis_id: str, ticker: str, workbook_path: Path) -> TenYearPeriodReport:
        warnings: list[str] = []
        cols = detect_ten_year_workbook_years(Path(workbook_path))
        fy_map = {k: v for k, v in cols.items() if str(k).startswith("FY")}
        years = sorted(fy_map, key=_fy_year)
        duplicate: list[str] = []
        seen_cols: dict[int, str] = {}
        for fy, col in fy_map.items():
            if col in seen_cols:
                duplicate.append(f"{seen_cols[col]}|{fy}")
            else:
                seen_cols[col] = fy

        chronology_ok = years == sorted(years, key=_fy_year) and not duplicate
        if len(years) != _REQUIRED_COUNT:
            warnings.append(
                "NEW_COMPANY_ANNUAL_PERIOD_RANGE_INVALID: "
                f"expected {_REQUIRED_COUNT} fiscal years, found {len(years)} ({years})."
            )
        if years:
            expected = [f"FY{y}" for y in range(_fy_year(years[0]), _fy_year(years[0]) + len(years))]
            if years != expected:
                chronology_ok = False
                warnings.append(
                    f"NEW_COMPANY_ANNUAL_PERIOD_RANGE_INVALID: non-contiguous years {years}."
                )

        family, version = self._classify_template(Path(workbook_path))
        quarter, q_fy = self._detect_quarter(Path(workbook_path), years)
        if quarter is None:
            warnings.append("NEW_COMPANY_QUARTER_NOT_IDENTIFIED")

        status = "ok"
        if any("ANNUAL_PERIOD_RANGE_INVALID" in w for w in warnings) or not years:
            status = "NEW_COMPANY_ANNUAL_PERIOD_RANGE_INVALID"
        elif quarter is None:
            status = "NEW_COMPANY_QUARTER_NOT_IDENTIFIED"

        return TenYearPeriodReport(
            analysis_id=analysis_id,
            ticker=ticker,
            fiscal_years=years,
            year_columns=fy_map,
            start_year=years[0] if years else None,
            end_year=years[-1] if years else None,
            latest_quarter=quarter,
            latest_quarter_label=f"Q{quarter}" if quarter else None,
            latest_quarter_fiscal_year=q_fy,
            chronology_ok=chronology_ok,
            duplicate_columns=duplicate,
            template_family=family,
            template_version=version,
            warnings=warnings,
            status=status,
            summary=(
                f"Periods: {len(years)} FY columns {years[0] if years else '—'}–"
                f"{years[-1] if years else '—'}; latest quarter="
                f"{'Q'+str(quarter) if quarter else 'unidentified'}; "
                f"template={family or 'unrecognized'}."
            ),
        )

    @staticmethod
    def _classify_template(path: Path) -> tuple[str | None, str | None]:
        wb = load_workbook(path, data_only=False, read_only=True)
        try:
            names = set(wb.sheetnames)
        finally:
            wb.close()
        industrial = {
            "Income - GAAP",
            "Balance Sheet - Standardized",
            "Cash Flow - Standardized",
            "Inputs",
        }
        if industrial <= names:
            version = "industrial_v1"
            if {"R&D", "Leases", "Tax"} <= names:
                version = "industrial_v1_full"
            return "industrial_template", version
        if {"Income - GAAP", "Inputs"} <= names:
            return "industrial_template", "industrial_v1"
        return "unrecognized", None

    @staticmethod
    def _detect_quarter(path: Path, years: list[str]) -> tuple[int | None, str | None]:
        wb = load_workbook(path, data_only=False)
        try:
            q = None
            q_fy = years[-1] if years else None
            if LQ_IS in wb.sheetnames:
                q = _detect_fiscal_quarter(wb[LQ_IS])
                ws = wb[LQ_IS]
                for row in range(1, 10):
                    for col in range(1, 10):
                        val = ws.cell(row, col).value
                        if val is None:
                            continue
                        text = str(val).upper()
                        for token in years:
                            if token.replace("FY", "") in text and "Q" in text:
                                q_fy = token
            if q is None:
                for sheet in wb.sheetnames:
                    if "last quarter" in sheet.lower() or sheet.lower().startswith("lq"):
                        q = _detect_fiscal_quarter(wb[sheet])
                        if q:
                            break
            return q, q_fy
        finally:
            wb.close()
