"""Annual Update output-readiness gates (tax/PE10/R&D/valuation/report authorization)."""

from __future__ import annotations

from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from models.annual_update import (
    AnnualInputsReport,
    AnnualOutputGateReport,
    AnnualRdReport,
    AnnualTaxReport,
    AnnualValuationOutputs,
)
from services.annual_period_service import detect_year_columns
from services.excel_recalc_service import ExcelRecalcReport

_FORMULA_ERRORS = ("#REF!", "#DIV/0!", "#VALUE!", "#NAME?", "#N/A", "#NUM!", "#NULL!")
_ETR_TOL = 0.015


class AnnualOutputGateService:
    def evaluate(
        self,
        *,
        analysis_id: str,
        ticker: str,
        workbook_path: Path,
        fiscal_year: str | None,
        tax: AnnualTaxReport | None,
        inputs: AnnualInputsReport | None,
        rd: AnnualRdReport | None,
        valuation: AnnualValuationOutputs | None,
        recalc: ExcelRecalcReport | None = None,
    ) -> AnnualOutputGateReport:
        blockers: list[str] = []
        warnings: list[str] = []
        gates: dict[str, str] = {}

        # Gate — workbook recalculation (blocking)
        if recalc is None or recalc.status != "ok":
            blockers.append("WORKBOOK_RECALCULATION_INCOMPLETE")
            if recalc and recalc.missing_cached_values:
                blockers.append(
                    "WORKBOOK_RECALCULATION_INCOMPLETE: missing cached values for "
                    + ", ".join(recalc.missing_cached_values[:8])
                )
            if recalc and recalc.formula_errors:
                blockers.append(
                    "REQUIRED_OUTPUT_FORMULA_ERROR: " + "; ".join(recalc.formula_errors[:6])
                )
            if recalc and recalc.error:
                blockers.append(f"WORKBOOK_RECALCULATION_INCOMPLETE: {recalc.error}")
            gates["workbook_recalculation"] = "fail"
        else:
            gates["workbook_recalculation"] = "pass"

        # Gate — tax inputs + Tax sheet calculated outputs
        if tax is None or not tax.schedule_populated or not tax.cells_written:
            blockers.append("ANNUAL_TAX_SCHEDULE_NOT_POPULATED")
            gates["tax_schedule"] = "fail"
        else:
            tax_issues = self._verify_tax_sheet(workbook_path, fiscal_year, tax)
            if tax_issues:
                blockers.extend(tax_issues)
                gates["tax_schedule"] = "fail"
            else:
                gates["tax_schedule"] = "pass"
            if tax.reconciliation_status == "TAX_RECONCILIATION_REVIEW_REQUIRED":
                warnings.append(tax.reconciliation_status)

        # Gate — PE10 fiscal required; FY vs current difference is informational
        if inputs is None or inputs.pe10 is None or inputs.pe10.source_value is None:
            blockers.append("ANNUAL_PE10_REQUIRED_INPUT_MISSING")
            gates["pe10"] = "fail"
        else:
            gates["pe10"] = "pass"
            if inputs.pe10_mismatch and "PE10_CROSS_SHEET_MISMATCH" in inputs.pe10_mismatch:
                warnings.append(inputs.pe10_mismatch)
            elif inputs.pe10_mismatch and "AS_OF" in inputs.pe10_mismatch:
                # Informational period note — not a blocker
                warnings.append(inputs.pe10_mismatch)

        # Gate — R&D life, expense, calculated asset/amortization
        if rd is None or rd.useful_life is None:
            blockers.append("RD_USEFUL_LIFE_MISSING")
            gates["rd"] = "fail"
        elif not rd.lookback_complete:
            blockers.append("RD_LOOKBACK_INCOMPLETE")
            gates["rd"] = "fail"
        elif rd.rd_expense is None:
            blockers.append("RD_INPUT_NOT_CONNECTED_TO_SCHEDULE")
            gates["rd"] = "fail"
        elif rd.rd_asset is None or rd.rd_amortization is None:
            blockers.append("RD_SCHEDULE_NOT_EXTENDED")
            blockers.append(
                "WORKBOOK_RECALCULATION_INCOMPLETE: FY R&D asset/amortization cached values missing."
            )
            gates["rd"] = "fail"
        else:
            gates["rd"] = "pass"

        # Gate — required valuation sheet outputs after recalc
        if valuation is None:
            blockers.append("VALUATION_OUTPUT_NOT_EXTRACTED")
            gates["valuation_outputs"] = "fail"
        else:
            required_attrs = [
                ("expected_annual_return", "Expected Returns!E14"),
                ("expected_return_with_dividends", "Expected Returns!F14"),
                ("current_graham_intrinsic_value", "Enterprise Value!B42"),
                ("graham_target_return_entry_price", "Enterprise Value!B48"),
                ("graham_margin_of_safety_entry_price", "Graham MOS entry"),
                ("graham_expected_annualized_return", "Enterprise Value!B47"),
                ("nopat", "IC & NOPAT & ROIC NOPAT"),
                ("invested_capital", "IC & NOPAT & ROIC Invested Capital"),
                ("roic", "IC & NOPAT & ROIC ROIC"),
                ("roic_wacc", "Final Metrics ROIC-WACC"),
                ("roce", "Final Metrics ROCE"),
            ]
            missing = [label for attr, label in required_attrs if getattr(valuation, attr) is None]
            for w in valuation.warnings:
                if "WORKBOOK_RECALCULATION_INCOMPLETE" in w:
                    blockers.append(w)
                elif "IMPLAUSIBLE" in w:
                    warnings.append(w)
                elif "Distinct metrics" in w:
                    warnings.append(w)
                else:
                    warnings.append(w)
            if missing:
                blockers.append("VALUATION_OUTPUT_NOT_EXTRACTED: " + ", ".join(missing))
                blockers.append("WORKBOOK_RECALCULATION_INCOMPLETE")
                gates["valuation_outputs"] = "fail"
            else:
                gates["valuation_outputs"] = "pass"

        # Report authorization
        # Deduplicate blockers while preserving order
        seen: set[str] = set()
        uniq_blockers: list[str] = []
        for b in blockers:
            if b not in seen:
                seen.add(b)
                uniq_blockers.append(b)
        blockers = uniq_blockers

        if blockers:
            status = "NEEDS_REVIEW"
            gates["report_authorization"] = "blocked"
            if "REPORT_GENERATION_BLOCKED_BY_VALIDATION" not in blockers:
                blockers.append("REPORT_GENERATION_BLOCKED_BY_VALIDATION")
        else:
            status = "ok"
            gates["report_authorization"] = "authorized"

        return AnnualOutputGateReport(
            analysis_id=analysis_id,
            ticker=ticker,
            status=status,
            blockers=blockers,
            warnings=warnings,
            gates=gates,
            summary=(
                f"Annual gates={status}; blockers={len(blockers)}; warnings={len(warnings)}; "
                f"detail={', '.join(blockers[:4]) or 'none'}."
            ),
        )

    def _verify_tax_sheet(
        self, path: Path, fiscal_year: str | None, tax: AnnualTaxReport
    ) -> list[str]:
        issues: list[str] = []
        if not fiscal_year:
            return ["ANNUAL_TAX_SCHEDULE_NOT_POPULATED"]
        if not Path(path).exists() or not Path(path).is_file():
            return [
                "WORKBOOK_RECALCULATION_INCOMPLETE: workbook missing for Tax-sheet verification."
            ]
        try:
            wb = load_workbook(path, data_only=True)
            wb_f = load_workbook(path, data_only=False)
        except Exception as exc:  # noqa: BLE001
            return [f"WORKBOOK_RECALCULATION_INCOMPLETE: cannot open workbook for Tax verify ({exc})"]
        try:
            if "Inputs" not in wb.sheetnames:
                return ["ANNUAL_TAX_SCHEDULE_NOT_POPULATED"]
            cols = detect_year_columns(wb_f["Inputs"], wb_f)
            token = fiscal_year if fiscal_year.startswith("FY") else f"FY{fiscal_year}"
            col = cols.get(token)
            if not col:
                return ["ANNUAL_TAX_SCHEDULE_NOT_POPULATED: FY column not detected"]
            for row in (107, 108, 111, 112):
                val = wb["Inputs"].cell(row, col).value
                if val is None or val == "":
                    # rates may be zero for some houses; require federal + total at minimum
                    if row in (107, 112):
                        issues.append(f"ANNUAL_TAX_SCHEDULE_NOT_POPULATED: Inputs tax row {row} blank")
                elif isinstance(val, str) and any(tok in val for tok in _FORMULA_ERRORS):
                    issues.append(f"REQUIRED_OUTPUT_FORMULA_ERROR: Inputs!{row}={val}")

            if "Tax" not in wb.sheetnames:
                issues.append("ANNUAL_TAX_SCHEDULE_NOT_POPULATED: Tax sheet missing")
                return issues
            tax_ws = wb["Tax"]
            # Total Effective Tax Rate row (rate-mode pulls Inputs total / ETR)
            etr = tax_ws.cell(15, col).value
            tot = tax_ws.cell(25, col).value  # Total Operating Taxes
            if isinstance(etr, str) and any(tok in etr for tok in _FORMULA_ERRORS):
                issues.append(f"REQUIRED_OUTPUT_FORMULA_ERROR: Tax! effective rate={etr}")
            if etr is None or etr == "":
                issues.append(
                    "WORKBOOK_RECALCULATION_INCOMPLETE: Tax sheet effective rate has no cached value."
                )
            else:
                etr_n = float(etr) if isinstance(etr, (int, float)) else None
                if etr_n is not None:
                    tax.tax_sheet_effective_rate = etr_n
                    tax.tax_sheet_verified = True
                    if tax.reported_effective_tax_rate is not None:
                        if abs(etr_n - float(tax.reported_effective_tax_rate)) > _ETR_TOL:
                            # Also accept Inputs total rate match as recon proof
                            inputs_total = wb["Inputs"].cell(112, col).value
                            if not (
                                isinstance(inputs_total, (int, float))
                                and abs(float(inputs_total) - float(tax.reported_effective_tax_rate))
                                <= _ETR_TOL
                            ):
                                issues.append("TAX_EFFECTIVE_RATE_RECONCILIATION_FAILED")
            if tot is None or tot == "":
                issues.append(
                    "WORKBOOK_RECALCULATION_INCOMPLETE: Tax sheet total operating taxes uncached."
                )
            elif isinstance(tot, (int, float)):
                tax.tax_sheet_total_operating_taxes = float(tot)

            # Confirm Tax formulas still present (not overwritten)
            tax_f = wb_f["Tax"]
            sample = tax_f.cell(15, col).value
            if not (isinstance(sample, str) and sample.startswith("=")):
                issues.append("REQUIRED_OUTPUT_FORMULA_ERROR: Tax sheet formula overwritten")
        finally:
            wb.close()
            wb_f.close()
        return issues
