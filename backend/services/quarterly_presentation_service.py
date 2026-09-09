"""Apply quarterly presentation decisions (Bloomberg preserve/gaps or SEC layout)."""

from __future__ import annotations

import shutil
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from openpyxl.workbook.workbook import Workbook

from models.quarterly_presentation import (
    STATEMENT_SHEETS,
    PresentationDecision,
    QuarterlyPresentationReport,
    QuarterlyStatementKind,
    QuarterlyStatementPresentation,
    SecLineItem,
)
from services.accounting_concept_matcher import resolve_workbook_gap
from services.quarterly_health_service import (
    BODY_END,
    BODY_START,
    LABEL_COL,
    VALUE_COL,
    YTD_COL,
    assess_all_quarterly_statements,
    iter_statement_rows,
)
from services.sec_10q_statement_service import (
    extract_sec_10q_statement,
    select_latest_10q_period,
)
from services.yahoo_quarterly_statement_service import (
    BASIC_CF_LAYOUT,
    BASIC_IS_LAYOUT,
    YahooQuarterlyStatementService,
)


def _is_formula(value: Any) -> bool:
    return isinstance(value, str) and value.startswith("=")


class QuarterlyPresentationService:
    """
    Inspect LQ statements, decide Bloomberg vs Yahoo basic template vs SEC fallback,
    and apply Yahoo basic layout when Bloomberg cumulative comparison is missing.
    """

    def __init__(self) -> None:
        self.yahoo = YahooQuarterlyStatementService()

    def plan_and_apply(
        self,
        *,
        analysis_id: str,
        ticker: str,
        source_workbook_path: Path,
        destination_workbook_path: Path,
        company_facts: dict[str, Any] | None,
        already_copied: bool = False,
    ) -> QuarterlyPresentationReport:
        """
        Copy source→destination (unless already_copied), assess each statement,
        apply SEC layout only where SEC_10Q_PRESENTATION_REQUIRED.
        Leaves source unchanged.
        """
        if not already_copied:
            destination_workbook_path.parent.mkdir(parents=True, exist_ok=True)
            shutil.copy2(source_workbook_path, destination_workbook_path)

        wb = load_workbook(destination_workbook_path)
        try:
            fy, fp = (None, None)
            if company_facts:
                fy, fp = select_latest_10q_period(company_facts)
            yahoo_bundle = self.yahoo.fetch(ticker)
            fiscal_q = self._detect_fiscal_quarter(wb)

            statements: list[QuarterlyStatementPresentation] = []
            for health, decision in assess_all_quarterly_statements(wb):
                entry = QuarterlyStatementPresentation(
                    statement=health.statement,
                    sheet=health.sheet,
                    health=health,
                    decision=decision,
                    reason=health.reason,
                    bloomberg_populated=health.populated_rows,
                    bloomberg_missing=health.missing_required_rows,
                )
                if decision == PresentationDecision.BLOOMBERG_PRESERVE:
                    rows = iter_statement_rows(wb[health.sheet]) if health.present else []
                    entry.rows_preserved = [r["cell_ref"] for r in rows if r["populated"] or r["formula"]]
                    entry.reason = f"BLOOMBERG_PRESERVE — {health.reason}"
                elif decision == PresentationDecision.BLOOMBERG_FILL_GAPS:
                    rows = iter_statement_rows(wb[health.sheet]) if health.present else []
                    entry.rows_preserved = [r["cell_ref"] for r in rows if r["populated"] or r["formula"]]
                    filled = self._fill_major_gaps(
                        wb,
                        health.statement,
                        health.major_totals_missing,
                        company_facts,
                        yahoo_bundle,
                        fy,
                        fp,
                        fiscal_q,
                    )
                    entry.rows_filled = filled
                    entry.data_source_primary = "yahoo"
                    entry.data_source_secondary = "sec"
                    entry.reason = f"BLOOMBERG_FILL_GAPS — {health.reason}"
                elif decision == PresentationDecision.BLOCKED:
                    entry.blocked_or_ambiguous = [health.reason]
                    entry.reason = f"BLOCKED — {health.reason}"
                elif decision == PresentationDecision.YAHOO_BASIC_TEMPLATE_REQUIRED:
                    yahoo_periods = self.yahoo.values_for_periods(
                        yahoo_bundle,
                        health.statement,
                        fiscal_quarter=fiscal_q,
                        company_facts=company_facts,
                        fiscal_year=fy,
                        fiscal_period=fp,
                    )
                    applied = self._apply_yahoo_basic_template(
                        wb,
                        health.statement,
                        yahoo_periods,
                        fiscal_quarter=fiscal_q,
                    )
                    entry.yahoo_rows_introduced = applied["introduced"]
                    entry.ytd_provenance = applied.get("ytd_provenance", {})
                    entry.rows_superseded = applied["superseded"]
                    entry.blocked_or_ambiguous = applied["ambiguous"]
                    entry.data_source_primary = "yahoo"
                    # SEC secondary for any still-missing majors
                    if company_facts and applied["ambiguous"]:
                        sec_fill = self._fill_major_gaps_from_sec(
                            wb,
                            health.statement,
                            health.major_totals_missing,
                            company_facts,
                            fy,
                            fp,
                        )
                        entry.rows_filled = sec_fill
                        entry.data_source_secondary = "sec"
                    entry.reason = (
                        f"YAHOO_BASIC_TEMPLATE — Bloomberg cumulative missing; "
                        f"populated basic template from Yahoo ({len(applied['introduced'])} rows)"
                    )
                elif decision == PresentationDecision.SEC_10Q_PRESENTATION_REQUIRED:
                    sec_items = (
                        extract_sec_10q_statement(
                            company_facts or {},
                            health.statement,
                            fiscal_year=fy,
                            fiscal_period=fp,
                        )
                        if company_facts
                        else []
                    )
                    entry.sec_line_items = sec_items
                    if sec_items:
                        entry.sec_filing_form = sec_items[0].form
                        entry.sec_filing_period = sec_items[0].fiscal_period
                        entry.sec_accession = sec_items[0].accession_number
                    ytd_items = []
                    if health.statement == QuarterlyStatementKind.INCOME and company_facts:
                        ytd_items = extract_sec_10q_statement(
                            company_facts,
                            health.statement,
                            fiscal_year=fy,
                            fiscal_period=fp,
                            duration_kind="ytd",
                        )
                    applied = self._apply_sec_layout(wb, health.statement, sec_items, ytd_items=ytd_items)
                    entry.rows_superseded = applied["superseded"]
                    entry.sec_rows_introduced = applied["introduced"]
                    entry.blocked_or_ambiguous = applied["ambiguous"]
                    entry.reason = (
                        f"SEC_10Q_PRESENTATION_REQUIRED — abandoned Bloomberg layout; "
                        f"reconstructed from 10-Q presentation "
                        f"({len(sec_items)} SEC lines, form={entry.sec_filing_form}, "
                        f"fp={entry.sec_filing_period})"
                    )
                    if not sec_items:
                        entry.blocked_or_ambiguous.append(
                            "No SEC 10-Q line items extracted; layout not rewritten"
                        )
                statements.append(entry)

            preserve = sum(
                1 for s in statements if s.decision == PresentationDecision.BLOOMBERG_PRESERVE
            )
            gaps = sum(
                1 for s in statements if s.decision == PresentationDecision.BLOOMBERG_FILL_GAPS
            )
            yahoo_n = sum(
                1
                for s in statements
                if s.decision == PresentationDecision.YAHOO_BASIC_TEMPLATE_REQUIRED
            )
            sec_n = sum(
                1
                for s in statements
                if s.decision == PresentationDecision.SEC_10Q_PRESENTATION_REQUIRED
            )
            blocked = sum(1 for s in statements if s.decision == PresentationDecision.BLOCKED)
            summary = (
                f"Quarterly presentation: PRESERVE={preserve}, FILL_GAPS={gaps}, "
                f"YAHOO_BASIC={yahoo_n}, SEC_10Q={sec_n}, BLOCKED={blocked}"
            )
            wb.save(destination_workbook_path)
            return QuarterlyPresentationReport(
                analysis_id=analysis_id,
                ticker=ticker,
                statements=statements,
                summary=summary,
            )
        finally:
            wb.close()

    @staticmethod
    def _detect_fiscal_quarter(wb) -> int:
        from services.quarterly_review_service import _detect_fiscal_quarter

        sheet = STATEMENT_SHEETS[QuarterlyStatementKind.INCOME]
        if sheet in wb.sheetnames:
            q = _detect_fiscal_quarter(wb[sheet])
            if q:
                return q
        return 3

    def _fill_major_gaps(
        self,
        workbook: Workbook,
        kind: QuarterlyStatementKind,
        missing_majors: list[str],
        company_facts: dict[str, Any] | None,
        yahoo_bundle,
        fy: int | None,
        fp: str | None,
        fiscal_q: int,
    ) -> list[str]:
        filled = self._fill_major_gaps_from_yahoo(
            workbook, kind, missing_majors, yahoo_bundle, fiscal_q, company_facts, fy, fp
        )
        still_missing = [m for m in missing_majors if not any(m in f for f in filled)]
        if still_missing and company_facts:
            filled.extend(
                self._fill_major_gaps_from_sec(
                    workbook, kind, still_missing, company_facts, fy, fp
                )
            )
        return filled

    def _fill_major_gaps_from_yahoo(
        self,
        workbook: Workbook,
        kind: QuarterlyStatementKind,
        missing_majors: list[str],
        yahoo_bundle,
        fiscal_q: int,
        company_facts: dict[str, Any] | None = None,
        fy: int | None = None,
        fp: str | None = None,
    ) -> list[str]:
        filled: list[str] = []
        if not missing_majors:
            return filled
        sheet_name = STATEMENT_SHEETS[kind]
        if sheet_name not in workbook.sheetnames:
            return filled
        ws = workbook[sheet_name]
        values = self.yahoo.values_for_periods(
            yahoo_bundle,
            kind,
            fiscal_quarter=fiscal_q,
            company_facts=company_facts,
            fiscal_year=fy,
            fiscal_period=fp,
        )
        items = self.yahoo.to_sec_line_items(values, kind, fiscal_quarter=fiscal_q)

        def _n(s: Any) -> str:
            return " ".join(str(s or "").lower().split())

        for needle in missing_majors:
            for r in range(BODY_START, BODY_END + 1):
                label = ws.cell(row=r, column=LABEL_COL).value
                value = ws.cell(row=r, column=VALUE_COL).value
                if label is None or needle not in _n(label):
                    continue
                if _is_formula(value) or isinstance(value, (int, float)):
                    continue
                match = resolve_workbook_gap(str(label).strip(), kind, items)
                if match.decision != "MATCHED" or match.value is None:
                    continue
                ws.cell(row=r, column=VALUE_COL).value = match.value
                filled.append(
                    f"{sheet_name}!C{r}:{label}={match.value} [yahoo concept={match.accounting_concept}]"
                )
                break
        return filled

    def _apply_yahoo_basic_template(
        self,
        workbook: Workbook,
        kind: QuarterlyStatementKind,
        yahoo_periods,
        *,
        fiscal_quarter: int,
    ) -> dict[str, list[str] | dict]:
        from services.yahoo_quarterly_statement_service import YahooPeriodValues

        period_values = (
            yahoo_periods
            if isinstance(yahoo_periods, YahooPeriodValues)
            else YahooPeriodValues(values=yahoo_periods)
        )
        yahoo_values = period_values.values
        ytd_prov = self.yahoo.provenance_report(period_values)

        sheet_name = STATEMENT_SHEETS[kind]
        result: dict[str, Any] = {"superseded": [], "introduced": [], "ambiguous": [], "ytd_provenance": ytd_prov}
        if sheet_name not in workbook.sheetnames:
            result["ambiguous"].append(f"Sheet {sheet_name} missing")
            return result

        layout = BASIC_IS_LAYOUT if kind == QuarterlyStatementKind.INCOME else BASIC_CF_LAYOUT
        ws = workbook[sheet_name]

        # Clear Bloomberg body (preserve formulas)
        for r in range(BODY_START, BODY_END + 1):
            lc = ws.cell(row=r, column=LABEL_COL)
            vc = ws.cell(row=r, column=VALUE_COL)
            if _is_formula(lc.value) or _is_formula(vc.value):
                continue
            if lc.value is not None or vc.value is not None:
                result["superseded"].append(f"{sheet_name}!A{r}/C{r}:{lc.value!r}")
                lc.value = None
                vc.value = None
                if kind == QuarterlyStatementKind.INCOME:
                    for col in (4, 7, 8):
                        if not _is_formula(ws.cell(row=r, column=col).value):
                            ws.cell(row=r, column=col).value = None

        for row, label, field_name, _ in layout:
            ws.cell(row=row, column=LABEL_COL).value = label
            fq = yahoo_values.get(f"fq_{field_name}")
            py = yahoo_values.get(f"py_{field_name}")
            if fq is not None:
                prov = period_values.ytd_provenance.get(field_name)
                prov_tag = f" [{prov.provenance}]" if prov else ""
                ws.cell(row=row, column=VALUE_COL).value = fq
                result["introduced"].append(f"{sheet_name}!C{row}:{label}={fq} [yahoo fq]{prov_tag}")
            else:
                result["ambiguous"].append(f"Missing Yahoo FQ value for {label}")
            if py is not None:
                ws.cell(row=row, column=4).value = py  # col D prior-year FQ
            if fiscal_quarter >= 2:
                ytd = yahoo_values.get(f"ytd_{field_name}")
                py_ytd = yahoo_values.get(f"py_ytd_{field_name}")
                if ytd is not None:
                    prov = period_values.ytd_provenance.get(field_name)
                    tag = prov.provenance if prov else "unknown"
                    comp = f" components={prov.components}" if prov and prov.components else ""
                    ws.cell(row=row, column=YTD_COL).value = ytd
                    result["introduced"].append(
                        f"{sheet_name}!G{row}:{label}={ytd} [ytd {tag}{comp}]"
                    )
                if py_ytd is not None:
                    ws.cell(row=row, column=8).value = py_ytd

        banner = ws.cell(row=7, column=LABEL_COL)
        if banner.value is None or not _is_formula(banner.value):
            banner.value = "Yahoo Finance basic template (fallback) — values in millions"

        return result

    def _fill_major_gaps_from_sec(
        self,
        workbook: Workbook,
        kind: QuarterlyStatementKind,
        missing_majors: list[str],
        company_facts: dict[str, Any] | None,
        fy: int | None,
        fp: str | None,
    ) -> list[str]:
        """Fill blank major-line VALUE cells via accounting-concept SEC matching."""
        filled: list[str] = []
        if not missing_majors or not company_facts:
            return filled
        sheet_name = STATEMENT_SHEETS[kind]
        if sheet_name not in workbook.sheetnames:
            return filled
        ws = workbook[sheet_name]
        items = extract_sec_10q_statement(
            company_facts, kind, fiscal_year=fy, fiscal_period=fp
        )
        if not items:
            return filled

        def _n(s: Any) -> str:
            return " ".join(str(s or "").lower().split())

        for needle in missing_majors:
            target_row = None
            workbook_label = needle
            for r in range(BODY_START, BODY_END + 1):
                label = ws.cell(row=r, column=LABEL_COL).value
                value = ws.cell(row=r, column=VALUE_COL).value
                if label is None:
                    continue
                if needle not in _n(label):
                    continue
                if _is_formula(value) or isinstance(value, (int, float)):
                    continue
                target_row = r
                workbook_label = str(label).strip()
                break
            if target_row is None:
                continue
            match = resolve_workbook_gap(workbook_label, kind, items)
            if match.decision != "MATCHED" or match.value is None:
                filled.append(
                    f"{sheet_name}!C{target_row}:{workbook_label} "
                    f"[{match.decision}] {match.reason}"
                )
                continue
            ws.cell(row=target_row, column=VALUE_COL).value = match.value
            filled.append(
                f"{sheet_name}!C{target_row}:{workbook_label}={match.value} "
                f"[concept={match.accounting_concept} xbrl={match.sec_xbrl_concept} "
                f"method={match.match_method} conf={match.confidence}]"
            )
        return filled

    def _apply_sec_layout(
        self,
        workbook: Workbook,
        kind: QuarterlyStatementKind,
        sec_items: list[SecLineItem],
        ytd_items: list[SecLineItem] | None = None,
    ) -> dict[str, list[str]]:
        sheet_name = STATEMENT_SHEETS[kind]
        result = {"superseded": [], "introduced": [], "ambiguous": []}
        if sheet_name not in workbook.sheetnames:
            result["ambiguous"].append(f"Sheet {sheet_name} missing")
            return result
        if not sec_items:
            return result

        ws = workbook[sheet_name]
        # Snapshot Bloomberg non-formula body rows to supersede.
        for r in range(BODY_START, BODY_END + 1):
            label_cell = ws.cell(row=r, column=LABEL_COL)
            value_cell = ws.cell(row=r, column=VALUE_COL)
            if _is_formula(label_cell.value) or _is_formula(value_cell.value):
                continue  # never overwrite formulas
            if label_cell.value is not None or value_cell.value is not None:
                result["superseded"].append(f"{sheet_name}!A{r}/C{r}:{label_cell.value!r}")
                label_cell.value = None
                value_cell.value = None

        # Write SEC presentation in order, skipping formula rows.
        write_row = BODY_START
        for item in sec_items:
            while write_row <= BODY_END:
                if _is_formula(ws.cell(row=write_row, column=LABEL_COL).value) or _is_formula(
                    ws.cell(row=write_row, column=VALUE_COL).value
                ):
                    write_row += 1
                    continue
                break
            if write_row > BODY_END:
                result["ambiguous"].append(f"No space for SEC line {item.label}")
                break
            # YTD cash-flow labels must remain explicitly YTD — never as quarter.
            label = item.label
            if item.duration_kind == "ytd":
                label = f"{item.label} (YTD)"
            ws.cell(row=write_row, column=LABEL_COL).value = label
            ws.cell(row=write_row, column=VALUE_COL).value = item.value
            if ytd_items:
                ytd_hit = next(
                    (y for y in ytd_items if y.label == item.label and y.value is not None),
                    None,
                )
                if ytd_hit is not None:
                    ytd_label = f"{ytd_hit.label} (YTD)" if ytd_hit.duration_kind == "ytd" else ytd_hit.label
                    _ = ytd_label
                    ws.cell(row=write_row, column=YTD_COL).value = ytd_hit.value
            result["introduced"].append(
                f"{sheet_name}!A{write_row}:{label}={item.value} "
                f"[{item.xbrl_concept}/{item.duration_kind}]"
            )
            write_row += 1

        # Annotate header banner if a free text cell exists at row 7
        banner = ws.cell(row=7, column=LABEL_COL)
        if banner.value is None or not _is_formula(banner.value):
            form = sec_items[0].form or "10-Q"
            fp = sec_items[0].fiscal_period or ""
            banner.value = f"SEC {form} presentation authority ({fp}) — not Bloomberg taxonomy"

        return result
