"""Fiscal-year detection and period alignment for Annual Update."""

from __future__ import annotations

import re
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.utils import column_index_from_string
from openpyxl.worksheet.worksheet import Worksheet

# Inputs puts FY headers on row 1; statement sheets use rows 5–8.
_HEADER_ROWS = (1, 2, 3, 4, 5, 6, 7, 8)
_DATA_START_COL = 3
_MAX_COL = 20
_YEAR_RE = re.compile(r"(20\d{2})")
_CROSS_REF_RE = re.compile(
    r"^=\s*'?([^'!]+)'?\s*!\s*([A-Z]+)([0-9]+)\s*$",
    re.IGNORECASE,
)
# Prefer annual "YYYY A" / FY tokens over Bloomberg side columns like "2026 Y".
_ANNUAL_MARKER_RE = re.compile(r"(20\d{2})\s*A\b|^FY\s*20\d{2}\b", re.IGNORECASE)


class AnnualPeriodAlignmentError(Exception):
  """Raised when template/previous fiscal windows cannot be aligned safely."""


@dataclass(frozen=True)
class AnnualPeriodAlignment:
    template_years: tuple[str, ...]
    previous_years: tuple[str, ...]
    overlap_years: tuple[str, ...]
    dropped_years: tuple[str, ...]
    new_fiscal_year: str
    template_end_year: str | None
    previous_end_year: str | None


def _fy_token(value: object) -> str | None:
    if value is None:
        return None
    if isinstance(value, datetime):
        return f"FY{value.year}"
    text = str(value).strip().upper()
    # FY 2025 / FY2025
    compact = text.replace(" ", "")
    if compact.startswith("FY") and len(compact) >= 6 and compact[2:6].isdigit():
        return f"FY{compact[2:6]}"
    m = _YEAR_RE.search(text)
    if m:
        return f"FY{m.group(1)}"
    return None


def _resolve_cell_year(ws: Worksheet, row: int, col: int, wb) -> str | None:
    cell = ws.cell(row, col)
    token = _fy_token(cell.value)
    if token:
        return token
    value = cell.value
    if not isinstance(value, str) or not value.startswith("="):
        return None
    ref = _CROSS_REF_RE.match(value.strip())
    if ref:
        sheet_name, col_letter, row_num = ref.group(1), ref.group(2), int(ref.group(3))
        if sheet_name in wb.sheetnames:
            target = wb[sheet_name]
            return _resolve_cell_year(target, row_num, column_index_from_string(col_letter), wb)
    # ="FY "&LEFT(C5,4) — read same-column row 5
    if "LEFT(" in value.upper() and row in (7, 8):
        return _fy_token(ws.cell(5, col).value)
    return None


def _column_annual_preference(ws: Worksheet, col: int) -> int:
    """Higher score = more likely a true annual FY column (not EQY_RECENT / YTD)."""
    score = 0
    for row in _HEADER_ROWS:
        raw = ws.cell(row, col).value
        if raw is None:
            continue
        text = str(raw).strip().upper()
        if _ANNUAL_MARKER_RE.search(text):
            score += 3
        if text.endswith(" Y") or "RECENT" in text or "EQY_" in text:
            score -= 5
        if isinstance(raw, str) and raw.startswith("=") and "LEFT(" in raw.upper():
            score += 1
    return score


def detect_year_columns(ws: Worksheet, wb=None) -> dict[str, int]:
    """Map FY tokens to column indices, resolving formula headers when possible.

    Prefers the leftmost annual column for each FY. Does not let Bloomberg
    side columns (e.g. ``2026 Y``) overwrite the primary FY window.
    """
    workbook = wb or ws.parent
    mapping: dict[str, int] = {}
    scores: dict[str, int] = {}
    max_col = min(ws.max_column or 1, _MAX_COL)
    for col in range(_DATA_START_COL, max_col + 1):
        for row in _HEADER_ROWS:
            token = _resolve_cell_year(ws, row, col, workbook)
            if not token:
                continue
            pref = _column_annual_preference(ws, col)
            if token not in mapping or pref > scores.get(token, -999):
                # Leftmost wins on ties.
                if token in mapping and pref == scores.get(token, -999) and col > mapping[token]:
                    break
                mapping[token] = col
                scores[token] = pref
            break
    if any(k.startswith("FY") for k in mapping):
        return {k: v for k, v in mapping.items() if k.startswith("FY")}
    for col in range(_DATA_START_COL, max_col + 1):
        mapping[f"COL{col}"] = col
    return mapping


def detect_workbook_years(path: Path, *, sheet_priority: tuple[str, ...] = ()) -> dict[str, int]:
    """Authoritative FY map for a workbook file."""
    default_priority = (
        "Balance Sheet - Standardized",
        "Income - GAAP",
        "Cash Flow - Standardized",
        "Inputs",
    )
    priority = sheet_priority or default_priority
    wb = load_workbook(path, data_only=False)
    try:
        for sheet in priority:
            if sheet not in wb.sheetnames:
                continue
            cols = detect_year_columns(wb[sheet], wb)
            if any(k.startswith("FY") for k in cols):
                return cols
        for sheet in wb.sheetnames:
            cols = detect_year_columns(wb[sheet], wb)
            if any(k.startswith("FY") for k in cols):
                return cols
        return {}
    finally:
        wb.close()


def _end_year_token(path: Path) -> str | None:
    wb = load_workbook(path, data_only=False)
    try:
        for sheet in ("Balance Sheet - Standardized", "Income - GAAP", "Inputs"):
            if sheet not in wb.sheetnames:
                continue
            ws = wb[sheet]
            token = _resolve_cell_year(ws, 3, 3, wb) or _fy_token(ws["C3"].value)
            if token:
                return token
        return None
    finally:
        wb.close()


def align_fiscal_periods(template_path: Path, previous_path: Path) -> AnnualPeriodAlignment:
    """Compare rolling windows; require exactly one new fiscal year in the template."""
    template_cols = detect_workbook_years(template_path)
    previous_cols = detect_workbook_years(previous_path)
    template_years = tuple(sorted(k for k in template_cols if k.startswith("FY")))
    previous_years = tuple(sorted(k for k in previous_cols if k.startswith("FY")))
    if not template_years:
        raise AnnualPeriodAlignmentError(
            "ANNUAL_PERIOD_ALIGNMENT_BLOCKED: could not detect fiscal years in current template."
        )
    if not previous_years:
        raise AnnualPeriodAlignmentError(
            "ANNUAL_PERIOD_ALIGNMENT_BLOCKED: could not detect fiscal years in previous workbook."
        )
    overlap = tuple(sorted(set(template_years) & set(previous_years)))
    dropped = tuple(sorted(set(previous_years) - set(template_years)))
    new_years = tuple(sorted(set(template_years) - set(previous_years)))
    if len(new_years) != 1:
        raise AnnualPeriodAlignmentError(
            "ANNUAL_PERIOD_ALIGNMENT_BLOCKED: expected exactly one new fiscal year in template; "
            f"found {new_years or 'none'} (template={template_years}, previous={previous_years})."
        )
    return AnnualPeriodAlignment(
        template_years=template_years,
        previous_years=previous_years,
        overlap_years=overlap,
        dropped_years=dropped,
        new_fiscal_year=new_years[0],
        template_end_year=_end_year_token(template_path),
        previous_end_year=_end_year_token(previous_path),
    )


def year_columns_for_sheet(path: Path, sheet: str) -> dict[str, int]:
    wb = load_workbook(path, data_only=False)
    try:
        if sheet not in wb.sheetnames:
            return {}
        return detect_year_columns(wb[sheet], wb)
    finally:
        wb.close()
