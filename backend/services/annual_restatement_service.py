"""Historical restatement: automatic change only when the annual filing shows explicit revised comparatives."""

from __future__ import annotations

from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

from models.annual_update import AnnualRestatementReport, ContinuityAction, RestatementChange
from services.annual_continuity_service import detect_year_columns
from services.accounting_concept_matcher import CONCEPT_ALIASES


def _num(v: Any) -> float | None:
    if v is None or v == "":
        return None
    if isinstance(v, bool):
        return None
    if isinstance(v, (int, float)):
        return float(v)
    try:
        return float(str(v).replace(",", ""))
    except (TypeError, ValueError):
        return None


_SHEET_FOR_CONCEPT = {
    "revenue": "Income - GAAP",
    "operating_income": "Income - GAAP",
    "net_income": "Income - GAAP",
    "total_assets": "Balance Sheet - Standardized",
    "total_liabilities": "Balance Sheet - Standardized",
    "equity": "Balance Sheet - Standardized",
    "cfo": "Cash Flow - Standardized",
    "cfi": "Cash Flow - Standardized",
    "cff": "Cash Flow - Standardized",
}


class AnnualRestatementService:
    """Apply historical changes only with explicit revised comparative figures."""

    def apply(
        self,
        *,
        analysis_id: str,
        ticker: str,
        workbook_path: Path,
        explicit_revisions: list[dict[str, Any]] | None = None,
        narrative_disclosures: list[dict[str, Any]] | None = None,
    ) -> AnnualRestatementReport:
        automatic: list[RestatementChange] = []
        review: list[RestatementChange] = []

        for note in narrative_disclosures or []:
            if note.get("revised_comparatives"):
                continue
            review.append(
                RestatementChange(
                    sheet=str(note.get("sheet") or ""),
                    cell=str(note.get("cell") or ""),
                    fiscal_year=str(note.get("fiscal_year") or ""),
                    accounting_concept=str(note.get("concept") or "unspecified"),
                    annual_report_source=note.get("source"),
                    page_or_section=note.get("section"),
                    reason=(
                        str(note.get("explanation") or "Disclosure suggests a possible historical change")
                        + " — revised comparatives were not explicitly presented, so HAP did not rewrite history."
                    ),
                    confidence=float(note.get("confidence") or 0.4),
                    automatic_change=False,
                    action=ContinuityAction.RESTATEMENT_REVIEW_REQUIRED,
                )
            )

        if explicit_revisions:
            wb = load_workbook(workbook_path, data_only=False)
            try:
                for rev in explicit_revisions:
                    concept = str(rev.get("concept") or "")
                    fy = str(rev.get("fiscal_year") or "")
                    new_val = _num(rev.get("revised_value"))
                    if new_val is None or not fy:
                        review.append(
                            RestatementChange(
                                sheet="",
                                cell="",
                                fiscal_year=fy,
                                accounting_concept=concept,
                                annual_report_source=rev.get("source"),
                                page_or_section=rev.get("section"),
                                reason="Revision mentioned without a usable comparative amount — REVIEW_REQUIRED.",
                                automatic_change=False,
                                action=ContinuityAction.RESTATEMENT_REVIEW_REQUIRED,
                            )
                        )
                        continue
                    change = self._write_revision(wb, concept, fy, new_val, rev)
                    if change.automatic_change:
                        automatic.append(change)
                    else:
                        review.append(change)
                wb.save(workbook_path)
            finally:
                wb.close()

        status = "ok"
        if review and automatic:
            status = "partial"
        elif review:
            status = "RESTATEMENT_REVIEW_REQUIRED"
        return AnnualRestatementReport(
            analysis_id=analysis_id,
            ticker=ticker,
            automatic_changes=automatic,
            review_required=review,
            status=status,
            summary=(
                f"Restatement: {len(automatic)} automatic explicit revision(s), "
                f"{len(review)} REVIEW_REQUIRED."
            ),
        )

    def _write_revision(self, wb, concept: str, fy: str, new_val: float, rev: dict) -> RestatementChange:
        sheet_name = str(rev.get("sheet") or _SHEET_FOR_CONCEPT.get(concept) or "")
        if sheet_name not in wb.sheetnames:
            return RestatementChange(
                sheet=sheet_name,
                cell="",
                fiscal_year=fy,
                revised_reported_value=new_val,
                accounting_concept=concept,
                annual_report_source=rev.get("source"),
                page_or_section=rev.get("section"),
                reason="Target sheet missing — REVIEW_REQUIRED.",
                automatic_change=False,
                action=ContinuityAction.RESTATEMENT_REVIEW_REQUIRED,
            )
        ws = wb[sheet_name]
        cols = detect_year_columns(ws)
        token = fy if fy.startswith("FY") else f"FY{fy}"
        col = cols.get(token) or cols.get(fy)
        row = rev.get("row")
        if row is None:
            row = self._find_row(ws, concept, rev.get("label"))
        if col is None or row is None:
            return RestatementChange(
                sheet=sheet_name,
                cell="",
                fiscal_year=token,
                revised_reported_value=new_val,
                accounting_concept=concept,
                annual_report_source=rev.get("source"),
                page_or_section=rev.get("section"),
                reason="Could not uniquely map concept/year to a cell — REVIEW_REQUIRED.",
                automatic_change=False,
                action=ContinuityAction.RESTATEMENT_REVIEW_REQUIRED,
            )
        addr = f"{get_column_letter(int(col))}{int(row)}"
        cell = ws[addr]
        old = cell.value
        if isinstance(old, str) and old.startswith("="):
            return RestatementChange(
                sheet=sheet_name,
                cell=addr,
                fiscal_year=token,
                old_workbook_value=old,
                revised_reported_value=new_val,
                accounting_concept=concept,
                annual_report_source=rev.get("source"),
                page_or_section=rev.get("section"),
                reason="Target historical cell is a formula — blocked; REVIEW_REQUIRED.",
                automatic_change=False,
                action=ContinuityAction.BLOCKED,
            )
        cell.value = new_val
        return RestatementChange(
            sheet=sheet_name,
            cell=addr,
            fiscal_year=token,
            old_workbook_value=old,
            revised_reported_value=new_val,
            accounting_concept=concept,
            annual_report_source=rev.get("source"),
            page_or_section=rev.get("section"),
            reason=str(rev.get("reason") or "Annual filing explicitly presents revised comparatives."),
            confidence=float(rev.get("confidence") or 0.9),
            automatic_change=True,
            action=ContinuityAction.EXPLICIT_RESTATEMENT_UPDATE,
        )

    @staticmethod
    def _find_row(ws, concept: str, label: str | None) -> int | None:
        aliases = list(CONCEPT_ALIASES.get(concept, ()))
        if label:
            aliases.append(str(label).lower())
        aliases = [a.lower() for a in aliases]
        hits: list[int] = []
        for row in range(1, min(ws.max_row or 1, 120) + 1):
            text = str(ws.cell(row, 1).value or "").strip().lower()
            if text and any(a in text or text in a for a in aliases):
                hits.append(row)
        if len(hits) == 1:
            return hits[0]
        return None
