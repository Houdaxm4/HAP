"""Write the cost-of-revenue recast into a New Company workbook, flagging every changed cell.

Provider rows are hard-coded values, so moving a cost between lines must keep every total intact: for each recast
year the cost, gross profit, operating-expense total, R&D and other-operating-expense cells are rewritten from the
filing so that revenue - cost - opex still equals operating income. A year is skipped (and reported) unless the
workbook and the filing reconcile first. Original values are kept in each cell's comment.
"""

from __future__ import annotations

from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from services.annual_period_service import detect_year_columns
from services.cost_recast import (
    TOLERANCE,
    load_filings,
    recast_cost,
    statement_rows_for_year,
    statement_total_for_year,
)
from services.workbook_flag_service import flag_recast

SHEET = "Income - GAAP"
ROW_KEYS = {
    "revenue": "SALES_REV_TURN",
    "cost": "IS_COGS_TO_FE_AND_PP_AND_G",
    "cost_detail": "IS_COG_AND_SERVICES_SOLD",
    "gross_profit": "GROSS_PROFIT",
    "opex": "IS_OPERATING_EXPN",
    "sga": "IS_SG&A_EXPENSE",
    "rd": "IS_OPERATING_EXPENSES_R&D",
    "other": "OTHER_OPERATING_EXPENSES_RATIO",
    "op_income": "IS_OPER_INC",
}
REQUIRED = {"revenue", "cost", "gross_profit", "opex", "rd", "other", "op_income"}


def _is_number(value: Any) -> bool:
    return isinstance(value, (int, float)) and not isinstance(value, bool)


def _is_formula(value: Any) -> bool:
    return isinstance(value, str) and value.startswith("=")


def _close(a: float, b: float, tolerance: float = TOLERANCE) -> bool:
    return abs(a - b) <= tolerance * max(abs(a), abs(b), 1.0)


def _find_rows(ws) -> dict[str, int]:
    found: dict[str, int] = {}
    for r in range(1, min(ws.max_row or 1, 120) + 1):
        key = str(ws.cell(r, 2).value or "").strip()
        for name, expected in ROW_KEYS.items():
            if key == expected and name not in found:
                found[name] = r
    return found


class CostRecastWriter:
    def apply(self, *, workbook_path: Path, filings_dir: Path | None) -> dict[str, Any]:
        report: dict[str, Any] = {"applied": False, "label": None, "years": [], "reason": None}
        if filings_dir is None or not filings_dir.exists():
            report["reason"] = "No cached 10-K filings."
            return report
        wb = load_workbook(workbook_path, data_only=False)
        try:
            if SHEET not in wb.sheetnames:
                report["reason"] = f"No '{SHEET}' sheet."
                return report
            ws = wb[SHEET]
            rows = _find_rows(ws)
            if not REQUIRED <= rows.keys():
                report["reason"] = "Income statement rows not found."
                return report
            columns = {
                int("".join(ch for ch in token if ch.isdigit())): col
                for token, col in detect_year_columns(ws, wb).items()
            }
            filings = load_filings(filings_dir)
            workbook_cost: dict[int, float | None] = {}
            for fy, col in columns.items():
                raw = ws.cell(rows["cost"], col).value
                workbook_cost[fy] = float(raw) if _is_number(raw) else None
            result = recast_cost(workbook_cost, filings)
            report["label"] = result.get("label")
            if not result.get("label"):
                report["reason"] = "No filing statement row matches the workbook's cost of revenue."
                return report
            for fy, info in sorted(result["years"].items()):
                if info["status"] == "recast":
                    report["years"].append(
                        self._recast_year(ws, rows, columns[fy], fy, info, filings, result["label"])
                    )
            if any(y["status"] == "written" for y in report["years"]):
                wb.save(workbook_path)
                report["applied"] = True
            return report
        finally:
            wb.close()

    def _recast_year(self, ws, rows, col, fy, info, filings, label) -> dict[str, Any]:
        entry: dict[str, Any] = {"fiscal_year": f"FY{fy}", "status": "skipped", "reason": None, "cells": []}

        def cell(name: str):
            return ws.cell(rows[name], col)

        def value(name: str):
            if name not in rows:
                return None
            raw = cell(name).value
            return float(raw) if _is_number(raw) else None

        filing_year = info["source_filing_year"]
        offset = filing_year - fy
        statement = statement_rows_for_year(filings, filing_year)
        total_row = statement_total_for_year(filings, filing_year)
        if not statement or len(total_row) <= offset or total_row[offset] is None:
            entry["reason"] = "Could not read the filing's expense lines."
            return entry
        divisor = 1000.0  # filings report in thousands; the workbook is in millions
        total = total_row[offset] / divisor
        cost_new = info["value"]
        rd_label = next((k for k in statement if "development" in k), None)
        if rd_label is None or statement[rd_label][offset] is None:
            entry["reason"] = "No R&D / development line to receive the reclassified amount."
            return entry
        rd_new = statement[rd_label][offset] / divisor

        for name in rows:
            if _is_formula(cell(name).value):
                entry["reason"] = f"Row '{name}' contains a formula; not overwritten."
                return entry
        revenue, op_income, old_opex = value("revenue"), value("op_income"), value("opex")
        old_cost = value("cost") or 0.0
        if revenue is None or op_income is None or old_opex is None:
            entry["reason"] = "Revenue, operating income or operating expenses missing."
            return entry
        if not _close(old_cost + old_opex, total) or not _close(revenue - total, op_income):
            entry["reason"] = "Workbook totals do not reconcile with the filing; left unchanged."
            return entry
        sga = value("sga") or 0.0
        other_new = total - cost_new - sga - rd_new
        if other_new < -TOLERANCE * max(total, 1.0):
            entry["reason"] = "Reclassified lines exceed total operating expenses."
            return entry

        new_values = {
            "cost": cost_new,
            "gross_profit": revenue - cost_new,
            "opex": total - cost_new,
            "rd": rd_new,
            "other": round(max(other_new, 0.0), 3),
        }
        if "cost_detail" in rows:
            new_values["cost_detail"] = cost_new
        source = f"{filing_year} Form 10-K income statement, line '{label.title()}'"
        issue = (
            "The data provider's cost-of-revenue line differs from the company's latest presentation; amounts moved "
            "between cost and operating expenses so that total costs and operating income are unchanged."
        )
        for name, new in new_values.items():
            target = cell(name)
            old = target.value
            if _is_number(old) and _close(float(old), float(new), 1e-6):
                continue
            flag_recast(ws, target.coordinate, old_value=old, new_value=round(new, 3), source=source, issue=issue)
            target.value = round(new, 3)
            entry["cells"].append({"cell": f"{SHEET}!{target.coordinate}", "row": name, "old": old, "new": round(new, 3)})
        entry["status"] = "written"
        entry["source"] = source
        return entry
