"""Primary quarterly deliverables: FA.xlsx + Quarterly Update.docx."""

from __future__ import annotations

import shutil
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from models.quarterly_update import (
    QuarterlyDeliverablesReport,
    QuarterlyProjectionReport,
    QuarterlyResearchReport,
    QuarterlyReviewReport,
)


def deliverable_stems(fiscal_year: int, fiscal_quarter: int, ticker: str) -> tuple[str, str]:
    t = ticker.upper()
    excel = f"{fiscal_year} Q{fiscal_quarter} {t} FA.xlsx"
    word = f"{fiscal_year} Q{fiscal_quarter} {t} Quarterly Update.docx"
    return excel, word


def _pct(a: float | None, b: float | None) -> str:
    if a is None or b is None or b == 0:
        return "n/a"
    return f"{(a / b - 1.0) * 100:+.1f}%"


def _fmt_num(v: float | None, *, money: bool = False, pct: bool = False) -> str:
    if v is None:
        return "n/a"
    if pct:
        return f"{v * 100:.2f}%"
    if money:
        if abs(v) >= 1000:
            return f"${v:,.1f}M"
        return f"${v:,.2f}"
    return f"{v:,.2f}"


def _find_comp(review: QuarterlyReviewReport | None, statement: str, metric: str, ctype: str):
    if review is None:
        return None
    for c in review.comparisons:
        if c.statement == statement and c.metric.lower() == metric.lower() and c.comparison_type == ctype:
            return c
    return None


class QuarterlyDeliverablesService:
    """Copy completed workbook to FA name and generate Word quarterly update."""

    def produce(
        self,
        *,
        analysis_id: str,
        ticker: str,
        completed_workbook_path: Path,
        output_dir: Path,
        fiscal_year: int | None,
        fiscal_quarter: int | None,
        projection: QuarterlyProjectionReport | None = None,
        review: QuarterlyReviewReport | None = None,
        research: QuarterlyResearchReport | None = None,
    ) -> QuarterlyDeliverablesReport:
        fy = fiscal_year or (projection.fiscal_year if projection else None) or 0
        q = fiscal_quarter or (projection.fiscal_quarter if projection else None) or 0
        if not fy or not q:
            fy, q = self._infer_fy_q(completed_workbook_path)
        excel_name, word_name = deliverable_stems(fy, q, ticker)
        output_dir.mkdir(parents=True, exist_ok=True)
        excel_path = output_dir / excel_name
        word_path = output_dir / word_name

        shutil.copy2(completed_workbook_path, excel_path)
        self._write_word(
            word_path,
            ticker=ticker,
            fiscal_year=fy,
            fiscal_quarter=q,
            workbook_path=excel_path,
            projection=projection,
            review=review,
            research=research,
        )
        return QuarterlyDeliverablesReport(
            analysis_id=analysis_id,
            ticker=ticker,
            fiscal_year=fy,
            fiscal_quarter=q,
            excel_filename=excel_name,
            excel_path=str(excel_path),
            word_filename=word_name,
            word_path=str(word_path),
            summary=f"Deliverables: {excel_name}; {word_name}",
        )

    def _infer_fy_q(self, path: Path) -> tuple[int, int]:
        wb = load_workbook(path, data_only=False)
        try:
            from services.quarterly_review_service import _detect_fiscal_quarter

            q = 3
            fy = 2026
            if "Last Quarter IS Standardized" in wb.sheetnames:
                ws = wb["Last Quarter IS Standardized"]
                q = _detect_fiscal_quarter(ws) or 3
                for r in range(1, 10):
                    for c in range(1, 8):
                        v = ws.cell(r, c).value
                        if isinstance(v, str):
                            for tok in v.replace("-", " ").split():
                                if tok.isdigit() and len(tok) == 4:
                                    fy = int(tok)
            return fy, q
        finally:
            wb.close()

    def _write_word(
        self,
        path: Path,
        *,
        ticker: str,
        fiscal_year: int,
        fiscal_quarter: int,
        workbook_path: Path,
        projection: QuarterlyProjectionReport | None,
        review: QuarterlyReviewReport | None,
        research: QuarterlyResearchReport | None = None,
    ) -> None:
        try:
            from docx import Document
            from docx.shared import Pt
        except ImportError as exc:  # pragma: no cover
            raise RuntimeError("python-docx is required for Word deliverables") from exc

        metrics = self._extract_word_metrics(workbook_path)
        doc = Document()
        style = doc.styles["Normal"]
        style.font.name = "Calibri"
        style.font.size = Pt(11)

        title = doc.add_heading(f"{ticker.upper()} — {fiscal_year} Q{fiscal_quarter} Quarterly Update", level=0)
        _ = title

        # Q2/Q3 projection section BEFORE Financial Highlights
        if fiscal_quarter in (2, 3) and projection and projection.status != "NOT_APPLICABLE":
            next_fy = projection.next_fiscal_year or (fiscal_year + 1)
            doc.add_heading(f"{next_fy} Projection", level=1)
            roic_w = projection.projected_roic_wacc
            prior_roic = projection.prior_fy_roic
            prior_wacc = projection.prior_fy_wacc
            prior_spread = None
            if prior_roic is not None and prior_wacc is not None:
                prior_spread = prior_roic - prior_wacc
            p = doc.add_paragraph()
            p.add_run(
                f"Projected ROIC - WACC at {_fmt_num(roic_w, pct=True)}"
                f" ({_fmt_num(prior_spread, pct=True)} in {fiscal_year}FY)"
                if prior_spread is not None
                else f"Projected ROIC - WACC at {_fmt_num(roic_w, pct=True)}"
            )
            p2 = doc.add_paragraph()
            p2.add_run(
                f"Projected ROCE at {_fmt_num(projection.projected_roce, pct=True)}"
                + (
                    f" ({_fmt_num(projection.prior_fy_roce, pct=True)} in {fiscal_year}FY)"
                    if projection.prior_fy_roce is not None
                    else ""
                )
            )
            doc.add_paragraph(
                "Projected values are derived from the HAP ROIC house methodology "
                "(latest-quarter OA/OL, YTD annualization, prior-FY lease/R&D/OpTax bridge)."
            )

        doc.add_heading("Financial Highlights", level=1)
        if research and research.earnings_call_status == "EARNINGS_CALL_SOURCE_UNAVAILABLE":
            doc.add_paragraph(
                f"Note: {research.earnings_call_status} — no verified earnings-call transcript was retrieved."
            )
        if research and research.reported_facts:
            doc.add_paragraph("Reported Fact", style="Heading 2")
            for fact in research.reported_facts:
                doc.add_paragraph(fact, style="List Bullet")
        if research and research.management_explanations:
            doc.add_paragraph("Management Explanation", style="Heading 2")
            for note in research.management_explanations:
                doc.add_paragraph(note, style="List Bullet")
        if research and research.hap_interpretations:
            doc.add_paragraph("HAP Analyst Interpretation", style="Heading 2")
            for note in research.hap_interpretations:
                doc.add_paragraph(note, style="List Bullet")
        if not research or (
            not research.reported_facts
            and not research.management_explanations
            and not research.hap_interpretations
        ):
            doc.add_paragraph(
                f"This quarterly update for {ticker.upper()} covers fiscal {fiscal_year} Q{fiscal_quarter}. "
                "Figures are drawn from the completed HAP workbook. External research was limited; "
                "analyst should attach official earnings release and call transcript."
            )
        else:
            doc.add_paragraph(
                f"Quarterly narrative for {ticker.upper()} fiscal {fiscal_year} Q{fiscal_quarter} "
                "combines workbook metrics with external sources listed below."
            )

        doc.add_heading(
            f"Comparison Quarter to Quarter — {'3 Months' if fiscal_quarter == 1 else '3 Months + YTD'}",
            level=1,
        )
        for metric in ("Revenue", "Net Income", "Operating Income", "Gross Margin", "Operating Margin", "Net Margin"):
            ctype = "yoy_quarter"
            comp = _find_comp(review, "income_statement", metric, ctype)
            if comp and comp.compare_value is not None:
                money = "margin" not in metric.lower()
                line = (
                    f"{metric} {_pct(comp.compare_value, comp.baseline_value)} "
                    f"({_fmt_num(comp.compare_value, money=money, pct=not money)} vs. "
                    f"{_fmt_num(comp.baseline_value, money=money, pct=not money)})"
                )
                doc.add_paragraph(line, style="List Bullet")
        # YTD for Q2/Q3
        if fiscal_quarter in (2, 3):
            doc.add_paragraph(
                f"YTD comparison uses prior-year equivalent "
                f"({'6M' if fiscal_quarter == 2 else '9M'})."
            )
            for metric in ("Revenue", "Operating Income", "Net Income"):
                comp = _find_comp(review, "income_statement", metric, "ytd")
                if comp and comp.compare_value is not None:
                    doc.add_paragraph(
                        f"YTD {metric} {_pct(comp.compare_value, comp.baseline_value)} "
                        f"({_fmt_num(comp.compare_value, money=True)} vs. "
                        f"{_fmt_num(comp.baseline_value, money=True)})",
                        style="List Bullet",
                    )

        cfo = _find_comp(review, "cash_flow", "CFO", "ytd")
        if cfo and cfo.compare_value is not None:
            doc.add_paragraph(
                f"Cash from Ops {_fmt_num(cfo.compare_value, money=True)} vs. "
                f"{_fmt_num(cfo.baseline_value, money=True)} (YTD)",
                style="List Bullet",
            )

        doc.add_heading("Comparison Last Quarter to Quarter Right Before", level=1)
        for metric in ("Cash", "Total assets", "Equity", "Inventory", "Receivables", "Accounts payable"):
            comp = _find_comp(review, "balance_sheet", metric, "qoq")
            if comp and comp.compare_value is not None:
                doc.add_paragraph(
                    f"{metric} {_pct(comp.compare_value, comp.baseline_value)} "
                    f"({_fmt_num(comp.compare_value, money=True)} vs. "
                    f"{_fmt_num(comp.baseline_value, money=True)})",
                    style="List Bullet",
                )

        doc.add_heading("Price & MoS", level=1)
        price = metrics.get("current_price")
        max_buy = metrics.get("max_price_to_buy")
        pe10_pct = metrics.get("pe10_percentile")
        doc.add_paragraph(
            f"At the current price of {_fmt_num(price, money=True)}, "
            f"PE10 percentile at {_fmt_num(pe10_pct, pct=True)}, "
            f"the max entry price is {_fmt_num(max_buy, money=True)}."
        )
        if metrics.get("ev_mos") is not None:
            doc.add_paragraph(f"EV MoS: {_fmt_num(metrics.get('ev_mos'), pct=True)}")
        else:
            doc.add_paragraph("EV MoS: not available in this workbook run.")
        if metrics.get("expected_return") is not None:
            doc.add_paragraph(
                f"Expected return at {_fmt_num(metrics.get('expected_return'), pct=True)}"
            )
        if metrics.get("graham_entry") is not None:
            doc.add_paragraph(
                f"Graham entry target price: {_fmt_num(metrics.get('graham_entry'), money=True)}"
            )

        doc.add_heading("Sources", level=1)
        doc.add_paragraph(
            "Primary: completed HAP Excel model (Bloomberg LQ statements, Yahoo basic fallback, "
            "SEC secondary gap fill, CRF current data, Yahoo live price)."
        )
        if research and research.sources:
            for src in research.sources[:10]:
                line = f"[{src.source_kind}] {src.title}"
                if src.url:
                    line += f" — {src.url}"
                doc.add_paragraph(line, style="List Bullet")
        else:
            doc.add_paragraph(
                f"Supporting: company IR earnings release, SEC 10-Q/8-K, earnings-call materials "
                f"for {fiscal_year} Q{fiscal_quarter}."
            )
        doc.save(path)

    def _extract_word_metrics(self, workbook_path: Path) -> dict[str, Any]:
        wb = load_workbook(workbook_path, data_only=False)
        out: dict[str, Any] = {}
        try:
            if "Inputs" in wb.sheetnames:
                inp = wb["Inputs"]
                out["current_price"] = _num_cell(inp["B63"].value)
                out["max_price_to_buy"] = _num_cell(inp["B67"].value)
                out["pe10_percentile"] = _num_cell(inp["B72"].value)
                out["expected_return"] = _num_cell(inp["B69"].value)
            if "Final Metrics" in wb.sheetnames:
                fm = wb["Final Metrics"]
                # Best-effort: B51 often mirrors current price; leave MoS if present nearby
                out.setdefault("current_price", _num_cell(fm["B51"].value if "B51" in fm else None))
            if "Enterprise Value" in wb.sheetnames:
                ev = wb["Enterprise Value"]
                # MoS often on EV sheet — scan for MOS label
                for r in range(1, 40):
                    lab = str(ev.cell(r, 1).value or "").lower()
                    if "mos" in lab or "margin of safety" in lab:
                        out["ev_mos"] = _num_cell(ev.cell(r, 2).value)
                        break
            # Graham entry — valuation sheets vary; leave None if absent
        finally:
            wb.close()
        return out


def _num_cell(v: Any) -> float | None:
    if v is None or isinstance(v, bool):
        return None
    if isinstance(v, (int, float)):
        return float(v)
    try:
        return float(v)
    except (TypeError, ValueError):
        return None
