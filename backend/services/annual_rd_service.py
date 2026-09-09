"""R&D annual update: carry useful life; ensure expense + schedule cover the new FY."""

from __future__ import annotations

import re
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

from models.annual_update import AnnualRdReport
from services.annual_period_service import detect_workbook_years, detect_year_columns

_LIFE_CELLS = ("B8", "C8", "B2")
_COL_LETTER_RE = re.compile(r"([A-Z]+)")


def _num(v: Any) -> float | None:
    if isinstance(v, (int, float)) and not isinstance(v, bool):
        return float(v)
    return None


def _fy_token(fy: str) -> str:
    return fy if str(fy).startswith("FY") else f"FY{fy}"


def _shift_formula(formula: str, from_col: int, to_col: int) -> str:
    """Shift simple same-row column references by (to_col - from_col)."""
    delta = to_col - from_col
    if delta == 0 or not isinstance(formula, str) or not formula.startswith("="):
        return formula

    def repl(match: re.Match[str]) -> str:
        letters = match.group(1)
        # Avoid shifting sheet names inside quotes by only touching bare refs later;
        # for Industrial Template R&D formulas are local A1 refs.
        from openpyxl.utils import column_index_from_string

        idx = column_index_from_string(letters)
        return get_column_letter(idx + delta)

    # Only shift column letters not inside sheet quotes.
    out = []
    i = 0
    text = formula
    while i < len(text):
        if text[i] == "'":
            j = text.find("'", i + 1)
            if j < 0:
                out.append(text[i:])
                break
            out.append(text[i : j + 1])
            i = j + 1
            continue
        m = re.match(r"[A-Z]+", text[i:])
        if m and (i + m.end() >= len(text) or text[i + m.end()].isdigit() or text[i + m.end()] in "+-*/(),"):
            # column letter possibly followed by row digits
            letters = m.group(0)
            rest_i = i + len(letters)
            if rest_i < len(text) and text[rest_i].isdigit():
                from openpyxl.utils import column_index_from_string

                idx = column_index_from_string(letters)
                out.append(get_column_letter(idx + delta))
                i = rest_i
                continue
        out.append(text[i])
        i += 1
    return "".join(out)


class AnnualRdService:
    def apply(
        self,
        *,
        analysis_id: str,
        ticker: str,
        workbook_path: Path,
        previous_workbook_path: Path,
        fiscal_year: str,
        filing_rd: float | None = None,
    ) -> AnnualRdReport:
        prev = load_workbook(previous_workbook_path, data_only=False)
        wb = load_workbook(workbook_path, data_only=False)
        life = None
        life_unchanged = True
        formulas_ok = True
        used_bb = True
        filled = False
        expense = None
        schedule_extended = False
        lookback: list[str] = []
        lookback_complete = True
        written: list[str] = []
        try:
            # Prefer Inputs R&D expense (formula-linked to Income) as Bloomberg/statement source.
            inputs_expense = self._inputs_rd_expense(wb, fiscal_year)
            if filing_rd is None and inputs_expense is not None:
                filing_rd = inputs_expense
                used_bb = True

            if "R&D" in prev.sheetnames and "R&D" in wb.sheetnames:
                pws, ows = prev["R&D"], wb["R&D"]
                for addr in _LIFE_CELLS:
                    pv = pws[addr].value
                    if _num(pv) is not None:
                        life = _num(pv)
                        ov = ows[addr].value
                        if ov != pv and not (isinstance(ov, str) and ov.startswith("=")):
                            ows[addr].value = pv
                            written.append(f"R&D!{addr}")
                        life_unchanged = True
                        break

                life_years = int(life) if life and life >= 1 else 3
                token = _fy_token(fiscal_year)
                year_num = int("".join(ch for ch in token if ch.isdigit()) or 0)
                lookback = [f"FY{y}" for y in range(year_num - life_years + 1, year_num + 1)]

                # Ensure Inputs historical R&D present for lookback when formula-backed.
                lookback_complete = self._verify_inputs_lookback(wb, lookback)

                # Extend R&D schedule formulas through the new FY column if missing.
                schedule_extended = self._ensure_schedule_formulas(
                    ows, wb, token, life_years=life_years, written=written
                )

                # Expense on R&D sheet is usually formula→Inputs; only fill blank constants.
                cols = detect_year_columns(ows, wb)
                if not any(k.startswith("FY") for k in cols):
                    # R&D headers are formula-driven; map via Inputs column alignment.
                    inp_cols = detect_year_columns(wb["Inputs"], wb) if "Inputs" in wb.sheetnames else {}
                    if not any(k.startswith("FY") for k in inp_cols):
                        inp_cols = detect_workbook_years(workbook_path)
                    # R&D expense row uses cols E..=Inputs!C.. so offset +2
                    cols = {fy: c + 2 for fy, c in inp_cols.items() if str(fy).startswith("FY")}

                col = cols.get(token)
                expense_row = 2
                for row in range(1, min(ows.max_row or 1, 40) + 1):
                    lab = str(ows.cell(row, 1).value or "").lower()
                    if "r&d" in lab and "expense" in lab:
                        expense_row = row
                        break
                if col:
                    cell = ows.cell(expense_row, col)
                    expense = _num(cell.value)
                    if expense is None and isinstance(cell.value, str) and cell.value.startswith("="):
                        expense = filing_rd  # formula present; value known from Inputs/Income
                        formulas_ok = True
                    elif expense is None and filing_rd is not None:
                        if isinstance(cell.value, str) and cell.value.startswith("="):
                            formulas_ok = True
                        else:
                            cell.value = filing_rd
                            expense = filing_rd
                            used_bb = False
                            filled = True
                            written.append(f"R&D!{get_column_letter(col)}{expense_row}")
                    elif expense is not None:
                        used_bb = True

                # Ensure Inputs R&D expense cell exists for new FY (do not overwrite formulas).
                if "Inputs" in wb.sheetnames and filing_rd is not None:
                    iws = wb["Inputs"]
                    icols = detect_year_columns(iws, wb)
                    if not any(k.startswith("FY") for k in icols):
                        icols = detect_workbook_years(workbook_path)
                    icol = icols.get(token)
                    if icol:
                        for row in range(100, 110):
                            lab = str(iws.cell(row, 1).value or "").lower()
                            if lab == "r&d expense" or (lab.startswith("r&d") and "expense" in lab):
                                cell = iws.cell(row, icol)
                                if cell.value in (None, ""):
                                    cell.value = float(filing_rd)
                                    filled = True
                                    used_bb = False
                                    written.append(f"Inputs!{get_column_letter(icol)}{row}")
                                elif isinstance(cell.value, str) and cell.value.startswith("="):
                                    expense = expense or filing_rd
                                break

            wb.save(workbook_path)
        finally:
            prev.close()
            wb.close()

        status_ok = life is not None and lookback_complete and (expense is not None or schedule_extended)
        return AnnualRdReport(
            analysis_id=analysis_id,
            ticker=ticker,
            useful_life_unchanged=life_unchanged,
            useful_life=life,
            bloomberg_rd_used=used_bb,
            filled_from_filing=filled,
            rd_expense=expense if expense is not None else filing_rd,
            formulas_preserved=formulas_ok,
            schedule_extended=schedule_extended,
            lookback_years=lookback,
            lookback_complete=lookback_complete,
            cells_written=written,
            summary=(
                f"R&D: useful life={life} (unchanged={life_unchanged}); "
                f"expense={expense if expense is not None else filing_rd}; "
                f"bloomberg={used_bb}; filing_fill={filled}; "
                f"schedule_extended={schedule_extended}; lookback_ok={lookback_complete}; "
                f"status={'ok' if status_ok else 'RD_LOOKBACK_OR_SCHEDULE_REVIEW'}."
            ),
        )

    def refresh_calculated_values(
        self, report: AnnualRdReport, workbook_path: Path, fiscal_year: str
    ) -> AnnualRdReport:
        """Read FY R&D expense/asset/amortization from Excel-cached values after recalc."""
        from openpyxl import load_workbook

        wb = load_workbook(workbook_path, data_only=True)
        wb_f = load_workbook(workbook_path, data_only=False)
        try:
            if "R&D" not in wb.sheetnames or "Inputs" not in wb_f.sheetnames:
                return report
            inp_cols = detect_year_columns(wb_f["Inputs"], wb_f)
            token = _fy_token(fiscal_year)
            inp_col = inp_cols.get(token)
            if not inp_col:
                return report
            rd_col = inp_col + 2  # R&D expense band offset
            expense = _num(wb["R&D"].cell(2, rd_col).value)
            asset = _num(wb["R&D"].cell(3, rd_col).value)
            amort = _num(wb["R&D"].cell(4, rd_col).value)
            return report.model_copy(
                update={
                    "rd_expense": expense if expense is not None else report.rd_expense,
                    "rd_asset": asset,
                    "rd_amortization": amort,
                    "summary": (
                        f"R&D: useful life={report.useful_life}; expense={expense}; "
                        f"asset={asset}; amortization={amort}; "
                        f"schedule_extended={report.schedule_extended}; lookback_ok={report.lookback_complete}."
                    ),
                }
            )
        finally:
            wb.close()
            wb_f.close()

    @staticmethod
    def _inputs_rd_expense(wb, fiscal_year: str) -> float | None:
        if "Income - GAAP" in wb.sheetnames:
            ws = wb["Income - GAAP"]
            cols = detect_year_columns(ws, wb)
            token = _fy_token(fiscal_year)
            col = cols.get(token)
            if col:
                for row in range(1, min(ws.max_row or 1, 80) + 1):
                    lab = str(ws.cell(row, 1).value or "").lower()
                    if "research" in lab and "development" in lab:
                        return _num(ws.cell(row, col).value)
        if "Inputs" not in wb.sheetnames:
            return None
        ws = wb["Inputs"]
        cols = detect_year_columns(ws, wb)
        token = _fy_token(fiscal_year)
        col = cols.get(token)
        if not col:
            return None
        for row in range(100, 110):
            lab = str(ws.cell(row, 1).value or "").lower()
            if "r&d expense" in lab:
                return _num(ws.cell(row, col).value)
        return None

    @staticmethod
    def _verify_inputs_lookback(wb, lookback: list[str]) -> bool:
        if "Inputs" not in wb.sheetnames or not lookback:
            return True
        ws = wb["Inputs"]
        cols = detect_year_columns(ws, wb)
        expense_row = None
        for row in range(100, 110):
            lab = str(ws.cell(row, 1).value or "").lower()
            if "r&d expense" in lab:
                expense_row = row
                break
        if expense_row is None:
            return True
        for fy in lookback:
            col = cols.get(fy)
            if not col:
                return False
            val = ws.cell(expense_row, col).value
            if val in (None, ""):
                return False
        return True

    def _ensure_schedule_formulas(
        self,
        ows,
        wb,
        token: str,
        *,
        life_years: int,
        written: list[str],
    ) -> bool:
        """Copy prior-column formula pattern into the new FY column when blank."""
        inp_cols = {}
        if "Inputs" in wb.sheetnames:
            inp_cols = detect_year_columns(wb["Inputs"], wb)
        if not any(k.startswith("FY") for k in inp_cols):
            return False
        # Map FY → R&D sheet column (expense band starts at E for Inputs!C).
        rd_cols = {fy: c + 2 for fy, c in inp_cols.items() if str(fy).startswith("FY")}
        new_col = rd_cols.get(token)
        if new_col is None:
            return False
        # Prior FY column
        prior_token = f"FY{int(token[2:]) - 1}"
        prior_col = rd_cols.get(prior_token)
        if prior_col is None:
            return False

        extended = False
        for row in (2, 3, 4):  # expense, asset, amortization
            cell = ows.cell(row, new_col)
            if cell.value not in (None, ""):
                if isinstance(cell.value, str) and cell.value.startswith("="):
                    extended = True
                continue
            prior = ows.cell(row, prior_col).value
            if isinstance(prior, str) and prior.startswith("="):
                cell.value = _shift_formula(prior, prior_col, new_col)
                written.append(f"R&D!{get_column_letter(new_col)}{row}")
                extended = True
            elif row == 2:
                # Expense sometimes constant; prefer Inputs formula pattern.
                inp_col = inp_cols.get(token)
                if inp_col:
                    letter = get_column_letter(inp_col)
                    cell.value = f'=IF(Inputs!{letter}103="",0,Inputs!{letter}103)'
                    written.append(f"R&D!{get_column_letter(new_col)}{row}")
                    extended = True
        return extended
