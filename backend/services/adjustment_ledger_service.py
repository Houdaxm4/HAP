"""HAP Adjustments tab: one auditable ledger of every value HAP changed on purpose.

Replaces ad-hoc colors and boxes with:
- a table (id, year, tab, linked cell, category, original, new, change, reason, source, method, confidence, status),
- a hyperlink from the ledger to the cell and a short, structured marker on the cell (ADJ-id, before -> after, one reason line),
- the original value or formula kept in full, so any adjustment can be undone by hand.

Used for lease-rate fixes and ROIC adjustments (operating assets, operating liabilities, one-time operating income items).
"""

from __future__ import annotations

from dataclasses import dataclass
from typing import Any

from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.worksheet import Worksheet

LEDGER_SHEET = "HAP Adjustments"
HEADERS = (
    "ID", "Year", "Tab", "Cell", "Category", "What changed", "Original (value or formula)", "New (value or formula)",
    "Amount", "Reason", "Source", "Method", "Confidence", "Status",
)
WIDTHS = (9, 8, 22, 10, 22, 34, 34, 34, 12, 60, 44, 20, 11, 14)
CELL_FILL = PatternFill("solid", fgColor="FFF3D6")
HEADER_FILL = PatternFill("solid", fgColor="1F3A5F")
CATEGORY_FILL = {
    "Lease rate": "DCEBFA",
    "Operating assets": "E3F4E1",
    "Operating liabilities": "FDE9D9",
    "One-time operating income": "EADCF4",
    "SEC override of Bloomberg": "FFF1C2",
}
THIN = Side(style="thin", color="C9CED6")
MARK = "◆"  # diamond: the marker shown in the comment title


@dataclass
class Adjustment:
    adj_id: str
    sheet: str
    cell: str
    row: int


class AdjustmentLedger:
    def __init__(self, wb) -> None:
        self.wb = wb
        self.ws = self._ensure_sheet()

    def _ensure_sheet(self) -> Worksheet:
        if LEDGER_SHEET in self.wb.sheetnames:
            return self.wb[LEDGER_SHEET]
        ws = self.wb.create_sheet(LEDGER_SHEET)
        ws["A1"] = "HAP Adjustments"
        ws["A1"].font = Font(bold=True, size=14, color="1F3A5F")
        ws["A2"] = (
            "Every number HAP changed on purpose. The Cell column links to the changed cell; the original value or formula is kept "
            "in full so any change can be undone by hand. Source is the filing or table the figure came from."
        )
        ws["A2"].alignment = Alignment(wrap_text=True, vertical="top")
        ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=len(HEADERS))
        ws.row_dimensions[2].height = 32
        for col, (head, width) in enumerate(zip(HEADERS, WIDTHS), start=1):
            cell = ws.cell(4, col, head)
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = HEADER_FILL
            cell.alignment = Alignment(vertical="center", wrap_text=True)
            ws.column_dimensions[cell.column_letter].width = width
        ws.freeze_panes = "A5"
        ws.sheet_properties.tabColor = "C9A227"
        return ws

    def _next_row(self) -> int:
        row = 5
        while self.ws.cell(row, 1).value not in (None, ""):
            row += 1
        return row

    def record(
        self,
        *,
        sheet: str,
        cell: str,
        fiscal_year: str | None,
        category: str,
        what: str,
        original: Any,
        new: Any,
        amount: float | None,
        reason: str,
        source: str,
        method: str,
        confidence: str = "medium",
        mark_cell: bool = True,
    ) -> Adjustment:
        row = self._next_row()
        adj_id = f"ADJ-{row - 4:03d}"
        values = (adj_id, fiscal_year or "", sheet, cell, category, what, str(original), str(new), amount, reason, source, method, confidence, "Applied")
        fill = PatternFill("solid", fgColor=CATEGORY_FILL.get(category, "FFFFFF"))
        for col, value in enumerate(values, start=1):
            c = self.ws.cell(row, col, value)
            c.alignment = Alignment(wrap_text=True, vertical="top")
            c.border = Border(top=THIN, bottom=THIN, left=THIN, right=THIN)
            if col in (1, 5):
                c.fill = fill
        for col in (7, 8):
            c = self.ws.cell(row, col)
            if str(c.value).startswith("="):
                c.data_type = "s"  # show the formula as text; never calculate it
        self.ws.cell(row, 9).number_format = "0.00%" if category == "Lease rate" else "#,##0.00"
        lines = max(len(reason) / 56, len(source) / 40, len(what) / 30, len(str(original)) / 30, len(str(new)) / 30, 1)
        self.ws.row_dimensions[row].height = max(30, 15 * (int(lines) + 1))
        link = self.ws.cell(row, 4)
        link.hyperlink = f"#'{sheet}'!{cell}"
        link.font = Font(color="0563C1", underline="single")
        self.ws.auto_filter.ref = f"A4:{self.ws.cell(4, len(HEADERS)).column_letter}{row}"
        if mark_cell and sheet in self.wb.sheetnames:
            self._mark(self.wb[sheet][cell], adj_id, category, original, new, reason, source, row)
        return Adjustment(adj_id, sheet, cell, row)

    @staticmethod
    def _short(value: Any, limit: int = 70) -> str:
        text = str(value)
        return text if len(text) <= limit else text[: limit - 1] + "…"

    def _mark(self, cell, adj_id: str, category: str, original: Any, new: Any, reason: str, source: str, ledger_row: int) -> None:
        cell.fill = CELL_FILL
        body = (
            f"{MARK} {adj_id} · {category}\n"
            f"Before: {self._short(original)}\n"
            f"After:  {self._short(new)}\n"
            f"Why: {self._short(reason, 160)}\n"
            f"Source: {self._short(source, 100)}\n"
            f"Full detail: '{LEDGER_SHEET}' row {ledger_row}"
        )
        existing = cell.comment.text.strip() if cell.comment and cell.comment.text else ""
        if adj_id in existing:
            return
        comment = Comment((existing + "\n\n" + body).strip() if existing else body, "HAP")
        comment.width, comment.height = 380, 170
        cell.comment = comment

    def summary_rows(self) -> dict[str, int]:
        counts: dict[str, int] = {}
        row = 5
        while self.ws.cell(row, 1).value not in (None, ""):
            cat = str(self.ws.cell(row, 5).value)
            counts[cat] = counts.get(cat, 0) + 1
            row += 1
        return counts
