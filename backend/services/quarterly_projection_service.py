"""Q2/Q3 projected ROIC-WACC and ROCE using house IC/NOPAT methodology."""

from __future__ import annotations

from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from models.quarterly_presentation import STATEMENT_SHEETS, QuarterlyStatementKind
from models.quarterly_update import QuarterlyProjectionReport
from services.quarterly_review_service import _detect_fiscal_quarter

IC_SHEET = "IC & NOPAT & ROIC "
LQ_BS = STATEMENT_SHEETS[QuarterlyStatementKind.BALANCE_SHEET]
LQ_IS = STATEMENT_SHEETS[QuarterlyStatementKind.INCOME]
LQ_CF = STATEMENT_SHEETS[QuarterlyStatementKind.CASH_FLOW]
LAST_FY_COL = 12  # L
PROJ_COL = 13  # M
ROCE_COL = 14  # N

# Same annual BS row classifications as Inputs!81/82/84/85 (house ROIC OA/OL)
_OA_CURRENT_ROWS = (11, 14, 20, 27)
_OA_NONCURRENT_ROWS = (37, 48, 51)
_OL_CURRENT_ROWS = (65, 68, 78)
_OL_NONCURRENT_ROWS = (98,)


def annualization_factor(fiscal_quarter: int) -> float | None:
    if fiscal_quarter == 2:
        return 2.0
    if fiscal_quarter == 3:
        return 4.0 / 3.0
    return None


def _num(v: Any) -> float | None:
    if v is None or isinstance(v, bool):
        return None
    if isinstance(v, (int, float)):
        return float(v)
    if isinstance(v, str) and v.startswith("="):
        return None
    try:
        return float(v)
    except (TypeError, ValueError):
        return None


def _sum_rows(ws, rows: tuple[int, ...], col: int = 3) -> float:
    total = 0.0
    for r in rows:
        v = _num(ws.cell(r, col).value)
        if v is not None:
            total += v
    return total


def _find_is_row(ws, needle: str) -> int | None:
    n = needle.lower()
    for r in range(11, 76):
        lab = ws.cell(r, 1).value
        if lab and n in str(lab).lower() and not str(lab).startswith(" "):
            return r
    for r in range(11, 76):
        lab = ws.cell(r, 1).value
        if lab and n in str(lab).lower():
            return r
    return None


def _find_cf_cfo_row(ws) -> int | None:
    # Prefer populated total "Cash from Operating Activities" (often row 27)
    hits: list[tuple[int, bool]] = []
    for r in range(11, 80):
        lab = ws.cell(r, 1).value
        if not lab:
            continue
        n = str(lab).lower()
        if "cash from operating" in n or "operating activities" in n:
            if str(lab).startswith(" "):
                continue
            hits.append((r, _num(ws.cell(r, 3).value) is not None))
    if not hits:
        return None
    hits.sort(key=lambda t: (not t[1], t[0]))
    return hits[0][0]


class QuarterlyProjectionService:
    """Write M-column ROIC projection and N-column ROCE for Q2/Q3."""

    def apply(
        self,
        *,
        analysis_id: str,
        ticker: str,
        workbook_path: Path,
        fiscal_quarter: int | None = None,
        previous_workbook_path: Path | None = None,
    ) -> QuarterlyProjectionReport:
        wb = load_workbook(workbook_path, data_only=False)
        prev_wb = None
        if previous_workbook_path and Path(previous_workbook_path).exists():
            try:
                prev_wb = load_workbook(previous_workbook_path, data_only=True)
            except Exception:  # noqa: BLE001
                prev_wb = None
        try:
            q = fiscal_quarter
            if q is None and LQ_IS in wb.sheetnames:
                q = _detect_fiscal_quarter(wb[LQ_IS])
            q = q or 1

            fy = self._detect_fiscal_year(wb, q)
            factor = annualization_factor(q)
            if factor is None:
                return QuarterlyProjectionReport(
                    analysis_id=analysis_id,
                    ticker=ticker,
                    status="NOT_APPLICABLE",
                    fiscal_quarter=q,
                    fiscal_year=fy,
                    next_fiscal_year=(fy + 1) if fy else None,
                    summary=f"Q{q}: projection NOT_APPLICABLE (Q2/Q3 only).",
                )

            if IC_SHEET not in wb.sheetnames:
                return QuarterlyProjectionReport(
                    analysis_id=analysis_id,
                    ticker=ticker,
                    status="BLOCKED",
                    fiscal_quarter=q,
                    fiscal_year=fy,
                    summary="IC & NOPAT & ROIC sheet missing.",
                )

            ic = wb[IC_SHEET]
            bs = wb[LQ_BS] if LQ_BS in wb.sheetnames else None
            is_ws = wb[LQ_IS] if LQ_IS in wb.sheetnames else None
            cf = wb[LQ_CF] if LQ_CF in wb.sheetnames else None
            fm = wb["Final Metrics"] if "Final Metrics" in wb.sheetnames else None

            # --- OA / OL from latest-quarter BS using house row classification ---
            oa = ol = total_assets = None
            if bs is not None:
                oa = _sum_rows(bs, _OA_CURRENT_ROWS) + _sum_rows(bs, _OA_NONCURRENT_ROWS)
                ol = _sum_rows(bs, _OL_CURRENT_ROWS) + _sum_rows(bs, _OL_NONCURRENT_ROWS)
                total_assets = _num(bs.cell(61, 3).value)

            # --- YTD Revenue / OI / CFO ---
            ytd_rev = ytd_oi = ytd_cfo = None
            if is_ws is not None:
                rev_r = _find_is_row(is_ws, "revenue")
                oi_r = _find_is_row(is_ws, "operating income")
                if rev_r:
                    ytd_rev = _num(is_ws.cell(rev_r, 7).value)  # col G = YTD
                if oi_r:
                    ytd_oi = _num(is_ws.cell(oi_r, 7).value)
            if cf is not None:
                cfo_r = _find_cf_cfo_row(cf)
                if cfo_r:
                    ytd_cfo = _num(cf.cell(cfo_r, 3).value)

            proj_rev = (ytd_rev * factor) if ytd_rev is not None else None
            proj_oi = (ytd_oi * factor) if ytd_oi is not None else None
            proj_cfo = (ytd_cfo * factor) if ytd_cfo is not None else None

            # --- Last FY lease/R&D/OpTax/OI from IC column L (and source sheets) ---
            prior_oi = self._last_fy_component(wb, ic, 13, prev_wb)
            prior_tax = self._last_fy_component(wb, ic, 19, prev_wb)

            lease_exp = self._last_fy_component(wb, ic, 14, prev_wb)
            lease_dep = self._last_fy_component(wb, ic, 15, prev_wb)
            rd_exp = self._last_fy_component(wb, ic, 16, prev_wb)
            rd_dep = self._last_fy_component(wb, ic, 17, prev_wb)
            cap_lease = self._last_fy_component(wb, ic, 5, prev_wb)
            cap_rd = self._last_fy_component(wb, ic, 6, prev_wb)

            # Projected taxes
            status = "ok"
            proj_tax = None
            if prior_oi is None or prior_oi == 0:
                status = "REVIEW_REQUIRED"
            elif proj_oi is not None and prior_tax is not None:
                proj_tax = prior_tax * (proj_oi / prior_oi)

            # IC = OA - OL + CapLeases + CapRD (same as L7)
            proj_ic = None
            if oa is not None and ol is not None:
                proj_ic = oa - ol + (cap_lease or 0.0) + (cap_rd or 0.0)

            # NOPAT = OI + LeaseExp - LeaseDep + RDExp - RDDep - OpTax
            proj_nopat = None
            if proj_oi is not None and proj_tax is not None:
                proj_nopat = (
                    proj_oi
                    + (lease_exp or 0.0)
                    - (lease_dep or 0.0)
                    + (rd_exp or 0.0)
                    - (rd_dep or 0.0)
                    - proj_tax
                )

            proj_roic = None
            if proj_nopat is not None and proj_ic not in (None, 0):
                proj_roic = proj_nopat / proj_ic

            # WACC from Final Metrics last FY (L7)
            wacc = None
            if fm is not None:
                wacc = self._resolve_value(wb, fm.cell(7, LAST_FY_COL).value, "Final Metrics")
            if wacc is None and prev_wb is not None and "Final Metrics" in prev_wb.sheetnames:
                wacc = _num(prev_wb["Final Metrics"].cell(7, LAST_FY_COL).value)

            roic_wacc = None
            if proj_roic is not None and wacc is not None:
                roic_wacc = proj_roic - wacc

            proj_roce = None
            if proj_cfo is not None and total_assets not in (None, 0):
                proj_roce = proj_cfo / total_assets

            prior_roic = self._last_fy_component(wb, ic, 23, prev_wb)
            prior_roce = None  # optional; Word may show N/A

            # --- Write projection cells (house structure from completed Q2 workbook) ---
            formulas: dict[str, str] = {}
            ic.cell(1, PROJ_COL).value = "Projected ROIC"
            ic.cell(1, ROCE_COL).value = "ROCE"

            if oa is not None:
                ic.cell(3, PROJ_COL).value = oa
            if ol is not None:
                ic.cell(4, PROJ_COL).value = ol
            if cap_lease is not None:
                ic.cell(5, PROJ_COL).value = cap_lease
            if cap_rd is not None:
                ic.cell(6, PROJ_COL).value = cap_rd
            ic.cell(7, PROJ_COL).value = "=M3-M4+M5+M6"
            formulas["M7"] = "=M3-M4+M5+M6"

            if proj_rev is not None:
                ic.cell(11, PROJ_COL).value = proj_rev
            if proj_oi is not None:
                ic.cell(13, PROJ_COL).value = proj_oi
            if lease_exp is not None:
                ic.cell(14, PROJ_COL).value = lease_exp
            if lease_dep is not None:
                ic.cell(15, PROJ_COL).value = lease_dep
            if rd_exp is not None:
                ic.cell(16, PROJ_COL).value = rd_exp
            if rd_dep is not None:
                ic.cell(17, PROJ_COL).value = rd_dep
            ic.cell(18, PROJ_COL).value = "=M13"
            formulas["M18"] = "=M13"
            ic.cell(19, PROJ_COL).value = "=IF(OR(L13=\"\",L13=0),\"\",M13/L13*L19)"
            formulas["M19"] = "=IF(OR(L13=\"\",L13=0),\"\",M13/L13*L19)"
            ic.cell(20, PROJ_COL).value = "=M13+M14-M15+M16-M17-M19"
            formulas["M20"] = "=M13+M14-M15+M16-M17-M19"
            ic.cell(23, PROJ_COL).value = "=M20/M7"
            formulas["M23"] = "=M20/M7"

            if wacc is not None:
                ic.cell(24, PROJ_COL).value = wacc  # M24 prior FY WACC
            ic.cell(25, PROJ_COL).value = "=M23-M24"
            formulas["M25"] = "=M23-M24"

            # N4 ROCE formula using LQ CFO × factor / LQ Total Assets
            cfo_row = _find_cf_cfo_row(cf) if cf is not None else 27
            factor_lit = "2" if abs(factor - 2.0) < 1e-9 else "4/3"
            roce_f = (
                f"='{LQ_CF}'!C{cfo_row}*{factor_lit}/'{LQ_BS}'!C61"
            )
            ic.cell(4, ROCE_COL).value = roce_f
            formulas["N4"] = roce_f

            wb.save(workbook_path)

            bs_date = None
            if bs is not None:
                bs_date = str(bs.cell(4, 3).value or bs.cell(5, 3).value or "")

            return QuarterlyProjectionReport(
                analysis_id=analysis_id,
                ticker=ticker,
                status=status,
                fiscal_quarter=q,
                fiscal_year=fy,
                next_fiscal_year=(fy + 1) if fy else None,
                latest_bs_date=bs_date or None,
                annualization_factor=factor,
                operating_assets=oa,
                operating_liabilities=ol,
                ytd_revenue=ytd_rev,
                ytd_operating_income=ytd_oi,
                projected_revenue=proj_rev,
                projected_operating_income=proj_oi,
                prior_fy_operating_income=prior_oi,
                prior_fy_operating_taxes=prior_tax,
                projected_operating_taxes=proj_tax,
                lease_expense=lease_exp,
                lease_depreciation=lease_dep,
                rd_expense=rd_exp,
                rd_depreciation=rd_dep,
                capitalized_leases=cap_lease,
                capitalized_rd=cap_rd,
                projected_nopat=proj_nopat,
                projected_invested_capital=proj_ic,
                projected_roic=proj_roic,
                prior_fy_wacc=wacc,
                projected_roic_wacc=roic_wacc,
                ytd_cfo=ytd_cfo,
                projected_cfo=proj_cfo,
                total_assets=total_assets,
                projected_roce=proj_roce,
                prior_fy_roic=prior_roic,
                prior_fy_roce=prior_roce,
                source_cells={
                    "oa": f"{LQ_BS}!C rows {_OA_CURRENT_ROWS}+{_OA_NONCURRENT_ROWS}",
                    "ol": f"{LQ_BS}!C rows {_OL_CURRENT_ROWS}+{_OL_NONCURRENT_ROWS}",
                    "ytd_revenue": f"{LQ_IS}!G (YTD)",
                    "ytd_oi": f"{LQ_IS}!G (YTD)",
                    "ytd_cfo": f"{LQ_CF}!C",
                    "total_assets": f"{LQ_BS}!C61",
                    "wacc": "Final Metrics!L7",
                    "prior_oi": f"{IC_SHEET}!L13",
                    "prior_tax": f"{IC_SHEET}!L19",
                },
                formulas_written=formulas,
                summary=(
                    (
                        f"Q{q} projection: factor={factor:.4g}, "
                        f"ROIC={proj_roic:.2%}"
                    )
                    if proj_roic is not None
                    else f"Q{q} projection incomplete"
                )
                + (f", ROIC-WACC={roic_wacc:.2%}" if roic_wacc is not None else "")
                + (f", ROCE={proj_roce:.2%}" if proj_roce is not None else "")
                + (f" [{status}]" if status != "ok" else ""),
            )
        finally:
            wb.close()
            if prev_wb is not None:
                prev_wb.close()

    def _last_fy_component(self, wb, ic, row: int, prev_wb=None) -> float | None:
        v = self._resolve_value(wb, ic.cell(row, LAST_FY_COL).value, IC_SHEET)
        if v is not None:
            return v
        if prev_wb is not None and IC_SHEET in prev_wb.sheetnames:
            return _num(prev_wb[IC_SHEET].cell(row, LAST_FY_COL).value)
        return None

    def _resolve_value(self, wb, value: Any, sheet_ctx: str, depth: int = 0) -> float | None:
        if depth > 25:
            return None
        n = _num(value)
        if n is not None:
            return n
        if not isinstance(value, str) or not value.startswith("="):
            return None
        return self._eval_expr(wb, value[1:].strip(), sheet_ctx, depth + 1)

    def _eval_expr(self, wb, expr: str, sheet_ctx: str, depth: int) -> float | None:
        expr = expr.strip()
        if not expr:
            return None

        # IF(cond, a, b)
        if expr.upper().startswith("IF(") and expr.endswith(")"):
            inner = expr[3:-1]
            parts = self._split_args(inner)
            if len(parts) == 3:
                cond = parts[0].strip()
                # Support Cell="" / Cell<>"" emptiness checks
                if '=""' in cond.replace(" ", "") or cond.endswith('=""'):
                    left = cond.split("=")[0].strip().rstrip("=").strip()
                    cell_raw = self._cell_raw(wb, left, sheet_ctx)
                    empty = cell_raw in (None, "")
                    branch = parts[1] if empty else parts[2]
                    return self._eval_expr(wb, branch.strip(), sheet_ctx, depth)
                if '<>""' in cond.replace(" ", ""):
                    left = cond.split("<>")[0].strip()
                    cell_raw = self._cell_raw(wb, left, sheet_ctx)
                    empty = cell_raw in (None, "")
                    branch = parts[2] if empty else parts[1]
                    return self._eval_expr(wb, branch.strip(), sheet_ctx, depth)

        # SUM(...)
        if expr.upper().startswith("SUM(") and expr.endswith(")"):
            inner = expr[4:-1]
            total = 0.0
            any_v = False
            for part in self._split_args(inner):
                part = part.strip()
                if ":" in part and "!" not in part.split(":")[0]:
                    # A1:A3 style on current sheet
                    start, end = part.split(":", 1)
                    vals = self._sum_range(wb, sheet_ctx, start.strip(), end.strip())
                    if vals is not None:
                        total += vals
                        any_v = True
                else:
                    v = self._eval_expr(wb, part, sheet_ctx, depth)
                    if v is not None:
                        total += v
                        any_v = True
            return total if any_v else None

        # Binary + - at top level (outside parens)
        for op in ("-", "+"):
            parts = self._split_top(expr, op)
            if len(parts) > 1:
                vals = [self._eval_expr(wb, p, sheet_ctx, depth) for p in parts]
                if any(v is None for v in vals):
                    # allow leading empty for unary? skip
                    if op == "-" and parts[0].strip() == "" and len(vals) == 2 and vals[1] is not None:
                        return -vals[1]
                    if all(v is not None or p.strip() == "0" for v, p in zip(vals, parts)):
                        pass
                    else:
                        # if IF returns 0 for blanks, parts should resolve
                        if any(v is None for v in vals):
                            # try treating unresolved as 0 only when IF pattern already handled
                            return None
                acc = vals[0] if vals[0] is not None else 0.0
                for v in vals[1:]:
                    if v is None:
                        return None
                    acc = acc + v if op == "+" else acc - v
                return acc

        for op in ("*", "/"):
            parts = self._split_top(expr, op)
            if len(parts) > 1:
                vals = [self._eval_expr(wb, p, sheet_ctx, depth) for p in parts]
                if any(v is None for v in vals):
                    return None
                acc = vals[0]
                for v in vals[1:]:
                    if op == "/" and v == 0:
                        return None
                    acc = acc * v if op == "*" else acc / v
                return acc

        # Parentheses
        if expr.startswith("(") and expr.endswith(")"):
            return self._eval_expr(wb, expr[1:-1], sheet_ctx, depth)

        # Sheet!Cell or Cell
        return self._eval_ref_token(wb, expr, sheet_ctx, depth)

    def _eval_ref_token(self, wb, token: str, sheet_ctx: str, depth: int) -> float | None:
        token = token.strip().replace("$", "")
        # bare number
        try:
            return float(token)
        except ValueError:
            pass
        sheet = sheet_ctx
        cell = token
        if "!" in token:
            sheet_part, cell = token.split("!", 1)
            sheet = sheet_part.strip().strip("'")
            cell = cell.strip()
        if sheet not in wb.sheetnames:
            return None
        try:
            raw = wb[sheet][cell].value
        except Exception:  # noqa: BLE001
            return None
        return self._resolve_value(wb, raw, sheet, depth)

    def _cell_raw(self, wb, token: str, sheet_ctx: str) -> Any:
        token = token.strip().replace("$", "")
        sheet = sheet_ctx
        cell = token
        if "!" in token:
            sheet_part, cell = token.split("!", 1)
            sheet = sheet_part.strip().strip("'")
            cell = cell.strip()
        if sheet not in wb.sheetnames:
            return None
        try:
            return wb[sheet][cell].value
        except Exception:  # noqa: BLE001
            return None

    @staticmethod
    def _split_args(inner: str) -> list[str]:
        parts: list[str] = []
        buf: list[str] = []
        depth = 0
        for ch in inner:
            if ch == "(":
                depth += 1
            elif ch == ")":
                depth -= 1
            if ch == "," and depth == 0:
                parts.append("".join(buf))
                buf = []
                continue
            buf.append(ch)
        if buf:
            parts.append("".join(buf))
        return parts

    @staticmethod
    def _split_top(expr: str, op: str) -> list[str]:
        parts: list[str] = []
        buf: list[str] = []
        depth = 0
        i = 0
        while i < len(expr):
            ch = expr[i]
            if ch == "(":
                depth += 1
            elif ch == ")":
                depth -= 1
            if depth == 0 and expr.startswith(op, i) and not (op == "-" and (not buf or buf[-1] in "+-*/(")):
                # avoid splitting sheet names / scientific; only split operator
                parts.append("".join(buf))
                buf = []
                i += len(op)
                continue
            buf.append(ch)
            i += 1
        if buf:
            parts.append("".join(buf))
        return parts if len(parts) > 1 else [expr]

    def _sum_range(self, wb, sheet: str, start: str, end: str) -> float | None:
        if sheet not in wb.sheetnames:
            return None
        ws = wb[sheet]
        try:
            cells = ws[f"{start}:{end}"]
        except Exception:  # noqa: BLE001
            return None
        total = 0.0
        any_v = False
        # cells may be tuple of tuples
        flat = cells if not isinstance(cells, tuple) else (
            [c for row in cells for c in (row if isinstance(row, tuple) else (row,))]
            if cells and isinstance(cells[0], tuple)
            else list(cells)
        )
        for cell in flat:
            v = self._resolve_value(wb, cell.value, sheet)
            if v is not None:
                total += v
                any_v = True
        return total if any_v else None

    # Back-compat for tests
    @classmethod
    def _eval_ref(cls, wb, formula: Any) -> float | None:
        return cls()._resolve_value(wb, formula, "Final Metrics")

    @staticmethod
    def _detect_fiscal_year(wb, fiscal_quarter: int) -> int | None:
        if LQ_IS in wb.sheetnames:
            ws = wb[LQ_IS]
            for r in range(1, 10):
                for c in range(1, 8):
                    v = ws.cell(r, c).value
                    if isinstance(v, str) and "Q" in v.upper():
                        for token in v.replace("-", " ").split():
                            if token.isdigit() and len(token) == 4:
                                return int(token)
        return None
