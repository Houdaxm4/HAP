"""Diagnostic analysis of the Operating Earnings valuation base.

Rate distortion (prospective growth) and base distortion (the dollar OE
starting point) are separate problems. This service does not write a
replacement base into the workbook and does not retune growth analysis.
"""

from __future__ import annotations

from statistics import median, pstdev
from typing import Any

from models.annual_update import (
    KEEP_REPORTED_BASE,
    USE_NORMALIZED_BASE,
    INSUFFICIENT_BASE_EVIDENCE,
    CandidateNormalizedBase,
    NormalizedEarningsPowerAnalysis,
)
from services.annual_growth_analysis_service import AnnualGrowthAnalysisService

KEEP = KEEP_REPORTED_BASE
USE = USE_NORMALIZED_BASE
INSUFFICIENT = INSUFFICIENT_BASE_EVIDENCE

_NEAR_ZERO_RATIO = 0.15
_SPIKE_MULTIPLE = 2.0
_TROUGH_RATIO = 0.50
_NI_STABLE_BAND = 0.35
_MARGIN_SHIFT = 0.05
_RELATIVE_CORROBORATION = 0.35


def _vals(pairs: list[tuple[str, float]]) -> list[float]:
    return [v for _, v in pairs]


def _mean(xs: list[float]) -> float | None:
    return sum(xs) / len(xs) if xs else None


def _med(xs: list[float]) -> float | None:
    return float(median(xs)) if xs else None


def _vol(xs: list[float]) -> float | None:
    if len(xs) < 2:
        return None
    return float(pstdev(xs))


def _ratio(a: float | None, b: float | None) -> float | None:
    if a is None or b is None or abs(b) < 1e-12:
        return None
    return a / b


def _aligned(a: list[tuple[str, float]], b: list[tuple[str, float]]) -> list[tuple[str, float, float]]:
    mb = dict(b)
    out = []
    for fy, va in a:
        if fy in mb:
            out.append((fy, va, mb[fy]))
    return out


def _intensity(numer: list[tuple[str, float]], denom: list[tuple[str, float]]) -> list[float]:
    out = []
    for _, n, d in _aligned(numer, denom):
        r = _ratio(abs(n), abs(d) if d else None)
        if r is not None:
            out.append(r)
    return out


def _component_stats(pairs: list[tuple[str, float]]) -> dict[str, Any]:
    xs = _vals(pairs)
    recent = xs[-5:] if len(xs) >= 5 else xs[-3:]
    return {
        "series": pairs,
        "n": len(xs),
        "latest": xs[-1] if xs else None,
        "mean_10y": _mean(xs),
        "median_10y": _med(xs),
        "median_5y": _med(xs[-5:] if len(xs) >= 5 else xs),
        "median_3y": _med(xs[-3:] if len(xs) >= 3 else xs),
        "recent_median": _med(recent) if recent else None,
        "volatility": _vol(xs),
        "latest_vs_10y_median": _ratio(xs[-1], _med(xs)) if xs else None,
        "latest_vs_recent_median": _ratio(xs[-1], _med(recent)) if xs and recent else None,
        "trend_last_over_first": _ratio(xs[-1], xs[0]) if len(xs) >= 2 else None,
    }


def _rel_close(a: float, b: float) -> bool:
    scale = max(abs(a), abs(b), 1.0)
    return abs(a - b) / scale <= _RELATIVE_CORROBORATION


def _cluster(values: list[float]) -> list[float] | None:
    if len(values) < 2:
        return None
    s = sorted(values)
    best: list[float] | None = None
    for i in range(len(s)):
        for j in range(i + 1, len(s)):
            subset = s[i : j + 1]
            mid = median(subset)
            scale = max(abs(mid), 1.0)
            if (subset[-1] - subset[0]) / scale > _RELATIVE_CORROBORATION:
                continue
            if best is None or len(subset) > len(best) or (
                len(subset) == len(best) and subset[-1] - subset[0] < best[-1] - best[0]
            ):
                best = subset
    if best and len(best) >= 2:
        return best
    return None


class AnnualNormalizedBaseService:
    """Build a NormalizedEarningsPowerAnalysis without mutating the workbook."""

    def analyze_workbooks(self, wb, wb_f) -> NormalizedEarningsPowerAnalysis:
        ev = AnnualGrowthAnalysisService().collect(wb, wb_f)
        return self.analyze_evidence(ev)

    def analyze_path(self, workbook_path) -> NormalizedEarningsPowerAnalysis:
        from openpyxl import load_workbook

        wb = load_workbook(workbook_path, data_only=True)
        wf = load_workbook(workbook_path, data_only=False)
        try:
            return self.analyze_workbooks(wb, wf)
        finally:
            wb.close()
            wf.close()

    def analyze_evidence(self, ev) -> NormalizedEarningsPowerAnalysis:
        ni, da, cx, oe = ev.ni, ev.da, ev.capex, ev.oe
        rev, oi = ev.revenue, ev.operating_income
        ni_s, da_s, cx_s, oe_s = _vals(ni), _vals(da), _vals(cx), _vals(oe)
        rev_s, oi_s = _vals(rev), _vals(oi)

        reported = oe_s[-1] if oe_s else None
        decomp = []
        ni_m, da_m, cx_m, oe_m, rev_m, oi_m = dict(ni), dict(da), dict(cx), dict(oe), dict(rev), dict(oi)
        for fy in ev.fy_labels or [k for k, _ in oe]:
            row = {
                "fy": fy,
                "ni": ni_m.get(fy),
                "da": da_m.get(fy),
                "capex": cx_m.get(fy),
                "oe": oe_m.get(fy),
                "revenue": rev_m.get(fy),
                "operating_income": oi_m.get(fy),
            }
            if row["revenue"] not in (None, 0) and row["oe"] is not None:
                row["oe_margin"] = row["oe"] / row["revenue"]
            if row["revenue"] not in (None, 0) and row["operating_income"] is not None:
                row["oi_margin"] = row["operating_income"] / row["revenue"]
            if row["revenue"] not in (None, 0) and row["capex"] is not None:
                row["capex_to_revenue"] = abs(row["capex"]) / abs(row["revenue"])
            if row["da"] not in (None, 0) and row["capex"] is not None:
                row["capex_to_da"] = abs(row["capex"]) / abs(row["da"])
            if row["ni"] not in (None, 0) and row["capex"] is not None:
                row["capex_to_ni"] = abs(row["capex"]) / abs(row["ni"])
            decomp.append(row)

        types, notes = self._classify_distortions(oe_s, ni_s, da_s, cx_s, rev_s, oi_s)
        candidates = self._candidates(reported, ni_s, da_s, cx_s, oe_s, rev_s, oi_s)
        independent = [c for c in candidates if c.method != "REPORTED_CURRENT_BASE" and c.result is not None]
        results = [c.result for c in independent if c.result is not None]
        lo = min(results) if results else None
        hi = max(results) if results else None

        family1 = [c.result for c in independent if c.family == "capex_intensity" and c.result is not None]
        family2 = [c.result for c in independent if c.family == "earnings_power" and c.result is not None]
        f1_cluster = _cluster(family1)
        f2_cluster = _cluster(family2)
        cross = None
        if f1_cluster and f2_cluster:
            a, b = median(f1_cluster), median(f2_cluster)
            if _rel_close(a, b):
                cross = [a, b]

        status, decision, selected, method, conf, basis = self._decide(
            reported=reported,
            types=types,
            notes=notes,
            ni_s=ni_s,
            oe_s=oe_s,
            family1_cluster=f1_cluster,
            family2_cluster=f2_cluster,
            cross=cross,
            candidates=independent,
        )

        hist = {
            "ni": _component_stats(ni),
            "da": _component_stats(da),
            "capex": _component_stats(cx),
            "oe": _component_stats(oe),
            "revenue": _component_stats(rev),
            "operating_income": _component_stats(oi),
            "capex_to_revenue_median": _med(_intensity(cx, rev)),
            "capex_to_da_median": _med(_intensity(cx, da)),
            "capex_to_ni_median": _med(_intensity(cx, ni)),
            "oe_margin_median": _med(_intensity(oe, rev)),
            "oi_margin_median": _med(_intensity(oi, rev)),
            "latest_capex_to_revenue": (_intensity(cx[-1:], rev[-1:]) or [None])[-1] if cx and rev else None,
            "latest_oe_margin": (_intensity(oe[-1:], rev[-1:]) or [None])[-1] if oe and rev else None,
        }

        growth_note = (
            "Base analysis is independent of prospective-growth KEEP/ADJUST/"
            "INSUFFICIENT. A DISTORTED_BASE flag on OE growth does not select a "
            "normalized dollar base. If a defensible normalized base is later "
            "authorized, OE growth may be re-evaluated from NI/OI/revenue evidence; "
            "that growth decision is not made here."
        )

        rationale = self._rationale(reported, types, candidates, status, decision, basis)
        evidence = [
            f"reported_oe={reported}",
            f"status={status}",
            f"decision={decision}",
            *[f"{c.method}={c.result}" for c in candidates],
            *types,
            *notes,
        ]
        return NormalizedEarningsPowerAnalysis(
            metric="operating_earnings_base",
            reported_current_base=reported,
            base_formula=getattr(ev, "ev_b9_formula", None) or "OE = NI + D&A + signed CapEx; EV forecast uses current OE base",
            component_series={"decomposition": decomp, "components": hist},
            distortions_identified=types + notes,
            distortion_type=types,
            historical_normalized_observations=hist,
            candidate_normalization_methods=[c.method for c in candidates],
            candidate_bases=candidates,
            candidate_base_low=lo,
            candidate_base_high=hi,
            selected_normalized_base=selected,
            selection_method=method,
            decision=decision,
            reported_base_status=status,
            rationale=rationale,
            evidence=evidence,
            provenance=list(ev.provenance) + ["OE identity NI+D&A+signed CapEx from Inputs; diagnostic only"],
            confidence=conf,
            valuation_impact="Diagnostic only; original OE base and EV formulas untouched",
            growth_interaction=growth_note,
            writes_to_workbook=False,
            implementation_status="DIAGNOSTIC_ONLY",
            valuation_bridge={
                "status": "not_computed",
                "reason": (
                    "Diagnostic phase only. Original OE base, EV forecast, EV/share, "
                    "and Margin of Safety remain the workbook authority. A HAP "
                    "normalized-base valuation bridge is not written until a later "
                    "implementation step authorizes it."
                ),
                "steps": [
                    "original_oe_base",
                    "normalization_adjustments",
                    "hap_normalized_oe_base",
                    "prospective_growth_assumption",
                    "hap_oe_projection",
                    "hap_enterprise_value",
                    "hap_value_per_share",
                    "hap_margin_of_safety",
                ],
            },
        )

    def _classify_distortions(
        self,
        oe: list[float],
        ni: list[float],
        da: list[float],
        capex: list[float],
        rev: list[float],
        oi: list[float],
    ) -> tuple[list[str], list[str]]:
        types: list[str] = []
        notes: list[str] = []
        if not oe:
            return ["UNKNOWN_BASE_DISTORTION"], ["No OE series"]

        last_oe = oe[-1]
        pos_oe = [x for x in oe[:-1] if x > 0]
        med_oe = _med(pos_oe) or _med([x for x in oe if x > 0])
        if last_oe <= 0:
            types.append("NEGATIVE_BASE")
        elif med_oe and abs(last_oe) < _NEAR_ZERO_RATIO * med_oe:
            types.append("NEAR_ZERO_BASE")

        if capex:
            abs_cx = [abs(x) for x in capex]
            prior = abs_cx[:-1] if len(abs_cx) > 1 else abs_cx
            med_cx = _med(prior)
            last_cx = abs_cx[-1]
            if med_cx and last_cx > _SPIKE_MULTIPLE * med_cx:
                types.append("CAPEX_SPIKE")
            elif med_cx and last_cx < _TROUGH_RATIO * med_cx:
                types.append("CAPEX_TROUGH")
            earlier = abs_cx[:-3] if len(abs_cx) >= 6 else []
            earlier_med = _med(earlier)
            if earlier_med and all(v > 1.5 * earlier_med for v in abs_cx[-3:]):
                types.append("STRUCTURAL_REINVESTMENT_CHANGE")
                notes.append(
                    "CapEx intensity has been elevated for multiple recent years versus "
                    "earlier history; a single-year temporary spike is not established."
                )

        if ni:
            last_ni, med_ni = ni[-1], _med([abs(x) for x in ni[:-1]] or [abs(x) for x in ni])
            if med_ni:
                if last_ni < _TROUGH_RATIO * med_ni and last_ni > 0:
                    types.append("ONE_TIME_EARNINGS_COLLAPSE")
                elif last_ni > 1.5 * med_ni:
                    types.append("ONE_TIME_EARNINGS_WINDFALL")
                if abs(last_ni - med_ni) / med_ni <= _NI_STABLE_BAND:
                    notes.append("NI is stable relative to its historical median.")

        if rev and oi and len(rev) >= 2:
            m_hist = _intensity(list(zip([str(i) for i in range(len(oi) - 1)], oi[:-1])), list(zip([str(i) for i in range(len(rev) - 1)], rev[:-1])))
            m_now = _ratio(oi[-1], rev[-1])
            med_m = _med(m_hist)
            if med_m is not None and m_now is not None and abs(m_now - med_m) > _MARGIN_SHIFT:
                types.append("MARGIN_DISLOCATION")

        if oe and len(oe) >= 6:
            def _cyclical(series: list[float], peak: bool) -> bool:
                if len(series) < 6:
                    return False
                diffs = [series[i] - series[i - 1] for i in range(1, len(series))]
                sign_flips = sum(
                    1 for i in range(1, len(diffs)) if diffs[i] * diffs[i - 1] < 0
                )
                med_all = _med(series)
                if sign_flips < 2 or not med_all:
                    return False
                if peak:
                    return series[-1] == max(series) and max(series) > 1.25 * med_all
                return series[-1] == min(series) and min(series) > 0 and min(series) < 0.75 * med_all

            # OE-only oscillation is often CapEx, not an earnings cycle.
            if _cyclical(oe, peak=True) and (_cyclical(oi, peak=True) or _cyclical(ni, peak=True)):
                types.append("CYCLICAL_PEAK")
            elif _cyclical(oe, peak=False) and (_cyclical(oi, peak=False) or _cyclical(ni, peak=False)):
                types.append("CYCLICAL_TROUGH")

        if not types:
            notes.append("No supported base-distortion type from workbook series.")
        return types, notes

    def _candidates(
        self,
        reported: float | None,
        ni: list[float],
        da: list[float],
        capex: list[float],
        oe: list[float],
        rev: list[float],
        oi: list[float],
    ) -> list[CandidateNormalizedBase]:
        out: list[CandidateNormalizedBase] = []
        last_ni = ni[-1] if ni else None
        last_da = da[-1] if da else None
        last_cx = capex[-1] if capex else None
        last_rev = rev[-1] if rev else None

        out.append(
            CandidateNormalizedBase(
                method="REPORTED_CURRENT_BASE",
                result=reported,
                calculation="Latest OE = latest NI + D&A + signed CapEx",
                economic_rationale="The workbook's current Operating Earnings endpoint, unadjusted.",
                strengths=["Matches the analyst model exactly", "No HAP reconstruction"],
                weaknesses=["Unusable if the endpoint is not earnings power"],
                uses_current_scale=True,
                assumes_capex_mean_reversion=False,
                family="reported",
            )
        )

        recent_oe = oe[-5:] if len(oe) >= 5 else oe
        excl = oe[:-1] if len(oe) >= 4 else oe
        out.append(
            CandidateNormalizedBase(
                method="RECENT_MEDIAN_OE",
                result=_med(recent_oe),
                calculation="Median of the last up-to-5 OE observations, including the endpoint",
                economic_rationale="A recent central OE observation, still in dollars.",
                strengths=["Simple", "Uses the company's own OE identity"],
                weaknesses=[
                    "Mixes the possibly distorted endpoint into the median",
                    "Historical dollar OE mixes scale if the firm has grown",
                ],
                uses_current_scale=False,
                assumes_capex_mean_reversion=False,
                family="mixed_scale",
            )
        )
        out.append(
            CandidateNormalizedBase(
                method="PRIOR_MEDIAN_OE_EXCLUDING_ENDPOINT",
                result=_med(excl) if excl else None,
                calculation="Median OE excluding the latest year",
                economic_rationale="Asks what OE looked like before the endpoint year.",
                strengths=["Removes the suspect endpoint"],
                weaknesses=["Still historical dollars; may mix scale", "May exclude a genuine new regime"],
                uses_current_scale=False,
                assumes_capex_mean_reversion=False,
                family="mixed_scale",
            )
        )

        # CapEx sign: OE = NI + DA + capex_signed. Normalize |capex| then re-apply outflow sign.
        cx_sign = -1.0 if (last_cx is not None and last_cx < 0) or (capex and sum(1 for x in capex if x < 0) >= sum(1 for x in capex if x > 0)) else 1.0
        med_abs_cx = _med([abs(x) for x in capex[:-1]]) if len(capex) > 1 else _med([abs(x) for x in capex])
        if last_ni is not None and last_da is not None and med_abs_cx is not None:
            norm_cx = cx_sign * med_abs_cx
            out.append(
                CandidateNormalizedBase(
                    method="NORMALIZED_CAPEX_CURRENT_NI_DA",
                    result=last_ni + last_da + norm_cx,
                    calculation=(
                        f"current NI ({last_ni}) + current D&A ({last_da}) + "
                        f"median prior |CapEx| with original outflow sign ({norm_cx})"
                    ),
                    economic_rationale=(
                        "Holds current earnings and depreciation fixed and replaces only the "
                        "CapEx term with a historical central CapEx dollar amount."
                    ),
                    strengths=["Preserves current NI/D&A scale", "Isolates the CapEx term"],
                    weaknesses=[
                        "Historical median CapEx dollars mix scale",
                        "Assumes CapEx mean-reverts and is not a new required intensity",
                    ],
                    uses_current_scale=True,
                    assumes_capex_mean_reversion=True,
                    family="capex_intensity",
                )
            )

        capex_da = _intensity(list(zip(map(str, range(len(capex))), capex)), list(zip(map(str, range(len(da))), da)))
        # exclude last year from median intensity when possible
        hist_cda = capex_da[:-1] if len(capex_da) > 2 else capex_da
        med_cda = _med(hist_cda)
        if last_ni is not None and last_da is not None and med_cda is not None:
            sust_abs = med_cda * abs(last_da)
            sust_cx = cx_sign * sust_abs
            out.append(
                CandidateNormalizedBase(
                    method="CAPEX_TO_DA",
                    result=last_ni + last_da + sust_cx,
                    calculation=(
                        f"sustainable |CapEx| = median historical |CapEx|/D&A ({med_cda:.2f}) "
                        f"× current D&A; OE = current NI + D&A + signed sustainable CapEx"
                    ),
                    economic_rationale=(
                        "Treats depreciation as a rough current-scale proxy for the capital "
                        "base and applies the company's own historical CapEx/D&A intensity."
                    ),
                    strengths=["Preserves current scale via current D&A", "Company-specific intensity"],
                    weaknesses=[
                        "Assumes historical CapEx/D&A is the sustainable intensity",
                        "Does not prove a CapEx spike is temporary",
                    ],
                    uses_current_scale=True,
                    assumes_capex_mean_reversion=True,
                    family="capex_intensity",
                )
            )

        capex_rev = _intensity(list(zip(map(str, range(len(capex))), capex)), list(zip(map(str, range(len(rev))), rev)))
        hist_cr = capex_rev[:-1] if len(capex_rev) > 2 else capex_rev
        med_cr = _med(hist_cr)
        if last_ni is not None and last_da is not None and last_rev is not None and med_cr is not None:
            sust_abs = med_cr * abs(last_rev)
            sust_cx = cx_sign * sust_abs
            out.append(
                CandidateNormalizedBase(
                    method="CAPEX_TO_REVENUE",
                    result=last_ni + last_da + sust_cx,
                    calculation=(
                        f"sustainable |CapEx| = median historical |CapEx|/Revenue ({med_cr:.3f}) "
                        f"× current revenue; OE = current NI + D&A + signed sustainable CapEx"
                    ),
                    economic_rationale=(
                        "Applies historical capital intensity to current revenue so the "
                        "CapEx reconstruction stays at current company scale."
                    ),
                    strengths=["Preserves current scale via current revenue"],
                    weaknesses=[
                        "Assumes historical capital intensity remains the sustainable rate",
                        "Does not prove a CapEx spike is temporary versus structural",
                    ],
                    uses_current_scale=True,
                    assumes_capex_mean_reversion=True,
                    family="capex_intensity",
                )
            )

        oe_m = _intensity(list(zip(map(str, range(len(oe))), oe)), list(zip(map(str, range(len(rev))), rev)))
        hist_om = oe_m[:-1] if len(oe_m) > 2 else oe_m
        med_om = _med(hist_om)
        if last_rev is not None and med_om is not None:
            out.append(
                CandidateNormalizedBase(
                    method="OE_MARGIN_TIMES_CURRENT_REVENUE",
                    result=med_om * last_rev,
                    calculation=(
                        f"median historical OE/Revenue excluding endpoint ({med_om:.3f}) "
                        f"× current revenue ({last_rev})"
                    ),
                    economic_rationale=(
                        "Reconstructs earnings power at current scale from the company's "
                        "own OE margin history, without copying historical OE dollars."
                    ),
                    strengths=["Preserves current scale", "Does not use historical average dollar OE"],
                    weaknesses=[
                        "Assumes historical OE margin is still attainable",
                        "OE margin embeds CapEx intensity, so it is not fully independent of CapEx regime change",
                    ],
                    uses_current_scale=True,
                    assumes_capex_mean_reversion=False,
                    family="earnings_power",
                )
            )

        oi_m = _intensity(list(zip(map(str, range(len(oi))), oi)), list(zip(map(str, range(len(rev))), rev)))
        hist_im = oi_m[:-1] if len(oi_m) > 2 else oi_m
        med_im = _med(hist_im)
        if last_rev is not None and med_im is not None:
            out.append(
                CandidateNormalizedBase(
                    method="OI_MARGIN_TIMES_CURRENT_REVENUE",
                    result=med_im * last_rev,
                    calculation=(
                        f"median historical Operating Income/Revenue excluding endpoint "
                        f"({med_im:.3f}) × current revenue ({last_rev})"
                    ),
                    economic_rationale=(
                        "Current-scale operating-income power. Independent of the "
                        "NI+D&A+CapEx identity, so it does not assume CapEx mean-reverts."
                    ),
                    strengths=[
                        "Preserves current scale",
                        "Not a CapEx-normalization method",
                    ],
                    weaknesses=[
                        "OI is not the workbook OE valuation base",
                        "Assumes historical OI margin is still attainable",
                    ],
                    uses_current_scale=True,
                    assumes_capex_mean_reversion=False,
                    family="earnings_power",
                )
            )

        ni_m = _intensity(list(zip(map(str, range(len(ni))), ni)), list(zip(map(str, range(len(rev))), rev)))
        hist_nm = ni_m[:-1] if len(ni_m) > 2 else ni_m
        med_nm = _med(hist_nm)
        da_m = _intensity(list(zip(map(str, range(len(da))), da)), list(zip(map(str, range(len(rev))), rev)))
        hist_dm = da_m[:-1] if len(da_m) > 2 else da_m
        med_dm = _med(hist_dm)
        if last_rev is not None and med_nm is not None and med_cr is not None:
            # NI margin × rev + DA margin × rev + signed capex/rev
            rec_ni = med_nm * last_rev
            rec_da = (med_dm * last_rev) if med_dm is not None else (last_da or 0.0)
            rec_cx = cx_sign * med_cr * abs(last_rev)
            out.append(
                CandidateNormalizedBase(
                    method="COMPONENT_MARGINS_AT_CURRENT_REVENUE",
                    result=rec_ni + rec_da + rec_cx,
                    calculation=(
                        "current-scale NI + D&A + CapEx reconstructed from historical "
                        "component-to-revenue intensities applied to current revenue"
                    ),
                    economic_rationale="Builds OE from component intensities at current revenue.",
                    strengths=["Current scale", "Decomposes the identity"],
                    weaknesses=["Each intensity still assumes historical ratios are sustainable"],
                    uses_current_scale=True,
                    assumes_capex_mean_reversion=True,
                    family="capex_intensity",
                )
            )

        out.append(
            CandidateNormalizedBase(
                method="TEN_YEAR_AVERAGE_OE_DOLLARS",
                result=_mean(oe) if oe else None,
                calculation="Arithmetic mean of historical OE dollars",
                economic_rationale="Shown as a contrast: a 10-year average is not automatically earnings power.",
                strengths=["Transparent"],
                weaknesses=[
                    "Mixes historical company scale",
                    "Not preferred when the firm has grown or changed intensity",
                    "Must not win automatically",
                ],
                uses_current_scale=False,
                assumes_capex_mean_reversion=False,
                family="mixed_scale",
            )
        )
        return out

    def _decide(
        self,
        *,
        reported: float | None,
        types: list[str],
        notes: list[str],
        ni_s: list[float],
        oe_s: list[float],
        family1_cluster: list[float] | None,
        family2_cluster: list[float] | None,
        cross: list[float] | None,
        candidates: list[CandidateNormalizedBase],
    ) -> tuple[str, str, float | None, str | None, float, str]:
        capex_story = "CAPEX_SPIKE" in types or "STRUCTURAL_REINVESTMENT_CHANGE" in types or "CAPEX_TROUGH" in types
        near_zero = "NEAR_ZERO_BASE" in types or "NEGATIVE_BASE" in types
        ni_event = "ONE_TIME_EARNINGS_COLLAPSE" in types or "ONE_TIME_EARNINGS_WINDFALL" in types
        pos_prior = [x for x in oe_s[:-1] if x > 0]
        med_prior = _med(pos_prior)

        unrepresentative = False
        if near_zero:
            unrepresentative = True
        if reported is not None and med_prior and abs(reported) < 0.5 * med_prior:
            unrepresentative = True
        if reported is not None and med_prior and abs(reported) > 1.5 * med_prior and ni_event:
            unrepresentative = True
        if ni_event and "ONE_TIME_EARNINGS_COLLAPSE" in types:
            unrepresentative = True

        material = False
        if reported is not None and med_prior:
            if abs(reported - med_prior) / med_prior >= 0.5:
                material = True
        if near_zero:
            material = True

        if unrepresentative and material:
            status = "materially_distorted"
        elif types:
            status = "potentially_distorted"
        else:
            status = "usable"

        # A CapEx regime shift does not prove the CapEx is temporary.
        # Capex-intensity candidates that cluster are still one assumption family.
        if capex_story and not ni_event and unrepresentative and material:
            return (
                status,
                INSUFFICIENT,
                None,
                "insufficient_evidence_capex_regime_not_shown_to_be_temporary",
                0.45,
                (
                    "The reported OE endpoint is economically unrepresentative and material "
                    "because reinvestment moved far from history while NI did not. Candidate "
                    "methods that replace CapEx with historical intensity assume mean-reversion. "
                    "Elevated CapEx can be temporary OR a structural change in required "
                    "reinvestment; workbook series alone do not distinguish those. "
                    "INSUFFICIENT_EVIDENCE. No method is averaged into a substitute, and no "
                    "base is written into the workbook."
                ),
            )

        if not (unrepresentative and material):
            return (
                status,
                KEEP,
                None,
                "reported_base_consistent_with_component_history"
                if status == "usable"
                else "reported_base_retained_distortion_not_material_and_replaceable",
                0.75 if status == "usable" else 0.6,
                (
                    "The reported OE base is not shown to be both economically unrepresentative "
                    "and material to a from-the-base valuation with a defensible replacement. "
                    "HAP keeps the reported base. Prospective growth remains a separate decision. "
                    "A CapEx flag without an unrepresentative OE endpoint does not trigger "
                    "normalization."
                ),
            )

        if ni_event and family2_cluster and not capex_story:
            selected = float(median(family2_cluster))
            return (
                status,
                USE,
                selected,
                "corroborating_earnings_power_methods",
                0.65,
                (
                    "A one-time earnings-level event is supported by the NI series, and "
                    "independent current-scale earnings-power methods corroborate an alternative "
                    "base. This is a diagnostic conclusion only; the workbook base is not overwritten."
                ),
            )

        if unrepresentative and material and cross and not capex_story:
            selected = float(median(cross))
            return (
                status,
                USE,
                selected,
                "cross_family_corroboration",
                0.6,
                (
                    "The reported base is economically unrepresentative and material, and "
                    "independent method families corroborate. Diagnostic only; original base "
                    "untouched."
                ),
            )

        return (
            status,
            INSUFFICIENT,
            None,
            "insufficient_evidence_no_defensible_alternative_base",
            0.4,
            (
                "The reported base appears unrepresentative and material to a from-the-base "
                "valuation, but available workbook evidence does not support a single "
                "defensible replacement. Conflicting or assumption-laden candidate methods "
                "are not averaged. INSUFFICIENT_EVIDENCE."
            ),
        )

    def _rationale(self, reported, types, candidates, status, decision, basis) -> str:
        parts = [
            f"Reported current OE base is {reported}.",
            "OE is reconstructed as NI + D&A + signed CapEx from Inputs.",
            "This is a base-distortion diagnostic, not a growth-rate decision.",
            f"Distortion types: {', '.join(types) or 'none'}.",
            f"Reported-base status: {status}. Decision: {decision}.",
            "No candidate is an automatic winner; a 10-year average of OE dollars is not treated as normalized earnings power.",
            basis,
            "Original workbook OE/EV cells are not written.",
        ]
        return " ".join(parts)
