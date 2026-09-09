"""M4 Excel fill engine — apply validated WRITE intents to a workbook copy."""

from __future__ import annotations

import shutil
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from openpyxl.worksheet.worksheet import Worksheet

from models.cell_diff import CellDiffEntry, CellDiffReport
from models.provenance import CellProvenance, CellTransformation, ProvenanceReport
from workbook_mapping.engine import IntentDecision, WriteIntent, WriteIntentReport


class ExcelFillError(ValueError):
    """Raised when fill cannot proceed safely."""


def _is_formula_cell(cell: Any) -> bool:
    if cell.data_type == "f":
        return True
    return isinstance(cell.value, str) and str(cell.value).startswith("=")


def _sheet_is_protected(sheet: Worksheet) -> bool:
    protection = getattr(sheet, "protection", None)
    if protection is None:
        return False
    return bool(getattr(protection, "sheet", False))


class ExcelFillService:
    """
    Copy the uploaded Industrial Template and apply only M3 WRITE intents.

    Never rebuilds sheets; only assigns ``cell.value`` so styles/number formats
    and structure are preserved by the openpyxl round-trip of a full-file copy.
    """

    def apply_write_intents(
        self,
        *,
        analysis_id: str,
        ticker: str,
        source_workbook_path: Path,
        destination_workbook_path: Path,
        intent_report: WriteIntentReport,
    ) -> tuple[list[CellProvenance], CellDiffReport, int, int, int]:
        """
        Returns (excel_provenance_entries, cell_diff, written, runtime_skip, runtime_block).
        Leaves ``source_workbook_path`` unchanged.
        """
        if not source_workbook_path.exists():
            raise ExcelFillError(f"Source workbook missing: {source_workbook_path}")

        destination_workbook_path.parent.mkdir(parents=True, exist_ok=True)
        shutil.copy2(source_workbook_path, destination_workbook_path)

        source_sheet_names = self._sheet_names(source_workbook_path)
        workbook = load_workbook(destination_workbook_path)
        try:
            diff_entries: list[CellDiffEntry] = []
            provenance: list[CellProvenance] = []
            written = 0
            runtime_skip = 0
            runtime_block = 0
            m3_skip = 0
            m3_block = 0

            for intent in intent_report.intents:
                if intent.decision == IntentDecision.SKIP:
                    m3_skip += 1
                    diff_entries.append(self._diff_from_intent(intent, write_decision="SKIP"))
                    provenance.append(self._prov_from_intent(intent, status="skipped", original=None))
                    continue

                if intent.decision == IntentDecision.BLOCK:
                    m3_block += 1
                    diff_entries.append(self._diff_from_intent(intent, write_decision="BLOCK"))
                    provenance.append(self._prov_from_intent(intent, status="blocked", original=None))
                    continue

                if intent.decision != IntentDecision.WRITE:
                    runtime_block += 1
                    reason = f"Unsupported intent decision {intent.decision!r}"
                    diff_entries.append(
                        self._diff_from_intent(intent, write_decision="BLOCK", reason=reason)
                    )
                    provenance.append(
                        self._prov_from_intent(
                            intent,
                            status="blocked",
                            original=None,
                            failure_reason=reason,
                        )
                    )
                    continue

                # --- Runtime safety for WRITE intents only ---
                if intent.sheet not in workbook.sheetnames:
                    runtime_block += 1
                    reason = f"Sheet '{intent.sheet}' not present in workbook at fill time"
                    diff_entries.append(
                        self._diff_from_intent(intent, write_decision="BLOCK", reason=reason)
                    )
                    provenance.append(
                        self._prov_from_intent(
                            intent, status="blocked", original=None, failure_reason=reason
                        )
                    )
                    continue

                sheet = workbook[intent.sheet]
                if _sheet_is_protected(sheet):
                    runtime_block += 1
                    reason = f"Worksheet '{intent.sheet}' is protected; write refused"
                    diff_entries.append(
                        self._diff_from_intent(intent, write_decision="BLOCK", reason=reason)
                    )
                    provenance.append(
                        self._prov_from_intent(
                            intent, status="blocked", original=None, failure_reason=reason
                        )
                    )
                    continue

                cell = sheet[intent.cell]
                original = cell.value

                if _is_formula_cell(cell):
                    runtime_block += 1
                    reason = (
                        f"Runtime conflict: {intent.sheet}!{intent.cell} contains a formula; "
                        "never overwrite"
                    )
                    diff_entries.append(
                        self._diff_from_intent(
                            intent,
                            write_decision="BLOCK",
                            reason=reason,
                            original_value=original,
                            new_value=original,
                        )
                    )
                    provenance.append(
                        self._prov_from_intent(
                            intent,
                            status="skipped_formula",
                            original=original,
                            failure_reason=reason,
                        )
                    )
                    continue

                if intent.value is None:
                    runtime_skip += 1
                    reason = "WRITE intent has null value at fill time; skipped"
                    diff_entries.append(
                        self._diff_from_intent(
                            intent,
                            write_decision="SKIP",
                            reason=reason,
                            original_value=original,
                        )
                    )
                    provenance.append(
                        self._prov_from_intent(
                            intent, status="blank", original=original, failure_reason=reason
                        )
                    )
                    continue

                # Assign value only — preserves number_format, font, fill, borders, etc.
                cell.value = intent.value
                written += 1
                changed = original != intent.value
                diff_entries.append(
                    self._diff_from_intent(
                        intent,
                        write_decision="WRITE",
                        reason=intent.reason,
                        original_value=original,
                        new_value=intent.value,
                        changed=changed,
                    )
                )
                provenance.append(
                    self._prov_from_intent(
                        intent,
                        status="filled",
                        original=original,
                        written_value=intent.value,
                    )
                )

            workbook.save(destination_workbook_path)
        finally:
            workbook.close()

        # Reload validation — structure intact
        self._assert_reload_ok(destination_workbook_path, source_sheet_names)

        changed_count = sum(1 for e in diff_entries if e.changed)
        cell_diff = CellDiffReport(
            analysis_id=analysis_id,
            ticker=ticker,
            source_workbook=str(source_workbook_path),
            completed_workbook=str(destination_workbook_path),
            entries=diff_entries,
            written_count=written,
            runtime_skip_count=runtime_skip,
            runtime_block_count=runtime_block,
            m3_skip_count=m3_skip,
            m3_block_count=m3_block,
            changed_count=changed_count,
        )
        return provenance, cell_diff, written, runtime_skip, runtime_block

    @staticmethod
    def _sheet_names(path: Path) -> list[str]:
        wb = load_workbook(path, read_only=True, data_only=False)
        try:
            return list(wb.sheetnames)
        finally:
            wb.close()

    @staticmethod
    def _assert_reload_ok(path: Path, expected_sheets: list[str]) -> None:
        wb = load_workbook(path, data_only=False)
        try:
            if list(wb.sheetnames) != expected_sheets:
                raise ExcelFillError(
                    "Reload validation failed: worksheet order/names differ from source "
                    f"(expected {expected_sheets!r}, got {list(wb.sheetnames)!r})"
                )
        finally:
            wb.close()

    @staticmethod
    def _diff_from_intent(
        intent: WriteIntent,
        *,
        write_decision: str,
        reason: str | None = None,
        original_value: Any = None,
        new_value: Any = None,
        changed: bool = False,
    ) -> CellDiffEntry:
        return CellDiffEntry(
            sheet=intent.sheet,
            cell=intent.cell,
            cell_ref=intent.cell_ref,
            original_value=original_value,
            new_value=new_value,
            metric=intent.metric,
            period=intent.period,
            source=intent.source,
            confidence=intent.confidence,
            transformation=intent.transformation,
            write_decision=write_decision,
            reason=reason or intent.reason,
            mapping_id=intent.mapping_id,
            intent_id=intent.intent_id,
            changed=changed,
        )

    @staticmethod
    def _prov_from_intent(
        intent: WriteIntent,
        *,
        status: str,
        original: Any,
        written_value: Any | None = None,
        failure_reason: str | None = None,
    ) -> CellProvenance:
        transformations: list[CellTransformation] = []
        if intent.transformation:
            transformations.append(
                CellTransformation(
                    type=intent.transformation,
                    description=f"Mapping transformation '{intent.transformation}'",
                    input_value=None,
                    output_value=written_value if written_value is not None else intent.value,
                )
            )
        return CellProvenance(
            cell_ref=intent.cell_ref,
            worksheet=intent.sheet,
            cell=intent.cell,
            concept=intent.metric,
            period=intent.period,
            value=written_value if status == "filled" else None,
            original_value=original,
            status=status,
            source_document=intent.source_document or intent.source,
            filing_type=intent.filing_type,
            accession_number=intent.accession_number,
            xbrl_tag=intent.xbrl_tag,
            confidence=intent.confidence,
            transformations=transformations,
            reasoning=intent.reason,
            failure_reason=failure_reason,
            write_decision=intent.decision.value if status != "filled" else IntentDecision.WRITE.value,
        )
