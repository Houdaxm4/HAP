"""ROIC adjustments, applied the way an analyst applies them by hand and logged in the HAP Adjustments tab.

Operating assets / operating liabilities (Inputs tab): the formulas add up Balance Sheet lines. HAP edits those formulas
- removes a line that is clearly non-operating (held for sale, discontinued operations, derivatives, investments), and
- adds a clearly operating line that the formula leaves out when it is material (5% of the section total).

Operating income (Income - GAAP tab): a one-time charge or gain that sits inside operating expenses is taken out by changing the
component line and its total by the same amount, so the tab still adds up. Evidence is the company's own SEC XBRL facts for
that fiscal year (restructuring, impairments, acquisition costs, litigation settlements, gains on sale of assets).

New company: every year in the table. Annual update: newest year only (earlier years are copied from the previous file).
"""

from __future__ import annotations

import re
from dataclasses import dataclass, field
from datetime import date
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from services.adjustment_ledger_service import AdjustmentLedger
from services.annual_period_service import detect_year_columns

BS = "Balance Sheet - Standardized"
INPUTS = "Inputs"
INCOME = "Income - GAAP"

# Inputs rows that hold the formulas
ROW_OP_ASSETS_CURRENT, ROW_OP_ASSETS_NONCURRENT = 81, 82
ROW_OP_LIAB_CURRENT, ROW_OP_LIAB_NONCURRENT = 84, 85
ROW_OP_ASSETS_TOTAL, ROW_OP_LIAB_TOTAL = 80, 83
# Income - GAAP rows
ROW_OPEX_TOTAL, ROW_OPEX_OTHER, ROW_OPEX_SGA, ROW_OPERATING_INCOME = 22, 28, 23, 30

BS_MATERIALITY = 0.05
OI_MATERIALITY = 0.02

# Balance-sheet lines that are clearly non-operating (removed from operating assets if they were included)
NON_OPERATING_ASSET_ROWS = {12: "ST investments", 28: "ST derivative and hedging assets", 29: "Assets held for sale",
                            30: "ST deferred tax assets", 32: "Assets of discontinued operations (ST)", 44: "LT investments",
                            45: "LT marketable securities", 52: "LT deferred tax assets", 53: "LT derivative and hedging assets",
                            55: "Investments in affiliates", 56: "Assets of discontinued operations (LT)"}
# Operating lines the template formula leaves out, added when material: (row, section, label)
OPERATING_ASSET_ADDITIONS = ((19, "current", "Unbilled revenues"),)
OPERATING_LIABILITY_ADDITIONS = ((66, "current", "Accrued taxes"), (93, "noncurrent", "Accrued liabilities (LT)"))
NON_OPERATING_LIABILITY_ROWS = {67: "Interest and dividends payable", 79: "ST derivatives and hedging", 81: "Liabilities of discontinued operations (ST)",
                                100: "LT derivatives and hedging", 101: "Liabilities of discontinued operations (LT)"}

# One-time items inside operating income, grouped in families because one event is often reported under several tags
# (for example an impairment under both AssetImpairmentCharges and ImpairmentOfIntangibleAssetsExcludingGoodwill).
# Each family is counted once: (key, label, kind, tags). "charge" lowers operating income, "gain" raises it.
ONE_TIME_FAMILIES: tuple[tuple[str, str, str, tuple[str, ...]], ...] = (
    ("restructuring", "Restructuring charges", "charge", ("RestructuringCharges", "RestructuringSettlementAndImpairmentProvisions")),
    ("impairment", "Impairment charges", "charge", ("AssetImpairmentCharges", "GoodwillImpairmentLoss",
                                                    "ImpairmentOfLongLivedAssetsHeldForUse", "ImpairmentOfIntangibleAssetsExcludingGoodwill")),
    ("acquisition_costs", "Acquisition-related costs", "charge", ("BusinessCombinationAcquisitionRelatedCosts",)),
    ("litigation", "Litigation settlement expense", "charge", ("LitigationSettlementExpense",)),
    ("asset_sale_gain", "Gain on sale of assets", "gain", ("GainLossOnSaleOfPropertyPlantEquipment", "GainLossOnDispositionOfAssets")),
)


def family_amount(key: str, values: dict[str, float]) -> tuple[float, list[str]]:
    """One amount per family from the tags found. Overlapping tags are not added together."""
    present = {tag: v for tag, v in values.items() if v}
    if not present:
        return 0.0, []
    if key == "impairment":
        if "AssetImpairmentCharges" in present:
            return present["AssetImpairmentCharges"], ["AssetImpairmentCharges"]
        goodwill = present.get("GoodwillImpairmentLoss", 0.0)
        others = {t: v for t, v in present.items() if t != "GoodwillImpairmentLoss"}
        used = ["GoodwillImpairmentLoss"] if goodwill else []
        other_val = 0.0
        if others:
            tag = max(others, key=lambda t: abs(others[t]))
            other_val = others[tag]
            used.append(tag)
        return goodwill + other_val, used
    tag = max(present, key=lambda t: abs(present[t]))
    return present[tag], [tag]


@dataclass
class RoicAdjustment:
    fiscal_year: str
    category: str
    cells: list[str]
    description: str
    amount: float
    reason: str
    source: str
    adj_id: str = ""


@dataclass
class RoicAdjustmentReport:
    adjustments: list[RoicAdjustment] = field(default_factory=list)
    skipped: list[str] = field(default_factory=list)

    @property
    def changed(self) -> bool:
        return bool(self.adjustments)


def _num(v: Any) -> float | None:
    if isinstance(v, (int, float)) and not isinstance(v, bool):
        return float(v)
    return None


def _term(col: str, row: int) -> str:
    ref = f"'{BS}'!{col}{row}"
    return f'IF({ref}="",0,{ref})'


def formula_rows(formula: Any, col: str) -> set[int]:
    if not isinstance(formula, str):
        return set()
    return {int(x) for x in re.findall(rf"'{re.escape(BS)}'!{col}(\d+)", formula)}


def remove_term(formula: str, col: str, row: int) -> str:
    term = re.escape(_term(col, row))
    out = re.sub(rf"\+{term}", "", formula, count=1)
    if out == formula:
        out = re.sub(rf"{term}\+", "", formula, count=1)
    return out


def add_term(formula: str, col: str, row: int) -> str:
    if formula.strip() == "=":
        return "=" + _term(col, row)
    return formula + "+" + _term(col, row)


def facts_for_year(company_facts: dict[str, Any] | None, tag: str, year_end: date | None) -> float | None:
    """Annual USD value of a us-gaap tag for the fiscal year ending near `year_end` (10-K, about 12 months)."""
    if not company_facts or year_end is None:
        return None
    entries = (((company_facts.get("facts") or {}).get("us-gaap") or {}).get(tag) or {}).get("units", {}).get("USD") or []
    best: tuple[str, float] | None = None
    for item in entries:
        if not str(item.get("form") or "").startswith("10-K"):
            continue
        try:
            end = date.fromisoformat(str(item.get("end")))
            start = date.fromisoformat(str(item.get("start")))
        except (TypeError, ValueError):
            continue
        if abs((end - year_end).days) > 7 or not 340 <= (end - start).days <= 380:
            continue
        filed = str(item.get("filed") or "")
        if best is None or filed > best[0]:
            best = (filed, float(item.get("val")))
    return best[1] if best else None


class RoicAdjustmentService:
    def apply(
        self,
        *,
        workbook_path: Path,
        company_facts: dict[str, Any] | None = None,
        newest_only_fy: str | None = None,
        filing_hint: str = "SEC 10-K (XBRL company facts)",
    ) -> RoicAdjustmentReport:
        report = RoicAdjustmentReport()
        path = Path(workbook_path)
        cached = load_workbook(path, data_only=True)
        wb = load_workbook(path, data_only=False)
        try:
            if not all(name in wb.sheetnames for name in (INPUTS, BS, INCOME)):
                report.skipped.append("Inputs, Balance Sheet or Income tab missing.")
                return report
            ledger = AdjustmentLedger(wb)
            cols = {k: c for k, c in detect_year_columns(wb[INPUTS], wb).items() if str(k).startswith("FY")}
            income_cols = {k: c for k, c in detect_year_columns(wb[INCOME], wb).items() if str(k).startswith("FY")}
            multiplier = self._unit_multiplier(wb[INCOME])
            for fy, col_idx in sorted(cols.items(), key=lambda kv: kv[1]):
                if newest_only_fy is not None and fy != newest_only_fy:
                    continue
                letter = wb[INPUTS].cell(1, col_idx).column_letter
                self._balance_sheet(wb, cached, ledger, report, fy, letter)
                if fy in income_cols:
                    self._operating_income(wb, cached, ledger, report, fy, income_cols[fy], company_facts, multiplier, filing_hint)
            if report.changed:
                wb.save(path)
        finally:
            wb.close()
            cached.close()
        return report

    # ------------------------------------------------------------------ balance sheet lines
    def _balance_sheet(self, wb, cached, ledger, report, fy: str, col: str) -> None:
        ws = wb[INPUTS]
        cws = cached[INPUTS]
        bs_cached = cached[BS]
        sections = (
            ("Operating assets", ROW_OP_ASSETS_CURRENT, ROW_OP_ASSETS_NONCURRENT, ROW_OP_ASSETS_TOTAL, NON_OPERATING_ASSET_ROWS, OPERATING_ASSET_ADDITIONS),
            ("Operating liabilities", ROW_OP_LIAB_CURRENT, ROW_OP_LIAB_NONCURRENT, ROW_OP_LIAB_TOTAL, NON_OPERATING_LIABILITY_ROWS, OPERATING_LIABILITY_ADDITIONS),
        )
        for category, cur_row, non_row, total_row, bad_rows, additions in sections:
            total = _num(cws[f"{col}{total_row}"].value)
            for row_no in (cur_row, non_row):
                cell = ws[f"{col}{row_no}"]
                formula = cell.value
                if not isinstance(formula, str) or not formula.startswith("="):
                    continue
                for bs_row in sorted(formula_rows(formula, col) & set(bad_rows)):
                    amount = _num(bs_cached[f"{col}{bs_row}"].value) or 0.0
                    if amount == 0:
                        continue
                    new_formula = remove_term(formula, col, bs_row)
                    if new_formula == formula:
                        continue
                    self._record(ledger, report, cell, fy, category, f"Removed {bad_rows[bs_row]} from {category.lower()}",
                                 formula, new_formula, -amount,
                                 f"{bad_rows[bs_row]} ({amount:,.1f}) is not an operating item, so it should not be in invested capital.",
                                 f"{BS}!{col}{bs_row} (balance sheet line)", "formula_term_removed", "medium")
                    formula = new_formula
            for bs_row, section, label in additions:
                row_no = cur_row if section == "current" else non_row
                cell = ws[f"{col}{row_no}"]
                formula = cell.value
                if not isinstance(formula, str) or not formula.startswith("="):
                    continue
                if any(bs_row in formula_rows(ws[f"{col}{r}"].value, col) for r in (cur_row, non_row)):
                    continue
                amount = _num(bs_cached[f"{col}{bs_row}"].value)
                if not amount or total is None or total <= 0 or abs(amount) / total < BS_MATERIALITY:
                    continue
                new_formula = add_term(formula, col, bs_row)
                sign = amount if category == "Operating assets" else amount
                self._record(ledger, report, cell, fy, category, f"Added {label} to {category.lower()}",
                             formula, new_formula, sign,
                             f"{label} ({amount:,.1f}) is {abs(amount) / total:.1%} of {category.lower()} and is an operating item the formula left out.",
                             f"{BS}!{col}{bs_row} (balance sheet line)", "formula_term_added", "medium")

    # ------------------------------------------------------------------ one-time operating items
    def _operating_income(self, wb, cached, ledger, report, fy, col_idx, company_facts, multiplier, filing_hint) -> None:
        ws = wb[INCOME]
        cws = cached[INCOME]
        letter = ws.cell(1, col_idx).column_letter
        year_end = self._year_end(cws, col_idx)
        oi = _num(cws[f"{letter}{ROW_OPERATING_INCOME}"].value)
        if oi in (None, 0) or year_end is None:
            return
        for key, label, kind, tags in ONE_TIME_FAMILIES:
            values = {t: facts_for_year(company_facts, t, year_end) for t in tags}
            raw, used_tags = family_amount(key, {t: v for t, v in values.items() if v is not None})
            if raw == 0:
                continue
            tag = " + ".join(used_tags)
            amount = raw / multiplier
            if abs(amount) / abs(oi) < OI_MATERIALITY:
                continue
            total_cell = ws[f"{letter}{ROW_OPEX_TOTAL}"]
            if not isinstance(_num(total_cell.value), float):
                report.skipped.append(f"{fy} {label}: operating expense total is a formula; not changed.")
                continue
            if kind == "charge":
                delta = -abs(amount)  # remove the charge: operating expenses fall
                placed = self._take_from_components(ws, letter, abs(amount))
                evidence = "operating expense lines contain the charge"
            else:
                gain = amount if raw > 0 else 0.0
                if gain <= 0:
                    continue
                other = _num(ws[f"{letter}{ROW_OPEX_OTHER}"].value) or 0.0
                if other >= 0 or abs(other) < 0.5 * gain:
                    report.skipped.append(f"{fy} {label}: gain {gain:,.1f} not visible in operating expense lines; not changed.")
                    continue
                delta = min(gain, abs(other))
                placed = [(f"{letter}{ROW_OPEX_OTHER}", delta)]
                evidence = "Other operating expense is negative, so the gain sits inside operating expenses"
            if not placed:
                report.skipped.append(f"{fy} {label}: no operating expense line large enough to remove it from; not changed.")
                continue
            total_delta = sum(d for _a, d in placed) if kind == "gain" else -sum(d for _a, d in placed)
            self._record_income(ledger, report, ws, fy, label, kind, placed, total_cell, total_delta, oi, evidence, filing_hint, tag)

    @staticmethod
    def _take_from_components(ws, letter: str, amount: float) -> list[tuple[str, float]]:
        placed: list[tuple[str, float]] = []
        left = amount
        for row in (ROW_OPEX_OTHER, ROW_OPEX_SGA):
            cell = ws[f"{letter}{row}"]
            value = _num(cell.value)
            if value is None or value <= 0 or isinstance(cell.value, str):
                continue
            take = min(value, left)
            placed.append((cell.coordinate, take))
            left -= take
            if left <= 1e-9:
                break
        return placed if left <= 1e-9 else []

    def _record_income(self, ledger, report, ws, fy, label, kind, placed, total_cell, total_delta, oi, evidence, filing_hint, tag) -> None:
        amount = sum(d for _a, d in placed)
        sign_word = "removed from" if kind == "charge" else "added back to"
        reason = (
            f"One-time item: {label} of {amount:,.1f} ({amount / abs(oi):.1%} of operating income) is {sign_word} operating "
            f"expenses so ROIC reflects recurring operations; {evidence}."
        )
        source = f"{filing_hint}, us-gaap:{tag}, fiscal year {fy}"
        cells: list[str] = []
        for addr, part in placed:
            cell = ws[addr]
            old = cell.value
            new = round(old - part, 6) if kind == "charge" else round(old + part, 6)
            cell.value = new
            ledger.record(sheet=INCOME, cell=addr, fiscal_year=fy, category="One-time operating income",
                          what=f"{label}: operating expense line", original=old, new=new, amount=-part if kind == "charge" else part,
                          reason=reason, source=source, method="one_time_item", confidence="medium")
            cells.append(f"{INCOME}!{addr}")
        old_total = total_cell.value
        new_total = round(old_total + total_delta, 6)
        total_cell.value = new_total
        ledger.record(sheet=INCOME, cell=total_cell.coordinate, fiscal_year=fy, category="One-time operating income",
                      what=f"{label}: operating expenses total", original=old_total, new=new_total, amount=total_delta,
                      reason=reason, source=source, method="one_time_item", confidence="medium")
        cells.append(f"{INCOME}!{total_cell.coordinate}")
        report.adjustments.append(RoicAdjustment(fy, "One-time operating income", cells, f"{label}", -total_delta, reason, source))

    def _record(self, ledger, report, cell, fy, category, what, original, new, amount, reason, source, method, confidence) -> None:
        adj = ledger.record(sheet=INPUTS, cell=cell.coordinate, fiscal_year=fy, category=category, what=what, original=original,
                            new=new, amount=amount, reason=reason, source=source, method=method, confidence=confidence)
        cell.value = new
        report.adjustments.append(RoicAdjustment(fy, category, [f"{INPUTS}!{cell.coordinate}"], what, amount, reason, source, adj.adj_id))

    @staticmethod
    def _year_end(cws, col_idx: int) -> date | None:
        for row in range(5, 12):
            value = cws.cell(row, col_idx).value
            if hasattr(value, "date"):
                try:
                    return value.date()
                except (AttributeError, TypeError):
                    pass
            if isinstance(value, date):
                return value
        return None

    @staticmethod
    def _unit_multiplier(ws) -> float:
        text = " ".join(str(ws.cell(r, 1).value or "") for r in range(1, 10)).lower()
        if "thousands" in text:
            return 1_000.0
        return 1_000_000.0
