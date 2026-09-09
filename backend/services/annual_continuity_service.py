"""Historical continuity for Annual Update — previous completed workbook is the historical authority."""

from __future__ import annotations

from copy import copy
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

from models.annual_update import AnnualModelContinuityReport, ContinuityAction, ContinuityEntry
from services.annual_period_service import (
    AnnualPeriodAlignment,
    AnnualPeriodAlignmentError,
    align_fiscal_periods,
    detect_year_columns,
)
from workbook_mapping.sheet_policies import OUTPUT_SHEETS

# Re-export for tests and downstream imports.
__all__ = ["AnnualContinuityService", "detect_year_columns", "align_fiscal_periods"]

_HISTORICAL_SHEETS = (
    "Income - GAAP",
    "Income - GAAP",
    "Balance Sheet - Standardized",
    "Balance Sheet - Standardized",
    "Cash Flow - Standardized",
    "Cash Flow - Standardized",
    "Inputs",
    "Tax",
    "R&D",
    "Leases",
    "IC & NOPAT & ROIC ",
    "IC & NOPAT & ROIC ",
    "Expected Returns & Buybacks",
    "Enterprise Value",
)

_FORMULA_ONLY_SHEETS = OUTPUT_SHEETS | {
    "All Ratios",
    "All Ratios",
    "Final Metrics",
    "Final Metrics",
    "IS%",
    "BS%",
    "CF%",
    "FCF",
}
_CURRENT_DATA_CELLS = {f"B{r}" for r in range(63, 76)}
_MAX_ROW = 140
_MAX_COL = 16


def _cell_kind(value: Any) -> str:
    if value is None or value == "":
        return "blank"
    if isinstance(value, str) and value.startswith("="):
        return "formula"
    if hasattr(value, "text"):
        return "formula"
    return "value"


def _formula_text(value: Any) -> str | None:
    if isinstance(value, str) and value.startswith("="):
        return value
    if hasattr(value, "text"):
        return str(value.text)
    return None


def _values_equal(a: Any, b: Any) -> bool:
    fa, fb = _formula_text(a), _formula_text(b)
    if fa is not None or fb is not None:
        return fa is not None and fb is not None and fa == fb
    if isinstance(a, (int, float)) and isinstance(b, (int, float)):
        return abs(float(a) - float(b)) < 1e-9
    return a == b


def _copy_style(src, dst) -> bool:
    if not src.has_style:
        return False
    dst.font = copy(src.font)
    dst.border = copy(src.border)
    dst.fill = copy(src.fill)
    dst.number_format = src.number_format
    dst.protection = copy(src.protection)
    dst.alignment = copy(src.alignment)
    return True


class AnnualContinuityService:
    """Reconcile persistent historical regions from the previous completed workbook."""

    def apply(
        self,
        *,
        analysis_id: str,
        ticker: str,
        template_path: Path,
        previous_workbook_path: Path,
        workbook_path: Path,
        new_fiscal_year: str | None = None,
        period_alignment: AnnualPeriodAlignment | None = None,
    ) -> AnnualModelContinuityReport:
        if period_alignment is None:
            period_alignment = align_fiscal_periods(template_path, previous_workbook_path)
        detected_new = new_fiscal_year or period_alignment.new_fiscal_year

        prev = load_workbook(previous_workbook_path, data_only=False)
        out = load_workbook(workbook_path, data_only=False)
        entries: list[ContinuityEntry] = []
        try:
            template_fy_cols = self._workbook_fy_columns(out)
            previous_fy_cols = self._workbook_fy_columns(prev)
            sheet_set = list(dict.fromkeys(list(_HISTORICAL_SHEETS) + list(_FORMULA_ONLY_SHEETS)))
            for sheet in sheet_set:
                if sheet not in prev.sheetnames or sheet not in out.sheetnames:
                    continue
                formula_guard = sheet in _FORMULA_ONLY_SHEETS
                pws, ows = prev[sheet], out[sheet]
                cols = detect_year_columns(ows, out)
                prev_cols = detect_year_columns(pws, prev)
                fy_cols = {fy: col for fy, col in cols.items() if fy.startswith("FY")}
                if not fy_cols:
                    fy_cols = {fy: template_fy_cols[fy] for fy in period_alignment.overlap_years if fy in template_fy_cols}
                if not fy_cols:
                    continue
                max_row = min(max(pws.max_row or 1, ows.max_row or 1, 1), _MAX_ROW)

                for row in range(1, max_row + 1):
                    for fy, col in fy_cols.items():
                        addr = f"{get_column_letter(col)}{row}"
                        if sheet == "Inputs" and addr in _CURRENT_DATA_CELLS:
                            entries.append(
                                ContinuityEntry(
                                    sheet=sheet,
                                    cell=addr,
                                    action=ContinuityAction.NEW_FY_REFRESH,
                                    reason="Current-data cell — refresh from live/CRF, do not carry stale market values.",
                                    previous_type=_cell_kind(pws[addr].value),
                                    template_type=_cell_kind(ows[addr].value),
                                )
                            )
                            continue

                        if fy == detected_new:
                            entries.append(
                                ContinuityEntry(
                                    sheet=sheet,
                                    cell=addr,
                                    fiscal_year=fy,
                                    action=ContinuityAction.NEW_FY_REFRESH,
                                    reason="Newly added fiscal year — not carried from previous workbook.",
                                    previous_type=_cell_kind(pws.cell(row, prev_cols.get(fy, col)).value),
                                    template_type=_cell_kind(ows[addr].value),
                                )
                            )
                            continue

                        prev_col = prev_cols.get(fy) or previous_fy_cols.get(fy)
                        if prev_col is None:
                            continue
                        prev_cell = pws.cell(row, prev_col)
                        out_cell = ows.cell(row, col)
                        entry = self._reconcile(
                            sheet=sheet,
                            addr=addr,
                            fy=fy,
                            prev_cell=prev_cell,
                            out_cell=out_cell,
                            formula_guard=formula_guard,
                            new_fiscal_year=detected_new,
                        )
                        if entry is not None:
                            entries.append(entry)

            out.save(workbook_path)
        finally:
            prev.close()
            out.close()

        counts: dict[str, int] = {}
        for e in entries:
            counts[e.action.value] = counts.get(e.action.value, 0) + 1
        return AnnualModelContinuityReport(
            analysis_id=analysis_id,
            ticker=ticker,
            phase="apply",
            new_fiscal_year=detected_new,
            entries=[e for e in entries if e.action != ContinuityAction.MATCH_ALREADY][:8000],
            action_counts=counts,
            summary=(
                f"Continuity apply: new_fy={detected_new}; overlap={len(period_alignment.overlap_years)}; "
                + ", ".join(f"{k}={v}" for k, v in sorted(counts.items()))
            ),
        )

    def verify_final(
        self,
        *,
        analysis_id: str,
        ticker: str,
        previous_workbook_path: Path,
        workbook_path: Path,
        new_fiscal_year: str | None = None,
    ) -> AnnualModelContinuityReport:
        """Sparse final snapshot: unauthorized historical mismatches only."""
        prev = load_workbook(previous_workbook_path, data_only=False)
        out = load_workbook(workbook_path, data_only=False)
        mismatches: list[ContinuityEntry] = []
        try:
            year_map = {}
            for sheet in _HISTORICAL_SHEETS:
                if sheet in out.sheetnames:
                    year_map = detect_year_columns(out[sheet], out)
                    if any(k.startswith("FY") for k in year_map):
                        break
            fy_keys = [k for k in year_map if k.startswith("FY")]
            detected_new = new_fiscal_year or (max(fy_keys) if fy_keys else None)

            for sheet in _HISTORICAL_SHEETS:
                if sheet not in prev.sheetnames or sheet not in out.sheetnames:
                    continue
                pws, ows = prev[sheet], out[sheet]
                cols = detect_year_columns(ows, out)
                prev_cols = detect_year_columns(pws, prev)
                fy_cols = {fy: col for fy, col in cols.items() if fy.startswith("FY")}
                max_row = min(max(pws.max_row or 1, 1), _MAX_ROW)
                for row in range(1, max_row + 1):
                    for fy, col in fy_cols.items():
                        if fy == detected_new:
                            continue
                        addr = f"{get_column_letter(col)}{row}"
                        if sheet == "Inputs" and addr in _CURRENT_DATA_CELLS:
                            continue
                        prev_col = prev_cols.get(fy, col)
                        pv, ov = pws.cell(row, prev_col).value, ows.cell(row, col).value
                        pk, ok = _cell_kind(pv), _cell_kind(ov)
                        if pk == "blank" and ok == "blank":
                            continue
                        if pk == "formula" and ok == "formula":
                            continue
                        if pk == "formula" and ok != "formula":
                            mismatches.append(
                                ContinuityEntry(
                                    sheet=sheet,
                                    cell=addr,
                                    fiscal_year=fy,
                                    action=ContinuityAction.REVIEW_REQUIRED,
                                    reason="Historical formula flattened in final workbook.",
                                    previous_type=pk,
                                    final_type=ok,
                                )
                            )
                        elif pk == "value" and ok == "value" and not _values_equal(pv, ov):
                            mismatches.append(
                                ContinuityEntry(
                                    sheet=sheet,
                                    cell=addr,
                                    fiscal_year=fy,
                                    action=ContinuityAction.REVIEW_REQUIRED,
                                    reason="Historical value differs from previous workbook without restatement action.",
                                    previous_value=pv,
                                    final_value=ov,
                                    previous_type=pk,
                                    final_type=ok,
                                )
                            )
        finally:
            prev.close()
            out.close()

        status = "ok" if not mismatches else "REVIEW_REQUIRED"
        return AnnualModelContinuityReport(
            analysis_id=analysis_id,
            ticker=ticker,
            phase="final",
            new_fiscal_year=detected_new,
            entries=mismatches[:200],
            unauthorized_historical_rewrites=len(mismatches),
            status=status,
            summary=f"Final continuity: unauthorized_historical_diffs={len(mismatches)}.",
        )

    def _reconcile(
        self,
        *,
        sheet,
        addr,
        fy,
        prev_cell,
        out_cell,
        formula_guard,
        new_fiscal_year: str | None = None,
    ) -> ContinuityEntry | None:
        if new_fiscal_year and fy == new_fiscal_year:
            return None
        pv, ov = prev_cell.value, out_cell.value
        pk, ok = _cell_kind(pv), _cell_kind(ov)
        if pk == "blank" and ok == "blank":
            return None

        if formula_guard:
            if ok == "formula":
                return ContinuityEntry(
                    sheet=sheet,
                    cell=addr,
                    fiscal_year=fy,
                    action=ContinuityAction.KEEP_NEW_FORMULA,
                    reason="Formula-driven output sheet — formulas left intact.",
                    previous_type=pk,
                    template_type=ok,
                    final_type=ok,
                )
            return ContinuityEntry(
                sheet=sheet,
                cell=addr,
                fiscal_year=fy,
                action=ContinuityAction.BLOCKED,
                reason="Blocked write on formula-driven Ratios/Final Metrics cell.",
                previous_type=pk,
                template_type=ok,
            )

        action = ContinuityAction.MATCH_ALREADY
        reason = "Already aligned."
        fmt = False
        changed = False

        if pk == "formula" and ok != "formula":
            out_cell.value = pv
            action = ContinuityAction.RESTORE_FORMULA
            reason = "Previous workbook formula restored on historical cell."
            changed = True
        elif pk == "formula" and ok == "formula":
            prev_f, out_f = _formula_text(pv), _formula_text(ov)
            if prev_f != out_f:
                action = ContinuityAction.KEEP_NEW_FORMULA
                reason = "New-template formula retained (deliberate template update)."
            else:
                action = ContinuityAction.MATCH_ALREADY
                reason = "Historical formula already matches."
        elif pk == "value" and ok == "formula":
            action = ContinuityAction.KEEP_NEW_FORMULA
            reason = "Template formula preserved; not flattened to prior value."
        elif pk == "value" and (ok == "blank" or (ok == "value" and not _values_equal(pv, ov))):
            out_cell.value = pv
            action = ContinuityAction.CARRY_FORWARD_VALUE
            reason = "Analyst-entered historical value carried from previous workbook."
            changed = True

        fmt = _copy_style(prev_cell, out_cell)
        if fmt and not changed and action == ContinuityAction.MATCH_ALREADY:
            action = ContinuityAction.CARRY_FORWARD_FORMAT
            reason = "Historical formatting copied from previous workbook."

        if action == ContinuityAction.MATCH_ALREADY and not fmt and not changed:
            return None

        return ContinuityEntry(
            sheet=sheet,
            cell=addr,
            fiscal_year=fy,
            previous_value="(formula)" if pk == "formula" else pv,
            previous_type=pk,
            template_value="(formula)" if ok == "formula" else ov,
            template_type=ok,
            final_value="(formula)" if _cell_kind(out_cell.value) == "formula" else out_cell.value,
            final_type=_cell_kind(out_cell.value),
            action=action,
            reason=reason,
            formatting_copied=fmt,
        )

    @staticmethod
    def _workbook_fy_columns(wb) -> dict[str, int]:
        for sheet in (
            "Balance Sheet - Standardized",
            "Income - GAAP",
            "Cash Flow - Standardized",
            "Inputs",
        ):
            if sheet not in wb.sheetnames:
                continue
            cols = detect_year_columns(wb[sheet], wb)
            fy_cols = {fy: col for fy, col in cols.items() if fy.startswith("FY")}
            if fy_cols:
                return fy_cols
        return {}
