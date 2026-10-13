"""Apply the corrections for Expected Return, Enterprise Value and Graham inputs that the review finds unrealistic.

Used by annual updates, quarter updates and new company analyses, before the valuation judgment step. The rule everywhere: when a model
gives a value that looks inaccurate because of noise in the data, look at the model's INPUTS first, and change the model only when the
inputs cannot be repaired.

Expected Return (tab "Expected Returns & Buybacks")
  The template projects ten years of earnings from book value growth (retention x ROE). When its return (price plus dividends, F14) is
  unrealistic (above 20% or negative):
  1. Look at the two inputs: the average ROE (A14) and the average retained earnings (C5, from the Inputs tab). When some years of one of
     them are clearly off, it is recalculated without those years. If that makes the return realistic, it stays.
  2. When the whole history is off (a severely cyclical business) the analyst's second model is used: ten years of EPS projected from the
     EPS growth rate (A14 becomes current EPS / book value per share, A11 the EPS growth rate). That is done only when EPS itself follows a
     steady trend. The EPS growth rate is the robust one when extraordinary years skew the first-to-last-year rate.
  3. When EPS is cyclical too, the model is kept as it is and every observation is written as a red flag.
Enterprise Value (tab "Enterprise Value")
  The annualized owner's earnings growth (C6) is checked. When it is off, two corrected rates are computed and compared in the notes: the
  average of the first three and last three years (Final Metrics M10) and the log-linear model through all years (N10). The one that looks
  more accurate is recommended and used.
Graham entry price (tab "Enterprise Value", EPS 10 Year CAGR in B41, from Final Metrics L31)
  The EPS growth rate is checked for being skewed by extraordinary EPS years. When it is, B41 is replaced by the rate of the other years.

The original formula or value of every replaced cell is kept in the HAP Adjustments tab. Notes are plain language (see tab_notes).
"""

from __future__ import annotations

import statistics
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from services.adjustment_ledger_service import AdjustmentLedger
from services.annual_period_service import detect_year_columns
from services.tab_notes import note
from services.valuation_history import (
    MAX_OFF_SHARE,
    MIN_YEARS,
    TREND_R2,
    eps_trend,
    loglinear,
    off_years,
    review_eps_growth,
    trimmed_mean,
)
from services.valuation_inputs_review_service import (
    EV_SHEET,
    ER_SHEET,
    FM_SHEET,
    ER_HIGH,
    ER_LOW,
    OE_HIGH,
    OE_LOW,
    AppliedCorrection,
    ReviewFinding,
    ValuationInputsReview,
    ValuationInputsReviewService,
    _pct,
)

EPS_GROWTH_LIMIT = 0.25       # an EPS growth rate beyond this (either way) is not projected for ten years
OE_CLEARLY_OFF = 0.25         # annualized owner's earnings growth at or above this is clearly off
OE_MATERIAL = 0.005           # a corrected growth rate must differ by at least half a point to matter
OE_AGREE = 0.01               # two corrected growth rates within a point of each other agree
ER_SOURCE = "the company's EPS history (Final Metrics tab)"
ER_INPUTS_SOURCE = "the Final Metrics and Inputs tab history"
EV_SOURCE = "the Final Metrics history (owner's earnings) of the company"
GRAHAM_SOURCE = "the company's EPS history (Final Metrics tab)"
INPUTS_SHEET = "Inputs"
RETENTION_LABEL = "change in retained earnings per share/ep"


def _cell(ws, addr: str) -> float | None:
    v = ws[addr].value
    return float(v) if isinstance(v, (int, float)) and not isinstance(v, bool) else None


def _years_text(years: list[str]) -> str:
    return ", ".join(years)


class ValuationCorrectionService:
    def apply(self, workbook_path: Path, company_facts: dict[str, Any] | None = None, review: ValuationInputsReview | None = None) -> ValuationInputsReview:
        path = Path(workbook_path)
        review = review or ValuationInputsReviewService().review(path, company_facts)
        wb = load_workbook(path, data_only=False)
        values = load_workbook(path, data_only=True)
        changed = False
        try:
            ledger = AdjustmentLedger(wb)
            hist = self._histories(wb, values)
            if ER_SHEET in wb.sheetnames and ER_SHEET in values.sheetnames:
                changed |= self._expected_return(wb[ER_SHEET], values[ER_SHEET], ledger, review, hist)
            if EV_SHEET in wb.sheetnames:
                changed |= self._owner_earnings(wb[EV_SHEET], ledger, review)
                if EV_SHEET in values.sheetnames:
                    changed |= self._graham(wb[EV_SHEET], values[EV_SHEET], ledger, review, hist)
            if changed:
                wb.save(path)
        finally:
            wb.close()
            values.close()
        return review

    # ------------------------------------------------------------------ the histories behind the inputs
    @staticmethod
    def _series(wb, values, sheet: str, match) -> dict[str, float]:
        """A yearly row of a tab, found by its label, as {FY label: value} in year order."""
        if sheet not in wb.sheetnames or sheet not in values.sheetnames:
            return {}
        ws = values[sheet]
        row = next((r for r in range(1, min(ws.max_row or 1, 200) + 1) if match(str(ws.cell(r, 1).value or "").strip().lower())), None)
        if not row:
            return {}
        try:
            cols = {k: c for k, c in detect_year_columns(wb[sheet], wb).items() if str(k).startswith("FY")}
        except Exception:  # noqa: BLE001 - a workbook without readable year headers simply has no history here
            return {}
        out: dict[str, float] = {}
        for fy, col in sorted(cols.items(), key=lambda kv: kv[1]):
            v = ws.cell(row, col).value
            if isinstance(v, (int, float)) and not isinstance(v, bool):
                out[fy] = float(v)
        return out

    def _histories(self, wb, values) -> dict[str, dict[str, float]]:
        return {
            "roe": self._series(wb, values, FM_SHEET, lambda s: s == "roe"),
            "eps": self._series(wb, values, FM_SHEET, lambda s: s == "eps"),
            "retention": self._series(wb, values, INPUTS_SHEET, lambda s: s.startswith(RETENTION_LABEL)),
        }

    # ------------------------------------------------------------------ expected return
    @staticmethod
    def _model_return(price, pe, bv0, growth, roe, payout) -> float | None:
        """The template's price-plus-dividends return for given growth (A11) and ROE (A14): the same arithmetic as cells B11 to F14."""
        if not price or price <= 0 or None in (pe, bv0, growth, roe):
            return None
        eps = [bv0 * (1 + growth) ** t * roe for t in range(1, 11)]
        total = eps[-1] * pe + (payout or 0.0) * sum(eps)
        return (total / price) ** 0.1 - 1 if total > 0 else None

    def _expected_return(self, ws, ws_values, ledger, review, hist) -> bool:
        review.findings = [f for f in review.findings if f.topic != "expected_return"]         # the earlier ROE reading is replaced by this check
        review.alternatives.pop("expected_return", None)
        model = _cell(ws_values, "F14") if _cell(ws_values, "F14") is not None else _cell(ws_values, "E14")
        if model is None:
            review.skipped.append("Expected return: the model's return is not calculated in this workbook.")
            return False
        if ER_LOW <= model <= ER_HIGH:
            review.findings.append(ReviewFinding(
                "expected_return", "model_return", "reasonable",
                f"The expected return of {_pct(model)} is within a realistic range of {_pct(ER_LOW)} to {_pct(ER_HIGH)}.", {"model": model},
            ))
            return False
        side = "too high" if model > ER_HIGH else "negative"
        price, pe, bv0 = _cell(ws_values, "A2"), _cell(ws_values, "E2"), _cell(ws_values, "B8")
        growth, eps0 = _cell(ws_values, "B5"), _cell(ws_values, "C30")
        d17, e17 = _cell(ws_values, "D17"), _cell(ws_values, "E17")
        payout = (e17 / d17) if d17 and e17 is not None else 0.0

        def keep(reason: str) -> bool:
            review.findings.append(ReviewFinding("expected_return", "model_return", "flag", f"The expected return of {_pct(model)} is {side}, but {reason}.", {"model": model}))
            review.left_as_is["expected_return"] = f"the EPS-growth model could not be used because {reason}"
            return False

        if None in (price, pe, bv0, growth, eps0) or not bv0 or bv0 <= 0:
            return keep("the EPS history and book value needed for it are not in the workbook")
        if eps0 <= 0:
            return keep("current EPS is not positive")

        # 1. the inputs of the model: the average ROE (A14) and the average retained earnings (C5)
        repaired, tried = self._repair_inputs(ws, ws_values, ledger, review, hist, model=model, side=side, price=price, pe=pe, bv0=bv0, payout=payout)
        if repaired:
            return True

        # 2. the other model, only when EPS follows a steady trend; 3. otherwise the model stays and the observations are red flags
        eps_hist = hist.get("eps") or {}
        trend = eps_trend(eps_hist) if eps_hist else None
        if trend is not None and not trend.reasonable and len(eps_hist) >= MIN_YEARS:
            return self._red_flags(ws, ledger, review, hist, model=model, side=side, trend_reason=trend.reason, tried=tried)

        if abs(growth) > EPS_GROWTH_LIMIT and not (eps_hist and review_eps_growth(eps_hist, growth).skewed):
            return keep(f"the EPS growth rate of {_pct(growth)} a year is too extreme to project for ten years")
        used_growth, growth_note = growth, ""
        if eps_hist:
            rg = review_eps_growth(eps_hist, growth)
            if rg.skewed and rg.robust is not None:
                used_growth = rg.robust
                growth_note = (f" (the rate of the years without extraordinary EPS: {_years_text(rg.extraordinary)} are left out; "
                               f"the first-to-last-year rate was {_pct(growth)})")
        if abs(used_growth) > EPS_GROWTH_LIMIT:
            return keep(f"the EPS growth rate of {_pct(used_growth)} a year is too extreme to project for ten years")
        new_roe = round(eps0 / bv0, 4)
        new_growth = round(used_growth, 4)
        new_return = self._model_return(price, pe, bv0, new_growth, new_roe, payout)
        if new_return is None:
            return keep("it gives no meaningful result")
        what = (
            f"The expected return model was switched to ten years of EPS projected from EPS growth of {_pct(used_growth)} a year{growth_note}, "
            f"starting from ${eps0:,.2f}; the return is now {_pct(new_return)} (it was {_pct(model)})"
        )
        why = (
            f"the original model, which grows book value by retention times return on equity, gave {_pct(model)}, which is {side} to rely on "
            f"(a realistic range is {_pct(ER_LOW)} to {_pct(ER_HIGH)})"
        )
        extras = review.extra_notes.setdefault("expected_return", [])           # separate short notes: one note holds at most 400 characters
        if tried:
            extras.append(note(tried[0].upper() + tried[1:].rstrip("."), None, ER_INPUTS_SOURCE, because=False))
        if trend is not None and trend.reasonable:
            extras.append(note(f"The EPS-growth model was chosen because {trend.reason}", None, ER_SOURCE, because=False))
        for cell, label, new_value, orig in (
            ("A14", "Average return on equity (now current EPS / book value per share)", new_roe, ws["A14"].value),
            ("A11", "Book value per share growth (now the EPS growth rate)", new_growth, ws["A11"].value),
        ):
            ws[cell].value = new_value
            ws[cell].number_format = "0.0%"
            ledger.record(sheet=ER_SHEET, cell=cell, fiscal_year=None, category="Expected return input", what=label, original=orig, new=new_value,
                          amount=None, reason=f"{what}, because {why}.", source=ER_SOURCE, method="eps_growth_model", confidence="medium")
            review.applied.append(AppliedCorrection("expected_return", f"{ER_SHEET}!{cell}", orig, new_value, "replaced", what, why, ER_SOURCE))
        review.findings.append(ReviewFinding("expected_return", "model_return", "flag", why[0].upper() + why[1:] + ".", {"model": model, "adjusted": new_return}))
        return True

    def _repair_inputs(self, ws, ws_values, ledger, review, hist, *, model, side, price, pe, bv0, payout) -> tuple[bool, str]:
        """Recalculate A14 and/or C5 without the years that are clearly off. Returns (done, a sentence about what was tried)."""
        roe_h, ret_h = hist.get("roe") or {}, hist.get("retention") or {}
        avg_roe, retention = _cell(ws_values, "A14"), _cell(ws_values, "C5")
        if avg_roe is None or retention is None or str(ws["A11"].value or "").replace(" ", "").upper() != "=C5*A14":
            return False, ""
        bad_roe, bad_ret = off_years(roe_h), off_years(ret_h)
        if not bad_roe and not bad_ret:
            return False, ""
        if len(bad_roe) > MAX_OFF_SHARE * len(roe_h) or len(bad_ret) > MAX_OFF_SHARE * len(ret_h):
            return False, ""                       # not a few years: the whole history is noisy
        new_roe = trimmed_mean(roe_h, bad_roe) if bad_roe else avg_roe
        new_ret = trimmed_mean(ret_h, bad_ret) if bad_ret else retention
        if new_roe is None or new_ret is None:
            return False, ""
        new_return = self._model_return(price, pe, bv0, new_ret * new_roe, new_roe, payout)
        parts = []
        if bad_roe:
            parts.append(f"the average return on equity without {_years_text(bad_roe)} is {_pct(new_roe)} (it was {_pct(avg_roe)})")
        if bad_ret:
            parts.append(f"the average retained earnings without {_years_text(bad_ret)} is {_pct(new_ret)} (it was {_pct(retention)})")
        recalculated = " and ".join(parts)
        if new_return is None or not (ER_LOW <= new_return <= ER_HIGH):
            outcome = "no meaningful result" if new_return is None else f"a return of {_pct(new_return)}, still {'too high' if new_return > ER_HIGH else 'negative'}"
            return False, f"the model's own inputs were recalculated first ({recalculated}), but that still gives {outcome}"
        what = f"The model's inputs were recalculated without the years that were clearly off: {recalculated}; the return is now {_pct(new_return)} (it was {_pct(model)})"
        why = (
            f"the return of {_pct(model)} was {side} to rely on (a realistic range is {_pct(ER_LOW)} to {_pct(ER_HIGH)}), and it rested on "
            f"{'return on equity' if bad_roe else ''}{' and ' if bad_roe and bad_ret else ''}{'retained earnings' if bad_ret else ''} "
            f"distorted by a few unusual years, so the model itself was kept"
        )
        for cell, label, use, new_value, orig in (
            ("A14", "Average return on equity (recalculated without the years that are off)", bool(bad_roe), round(new_roe, 4), ws["A14"].value),
            ("C5", "Average retained earnings (recalculated without the years that are off)", bool(bad_ret), round(new_ret, 4), ws["C5"].value),
        ):
            if not use:
                continue
            ws[cell].value = new_value
            ws[cell].number_format = "0.0%"
            ledger.record(sheet=ER_SHEET, cell=cell, fiscal_year=None, category="Expected return input", what=label, original=orig, new=new_value,
                          amount=None, reason=f"{what}, because {why}.", source=ER_INPUTS_SOURCE, method="normalized_input", confidence="medium")
            review.applied.append(AppliedCorrection("expected_return", f"{ER_SHEET}!{cell}", orig, new_value, "replaced", what, why, ER_INPUTS_SOURCE))
        review.findings.append(ReviewFinding("expected_return", "model_return", "flag", why[0].upper() + why[1:] + ".", {"model": model, "adjusted": new_return}))
        return True, ""

    def _red_flags(self, ws, ledger, review, hist, *, model, side, trend_reason, tried) -> bool:
        """The history is cyclical and EPS is too: keep the model as it is and write every observation as a red flag."""
        roe_h, ret_h = hist.get("roe") or {}, hist.get("retention") or {}
        flags = [f"The expected return of {_pct(model)} is {side} to rely on (a realistic range is {_pct(ER_LOW)} to {_pct(ER_HIGH)})"]
        if roe_h:
            flags.append(f"Return on equity ranged from {_pct(min(roe_h.values()))} to {_pct(max(roe_h.values()))} over {len(roe_h)} years, so its average is not a stable base")
        if ret_h:
            flags.append(f"Retained earnings ranged from {_pct(min(ret_h.values()))} to {_pct(max(ret_h.values()))} of EPS, so growth from retention is not stable either")
        flags.append(f"The EPS-growth model is not a better base because {trend_reason}")
        if tried:
            flags.append(tried[0].upper() + tried[1:])
        text = ". ".join(f.rstrip(".") for f in flags) + "."
        for f in flags:
            review.findings.append(ReviewFinding("expected_return", "red_flag", "flag", f.rstrip(".") + ".", {"model": model}))
        review.left_as_is["expected_return"] = "the business looks cyclical in return on equity, retention and EPS, so no other model is better; the red flags are for your judgment"
        ledger.record(sheet=ER_SHEET, cell="F14", fiscal_year=None, category="Expected return input", what="Expected return left as the model gives it (red flags)",
                      original=_pct(model), new="unchanged", amount=None, reason=text, source=ER_INPUTS_SOURCE, method="flag_only",
                      confidence="low", mark_cell=False, status="Red flag")
        return True

    # ------------------------------------------------------------------ owner's earnings
    def _owner_earnings(self, ws, ledger, review) -> bool:
        finding = next((f for f in review.findings if f.topic == "owner_earnings"), None)
        if finding is None or finding.verdict != "flag":
            return False
        annual = finding.values.get("annualized_growth")
        series: dict[str, float] = finding.values.get("series") or {}
        alt = finding.values.get("three_year_average_cagr")
        if annual is None or alt is None or len(series) < 6:
            review.left_as_is["owner_earnings"] = "The flagged rate was left as it is because there is not enough history to compute a corrected one"
            return False
        # the two ways to correct the rate: the 3-year averages (Final Metrics M10) and the log-linear model (N10)
        fit = loglinear(series)
        g3 = alt
        gll = fit.growth if fit else None
        if fit is not None and fit.r2 >= TREND_R2:
            chosen, chosen_name = gll, "log-linear model"
            recommendation = (f"The log-linear model is recommended because owner's earnings follow a steady trend (the fit explains {fit.r2:.0%} of their movement) "
                              f"and it uses all {len(series)} years instead of six")
        else:
            chosen, chosen_name = g3, "average of the first three and last three years"
            recommendation = ("The average of the first three and last three years is recommended because "
                              + (f"owner's earnings are irregular (a straight-line fit explains only {fit.r2:.0%} of their movement)" if fit
                                 else "owner's earnings include zero or negative years, which a log-linear model cannot use"))
        numbers = (
            f"Two corrected owner's earnings growth rates were compared: the average of the first three and last three years (Final Metrics M10) gives {_pct(g3)} a year"
            + (f", and the log-linear model through all {len(series)} years (Final Metrics N10) gives {_pct(gll)} a year" if gll is not None else "")
            + (" and they agree within a point" if gll is not None and abs(g3 - gll) <= OE_AGREE else "")
        )
        review.extra_notes.setdefault("owner_earnings", []).extend([
            note(numbers, None, EV_SOURCE, because=False),
            note(recommendation, None, None, because=False),
        ])
        corrected = round(min(max(chosen, OE_LOW), OE_HIGH), 4)
        if abs(corrected - annual) < OE_MATERIAL:
            review.left_as_is["owner_earnings"] = "The flagged rate was left as it is because the corrected rate is almost the same"
            return False
        years = list(series)
        typical = statistics.median(series.values())
        base_low = series[years[0]] < 0.6 * typical
        clearly_off = annual >= OE_CLEARLY_OFF or annual < 0 or (annual > OE_HIGH and base_low)
        if annual > OE_HIGH:
            reason = f"{_pct(annual)} a year is very high to project forward"
        elif annual < 0:
            reason = f"{_pct(annual)} a year is negative"
        else:
            reason = f"{_pct(annual)} a year is unrealistically low"
        if base_low and annual > OE_HIGH:
            reason += f", and it rests on {years[0]}, when owner's earnings were only {series[years[0]]:,.0f} against a typical {typical:,.0f}"
        clamp = f" (limited to {_pct(OE_HIGH if chosen > OE_HIGH else OE_LOW)})" if corrected != round(chosen, 4) else ""
        if clearly_off:
            original = ws["C6"].value
            what = f"Owner's earnings growth was replaced with {_pct(corrected)} a year{clamp} (it was {_pct(annual)})"
            why = f"{reason}; the corrected rate is from the {chosen_name}, the more accurate of the two methods compared"
            ws["C6"].value = corrected
            ws["C6"].number_format = "0.0%"
            ledger.record(sheet=EV_SHEET, cell="C6", fiscal_year=None, category="Enterprise value input", what="Owner's earnings growth (annualized)",
                          original=original, new=corrected, amount=None, reason=f"{what}, because {why}.", source=EV_SOURCE,
                          method="normalized_input", confidence="medium")
            review.applied.append(AppliedCorrection("owner_earnings", f"{EV_SHEET}!C6", original, corrected, "replaced", what, why, EV_SOURCE))
            return True
        from services.annual_judgment_service import AnnualJudgmentService

        written = AnnualJudgmentService._clone_ev_projections(ws, corrected)
        if not written:
            return False
        what = f"A parallel calculation with owner's earnings growth of {_pct(corrected)} a year{clamp} was added below the original (rows 54 to 57)"
        why = f"{reason}, but the original is not clearly wrong, so both are shown; the corrected rate is from the {chosen_name}"
        ledger.record(sheet=EV_SHEET, cell="A54", fiscal_year=None, category="Enterprise value input", what="Parallel owner's earnings projection",
                      original="not present", new=f"{_pct(corrected)} growth", amount=None, reason=f"{what}, because {why}.", source=EV_SOURCE,
                      method="normalized_input", confidence="medium", mark_cell=False)
        review.applied.append(AppliedCorrection("owner_earnings", f"{EV_SHEET}!A54", None, corrected, "parallel", what, why, EV_SOURCE))
        return True

    # ------------------------------------------------------------------ Graham: the EPS growth rate (EV tab B41)
    def _graham(self, ws, ws_values, ledger, review, hist) -> bool:
        eps_hist = hist.get("eps") or {}
        endpoint = _cell(ws_values, "B41")
        if endpoint is None or len(eps_hist) < MIN_YEARS:
            return False
        rg = review_eps_growth(eps_hist, endpoint)
        if not rg.skewed or rg.robust is None:
            review.findings.append(ReviewFinding("graham", "eps_growth", "reasonable",
                                                 f"The EPS growth rate of {_pct(endpoint)} a year used for the Graham entry price is not skewed by extraordinary years.",
                                                 {"endpoint": endpoint, "robust": rg.robust}))
            return False
        corrected = round(rg.robust, 4)
        original = ws["B41"].value
        what = f"The EPS growth rate used for the Graham entry price was replaced with {_pct(corrected)} a year (it was {_pct(endpoint)})"
        why = (f"EPS was extraordinary in {_years_text(rg.extraordinary)}, which skews the first-to-last-year rate; "
               f"the corrected rate is the steady growth of the other years (the fit explains {rg.r2:.0%} of their movement)"
               if rg.r2 is not None else f"EPS was extraordinary in {_years_text(rg.extraordinary)}, which skews the first-to-last-year rate")
        ws["B41"].value = corrected
        ws["B41"].number_format = "0.0%"
        ledger.record(sheet=EV_SHEET, cell="B41", fiscal_year=None, category="Enterprise value input", what="EPS growth used for the Graham entry price",
                      original=original, new=corrected, amount=None, reason=f"{what}, because {why}.", source=GRAHAM_SOURCE,
                      method="normalized_input", confidence="medium")
        review.applied.append(AppliedCorrection("graham", f"{EV_SHEET}!B41", original, corrected, "replaced", what, why, GRAHAM_SOURCE))
        review.findings.append(ReviewFinding("graham", "eps_growth", "flag", why[0].upper() + why[1:] + ".", {"endpoint": endpoint, "robust": rg.robust}))
        return True
