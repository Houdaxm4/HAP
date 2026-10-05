"""Leases tab: replace unrealistic values in the "Estimated Long-Term Rate" row.

The template computes the rate for each year as interest expense / total debt. That is blank or N/A when the company has no debt
and is often too high or too low when interest includes non-lease items. An analyst then types in a number they can justify.
This service does the same and flags every cell it changes.

Order of choice for each bad year:
1. The company's own weighted-average discount rate for that year (10-K lease note), if realistic.
2. The median of the nearest realistic years in the same row (matches earlier years).
3. A business-based estimate: 4-6% for investment-grade-like companies, 6-9% otherwise.

It reads cached values, so the workbook must have been recalculated first, and it needs another recalculation afterwards.
"""

from __future__ import annotations

import statistics
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from services.annual_period_service import detect_year_columns
from services.adjustment_ledger_service import AdjustmentLedger

RATE_MAX = 0.10
RATE_MIN = 0.01
IG_RANGE = (0.04, 0.06)
NON_IG_RANGE = (0.06, 0.09)
IG_RATE = 0.05
NON_IG_RATE = 0.075
STRESSED_RATE = 0.085
_RATE_LABELS = ("estimated long-term rate", "long-term rate", "long term rate", "lease rate")
NOTES_START_ROW = 23


@dataclass
class LeaseRateFix:
    fiscal_year: str
    cell: str
    original: Any
    new: float
    method: str
    reason: str
    source: str


@dataclass
class LeaseRowReport:
    checked: int = 0
    fixes: list[LeaseRateFix] = field(default_factory=list)
    credit_profile: dict[str, Any] = field(default_factory=dict)
    skipped: list[str] = field(default_factory=list)
    notes_rows: list[int] = field(default_factory=list)

    @property
    def changed(self) -> bool:
        return bool(self.fixes)


def is_realistic(value: Any) -> bool:
    return isinstance(value, (int, float)) and not isinstance(value, bool) and RATE_MIN <= float(value) <= RATE_MAX


def _num(value: Any) -> float | None:
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        return float(value)
    return None


def _as_rate(value: float) -> float:
    return value / 100.0 if abs(value) > 1.5 else value


def credit_profile(inputs_cached: dict[str, float | None]) -> dict[str, Any]:
    """Investment-grade-like or not, from the latest year's debt, interest and EBIT."""
    debt = inputs_cached.get("debt")
    interest = inputs_cached.get("interest")
    ebit = inputs_cached.get("ebit")
    profile: dict[str, Any] = {"debt": debt, "interest": interest, "ebit": ebit}
    coverage = ebit / interest if ebit is not None and interest not in (None, 0) else None
    leverage = debt / ebit if debt is not None and ebit not in (None, 0) and ebit > 0 else None
    profile.update({"interest_coverage": coverage, "debt_to_ebit": leverage})
    if debt is not None and debt <= 0:
        profile.update(
            {"class": "investment_grade_like", "reason": "the company reports no debt, so a strong credit profile is assumed"}
        )
    elif coverage is not None and leverage is not None and (coverage < 2.5 or leverage > 5):
        profile.update(
            {"class": "stressed", "reason": f"interest coverage {coverage:.1f}x and debt/EBIT {leverage:.1f}x are weak"}
        )
    elif coverage is not None and leverage is not None and coverage >= 6 and leverage <= 3:
        profile.update(
            {"class": "investment_grade_like", "reason": f"interest coverage {coverage:.1f}x and debt/EBIT {leverage:.1f}x are strong"}
        )
    elif coverage is not None and leverage is not None:
        profile.update(
            {"class": "non_investment_grade_like", "reason": f"interest coverage {coverage:.1f}x and debt/EBIT {leverage:.1f}x are middling"}
        )
    else:
        profile.update({"class": "unknown", "reason": "debt, interest or EBIT were not available"})
    return profile


def business_estimate(profile: dict[str, Any]) -> tuple[float, str]:
    cls = profile.get("class")
    if cls == "investment_grade_like":
        return IG_RATE, f"investment-grade-like profile ({profile['reason']}); range 4-6%"
    if cls == "stressed":
        return STRESSED_RATE, f"higher-risk profile ({profile['reason']}); range 6-9%"
    if cls == "non_investment_grade_like":
        return NON_IG_RATE, f"non-investment-grade-like profile ({profile['reason']}); range 6-9%"
    return NON_IG_RATE, "credit profile could not be measured, so the middle of the 6-9% range is used"


def choose_rate(
    fy: str,
    ordered: list[tuple[str, Any]],
    disclosed: dict[str, float],
    profile: dict[str, Any],
) -> tuple[float, str, str]:
    """Return (rate, method, reason) for one bad year. `ordered` is [(fy, current value)] oldest to newest."""
    if fy in disclosed:
        return round(disclosed[fy], 4), "company_disclosed", f"{fy} 10-K reports a weighted-average operating-lease discount rate of {disclosed[fy] * 100:.2f}%"
    index = next(i for i, (k, _v) in enumerate(ordered) if k == fy)
    realistic = [(abs(i - index), float(v), k) for i, (k, v) in enumerate(ordered) if i != index and is_realistic(v)]
    if realistic:
        realistic.sort(key=lambda item: item[0])
        nearest = realistic[:3]
        rate = round(statistics.median(v for _d, v, _k in nearest), 4)
        years = ", ".join(k for _d, _v, k in nearest)
        return rate, "nearest_realistic_years", f"median of the nearest realistic years ({years}) in the same row"
    rate, why = business_estimate(profile)
    return rate, "business_estimate", why


class LeaseRateRowService:
    def apply(
        self,
        *,
        workbook_path: Path,
        lease_years: list[Any] | None = None,
        newest_only_fy: str | None = None,
    ) -> LeaseRowReport:
        """Fix the rate row. `newest_only_fy` limits changes to one year (annual update: earlier years are copied)."""
        report = LeaseRowReport()
        path = Path(workbook_path)
        cached_wb = load_workbook(path, data_only=True)
        try:
            if "Leases" not in cached_wb.sheetnames:
                report.skipped.append("No Leases tab.")
                return report
            cached_ws = cached_wb["Leases"]
            rate_row = self._rate_row(cached_ws)
            if rate_row is None:
                report.skipped.append("No long-term rate row found on the Leases tab.")
                return report
            cols = {k: c for k, c in detect_year_columns(cached_ws, cached_wb).items() if str(k).startswith("FY")}
            ordered = [(k, cached_ws.cell(rate_row, c).value) for k, c in sorted(cols.items(), key=lambda kv: kv[1])]
            report.credit_profile = credit_profile(self._latest_inputs(cached_wb))
        finally:
            cached_wb.close()

        disclosed: dict[str, float] = {}
        for item in lease_years or []:
            raw = _num(getattr(item, "reported_discount_rate", None))
            if raw is not None and is_realistic(_as_rate(raw)):
                disclosed[str(item.fiscal_year)] = _as_rate(raw)

        wb = load_workbook(path, data_only=False)
        try:
            ws = wb["Leases"]
            ledger = AdjustmentLedger(wb)
            report.checked = len(ordered)
            for fy, current in ordered:
                if newest_only_fy is not None and fy != newest_only_fy:
                    continue
                if is_realistic(current):
                    continue
                cell = ws.cell(rate_row, cols[fy])
                rate, method, reason = choose_rate(fy, ordered, disclosed, report.credit_profile)
                original = current if current is not None else "blank or N/A"
                source = {
                    "company_disclosed": "10-K operating-lease note (SEC XBRL)",
                    "nearest_realistic_years": f"Leases!{ws.cell(rate_row, 1).coordinate} neighbors",
                    "business_estimate": "HAP business-based estimate (Inputs: debt, interest, EBIT)",
                }[method]
                why = (
                    f"the formula (interest expense / total debt) gave {self._fmt(current)}, "
                    f"outside the realistic range {RATE_MIN * 100:.0f}%-{RATE_MAX * 100:.0f}%; {reason}"
                )
                original_formula = cell.value
                cell.value = rate
                ledger.record(
                    sheet="Leases", cell=cell.coordinate, fiscal_year=fy, category="Lease rate",
                    what="Estimated Long-Term Rate", original=original_formula, new=rate, amount=rate,
                    reason=why, source=source, method=method,
                    confidence="high" if method == "company_disclosed" else ("medium" if method == "nearest_realistic_years" else "low"),
                )
                report.fixes.append(
                    LeaseRateFix(fy, f"Leases!{cell.coordinate}", original, rate, method, why, source)
                )
            if report.fixes:
                report.notes_rows = self._write_notes(ws, report, rate_row)
                wb.save(path)
        finally:
            wb.close()
        return report

    @staticmethod
    def _fmt(value: Any) -> str:
        number = _num(value)
        return f"{number * 100:.2f}%" if number is not None else "blank or N/A"

    @staticmethod
    def _rate_row(ws) -> int | None:
        for row in range(1, min(ws.max_row or 1, 40) + 1):
            label = str(ws.cell(row, 1).value or "").strip().lower()
            if any(name in label for name in _RATE_LABELS):
                return row
        return None

    @staticmethod
    def _latest_inputs(wb) -> dict[str, float | None]:
        out: dict[str, float | None] = {"debt": None, "interest": None, "ebit": None}
        if "Inputs" not in wb.sheetnames:
            return out
        ws = wb["Inputs"]
        cols = {k: c for k, c in detect_year_columns(ws, wb).items() if str(k).startswith("FY")}
        if not cols:
            return out
        col = max(cols.values())
        wanted = {"debt": ("total debt",), "interest": ("interest expense",), "ebit": ("operating income before interest", "(ebit)")}
        for row in range(1, min(ws.max_row or 1, 60) + 1):
            label = str(ws.cell(row, 1).value or "").strip().lower()
            for key, needles in wanted.items():
                if out[key] is None and any(n in label for n in needles):
                    out[key] = _num(ws.cell(row, col).value)
        return out

    @staticmethod
    def _write_notes(ws, report: LeaseRowReport, rate_row: int) -> list[int]:
        """One plain-language note in the tab's single Notes block; the detail is in the HAP Adjustments tab."""
        from services.tab_notes import add_notes, note

        years = ", ".join(fix.fiscal_year for fix in report.fixes)
        written = add_notes(
            ws,
            [
                note(
                    f"The long-term lease rate for {years} was replaced",
                    f"the template formula gave a blank or unrealistic result (outside {RATE_MIN * 100:.0f}% to {RATE_MAX * 100:.0f}%)",
                    "the company's lease note, neighbouring years, or a credit-based estimate; details are in the HAP Adjustments tab",
                )
            ],
            replace_containing=("The long-term lease rate for",),
        )
        return [int("".join(ch for ch in cell.split("!")[-1] if ch.isdigit())) for cell in written]
