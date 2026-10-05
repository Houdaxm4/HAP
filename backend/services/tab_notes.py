"""One notes block per tab, in plain language.

Every tab has a single designated spot for HAP notes: a block headed "Notes" below the tab's data (never in a hidden column).
- Tabs where column B is hidden (the Last Quarter tabs): notes are in column D.
- All other tabs: notes are in column A.

Each note is one or two short sentences that answer only: what was done, why it was done, and which source was used.
No codes, no cell formulas, no repetition. Any older HAP notes block ("HAP ANALYSIS — NOTES") on the tab is removed when the
new block is written, so notes never end up in two places.
"""

from __future__ import annotations

from typing import Any

from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.worksheet import Worksheet

NOTES_TITLE = "Notes"
LEGACY_HEADERS = ("HAP ANALYSIS — NOTES", "HAP ANALYSIS - NOTES")
NOTE_FILL = PatternFill("solid", fgColor="F7F7F2")
MAX_NOTE_LENGTH = 400


def note(what: str, why: str | None = None, source: str | None = None, *, because: bool = True) -> str:
    """One note: what was done, why, and the source, as plain sentences."""
    def sentence(text: str) -> str:
        text = " ".join(str(text).split()).rstrip(".")
        return (text[:1].upper() + text[1:] + ".") if text else ""

    what = " ".join(str(what).split()).rstrip(".")
    why = " ".join(str(why).split()).rstrip(".") if why else ""
    if why and because:
        parts = [sentence(f"{what} because {why}")]
    else:
        parts = [sentence(what)] + ([sentence(why)] if why else [])
    if source:
        parts.append(f"Source: {' '.join(str(source).split()).rstrip('.')}.")
    return " ".join(p for p in parts if p)[:MAX_NOTE_LENGTH]


def notes_column(ws: Worksheet) -> int:
    hidden_b = "B" in ws.column_dimensions and bool(ws.column_dimensions["B"].hidden)
    return 4 if hidden_b else 1


def _is_blank(ws: Worksheet, row: int, max_col: int = 8) -> bool:
    return all(ws.cell(row, c).value in (None, "") for c in range(1, max_col + 1))


def _clear_row(ws: Worksheet, row: int, max_col: int = 8) -> None:
    for c in range(1, max_col + 1):
        cell = ws.cell(row, c)
        cell.value = None
        cell.fill = PatternFill(fill_type=None)
        cell.comment = None


def _legacy_header_rows(ws: Worksheet) -> list[int]:
    rows = []
    for row in range(1, (ws.max_row or 1) + 1):
        value = ws.cell(row, 1).value
        if isinstance(value, str) and value.strip() in LEGACY_HEADERS:
            rows.append(row)
    return rows


def remove_legacy_blocks(ws: Worksheet) -> int:
    """Delete older HAP notes blocks (header in column A and the rows that follow until the first empty row)."""
    removed = 0
    for header in sorted(_legacy_header_rows(ws), reverse=True):
        row = header
        _clear_row(ws, row)
        row += 1
        while row <= (ws.max_row or 1) and not _is_blank(ws, row):
            _clear_row(ws, row)
            row += 1
            removed += 1
    return removed


def _notes_header_row(ws: Worksheet, col: int) -> int | None:
    for row in range(1, (ws.max_row or 1) + 1):
        value = ws.cell(row, col).value
        if isinstance(value, str) and value.strip() == NOTES_TITLE:
            return row
    return None


def read_notes(ws: Worksheet) -> list[str]:
    col = notes_column(ws)
    header = _notes_header_row(ws, col)
    if header is None:
        return []
    out: list[str] = []
    row = header + 1
    while row <= (ws.max_row or 1) and ws.cell(row, col).value not in (None, ""):
        out.append(str(ws.cell(row, col).value))
        row += 1
    return out


def _last_data_row(ws: Worksheet) -> int:
    last = 0
    for row in range(1, (ws.max_row or 1) + 1):
        if not _is_blank(ws, row, max_col=min(ws.max_column or 1, 26)):
            last = row
    return last


def add_notes(ws: Worksheet, notes: list[str], *, replace_containing: tuple[str, ...] = ()) -> list[str]:
    """Add notes to the tab's single Notes block (created below the data when it does not exist yet).

    Notes that are already there are not repeated. `replace_containing` drops older notes that contain any of the given phrases, so a
    re-run replaces its own earlier note instead of adding a second one. Returns the cells written."""
    col = notes_column(ws)
    remove_legacy_blocks(ws)
    existing = read_notes(ws)
    kept = [n for n in existing if not any(p in n for p in replace_containing)]
    merged = list(kept)
    for text in notes:
        text = " ".join(str(text).split())
        if text and text not in merged:
            merged.append(text)
    header = _notes_header_row(ws, col)
    if header is None:
        header = _last_data_row(ws) + 2
    else:
        row = header + 1
        while row <= (ws.max_row or 1) and ws.cell(row, col).value not in (None, ""):
            ws.cell(row, col).value = None
            row += 1
    title = ws.cell(header, col)
    title.value = NOTES_TITLE
    title.font = Font(bold=True)
    title.fill = NOTE_FILL
    written = [f"{ws.title}!{title.coordinate}"]
    for offset, text in enumerate(merged, start=1):
        cell = ws.cell(header + offset, col)
        cell.value = text
        cell.fill = NOTE_FILL
        cell.alignment = Alignment(wrap_text=False, vertical="top")
        written.append(f"{ws.title}!{cell.coordinate}")
    return written


def fmt_period(label: Any) -> str:
    """'FY2026 Q3' -> 'fiscal Q3 2026' for sentences."""
    text = str(label or "").strip()
    if text.upper().startswith("FY") and " " in text:
        year, quarter = text.split(" ", 1)
        return f"fiscal {quarter} {year[2:]}"
    return text or "the latest quarter"
