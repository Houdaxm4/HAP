"""Make sure a finished workbook carries calculated values before anything is read from it (the emails, the business document).

Steps that save a workbook with openpyxl drop the values Excel had calculated, so a formula cell can read as empty. When the Final Metrics
tab shows that, Excel recalculates the deliverable once (the same full rebuild the pipeline already uses).
"""

from __future__ import annotations

import os
from pathlib import Path

from openpyxl import load_workbook


def needs_recalculation(path: Path) -> bool:
    try:
        formulas = load_workbook(path, data_only=False)
        values = load_workbook(path, data_only=True)
    except Exception:  # noqa: BLE001 - unreadable: nothing to do here
        return False
    try:
        if "Final Metrics" not in formulas.sheetnames:
            return False
        probe = formulas["Final Metrics"]["B5"].value
        return isinstance(probe, str) and probe.startswith("=") and values["Final Metrics"]["B5"].value is None
    finally:
        formulas.close()
        values.close()


def ensure_calculated(path: Path, *, analysis_id: str, ticker: str, fiscal_year: str | int | None = None) -> str | None:
    """Recalculate the workbook with Excel when its formulas have no values. Returns the recalculation status, or None when not needed."""
    if os.environ.get("HAP_EMAIL_RECALC", "1").strip() in {"0", "false", "no"}:
        return None
    if not needs_recalculation(Path(path)):
        return None
    from services.excel_recalc_service import ExcelRecalcService

    token = f"FY{fiscal_year}" if fiscal_year and not str(fiscal_year).startswith("FY") else (str(fiscal_year) if fiscal_year else None)
    try:
        report = ExcelRecalcService().recalculate(analysis_id=analysis_id, ticker=ticker, workbook_path=Path(path), fiscal_year=token)
    except Exception as exc:  # noqa: BLE001 - advisory: the email then shows n/a for what has no value
        return f"FAILED: {type(exc).__name__}"
    return report.status
