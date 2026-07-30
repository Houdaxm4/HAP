"""Workbook Scanner — read-only openpyxl inspection."""

from __future__ import annotations

import hashlib
import re
from collections import Counter
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
from openpyxl.workbook.workbook import Workbook
from openpyxl.worksheet.worksheet import Worksheet

# Hard caps: Industrial Template used ranges are modest; declared max_column can be 16382.
MAX_SCAN_ROWS = 400
MAX_SCAN_COLS = 80

SHEET_REF_RE = re.compile(r"'([^']+)'!")


@dataclass
class RawCell:
    row: int
    column: int
    address: str
    value: Any
    is_formula: bool
    number_format: str | None = None


@dataclass
class RawSheet:
    name: str
    index: int
    state: str
    declared_max_row: int | None
    declared_max_column: int | None
    used_max_row: int
    used_max_column: int
    cells: list[RawCell] = field(default_factory=list)
    merged_cells: list[str] = field(default_factory=list)
    tables: list[dict[str, Any]] = field(default_factory=list)
    charts: list[dict[str, Any]] = field(default_factory=list)
    data_validations: list[dict[str, Any]] = field(default_factory=list)
    conditional_formatting: list[dict[str, Any]] = field(default_factory=list)
    freeze_panes: str | None = None
    hidden_rows: list[int] = field(default_factory=list)
    hidden_columns: list[str] = field(default_factory=list)
    protection: dict[str, Any] = field(default_factory=dict)
    outbound_deps: Counter[str] = field(default_factory=Counter)


@dataclass
class RawWorkbook:
    path: Path
    filename: str
    sha256: str
    sheetnames: list[str]
    named_ranges: list[dict[str, Any]]
    workbook_protection: dict[str, Any]
    sheets: list[RawSheet]
    template_version_sheet: str | None = None


def _is_formula(v: Any) -> bool:
    return isinstance(v, str) and v.startswith("=")


def _preview(v: Any, limit: int = 120) -> str | None:
    if v is None:
        return None
    if _is_formula(v):
        return None
    s = str(v).strip()
    return s[:limit] if s else None


def _sheet_refs(formula: str, known: set[str]) -> set[str]:
    refs: set[str] = set()
    for m in SHEET_REF_RE.finditer(formula):
        refs.add(m.group(1))
    for ks in known:
        pat = re.compile(rf"(?<![A-Za-z0-9_]){re.escape(ks)}!")
        if pat.search(formula):
            refs.add(ks)
    return refs


def _compute_used_bounds(ws: Worksheet) -> tuple[int, int]:
    """Find last row/col with content within scan caps (ignore spurious XFB dims)."""
    max_r = min(ws.max_row or 1, MAX_SCAN_ROWS)
    max_c = min(ws.max_column or 1, MAX_SCAN_COLS)
    used_r, used_c = 0, 0
    for r in range(1, max_r + 1):
        for c in range(1, max_c + 1):
            v = ws.cell(r, c).value
            if v is None:
                continue
            if isinstance(v, str) and not v.strip():
                continue
            used_r = max(used_r, r)
            used_c = max(used_c, c)
    return used_r, used_c


def _extract_tables(ws: Worksheet) -> list[dict[str, Any]]:
    out = []
    try:
        for tname, table in (ws.tables or {}).items():
            out.append(
                {
                    "name": tname,
                    "ref": getattr(table, "ref", None),
                    "displayName": getattr(table, "displayName", tname),
                }
            )
    except Exception as exc:  # noqa: BLE001
        out.append({"error": str(exc)})
    return out


def _extract_charts(ws: Worksheet) -> list[dict[str, Any]]:
    out = []
    try:
        charts = getattr(ws, "_charts", None) or []
        for i, ch in enumerate(charts):
            out.append(
                {
                    "index": i,
                    "type": type(ch).__name__,
                    "title": str(getattr(getattr(ch, "title", None), "tx", None) or "")[:80],
                }
            )
    except Exception as exc:  # noqa: BLE001
        out.append({"error": str(exc)})
    return out


def _extract_data_validations(ws: Worksheet) -> list[dict[str, Any]]:
    dvs = getattr(ws, "data_validations", None)
    if not dvs or not getattr(dvs, "dataValidation", None):
        return []
    out = []
    for dv in dvs.dataValidation:
        out.append(
            {
                "sqref": str(getattr(dv, "sqref", "")),
                "type": getattr(dv, "type", None),
                "formula1": getattr(dv, "formula1", None),
                "formula2": getattr(dv, "formula2", None),
            }
        )
    return out


def _extract_cf(ws: Worksheet) -> list[dict[str, Any]]:
    cf = getattr(ws, "conditional_formatting", None)
    if not cf:
        return []
    out = []
    try:
        for sqref in cf:
            rules = cf[sqref]
            out.append({"sqref": str(sqref), "rule_count": len(rules)})
    except Exception as exc:  # noqa: BLE001
        out.append({"error": str(exc)})
    return out


def _hidden_rows_cols(ws: Worksheet) -> tuple[list[int], list[str]]:
    rows = []
    for idx, dim in (ws.row_dimensions or {}).items():
        if getattr(dim, "hidden", False):
            rows.append(int(idx))
    cols = []
    for letter, dim in (ws.column_dimensions or {}).items():
        if getattr(dim, "hidden", False):
            cols.append(str(letter))
    return rows, cols


def _protection(ws: Worksheet) -> dict[str, Any]:
    p = ws.protection
    return {
        "sheet": bool(getattr(p, "sheet", False)),
        "password_set": bool(getattr(p, "password", None)),
    }


def _workbook_protection(wb: Workbook) -> dict[str, Any]:
    sec = getattr(wb, "security", None)
    if sec is None:
        return {
            "workbook_locked": False,
            "structure_locked": False,
            "windows_locked": False,
            "detail": {},
        }
    return {
        "workbook_locked": bool(getattr(sec, "workbookPassword", None) or getattr(sec, "workbook_password", None)),
        "structure_locked": bool(getattr(sec, "lockStructure", False) or getattr(sec, "lock_structure", False)),
        "windows_locked": bool(getattr(sec, "lockWindows", False) or getattr(sec, "lock_windows", False)),
        "detail": {
            "revisionsPassword_set": bool(getattr(sec, "revisionsPassword", None)),
        },
    }


def _named_ranges(wb: Workbook) -> list[dict[str, Any]]:
    out = []
    defined = wb.defined_names
    for name in defined:
        dn = defined[name]
        attr = getattr(dn, "attr_text", None)
        dests = []
        broken = bool(attr and "#REF!" in str(attr))
        try:
            for sheet, coord in list(dn.destinations) if attr else []:
                dests.append({"sheet": sheet, "coord": str(coord)})
        except Exception:  # noqa: BLE001
            broken = True
        out.append(
            {
                "name": name,
                "attr_text": attr,
                "destinations": dests,
                "broken": broken,
            }
        )
    return out


class WorkbookScanner:
    """Scan an .xlsx into a RawWorkbook (no classification)."""

    def scan(self, path: Path | str) -> RawWorkbook:
        path = Path(path)
        sha = hashlib.sha256(path.read_bytes()).hexdigest()
        wb = load_workbook(path, data_only=False, read_only=False)
        known = set(wb.sheetnames)
        sheets: list[RawSheet] = []

        for idx, name in enumerate(wb.sheetnames):
            ws = wb[name]
            state = getattr(ws, "sheet_state", "visible") or "visible"
            used_r, used_c = _compute_used_bounds(ws)
            cells: list[RawCell] = []
            deps: Counter[str] = Counter()

            for r in range(1, used_r + 1):
                for c in range(1, used_c + 1):
                    cell = ws.cell(r, c)
                    v = cell.value
                    if v is None or (isinstance(v, str) and not str(v).strip()):
                        continue
                    formula = _is_formula(v)
                    addr = f"{get_column_letter(c)}{r}"
                    cells.append(
                        RawCell(
                            row=r,
                            column=c,
                            address=addr,
                            value=v,
                            is_formula=formula,
                            number_format=str(cell.number_format) if cell.number_format else None,
                        )
                    )
                    if formula:
                        for ref in _sheet_refs(str(v), known):
                            if ref != name:
                                deps[ref] += 1

            hidden_rows, hidden_cols = _hidden_rows_cols(ws)
            sheets.append(
                RawSheet(
                    name=name,
                    index=idx,
                    state=str(state),
                    declared_max_row=ws.max_row,
                    declared_max_column=ws.max_column,
                    used_max_row=used_r,
                    used_max_column=used_c,
                    cells=cells,
                    merged_cells=[str(r) for r in ws.merged_cells.ranges],
                    tables=_extract_tables(ws),
                    charts=_extract_charts(ws),
                    data_validations=_extract_data_validations(ws),
                    conditional_formatting=_extract_cf(ws),
                    freeze_panes=str(ws.freeze_panes) if ws.freeze_panes else None,
                    hidden_rows=hidden_rows,
                    hidden_columns=hidden_cols,
                    protection=_protection(ws),
                    outbound_deps=deps,
                )
            )

        version = None
        if "Template Version" in wb.sheetnames:
            version = wb["Template Version"]["A2"].value
            version = str(version).strip() if version is not None else None

        protection = _workbook_protection(wb)
        named = _named_ranges(wb)
        sheetnames = list(wb.sheetnames)
        wb.close()

        m = re.search(r"v\d+\.\d+", path.name)
        _ = m.group(0) if m else None  # filename version available to caller via filename

        return RawWorkbook(
            path=path,
            filename=path.name,
            sha256=sha,
            sheetnames=sheetnames,
            named_ranges=named,
            workbook_protection=protection,
            sheets=sheets,
            template_version_sheet=version,
        )
