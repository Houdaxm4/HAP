"""Daily batch: run several HAP analyses from the HAP work folder and put the finished files in HAP results.

    python scripts/run_batch.py TKR:new_company SXI:new_company WWD:new_company
    python scripts/run_batch.py AMAT:quarterly_update:2025Q1      # optional third part: the period to update, when the folder holds several
    python scripts/run_batch.py TPL:new_company:2019Q4 TPL:new_company:2020Q3 --lease-rate TPL=0.05
                                  # a new company can also name its period when the folder holds several workbooks of it (one run each);
                                  # --lease-rate: an analyst-confirmed lease discount rate, used when HAP finds no evidence for one (no debt)

Files are matched by name in the work folder (default Downloads\\HAP work):
  workbook      a ticker and a period in any order or style: "RPM 2027 Q1 - Industrial Template v28.2.xlsx", "RPM_2027Q1.xlsx",
                "RPM FY2026.xlsx", "Q1 2027 RPM.xlsx"     (new company: the template; updates: newer period = new, older = previous)
                a finished result "2026 FY RPM FA.xlsx" also counts, so last period's result can be dropped in as is
  custom filter "custom_run_filter_<date>_<TICKER>.xlsx"       (spaces or no date also work; latest date and time wins)
  no filter     when a company has no custom filter, a stand-in is built from the "Current Filter Summary Sheet" (an .xlsx with a Summary tab,
                one row per company). It carries today's market data only, so price, entry prices, margin of safety and expected return are
                today's, not the period's; the run is marked "stand-in custom filter".
  sheet         the newest .xlsx with an S&P tab               (email fields: Classified as, Status with PE10, TBV/P)
Results (flat folder, default Downloads\\HAP results): "<period> <TICKER> FA.xlsx" and "<period> <TICKER> Email.txt / .eml".
The lease-rate review is approved automatically and the run is marked "lease rate auto-approved, please check".
A company with missing or ambiguous files is skipped and reported; the others still run.
"""

from __future__ import annotations

import argparse
import json
import os
import re
import shutil
import sys
import traceback
from dataclasses import dataclass, field
from pathlib import Path

BACKEND_ROOT = Path(__file__).resolve().parents[1]
if str(BACKEND_ROOT) not in sys.path:
    sys.path.insert(0, str(BACKEND_ROOT))

DOWNLOADS = Path.home() / "Downloads"
# A workbook name holds a period ("2027 Q1", "2027Q1", "Q1 2027", "FY2026", "2026_FY") and the ticker, in either order:
# "RPM 2027 Q1 - Industrial Template v28.2", "RPM_2027_Q1", "RPM FY2026", "Q1 2027 RPM", and a finished result "2026 FY RPM FA".
_PERIOD_YEAR_FIRST = re.compile(r"(?<!\d)(?P<y>20\d{2})[ _\-]*(?P<p>Q[1-4]|FY)(?![A-Za-z0-9])", re.I)
_PERIOD_PERIOD_FIRST = re.compile(r"(?<![A-Za-z0-9])(?P<p>Q[1-4]|FY)[ _\-]*(?P<y>20\d{2})(?!\d)", re.I)
_TOKEN = re.compile(r"[A-Za-z0-9][A-Za-z0-9.\-]*")
# Custom run filter: "Custom_Run_Filter_2026-10-12-(16-47)-RPM", "custom run filter 2026-10-12 RPM", "Custom_Run_Filter_RPM", "... (1)" copies
_CRF_PREFIX = re.compile(r"^custom[ _\-]*run[ _\-]*filter", re.I)
_CRF_NAME = re.compile(
    r"^custom[ _\-]*run[ _\-]*filter[ _\-]*"
    r"(?:(?P<d>\d{4}[ _\-]\d{2}[ _\-]\d{2}(?:[ _\-]+\d{2}[ _\-]\d{2})?)[ _\-]+)?"
    r"(?P<t>[A-Za-z0-9][A-Za-z0-9.\-]*)(?:[ _]+.*)?$", re.I)
TYPES = {"new_company", "annual_update", "quarterly_update"}
PERIOD_ORDER = {"Q1": 1, "Q2": 2, "Q3": 3, "Q4": 4, "FY": 5}


@dataclass
class Match:
    workbook: Path | None = None
    previous: Path | None = None
    crf: Path | None = None
    standin: dict | None = None
    pit: dict | None = None
    notes: list[str] = field(default_factory=list)
    problems: list[str] = field(default_factory=list)


def _tokens(text: str) -> list[str]:
    return [tok.strip(".-").upper() for tok in _TOKEN.findall(text) if tok.strip(".-")]


def _parse_workbook_name(path: Path) -> tuple[set[str], tuple[int, int]] | None:
    """(possible tickers, (year, period order)) of a workbook name, or None when it holds no period.

    The ticker is a word before the period ("RPM 2027 Q1 ...", "Industrial Template RPM Q1 2027"); when nothing comes
    before the period it is the first word after it (a finished result: "2026 FY RPM FA")."""
    text = path.stem
    hits = [m for m in (_PERIOD_YEAR_FIRST.search(text), _PERIOD_PERIOD_FIRST.search(text)) if m]
    if not hits:
        return None
    m = min(hits, key=lambda x: x.start())
    before, after = _tokens(text[: m.start()]), _tokens(text[m.end():])
    tickers = set(before) if before else set(after[:1])
    if not tickers:
        return None
    return tickers, (int(m.group("y")), PERIOD_ORDER[m.group("p").upper()])


def _period_key(path: Path) -> tuple[int, int] | None:
    parsed = _parse_workbook_name(path)
    return parsed[1] if parsed else None


def _parse_crf_name(path: Path) -> tuple[str, tuple[int, ...]] | None:
    """(ticker, date/time numbers) of a custom run filter, or None. A '(1)' copy suffix and the brackets of the time are ignored."""
    stem = re.sub(r"\s*\(\d{1,2}\)$", "", path.stem)
    m = _CRF_NAME.match(stem.replace("(", "").replace(")", "").strip())
    if not m:
        return None
    return m.group("t").upper(), tuple(int(x) for x in re.findall(r"\d+", m.group("d") or ""))


def parse_period(text: str) -> tuple[int, int]:
    """'2025Q1', '2025 Q1', 'Q1 2025', 'FY2024' -> (year, period order)."""
    for rx in (_PERIOD_YEAR_FIRST, _PERIOD_PERIOD_FIRST):
        m = rx.search(text)
        if m:
            return int(m.group("y")), PERIOD_ORDER[m.group("p").upper()]
    raise ValueError(f"cannot read a period from '{text}' (use for example 2025Q1 or 2024FY)")


def _workbook_ticker_and_version(path: Path) -> tuple[str | None, str | None]:
    """Ticker on the statement tab and the template version of a workbook (None when not readable)."""
    from openpyxl import load_workbook

    try:
        wb = load_workbook(path, read_only=True, data_only=True)
    except Exception:  # noqa: BLE001 - not a readable workbook
        return None, None
    try:
        ticker = version = None
        if "Balance Sheet - Standardized" in wb.sheetnames:
            row = next(wb["Balance Sheet - Standardized"].iter_rows(min_row=1, max_row=1, max_col=4, values_only=True), ())
            ticker = str(row[2]).strip().upper() if len(row) > 2 and row[2] else None
        if "Template Version" in wb.sheetnames:
            for r in wb["Template Version"].iter_rows(min_row=1, max_row=4, max_col=3, values_only=True):
                if r and str(r[0] or "").strip().lower() == "version" and r[1]:
                    version = str(r[1]).strip()
        return ticker, version
    finally:
        wb.close()


def match_files(folder: Path, ticker: str, analysis_type: str, period: tuple[int, int] | None = None, standin_dir: Path | None = None) -> Match:
    t = ticker.upper()
    out = Match()
    books: list[tuple[tuple[int, int], Path]] = []
    crfs: list[tuple[tuple[int, ...], str, Path]] = []
    for p in sorted(folder.glob("*.xlsx")):
        if p.name.startswith("~$"):
            continue
        if _CRF_PREFIX.match(p.name):          # a custom run filter is never a workbook, whatever its name holds
            crf = _parse_crf_name(p)
            if crf and crf[0] == t:
                crfs.append((crf[1], p.name, p))
            continue
        parsed = _parse_workbook_name(p)
        if parsed and t in parsed[0]:
            books.append((parsed[1], p))
    if crfs:
        out.crf = max(crfs)[2]            # the latest date and time wins
    books.sort(key=lambda x: x[0])
    if not books:
        out.problems.append(f"no workbook for {t} (expected '{t} <year> <Q#|FY> ...xlsx')")
    elif analysis_type == "new_company" and period is not None:
        chosen = [b for b in books if b[0] == period]
        if len(chosen) == 1:
            out.workbook = chosen[0][1]
        elif not chosen:
            out.problems.append(f"no workbook for {t} for the requested period; found: " + ", ".join(p.name for _k, p in books))
        else:
            out.problems.append(f"{len(chosen)} workbooks for {t} for the requested period: " + ", ".join(p.name for _k, p in chosen))
    elif analysis_type == "new_company":
        if len(books) > 1:
            out.problems.append(f"{len(books)} workbooks for {t}; a new company needs exactly one: " + ", ".join(p.name for _k, p in books))
        else:
            out.workbook = books[0][1]
    else:
        if period is not None:
            newer = [b for b in books if b[0] == period]
            if not newer:
                out.problems.append(f"no workbook for {t} for the requested period; found: " + ", ".join(p.name for _k, p in books))
                books = []
            else:
                older = [b for b in books if b[0] < period]
                books = older[-1:] + newer[-1:]
        elif len(books) > 2:
            out.notes.append(
                f"{len(books)} workbooks for {t}: using the newest two ({books[-1][1].name} as new, {books[-2][1].name} as previous); "
                f"to pick another period write {t}:{analysis_type}:{_period_label(books[0][0])}"
            )
        if len(books) < 2:
            if not out.problems:
                near = sorted({p.name for p in folder.glob("*.xlsx") if (info := _parse_workbook_name(p)) and any(x != t and t.startswith(x) for x in info[0])})
                out.problems.append(
                    f"an update needs two workbooks for {t} (new and previous period); found {len(books)}"
                    + (f". Files for a similar ticker are ignored: {', '.join(near)}" if near else "")
                )
        elif books[-1][0] == books[-2][0]:
            out.problems.append(f"the two newest workbooks for {t} are both for the same period: {books[-2][1].name}, {books[-1][1].name}")
        else:
            out.workbook, out.previous = books[-1][1], books[-2][1]
    # the files must really be this company, and on the same template version (or HAP maps the older one)
    versions: dict[str, str | None] = {}
    for label, path in (("workbook", out.workbook), ("previous", out.previous)):
        if path is None:
            continue
        found, version = _workbook_ticker_and_version(path)
        versions[label] = version
        if found and found != t:
            out.problems.append(f"{path.name} holds {found}, not {t}: wrong company or a mislabeled template")
            out.workbook = out.previous = None
            break
    if out.workbook and out.previous and versions.get("workbook") and versions.get("previous") and versions["workbook"] != versions["previous"]:
        out.notes.append(
            f"previous workbook is on template {versions['previous']}, the new one on {versions['workbook']}: "
            "statement rows differ, HAP keeps the new layout and carries analyst values by row label"
        )
    if out.crf is None and out.workbook is None:
        out.problems.append(f"no custom run filter for {t} (expected custom_run_filter_<date>_{t}.xlsx)")
    elif out.crf is None:
        from services.custom_run_standin_service import StandInError, build_standin, find_filter_summary

        summary = find_filter_summary(folder)
        if summary is None:
            out.problems.append(f"no custom run filter for {t} (expected custom_run_filter_<date>_{t}.xlsx) and no filter summary sheet to build one from")
        else:
            try:
                target = (standin_dir or folder) / f"custom_run_filter_standin_{t}.xlsx"
                out.standin = build_standin(t, summary, target)
                out.crf = target
                out.notes.append(
                    f"stand-in custom filter built from {summary.name}: price, entry prices, margin of safety, expected return and PE10 status "
                    "are today's values, not the period's; no PE10 history"
                )
            except StandInError as exc:
                out.problems.append(f"no custom run filter for {t} and no stand-in possible: {exc}")
    return out


def _period_label(key: tuple[int, int]) -> str:
    return f"{key[0]}{'FY' if key[1] == 5 else 'Q' + str(key[1])}"


def apply_point_in_time(match: Match, ticker: str, results: Path, mode: str) -> None:
    """Back-test: when the custom run filter holds history later than the period being updated, use the filter as it stood at the end of
    that period (scalars from the quarter's history column, nothing later) and price the run with that quarter's average price."""
    if mode == "off" or match.standin or match.crf is None or match.workbook is None:
        return
    key = _period_key(match.workbook)
    if key is None:
        return
    from services.market_price_service import MarketPriceService
    from services.point_in_time_filter_service import (
        PointInTimeError, available_quarters, build_point_in_time, infer_share_factor, quarter_key,
    )

    year, order = key
    label = f"{year} Q{4 if order == 5 else order}"
    try:
        quarters = available_quarters(match.crf)
        if label not in quarters:
            match.notes.append(f"point-in-time not applied: {match.crf.name} has no history column for {label}")
            return
        if quarter_key(label) >= quarter_key(quarters[-1]):
            return                                           # the filter's latest quarter: it is already current for this period
        factor = infer_share_factor(match.workbook, match.crf, label)
        info = build_point_in_time(match.crf, label, results / "_point_in_time" / f"custom_run_filter_pit_{ticker}_{label.replace(' ', '')}.xlsx",
                                   share_factor=factor)
    except PointInTimeError as exc:
        match.notes.append(f"point-in-time not applied: {exc}")
        return
    match.pit = info
    match.crf = Path(info["path"])
    MarketPriceService.set_point_in_time(ticker, info["price_quarter_average"], f"custom run filter quarter average ({label})")
    if info.get("share_factor", 1.0) != 1.0:
        match.notes.append(
            f"share basis: the workbook is adjusted for later splits (diluted shares x{info['share_factor']:g} the filter's), so the filter's "
            f"per-share figures were divided by {info['share_factor']:g} before use"
        )
    match.notes.append(
        f"point-in-time back-test as of {label} (quarter end {info['quarter_end']}) from {info['source']}: price = quarter average "
        f"{info['price_quarter_average']:.2f}, max entry price {info['max_price_to_buy']:.2f}; PE10 percentile is approximate (quarterly history)"
    )


def _copy_upload(src: Path, dest: Path):
    from models.analysis import UploadedFileMetadata
    from models.common import utc_now_iso

    dest.parent.mkdir(parents=True, exist_ok=True)
    shutil.copy2(src, dest)
    return UploadedFileMetadata(filename=src.name, stored_filename=dest.name, size_bytes=dest.stat().st_size, uploaded_at=utc_now_iso())


def _fill_expected_return_rows(excel: Path, analysis_id: str, ticker: str, notes: list[str]) -> None:
    """Back-test: Bloomberg has no expected return for an earlier date, so Inputs B69:B71 take the workbook's own model.

    B69 price-only and B70 price + dividends come from the Expected Returns & Buybacks tab (at the current price); B71 is the same
    price + dividends return measured from the maximum entry price (Inputs B67). Blank cells only; the workbook is recalculated after."""
    from openpyxl import load_workbook

    tab = "'Expected Returns & Buybacks'"
    formulas = {
        "B69": f"={tab}!E14",
        "B70": f"={tab}!F14",
        "B71": f"=({tab}!D14/Inputs!B67)^(1/10)-1",
    }
    wb = load_workbook(excel)
    try:
        if "Inputs" not in wb.sheetnames or "Expected Returns & Buybacks" not in wb.sheetnames:
            return
        ws = wb["Inputs"]
        todo = {cell: f for cell, f in formulas.items() if ws[cell].value in (None, "")}
        if not todo:
            return
        for cell, formula in todo.items():
            ws[cell] = formula
        wb.save(excel)
    finally:
        wb.close()
    from services.excel_recalc_service import ExcelRecalcService

    report = ExcelRecalcService().recalculate(analysis_id=analysis_id, ticker=ticker, workbook_path=excel)
    notes.append("Expected return rows (Inputs B69:B71) come from the workbook's own model, not Bloomberg (back-test)"
                 + (f"; recalculation {report.status}, please open and recalculate" if report.status in ("UNAVAILABLE", "FAILED") else ""))


def _apply_analyst_lease_rate(orchestrator, output_service, analysis, rate: float, notes: list[str]):
    """An analyst-confirmed lease rate, for a company where HAP finds no evidence for one (no debt, no disclosed lease rate)."""
    try:
        review = output_service.read_json(analysis.analysis_id, "lease_rate_review.json")
    except FileNotFoundError:
        return analysis
    if review.get("selected_rate") is not None or review.get("approved_rate") is not None:
        return analysis                                       # HAP found its own rate: the analyst's is not needed
    analysis = orchestrator.finalize_new_company_review(
        analysis, action="correct", rate=rate, reason="Analyst-confirmed lease discount rate (HAP found no company evidence for one)"
    )
    notes.append(f"lease rate {rate:.2%} set by the analyst (HAP had no evidence for one)")
    return analysis


def _period_of(excel: Path, deliv: dict) -> tuple[int | None, int | None]:
    """Year and quarter of the deliverable, read from the Excel name ('2026 Q2 TKR FA.xlsx'); the report carries no quarter for new companies."""
    m = re.match(r"^(\d{4}) (Q[1-3]|FY) ", excel.name)
    if m:
        return int(m.group(1)), (int(m.group(2)[1]) if m.group(2).startswith("Q") else None)
    return deliv.get("fiscal_year"), deliv.get("fiscal_quarter")


def run_one(ticker: str, analysis_type: str, match: Match, sheet, results: Path, lease_rate: float | None = None, analyst: dict | None = None) -> dict:
    import main as hap_main
    from models.analysis import AnalysisFiles, CreateAnalysisRequest
    from services.deliverable_naming import business_deliverable_name, email_deliverable_name
    from services.business_doc_service import BusinessDocService
    from services.email_draft_service import EmailDraftService, SheetFields
    from services.daily_sheet_service import custom_filter_status

    analysis_service, file_service, output_service = hap_main.analysis_service, hap_main.file_service, hap_main.output_service
    orchestrator = hap_main.pipeline_orchestrator
    company = sheet.company_name(ticker) if sheet else ""
    created = analysis_service.create(CreateAnalysisRequest(company=company or ticker, ticker=ticker, analysis_type=analysis_type))
    analysis_id = created.analysis_id
    upload_dir = file_service.analysis_upload_dir(analysis_id)
    files = AnalysisFiles(
        prefilled_workbook=_copy_upload(match.workbook, upload_dir / "prefilled_workbook.xlsx"),
        custom_run_filter=_copy_upload(match.crf, upload_dir / "custom_run_filter.xlsx"),
        previous_workbook=_copy_upload(match.previous, upload_dir / "previous_workbook.xlsx") if match.previous else None,
    )
    created.files = files
    created.status = "uploaded"
    analysis_service.save(created)

    notes: list[str] = list(match.notes)
    analysis = orchestrator.run(analysis_id)
    if lease_rate is not None and analysis_type == "new_company":
        analysis = _apply_analyst_lease_rate(orchestrator, output_service, analysis, lease_rate, notes)
    if analysis.status == "awaiting_analyst_review":
        review = {}
        try:
            review = output_service.read_json(analysis_id, "lease_rate_review.json")
        except FileNotFoundError:
            pass
        analysis = orchestrator.finalize_new_company_review(
            analysis, action="approve", rate=None, reason="Daily batch: lease rate auto-approved, please check"
        )
        notes.append(f"lease rate auto-approved, please check (proposed rate {review.get('proposed_rate')})")

    summary = {"ticker": ticker, "type": analysis_type, "analysis_id": analysis_id, "status": analysis.status, "notes": notes}
    key = {"new_company": "new_company_deliverables_report.json", "annual_update": "annual_deliverables_report.json",
           "quarterly_update": "quarterly_deliverables_report.json"}[analysis_type]
    try:
        deliv = output_service.read_json(analysis_id, key)
    except FileNotFoundError:
        summary["problem"] = f"no deliverables were produced (status {analysis.status}); see storage/outputs/{analysis_id}"
        return summary

    out_dir = output_service.analysis_output_dir(analysis_id)
    excel = Path(deliv["excel_path"])
    fy, fq = _period_of(excel, deliv)
    if match.pit:
        try:
            _fill_expected_return_rows(excel, analysis_id, ticker, notes)
        except Exception as exc:  # noqa: BLE001 - the rows stay blank; the run is not lost
            notes.append(f"expected return rows not filled: {exc}")
    projection = None
    if analysis_type == "quarterly_update":
        try:
            from models.quarterly_update import QuarterlyProjectionReport
            from services.quarterly_deliverables_service import QuarterlyDeliverablesService

            projection = QuarterlyDeliverablesService._email_projection(
                QuarterlyProjectionReport.model_validate(output_service.read_json(analysis_id, "quarterly_projection_report.json"))
            )
        except Exception:  # noqa: BLE001 - the projection lines are optional (Q2 and Q3 only)
            projection = None
    # email again with the daily sheet's fields (the pipeline's own draft has no access to the sheet)
    if analysis_type == "new_company":
        try:
            from models.new_company import NewCompanyProjectionReport
            from services.new_company_deliverables_service import NewCompanyDeliverablesService

            projection = NewCompanyDeliverablesService.email_projection(
                NewCompanyProjectionReport.model_validate(output_service.read_json(analysis_id, "new_company_projection_report.json")), fy, fq
            )
        except Exception:  # noqa: BLE001 - the projection lines are optional (Q2 and Q3 only)
            projection = None

    fields = sheet.fields(ticker) if sheet else SheetFields()
    if match.pit and sheet:
        # a back-test keeps today's classification only: market cap, TBV/P and status are the period's own (workbook and filter)
        today = fields
        fields = SheetFields(classified_as=today.classified_as)
        notes.append("Classified as is today's value from the daily sheet (the period's own is not kept)")
    fields.pe10_status = custom_filter_status(match.crf, ticker)       # "Current Status with PE10 Percentile" of the custom run filter
    base = email_deliverable_name(fiscal_year=fy or 0, ticker=ticker, fiscal_quarter=fq)
    biz = None
    if analysis_type == "new_company":
        biz = BusinessDocService().produce_safe(
            ticker=ticker, company=company, workbook_path=excel, output_dir=out_dir,
            base_name=business_deliverable_name(fiscal_year=fy or 0, ticker=ticker, fiscal_quarter=fq),
            search_dir=out_dir, sheet_fields=(SheetFields(classified_as=sheet.fields(ticker).classified_as) if match.pit else sheet.fields(ticker)) if sheet else None,
            fiscal_year=fy, fiscal_quarter=fq, analyst_notes=analyst,
        )
    email = EmailDraftService().produce_safe(
        analysis_type=analysis_type, ticker=ticker, company=company, workbook_path=excel, output_dir=out_dir,
        sheet=fields, base_name=base, fiscal_year=fy, fiscal_quarter=fq, projection=projection,
        highlights=(analyst or {}).get("highlights"),
        basis_note=(f"Per-share figures (price, entry prices, EPS) are on the workbook's share basis, adjusted for later splits "
                    f"(x{match.pit['share_factor']:g}); the prices seen at the time were x{match.pit['share_factor']:g} these.")
        if match.pit and match.pit.get("share_factor", 1.0) != 1.0 else None,
        attachments=[excel] + ([Path(biz["path"])] if biz and biz.get("path") else []),
    )
    results.mkdir(parents=True, exist_ok=True)
    final = {}
    files = [("excel", excel)]
    if email.get("error"):
        notes.append(f"email draft not produced: {email['error']}")
    else:
        files += [("email_txt", Path(email["text_path"])), ("email_eml", Path(email["eml_path"]))]
    if biz is not None:
        if biz.get("path"):
            files.append(("business_doc", Path(biz["path"])))
        else:
            notes.append(f"business document not produced: {biz.get('error')}")
    for label, src in files:
        dest = results / src.name
        shutil.copy2(src, dest)
        final[label] = str(dest)
    summary.update(final)
    summary["authorized"] = deliv.get("authorized")
    if sheet and not sheet.has(ticker):
        notes.append("ticker not found in the daily sheet: Classified as and Status with PE10 are marked in the email")
    return summary


def main() -> int:
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("jobs", nargs="+", help="TICKER:analysis_type[:period], for example WWD:new_company or AMAT:quarterly_update:2025Q1")
    ap.add_argument("--work", default=str(DOWNLOADS / "HAP work"))
    ap.add_argument("--results", default=str(DOWNLOADS / "HAP results"))
    ap.add_argument("--point-in-time", choices=("auto", "off"), default="auto",
                    help="auto: when the custom run filter has history later than the period, back-test as of that period (default)")
    ap.add_argument("--lease-rate", action="append", default=[], metavar="TICKER=RATE",
                    help="analyst-confirmed lease discount rate (for example TPL=0.05), used only when HAP finds no evidence for one")
    ap.add_argument("--notes", default=None, metavar="FILE.json",
                    help='analyst commentary per run, {"TPL 2020Q3": {"highlights": "...", "quarter_title": "...", "quarter_notes": ["..."]}}: '
                         "the highlights go into the email, the quarter notes into the Business document")
    ap.add_argument("--sheet-url", default=None, help="published .xlsx link of the Google Sheet (or set HAP_SHEET_URL); falls back to the newest download")
    args = ap.parse_args()
    work, results = Path(args.work), Path(args.results)
    notes = json.loads(Path(args.notes).read_text(encoding="utf-8")) if args.notes else {}
    lease_rates: dict[str, float] = {}
    for item in args.lease_rate:
        key, _, val = item.partition("=")
        lease_rates[key.strip().upper()] = float(val)

    from services.daily_sheet_service import open_daily_sheet

    sheet, source = open_daily_sheet(work, args.sheet_url or os.environ.get("HAP_SHEET_URL"))
    print(f"Daily sheet: {source if sheet else 'NOT FOUND (email fields will be marked)'}")

    report = []
    for job in args.jobs:
        ticker, _, rest = job.partition(":")
        kind, _, period_text = rest.partition(":")
        ticker, kind = ticker.strip().upper(), kind.strip().lower()
        period = None
        if period_text.strip():
            try:
                period = parse_period(period_text)
            except ValueError as exc:
                report.append({"ticker": ticker, "problem": str(exc)})
                continue
        if kind not in TYPES:
            report.append({"ticker": ticker, "problem": f"unknown analysis type '{kind}' (use {', '.join(sorted(TYPES))})"})
            continue
        match = match_files(work, ticker, kind, period, standin_dir=results / "_standin_filters")
        if match.problems:
            report.append({"ticker": ticker, "type": kind, "problem": "; ".join(match.problems)})
            print(f"[{ticker}] skipped: {report[-1]['problem']}")
            continue
        print(f"[{ticker}] {kind}: {match.workbook.name} + {match.crf.name}" + (f" (previous {match.previous.name})" if match.previous else ""), flush=True)
        apply_point_in_time(match, ticker, results, args.point_in_time)
        try:
            result = run_one(ticker, kind, match, sheet, results, lease_rates.get(ticker),
                             notes.get(f"{ticker} {_period_label(_period_key(match.workbook))}") if match.workbook else None)     # the daily sheet holds today's fields
        except Exception as exc:  # noqa: BLE001 - one company must not stop the batch
            traceback.print_exc()
            result = {"ticker": ticker, "type": kind, "problem": f"{type(exc).__name__}: {exc}"}
        finally:
            from services.market_price_service import MarketPriceService

            MarketPriceService.clear_point_in_time(ticker)
        if match.pit:
            result["point_in_time"] = match.pit
        report.append(result)
        print(f"[{ticker}] {result.get('status', 'FAILED')}: {result.get('problem') or ', '.join(result.get('notes') or ['done'])}", flush=True)

    results.mkdir(parents=True, exist_ok=True)
    from datetime import datetime

    summary_path = results / f"_batch_summary_{datetime.now():%Y-%m-%d_%H%M%S}.json"
    summary_path.write_text(json.dumps(report, indent=2, default=str), encoding="utf-8")
    print("\nSummary written to", summary_path)
    return 0 if all("problem" not in r for r in report) else 1


if __name__ == "__main__":
    raise SystemExit(main())
