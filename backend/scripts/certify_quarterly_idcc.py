"""IDCC Q2 Quarterly Update certification (retrieval, carry-forward, ROIC, valuation)."""

from __future__ import annotations

import json
import shutil
import sys
from datetime import datetime, timezone
from pathlib import Path

BACKEND = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(BACKEND))

from openpyxl import load_workbook

from services.circular_reference_service import CircularReferenceService
from services.excel_recalc_service import ExcelRecalcService, genuine_excel_com_recalc
from services.quarterly_carry_forward_service import QuarterlyCarryForwardService
from services.quarterly_presentation_service import QuarterlyPresentationService
from services.quarterly_projection_service import QuarterlyProjectionService
from services.quarterly_valuation_service import QuarterlyValuationService
from services.formula_utils import is_formula

SOURCE_ANALYSIS = "275d5e21-42e2-4a39-9f78-be9fa48c1b44"
CERT_ID = "quarterly-idcc-q2-hardening"


def _json(path: Path):
    if not path.exists():
        return None
    return json.loads(path.read_text(encoding="utf-8"))


def _blank_statement_body(ws, start=11, end=130) -> int:
    cleared = 0
    for r in range(start, end + 1):
        for c in range(1, 12):
            cell = ws.cell(r, c)
            if is_formula(cell.value):
                continue
            if cell.value is not None:
                cell.value = None
                cleared += 1
    return cleared


def _fill_rgb(cell) -> str:
    fill = cell.fill
    if fill is None or fill.fgColor is None:
        return ""
    return str(fill.fgColor.rgb or "")


def run() -> dict:
    root = BACKEND / "storage"
    src_uploads = root / "uploads" / SOURCE_ANALYSIS
    src_outputs = root / "outputs" / SOURCE_ANALYSIS
    cert_dir = root / "outputs" / CERT_ID
    cert_dir.mkdir(parents=True, exist_ok=True)

    tmpl = src_uploads / "prefilled_workbook.xlsx"
    prev = src_uploads / "previous_workbook.xlsx"
    facts_path = src_outputs / "company_facts.json"
    report: dict = {
        "analysis_id": CERT_ID,
        "source_analysis": SOURCE_ANALYSIS,
        "ticker": "IDCC",
        "generated_at": datetime.now(timezone.utc).isoformat(),
        "gates": {},
    }

    if not tmpl.exists() or not prev.exists():
        report["status"] = "BLOCKED"
        report["error"] = "IDCC template or previous workbook missing"
        return report

    working = cert_dir / "working_template.xlsx"
    shutil.copy2(tmpl, working)
    wb = load_workbook(working)
    try:
        cleared = 0
        for name in ("Last Quarter IS Standardized", "Last Quarter CF Standardized"):
            if name in wb.sheetnames:
                cleared += _blank_statement_body(wb[name])
        report["bloomberg_is_cf_cleared_cells"] = cleared
        wb.save(working)
    finally:
        wb.close()

    company_facts = _json(facts_path) or {}
    carried = cert_dir / "carried.xlsx"
    carry = QuarterlyCarryForwardService().apply(
        analysis_id=CERT_ID,
        ticker="IDCC",
        new_template_path=working,
        previous_workbook_path=prev,
        destination_path=carried,
    )
    (cert_dir / "quarterly_carry_forward_report.json").write_text(
        carry.model_dump_json(indent=2), encoding="utf-8"
    )
    i_entries = [
        e.model_dump()
        for e in carry.entries
        if e.sheet == "Last Quarter BS Standardized" and e.cell in {f"I{r}" for r in range(27, 33)}
    ]
    report["i27_i32"] = i_entries
    carry_ok = all(e["final_action"] == "CARRY_FORWARD" for e in i_entries) and len(i_entries) == 6
    report["gates"]["carry_forward"] = {
        "result": "PASS" if carry_ok else "FAIL",
        "evidence": str(cert_dir / "quarterly_carry_forward_report.json"),
    }

    presented = cert_dir / "completed_workbook.xlsx"
    shutil.copy2(carried, presented)
    pres = QuarterlyPresentationService().plan_and_apply(
        analysis_id=CERT_ID,
        ticker="IDCC",
        source_workbook_path=carried,
        destination_workbook_path=presented,
        company_facts=company_facts,
        already_copied=True,
    )
    (cert_dir / "quarterly_presentation_report.json").write_text(
        pres.model_dump_json(indent=2), encoding="utf-8"
    )
    retrieval_ok = True
    retrieval_notes = []
    for stmt in pres.statements:
        if stmt.sheet.endswith("IS Standardized") or stmt.sheet.endswith("CF Standardized"):
            populated = bool(stmt.sec_rows_introduced or stmt.yahoo_rows_introduced or stmt.rows_filled)
            if not populated and stmt.unresolved_facts:
                retrieval_notes.append(f"{stmt.sheet}: unresolved {len(stmt.unresolved_facts)}")
            elif not populated:
                retrieval_ok = False
                retrieval_notes.append(f"{stmt.sheet}: blank with no unresolved registry")
            else:
                retrieval_notes.append(
                    f"{stmt.sheet}: {stmt.decision} primary={stmt.data_source_primary} "
                    f"introduced={len(stmt.sec_rows_introduced or stmt.yahoo_rows_introduced)}"
                )
    report["gates"]["retrieval"] = {
        "result": "PASS" if retrieval_ok else "FAIL",
        "notes": retrieval_notes,
        "evidence": str(cert_dir / "quarterly_presentation_report.json"),
    }

    proj = QuarterlyProjectionService().apply(
        analysis_id=CERT_ID,
        ticker="IDCC",
        workbook_path=presented,
        previous_workbook_path=prev,
        fiscal_quarter=2,
    )
    (cert_dir / "quarterly_projection_report.json").write_text(
        proj.model_dump_json(indent=2), encoding="utf-8"
    )
    wb = load_workbook(presented)
    try:
        ic = wb["IC & NOPAT & ROIC "] if "IC & NOPAT & ROIC " in wb.sheetnames else None
        yellow_ok = True
        pct_ok = True
        if ic is not None:
            for row in range(1, 26):
                for col in (13, 14):
                    if "FFFF00" not in _fill_rgb(ic.cell(row, col)).upper():
                        yellow_ok = False
            for addr in ("M23", "M24", "M25", "N4"):
                if ic[addr].number_format != "0.00%":
                    pct_ok = False
            report["m24_stored"] = ic["M24"].value
        else:
            yellow_ok = pct_ok = False
        bs = wb["Last Quarter BS Standardized"]
        report["completed_i27_i32"] = {f"I{r}": bs[f"I{r}"].value for r in range(27, 33)}
        report["i22_formula_preserved"] = bs["I22"].value if "I22" in bs else None
    finally:
        wb.close()
    report["gates"]["roic_format"] = {
        "result": "PASS" if yellow_ok and pct_ok else "FAIL",
        "yellow": yellow_ok,
        "percent": pct_ok,
        "evidence": str(cert_dir / "quarterly_projection_report.json"),
    }

    val_report, er_rep, judge, circular, recalc, analytical = QuarterlyValuationService().apply(
        analysis_id=CERT_ID,
        ticker="IDCC",
        workbook_path=presented,
        previous_workbook_path=prev,
        template_path=tmpl,
        company_facts=company_facts,
        fiscal_year=2026,
        fiscal_quarter=2,
        cache_dir=src_outputs / "sec_cache",
    )
    (cert_dir / "quarterly_valuation_report.json").write_text(
        val_report.model_dump_json(indent=2), encoding="utf-8"
    )
    (cert_dir / "quarterly_analyst_judgment_report.json").write_text(
        judge.model_dump_json(indent=2), encoding="utf-8"
    )
    (cert_dir / "quarterly_expected_return_judgment_report.json").write_text(
        er_rep.model_dump_json(indent=2), encoding="utf-8"
    )
    (cert_dir / "quarterly_circular_reference_report.json").write_text(
        circular.model_dump_json(indent=2), encoding="utf-8"
    )
    (cert_dir / "quarterly_excel_recalc_report.json").write_text(
        json.dumps(
            {
                "status": recalc.status,
                "method": recalc.method,
                "com_invoked": recalc.com_invoked,
                "summary": recalc.summary,
                "error": recalc.error,
            },
            indent=2,
        ),
        encoding="utf-8",
    )
    if analytical is not None:
        (cert_dir / "quarterly_analytical_research_report.json").write_text(
            analytical.model_dump_json(indent=2), encoding="utf-8"
        )

    orig_ok = val_report.original_assumptions_preserved
    report["gates"]["valuation_parity"] = {
        "result": "PASS" if orig_ok and val_report.er_decision else "FAIL",
        "er": val_report.er_decision,
        "oe": val_report.oe_decision,
        "graham": val_report.graham_decision,
        "oe_base": val_report.normalized_base_decision,
        "evidence": str(cert_dir / "quarterly_valuation_report.json"),
    }
    com_ok = genuine_excel_com_recalc(recalc)
    report["gates"]["excel_recalc"] = {
        "result": "PASS" if com_ok else ("BLOCKED" if recalc.status == "UNAVAILABLE" else "FAIL"),
        "status": recalc.status,
        "method": recalc.method,
        "evidence": str(cert_dir / "quarterly_excel_recalc_report.json"),
    }
    circ_ok = circular.hap_introduced == []
    report["gates"]["no_hap_circular"] = {
        "result": "PASS" if circ_ok else "FAIL",
        "hap_introduced": len(circular.hap_introduced),
        "pre_existing": len(circular.pre_existing),
        "evidence": str(cert_dir / "quarterly_circular_reference_report.json"),
    }
    report["gates"]["provenance"] = {
        "result": "PASS" if any(s.data_source_primary for s in pres.statements) else "FAIL",
        "evidence": str(cert_dir / "quarterly_presentation_report.json"),
    }

    # Cell-diff: template formulas on historical IC L column vs completed
    src_wb = load_workbook(tmpl, data_only=False)
    out_wb = load_workbook(presented, data_only=False)
    diffs = []
    preserved = 0
    try:
        if "Expected Returns & Buybacks" in src_wb.sheetnames:
            for addr in ("A11", "C5", "E14"):
                if src_wb["Expected Returns & Buybacks"][addr].value != out_wb["Expected Returns & Buybacks"][addr].value:
                    # HAP_ANALYSIS may add adjacent columns; original cells should match unless blank fill
                    sv = src_wb["Expected Returns & Buybacks"][addr].value
                    ov = out_wb["Expected Returns & Buybacks"][addr].value
                    if is_formula(sv) and sv != ov:
                        diffs.append({"sheet": "Expected Returns & Buybacks", "cell": addr, "src": sv, "out": ov})
                    elif sv == ov:
                        preserved += 1
                else:
                    preserved += 1
        if "IC & NOPAT & ROIC " in src_wb.sheetnames:
            for row in range(3, 26):
                src_l = src_wb["IC & NOPAT & ROIC "].cell(row, 12).value
                out_l = out_wb["IC & NOPAT & ROIC "].cell(row, 12).value
                if src_l != out_l:
                    diffs.append({"sheet": "IC & NOPAT & ROIC ", "cell": f"L{row}", "src": src_l, "out": out_l})
                else:
                    preserved += 1
    finally:
        src_wb.close()
        out_wb.close()
    (cert_dir / "cell_diff_original_assumptions.json").write_text(
        json.dumps({"preserved": preserved, "diffs": diffs}, indent=2, default=str),
        encoding="utf-8",
    )
    report["gates"]["no_regression_originals"] = {
        "result": "PASS" if not diffs else "FAIL",
        "preserved": preserved,
        "diff_count": len(diffs),
        "evidence": str(cert_dir / "cell_diff_original_assumptions.json"),
    }

    results = [g["result"] for g in report["gates"].values()]
    if "FAIL" in results:
        report["status"] = "NOT_READY"
    elif "BLOCKED" in results:
        report["status"] = "BLOCKED"
    else:
        report["status"] = "READY_FOR_PRODUCTION" if all(r == "PASS" for r in results) else "BLOCKED"
    (cert_dir / "quarterly_certification_report.json").write_text(
        json.dumps(report, indent=2, default=str), encoding="utf-8"
    )
    md = [f"# Quarterly Update certification — IDCC Q2", "", f"Status: **{report['status']}**", ""]
    for name, gate in report["gates"].items():
        md.append(f"- {name}: **{gate['result']}** — `{gate.get('evidence', '')}`")
    (cert_dir / "QUARTERLY_UPDATE_PRODUCTION_CERTIFICATION.md").write_text(
        "\n".join(md) + "\n", encoding="utf-8"
    )
    report["markdown"] = str(cert_dir / "QUARTERLY_UPDATE_PRODUCTION_CERTIFICATION.md")
    return report


if __name__ == "__main__":
    out = run()
    print(json.dumps({"status": out.get("status"), "gates": out.get("gates")}, indent=2, default=str))
