"""IDCC Quarterly Update fill-in-place certification.

Uses the untouched Bloomberg workbook. Does not clear quarterly statement
tabs. Writes a new analysis directory and leaves quarterly-idcc-q2-hardening
unchanged.
"""

from __future__ import annotations

import hashlib
import json
import shutil
import subprocess
import sys
from datetime import datetime, timezone
from pathlib import Path

BACKEND = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(BACKEND))

from openpyxl import load_workbook

from models.analysis import AnalysisFiles, CreateAnalysisRequest, UploadedFileMetadata
from models.common import utc_now_iso
from models.quarterly_presentation import QuarterlyStatementKind
from pipeline.orchestrator import PipelineOrchestrator
from services.analysis_service import AnalysisService
from services.excel_recalc_service import genuine_excel_com_recalc
from services.file_service import FileService
from services.formula_utils import is_formula
from services.quarterly_health_service import assess_statement_health, decide_presentation

SOURCE_ANALYSIS = "275d5e21-42e2-4a39-9f78-be9fa48c1b44"
SOURCE_UPLOADS = BACKEND / "storage" / "uploads" / SOURCE_ANALYSIS
SOURCE_WORKBOOK = SOURCE_UPLOADS / "prefilled_workbook.xlsx"
PREVIOUS_WORKBOOK = SOURCE_UPLOADS / "previous_workbook.xlsx"
SOURCE_CRF = SOURCE_UPLOADS / "Custom_Run_Filter_2026-09-09-_14-38_-IDCC.xlsx"
RECONSTRUCTION_DIR = BACKEND / "storage" / "outputs" / "quarterly-idcc-q2-hardening"
STATEMENT_SHEETS = ("Last Quarter IS Standardized", "Last Quarter CF Standardized")


def _sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def _git_sha() -> str:
    try:
        return subprocess.check_output(
            ["git", "rev-parse", "HEAD"], cwd=BACKEND.parent, text=True
        ).strip()
    except (OSError, subprocess.CalledProcessError):
        return "unknown"


def _close(value, expected: float, tolerance: float) -> bool:
    if isinstance(value, bool) or not isinstance(value, (int, float)):
        return False
    return abs(float(value) - expected) <= tolerance


def _current_period_values(path: Path) -> dict:
    """FY2026 Q2 current-column amounts, not prior-year comparatives."""
    book = load_workbook(path, data_only=False)
    try:
        ws = book["Last Quarter IS Standardized"]
        actual = {coord: ws[coord].value for coord in (
            "C11", "C12", "C32", "C60", "C67", "C71", "C73",
            "G11", "G32", "G60", "G67", "G71",
        )}
    finally:
        book.close()
    checks = {
        "revenue_standalone_not_prior_year": _close(actual["C11"], 260.17, 0.02) and not _close(actual["C11"], 300.596, 0.05),
        "sales_component_preserved": _close(actual["C12"], 110, 0.02),
        "operating_income_not_prior_year": _close(actual["C32"], 139.24, 0.02) and not _close(actual["C32"], 205.427, 0.05),
        "net_income_current_quarter": _close(actual["C60"], 116.37, 0.02) and not _close(actual["C60"], 115.602, 0.05),
        "basic_eps_gaap": _close(actual["C67"], 4.51, 0.001),
        "diluted_eps_gaap_not_prior_year": _close(actual["C71"], 3.4, 0.001) and not _close(actual["C71"], 5.35, 0.001),
        "adjusted_diluted_eps_preserved": _close(actual["C73"], 3.38, 0.001),
        "revenue_ytd": _close(actual["G11"], 465.586, 0.02),
        "operating_income_ytd": _close(actual["G32"], 221.5, 0.02),
        "net_income_ytd": _close(actual["G60"], 191.701, 0.02),
        "basic_eps_ytd": _close(actual["G67"], 7.44, 0.001),
        "diluted_eps_ytd": _close(actual["G71"], 5.51, 0.001),
    }
    return {"actual": actual, "checks": checks, "ok": all(checks.values())}


def _num(value) -> float | None:
    if isinstance(value, bool) or not isinstance(value, (int, float)):
        return None
    return float(value)


def _preflight(path: Path) -> dict:
    """Stop unless the untouched workbook still matches the investigated state."""
    wb = load_workbook(path, data_only=False)
    try:
        income = wb["Last Quarter IS Standardized"]
        cash_flow = wb["Last Quarter CF Standardized"]
        is_health = assess_statement_health(wb, QuarterlyStatementKind.INCOME)
        cf_health = assess_statement_health(wb, QuarterlyStatementKind.CASH_FLOW)
        is_decision = decide_presentation(is_health)
        cf_decision = decide_presentation(cf_health)
        a11 = str(income["A11"].value or "")
        a12 = str(income["A12"].value or "")
        cf_a11 = str(cash_flow["A11"].value or "")
        c11 = _num(income["C11"].value)
        c12 = _num(income["C12"].value)
        problems: list[str] = []
        if "revenue" not in a11.lower() or a11.strip().startswith("+"):
            problems.append(f"A11 is {a11!r}, expected the Revenue total")
        if c11 is None or abs(c11 - 260.17) > 0.02:
            problems.append(f"C11 is {income['C11'].value!r}, expected 260.17")
        if "sales" not in a12.lower():
            problems.append(f"A12 is {a12!r}, expected Sales & Services Revenue")
        if c12 is None or abs(c12 - 110) > 0.02:
            problems.append(f"C12 is {income['C12'].value!r}, expected 110")
        if income["G11"].value not in (None, ""):
            problems.append(f"G11 is {income['G11'].value!r}, expected blank YTD")
        if "operating" not in cf_a11.lower():
            problems.append(f"CF A11 is {cf_a11!r}")
        if cash_flow["C11"].value not in (None, ""):
            problems.append(f"CF C11 is {cash_flow['C11'].value!r}, expected blank")
        if is_decision.value != "BLOOMBERG_FILL_GAPS":
            problems.append(f"IS decision {is_decision.value}")
        if cf_decision.value != "BLOOMBERG_FILL_GAPS":
            problems.append(f"CF decision {cf_decision.value}")
        if is_health.populated_rows < 35 or is_health.expected_mapped_rows < 50:
            problems.append(
                f"IS inventory {is_health.populated_rows}/{is_health.expected_mapped_rows}"
            )
        if cf_health.populated_rows != 0:
            problems.append(f"CF populated {cf_health.populated_rows}, expected 0")
        labels = {
            "income": [
                str(income.cell(row, 1).value)
                for row in range(11, 81)
                if income.cell(row, 1).value not in (None, "")
            ],
            "cash_flow": [
                str(cash_flow.cell(row, 1).value)
                for row in range(11, 81)
                if cash_flow.cell(row, 1).value not in (None, "")
            ],
        }
        formulas = _formula_map(wb, STATEMENT_SHEETS)
        period = {
            "C5": income["C5"].value,
            "C9": income["C9"].value,
            "A7": income["A7"].value,
            "A10": income["A10"].value,
            "C10": income["C10"].value,
            "G10": income["G10"].value,
        }
        inventory = {
            "sha256": _sha256(path),
            "is_decision": is_decision.value,
            "cf_decision": cf_decision.value,
            "is_populated": is_health.populated_rows,
            "is_expected": is_health.expected_mapped_rows,
            "cf_populated": cf_health.populated_rows,
            "cf_expected": cf_health.expected_mapped_rows,
            "is_reason": is_health.reason,
            "cf_reason": cf_health.reason,
            "fingerprint": {
                "A11": a11,
                "C11": c11,
                "A12": a12,
                "C12": c12,
                "G11": income["G11"].value,
                "CF_A11": cf_a11,
                "CF_C11": cash_flow["C11"].value,
            },
            "period_cells": {key: _jsonable(value) for key, value in period.items()},
            "labels": labels,
            "formula_count": len(formulas),
            "problems": problems,
        }
        return inventory
    finally:
        wb.close()


def _jsonable(value):
    if is_formula(value) or isinstance(value, str):
        return value
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        return value
    return None if value is None else str(value)


def _formula_text(value) -> str:
    text = getattr(value, "text", None)
    return str(text if text is not None else value)


def _formula_map(wb, sheets: tuple[str, ...]) -> dict[str, str]:
    found: dict[str, str] = {}
    for name in sheets:
        ws = wb[name]
        for row in ws.iter_rows(min_row=1, max_row=ws.max_row or 1, max_col=min(ws.max_column or 1, 16)):
            for cell in row:
                if is_formula(cell.value) or getattr(cell.value, "text", None):
                    found[f"{name}!{cell.coordinate}"] = _formula_text(cell.value)
    return found


def _statement_value_changes(source: Path, completed: Path) -> list[dict]:
    """Numeric body changes on the quarterly statement tabs."""
    before = load_workbook(source, data_only=False)
    after = load_workbook(completed, data_only=False)
    changes: list[dict] = []
    try:
        for name in STATEMENT_SHEETS:
            left, right = before[name], after[name]
            for row in range(11, 131):
                label = left.cell(row, 1).value
                for col in (3, 7):
                    old = left.cell(row, col).value
                    new = right.cell(row, col).value
                    if old == new or is_formula(old) or is_formula(new):
                        continue
                    changes.append(
                        {
                            "cell": f"{name}!{left.cell(row, col).coordinate}",
                            "label": label,
                            "before": old,
                            "after": new,
                        }
                    )
    finally:
        before.close()
        after.close()
    return changes


def _label_map(path: Path) -> dict[str, str]:
    wb = load_workbook(path, data_only=False)
    try:
        found: dict[str, str] = {}
        for name in STATEMENT_SHEETS:
            ws = wb[name]
            for row in range(11, 131):
                value = ws.cell(row, 1).value
                if value not in (None, ""):
                    found[f"{name}!A{row}"] = str(value)
        return found
    finally:
        wb.close()


def _operating_income_trace(path: Path) -> dict:
    wb = load_workbook(path, data_only=False)
    try:
        income = wb["Last Quarter IS Standardized"]
        target = None
        for row in range(11, 131):
            label = " ".join(str(income.cell(row, 1).value or "").lower().split())
            if label.startswith("+") or label.startswith("-"):
                continue
            if "operating income" in label:
                target = f"C{row}"
                break
        refs: list[str] = []
        if target:
            needle = target
            for ws in wb.worksheets:
                for row in ws.iter_rows(
                    min_row=1, max_row=min(ws.max_row or 1, 200), max_col=min(ws.max_column or 1, 20)
                ):
                    for cell in row:
                        if is_formula(cell.value) and needle in str(cell.value) and "Last Quarter IS" in str(cell.value):
                            refs.append(f"{ws.title}!{cell.coordinate}={cell.value}")
        return {"value_cell": target, "cross_sheet_formulas": refs[:40]}
    finally:
        wb.close()


def _notes_present(path: Path) -> dict[str, bool]:
    wb = load_workbook(path, data_only=False)
    try:
        found = {}
        for name in (*STATEMENT_SHEETS, "Leases", "R&D", "Expected Returns & Buybacks", "Enterprise Value"):
            if name not in wb.sheetnames:
                found[name] = False
                continue
            ws = wb[name]
            found[name] = any(
                ws.cell(row, 1).value == "HAP ANALYSIS — NOTES"
                for row in range(1, (ws.max_row or 1) + 1)
            )
        return found
    finally:
        wb.close()


def _copy_upload(src: Path, dest: Path) -> UploadedFileMetadata:
    dest.parent.mkdir(parents=True, exist_ok=True)
    shutil.copy2(src, dest)
    return UploadedFileMetadata(
        filename=src.name,
        stored_filename=dest.name,
        size_bytes=dest.stat().st_size,
        uploaded_at=utc_now_iso(),
    )


def _gate(ok: bool, evidence: str, **extra) -> dict:
    payload = {"result": "PASS" if ok else "FAIL", "evidence": evidence}
    payload.update(extra)
    return payload


def main() -> int:
    if not SOURCE_WORKBOOK.is_file() or not PREVIOUS_WORKBOOK.is_file() or not SOURCE_CRF.is_file():
        print("IDCC source workbook, previous workbook, or CRF is missing.", file=sys.stderr)
        return 2
    source_sha = _sha256(SOURCE_WORKBOOK)
    inventory = _preflight(SOURCE_WORKBOOK)
    if inventory["problems"]:
        print(json.dumps({"status": "STOPPED", "problems": inventory["problems"]}, indent=2))
        return 2
    if _sha256(SOURCE_WORKBOOK) != source_sha:
        print("Source workbook changed during preflight.", file=sys.stderr)
        return 2

    before_labels = _label_map(SOURCE_WORKBOOK)
    before_formulas = {}
    before_book = load_workbook(SOURCE_WORKBOOK, data_only=False)
    try:
        before_formulas = _formula_map(before_book, STATEMENT_SHEETS)
        before_oi = _operating_income_trace(SOURCE_WORKBOOK)
    finally:
        before_book.close()

    analysis_service = AnalysisService()
    file_service = FileService()
    created = analysis_service.create(
        CreateAnalysisRequest(
            company="InterDigital",
            ticker="IDCC",
            analysis_type="quarterly_update",
        )
    )
    analysis_id = created.analysis_id
    upload_dir = file_service.analysis_upload_dir(analysis_id)
    created.files = AnalysisFiles(
        prefilled_workbook=_copy_upload(SOURCE_WORKBOOK, upload_dir / "prefilled_workbook.xlsx"),
        previous_workbook=_copy_upload(PREVIOUS_WORKBOOK, upload_dir / "previous_workbook.xlsx"),
        custom_run_filter=_copy_upload(SOURCE_CRF, upload_dir / SOURCE_CRF.name),
    )
    created.status = "uploaded"
    analysis_service.save(created)
    if _sha256(SOURCE_WORKBOOK) != source_sha:
        print("Source workbook was modified while copying.", file=sys.stderr)
        return 2

    print(f"Created analysis {analysis_id}. Running Quarterly Update without clearing tabs...")
    analysis = PipelineOrchestrator().run(analysis_id)
    out_dir = BACKEND / "storage" / "outputs" / analysis_id

    def _read(name: str):
        path = out_dir / name
        if not path.exists():
            return None
        return json.loads(path.read_text(encoding="utf-8"))

    presentation = _read("quarterly_presentation_report.json") or {}
    statements = {item.get("statement"): item for item in presentation.get("statements") or []}
    income = statements.get("quarterly_income_statement") or {}
    cash_flow = statements.get("quarterly_cash_flow") or {}
    recalc = _read("quarterly_excel_recalc_report.json") or {}
    circular = _read("quarterly_circular_reference_report.json") or {}
    valuation = _read("quarterly_valuation_report.json") or {}
    projection = _read("quarterly_projection_report.json") or {}
    deliverables = _read("quarterly_deliverables_report.json") or {}
    completed = out_dir / "completed_workbook.xlsx"
    after_labels = _label_map(completed) if completed.exists() else {}
    label_changes = [
        {"cell": key, "before": before, "after": after_labels.get(key)}
        for key, before in before_labels.items()
        if after_labels.get(key) != before
    ]
    after_formulas = {}
    if completed.exists():
        book = load_workbook(completed, data_only=False)
        try:
            after_formulas = _formula_map(book, STATEMENT_SHEETS)
        finally:
            book.close()
    formula_changes = [
        {"cell": key, "before": before, "after": after_formulas.get(key)}
        for key, before in before_formulas.items()
        if after_formulas.get(key) != before
    ]
    value_changes = _statement_value_changes(SOURCE_WORKBOOK, completed) if completed.exists() else []
    after_oi = _operating_income_trace(completed) if completed.exists() else {}
    notes = _notes_present(completed) if completed.exists() else {}
    unresolved = presentation.get("unresolved_dependencies") or []
    dependency_diff = presentation.get("dependency_diff") or []
    component_fills = [
        row
        for row in list(income.get("rows_filled") or []) + list(cash_flow.get("rows_filled") or [])
        if "+" in str(row) and any(token in str(row).lower() for token in ("concept=revenue", "concept=operating_income"))
    ]
    excel_name = deliverables.get("excel_filename") or ""
    fiscal_year = projection.get("fiscal_year")
    fiscal_quarter = projection.get("fiscal_quarter")
    expected_excel = (
        f"{fiscal_year} Q{fiscal_quarter} IDCC Quarterly Update.xlsx"
        if fiscal_year and fiscal_quarter
        else ""
    )
    com_ok = genuine_excel_com_recalc(
        type("Recalc", (), {"status": recalc.get("status"), "method": recalc.get("method"), "com_invoked": recalc.get("com_invoked")})()
    )
    word_path = Path(str(deliverables.get("word_path") or ""))
    word_ok = bool(deliverables.get("word_path") and word_path.is_file() and not unresolved)
    period_values = _current_period_values(completed) if completed.exists() else {"ok": False, "actual": {}, "checks": {}}
    gates = {
        "pre_run_taxonomy": _gate(True, "preflight fingerprint"),
        "income_fill_gaps": _gate(str(income.get("decision")) == "BLOOMBERG_FILL_GAPS", str(out_dir / "quarterly_presentation_report.json"), decision=income.get("decision")),
        "cash_flow_fill_gaps": _gate(str(cash_flow.get("decision")) == "BLOOMBERG_FILL_GAPS", str(out_dir / "quarterly_presentation_report.json"), decision=cash_flow.get("decision")),
        "no_taxonomy_replacement": _gate(
            not income.get("rows_superseded") and not cash_flow.get("rows_superseded") and not label_changes,
            str(out_dir / "cell_diff_fill_in_place.json"),
            label_changes=len(label_changes),
        ),
        "component_rows": _gate(not component_fills, str(out_dir / "quarterly_presentation_report.json"), hits=component_fills[:8]),
        "dependencies": _gate(not unresolved and not dependency_diff, str(out_dir / "quarterly_presentation_report.json"), unresolved=len(unresolved), changed=len(dependency_diff)),
        "statement_formulas": _gate(not formula_changes, str(out_dir / "cell_diff_fill_in_place.json"), changed=len(formula_changes)),
        "operating_income_trace": _gate(
            bool(after_oi.get("value_cell")) and after_oi.get("value_cell") == before_oi.get("value_cell"),
            str(out_dir / "operating_income_trace.json"),
            before=before_oi.get("value_cell"),
            after=after_oi.get("value_cell"),
        ),
        "excel_recalc": _gate(com_ok, str(out_dir / "quarterly_excel_recalc_report.json"), method=recalc.get("method"), status=recalc.get("status")),
        "no_hap_circular": _gate(circular.get("hap_introduced") == [], str(out_dir / "quarterly_circular_reference_report.json"), hap=len(circular.get("hap_introduced") or [])),
        "assumptions": _gate(bool(valuation.get("original_assumptions_preserved")), str(out_dir / "quarterly_valuation_report.json")),
        "notes": _gate(notes.get("Last Quarter IS Standardized") and notes.get("Last Quarter CF Standardized"), str(completed), sheets=notes),
        "excel_name": _gate(bool(expected_excel) and excel_name == expected_excel, str(out_dir / "quarterly_deliverables_report.json"), actual=excel_name, expected=expected_excel),
        "word": _gate(word_ok, deliverables.get("word_path") or str(out_dir / "quarterly_deliverables_report.json")),
        "current_period_eps_and_net_income": _gate(
            bool(period_values.get("ok")),
            str(completed),
            checks=period_values.get("checks"),
            actual=period_values.get("actual"),
        ),
    }
    results = [gate["result"] for gate in gates.values()]
    status = "READY_FOR_PRODUCTION" if results and all(item == "PASS" for item in results) else "NOT_READY"
    if analysis.status not in {"complete", "completed"}:
        status = "NOT_READY"
        gates["pipeline"] = _gate(False, str(out_dir), analysis_status=analysis.status)
    report = {
        "certification": "quarterly_fill_in_place",
        "distinguished_from": str(RECONSTRUCTION_DIR),
        "reconstruction_preserved": RECONSTRUCTION_DIR.is_dir(),
        "source_commit": _git_sha(),
        "generated_at": datetime.now(timezone.utc).isoformat(),
        "analysis_id": analysis_id,
        "analysis_status": analysis.status,
        "source_workbook": str(SOURCE_WORKBOOK),
        "source_sha256": source_sha,
        "source_sha256_after_run": _sha256(SOURCE_WORKBOOK),
        "source_unchanged": _sha256(SOURCE_WORKBOOK) == source_sha,
        "preflight": inventory,
        "presentation_summary": presentation.get("summary"),
        "fiscal_year": fiscal_year,
        "fiscal_quarter": fiscal_quarter,
        "sec_filing_form": income.get("sec_filing_form") or cash_flow.get("sec_filing_form"),
        "sec_filing_period": income.get("sec_filing_period") or cash_flow.get("sec_filing_period"),
        "sec_accession": income.get("sec_accession") or cash_flow.get("sec_accession"),
        "current_period_values": period_values,
        "income_rows_filled": income.get("rows_filled") or [],
        "cash_flow_rows_filled": cash_flow.get("rows_filled") or [],
        "income_discrepancies": income.get("source_discrepancies") or [],
        "cash_flow_discrepancies": cash_flow.get("source_discrepancies") or [],
        "income_derivation": income.get("derivation_notes") or [],
        "cash_flow_derivation": cash_flow.get("derivation_notes") or [],
        "value_changes": value_changes,
        "dependency_diff_count": len(dependency_diff),
        "unresolved_dependencies": len(unresolved),
        "formula_changes": formula_changes[:40],
        "label_changes": label_changes[:40],
        "operating_income_before": before_oi,
        "operating_income_after": after_oi,
        "recalc": recalc,
        "deliverables": deliverables,
        "gates": gates,
        "status": status,
    }
    (out_dir / "cell_diff_fill_in_place.json").write_text(
        json.dumps(
            {
                "label_changes": label_changes,
                "formula_changes": formula_changes,
                "value_changes": value_changes,
            },
            indent=2,
            default=str,
        ),
        encoding="utf-8",
    )
    (out_dir / "operating_income_trace.json").write_text(
        json.dumps({"before": before_oi, "after": after_oi}, indent=2, default=str),
        encoding="utf-8",
    )
    (out_dir / "quarterly_fill_in_place_certification.json").write_text(
        json.dumps(report, indent=2, default=str),
        encoding="utf-8",
    )
    lines = [
        "# IDCC Quarterly Update — fill-in-place certification",
        "",
        "This run did not clear the Bloomberg quarterly tabs.",
        f"The reconstruction certification remains at `{RECONSTRUCTION_DIR}`.",
        "",
        f"Status: **{status}**",
        f"Analysis: `{analysis_id}`",
        f"Source SHA-256: `{source_sha}`",
        f"HEAD: `{report['source_commit']}`",
        f"Period: FY{fiscal_year} Q{fiscal_quarter}",
        f"SEC: {report['sec_filing_form']} {report['sec_filing_period']} {report['sec_accession']}",
        "",
        "This certification is fill-in-place. It is separate from the reconstruction certification, which cleared the statement bodies before the pipeline ran.",
        "The prior fill-in-place certification e6aef478-1fa4-4839-af1d-4cb8a8fcc64c is preserved.",
        "",
    ]
    for name, gate in gates.items():
        lines.append(f"- {name}: **{gate['result']}**")
    (out_dir / "QUARTERLY_FILL_IN_PLACE_CERTIFICATION.md").write_text("\n".join(lines) + "\n", encoding="utf-8")
    print(json.dumps({"status": status, "analysis_id": analysis_id, "gates": {k: v["result"] for k, v in gates.items()}}, indent=2))
    return 0 if status == "READY_FOR_PRODUCTION" else 1


if __name__ == "__main__":
    raise SystemExit(main())
