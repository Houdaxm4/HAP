"""Focused tests for statement validation + analyst review."""

from __future__ import annotations

import hashlib
from pathlib import Path

import pytest
from openpyxl import Workbook

from models.analyst_review import ExplanationStatus, FindingSeverity
from models.statement_validation import StatementValidationDecision
from services.analyst_review_service import AnalystReviewService
from services.statement_validation_service import (
    StatementValidationService,
    compare_values,
)


def _sha(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def _facts_revenue(fy2018: float = 265_595_000_000.0) -> dict:
    return {
        "facts": {
            "us-gaap": {
                "RevenueFromContractWithCustomerExcludingAssessedTax": {
                    "label": "Revenue",
                    "units": {
                        "USD": [
                            {
                                "val": fy2018,
                                "fy": 2018,
                                "fp": "FY",
                                "form": "10-K",
                                "filed": "2018-11-05",
                                "accn": "0000320193-18-000145",
                                "frame": "CY2018",
                                "end": "2018-09-29",
                                "start": "2017-10-01",
                            },
                            {
                                "val": 260_174_000_000.0,
                                "fy": 2019,
                                "fp": "FY",
                                "form": "10-K",
                                "filed": "2019-10-31",
                                "accn": "0000320193-19-000119",
                                "frame": "CY2019",
                                "end": "2019-09-28",
                                "start": "2018-09-30",
                            },
                        ]
                    },
                },
                "GrossProfit": {
                    "units": {
                        "USD": [
                            {
                                "val": 101_839_000_000.0,
                                "fy": 2018,
                                "fp": "FY",
                                "form": "10-K",
                                "filed": "2018-11-05",
                                "accn": "0000320193-18-000145",
                                "frame": "CY2018",
                                "end": "2018-09-29",
                            },
                            {
                                "val": 98_392_000_000.0,
                                "fy": 2019,
                                "fp": "FY",
                                "form": "10-K",
                                "filed": "2019-10-31",
                                "accn": "0000320193-19-000119",
                                "frame": "CY2019",
                                "end": "2019-09-28",
                            },
                        ]
                    }
                },
                "OperatingIncomeLoss": {
                    "units": {
                        "USD": [
                            {
                                "val": 70_898_000_000.0,
                                "fy": 2018,
                                "fp": "FY",
                                "form": "10-K",
                                "filed": "2018-11-05",
                                "accn": "0000320193-18-000145",
                                "frame": "CY2018",
                                "end": "2018-09-29",
                            },
                            {
                                "val": 63_930_000_000.0,
                                "fy": 2019,
                                "fp": "FY",
                                "form": "10-K",
                                "filed": "2019-10-31",
                                "accn": "0000320193-19-000119",
                                "frame": "CY2019",
                                "end": "2019-09-28",
                            },
                        ]
                    }
                },
                "NetIncomeLoss": {
                    "units": {
                        "USD": [
                            {
                                "val": 59_531_000_000.0,
                                "fy": 2018,
                                "fp": "FY",
                                "form": "10-K",
                                "filed": "2018-11-05",
                                "accn": "0000320193-18-000145",
                                "frame": "CY2018",
                                "end": "2018-09-29",
                            },
                            {
                                "val": 55_256_000_000.0,
                                "fy": 2019,
                                "fp": "FY",
                                "form": "10-K",
                                "filed": "2019-10-31",
                                "accn": "0000320193-19-000119",
                                "frame": "CY2019",
                                "end": "2019-09-28",
                            },
                        ]
                    }
                },
                "RestructuringCharges": {
                    "units": {
                        "USD": [
                            {
                                "val": 2_000_000_000.0,
                                "fy": 2019,
                                "fp": "FY",
                                "form": "10-K",
                                "filed": "2019-10-31",
                                "accn": "0000320193-19-000119",
                                "frame": "CY2019",
                                "end": "2019-09-28",
                            }
                        ]
                    }
                },
            }
        }
    }


@pytest.fixture
def mini_workbook(tmp_path: Path) -> Path:
    wb = Workbook()
    ws = wb.active
    ws.title = "Income - GAAP"
    for i, col in enumerate(list("CDEFGHIJKL")):
        ws[f"{col}7"] = f"FY{2016 + i}"
    ws["A9"] = "Revenue"
    ws["E9"] = 265595  # FY2018 exact
    ws["F9"] = 260174  # FY2019 exact
    ws["A19"] = "Gross Profit"
    ws["E19"] = 101839
    ws["F19"] = 98392
    ws["A30"] = "Operating Income (Loss)"
    ws["E30"] = 70898
    ws["F30"] = 63930
    ws["A58"] = "Net Income, GAAP"
    ws["E58"] = 59531
    ws["F58"] = 55256
    # Other required sheets empty-ish
    for name in (
        "Balance Sheet - Standardized",
        "Cash Flow - Standardized",
    ):
        wb.create_sheet(name)
    path = tmp_path / "wb.xlsx"
    wb.save(path)
    wb.close()
    return path


def test_exact_match_validated():
    assert (
        compare_values(265595.0, 265595.0, scale="millions")
        == StatementValidationDecision.VALIDATED
    )


def test_rounding_difference_validated():
    # $0.4M rounding
    assert (
        compare_values(265595.0, 265595.4, scale="millions")
        == StatementValidationDecision.VALIDATED
    )


def test_material_mismatch_discrepancy():
    assert (
        compare_values(265595.0, 215639.0, scale="millions")
        == StatementValidationDecision.DISCREPANCY
    )


def test_workbook_sec_exact_match(mini_workbook: Path):
    report = StatementValidationService().validate(
        analysis_id="v1",
        ticker="AAPL",
        workbook_path=mini_workbook,
        company_facts=_facts_revenue(),
    )
    e2018 = next(
        e
        for e in report.entries
        if e.metric == "Revenue" and e.fiscal_period == "FY2018"
    )
    assert e2018.decision == StatementValidationDecision.VALIDATED
    assert e2018.sec_concept
    assert e2018.accession_number
    assert e2018.fiscal_period == "FY2018"


def test_material_mismatch_in_workbook(mini_workbook: Path, tmp_path: Path):
    from openpyxl import load_workbook

    wb = load_workbook(mini_workbook)
    wb["Income - GAAP"]["E9"] = 215639  # wrong FY
    wb.save(mini_workbook)
    wb.close()
    report = StatementValidationService().validate(
        analysis_id="v2",
        ticker="AAPL",
        workbook_path=mini_workbook,
        company_facts=_facts_revenue(),
    )
    e2018 = next(
        e for e in report.entries if e.metric == "Revenue" and e.fiscal_period == "FY2018"
    )
    assert e2018.decision == StatementValidationDecision.DISCREPANCY


def test_large_yoy_is_analyst_finding_not_discrepancy(mini_workbook: Path):
    # Correct values that decline ~2% — small; force larger drop in workbook+facts
    facts = _facts_revenue(fy2018=300_000_000_000.0)
    facts["facts"]["us-gaap"]["RevenueFromContractWithCustomerExcludingAssessedTax"]["units"][
        "USD"
    ][1]["val"] = 200_000_000_000.0
    from openpyxl import load_workbook

    wb = load_workbook(mini_workbook)
    wb["Income - GAAP"]["E9"] = 300000
    wb["Income - GAAP"]["F9"] = 200000
    wb.save(mini_workbook)
    wb.close()

    validation = StatementValidationService().validate(
        analysis_id="v3",
        ticker="AAPL",
        workbook_path=mini_workbook,
        company_facts=facts,
    )
    # Both periods validate against their SEC values
    assert all(
        e.decision == StatementValidationDecision.VALIDATED
        for e in validation.entries
        if e.metric == "Revenue" and e.fiscal_period in {"FY2018", "FY2019"}
    )
    review = AnalystReviewService().review(
        analysis_id="v3",
        ticker="AAPL",
        validation=validation,
        company_facts=facts,
    )
    rev_findings = [f for f in review.findings if f.metric == "Revenue" and not f.is_data_discrepancy]
    assert rev_findings
    assert rev_findings[0].severity in {FindingSeverity.WATCH, FindingSeverity.MATERIAL}


def test_unusual_expense_analyst_finding(mini_workbook: Path):
    validation = StatementValidationService().validate(
        analysis_id="v4",
        ticker="AAPL",
        workbook_path=mini_workbook,
        company_facts=_facts_revenue(),
    )
    review = AnalystReviewService().review(
        analysis_id="v4",
        ticker="AAPL",
        validation=validation,
        company_facts=_facts_revenue(),
    )
    restructuring = [
        f for f in review.findings if "Restructuring" in f.metric or "restructuring" in f.observation.lower()
    ]
    assert restructuring
    assert restructuring[0].severity == FindingSeverity.MATERIAL
    assert restructuring[0].explanation_status == ExplanationStatus.EXPLAINED_BY_FILING


def test_unexplained_margin_move(mini_workbook: Path):
    from openpyxl import load_workbook

    wb = load_workbook(mini_workbook)
    # Collapse gross margin sharply while keeping SEC-aligned revenue/GP for validation
    # Use workbook-only distortion on GP FY2019 to create margin finding from workbook series
    wb["Income - GAAP"]["F19"] = 20000  # vs revenue 260174 → huge margin drop
    wb.save(mini_workbook)
    wb.close()
    # SEC still has correct GP — will be DISCREPANCY for GP FY2019; still get margin finding from workbook series
    validation = StatementValidationService().validate(
        analysis_id="v5",
        ticker="AAPL",
        workbook_path=mini_workbook,
        company_facts=_facts_revenue(),
    )
    review = AnalystReviewService().review(
        analysis_id="v5",
        ticker="AAPL",
        validation=validation,
        company_facts=_facts_revenue(),
    )
    margins = [f for f in review.findings if f.metric == "Gross Margin"]
    assert margins
    assert margins[0].explanation_status in {
        ExplanationStatus.UNEXPLAINED_REVIEW_REQUIRED,
        ExplanationStatus.LIKELY_EXPLANATION,
    }


def test_validation_does_not_modify_workbook(mini_workbook: Path):
    before = _sha(mini_workbook)
    StatementValidationService().validate(
        analysis_id="v6",
        ticker="AAPL",
        workbook_path=mini_workbook,
        company_facts=_facts_revenue(),
    )
    assert _sha(mini_workbook) == before


def test_review_required_mid_tier_difference():
    # ~$10M on a large base — above $1M rounding, below $25M hard discrepancy
    assert (
        compare_values(265595.0, 265605.0, scale="millions")
        == StatementValidationDecision.REVIEW_REQUIRED
    )


def test_quarterly_period_identity_preserved(tmp_path: Path):
    wb = Workbook()
    wb.active.title = "Income - GAAP"
    for name in (
        "Balance Sheet - Standardized",
        "Cash Flow - Standardized",
        "Last Quarter IS Standardized",
        "Last Quarter BS Standardized",
        "Last Quarter CF Standardized",
    ):
        wb.create_sheet(name)
    ws = wb["Last Quarter IS Standardized"]
    ws["C5"] = "2026 Q3"
    ws["A11"] = "Revenue"
    ws["C11"] = 94930
    path = tmp_path / "q.xlsx"
    wb.save(path)
    wb.close()

    facts = {
        "facts": {
            "us-gaap": {
                "RevenueFromContractWithCustomerExcludingAssessedTax": {
                    "units": {
                        "USD": [
                            {
                                "val": 94_930_000_000.0,
                                "fy": 2026,
                                "fp": "Q3",
                                "form": "10-Q",
                                "filed": "2026-08-01",
                                "accn": "0000320193-26-000001",
                                "start": "2026-03-29",
                                "end": "2026-06-27",
                            }
                        ]
                    }
                }
            }
        }
    }
    report = StatementValidationService().validate(
        analysis_id="q1",
        ticker="AAPL",
        workbook_path=path,
        company_facts=facts,
        include_quarterly=True,
    )
    q = [e for e in report.entries if e.statement == "quarterly_income_statement"]
    assert q
    assert q[0].fiscal_period == "2026 Q3"


def test_period_end_attached_on_validated(mini_workbook: Path):
    report = StatementValidationService().validate(
        analysis_id="v8",
        ticker="AAPL",
        workbook_path=mini_workbook,
        company_facts=_facts_revenue(),
    )
    e = next(x for x in report.entries if x.metric == "Revenue" and x.fiscal_period == "FY2018")
    assert e.decision == StatementValidationDecision.VALIDATED
    assert e.period_end == "2018-09-29"
    assert e.period_start == "2017-10-01"


def test_source_missing_when_no_sec_fact(mini_workbook: Path):
    report = StatementValidationService().validate(
        analysis_id="v9",
        ticker="AAPL",
        workbook_path=mini_workbook,
        company_facts={"facts": {"us-gaap": {}}},
    )
    assert any(e.decision == StatementValidationDecision.SOURCE_MISSING for e in report.entries)


def test_eps_split_marked_not_comparable():
    # As-reported 11.91 vs post-split restated ~2.98 (4-for-1)
    assert (
        compare_values(2.9775, 11.91, scale="per_share")
        == StatementValidationDecision.NOT_COMPARABLE
    )


def test_capex_outflow_sign_normalizes(tmp_path: Path):
    wb = Workbook()
    wb.active.title = "Income - GAAP"
    for col, fy in zip(list("CDEFGHIJKL"), range(2016, 2026)):
        wb.active[f"{col}7"] = f"FY{fy}"
    bs = wb.create_sheet("Balance Sheet - Standardized")
    for col, fy in zip(list("CDEFGHIJKL"), range(2016, 2026)):
        bs[f"{col}7"] = f"FY{fy}"
    cf = wb.create_sheet("Cash Flow - Standardized")
    for col, fy in zip(list("CDEFGHIJKL"), range(2016, 2026)):
        cf[f"{col}7"] = f"FY{fy}"
    cf["A32"] = "+ Acq of Fixed Prod Assets"
    cf["E32"] = -13313  # FY2018 workbook outflow
    path = tmp_path / "capex.xlsx"
    wb.save(path)
    wb.close()

    facts = {
        "facts": {
            "us-gaap": {
                "PaymentsToAcquirePropertyPlantAndEquipment": {
                    "units": {
                        "USD": [
                            {
                                "val": 13_313_000_000.0,
                                "fy": 2018,
                                "fp": "FY",
                                "form": "10-K",
                                "filed": "2018-11-05",
                                "accn": "0000320193-18-000145",
                                "frame": "CY2018",
                                "end": "2018-09-29",
                                "start": "2017-10-01",
                            }
                        ]
                    }
                }
            }
        }
    }
    report = StatementValidationService().validate(
        analysis_id="capex",
        ticker="AAPL",
        workbook_path=path,
        company_facts=facts,
    )
    e = next(
        x for x in report.entries if x.metric == "Capital Expenditures" and x.fiscal_period == "FY2018"
    )
    assert e.decision == StatementValidationDecision.VALIDATED
