"""Updates must cope with a previous workbook on an older template, a missing custom filter and look-alike files."""

from __future__ import annotations

import datetime as dt
import sys
from pathlib import Path

from openpyxl import Workbook, load_workbook

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "scripts"))
import run_batch  # noqa: E402

from services.custom_run_standin_service import build_standin, find_filter_summary  # noqa: E402
from services.quarterly_model_continuity_service import (  # noqa: E402
    QuarterlyModelContinuityService,
    carry_constants_by_label,
    layout_matches,
)


def _statement(path: Path, ticker: str, version: str, labels: list[str], values: dict[str, float] | None = None) -> Path:
    wb = Workbook()
    ws = wb.active
    ws.title = "Balance Sheet - Standardized"
    ws["C1"] = ticker
    wt = wb.create_sheet("Template Version")
    wt["A1"], wt["B1"] = "Version", version
    sheet = wb.create_sheet("Final Metrics")
    for i, lab in enumerate(labels, start=3):
        sheet.cell(i, 1).value = lab
        if values and lab in values:
            sheet.cell(i, 3).value = values[lab]
    wb.save(path)
    return path


def test_layout_guard_and_label_carry(tmp_path: Path):
    prev = _statement(tmp_path / "p.xlsx", "ABC", "v25", ["ROA", "Discount Rate"], {"Discount Rate": 0.12})
    new = _statement(tmp_path / "n.xlsx", "ABC", "v28", ["ROA", "Extra row", "Discount Rate"])
    pw, nw = load_workbook(prev)["Final Metrics"], load_workbook(new)["Final Metrics"]
    assert not layout_matches(pw, nw)
    assert layout_matches(nw, nw)
    assert carry_constants_by_label(pw, nw) == 1
    assert nw["C5"].value == 0.12 and nw["C4"].value is None        # landed on the 'Discount Rate' row, not on row 4


def _lq(path: Path, cur: dt.datetime, d: dt.datetime, values: dict[str, tuple[float, float]]) -> Path:
    wb = Workbook()
    ws = wb.active
    ws.title = "Last Quarter BS Standardized"
    ws["C4"], ws["D4"] = cur, d
    for i, (code, (c, dd)) in enumerate(values.items(), start=11):
        ws.cell(i, 2).value, ws.cell(i, 3).value, ws.cell(i, 4).value = code, c, dd
    wb.save(path)
    return path


def test_prior_quarter_column_replaced_from_previous_workbook(tmp_path: Path):
    prev = _lq(tmp_path / "p.xlsx", dt.datetime(2024, 10, 27), dt.datetime(2024, 7, 28), {"BS_TOT_ASSET": (34409, 33647)})
    out = _lq(tmp_path / "o.xlsx", dt.datetime(2025, 1, 26), dt.datetime(2026, 4, 26), {"BS_TOT_ASSET": (33338, 40286)})
    rep = QuarterlyModelContinuityService().align_prior_quarter_balance_sheet(previous_workbook_path=prev, workbook_path=out)
    assert rep["status"] == "ok"
    ws = load_workbook(out)["Last Quarter BS Standardized"]
    assert ws["D11"].value == 34409 and ws["D4"].value == dt.datetime(2024, 10, 27)


def test_prior_quarter_cleared_when_previous_is_not_the_quarter_before(tmp_path: Path):
    prev = _lq(tmp_path / "p.xlsx", dt.datetime(2023, 12, 31), dt.datetime(2023, 9, 30), {"BS_TOT_ASSET": (2434, 2328)})
    out = _lq(tmp_path / "o.xlsx", dt.datetime(2024, 9, 30), dt.datetime(2026, 3, 31), {"BS_TOT_ASSET": (2854, 4448)})
    rep = QuarterlyModelContinuityService().align_prior_quarter_balance_sheet(previous_workbook_path=prev, workbook_path=out)
    assert rep["status"] == "not_aligned"
    assert load_workbook(out)["Last Quarter BS Standardized"]["D11"].value is None


def _summary(path: Path, ticker: str) -> Path:
    wb = Workbook()
    ws = wb.active
    ws.title = "Summary"
    ws.append(["Company", "Ticker", "Current Price (Live Price)", "Max Current Price to Buy", "Current Market Capitalization"])
    ws.append(["ABC INC", ticker, 100.0, 80.0, 5000.0])
    wb.save(path)
    return path


def test_standin_filter_is_built_from_the_summary_row(tmp_path: Path):
    summary = _summary(tmp_path / "Current Filter Summary Sheet.xlsx", "ABC")
    assert find_filter_summary(tmp_path) == summary
    rep = build_standin("ABC", summary, tmp_path / "out" / "f.xlsx")
    ws = load_workbook(rep["path"])["ABC"]
    assert ws["B5"].value == 100.0 and "STAND-IN" in ws["A13"].value


def test_batch_matches_with_standin_and_rejects_a_look_alike_template(tmp_path: Path):
    _summary(tmp_path / "Current Filter Summary Sheet.xlsx", "ABC")
    _statement(tmp_path / "ABC 2025 Q2 - Template.xlsx", "ABC", "v28", ["ROA"])
    _statement(tmp_path / "2025 FY ABC FA.xlsx", "ABC", "v25", ["ROA"])
    m = run_batch.match_files(tmp_path, "ABC", "quarterly_update", standin_dir=tmp_path / "s")
    assert not m.problems and m.standin and any("stand-in" in n for n in m.notes) and any("v25" in n for n in m.notes)
    # a template that holds another company (EXP versus EXPD) is refused
    _summary(tmp_path / "Current Filter Summary Sheet.xlsx", "ABCD")
    _statement(tmp_path / "ABCD 2025 Q4 - Template.xlsx", "ABC", "v28", ["ROA"])
    _statement(tmp_path / "2024 FY ABCD FA.xlsx", "ABCD", "v25", ["ROA"])
    bad = run_batch.match_files(tmp_path, "ABCD", "annual_update", standin_dir=tmp_path / "s")
    assert any("holds ABC, not ABCD" in p for p in bad.problems)
