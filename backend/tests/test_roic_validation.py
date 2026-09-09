"""Focused tests for ROIC / NOPAT / Invested Capital review."""

from __future__ import annotations

from pathlib import Path

import pytest
from openpyxl import Workbook

from models.roic_validation import RoicDecision
from services.roic_validation_service import (
    RoicValidationService,
    average_invested_capital,
    classify_asset,
    classify_liability,
    house_invested_capital,
    house_nopat,
    house_roic,
    interpret_roic_wacc_spread,
    roic_within_tolerance,
    wacc_plausibility,
)


def test_nopat_calculation_house_method():
    nopat = house_nopat(
        revenue=1000,
        operating_expenses=600,
        lease_expense=50,
        lease_depreciation=40,
        rd_expense=80,
        rd_amortization=30,
        operating_taxes=70,
    )
    # 1000-600+50-40+80-30-70 = 390
    assert nopat == 390


def test_tax_rate_handling_zero_vs_economic():
    pretax_op = house_nopat(
        revenue=1000,
        operating_expenses=600,
        operating_taxes=0,
    )
    assert pretax_op == 400
    economic_rate = 0.21
    economic_nopat = pretax_op * (1 - economic_rate)
    assert economic_nopat == pytest.approx(316.0)
    # Zero house tax with material economic rate → review, not silent accept as after-tax
    assert pretax_op != economic_nopat


def test_operating_asset_classification():
    assert classify_asset("Accounts Receivable") == ("operating_asset", RoicDecision.VALIDATED)
    assert classify_asset("Inventories") == ("operating_asset", RoicDecision.VALIDATED)
    assert classify_asset("ST Investments") == ("non_operating", RoicDecision.VALIDATED)
    assert classify_asset("LT Marketable Securities") == ("non_operating", RoicDecision.VALIDATED)


def test_operating_liability_classification():
    assert classify_liability("Accounts Payable") == (
        "operating_liability",
        RoicDecision.VALIDATED,
    )
    assert classify_liability("ST Deferred Revenue") == (
        "operating_liability",
        RoicDecision.VALIDATED,
    )


def test_debt_excluded_from_operating_liabilities():
    cls, status = classify_liability("ST Debt")
    assert cls == "financing"
    assert status == RoicDecision.VALIDATED
    cls2, _ = classify_liability("LT Debt / Borrowings")
    assert cls2 == "financing"
    ic = house_invested_capital(
        operating_assets=500,
        operating_liabilities=100,  # no debt in OL
        capitalized_leases=0,
        capitalized_rd=0,
    )
    assert ic == 400


def test_ambiguous_classification_review_not_guess():
    cls, status = classify_liability("Interest & Dividends Payable")
    assert cls == "ambiguous"
    assert status == RoicDecision.REVIEW_REQUIRED
    cls2, status2 = classify_liability("ST Lease Liabilities")
    assert cls2 == "ambiguous"
    assert status2 == RoicDecision.REVIEW_REQUIRED
    cls3, status3 = classify_asset("Other Assets Misc")
    assert cls3 == "ambiguous"
    assert status3 == RoicDecision.REVIEW_REQUIRED


def test_invested_capital_reconciliation():
    ic = house_invested_capital(
        operating_assets=200,
        operating_liabilities=80,
        capitalized_leases=10,
        capitalized_rd=15,
    )
    assert ic == 145
    workbook_ic = 145.0
    assert abs(workbook_ic - ic) <= 1.0


def test_average_vs_ending_invested_capital_convention():
    beg, end = 100.0, 200.0
    ending_roic = house_roic(nopat=50, invested_capital=end)
    avg_roic = house_roic(nopat=50, invested_capital=average_invested_capital(beg, end))
    assert ending_roic == pytest.approx(0.25)
    assert avg_roic == pytest.approx(50 / 150)
    assert ending_roic != avg_roic
    # Must not silently switch
    assert average_invested_capital(beg, end) != end


def test_roic_comparison_tolerance():
    assert roic_within_tolerance(1.160, 1.158)
    assert not roic_within_tolerance(1.160, 1.140)


def test_wacc_missing_implausible_handling():
    assert wacc_plausibility(None) == RoicDecision.SOURCE_MISSING
    assert wacc_plausibility(0.0) == RoicDecision.REVIEW_REQUIRED
    assert wacc_plausibility(0.55) == RoicDecision.REVIEW_REQUIRED
    assert wacc_plausibility(0.09) == RoicDecision.VALIDATED


def test_roic_wacc_interpretation():
    assert interpret_roic_wacc_spread(0.10) == "value_creation"
    assert interpret_roic_wacc_spread(-0.05) == "value_destruction"
    assert interpret_roic_wacc_spread(0.005) == "weak_neutral"


def _minimal_roic_workbook(path: Path) -> Path:
    """Build a tiny workbook with HAP-like ROIC architecture (cached values)."""
    wb = Workbook()
    # remove default
    default = wb.active
    wb.remove(default)

    bs = wb.create_sheet(BS_SHEET := "Balance Sheet - Standardized")
    income = wb.create_sheet("Income - GAAP")
    inputs = wb.create_sheet("Inputs")
    tax = wb.create_sheet("Tax")
    ic = wb.create_sheet("IC & NOPAT & ROIC ")
    fm = wb.create_sheet("Final Metrics")

    # One period in column C = FY2024
    bs["C8"] = "FY2024"
    # OA current
    bs["C11"] = 20  # cash
    bs["C14"] = 30  # AR
    bs["C20"] = 10  # inv
    bs["C27"] = 5  # prepaid
    # OA noncurrent
    bs["C37"] = 100  # ppe
    bs["C48"] = 15  # intangibles
    bs["C51"] = 0
    # non-op
    bs["C12"] = 200  # STI excluded
    bs["C45"] = 300  # LT mkt securities
    # OL
    bs["C65"] = 40
    bs["C68"] = 5
    bs["C78"] = 10
    bs["C98"] = 5
    # debt financing
    bs["C70"] = 50
    bs["C86"] = 80

    income["C9"] = 500
    income["C14"] = 200
    income["C22"] = 100
    income["C30"] = 200
    income["C42"] = 210
    income["C44"] = 42  # 20% tax

    inputs["C54"] = 9.0  # WACC percent
    inputs["C80"] = 180  # OA
    inputs["C83"] = 60  # OL
    inputs["C87"] = 500
    inputs["C88"] = 300
    inputs["C89"] = 200

    tax["C25"] = 40
    tax["C26"] = 0.2

    # IC sheet values (as Excel would cache)
    ic["C3"] = 180
    ic["C4"] = 60
    ic["C5"] = 10
    ic["C6"] = 20
    ic["C7"] = 150  # 180-60+10+20
    ic["C11"] = 500
    ic["C12"] = 300
    ic["C13"] = 200
    ic["C14"] = 8
    ic["C15"] = 6
    ic["C16"] = 12
    ic["C17"] = 4
    ic["C19"] = 40
    # NOPAT = 500-300+8-6+12-4-40 = 170
    ic["C20"] = 170
    ic["C23"] = 170 / 150
    ic["C7"].value = 150
    for col in range(4, 13):
        # leave other years empty / zero-ish for unused cols
        pass

    # Formulas for methodology discovery
    ic["C7"] = "=C3-C4+C5+C6"
    # But data_only load won't have formula results — write numeric via a second save trick:
    # openpyxl data_only reads cached; without Excel cache formulas return None.
    # So store numbers only (no formulas) for test workbook values path.
    ic["C7"] = 150
    ic["C20"] = 170
    ic["C23"] = 170 / 150

    fm["C6"] = 170 / 150
    fm["C7"] = 0.09
    fm["C8"] = (170 / 150) - 0.09
    fm["A6"] = "ROIC Including Goodwill"
    fm["A7"] = "WACC  Bloomberg"
    fm["A8"] = "ROIC in - WACC"

    # Labels
    ic["A3"] = "Operating Assets"
    ic["A4"] = "Operating Liabilities"
    ic["A7"] = "Invested Capital "
    ic["A20"] = "NOPAT"
    ic["A23"] = "ROIC in "

    wb.save(path)
    wb.close()
    return path


def test_workbook_roic_validation_end_to_end(tmp_path: Path):
    path = _minimal_roic_workbook(tmp_path / "roic.xlsx")
    report = RoicValidationService().validate(
        analysis_id="t-roic",
        ticker="TEST",
        workbook_path=path,
    )
    assert report.methodology["invested_capital_convention"] == "ending"
    assert len(report.periods) == 10
    p = next(x for x in report.periods if x.column == "C")
    assert p.nopat_workbook == pytest.approx(170)
    assert p.nopat_independent_house == pytest.approx(170)
    assert p.invested_capital_independent == pytest.approx(150)
    assert p.roic_workbook == pytest.approx(170 / 150)
    assert p.roic_decision == RoicDecision.VALIDATED
    assert p.wacc == pytest.approx(0.09)
    assert p.roic_minus_wacc == pytest.approx((170 / 150) - 0.09)
    assert p.spread_interpretation == "value_creation"

    # Debt not in OL: independent OL from BS = 40+5+10+5=60
    assert p.operating_liabilities == pytest.approx(60)

    # Classifications: financing debt validated; ambiguous interest/lease review
    classes = {c.workbook_line: c for c in report.classifications}
    assert any(c.classification == "financing" for c in report.classifications)
    assert any(
        c.classification == "ambiguous" and c.status == RoicDecision.REVIEW_REQUIRED
        for c in report.classifications
    )

    review = RoicValidationService().build_analyst_review(report)
    assert review.items
    assert "roic_minus_wacc_trend" in {i.topic for i in review.items}


def test_average_convention_metadata_preserved(tmp_path: Path):
    path = _minimal_roic_workbook(tmp_path / "roic_avg.xlsx")
    report = RoicValidationService().validate(
        analysis_id="t-avg",
        ticker="TEST",
        workbook_path=path,
        invested_capital_convention="average",
    )
    assert report.methodology["invested_capital_convention"] == "average"
    p = next(x for x in report.periods if x.column == "C")
    assert p.evidence["invested_capital_convention"] == "average"
