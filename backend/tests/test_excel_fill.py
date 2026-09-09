"""M4 Excel fill engine tests."""

from __future__ import annotations

import hashlib
from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, numbers

from models.analysis import Analysis
from models.custom_run import CustomRunData
from models.workbook_schema import WorkbookStructure
from pipeline.stages.fill_workbook import FillWorkbookStage
from services.excel_fill_service import ExcelFillService
from services.output_service import OutputService
from workbook_mapping.engine import IntentDecision, WriteIntent, WriteIntentReport


def _sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def _intent(
    *,
    intent_id: str,
    sheet: str,
    cell: str,
    value,
    decision: IntentDecision,
    metric: str = "Revenue",
    period: str = "FY2024",
    reason: str = "test",
    write_policy_result: str = "writable:Hybrid",
) -> WriteIntent:
    return WriteIntent(
        intent_id=intent_id,
        mapping_id=intent_id,
        sheet=sheet,
        cell=cell,
        metric=metric,
        period=period,
        value=value,
        source="test",
        source_document="test.xlsx",
        transformation="divide_by_1_000_000" if isinstance(value, float) else None,
        confidence=0.9,
        existing_cell_classification="Writable Input",
        write_policy_result=write_policy_result,
        decision=decision,
        reason=reason,
        cfm_path="income_statement.revenue",
    )


@pytest.fixture
def industrial_mini(tmp_path: Path) -> Path:
    wb = Workbook()
    ws = wb.active
    ws.title = "Income - GAAP"
    ws["C3"] = "FY2024"
    ws["A9"] = "Revenue"
    ws["K9"] = None
    ws["K9"].number_format = numbers.FORMAT_NUMBER_COMMA_SEPARATED1
    ws["K9"].font = Font(name="Calibri", size=11, bold=False)
    ws["M9"] = "=K9*2"
    ws["C1"] = None
    path = tmp_path / "upload.xlsx"
    wb.save(path)
    wb.close()
    return path


def test_write_intent_changes_target_and_preserves_formula_and_format(
    industrial_mini: Path,
    tmp_path: Path,
):
    dest = tmp_path / "completed.xlsx"
    before = _sha256(industrial_mini)
    report = WriteIntentReport(
        analysis_id="m4",
        ticker="AAPL",
        intents=[
            _intent(
                intent_id="is.revenue:K9",
                sheet="Income - GAAP",
                cell="K9",
                value=391035.0,
                decision=IntentDecision.WRITE,
            ),
            _intent(
                intent_id="bad.formula:M9",
                sheet="Income - GAAP",
                cell="M9",
                value=999.0,
                decision=IntentDecision.WRITE,
                reason="should be blocked at runtime",
            ),
            _intent(
                intent_id="skip.me:C1",
                sheet="Income - GAAP",
                cell="C1",
                value="AAPL",
                decision=IntentDecision.SKIP,
                metric="Ticker",
            ),
            _intent(
                intent_id="block.me:C1",
                sheet="Income - GAAP",
                cell="C1",
                value="NOPE",
                decision=IntentDecision.BLOCK,
                metric="Ticker",
                write_policy_result="protected_cell",
            ),
        ],
        write_count=2,
        skip_count=1,
        block_count=1,
    )

    provenance, cell_diff, written, runtime_skip, runtime_block = ExcelFillService().apply_write_intents(
        analysis_id="m4",
        ticker="AAPL",
        source_workbook_path=industrial_mini,
        destination_workbook_path=dest,
        intent_report=report,
    )

    assert _sha256(industrial_mini) == before  # original unchanged
    assert written == 1
    assert runtime_block >= 1
    assert cell_diff.changed_count >= 1
    assert cell_diff.written_count == 1

    src = load_workbook(industrial_mini, data_only=False)
    try:
        assert src["Income - GAAP"]["K9"].value is None
        assert src["Income - GAAP"]["M9"].value == "=K9*2"
    finally:
        src.close()

    out = load_workbook(dest, data_only=False)
    try:
        cell = out["Income - GAAP"]["K9"]
        assert cell.value == pytest.approx(391035.0)
        assert cell.number_format == numbers.FORMAT_NUMBER_COMMA_SEPARATED1
        assert cell.font.name == "Calibri"
        assert out["Income - GAAP"]["M9"].value == "=K9*2"
        assert out["Income - GAAP"]["M9"].data_type == "f"
        assert out["Income - GAAP"]["C1"].value is None  # SKIP/BLOCK not applied
        assert list(out.sheetnames) == ["Income - GAAP"]
    finally:
        out.close()

    filled = [e for e in provenance if e.status == "filled" and e.cell == "K9"]
    assert filled
    assert filled[0].original_value is None
    assert filled[0].value == pytest.approx(391035.0)
    assert filled[0].cell_ref == "Income - GAAP!K9"

    formula_blocks = [
        e for e in cell_diff.entries if e.cell == "M9" and e.write_decision == "BLOCK"
    ]
    assert formula_blocks


def test_runtime_missing_sheet_is_blocked(tmp_path: Path, industrial_mini: Path):
    dest = tmp_path / "completed.xlsx"
    report = WriteIntentReport(
        analysis_id="m4",
        ticker="AAPL",
        intents=[
            _intent(
                intent_id="x",
                sheet="Does Not Exist",
                cell="A1",
                value=1,
                decision=IntentDecision.WRITE,
            )
        ],
        write_count=1,
    )
    _, cell_diff, written, _, runtime_block = ExcelFillService().apply_write_intents(
        analysis_id="m4",
        ticker="AAPL",
        source_workbook_path=industrial_mini,
        destination_workbook_path=dest,
        intent_report=report,
    )
    assert written == 0
    assert runtime_block == 1
    assert cell_diff.entries[0].write_decision == "BLOCK"


def test_fill_stage_loads_write_intents_artifact(industrial_mini: Path, tmp_path: Path):
    output_service = OutputService(outputs_dir=tmp_path / "outputs")
    analysis = Analysis(
        analysis_id="stage-m4",
        company="Apple Inc.",
        ticker="AAPL",
        analysis_type="annual_update",
        status="uploaded",
    )
    report = WriteIntentReport(
        analysis_id="stage-m4",
        ticker="AAPL",
        milestone="M3",
        intents=[
            _intent(
                intent_id="is.revenue:K9",
                sheet="Income - GAAP",
                cell="K9",
                value=100.0,
                decision=IntentDecision.WRITE,
                period="FY2024",
            )
        ],
        write_count=1,
    )
    intents_rel = output_service.write_json("stage-m4", "write_intents.json", report)
    upload_hash = _sha256(industrial_mini)

    stage = FillWorkbookStage(output_service=output_service)
    custom_run = CustomRunData(
        source_filename="Custom_Run_Filter_AAPL.xlsx",
        ticker="AAPL",
        ticker_sheet_name="AAPL",
        summary_field_count=60,
    )
    provenance, completion, workbook_rel, prov_path, diff_path, completion_path, log = stage.run(
        analysis,
        industrial_mini,
        custom_run,
        WorkbookStructure(workbook_filename="upload.xlsx"),
        company_facts={},
        filings_manifest={},
        write_intents_path=intents_rel,
    )

    assert _sha256(industrial_mini) == upload_hash
    completed = output_service.artifact_path("stage-m4", "completed_workbook.xlsx")
    assert completed.exists()
    assert _sha256(completed) != upload_hash
    assert (tmp_path / "outputs" / "stage-m4" / "cell_diff_report.json").exists()
    assert (tmp_path / "outputs" / "stage-m4" / "completion_report.json").exists()
    assert completion.fill_count == 1
    assert "FILL=" in log.detail or "excel_wrote=" in log.detail

    wb = load_workbook(completed, data_only=False)
    try:
        assert wb["Income - GAAP"]["K9"].value == pytest.approx(100.0)
    finally:
        wb.close()

    assert any(e.cell_ref == "Income - GAAP!K9" and e.status == "filled" for e in provenance.entries)
    assert Path(prov_path).name == "provenance_report.json" or "provenance_report" in prov_path
    assert "cell_diff_report" in diff_path
    assert "completion_report" in completion_path
    assert workbook_rel.endswith("completed_workbook.xlsx")
