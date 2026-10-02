"""Blank statement cells: filled from SEC and flagged; blanks no metric uses are ignored; needed-but-unavailable are flagged."""

from openpyxl import Workbook, load_workbook

from services.formula_dependencies import MetricDependencies
from services.new_company_statement_validation_service import NewCompanyStatementValidationService


def test_metric_dependencies_follow_ranges_and_chains():
    wb = Workbook()
    inputs = wb.active
    inputs.title = "Inputs"
    helper = wb.create_sheet("Helper")
    metrics = wb.create_sheet("Final Metrics")
    helper["A1"] = "=SUM(Inputs!B2:B3)"
    metrics["A1"] = "=Helper!A1*2"
    metrics["A2"] = "='Inputs'!$D$9"
    deps = MetricDependencies(wb)
    assert deps.feeds_metrics("Inputs", 2, 2) and deps.feeds_metrics("Inputs", 3, 2)  # via the range and the chain
    assert deps.feeds_metrics("Inputs", 9, 4)  # direct, quoted sheet and absolute reference
    assert not deps.feeds_metrics("Inputs", 4, 2)  # nothing reads it


def build_workbook(path):
    wb = Workbook()
    ws = wb.active
    ws.title = "Income - GAAP"
    ws.cell(3, 3, "FY 2025")
    ws.cell(9, 1, "Revenue")          # blank, needed, SEC has it -> filled
    ws.cell(30, 1, "Operating Income")  # blank, needed, SEC lacks it -> flagged
    ws.cell(58, 1, "Net Income")        # supplied value that differs materially -> flagged, never overwritten
    ws.cell(58, 3, 100.0)
    ws.cell(19, 1, "Gross Profit")      # blank, nothing uses it -> ignored
    metrics = wb.create_sheet("Final Metrics")
    metrics["A1"] = "='Income - GAAP'!C9+'Income - GAAP'!C30+'Income - GAAP'!C58"
    wb.save(path)


def test_blank_cells_are_filled_flagged_or_ignored_by_importance(tmp_path):
    path = tmp_path / "w.xlsx"
    build_workbook(path)
    report = NewCompanyStatementValidationService().validate(
        analysis_id="a", ticker="T", workbook_path=path, fiscal_years=["FY2025"],
        filing_overrides={"FY2025": {"revenue": 500.0, "net_income": 200.0}},
    )
    ws = load_workbook(path)["Income - GAAP"]
    # filled from SEC, shaded, source in the comment
    assert ws["C9"].value == 500.0 and "HAP FILLED" in ws["C9"].comment.text
    assert [f.concept for f in report.filled_missing] == ["revenue"]
    # needed by a metric but unavailable anywhere -> flagged as missing data (not invented)
    assert ws["C30"].value is None and "DATA MISSING" in ws["C30"].comment.text
    assert report.flagged_missing == ["FY2025:operating_income:Income - GAAP!C30"]
    # not used by any metric -> left alone and not flagged
    assert ws["C19"].value is None and ws["C19"].comment is None
    # a supplied value is never overwritten, even when it differs from SEC
    assert ws["C58"].value == 100.0
    assert any("net_income" in u for u in report.unresolved_material)


def test_liabilities_and_equity_row_is_never_filled_from_liabilities():
    from services.quarterly_presentation_service import _LIABILITIES_AND_EQUITY as pattern

    for label in ("Liabilities & Shareholders' Equity", "Liabilities and Stockholders Equity", "Total Liabilities & Equity"):
        assert pattern.search(label), label
    assert not pattern.search("Total Liabilities") and not pattern.search("Accounts Payable")
