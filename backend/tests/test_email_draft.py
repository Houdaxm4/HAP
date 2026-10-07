from email import message_from_bytes, policy
from pathlib import Path

from services.email_draft_service import (
    EmailDraftService,
    Facts,
    SheetFields,
    change,
    money_b,
    money_m,
    pct,
    pe10_status,
    render_annual,
    render_new_company,
    render_quarter,
    usd,
)


def _facts(**kw) -> Facts:
    f = Facts(ticker="LNN", company="Lindsay Corporation", fy="FY2025", quarter_label="2026 Q3", quarter_number=3, quarter_year=2026,
              ytd_label="9M", price=93.34, market_cap_m=8920.0, tbv_p=0.0358, pe10_percentile=0.1687, max_entry=94.67, ev_mos=-1.0352,
              expected_return=0.0468, expected_return_adjusted=True, graham_entry=28.72)
    f.fy_metrics = {
        "roce": (0.0729, 0.0837, 0.1097), "roic_wacc": (-0.0157, 0.0397, 0.0377), "a_l": (2.90, 2.89, None),
        "gross_margin": (0.612, 0.605, None), "op_margin": (0.0777, 0.08, None), "net_margin": (0.0688, 0.09, None),
        "debt_assets": (0.0004, 0.0, None), "interest_cov": (198.31, 0.0, None),
        "revenue": (579.5, 614.6, None), "net_income": (39.9, 51.6, None), "eps": (1.56, 2.01, None),
    }
    f.fy_changes = {"cfo": (100.0, 117.6), "retained": (990.0, 998.2), "equity": (700.0, 716.0)}
    f.q_yoy = {
        "revenue": {"q": (1200.0, 1100.0), "ytd": (3686.8, 3444.2)},
        "net_income": {"q": (100.0, 80.0), "ytd": (290.3, 243.1)},
        "net_common": {"ytd": (290.3, 243.1), "q": (100.0, 80.0)},
        "eps": {"q": (1.0, 0.8), "ytd": (2.99, 2.42)},
        "op_income": {"ytd": (397.4, 317.6)},
        "gross_profit": {"ytd": (1500.0, 1400.0)},
        "cfo": {"ytd": (476.2, 348.9)},
    }
    f.q_qoq = {"assets": (3200.0, 3140.0), "liabilities": (1988.0, 1975.0), "retained": (980.0, 999.0), "equity": (1212.0, 1241.0)}
    for k, v in kw.items():
        setattr(f, k, v)
    return f


def test_number_formats_follow_the_templates():
    assert money_b(8920.0) == "~$8.92B"
    assert money_b(550.0) == "~$0.55B"
    assert pct(0.0622) == "6.22%" and pct(0.0704, sign=True) == "+7.04%" and pct(-1.0352, sign=True) == "-103.52%"
    assert money_m(3686.8) == "$3,686.8M" and money_m(-8.6) == "-$8.6M"
    assert usd(93.34) == "$93.34" and usd(-0.15) == "-$0.15"
    assert change(3686.8, 3444.2) > 0.0704 and change(1.0, 0) is None


def test_status_with_pe10_is_never_guessed_from_the_price():
    assert pe10_status(_facts(price=93.34, max_entry=94.67)) == "[Google Sheet]"
    assert pe10_status(_facts(price=55.84, max_entry=55.75)) == "[Google Sheet]"


def test_quarter_email_matches_the_template():
    subject, body = render_quarter(_facts(projection={"next_fy": "2026", "fy": "FY2025", "roic_wacc": 0.0622, "prior_spread": 0.0188,
                                                      "roce": 0.18, "prior_roce": 0.1925}),
                                   sheet=SheetFields(classified_as="Q", pe10_status="Buy"))
    assert subject == "LNN quarter update Q3 2026"
    assert body.startswith("Dear All,\n\nPlease find attached the quarter update of LNN.\nMarket Cap: ~$8.92B\nCurrent TBV/P: 3.58%\nClassified as: Q\nStatus with PE10: Buy")
    assert "2026 Q3 Financial Highlights" in body
    assert "2026 Projection\nProjected ROIC - WACC at 6.22% (1.88% in FY2025)\nProjected ROCE at 18.00% (19.25% in FY2025)" in body
    assert "Comparison Q3 2026 to Q3 2025 (9M to 9M)" in body
    assert "Revenue +7.04% ($3,686.8M vs. $3,444.2M)" in body
    assert "Net Income (GAAP) +19.42% ($290.3M vs. $243.1M)" in body
    assert "EPS Diluted (GAAP) +23.55% ($2.99 vs. $2.42)" in body
    assert "Cash from Ops +36.49% ($476.2M vs. $348.9M)" in body
    assert "Comparison Q3 2026 to Q2 2026\nA/L: 1.61 vs. 1.59" in body
    assert "Expected Return: 4.68% (adjusted)" in body
    assert "Graham Entry Price: $28.72" in body
    assert body.rstrip().endswith("Very best,")


def test_annual_email_matches_the_template():
    f = _facts(ticker="ETD", price=21.71, max_entry=25.35, market_cap_m=550.0, expected_return=0.0923, expected_return_adjusted=False)
    subject, body = render_annual(f, sheet=SheetFields(classified_as="K", pe10_status="Out"), description="Ethan Allen is a manufacturer")
    assert subject == "ETD annual update FY2025"
    assert "Please find the annual update for ETD below (Ethan Allen is a manufacturer)" in body
    assert "Classified as: K\nStatus with PE10: Out" in body
    assert "FY2025 Highlights\nNet Revenue: $579.5M\nNet Income: $39.9M\nEPS: $1.56" in body
    assert "ROCE: 7.29% (8.37% in FY2024, 10-year average 10.97%)" in body
    assert "ROIC-WACC: -1.57% (3.97% in FY2024, 10-year average 3.77%)" in body
    assert "Margins: (Gross 61.2% — Op 7.77% — Net 6.88%)" in body
    assert "Interest Coverage ratio: 198.31x" in body
    assert "A/L: 2.90 from 2.89" in body
    assert "Expected Return: 9.23%\n" in body and "(adjusted)" not in body.split("Expected Return")[1].split("\n")[0]


def test_new_company_email_uses_net_income_available_to_common_when_preferred_dividends_exist():
    f = _facts(ticker="BRKR", company="Bruker Corporation", quarter_label="2026 Q1", quarter_number=1, quarter_year=2026, ytd_label="3M")
    f.q_yoy["net_income"] = {"q": (7.0, 17.4), "ytd": (7.0, 17.4)}
    f.q_yoy["net_common"] = {"q": (3.5, 17.4), "ytd": (3.5, 17.4)}
    subject, body = render_new_company(f, sheet=SheetFields(classified_as="U"), description="Bruker Corporation, headquartered in Billerica, Massachusetts, develops scientific instruments.")
    assert "Please find attached the analysis of BRKR: Bruker Corporation, headquartered in Billerica" in body
    assert "Classified as: U" in body
    assert "Q1 2026 vs. Q1 2025" in body and "(3M to 3M)" not in body          # no YTD note in Q1
    assert "Q1 2026 vs. Q4 2025" in body
    assert "Net Income available to Common (after preferred dividends)" in body
    assert "Gross Margin" in body
    assert body.rstrip().endswith("Let me know if you have any questions.\n\nVery best")


def test_unfilled_fields_are_marked_not_invented():
    _s, body = render_annual(Facts(ticker="XYZ"))
    assert "n/a" in body and "[Google Sheet]" in body and "[one-line description" in body


def test_eml_is_an_unsent_draft_addressed_to_the_analyst(tmp_path: Path, monkeypatch):
    from openpyxl import Workbook

    wb = Workbook()
    wb.active.title = "Inputs"
    wb["Inputs"]["B63"] = 50.0
    path = tmp_path / "w.xlsx"
    wb.save(path)
    monkeypatch.delenv("HAP_EMAIL_DRAFT_TO", raising=False)
    out = EmailDraftService().produce(analysis_type="annual_update", ticker="XYZ", workbook_path=path, output_dir=tmp_path, base_name="draft")
    msg = message_from_bytes(Path(out["eml_path"]).read_bytes())
    assert msg["To"] == "houda@vlixes.us" and msg["X-Unsent"] == "1"
    assert Path(out["text_path"]).read_text(encoding="utf-8").startswith("Subject: ")
    assert not list(tmp_path.glob("*.docx"))


def test_missing_gross_profit_or_revenue_does_not_break_the_comparison_block():
    f = _facts()
    f.q_yoy["gross_profit"] = {"ytd": (None, 1400.0)}
    f.q_yoy["op_income"] = {"ytd": (None, None)}
    _s, body = render_new_company(f)
    assert "Gross Margin" not in body and "Operating Margin" not in body and "Net Margin" in body


def test_email_problems_never_stop_an_analysis(tmp_path: Path):
    out = EmailDraftService().produce_safe(analysis_type="annual_update", ticker="XYZ", workbook_path=tmp_path / "missing.xlsx", output_dir=tmp_path)
    assert out["text_path"] is None and out["error"]


def test_a_quarter_block_with_no_data_is_left_out_not_filled_with_na():
    f = _facts()
    f.q_yoy = {k: {"q": (None, None), "ytd": (None, None)} for k in ("revenue", "net_income", "eps", "op_income", "gross_profit")}
    f.q_yoy["cfo"] = {"ytd": (None, None)}
    _s, body = render_new_company(f)
    assert "Q3 2026 vs. Q3 2025" not in body and "n/a ($" not in body and "Revenue n/a" not in body


def _with_series(f: Facts) -> Facts:
    f.fy_series = {"roic_wacc": [(f"FY{2016 + i}", v) for i, v in enumerate([-0.0477, -0.01, 0.0238, 0.01, -0.02, -0.03, 0.0, -0.01, -0.0367, -0.0477])]}
    return f


def test_annual_highlights_are_written_from_the_numbers_without_claiming_causes():
    f = _facts()
    _s, body = render_annual(f, description="Ethan Allen is a manufacturer")
    text = body.split("EPS: $1.56")[1].split("FY2025 Metrics")[0]
    assert "[What drove the year" not in body
    assert "Revenue declined 5.7% to $579.5M." in text and "Net income fell 22.7%" in text
    assert "Gross margin improved to 61.2% from 60.5%." in text
    assert "ROIC-WACC is negative at -1.57% (3.77% on a 10-year average)" in text and "ROCE at 7.29%, below its 10-year average of 10.97%" in text
    assert "because" not in text.lower() and "demand" not in text.lower()


def test_quarter_highlights_use_the_year_to_date_comparison():
    _s, body = render_quarter(_facts())
    text = body.split("2026 Q3 Financial Highlights")[1].split("Comparison")[0]
    assert "[What drove the quarter" not in body
    assert "Revenue for the 9M grew 7.0% to $3,686.8M (from $3,444.2M)." in text and "Cash from operations rose 36.5%" in text


def test_new_company_highlights_and_conclusion_state_the_decade_range_and_the_price_position():
    f = _with_series(_facts(ticker="BRKR", company="Bruker Corporation", price=55.84, max_entry=55.75, pe10_percentile=0.2154, ev_mos=-2.197,
                            expected_return=-0.0362, graham_entry=6.42))
    _s, body = render_new_company(f, sheet=SheetFields(pe10_status="Out"), description="Bruker Corporation, headquartered in Billerica, Massachusetts, develops scientific instruments.")
    assert "ROIC-WACC has ranged between -4.77% and 2.38% over the last 10 years and was -4.77% in FY2025." in body
    assert "At $55.84 the stock trades above its maximum entry price of $55.75, with the PE10 percentile at 21.54%." in body
    assert "The enterprise-value margin of safety is -219.7%" in body and "The expected annual return is -3.62% (adjusted)." in body
    assert "Status with PE10: Out" in body.split("Highlights")[0] and "[Overall view" not in body.split("Conclusion")[1]


def test_too_few_figures_keep_the_bracketed_line_for_the_analyst():
    _s, body = render_new_company(Facts(ticker="XYZ"))
    assert "[Competitive position" in body and "[Overall view" in body and "[one-line description of the company]" in body


def test_expected_return_is_the_price_plus_dividends_column_of_the_expected_returns_tab(tmp_path: Path):
    from openpyxl import Workbook

    from services.email_draft_service import read_facts

    wb = Workbook()
    wb.active.title = "Inputs"
    er = wb.create_sheet("Expected Returns & Buybacks")
    headers = ["Average ROE", "EPS in ten years", "Price in ten years", "Price + Dividends", "Expected Annual Return",
               "Expected Return Price + Dividends", None, "HAP Alternative Expected Return Price + Dividends"]
    values = [0.1255, 2.38, 25.2, 52.5, 0.0182, 0.0923, None, -0.1119]
    for c, (h, v) in enumerate(zip(headers, values), start=1):
        er.cell(13, c).value, er.cell(14, c).value = h, v
    path = tmp_path / "w.xlsx"
    wb.save(path)
    assert read_facts(path, ticker="ETD").expected_return == 0.0923            # F14, not E14 (1.82%) and not the HAP alternative
    er["F13"].value = None
    wb.save(path)
    assert read_facts(path, ticker="ETD").expected_return == 0.0182            # falls back to E14 when the template has no F column


def test_eml_carries_the_excel_and_the_business_doc_as_attachments(tmp_path: Path, monkeypatch):
    from openpyxl import Workbook

    wb = Workbook()
    wb.active.title = "Inputs"
    wb["Inputs"]["B63"] = 50.0
    xlsx = tmp_path / "2026 Q2 XYZ FA.xlsx"
    wb.save(xlsx)
    docx = tmp_path / "2026 Q2 XYZ Business.docx"
    docx.write_bytes(b"PK-business")
    monkeypatch.delenv("HAP_EMAIL_DRAFT_TO", raising=False)
    out = EmailDraftService().produce(
        analysis_type="new_company", ticker="XYZ", workbook_path=xlsx, output_dir=tmp_path, base_name="draft",
        attachments=[xlsx, docx, tmp_path / "missing.pdf"],
    )
    msg = message_from_bytes(Path(out["eml_path"]).read_bytes(), policy=policy.default)
    names = [part.get_filename() for part in msg.iter_attachments()]
    assert names == ["2026 Q2 XYZ FA.xlsx", "2026 Q2 XYZ Business.docx"]       # a missing file is skipped, not an error
    assert msg["X-Unsent"] == "1" and "Please find attached" in msg.get_body().get_content()
