"""Annual and Quarterly Word reports start with the Flags section and end the analysis with the opinion sections."""

from openpyxl import Workbook
from docx import Document

from models.annual_update import AnnualPerformanceReport
from services.annual_deliverables_service import AnnualDeliverablesService
from services.quarterly_deliverables_service import QuarterlyDeliverablesService


def texts(path):
    return [p.text for p in Document(str(path)).paragraphs if p.text.strip()]


def test_quarterly_report_starts_with_flags_and_says_when_not_authorized(tmp_path):
    workbook = tmp_path / "w.xlsx"
    Workbook().save(workbook)
    out = tmp_path / "q.docx"
    QuarterlyDeliverablesService()._write_word(
        out, ticker="acme", fiscal_year=2026, fiscal_quarter=2, workbook_path=workbook,
        projection=None, review=None, research=None, authorized=False,
    )
    lines = texts(out)
    assert lines[1] == "Flags" and lines[2].startswith("Status: NOT AUTHORIZED")
    assert lines.index("Flags") < lines.index("Financial Highlights") < lines.index("Fundamentals: how strong are they?") < lines.index("Sources")
    assert "Rules-based assessment: not available" in "\n".join(lines)  # a quarterly update has no business-quality score


def test_annual_report_starts_with_flags_and_places_opinions_before_validation(tmp_path):
    out = tmp_path / "a.docx"
    perf = AnnualPerformanceReport(analysis_id="x", ticker="ACME", fiscal_year=2026)
    AnnualDeliverablesService()._write_word(out, "ACME", 2026, perf, None)
    lines = texts(out)
    assert lines[1] == "Flags" and lines[2].startswith("Status: report authorized")
    assert (lines.index("Flags") < lines.index("1. Executive Investment Conclusion")
            < lines.index("Fundamentals: how strong are they?") < lines.index("9. Validation and Open Issues"))
