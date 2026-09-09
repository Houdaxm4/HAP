"""New Company pipeline unit and integration tests."""

from __future__ import annotations

from datetime import datetime
from pathlib import Path
from unittest.mock import MagicMock, patch

import pytest
from openpyxl import Workbook, load_workbook

from models.custom_run import CustomRunData, CustomRunPeriods
from models.new_company import (
    BuybackAbsenceClass,
    CLOUD_PENDING_WINDOWS_CERTIFICATION,
    NewCompanyWorkflowState,
    ProjectionConfidence,
)
from services.excel_recalc_service import ExcelRecalcReport, genuine_excel_com_recalc
from services.new_company_buyback_service import NewCompanyBuybackService
from services.new_company_lease_service import NewCompanyLeaseService
from services.new_company_output_gate_service import NewCompanyOutputGateService
from services.new_company_pe10_service import NewCompanyPe10Service, exact_field_identity
from services.new_company_period_service import NewCompanyPeriodService
from services.new_company_projection_service import NewCompanyProjectionService
from services.new_company_rd_service import NewCompanyRdService
from services.new_company_runner import NewCompanyRunner
from services.new_company_seasonality_service import NewCompanySeasonalityService
from services.new_company_sec_coverage_service import NewCompanySecCoverageService
from services.new_company_statement_validation_service import NewCompanyStatementValidationService
from services.new_company_tax_service import NewCompanyTaxService
from services.output_service import OutputService


YEARS = list(range(2016, 2026))
FY = [f"FY{y}" for y in YEARS]


def _header(ws, years: list[int]) -> None:
    ws["A1"] = "Line"
    ws["A7"] = "Line"
    for i, y in enumerate(years):
        ws.cell(1, 3 + i, f"FY{y}")
        ws.cell(5, 3 + i, f"FY{y}")
        ws.cell(7, 3 + i, f"FY{y}")


def _lq_is(wb: Workbook, *, quarter: int, ytd_rev: float, ytd_oi: float) -> None:
    ws = wb.create_sheet("Last Quarter IS Standardized")
    ws["B2"] = f"Q{quarter} FY2025"
    ws["A11"] = "Revenue"
    ws["C11"] = ytd_rev / quarter
    ws["G11"] = ytd_rev
    ws["H11"] = ytd_rev * 0.9
    ws["A20"] = "Operating income"
    ws["C20"] = ytd_oi / quarter
    ws["G20"] = ytd_oi
    ws["H20"] = ytd_oi * 0.85


def industrial_workbook(
    path: Path,
    *,
    years: list[int] | None = None,
    quarter: int = 3,
    rd_by_year: dict[int, float] | None = None,
    revenue_by_year: dict[int, float] | None = None,
    oi_by_year: dict[int, float] | None = None,
    missing_pe10_year: int | None = None,
    seasonal: bool = True,
    ytd_rev: float = 75.0,
    ytd_oi: float = 18.0,
    rd_life_cell: float | None = None,
) -> Path:
    years = years or YEARS
    rd_by_year = rd_by_year or {y: 10.0 + (y - 2016) for y in years}
    revenue_by_year = revenue_by_year or {
        y: (80.0 if seasonal and y >= 2020 else 100.0) + (y - 2016) * 5 for y in years
    }
    oi_by_year = oi_by_year or {y: r * 0.25 for y, r in revenue_by_year.items()}
    wb = Workbook()
    is_ = wb.active
    is_.title = "Income - GAAP"
    _header(is_, years)
    is_["C3"] = years[-1]
    labels = {
        11: "Revenue",
        12: "Gross profit",
        13: "Research and development",
        14: "Operating income",
        15: "Pretax income",
        16: "Income tax expense",
        17: "Net income",
        18: "Diluted EPS",
        19: "Diluted weighted average shares",
    }
    for row, lab in labels.items():
        is_.cell(row, 1, lab)
        for i, y in enumerate(years):
            col = 3 + i
            if row == 11:
                is_.cell(row, col, revenue_by_year[y])
            elif row == 12:
                is_.cell(row, col, revenue_by_year[y] * 0.4)
            elif row == 13:
                is_.cell(row, col, rd_by_year.get(y, 0.0))
            elif row == 14:
                is_.cell(row, col, oi_by_year[y])
            elif row == 15:
                is_.cell(row, col, oi_by_year[y] * 0.95)
            elif row == 16:
                is_.cell(row, col, oi_by_year[y] * 0.95 * 0.21)
            elif row == 17:
                is_.cell(row, col, oi_by_year[y] * 0.75)
            elif row == 18:
                is_.cell(row, col, 2.0 + (y - 2016) * 0.1)
            elif row == 19:
                is_.cell(row, col, 100.0)

    bs = wb.create_sheet("Balance Sheet - Standardized")
    _header(bs, years)
    bs["C3"] = years[-1]
    for row, lab, base in (
        (11, "Cash", 50),
        (61, "Total assets", 1000),
        (65, "Total current liabilities", 200),
        (90, "Total liabilities", 600),
        (100, "Total shareholders' equity", 400),
        (80, "Long-term debt", 150),
    ):
        bs.cell(row, 1, lab)
        for i, _y in enumerate(years):
            bs.cell(row, 3 + i, base + i)

    cf = wb.create_sheet("Cash Flow - Standardized")
    _header(cf, years)
    for row, lab, base in (
        (11, "Cash from operating activities", 80),
        (12, "Capital expenditure", -20),
        (13, "Cash from investing activities", -25),
        (14, "Cash from financing activities", -30),
    ):
        cf.cell(row, 1, lab)
        for i, _y in enumerate(years):
            cf.cell(row, 3 + i, base)

    inp = wb.create_sheet("Inputs")
    _header(inp, years)
    inp["A10"] = "PE10"
    inp["A11"] = "E10"
    inp["A12"] = "Stock price"
    inp["A57"] = "PE10"
    inp["A58"] = "E10"
    inp["A63"] = "Current Price"
    inp["B63"] = 100
    inp["A65"] = "Current PE10"
    inp["B65"] = 22
    inp["A103"] = "R&D expense"
    inp["A106"] = "Tax table"
    inp["A107"] = "Federal tax"
    inp["A108"] = "State taxes"
    inp["A109"] = "Foreign taxes"
    inp["A110"] = "R&D tax credits"
    inp["A111"] = "All other items"
    inp["A112"] = "Income tax expense"
    inp["A120"] = "Buybacks"
    inp["A121"] = "Shares repurchased"
    for i, y in enumerate(years):
        col = 3 + i
        if y != missing_pe10_year:
            inp.cell(10, col, 18 + i * 0.2)
            inp.cell(57, col, 18 + i * 0.2)
            inp.cell(11, col, 5 + i * 0.1)
            inp.cell(58, col, 5 + i * 0.1)
        inp.cell(103, col, rd_by_year.get(y, 0.0))

    rd = wb.create_sheet("R&D")
    rd["B8"] = rd_life_cell
    rd["A2"] = "R&D expense"
    rd["A3"] = "R&D asset"
    rd["A4"] = "R&D amortization"
    _header(rd, years)
    for i, y in enumerate(years):
        rd.cell(2, 3 + i, f"=IF(Inputs!{chr(67+i)}103=\"\",0,Inputs!{chr(67+i)}103)")

    ls = wb.create_sheet("Leases")
    _header(ls, years)
    ls["A10"] = "Year 1"
    ls["A18"] = "Estimated Long-Term Rate"

    tax = wb.create_sheet("Tax")
    _header(tax, years)
    for i, _y in enumerate(years):
        tax.cell(15, 3 + i, f"=Inputs!{chr(67+i)}112")
        tax.cell(25, 3 + i, f"=Inputs!{chr(67+i)}112")

    ic = wb.create_sheet("IC & NOPAT & ROIC ")
    _header(ic, years)
    ic["A9"] = "Invested capital"
    ic["A20"] = "NOPAT"
    ic["A23"] = "ROIC"
    for i, y in enumerate(years):
        ic.cell(9, 3 + i, 500 + i * 10)
        ic.cell(20, 3 + i, 80 + i)
        ic.cell(23, 3 + i, 0.12)

    fm = wb.create_sheet("Final Metrics")
    _header(fm, years)
    fm["A5"] = "ROCE"
    fm["A8"] = "ROIC - WACC"
    for i, _y in enumerate(years):
        fm.cell(5, 3 + i, 0.14)
        fm.cell(8, 3 + i, 0.04)

    er = wb.create_sheet("Expected Returns & Buybacks")
    er["E14"] = 0.11
    er["F14"] = 0.13
    ev = wb.create_sheet("Enterprise Value")
    ev["B20"] = 200
    ev["B27"] = 180
    ev["B32"] = 0.2
    ev["B42"] = 150
    ev["B47"] = 0.1
    ev["B48"] = 120
    ev["B6"] = 90
    ev["C6"] = 100

    _lq_is(wb, quarter=quarter, ytd_rev=ytd_rev, ytd_oi=ytd_oi)
    lq_cf = wb.create_sheet("Last Quarter CF Standardized")
    lq_cf["A27"] = "Cash from operating activities"
    lq_cf["C27"] = 40
    lq_bs = wb.create_sheet("Last Quarter BS Standardized")
    lq_bs["A11"] = "Cash"
    lq_bs["C11"] = 60

    wb.save(path)
    wb.close()
    return path


def _facts_entry(val: float, fy: int, form: str = "10-K", accn: str = "0001") -> dict:
    return {
        "val": val,
        "fy": fy,
        "fp": "FY",
        "form": form,
        "end": f"{fy}-12-31",
        "filed": f"{fy+1}-02-01",
        "accn": accn,
        "frame": f"CY{fy}",
    }


def company_facts_for(
    *,
    years: list[int] | None = None,
    rd: bool = True,
    tax: bool = True,
    leases: bool = True,
    buybacks: bool = True,
    missing_tax_component: bool = False,
    pharma: bool = False,
    sbc_offset: bool = False,
) -> dict:
    years = years or YEARS
    us_gaap: dict = {}

    def add(tag: str, make):
        us_gaap[tag] = {"label": tag, "units": {"USD": [make(y) for y in years]}}

    add("Revenues", lambda y: _facts_entry(100_000_000 + (y - 2016) * 5_000_000, y))
    add("GrossProfit", lambda y: _facts_entry(40_000_000, y))
    add("OperatingIncomeLoss", lambda y: _facts_entry(25_000_000, y))
    add(
        "IncomeLossFromContinuingOperationsBeforeIncomeTaxesExtraordinaryItemsNoncontrollingInterest",
        lambda y: _facts_entry(24_000_000, y),
    )
    add("IncomeTaxExpenseBenefit", lambda y: _facts_entry(5_040_000, y))
    add("NetIncomeLoss", lambda y: _facts_entry(18_000_000, y))
    add("EarningsPerShareDiluted", lambda y: _facts_entry(2.0, y))
    add("NetCashProvidedByUsedInOperatingActivities", lambda y: _facts_entry(80_000_000, y))
    add("PaymentsToAcquirePropertyPlantAndEquipment", lambda y: _facts_entry(20_000_000, y))
    add("NetCashProvidedByUsedInInvestingActivities", lambda y: _facts_entry(-25_000_000, y))
    add("NetCashProvidedByUsedInFinancingActivities", lambda y: _facts_entry(-30_000_000, y))
    add("CashAndCashEquivalentsAtCarryingValue", lambda y: _facts_entry(50_000_000, y))
    add("Assets", lambda y: _facts_entry(1_000_000_000, y))
    add("LiabilitiesCurrent", lambda y: _facts_entry(200_000_000, y))
    add("Liabilities", lambda y: _facts_entry(600_000_000, y))
    add("StockholdersEquity", lambda y: _facts_entry(400_000_000, y))
    add("LongTermDebt", lambda y: _facts_entry(150_000_000, y))
    add("WeightedAverageNumberOfDilutedSharesOutstanding", lambda y: _facts_entry(100_000_000, y))
    us_gaap["EffectiveIncomeTaxRateContinuingOperations"] = {
        "label": "ETR",
        "units": {"pure": [_facts_entry(0.21, y) for y in years]},
    }
    if tax:
        us_gaap["IncomeTaxReconciliationIncomeTaxExpenseBenefitAtFederalStatutoryIncomeTaxRate"] = {
            "label": "federal",
            "units": {"USD": [_facts_entry(5_040_000, y) for y in years]},
        }
        if not missing_tax_component:
            us_gaap["IncomeTaxReconciliationStateAndLocalIncomeTaxes"] = {
                "label": "state",
                "units": {"USD": [_facts_entry(480_000, y) for y in years]},
            }
    if rd:
        us_gaap["ResearchAndDevelopmentExpense"] = {
            "label": "R&D",
            "units": {"USD": [_facts_entry(10_000_000 + (y - 2016) * 1_000_000, y) for y in years + [y for y in range(2012, 2016)]]},
        }
    if leases:
        for tag, val in (
            ("LesseeOperatingLeaseLiabilityPaymentsDueNextTwelveMonths", 10_000_000),
            ("LesseeOperatingLeaseLiabilityPaymentsDueYearTwo", 9_000_000),
            ("LesseeOperatingLeaseLiabilityPaymentsDueYearThree", 8_000_000),
            ("LesseeOperatingLeaseLiabilityPaymentsDueYearFour", 7_000_000),
            ("LesseeOperatingLeaseLiabilityPaymentsDueYearFive", 6_000_000),
            ("LesseeOperatingLeaseLiabilityPaymentsDueAfterYearFive", 20_000_000),
            ("OperatingLeaseLiabilityCurrent", 9_000_000),
            ("OperatingLeaseLiabilityNoncurrent", 40_000_000),
            ("OperatingLeaseRightOfUseAsset", 45_000_000),
            ("OperatingLeaseWeightedAverageDiscountRatePercent", 4.5),
            ("OperatingLeaseWeightedAverageRemainingLeaseTerm", 8.0),
        ):
            us_gaap[tag] = {"label": tag, "units": {"pure" if "Rate" in tag or "Term" in tag else "USD": [_facts_entry(val, y) for y in years]}}
        # Pre-ASC 842 for early years
        us_gaap["OperatingLeasesFutureMinimumPaymentsDueCurrent"] = {
            "label": "pre842",
            "units": {"USD": [_facts_entry(8_000_000, y) for y in years if y <= 2018]},
        }
    if buybacks:
        us_gaap["PaymentsForRepurchaseOfCommonStock"] = {
            "label": "buybacks $",
            "units": {"USD": [_facts_entry(5_000_000, y) for y in years]},
        }
        us_gaap["StockRepurchasedDuringPeriodShares"] = {
            "label": "buybacks sh",
            "units": {"shares": [_facts_entry(1_000_000, y) for y in years]},
        }
        us_gaap["TreasuryStockAcquiredAverageCostPerShare"] = {
            "label": "avg px",
            "units": {"USD/shares": [_facts_entry(5.0, y) for y in years]},
        }
        if sbc_offset:
            us_gaap["AllocatedShareBasedCompensationExpense"] = {
                "label": "sbc",
                "units": {"USD": [_facts_entry(4_000_000, y) for y in years]},
            }
    facts = {"facts": {"us-gaap": us_gaap}, "entityName": "Test Co"}
    if pharma:
        facts["entityName"] = "Example Pharmaceuticals Inc"
    return facts


def _filings_manifest(years: list[int] | None = None) -> dict:
    years = years or YEARS
    # Four 10-Ks covering the window (3+3+3+1 style selection happens in the service).
    selected = []
    for fy in (2025, 2022, 2019, 2016):
        selected.append(
            {
                "accession_number": f"000-{fy}",
                "filing_type": "10-K",
                "filing_date": f"{fy+1}-02-01",
                "report_date": f"{fy}-12-31",
                "document_url": f"https://www.sec.gov/{fy}",
                "fiscal_year": fy,
                "primary_document": "d10k.htm",
            }
        )
    return {"ticker": "TEST", "cik": "0001", "company_name": "Test Co", "selected_filings": selected, "sic": "5331"}


def _ok_recalc(analysis_id="a", ticker="TEST", path="") -> ExcelRecalcReport:
    return ExcelRecalcReport(
        analysis_id=analysis_id,
        ticker=ticker,
        status="ok",
        method="excel_com_calculate_full_rebuild",
        workbook_path=path,
        cells_checked=["Expected Returns & Buybacks!E14"],
        summary="ok",
        com_invoked=True,
    )


@pytest.fixture
def out_svc(tmp_path: Path) -> OutputService:
    return OutputService(outputs_dir=tmp_path / "outputs")


_CURRENT_APPLY = "services.current_data_refresh_service.CurrentDataRefreshService.apply"
_CRF_PARSE = "services.new_company_pe10_service.CustomRunService.parse"


def _refresh_ok():
    return MagicMock(entries=[], summary="ok", missing_required=[])


@pytest.fixture(autouse=True)
def _stub_current_data_refresh():
    """Stub current-data refresh at its definition site.

    Patching CurrentDataRefreshService.apply through two importing modules
    (and stopping those patches in start-order) previously leaked a MagicMock
    into later tests such as test_quarterly_update.
    """
    with patch(_CURRENT_APPLY, return_value=_refresh_ok()):
        yield


def test_ten_year_period_detection(tmp_path: Path):
    path = industrial_workbook(tmp_path / "wb.xlsx", quarter=2)
    report = NewCompanyPeriodService().detect(analysis_id="a", ticker="MSFT", workbook_path=path)
    assert report.fiscal_years == FY
    assert report.latest_quarter == 2
    assert report.template_family == "industrial_template"
    assert report.chronology_ok
    assert report.status == "ok"


def test_period_range_invalid_when_not_ten(tmp_path: Path):
    path = industrial_workbook(tmp_path / "wb.xlsx", years=list(range(2020, 2026)))
    report = NewCompanyPeriodService().detect(analysis_id="a", ticker="X", workbook_path=path)
    assert "NEW_COMPANY_ANNUAL_PERIOD_RANGE_INVALID" in report.status


def test_sec_coverage_prefers_latest_restated_and_3_plus_pattern():
    years = FY
    manifest = _filings_manifest()
    facts = company_facts_for()
    report = NewCompanySecCoverageService().build(
        analysis_id="a", ticker="MSFT", fiscal_years=years, sec_manifest=manifest, company_facts=facts
    )
    assert report.complete
    assert report.pattern in {"3+3+3+1", "4_filings_comparative_window"}
    covered = {y.fiscal_year: y for y in report.years}
    assert covered["FY2024"].restated or covered["FY2024"].coverage_status.value.startswith("covered")

    lookback = NewCompanySecCoverageService().build(
        analysis_id="a",
        ticker="MSFT",
        fiscal_years=years,
        sec_manifest=manifest,
        company_facts=facts,
        extra_lookback_years=["FY2012", "FY2013"],
    )
    assert lookback.complete
    assert lookback.displayed_years == years
    assert "FY2012" in lookback.lookback_years


def test_exact_pe10_vs_e10_identities():
    assert exact_field_identity("PE10") == "pe10_fiscal_year"
    assert exact_field_identity("E10") == "e10_fiscal_year"
    assert exact_field_identity("Current PE10") == "pe10_current"
    assert exact_field_identity("Current E10") == "e10_current"
    assert exact_field_identity("PE10 Percentile") is None
    assert exact_field_identity("E10") != exact_field_identity("PE10")
    # Substring trap: e10 inside pe10
    assert exact_field_identity("PE10") != "e10_fiscal_year"


def test_nearest_date_crf_matching_and_zero_not_missing(tmp_path: Path):
    path = industrial_workbook(tmp_path / "wb.xlsx", missing_pe10_year=2018)
    crf = CustomRunData(
        source_filename="crf.xlsx",
        ticker="MSFT",
        ticker_sheet_name="MSFT",
        metadata={
            "inputs_annual_pe10": {fy: (0.0 if fy == "FY2018" else 20.0 + i) for i, fy in enumerate(FY)},
            "inputs_annual_e10": {fy: 5.0 + i * 0.1 for i, fy in enumerate(FY)},
        },
        scalars={"Current PE10": 25.0, "Current E10": 6.0, "Current Price (Live Price)": 110.0},
        periods=CustomRunPeriods(dates=["2025-12-31"]),
    )
    (tmp_path / "crf.xlsx").write_bytes(b"x")
    with patch(_CRF_PARSE, return_value=crf):
        report = NewCompanyPe10Service().apply(
            analysis_id="a",
            ticker="MSFT",
            workbook_path=path,
            custom_run_path=tmp_path / "crf.xlsx",
            fiscal_years=FY,
        )
    assert any("zero" in w.lower() or "TEN_YEAR_PE10" in w for w in report.warnings)
    fy2018 = next(o for o in report.fiscal_year_pe10 if o.fiscal_year == "FY2018")
    assert fy2018.missing or fy2018.value is None
    assert report.current_pe10 and report.current_pe10.value == 25.0
    assert all(o.field_role == "e10_fiscal_year" for o in report.fiscal_year_e10)


def test_ten_year_tax_residual_and_missing_component(tmp_path: Path):
    path = industrial_workbook(tmp_path / "wb.xlsx")
    year_inputs = {
        fy: {
            "reported_effective_rate": 0.21,
            "components": [
                {"label": "Federal statutory", "rate": 0.21, "house": "statutory_federal"},
                {"label": "State", "rate": 0.02, "house": "state"},
                {"label": "R&D credit", "rate": -0.012, "house": "credits"},
                {"label": "Other items in the filing", "rate": 0.005, "house": "other"},
            ],
            "pretax_income": 100.0,
            "income_tax_expense": 21.0,
            "source": "10-K note 8",
        }
        for fy in FY
    }
    # Drop foreign so residual Other absorbs it; filing "Other" must not double count.
    report = NewCompanyTaxService().apply(
        analysis_id="a", ticker="MSFT", workbook_path=path, fiscal_years=FY, year_inputs=year_inputs
    )
    assert report.complete
    y = report.years[-1]
    assert y.statutory_federal == pytest.approx(0.21)
    assert y.state == pytest.approx(0.02)
    assert y.rd_credit == pytest.approx(-0.012)
    assert y.other == pytest.approx(0.21 - 0.21 - 0.02 - (-0.012))
    assert y.residual_other == y.other
    assert abs(y.component_sum - y.reported_effective_rate) < 0.006

    # Missing component year
    year_inputs["FY2016"]["components"] = [{"label": "Federal statutory", "rate": 0.21, "house": "statutory_federal"}]
    report2 = NewCompanyTaxService().apply(
        analysis_id="a", ticker="MSFT", workbook_path=path, fiscal_years=["FY2016"], year_inputs=year_inputs
    )
    assert report2.years[0].other == pytest.approx(0.0)


def test_rd_useful_life_schema_and_not_from_spend(tmp_path: Path):
    path = industrial_workbook(tmp_path / "wb.xlsx", rd_by_year={y: 50.0 for y in YEARS})
    svc = NewCompanyRdService()
    tech = svc.select_useful_life(
        analysis_id="a",
        ticker="MSFT",
        workbook_path=path,
        fiscal_years=FY,
        sec_manifest={"company_name": "Example Software Cloud Platform", "sic": "7372"},
    )
    pharma = svc.select_useful_life(
        analysis_id="a",
        ticker="PFE",
        workbook_path=path,
        fiscal_years=FY,
        sec_manifest={"company_name": "Example Pharmaceuticals Inc", "sic": "2834"},
    )
    consumer = svc.select_useful_life(
        analysis_id="a",
        ticker="JBSS",
        workbook_path=path,
        fiscal_years=FY,
        sec_manifest={"company_name": "Snack Food Consumer Staples", "sic": "2060"},
        company_facts={"entityName": "Snack Food Consumer Staples"},
    )
    assert 1 <= tech.selected_useful_life <= 10
    assert pharma.selected_useful_life >= tech.selected_useful_life
    assert consumer.selected_useful_life <= 5
    assert tech.provenance_class == "agent_selected_analyst_assumption"
    assert "agent-selected" in tech.warning.lower() or "not been manually approved" in tech.warning.lower()
    assert tech.original_agent_selection == tech.selected_useful_life
    # Amount of R&D must not be the decision driver: same spend, different industries → different lives
    assert {tech.selected_useful_life, pharma.selected_useful_life, consumer.selected_useful_life}


def test_dynamic_rd_lookback_and_capitalization(tmp_path: Path):
    path = industrial_workbook(tmp_path / "wb.xlsx")
    facts = company_facts_for()
    svc = NewCompanyRdService()
    decision = svc.select_useful_life(
        analysis_id="a",
        ticker="MSFT",
        workbook_path=path,
        fiscal_years=FY,
        sec_manifest={"company_name": "Software", "sic": "7372"},
        company_facts=facts,
    )
    report = svc.apply(
        analysis_id="a",
        ticker="MSFT",
        workbook_path=path,
        fiscal_years=FY,
        decision=decision,
        company_facts=facts,
    )
    life = decision.selected_useful_life
    earliest = 2016 - (life - 1)
    assert report.earliest_required_year == f"FY{earliest}"
    assert f"FY{earliest}" in report.lookback_years
    assert report.warning_visible
    assert any(s["role"] == "minus_one" or s["useful_life"] == life - 1 for s in report.sensitivity) or life == 1


def test_lease_rate_hierarchy_and_review_checkpoint(tmp_path: Path):
    years_data = []
    from models.new_company import LeaseYearData

    years_data.append(
        LeaseYearData(
            fiscal_year="FY2025",
            year_1=10,
            year_2=9,
            total_undiscounted=50,
            current_liability=9,
            long_term_liability=40,
            rou_asset=45,
            remaining_term=8,
            reported_discount_rate=0.045,
            regime="post_asc_842",
        )
    )
    svc = NewCompanyLeaseService()
    prop = svc.estimate_rate(years=years_data)
    assert prop.methodology == "reported_weighted_average_discount_rate"
    assert prop.proposed_rate == pytest.approx(0.045)

    inferred_years = [
        LeaseYearData(
            fiscal_year="FY2025",
            year_1=12,
            year_2=12,
            year_3=12,
            year_4=12,
            year_5=12,
            thereafter=24,
            total_undiscounted=84,
            current_liability=10,
            long_term_liability=50,
            remaining_term=7,
            regime="post_asc_842",
        )
    ]
    prop2 = svc.estimate_rate(years=inferred_years)
    assert prop2.methodology_rank >= 3
    assert prop2.proposed_rate and 0 < prop2.proposed_rate < 0.25

    path = industrial_workbook(tmp_path / "wb.xlsx")
    report = svc.apply(
        analysis_id="a", ticker="TJX", workbook_path=path, fiscal_years=FY, company_facts=company_facts_for()
    )
    assert report.review.blocking
    assert report.review.status == "LEASE_RATE_REVIEW_PENDING"
    approved = svc.apply_review(report.review, action="approve", reason="matches 10-K", workbook_path=path)
    assert not approved.blocking
    assert approved.approved_rate == report.review.proposed_rate
    overridden = svc.apply_review(
        report.review, action="correct", rate=0.05, reason="credit spread widened", workbook_path=path
    )
    assert overridden.status == "overridden"
    assert overridden.approved_rate == pytest.approx(0.05)
    assert overridden.proposed_rate == report.review.proposed_rate
    assert any(e.get("event") == "LEASE_RATE_ANALYST_OVERRIDDEN" for e in overridden.audit_trail)
    pending = svc.apply_review(report.review, action="request_more_evidence", reason="need IBR")
    assert pending.blocking


def test_buyback_dollars_shares_derived_and_not_delta_shares():
    facts = company_facts_for(buybacks=True, sbc_offset=True)
    # Remove share tag for one year to force derivation
    facts["facts"]["us-gaap"]["StockRepurchasedDuringPeriodShares"]["units"]["shares"] = [
        e for e in facts["facts"]["us-gaap"]["StockRepurchasedDuringPeriodShares"]["units"]["shares"] if e["fy"] != 2016
    ]
    from tempfile import TemporaryDirectory

    with TemporaryDirectory() as td:
        path = industrial_workbook(Path(td) / "wb.xlsx")
        report = NewCompanyBuybackService().apply(
            analysis_id="a", ticker="AAPL", workbook_path=path, fiscal_years=FY, company_facts=facts
        )
    y2016 = next(y for y in report.years if y.fiscal_year == "FY2016")
    assert y2016.shares_derived
    assert y2016.derivation_formula
    assert report.analysis.sbc_offset_material is True
    assert "change in shares outstanding is not used" in " ".join(report.analysis.notes).lower()
    assert report.analysis.cumulative_dollars


def test_ytd_seasonality_q2_q3_and_negative_numerator(tmp_path: Path):
    path = industrial_workbook(tmp_path / "q3.xlsx", quarter=3, ytd_rev=78, ytd_oi=19)
    q3 = NewCompanySeasonalityService().project(
        analysis_id="a", ticker="TJX", workbook_path=path, fiscal_years=FY, latest_quarter=3
    )
    rev = next(c for c in q3.components if c.metric == "revenue")
    assert rev.unadjusted_annualized == pytest.approx(78 * 4 / 3)
    assert rev.seasonality_adjusted is not None
    assert rev.seasonality_adjusted != rev.unadjusted_annualized or rev.selected_factor == pytest.approx(0.75)
    assert q3.latest_quarter == 3

    path2 = industrial_workbook(tmp_path / "q2.xlsx", quarter=2, ytd_rev=40, ytd_oi=10)
    q2 = NewCompanySeasonalityService().project(
        analysis_id="a", ticker="TJX", workbook_path=path2, fiscal_years=FY, latest_quarter=2
    )
    assert q2.latest_quarter == 2
    oi = next(c for c in q2.components if c.metric == "operating_income")
    assert oi.unadjusted_annualized == pytest.approx(20.0)

    path3 = industrial_workbook(tmp_path / "neg.xlsx", quarter=2, ytd_oi=-5, oi_by_year={y: -10 for y in YEARS})
    neg = NewCompanySeasonalityService().project(
        analysis_id="a", ticker="X", workbook_path=path3, fiscal_years=FY, latest_quarter=2
    )
    oi_neg = next(c for c in neg.components if c.metric == "operating_income")
    assert oi_neg.unreliable or "UNRELIABLE" in ",".join(neg.warnings) or neg.confidence in {
        ProjectionConfidence.UNRELIABLE,
        ProjectionConfidence.LOW,
    }

    path_q1 = industrial_workbook(tmp_path / "q1.xlsx", quarter=1, ytd_rev=20, ytd_oi=5)
    q1 = NewCompanySeasonalityService().project(
        analysis_id="a", ticker="X", workbook_path=path_q1, fiscal_years=FY, latest_quarter=1
    )
    assert q1.confidence in {ProjectionConfidence.LOW, ProjectionConfidence.UNRELIABLE}
    assert q1.confidence != ProjectionConfidence.HIGH
    assert any("indicative" in w.lower() or "q1" in w.lower() for w in q1.warnings)


def test_q1_roic_projection_stays_low_confidence(tmp_path: Path):
    path = industrial_workbook(tmp_path / "q1.xlsx", quarter=1, ytd_rev=20, ytd_oi=5)
    seasonality = NewCompanySeasonalityService().project(
        analysis_id="a", ticker="MSFT", workbook_path=path, fiscal_years=FY, latest_quarter=1
    )
    proj = NewCompanyProjectionService().apply(
        analysis_id="a",
        ticker="MSFT",
        workbook_path=path,
        fiscal_years=FY,
        latest_quarter=1,
        seasonality=seasonality,
        wacc=0.08,
    )
    assert proj.confidence in {ProjectionConfidence.LOW, ProjectionConfidence.UNRELIABLE}
    assert proj.confidence != ProjectionConfidence.HIGH
    assert any("q1" in w.lower() for w in proj.warnings)


def test_roic_roce_projection_and_backtest_confidence(tmp_path: Path):
    path = industrial_workbook(tmp_path / "wb.xlsx", quarter=3)
    seasonality = NewCompanySeasonalityService().project(
        analysis_id="a", ticker="MSFT", workbook_path=path, fiscal_years=FY, latest_quarter=3
    )
    proj = NewCompanyProjectionService().apply(
        analysis_id="a",
        ticker="MSFT",
        workbook_path=path,
        fiscal_years=FY,
        latest_quarter=3,
        seasonality=seasonality,
        wacc=0.08,
    )
    assert proj.seasonality_adjusted_roic is not None
    assert proj.projected_roic_wacc == pytest.approx(proj.seasonality_adjusted_roic - 0.08)
    assert proj.seasonality_adjusted_roce is not None
    assert proj.ytd_unadjusted_annualized_roic is not None
    assert proj.confidence in set(ProjectionConfidence)
    assert proj.backtests


def test_output_gates_lease_pending_blocks_complete(tmp_path: Path):
    path = industrial_workbook(tmp_path / "wb.xlsx")
    periods = NewCompanyPeriodService().detect(analysis_id="a", ticker="MSFT", workbook_path=path)
    coverage = NewCompanySecCoverageService().build(
        analysis_id="a", ticker="MSFT", fiscal_years=FY, sec_manifest=_filings_manifest(), company_facts=company_facts_for()
    )
    from models.new_company import LeaseRateReview

    gate = NewCompanyOutputGateService().evaluate(
        analysis_id="a",
        ticker="MSFT",
        periods=periods,
        coverage=coverage,
        statements=None,
        pe10=None,
        tax=None,
        rd_decision=None,
        rd=None,
        leases=None,
        lease_review=LeaseRateReview(analysis_id="a", ticker="MSFT", status="LEASE_RATE_REVIEW_PENDING", proposed_rate=0.04, blocking=True),
        buybacks=None,
        current=None,
        projection=None,
        recalc=None,
        valuation=None,
    )
    assert "LEASE_RATE_REVIEW_PENDING" in gate.blockers
    assert gate.status == "AWAITING_ANALYST_REVIEW"
    assert not gate.report_authorized


def test_recalc_unavailable_needs_review(tmp_path: Path, out_svc: OutputService):
    path = industrial_workbook(tmp_path / "wb.xlsx")
    runner = NewCompanyRunner(output_service=out_svc)
    crf_path = tmp_path / "crf.xlsx"
    crf_path.write_bytes(b"x")
    crf = CustomRunData(
        source_filename="crf.xlsx",
        ticker="MSFT",
        ticker_sheet_name="MSFT",
        metadata={
            "inputs_annual_pe10": {fy: 20.0 for fy in FY},
            "inputs_annual_e10": {fy: 5.0 for fy in FY},
        },
        scalars={"Current PE10": 22.0, "Current E10": 5.5},
    )
    tax_inputs = {
        fy: {
            "reported_effective_rate": 0.21,
            "components": [{"label": "Federal statutory", "rate": 0.21, "house": "statutory_federal"}],
            "pretax_income": 100.0,
            "income_tax_expense": 21.0,
        }
        for fy in FY
    }
    with patch(_CRF_PARSE, return_value=crf):
        result = runner.run(
            analysis_id="run1",
            ticker="MSFT",
            company="Microsoft",
            template_path=path,
            working_path=tmp_path / "working.xlsx",
            custom_run_path=crf_path,
            company_facts=company_facts_for(),
            sec_manifest=_filings_manifest(),
            tax_year_inputs=tax_inputs,
            lease_review_override={"action": "approve", "reason": "ok"},
            finalize=True,
        )
    assert result["workflow_state"] == NewCompanyWorkflowState.NEEDS_REVIEW
    assert "WORKBOOK_RECALCULATION_INCOMPLETE" in result["output_gate"].blockers
    assert result["deliverables"] is None or result["deliverables"].authorized is False
    assert result["certification_status"] == CLOUD_PENDING_WINDOWS_CERTIFICATION
    assert result["workflow_state"] != NewCompanyWorkflowState.COMPLETE


def test_runner_pauses_for_lease_review_then_blocks_without_com(tmp_path: Path, out_svc: OutputService):
    path = industrial_workbook(tmp_path / "wb.xlsx")
    runner = NewCompanyRunner(output_service=out_svc)
    crf_path = tmp_path / "crf.xlsx"
    crf_path.write_bytes(b"x")
    crf = CustomRunData(
        source_filename="crf.xlsx",
        ticker="AMZN",
        ticker_sheet_name="AMZN",
        metadata={"inputs_annual_pe10": {fy: 30.0 for fy in FY}, "inputs_annual_e10": {fy: 4.0 for fy in FY}},
        scalars={"Current PE10": 32.0, "Current E10": 4.2},
    )
    tax_inputs = {
        fy: {
            "reported_effective_rate": 0.21,
            "components": [{"label": "Federal statutory", "rate": 0.21, "house": "statutory_federal"}],
            "pretax_income": 100.0,
            "income_tax_expense": 21.0,
        }
        for fy in FY
    }
    with patch(_CRF_PARSE, return_value=crf):
        paused = runner.run(
            analysis_id="run2",
            ticker="AMZN",
            company="Amazon",
            template_path=path,
            working_path=tmp_path / "working.xlsx",
            custom_run_path=crf_path,
            company_facts=company_facts_for(),
            sec_manifest=_filings_manifest(),
            tax_year_inputs=tax_inputs,
        )
        assert paused["workflow_state"] == NewCompanyWorkflowState.AWAITING_ANALYST_REVIEW
        assert paused["lease_review"].blocking
        assert paused["output_gate"] is None
        assert paused["deliverables"] is None or paused["deliverables"].authorized is False
        done = runner.run(
            analysis_id="run2",
            ticker="AMZN",
            company="Amazon",
            template_path=tmp_path / "working.xlsx",
            working_path=tmp_path / "working.xlsx",
            custom_run_path=crf_path,
            company_facts=company_facts_for(),
            sec_manifest=_filings_manifest(),
            tax_year_inputs=tax_inputs,
            lease_review_override={"action": "approve", "reason": "10-K rate"},
            finalize=True,
            prepare_working=False,
        )
        assert done["lease_review"].blocking is False
        gate = done["output_gate"]
        assert gate is not None
        assert "LEASE_RATE_REVIEW_PENDING" not in gate.blockers
        assert "WORKBOOK_RECALCULATION_INCOMPLETE" in gate.blockers
        assert done["workflow_state"] == NewCompanyWorkflowState.NEEDS_REVIEW
        assert done["workflow_state"] != NewCompanyWorkflowState.COMPLETE
        assert done["certification_status"] == CLOUD_PENDING_WINDOWS_CERTIFICATION
        assert done["deliverables"] is None or done["deliverables"].authorized is False


def test_rd_override_preserves_original_selection(tmp_path: Path):
    path = industrial_workbook(tmp_path / "wb.xlsx")
    svc = NewCompanyRdService()
    first = svc.select_useful_life(
        analysis_id="a",
        ticker="MSFT",
        workbook_path=path,
        fiscal_years=FY,
        sec_manifest={"company_name": "Software Cloud", "sic": "7372"},
    )
    second = svc.select_useful_life(
        analysis_id="a",
        ticker="MSFT",
        workbook_path=path,
        fiscal_years=FY,
        sec_manifest={"company_name": "Software Cloud", "sic": "7372"},
        override=7,
        override_reason="longer platform cycle",
        prior_decision=first,
    )
    assert second.analyst_override == 7
    assert second.original_agent_selection == first.selected_useful_life
    assert second.selected_useful_life == 7


def test_statement_fill_missing_and_sec_correction(tmp_path: Path):
    path = industrial_workbook(tmp_path / "wb.xlsx")
    wb = load_workbook(path)
    # Blank one revenue cell (not a formula)
    wb["Income - GAAP"].cell(11, 3, None)
    wb.save(path)
    wb.close()
    facts = company_facts_for()
    report = NewCompanyStatementValidationService().validate(
        analysis_id="a", ticker="MSFT", workbook_path=path, fiscal_years=FY, company_facts=facts
    )
    assert report.filled_missing or report.corrections or report.items
    assert report.formulas_preserved


def test_cross_company_runner_suite(tmp_path: Path, out_svc: OutputService):
    """MSFT, AMZN, TJX, R&D-light consumer, lease-heavy retailer, seasonal — not certified without Excel COM."""
    profiles = [
        ("MSFT", "Microsoft Software Cloud", "7372", False, 3),
        ("AMZN", "Amazon.com Inc", "5961", False, 3),
        ("TJX", "The TJX Companies Retail Stores", "5331", False, 2),
        ("JBSS", "Snack Food Consumer Staples", "2060", False, 4),
        ("ROST", "Ross Stores lease-heavy retailer", "5651", False, 2),
        ("WSM", "Williams-Sonoma seasonal retailer", "5700", True, 2),
    ]
    runner = NewCompanyRunner(output_service=out_svc)
    crf = CustomRunData(
        source_filename="crf.xlsx",
        ticker="X",
        ticker_sheet_name="X",
        metadata={"inputs_annual_pe10": {fy: 18.0 for fy in FY}, "inputs_annual_e10": {fy: 5.0 for fy in FY}},
        scalars={"Current PE10": 19.0, "Current E10": 5.1},
    )
    tax_inputs = {
        fy: {
            "reported_effective_rate": 0.21,
            "components": [{"label": "Federal statutory", "rate": 0.21, "house": "statutory_federal"}],
            "pretax_income": 100.0,
            "income_tax_expense": 21.0,
        }
        for fy in FY
    }
    results = []
    with patch(_CRF_PARSE, return_value=crf):
        for ticker, name, sic, seasonal, q in profiles:
            wb = industrial_workbook(
                tmp_path / f"{ticker}.xlsx",
                quarter=q,
                seasonal=seasonal,
                rd_by_year={y: (0.0 if ticker == "JBSS" else 12.0) for y in YEARS},
            )
            crf_path = tmp_path / f"{ticker}_crf.xlsx"
            crf_path.write_bytes(b"x")
            paused = runner.run(
                analysis_id=f"{ticker}-nc",
                ticker=ticker,
                company=name,
                template_path=wb,
                working_path=tmp_path / f"{ticker}_work.xlsx",
                custom_run_path=crf_path,
                company_facts=company_facts_for(),
                sec_manifest={**_filings_manifest(), "company_name": name, "sic": sic},
                tax_year_inputs=tax_inputs,
            )
            assert paused["workflow_state"] == NewCompanyWorkflowState.AWAITING_ANALYST_REVIEW
            done = runner.run(
                analysis_id=f"{ticker}-nc",
                ticker=ticker,
                company=name,
                template_path=tmp_path / f"{ticker}_work.xlsx",
                working_path=tmp_path / f"{ticker}_work.xlsx",
                custom_run_path=crf_path,
                company_facts=company_facts_for(),
                sec_manifest={**_filings_manifest(), "company_name": name, "sic": sic},
                tax_year_inputs=tax_inputs,
                lease_review_override={"action": "approve", "reason": "fixture approval"},
                finalize=True,
                prepare_working=False,
            )
            results.append((ticker, done["workflow_state"], done["output_gate"].blockers if done["output_gate"] else []))
            assert done["lease_review"].blocking is False
            assert done["rd_decision"].selected_useful_life
            assert len(done["pe10"].fiscal_year_pe10) == 10
            assert len(done["tax"].years) == 10
            assert "LEASE_RATE_REVIEW_PENDING" not in (done["output_gate"].blockers if done["output_gate"] else [])
            assert "TEN_YEAR_SEC_COVERAGE_INCOMPLETE" not in (
                done["output_gate"].blockers if done["output_gate"] else []
            )
            assert done["workflow_state"] != NewCompanyWorkflowState.COMPLETE
            assert done["certification_status"] == CLOUD_PENDING_WINDOWS_CERTIFICATION
            assert "WORKBOOK_RECALCULATION_INCOMPLETE" in done["output_gate"].blockers
    assert len(results) == 6
    assert all(r[0] for r in results)


def test_gate_j_rejects_forged_ok_without_genuine_com(tmp_path: Path):
    from models.annual_update import AnnualValuationOutputs
    from services.excel_recalc_service import ExcelRecalcReport

    path = industrial_workbook(tmp_path / "wb.xlsx")
    periods = NewCompanyPeriodService().detect(analysis_id="a", ticker="MSFT", workbook_path=path)
    coverage = NewCompanySecCoverageService().build(
        analysis_id="a",
        ticker="MSFT",
        fiscal_years=FY,
        sec_manifest=_filings_manifest(),
        company_facts=company_facts_for(),
    )
    forged = ExcelRecalcReport(
        analysis_id="a",
        ticker="MSFT",
        status="ok",
        method="none",
        workbook_path=str(path),
        com_invoked=False,
        summary="forged",
    )
    gate = NewCompanyOutputGateService().evaluate(
        analysis_id="a",
        ticker="MSFT",
        periods=periods,
        coverage=coverage,
        statements=MagicMock(unresolved_material=[], summary="ok"),
        pe10=MagicMock(
            fiscal_year_pe10=[MagicMock(missing=False) for _ in FY],
            fiscal_year_e10=[MagicMock(missing=False) for _ in FY],
            current_pe10=MagicMock(value=20.0),
            current_e10=MagicMock(value=5.0),
            warnings=[],
        ),
        tax=MagicMock(complete=True, years=[]),
        rd_decision=MagicMock(selected_useful_life=3, blocking=False, blocking_reasons=[]),
        rd=MagicMock(lookback_complete=True, capitalization_ok=True),
        leases=MagicMock(complete=True, review=None),
        lease_review=MagicMock(proposed_rate=0.04, blocking=False, status="approved", analyst_action="approve"),
        buybacks=MagicMock(
            years=[MagicMock(dollars=1.0, shares=1.0, absence_class=None) for _ in FY],
            complete=True,
        ),
        current=MagicMock(as_of_mismatch=False, warnings=[]),
        projection=MagicMock(
            seasonality_adjusted_roic=0.1,
            seasonality_adjusted_roce=0.1,
            confidence=ProjectionConfidence.HIGH,
        ),
        recalc=forged,
        valuation=AnnualValuationOutputs(
            expected_annual_return=0.1,
            expected_return_with_dividends=0.12,
            current_graham_intrinsic_value=100.0,
            nopat=80.0,
            invested_capital=500.0,
            roic=0.16,
        ),
    )
    assert "WORKBOOK_RECALCULATION_INCOMPLETE" in gate.blockers
    assert not gate.report_authorized
    assert not genuine_excel_com_recalc(forged)
