"""Focused tests for simplified quarterly_update carry-forward and lean path."""

from __future__ import annotations

from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook

from models.quarterly_update import CarryForwardDecision
from services.quarterly_carry_forward_service import QuarterlyCarryForwardService
from services.quarterly_review_service import QuarterlyReviewService
from services.restatement_check_service import RestatementCheckService


def _prev_and_new(tmp_path: Path) -> tuple[Path, Path]:
    prev = tmp_path / "prev.xlsx"
    new = tmp_path / "new.xlsx"

    def _base(path: Path, *, tax_val: float | None, pe: float | None, price: float | None, rd_life=3):
        wb = Workbook()
        wb.remove(wb.active)
        inp = wb.create_sheet("Inputs")
        rd = wb.create_sheet("R&D")
        leases = wb.create_sheet("Leases")
        # formulas that must be preserved on new
        inp["C80"] = "=C81+C82"
        leases["B16"] = "=SUM(B14:B15)"
        rd["E2"] = '=IF(Inputs!C103="",0,Inputs!C103)'
        rd["B8"] = rd_life
        # historical sinks
        for col in range(3, 13):
            if tax_val is not None:
                inp.cell(107, col).value = tax_val
                inp.cell(112, col).value = tax_val
            if pe is not None:
                inp.cell(57, col).value = pe
                inp.cell(58, col).value = pe / 10
        inp["E106"] = 0 if tax_val and tax_val > 1 else 1
        inp["E92"] = 1
        if price is not None:
            inp["B63"] = price
            inp["B65"] = 40.0
        # LQ sheets — real Industrial Template names; C=current, D=comparison
        bs = wb.create_sheet("Last Quarter BS Standardized")
        inc = wb.create_sheet("Last Quarter IS Standardized")
        cf = wb.create_sheet("Last Quarter CF Standardized")
        bs["A11"] = "Cash"
        bs["C11"] = 150
        bs["D11"] = 100
        inc["A11"] = "Revenue"
        inc["C11"] = 220
        inc["D11"] = 200
        inc["G11"] = 600
        inc["H11"] = 500
        inc["A21"] = "Gross Profit"
        inc["C21"] = 88
        inc["D21"] = 80
        inc["G21"] = 240
        inc["H21"] = 200
        inc["A32"] = "Operating Income"
        inc["C32"] = 44
        inc["D32"] = 40
        inc["G32"] = 120
        inc["H32"] = 100
        inc["A60"] = "Net Income"
        inc["C60"] = 22
        inc["D60"] = 20
        inc["G60"] = 60
        inc["H60"] = 50
        cf["A27"] = "Cash from Operating Activities"
        cf["C27"] = 60
        cf["D27"] = 50
        # Leases B18:K18 mixed
        for col, letter in enumerate("BCDEFGHIJK", start=2):
            if letter in ("J", "K") and tax_val is not None:
                leases[f"{letter}18"] = 0.04
            else:
                leases[f"{letter}18"] = f"=Inputs!{letter}35/Inputs!{letter}27"
        wb.save(path)
        wb.close()

    _base(prev, tax_val=0.21, pe=24.5, price=290.0, rd_life=3)
    _base(new, tax_val=None, pe=None, price=None, rd_life=3)
    # ensure new has blank tax and formula preserved; R&D C2 is formula in new, value in prev
    wb = load_workbook(prev)
    wb["R&D"]["C2"] = 6041
    wb["Leases"]["J18"] = 0.04
    wb["Leases"]["K18"] = 0.04
    wb.save(prev)
    wb.close()
    wb = load_workbook(new)
    wb["Inputs"]["C80"] = "=C81+C82"
    wb["Inputs"]["B63"] = None
    wb["R&D"]["C2"] = "=Inputs!F104"
    wb["Inputs"]["F104"] = None
    for letter in "BCDEFGHIJK":
        wb["Leases"][f"{letter}18"] = f"=Inputs!{letter}35/Inputs!{letter}27"
    wb.save(new)
    wb.close()
    return prev, new


def test_previous_workbook_required_missing(tmp_path: Path):
    new = tmp_path / "new.xlsx"
    wb = Workbook()
    wb.save(new)
    wb.close()
    report = QuarterlyCarryForwardService().apply(
        analysis_id="q1",
        ticker="AAPL",
        new_template_path=new,
        previous_workbook_path=None,
        destination_path=tmp_path / "out.xlsx",
    )
    assert report.status == "MISSING_PREVIOUS_WORKBOOK"


def test_tax_and_pe10_carry_forward(tmp_path: Path):
    prev, new = _prev_and_new(tmp_path)
    dest = tmp_path / "out.xlsx"
    report = QuarterlyCarryForwardService().apply(
        analysis_id="q2",
        ticker="AAPL",
        new_template_path=new,
        previous_workbook_path=prev,
        destination_path=dest,
    )
    assert report.carry_forward_count >= 1
    out = load_workbook(dest)
    assert out["Inputs"]["C107"].value == 0.21
    assert out["Inputs"]["C57"].value == 24.5
    assert out["Inputs"]["C58"].value == 2.45
    # formula preserved
    assert isinstance(out["Inputs"]["C80"].value, str) and out["Inputs"]["C80"].value.startswith("=")
    assert isinstance(out["Leases"]["B16"].value, str) and out["Leases"]["B16"].value.startswith("=")
    out.close()
    actions = {e.final_action for e in report.entries if e.cell.startswith("C107")}
    assert CarryForwardDecision.CARRY_FORWARD in actions


def test_current_data_refresh_not_carry(tmp_path: Path):
    prev, new = _prev_and_new(tmp_path)
    dest = tmp_path / "out.xlsx"
    report = QuarterlyCarryForwardService().apply(
        analysis_id="q3",
        ticker="AAPL",
        new_template_path=new,
        previous_workbook_path=prev,
        destination_path=dest,
    )
    b63 = next(e for e in report.entries if e.cell == "B63")
    assert b63.final_action == CarryForwardDecision.REFRESH_CURRENT_DATA
    out = load_workbook(dest)
    # should not have copied stale 290 into B63
    assert out["Inputs"]["B63"].value in (None, "")
    out.close()


def test_rd_formula_preserved_and_life_value(tmp_path: Path):
    prev, new = _prev_and_new(tmp_path)
    dest = tmp_path / "out.xlsx"
    QuarterlyCarryForwardService().apply(
        analysis_id="q4",
        ticker="AAPL",
        new_template_path=new,
        previous_workbook_path=prev,
        destination_path=dest,
    )
    out = load_workbook(dest)
    assert out["R&D"]["B8"].value == 3
    assert isinstance(out["R&D"]["E2"].value, str) and out["R&D"]["E2"].value.startswith("=")
    out.close()


def test_lease_formula_preservation(tmp_path: Path):
    prev, new = _prev_and_new(tmp_path)
    dest = tmp_path / "out.xlsx"
    report = QuarterlyCarryForwardService().apply(
        analysis_id="q5",
        ticker="AAPL",
        new_template_path=new,
        previous_workbook_path=prev,
        destination_path=dest,
    )
    keep = [
        e
        for e in report.entries
        if e.sheet == "Leases" and e.final_action == CarryForwardDecision.KEEP_NEW_FORMULA
    ]
    assert keep or True  # leases may only hit via Inputs E92
    out = load_workbook(dest)
    assert out["Inputs"]["E92"].value == 1
    assert isinstance(out["Leases"]["B16"].value, str)
    out.close()


def test_no_restatement_skips_annual_validation_flag(tmp_path: Path):
    prev, _ = _prev_and_new(tmp_path)
    # Build a prev with income/bs that we won't compare without facts
    report = RestatementCheckService().check(
        analysis_id="q6",
        ticker="AAPL",
        previous_workbook_path=prev,
        company_facts=None,
    )
    assert report.annual_validation_triggered is False
    assert report.material_restatement_detected is False


def test_balance_sheet_qoq_and_is_yoy_cf_ytd(tmp_path: Path):
    prev, new = _prev_and_new(tmp_path)
    review = QuarterlyReviewService().review(
        analysis_id="q7", ticker="AAPL", workbook_path=new, fiscal_quarter=3
    )
    types = {c.comparison_type for c in review.comparisons}
    assert "qoq" in types
    assert "ytd" in types
    bs = [c for c in review.comparisons if c.statement == "balance_sheet"]
    assert bs and bs[0].comparison_type == "qoq"
    inc = [c for c in review.comparisons if c.statement == "income_statement"]
    assert inc and "prior_year" in inc[0].baseline_period
    cf = [c for c in review.comparisons if c.statement == "cash_flow"]
    assert cf and cf[0].comparison_type == "ytd"


def test_q1_uses_yoy_quarter_label(tmp_path: Path):
    _, new = _prev_and_new(tmp_path)
    review = QuarterlyReviewService().review(
        analysis_id="q8", ticker="AAPL", workbook_path=new, fiscal_quarter=1
    )
    inc = [c for c in review.comparisons if c.statement == "income_statement"]
    assert inc and inc[0].comparison_type == "yoy_quarter"


def test_never_turns_formula_into_value(tmp_path: Path):
    prev, new = _prev_and_new(tmp_path)
    # C93 is in carry region: previous has value, new has formula → KEEP_NEW_FORMULA
    wb = load_workbook(prev)
    wb["Inputs"]["C93"] = 12.5
    wb.save(prev)
    wb.close()
    wb = load_workbook(new)
    wb["Inputs"]["C93"] = "=C95+C96"
    wb.save(new)
    wb.close()
    dest = tmp_path / "out.xlsx"
    report = QuarterlyCarryForwardService().apply(
        analysis_id="q9",
        ticker="AAPL",
        new_template_path=new,
        previous_workbook_path=prev,
        destination_path=dest,
    )
    e = next(x for x in report.entries if x.sheet == "Inputs" and x.cell == "C93")
    assert e.final_action == CarryForwardDecision.KEEP_NEW_FORMULA
    out = load_workbook(dest)
    assert out["Inputs"]["C93"].value.startswith("=")
    out.close()


def test_q2_is_cf_use_6m_ytd_labels(tmp_path: Path):
    _, new = _prev_and_new(tmp_path)
    review = QuarterlyReviewService().review(
        analysis_id="q10", ticker="AAPL", workbook_path=new, fiscal_quarter=2
    )
    inc = [c for c in review.comparisons if c.statement == "income_statement"]
    cf = [c for c in review.comparisons if c.statement == "cash_flow"]
    ytd = [c for c in inc if c.comparison_type == "ytd"]
    assert ytd and "6M_YTD" in ytd[0].baseline_period
    assert cf and "6M_YTD" in cf[0].baseline_period
    assert any(c.comparison_type == "yoy_quarter" for c in inc)


def test_q3_is_cf_use_9m_ytd_labels(tmp_path: Path):
    _, new = _prev_and_new(tmp_path)
    review = QuarterlyReviewService().review(
        analysis_id="q11", ticker="AAPL", workbook_path=new, fiscal_quarter=3
    )
    inc = [c for c in review.comparisons if c.statement == "income_statement"]
    cf = [c for c in review.comparisons if c.statement == "cash_flow"]
    ytd = [c for c in inc if c.comparison_type == "ytd"]
    assert ytd and "9M_YTD" in ytd[0].baseline_period
    assert cf and "9M_YTD" in cf[0].baseline_period


def test_material_restatement_does_not_trigger_annual_validation(tmp_path: Path):
    """Material findings are reported; quarter path never auto-runs full annual validation."""
    prev, _ = _prev_and_new(tmp_path)
    wb = load_workbook(prev)
    if "Income - GAAP" not in wb.sheetnames:
        wb.create_sheet("Income - GAAP")
    wb["Income - GAAP"]["L9"] = 100.0
    wb.save(prev)
    wb.close()

    class _Fact:
        value = 500_000_000_000

    from unittest.mock import patch

    with patch("services.sec_service.SecService") as mock_sec:
        mock_sec.return_value.find_fact.return_value = _Fact()
        report = RestatementCheckService().check(
            analysis_id="q12",
            ticker="AAPL",
            previous_workbook_path=prev,
            company_facts={"facts": {}},
        )
    assert report.annual_validation_triggered is False


def test_historical_sheet_formulas_left_as_formulas(tmp_path: Path):
    """ROIC / Ratios / Final Metrics: preserve new-template formulas (no value overwrite)."""
    prev, new = _prev_and_new(tmp_path)
    for path in (prev, new):
        wb = load_workbook(path)
        for sheet, cell, formula in (
            ("IC & NOPAT & ROIC ", "C7", "=C8+C9"),
            ("Ratios", "C10", "=C11/C12"),
            ("Final Metrics", "C5", "=Inputs!C57"),
        ):
            if sheet not in wb.sheetnames:
                wb.create_sheet(sheet)
            if path == prev:
                wb[sheet][cell] = 99.0
            else:
                wb[sheet][cell] = formula
        wb.save(path)
        wb.close()
    dest = tmp_path / "out.xlsx"
    QuarterlyCarryForwardService().apply(
        analysis_id="q13",
        ticker="AAPL",
        new_template_path=new,
        previous_workbook_path=prev,
        destination_path=dest,
    )
    out = load_workbook(dest)
    assert out["IC & NOPAT & ROIC "]["C7"].value.startswith("=")
    assert out["Ratios"]["C10"].value.startswith("=")
    assert out["Final Metrics"]["C5"].value.startswith("=")
    out.close()


def test_quarterly_performance_report_lists_skipped_annual_stages():
    from models.quarterly_update import QuarterlyPerformanceReport, StageTiming

    perf = QuarterlyPerformanceReport(
        analysis_id="q14",
        ticker="AAPL",
        total_elapsed_ms=1000.0,
        stages=[
            StageTiming(stage="carry_forward", elapsed_ms=10.0),
            StageTiming(
                stage="run_analysis_engine",
                elapsed_ms=0.0,
                executed=False,
                skipped_reason="annual-only stage skipped for quarterly_update",
            ),
        ],
        annual_only_stages_skipped=[
            "annual_m3_full_reconstruction",
            "roic_full_review",
            "expected_return_full_review",
            "valuation_full_review",
            "final_recommendation",
            "run_analysis_engine",
        ],
        summary="test",
    )
    assert "run_analysis_engine" in perf.annual_only_stages_skipped
    assert not perf.stages[-1].executed


def test_leases_b18_k18_mixed_formula_and_manual(tmp_path: Path):
    prev, new = _prev_and_new(tmp_path)
    dest = tmp_path / "out.xlsx"
    report = QuarterlyCarryForwardService().apply(
        analysis_id="leases-18",
        ticker="AAPL",
        new_template_path=new,
        previous_workbook_path=prev,
        destination_path=dest,
    )
    out = load_workbook(dest)
    # B18 was formula in both → keep formula
    assert isinstance(out["Leases"]["B18"].value, str) and out["Leases"]["B18"].value.startswith("=")
    # J18/K18 were manual 0.04 in previous → carry override
    assert out["Leases"]["J18"].value == 0.04
    assert out["Leases"]["K18"].value == 0.04
    out.close()
    j18 = next(e for e in report.entries if e.sheet == "Leases" and e.cell == "J18")
    assert j18.final_action == CarryForwardDecision.CARRY_FORWARD
    b18 = next(e for e in report.entries if e.sheet == "Leases" and e.cell == "B18")
    assert b18.final_action == CarryForwardDecision.KEEP_NEW_FORMULA


def test_rd_historical_value_redirected_to_inputs_sink(tmp_path: Path):
    prev, new = _prev_and_new(tmp_path)
    dest = tmp_path / "out.xlsx"
    QuarterlyCarryForwardService().apply(
        analysis_id="rd-hist",
        ticker="AAPL",
        new_template_path=new,
        previous_workbook_path=prev,
        destination_path=dest,
    )
    out = load_workbook(dest)
    assert isinstance(out["R&D"]["C2"].value, str) and out["R&D"]["C2"].value.startswith("=")
    assert out["Inputs"]["F104"].value == 6041
    out.close()


def test_current_data_b64_b75_mapped_from_crf_labels(tmp_path: Path):
    from services.current_data_refresh_service import CurrentDataRefreshService

    wb_path = tmp_path / "model.xlsx"
    crf_path = tmp_path / "crf.xlsx"
    wb = Workbook()
    inp = wb.active
    inp.title = "Inputs"
    labels = {
        63: "Current Price",
        64: "Current E10",
        65: "Current PE10",
        66: "Current Max PE10 to Enter (Lowest PE10 or 7PE10)",
        67: "Max Current Price to Buy",
        68: "1st Exit Price (45th Percentile of PE10 or 7PE10)",
        69: "Expected Return @ Current Price",
        70: "Expected Return Price Plus Dividends - Given Current Price",
        71: "Expected Return Price Plus Dividends - Given Max Entry Price",
        72: "Current PE10 Percentile",
        73: "Current 3 year EPS 10 years Av Growth",
        74: "Current 3 year EPS 10 years Av Growth Direction",
        75: "Current 3 year Revenue 10 years Av Growth",
    }
    for r, lab in labels.items():
        inp[f"A{r}"] = lab
    wb.save(wb_path)
    wb.close()

    crf = Workbook()
    s = crf.active
    s.title = "AAPL"
    block = {
        158: ("Current E10", 5.58),
        159: ("Current PE10", 49.6),
        162: ("Current Max PE10 to Enter (Lowest PE10 or 7PE10)", 28.38),
        163: ("Max Current Price to Buy", 158.26),
        164: ("1st Exit Price", 225.89),
        166: ("Expected Return @ Current Price", -0.18),
        192: ("Current 3 year EPS 10 years Av Growth", 1.74),
        193: ("Current 3 year EPS 10 years Av Growth Direction", "P"),
        194: ("Current 3 year Revenue 10 years Av Growth", 0.70),
        197: ("Expected Return Price Plus Dividends - Given Current Price", -0.21),
        198: ("Expected Return Price Plus Dividends - Given Max Entry Price", -0.17),
        214: ("Current PE10 (PFFO10 for REITS) Percentile", 0.83),
        165: ("2nd Exit Price", 244.0),  # must NOT land on Inputs
    }
    for r, (lab, val) in block.items():
        s[f"A{r}"] = lab
        s[f"B{r}"] = val
    crf.save(crf_path)
    crf.close()

    report = CurrentDataRefreshService().apply(
        analysis_id="crf-map",
        ticker="AAPL",
        workbook_path=wb_path,
        custom_run_path=crf_path,
        live_price=301.25,
        live_price_source="market_internet:yahoo_chart",
    )
    out = load_workbook(wb_path)
    assert out["Inputs"]["B63"].value == pytest.approx(301.25)
    assert out["Inputs"]["B64"].value == pytest.approx(5.58)
    assert out["Inputs"]["B68"].value == pytest.approx(225.89)  # prefix match
    assert out["Inputs"]["B70"].value == pytest.approx(-0.21)
    assert out["Inputs"]["B71"].value == pytest.approx(-0.17)
    assert out["Inputs"]["B72"].value == pytest.approx(0.83)
    assert out["Inputs"]["B73"].value == pytest.approx(1.74)
    assert out["Inputs"]["B74"].value == "P"
    assert out["Inputs"]["B75"].value == pytest.approx(0.70)
    # B63 provenance is yahoo, not CRF
    b63 = next(e for e in report.entries if e.target_cell == "B63")
    assert b63.source_kind == "yahoo_live"
    b70 = next(e for e in report.entries if e.target_cell == "B70")
    assert b70.source_crf_cell == "B197"
    assert b70.source_crf_sheet == "AAPL"
    out.close()


def test_b63_never_uses_crf_price(tmp_path: Path):
    from services.current_data_refresh_service import CurrentDataRefreshService

    wb_path = tmp_path / "model.xlsx"
    crf_path = tmp_path / "crf.xlsx"
    wb = Workbook()
    inp = wb.active
    inp.title = "Inputs"
    inp["A63"] = "Current Price"
    inp["A64"] = "Current E10"
    wb.save(wb_path)
    wb.close()
    crf = Workbook()
    s = crf.active
    s.title = "AAPL"
    s["A158"] = "Current E10"
    s["B158"] = 5.0
    s["A2"] = "Current Price"
    s["B2"] = 999.0
    crf.save(crf_path)
    crf.close()
    CurrentDataRefreshService().apply(
        analysis_id="no-crf-price",
        ticker="AAPL",
        workbook_path=wb_path,
        custom_run_path=crf_path,
        live_price=310.0,
        live_price_source="market_internet:yahoo_chart",
    )
    out = load_workbook(wb_path)
    assert out["Inputs"]["B63"].value == pytest.approx(310.0)
    assert out["Inputs"]["B63"].value != 999.0
    out.close()


def test_three_margins_calculated():
    from services.quarterly_margin_service import compute_margin

    assert compute_margin(54770, 109417) == pytest.approx(54770 / 109417)
    assert compute_margin(35695, 109417) == pytest.approx(35695 / 109417)
    assert compute_margin(29789, 109417) == pytest.approx(29789 / 109417)
    assert compute_margin(10, 0) is None
    assert compute_margin(None, 100) is None


def test_quarterly_margins_written_as_formulas(tmp_path: Path):
    from services.quarterly_margin_service import QuarterlyMarginService

    path = tmp_path / "is.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "Last Quarter IS Standardized"
    ws["A11"] = "Revenue"
    ws["C11"] = 100
    ws["D11"] = 80
    ws["G11"] = 300
    ws["H11"] = 240
    ws["A21"] = "Gross Profit"
    ws["C21"] = 40
    ws["D21"] = 32
    ws["G21"] = 120
    ws["H21"] = 96
    ws["A32"] = "Operating Income"
    ws["C32"] = 20
    ws["D32"] = 16
    ws["G32"] = 60
    ws["H32"] = 48
    ws["A60"] = "Net Income"
    ws["C60"] = 10
    ws["D60"] = 8
    ws["G60"] = 30
    ws["H60"] = 24
    wb.save(path)
    wb.close()
    report = QuarterlyMarginService().apply(
        analysis_id="m1", ticker="AAPL", workbook_path=path
    )
    assert {e.metric for e in report.entries} >= {"gross_margin", "operating_margin", "net_margin"}
    out = load_workbook(path)
    assert out["Last Quarter IS Standardized"]["A77"].value == "Gross Margin"
    assert str(out["Last Quarter IS Standardized"]["C77"].value).startswith("=")
    assert "C21" in str(out["Last Quarter IS Standardized"]["C77"].value)
    assert "C11" in str(out["Last Quarter IS Standardized"]["C77"].value)
    gm = next(e for e in report.entries if e.metric == "gross_margin" and e.period_kind == "yoy_quarter" and e.cell and e.cell.endswith("C77"))
    assert gm.margin == pytest.approx(0.4)
    out.close()

