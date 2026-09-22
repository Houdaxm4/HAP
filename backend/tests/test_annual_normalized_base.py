"""Normalized earnings-power diagnostic — analysis only, no valuation writes."""

from __future__ import annotations

from pathlib import Path
from types import SimpleNamespace

import pytest
from openpyxl import load_workbook

from models.annual_update import (
    KEEP_REPORTED_BASE,
    USE_NORMALIZED_BASE,
    INSUFFICIENT_BASE_EVIDENCE,
)
from services.annual_growth_analysis_service import AnnualGrowthAnalysisService
from services.annual_judgment_service import AnnualJudgmentService
from services.annual_normalized_base_service import AnnualNormalizedBaseService
from tests.test_annual_growth_analysis import _geo, _valuation_workbook


def _pairs(xs: list[float], start: int = 2016) -> list[tuple[str, float]]:
    return [(f"FY {start + i}", float(v)) for i, v in enumerate(xs)]


def _ev(
    *,
    ni: list[float],
    da: list[float],
    capex: list[float],
    revenue: list[float],
    operating_income: list[float],
) -> SimpleNamespace:
    n = len(ni)
    fy = [f"FY {2016 + i}" for i in range(n)]
    oe = [(fy[i], ni[i] + da[i] + capex[i]) for i in range(n)]
    return SimpleNamespace(
        fy_labels=fy,
        ni=_pairs(ni),
        da=_pairs(da),
        capex=_pairs(capex),
        oe=oe,
        revenue=_pairs(revenue),
        operating_income=_pairs(operating_income),
        ev_b9_formula="='Final Metrics'!$L$10*(1+$C$6)^B8",
        provenance=["unit-test-evidence"],
    )


def _analyze(**kwargs) -> object:
    return AnnualNormalizedBaseService().analyze_evidence(_ev(**kwargs))


def _snapshot_ev(path: Path) -> dict:
    wb = load_workbook(path, data_only=False)
    try:
        ev = wb["Enterprise Value"]
        return {
            "B6": ev["B6"].value,
            "C6": ev["C6"].value,
            "B9": ev["B9"].value,
            "B20": ev["B20"].value,
            "B27": ev["B27"].value,
        }
    finally:
        wb.close()


def test_reported_oe_base_is_never_overwritten(tmp_path: Path):
    path = tmp_path / "wb.xlsx"
    _valuation_workbook(path)
    before = _snapshot_ev(path)
    rec = AnnualNormalizedBaseService().analyze_path(path)
    after = _snapshot_ev(path)
    assert after == before
    assert rec.writes_to_workbook is False
    assert rec.implementation_status == "DIAGNOSTIC_ONLY"
    assert rec.reported_current_base is not None


def test_distorted_base_does_not_automatically_use_normalized_base():
    ni = [62.0] * 10
    da = [18.0] * 10
    capex = [-26.0] * 9 + [-80.0]
    rec = _analyze(
        ni=ni,
        da=da,
        capex=capex,
        revenue=[900.0 + 20 * i for i in range(10)],
        operating_income=[80.0] * 10,
    )
    assert "NEAR_ZERO_BASE" in rec.distortion_type or "CAPEX_SPIKE" in rec.distortion_type
    assert rec.decision == INSUFFICIENT_BASE_EVIDENCE
    assert rec.selected_normalized_base is None
    assert rec.writes_to_workbook is False


def test_capex_spike_does_not_imply_temporary_capex():
    ni = [62.0] * 10
    da = [20.0] * 10
    capex = [-25.0] * 9 + [-95.0]
    rec = _analyze(
        ni=ni,
        da=da,
        capex=capex,
        revenue=[1000.0] * 10,
        operating_income=[85.0] * 10,
    )
    assert "CAPEX_SPIKE" in rec.distortion_type
    capex_methods = [
        c for c in rec.candidate_bases if c.assumes_capex_mean_reversion and c.result is not None
    ]
    assert len(capex_methods) >= 2
    results = [c.result for c in capex_methods]
    # Mean-reversion methods may cluster; that still does not prove CapEx is temporary.
    assert rec.decision == INSUFFICIENT_BASE_EVIDENCE
    assert rec.selected_normalized_base is None
    assert rec.selection_method and "temporary" in rec.selection_method
    assert max(results) != min(results) or rec.decision != USE_NORMALIZED_BASE


def test_candidate_methods_remain_separate():
    rec = _analyze(
        ni=_geo(36.0, 0.04, 10),
        da=[15.0 + i * 0.5 for i in range(10)],
        capex=[-12.0 - i * 0.4 for i in range(10)],
        revenue=_geo(800.0, 0.037, 10),
        operating_income=_geo(60.0, 0.04, 10),
    )
    methods = [c.method for c in rec.candidate_bases]
    assert "REPORTED_CURRENT_BASE" in methods
    assert "OE_MARGIN_TIMES_CURRENT_REVENUE" in methods
    assert "CAPEX_TO_REVENUE" in methods
    assert "TEN_YEAR_AVERAGE_OE_DOLLARS" in methods
    assert len(set(methods)) == len(methods)
    blended = [c for c in rec.candidate_bases if "average of methods" in c.method.lower()]
    assert blended == []
    if rec.decision != USE_NORMALIZED_BASE:
        assert rec.selected_normalized_base is None


def test_conflicting_candidate_methods_yield_insufficient_evidence():
    # NI collapses; operating income does not. Earnings-power methods disagree.
    ni = [40.0] * 9 + [5.0]
    da = [15.0] * 10
    capex = [-12.0] * 10
    rec = _analyze(
        ni=ni,
        da=da,
        capex=capex,
        revenue=[400.0] * 10,
        operating_income=[80.0 + i for i in range(10)],
    )
    results = [c.result for c in rec.candidate_bases if c.method != "REPORTED_CURRENT_BASE" and c.result is not None]
    assert rec.candidate_base_low is not None and rec.candidate_base_high is not None
    assert rec.candidate_base_high > rec.candidate_base_low * 1.3
    assert rec.decision == INSUFFICIENT_BASE_EVIDENCE
    assert rec.selected_normalized_base is None
    # Do not average the conflict into a substitute.
    avg = sum(results) / len(results)
    assert rec.selected_normalized_base != pytest.approx(avg)


def test_corroborating_earnings_power_methods_can_support_normalization():
    # One-time NI and OI collapse; CapEx regime unchanged. OE ≈ NI + small DA-CapEx gap.
    ni = [40.0] * 9 + [6.0]
    da = [12.0] * 10
    capex = [-11.0] * 10
    oi = [48.0] * 9 + [7.0]
    rec = _analyze(
        ni=ni,
        da=da,
        capex=capex,
        revenue=[400.0] * 10,
        operating_income=oi,
    )
    assert "ONE_TIME_EARNINGS_COLLAPSE" in rec.distortion_type
    assert "CAPEX_SPIKE" not in rec.distortion_type
    assert rec.decision == USE_NORMALIZED_BASE
    assert rec.selected_normalized_base is not None
    assert rec.selection_method == "corroborating_earnings_power_methods"
    assert rec.writes_to_workbook is False
    assert rec.implementation_status == "DIAGNOSTIC_ONLY"
    assert rec.selection_method != "TEN_YEAR_AVERAGE_OE_DOLLARS"


def test_current_scale_methods_use_current_revenue_not_historical_dollars():
    ni = [10.0 * (1.15 ** i) for i in range(10)]
    da = [4.0 * (1.12 ** i) for i in range(10)]
    capex = [-3.5 * (1.12 ** i) for i in range(10)]
    rev = [100.0 * (1.15 ** i) for i in range(10)]
    oi = [16.0 * (1.15 ** i) for i in range(10)]
    rec = _analyze(ni=ni, da=da, capex=capex, revenue=rev, operating_income=oi)
    by_method = {c.method: c for c in rec.candidate_bases}
    margin = by_method["OE_MARGIN_TIMES_CURRENT_REVENUE"]
    capex_rev = by_method["CAPEX_TO_REVENUE"]
    ten = by_method["TEN_YEAR_AVERAGE_OE_DOLLARS"]
    assert margin.uses_current_scale is True
    assert capex_rev.uses_current_scale is True
    assert ten.uses_current_scale is False
    assert "current revenue" in margin.calculation.lower()
    assert ten.result is not None and margin.result is not None
    # Growing firm: current-scale margin reconstruction is not the 10y dollar average.
    assert margin.result != pytest.approx(ten.result, rel=0.02)


def test_base_analysis_and_growth_analysis_remain_independent(tmp_path: Path):
    path = tmp_path / "wb.xlsx"
    ni = [62.0] * 10
    da = [20.0] * 10
    capex = [-25.0] * 9 + [-95.0]
    _valuation_workbook(
        path,
        ni=ni,
        da=da,
        capex=capex,
        oi=[85.0] * 10,
        revenue=[1000.0] * 10,
        b6_total=-0.95,
        c6_annual=-0.28,
    )
    wb = load_workbook(path, data_only=True)
    wf = load_workbook(path, data_only=False)
    try:
        growth = AnnualGrowthAnalysisService().analyze_workbooks(wb, wf)
        base = AnnualNormalizedBaseService().analyze_workbooks(wb, wf)
    finally:
        wb.close()
        wf.close()
    assert base.decision == INSUFFICIENT_BASE_EVIDENCE
    # Distorted reported OE does not auto-select a growth rate, and a growth
    # INSUFFICIENT/ADJUST decision does not select a normalized dollar base.
    assert base.selected_normalized_base is None
    assert growth["oe"].decision in {"KEEP_EXISTING", "ADJUST", "INSUFFICIENT_EVIDENCE"}
    if base.decision == USE_NORMALIZED_BASE:
        raise AssertionError("capex-spike book must not auto-normalize the base")
    if growth["oe"].decision == "ADJUST":
        assert base.decision != USE_NORMALIZED_BASE or base.selection_method != growth["oe"].selection_method


def test_original_valuation_formulas_remain_untouched_when_base_is_use(tmp_path: Path):
    path = tmp_path / "wb.xlsx"
    ni = [40.0] * 9 + [6.0]
    da = [12.0] * 10
    capex = [-11.0] * 10
    oi = [48.0] * 9 + [7.0]
    _valuation_workbook(
        path,
        ni=ni,
        da=da,
        capex=capex,
        oi=oi,
        revenue=[400.0] * 10,
        c6_annual=0.04,
        b6_total=0.4,
    )
    rec = AnnualNormalizedBaseService().analyze_path(path)
    assert rec.decision == USE_NORMALIZED_BASE
    before = _snapshot_ev(path)
    _, judge = AnnualJudgmentService().apply(
        analysis_id="a1",
        ticker="ZZ",
        workbook_path=path,
        context={
            "oe_base_analysis": rec,
            "book_value_growth": 0.03,
            "owner_earnings_growth": 0.04,
            "eps_growth": 0.04,
            "bv_classification": "REASONABLE",
            "oe_classification": "REASONABLE",
            "eps_classification": "REASONABLE",
        },
    )
    after = _snapshot_ev(path)
    assert after == before
    assert after["B9"] == "='Final Metrics'!$L$10*(1+$C$6)^B8"
    assert judge.oe_base_analysis is not None
    assert judge.oe_base_analysis.writes_to_workbook is False


def test_no_hard_coded_normalized_oe_fallback():
    src = (Path(__file__).resolve().parents[1] / "services" / "annual_normalized_base_service.py").read_text(
        encoding="utf-8"
    )
    lowered = src.lower()
    for token in ("jbss", "etd", "csco", "mzti", "56.5", "35.2"):
        assert token not in lowered
    assert "fallback" not in lowered
    assert "normalized_oe = " not in lowered
    assert "return 0.08" not in lowered


def test_clean_history_keeps_reported_base():
    rec = _analyze(
        ni=_geo(36.0, 0.04, 10),
        da=[15.0 + i * 0.5 for i in range(10)],
        capex=[-12.0 - i * 0.4 for i in range(10)],
        revenue=_geo(800.0, 0.037, 10),
        operating_income=_geo(60.0, 0.04, 10),
    )
    assert rec.decision == KEEP_REPORTED_BASE
    assert rec.reported_base_status in {"usable", "potentially_distorted"}
    assert rec.selected_normalized_base is None
    assert rec.writes_to_workbook is False
