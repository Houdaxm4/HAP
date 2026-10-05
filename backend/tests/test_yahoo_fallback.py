"""Yahoo is a last resort: SEC first, proven year alignment, indicative label, never over a supplied value."""

import json

from openpyxl import Workbook, load_workbook

from research.http import GuardedFetcher
from research.yahoo_fundamentals import YahooFallback, yahoo_fallback_enabled
from services.annual_statement_validation_service import AnnualStatementValidationService
from services.new_company_statement_validation_service import NewCompanyStatementValidationService
from services.report_flags import collect_flags
from tests.ledger_util import entry, notes_text


def payload(**series):
    return json.dumps({"timeseries": {"result": [
        {"meta": {"type": [kind]}, kind: [{"asOfDate": f"{year}-12-31", "reportedValue": {"raw": value * 1e6}} for year, value in rows.items()]}
        for kind, rows in series.items()
    ], "error": None}})


class FakeFetcher(GuardedFetcher):
    def __init__(self, text=None, boom=False):
        super().__init__()
        self.text, self.boom, self.calls = text, boom, 0

    def get(self, url, **kw):
        self.calls += 1
        if self.boom:
            raise RuntimeError("network down")
        return url, self.text


YAHOO = payload(annualTotalRevenue={2024: 800.0, 2025: 900.0}, annualNetIncome={2024: 80.0, 2025: 90.0},
                annualTotalAssets={2024: 5000.0, 2025: 5500.0}, annualOperatingIncome={2025: 123.0})


def test_value_requires_a_trusted_concept_and_proven_alignment():
    y = YahooFallback(FakeFetcher(YAHOO))
    anchors = {"revenue": 900.0, "net_income": 90.0}
    value, source = y.value("T", "total_assets", "FY2025", anchors)
    assert value == 5500.0 and "indicative" in source and "2025-12-31" in source
    assert y.value("T", "operating_income", "FY2025", anchors) is None          # definition differs from the workbook: excluded
    assert y.value("T", "total_assets", "FY2025", {"revenue": 700.0}) is None     # revenue disagrees: the year does not line up
    assert y.value("T", "total_assets", "FY2025", {}) is None                     # nothing to prove alignment with
    assert y.value("T", "total_assets", "FY2023", anchors) is None                # Yahoo has no such year


def test_an_outage_means_no_fallback_not_an_error():
    y = YahooFallback(FakeFetcher(boom=True))
    assert y.value("T", "total_assets", "FY2025", {"revenue": 900.0}) is None and "T" in y.errors


def build(path):
    wb = Workbook()
    ws = wb.active
    ws.title = "Income - GAAP"
    ws.cell(7, 3, "FY2025")
    ws.cell(11, 1, "Revenue")
    ws.cell(11, 3, 900.0)
    ws.cell(13, 1, "Net Income")
    ws.cell(13, 3, 90.0)
    bs = wb.create_sheet("Balance Sheet - Standardized")
    bs.cell(7, 3, "FY2025")
    bs.cell(9, 1, "Total Assets")
    wb.create_sheet("Cash Flow - Standardized")
    wb.create_sheet("Final Metrics")["A1"] = "='Balance Sheet - Standardized'!C9"
    wb.save(path)


def test_new_company_check_uses_yahoo_only_when_sec_has_nothing(tmp_path):
    path = tmp_path / "w.xlsx"
    build(path)
    fetcher = FakeFetcher(YAHOO)
    report = NewCompanyStatementValidationService().validate(
        analysis_id="a", ticker="T", workbook_path=path, fiscal_years=["FY2025"], yahoo_fallback=YahooFallback(fetcher))
    cell = load_workbook(path)["Balance Sheet - Standardized"]["C9"]
    wbk = load_workbook(path)
    logged = entry(wbk, "Balance Sheet - Standardized", "C9", "Filled from filing")
    assert cell.value == 5500.0 and cell.comment is None and logged and "Yahoo Finance" in logged["source"]
    assert "Yahoo Finance (unofficial, indicative)" in notes_text(wbk["Balance Sheet - Standardized"])
    assert [(f.concept, f.source.startswith("Yahoo")) for f in report.filled_missing] == [("total_assets", True)]
    # SEC wins when it has the figure: Yahoo is never consulted for that cell
    path2 = tmp_path / "w2.xlsx"
    build(path2)
    sec_first = NewCompanyStatementValidationService().validate(
        analysis_id="a", ticker="T", workbook_path=path2, fiscal_years=["FY2025"],
        filing_overrides={"FY2025": {"total_assets": 4999.0}}, yahoo_fallback=YahooFallback(FakeFetcher(YAHOO)))
    assert load_workbook(path2)["Balance Sheet - Standardized"]["C9"].value == 4999.0
    assert sec_first.filled_missing[0].source == "filing_override"


def test_a_misaligned_year_is_never_filled(tmp_path):
    path = tmp_path / "w.xlsx"
    build(path)
    wrong_year = payload(annualTotalRevenue={2025: 400.0}, annualNetIncome={2025: 40.0}, annualTotalAssets={2025: 5500.0})
    NewCompanyStatementValidationService().validate(
        analysis_id="a", ticker="T", workbook_path=path, fiscal_years=["FY2025"], yahoo_fallback=YahooFallback(FakeFetcher(wrong_year)))
    assert load_workbook(path)["Balance Sheet - Standardized"]["C9"].value is None


def test_annual_check_and_flags(tmp_path):
    path = tmp_path / "w.xlsx"
    build(path)
    wb = load_workbook(path)
    wb["Balance Sheet - Standardized"].cell(7, 3).value = "FY2025"
    wb.save(path)
    report = AnnualStatementValidationService().validate(
        analysis_id="a", ticker="T", workbook_path=path, fiscal_year="FY2025", yahoo_fallback=YahooFallback(FakeFetcher(YAHOO)))
    item = next(i for i in report.items if i.concept == "total_assets")
    assert item.status == "FILLED_FROM_YAHOO" and report.filled >= 1
    (tmp_path / "annual_statement_validation_report.json").write_text(report.model_dump_json())
    flags = collect_flags(tmp_path)["flags"]
    assert any(f["source"] == "Yahoo Finance (indicative)" for f in flags["filled"])
    assert any("came from Yahoo Finance" in f["title"] for f in flags["notes"])


def test_switch_off(monkeypatch):
    monkeypatch.setenv("HAP_YAHOO_FALLBACK", "0")
    assert not yahoo_fallback_enabled()
