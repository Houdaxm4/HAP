"""Certification-focused tests: concept matching, projection, deliverables."""

from __future__ import annotations

from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook

from models.quarterly_presentation import QuarterlyStatementKind, SecLineItem
from models.quarterly_update import QuarterlyProjectionReport, QuarterlyReviewReport
from services.accounting_concept_matcher import (
    interpret_workbook_label,
    match_sec_item_to_concept,
    resolve_workbook_gap,
)
from services.quarterly_deliverables_service import (
    QuarterlyDeliverablesService,
    deliverable_stems,
)
from services.quarterly_projection_service import (
    QuarterlyProjectionService,
    annualization_factor,
)


def test_revenue_synonyms_same_concept():
    for lab in ("Revenue", "Net Sales", "Total Net Revenues", "Sales"):
        concepts = interpret_workbook_label(lab, QuarterlyStatementKind.INCOME)
        assert "revenue" in concepts


def test_statement_context_blocks_bs_concept_on_income():
    concepts = interpret_workbook_label("Total Assets", QuarterlyStatementKind.INCOME)
    assert "total_assets" not in concepts


def test_xbrl_canonical_match():
    items = [
        SecLineItem(
            statement="quarterly_income_statement",
            label="Net sales",
            xbrl_concept="RevenueFromContractWithCustomerExcludingAssessedTax",
            value=100.0,
            duration_kind="standalone_quarter",
        )
    ]
    m = match_sec_item_to_concept("revenue", items, kind=QuarterlyStatementKind.INCOME)
    assert m.decision == "MATCHED"
    assert m.match_method == "canonical_xbrl"
    assert m.value == 100.0


def test_ambiguous_concepts_review_required():
    items = [
        SecLineItem(
            statement="quarterly_balance_sheet",
            label="Cash",
            xbrl_concept="CashAndCashEquivalentsAtCarryingValue",
            value=10.0,
            duration_kind="instant",
        ),
        SecLineItem(
            statement="quarterly_balance_sheet",
            label="Cash and cash equivalents",
            xbrl_concept="CashAndCashEquivalentsAtCarryingValue",
            value=99.0,
            duration_kind="instant",
        ),
    ]
    # Two XBRL hits with different values → REVIEW
    m = match_sec_item_to_concept("cash", items, kind=QuarterlyStatementKind.BALANCE_SHEET)
    assert m.decision == "REVIEW_REQUIRED"
    assert m.alternatives


def test_no_blind_literal_only_similarity():
    items = [
        SecLineItem(
            statement="quarterly_income_statement",
            label="Research and development",
            xbrl_concept="ResearchAndDevelopmentExpense",
            value=5.0,
            duration_kind="standalone_quarter",
        )
    ]
    # Workbook "Revenue" must not match R&D just because both are IS lines
    m = resolve_workbook_gap("Revenue", QuarterlyStatementKind.INCOME, items)
    assert m.decision in {"SOURCE_MISSING", "REVIEW_REQUIRED", "BLOCKED"}
    assert m.value is None


def test_goodwill_source_missing_when_sec_absent():
    items = [
        SecLineItem(
            statement="quarterly_balance_sheet",
            label="Total assets",
            xbrl_concept="Assets",
            value=1000.0,
            duration_kind="instant",
        )
    ]
    m = resolve_workbook_gap("Goodwill", QuarterlyStatementKind.BALANCE_SHEET, items)
    assert m.decision == "SOURCE_MISSING"


def test_annualization_factors():
    assert annualization_factor(1) is None
    assert annualization_factor(2) == 2.0
    assert annualization_factor(3) == pytest.approx(4 / 3)


def test_eval_ref_wacc_inputs_over_100(tmp_path: Path):
    path = tmp_path / "wacc.xlsx"
    wb = Workbook()
    wb.active.title = "Final Metrics"
    inputs = wb.create_sheet("Inputs")
    inputs["L54"] = 8.5  # percent
    wb["Final Metrics"]["L7"] = "=Inputs!L54/100"
    wb.save(path)
    wb.close()
    wb2 = load_workbook(path)
    try:
        got = QuarterlyProjectionService._eval_ref(wb2, wb2["Final Metrics"]["L7"].value)
        assert got == pytest.approx(0.085)
    finally:
        wb2.close()


def test_q1_projection_not_applicable(tmp_path: Path):
    path = tmp_path / "wb.xlsx"
    wb = Workbook()
    wb.active.title = "IC & NOPAT & ROIC "
    is_ = wb.create_sheet("Last Quarter IS Standardized")
    is_["C5"] = "2026 Q1"
    wb.create_sheet("Last Quarter BS Standardized")
    wb.create_sheet("Last Quarter CF Standardized")
    wb.save(path)
    wb.close()
    report = QuarterlyProjectionService().apply(
        analysis_id="p1", ticker="AAPL", workbook_path=path, fiscal_quarter=1
    )
    assert report.status == "NOT_APPLICABLE"


def _mini_proj_wb(path: Path, *, q: int = 3) -> Path:
    wb = Workbook()
    ic = wb.active
    ic.title = "IC & NOPAT & ROIC "
    for r, lab in [
        (3, "Operating Assets"),
        (4, "Operating Liabilities"),
        (5, "Capitalized Leases"),
        (6, "Capitalized R&D"),
        (7, "Invested Capital"),
        (11, "Revenue from Operations"),
        (13, "Operating Income"),
        (14, "Lease Expense"),
        (15, "Lease Depreciation"),
        (16, "R&D Expense"),
        (17, "R&D Amortization"),
        (19, "Operating Taxes"),
        (20, "NOPAT"),
        (23, "ROIC in"),
    ]:
        ic.cell(r, 1).value = lab
    # Last FY L column values
    ic["L5"] = 100.0
    ic["L6"] = 200.0
    ic["L13"] = 1000.0
    ic["L14"] = 50.0
    ic["L15"] = 10.0
    ic["L16"] = 80.0
    ic["L17"] = 40.0
    ic["L19"] = 200.0
    ic["L20"] = 880.0
    ic["L7"] = 500.0
    ic["L23"] = 0.2

    bs = wb.create_sheet("Last Quarter BS Standardized")
    # OA rows
    for r, v in [(11, 50), (14, 20), (20, 30), (27, 0), (37, 100), (48, 40), (51, 10)]:
        bs.cell(r, 1).value = f"line {r}"
        bs.cell(r, 3).value = v
    # OL rows
    for r, v in [(65, 40), (68, 10), (78, 5), (98, 25)]:
        bs.cell(r, 1).value = f"ol {r}"
        bs.cell(r, 3).value = v
    bs["A61"] = "Total Assets"
    bs["C61"] = 1000.0
    bs["C5"] = f"2026 Q{q}"

    is_ = wb.create_sheet("Last Quarter IS Standardized")
    is_["A11"] = "Revenue"
    is_["C11"] = 90
    is_["G11"] = 270 if q == 3 else 180  # YTD
    is_["A32"] = "Operating Income"
    is_["C32"] = 30
    is_["G32"] = 90 if q == 3 else 60
    is_["C5"] = f"2026 Q{q}"

    cf = wb.create_sheet("Last Quarter CF Standardized")
    cf["A27"] = "Cash from Operating Activities"
    cf["C27"] = 120.0

    fm = wb.create_sheet("Final Metrics")
    fm["A7"] = "WACC  Bloomberg"
    fm["L7"] = 0.10

    wb.create_sheet("Inputs")
    wb["Inputs"]["B63"] = 300.0
    wb["Inputs"]["B67"] = 150.0
    wb["Inputs"]["B69"] = -0.2
    wb["Inputs"]["B72"] = 0.8

    wb.save(path)
    wb.close()
    return path


def test_q2_projection_factor_and_bridge(tmp_path: Path):
    path = _mini_proj_wb(tmp_path / "q2.xlsx", q=2)
    report = QuarterlyProjectionService().apply(
        analysis_id="p2", ticker="AAPL", workbook_path=path, fiscal_quarter=2
    )
    assert report.status == "ok"
    assert report.annualization_factor == 2.0
    assert report.ytd_revenue == 180
    assert report.projected_revenue == 360
    assert report.projected_operating_income == 120
    # taxes = 200 * (120/1000) = 24
    assert report.projected_operating_taxes == pytest.approx(24.0)
    # OA = 50+20+30+0+100+40+10 = 250; OL = 40+10+5+25 = 80
    # IC = 250-80+100+200 = 470
    assert report.operating_assets == pytest.approx(250)
    assert report.operating_liabilities == pytest.approx(80)
    assert report.projected_invested_capital == pytest.approx(470)
    # NOPAT = 120+50-10+80-40-24 = 176
    assert report.projected_nopat == pytest.approx(176)
    assert report.projected_roic == pytest.approx(176 / 470)
    assert report.prior_fy_wacc == pytest.approx(0.10)
    assert report.projected_roic_wacc == pytest.approx(176 / 470 - 0.10)
    assert report.projected_cfo == pytest.approx(240)
    assert report.projected_roce == pytest.approx(240 / 1000)

    out = load_workbook(path)
    ic = out["IC & NOPAT & ROIC "]
    assert ic["M23"].value == "=M20/M7"
    assert ic["M24"].value == pytest.approx(0.10)
    assert ic["M25"].value == "=M23-M24"
    assert "C27*2" in str(ic["N4"].value)
    out.close()


def test_q3_projection_factor_four_thirds(tmp_path: Path):
    path = _mini_proj_wb(tmp_path / "q3.xlsx", q=3)
    report = QuarterlyProjectionService().apply(
        analysis_id="p3", ticker="AAPL", workbook_path=path, fiscal_quarter=3
    )
    assert report.annualization_factor == pytest.approx(4 / 3)
    assert report.projected_revenue == pytest.approx(270 * 4 / 3)
    assert report.projected_operating_income == pytest.approx(90 * 4 / 3)
    assert "4/3" in report.formulas_written.get("N4", "")


def test_zero_prior_oi_blocks_tax(tmp_path: Path):
    path = _mini_proj_wb(tmp_path / "z.xlsx", q=2)
    wb = load_workbook(path)
    wb["IC & NOPAT & ROIC "]["L13"] = 0
    wb.save(path)
    wb.close()
    report = QuarterlyProjectionService().apply(
        analysis_id="pz", ticker="AAPL", workbook_path=path, fiscal_quarter=2
    )
    assert report.status == "REVIEW_REQUIRED"
    assert report.projected_operating_taxes is None


def test_deliverable_filenames():
    x, w = deliverable_stems(2026, 2, "aapl")
    assert x == "2026 Q2 AAPL FA.xlsx"
    assert w == "2026 Q2 AAPL Quarterly Update.docx"


def test_word_q1_excludes_projection_q3_includes(tmp_path: Path):
    path = _mini_proj_wb(tmp_path / "d.xlsx", q=3)
    svc = QuarterlyDeliverablesService()
    proj = QuarterlyProjectionReport(
        analysis_id="d",
        ticker="AAPL",
        status="ok",
        fiscal_quarter=3,
        fiscal_year=2026,
        next_fiscal_year=2027,
        projected_roic_wacc=0.05,
        projected_roce=0.15,
        prior_fy_roic=0.2,
        prior_fy_wacc=0.1,
    )
    out = svc.produce(
        analysis_id="d",
        ticker="AAPL",
        completed_workbook_path=path,
        output_dir=tmp_path / "out",
        fiscal_year=2026,
        fiscal_quarter=3,
        projection=proj,
        review=QuarterlyReviewReport(analysis_id="d", ticker="AAPL"),
    )
    assert out.excel_filename == "2026 Q3 AAPL FA.xlsx"
    assert out.word_filename == "2026 Q3 AAPL Quarterly Update.docx"
    assert (tmp_path / "out" / out.excel_filename).exists()
    assert (tmp_path / "out" / out.word_filename).exists()
    from docx import Document

    doc = Document(tmp_path / "out" / out.word_filename)
    texts = [p.text for p in doc.paragraphs]
    assert any("2027 Projection" in t for t in texts)
    assert any("Financial Highlights" in t for t in texts)
    # Projection heading appears before Financial Highlights
    i_proj = next(i for i, t in enumerate(texts) if "2027 Projection" in t)
    i_fh = next(i for i, t in enumerate(texts) if "Financial Highlights" in t)
    assert i_proj < i_fh

    # Q1 excludes projection
    out1 = svc.produce(
        analysis_id="d1",
        ticker="AAPL",
        completed_workbook_path=path,
        output_dir=tmp_path / "out1",
        fiscal_year=2026,
        fiscal_quarter=1,
        projection=QuarterlyProjectionReport(
            analysis_id="d1", ticker="AAPL", status="NOT_APPLICABLE", fiscal_quarter=1
        ),
        review=None,
    )
    doc1 = Document(tmp_path / "out1" / out1.word_filename)
    texts1 = [p.text for p in doc1.paragraphs]
    assert not any("Projection" in t and "2027" in t for t in texts1)
    assert any("Financial Highlights" in t for t in texts1)
