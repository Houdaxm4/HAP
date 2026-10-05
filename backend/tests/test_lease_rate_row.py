"""Leases tab: unrealistic long-term rates are replaced with a justified, flagged number."""

from __future__ import annotations

import shutil
from pathlib import Path
from types import SimpleNamespace

import pytest
from openpyxl import Workbook, load_workbook

from services.lease_rate_row_service import (
    IG_RANGE,
    NON_IG_RANGE,
    LeaseRateRowService,
    business_estimate,
    choose_rate,
    credit_profile,
    is_realistic,
)

IDCC = Path(__file__).resolve().parents[1] / "storage" / "certifications" / "new-company-idcc-edd8ea8f" / "completed_workbook.xlsx"


def test_realistic_bounds():
    assert is_realistic(0.05) and is_realistic(0.10) and is_realistic(0.01)
    assert not is_realistic(0.1001)
    assert not is_realistic(0.004)
    assert not is_realistic(None) and not is_realistic("#DIV/0!") and not is_realistic(True)


def test_choose_rate_prefers_disclosed_then_neighbors_then_business():
    ordered = [("FY2020", 0.06), ("FY2021", 0.13), ("FY2022", 0.07), ("FY2023", None)]
    profile = credit_profile({"debt": 100.0, "interest": 2.0, "ebit": 80.0})
    rate, method, _ = choose_rate("FY2021", ordered, {"FY2021": 0.045}, profile)
    assert (rate, method) == (0.045, "company_disclosed")
    rate, method, _ = choose_rate("FY2021", ordered, {}, profile)
    assert method == "nearest_realistic_years" and 0.06 <= rate <= 0.07
    rate, method, _ = choose_rate("FY2023", [("FY2023", None)], {}, profile)
    assert method == "business_estimate"


def test_business_ranges_follow_the_credit_profile():
    strong = business_estimate(credit_profile({"debt": 100.0, "interest": 3.0, "ebit": 200.0}))[0]
    middling = business_estimate(credit_profile({"debt": 400.0, "interest": 20.0, "ebit": 100.0}))[0]
    weak = business_estimate(credit_profile({"debt": 900.0, "interest": 60.0, "ebit": 100.0}))[0]
    nodebt = business_estimate(credit_profile({"debt": 0.0, "interest": 0.0, "ebit": 100.0}))[0]
    unknown = business_estimate(credit_profile({}))[0]
    assert IG_RANGE[0] <= strong <= IG_RANGE[1]
    assert IG_RANGE[0] <= nodebt <= IG_RANGE[1]
    assert NON_IG_RANGE[0] <= middling <= NON_IG_RANGE[1]
    assert NON_IG_RANGE[0] <= weak <= NON_IG_RANGE[1]
    assert NON_IG_RANGE[0] <= unknown <= NON_IG_RANGE[1]


@pytest.mark.skipif(not IDCC.exists(), reason="IDCC certification workbook not present")
def test_idcc_years_above_ten_percent_are_fixed_and_flagged(tmp_path: Path):
    path = tmp_path / "wb.xlsx"
    shutil.copy2(IDCC, path)
    before = load_workbook(path, data_only=True)["Leases"]
    bad_before = [c.coordinate for c in before[18][1:] if isinstance(c.value, float) and not is_realistic(c.value)]
    assert bad_before, "fixture should contain at least one unrealistic year"
    report = LeaseRateRowService().apply(workbook_path=path)
    assert {fix.cell.split("!")[1] for fix in report.fixes} == set(bad_before)
    after_formulas = load_workbook(path)["Leases"]
    for fix in report.fixes:
        cell = after_formulas[fix.cell.split("!")[1]]
        assert is_realistic(cell.value) and not str(cell.value).startswith("=")
        assert cell.comment is None and cell.fill.fill_type == "solid" and fix.method in {"nearest_realistic_years", "company_disclosed"}
    ledger = load_workbook(path)["HAP Adjustments"]
    assert ledger["A5"].value == "ADJ-001" and ledger["E5"].value == "Lease rate"
    assert str(ledger["G5"].value).startswith("=")  # original formula kept as text
    assert ledger["D5"].hyperlink is not None
    assert any("HAP Adjustments tab" in str(c.value) for c in after_formulas["A"])
    # untouched years keep their formulas
    untouched = [c for c in after_formulas[18][1:] if c.coordinate not in bad_before and c.value is not None]
    assert untouched and all(str(c.value).startswith("=") for c in untouched)


def test_annual_update_only_changes_the_newest_year(tmp_path: Path):
    wb = Workbook()
    ws = wb.active
    ws.title = "Leases"
    ws["A1"] = "Year"
    for i, fy in enumerate(("FY2022", "FY2023", "FY2024"), start=2):
        ws.cell(1, i).value = fy
        ws.cell(3, i).value = fy
    ws["A18"] = "Estimated Long-Term Rate"
    ws["B18"], ws["C18"], ws["D18"] = 0.20, 0.06, 0.15
    path = tmp_path / "annual.xlsx"
    wb.save(path)
    # Year headers are plain text here, so detect_year_columns reads row 1.
    report = LeaseRateRowService().apply(workbook_path=path, newest_only_fy="FY2024")
    out = load_workbook(path)["Leases"]
    assert out["B18"].value == 0.20  # earlier year copied, not touched
    assert out["C18"].value == 0.06
    assert is_realistic(out["D18"].value) and out["D18"].value != 0.15
    assert len(report.fixes) == 1


def test_lease_year_data_disclosed_rate_is_used(tmp_path: Path):
    wb = Workbook()
    ws = wb.active
    ws.title = "Leases"
    ws["A1"] = "Year"
    for i, fy in enumerate(("FY2023", "FY2024"), start=2):
        ws.cell(1, i).value = fy
    ws["A18"] = "Estimated Long-Term Rate"
    ws["B18"], ws["C18"] = 0.06, None
    path = tmp_path / "d.xlsx"
    wb.save(path)
    report = LeaseRateRowService().apply(
        workbook_path=path, lease_years=[SimpleNamespace(fiscal_year="FY2024", reported_discount_rate=4.2)]
    )
    assert report.fixes[0].method == "company_disclosed" and report.fixes[0].new == pytest.approx(0.042)


def test_nothing_to_fix_means_the_file_is_not_rewritten(tmp_path: Path):
    """Re-saving with openpyxl drops Excel's calculated values, which the valuation step needs: only save when a rate was replaced."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Leases"
    ws["A1"] = "Year"
    for i, fy in enumerate(("FY2023", "FY2024"), start=2):
        ws.cell(1, i).value = fy
    ws["A18"] = "Estimated Long-Term Rate"
    ws["B18"], ws["C18"] = 0.05, 0.06
    path = tmp_path / "ok.xlsx"
    wb.save(path)
    before = path.read_bytes()
    report = LeaseRateRowService().apply(workbook_path=path)
    assert not report.changed and path.read_bytes() == before
