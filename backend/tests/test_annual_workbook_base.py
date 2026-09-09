"""Annual base-workbook, period alignment, and anti-previous-file guard tests."""

from __future__ import annotations

import shutil
from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook

from services.annual_continuity_service import AnnualContinuityService
from services.annual_period_service import (
    AnnualPeriodAlignmentError,
    align_fiscal_periods,
    detect_year_columns,
)
from services.annual_update_runner import sha256_file
from services.annual_workbook_guard_service import assert_base_workbook_guard


def _header(ws, years: list[int]) -> None:
    ws["A7"] = "Line"
    for i, y in enumerate(years):
        ws.cell(7, 3 + i, f"FY{y}")


def _rolling_pair(tmp: Path) -> tuple[Path, Path]:
    prev = tmp / "prev.xlsx"
    tmpl = tmp / "tmpl.xlsx"
    years_prev = list(range(2016, 2026))
    years_new = list(range(2017, 2027))

    def _book(path: Path, years: list[int], *, rev: float) -> None:
        wb = Workbook()
        bs = wb.active
        bs.title = "Balance Sheet - Standardized"
        _header(bs, years)
        bs["A11"] = "Revenue proxy"
        for i, y in enumerate(years):
            bs.cell(11, 3 + i, rev + i)
        is_ = wb.create_sheet("Income - GAAP")
        for col in range(3, 3 + len(years)):
            is_.cell(7, col, f"='Balance Sheet - Standardized'!{chr(64+col)}7")
            is_.cell(11, col, f"='Balance Sheet - Standardized'!{chr(64+col)}11")
        wb.save(path)
        wb.close()

    _book(prev, years_prev, rev=100)
    _book(tmpl, years_new, rev=500)
    return prev, tmpl


def test_align_detects_single_new_fiscal_year(tmp_path: Path):
    prev, tmpl = _rolling_pair(tmp_path)
    alignment = align_fiscal_periods(tmpl, prev)
    assert alignment.new_fiscal_year == "FY2026"
    assert alignment.dropped_years == ("FY2016",)
    assert "FY2025" in alignment.overlap_years


def test_working_copy_seeds_from_template_not_previous(tmp_path: Path):
    prev, tmpl = _rolling_pair(tmp_path)
    out = tmp_path / "working.xlsx"
    shutil.copy2(tmpl, out)
    assert sha256_file(out) == sha256_file(tmpl)
    assert sha256_file(out) != sha256_file(prev)


def test_continuity_preserves_template_new_year_values(tmp_path: Path):
    prev, tmpl = _rolling_pair(tmp_path)
    out = tmp_path / "out.xlsx"
    shutil.copy2(tmpl, out)
    alignment = align_fiscal_periods(tmpl, prev)
    AnnualContinuityService().apply(
        analysis_id="a1",
        ticker="ZZ",
        template_path=tmpl,
        previous_workbook_path=prev,
        workbook_path=out,
        period_alignment=alignment,
    )
    wb_t = load_workbook(tmpl, data_only=True)
    wb_o = load_workbook(out, data_only=True)
    cols = detect_year_columns(wb_t["Balance Sheet - Standardized"], wb_t)
    new_col = cols["FY2026"]
    assert wb_o["Balance Sheet - Standardized"].cell(11, new_col).value == wb_t[
        "Balance Sheet - Standardized"
    ].cell(11, new_col).value
    overlap_col = cols["FY2025"]
    prev_cols = detect_year_columns(load_workbook(prev)["Balance Sheet - Standardized"], load_workbook(prev))
    wb_p = load_workbook(prev, data_only=True)
    assert wb_o["Balance Sheet - Standardized"].cell(11, overlap_col).value == wb_p[
        "Balance Sheet - Standardized"
    ].cell(11, prev_cols["FY2025"]).value
    wb_t.close()
    wb_o.close()
    wb_p.close()


def test_final_not_previous_when_template_differs(tmp_path: Path):
    prev, tmpl = _rolling_pair(tmp_path)
    out = tmp_path / "out.xlsx"
    shutil.copy2(tmpl, out)
    alignment = align_fiscal_periods(tmpl, prev)
    AnnualContinuityService().apply(
        analysis_id="a1",
        ticker="ZZ",
        template_path=tmpl,
        previous_workbook_path=prev,
        workbook_path=out,
        period_alignment=alignment,
    )
    guard = assert_base_workbook_guard(
        analysis_id="a1",
        ticker="ZZ",
        template_path=tmpl,
        previous_path=prev,
        final_path=out,
        alignment=alignment,
    )
    assert guard.status == "ok"
    assert sha256_file(out) != sha256_file(prev)


def test_blocked_when_no_unique_new_year(tmp_path: Path):
    prev, tmpl = _rolling_pair(tmp_path)
    same = tmp_path / "same.xlsx"
    shutil.copy2(prev, same)
    with pytest.raises(AnnualPeriodAlignmentError, match="ANNUAL_PERIOD_ALIGNMENT_BLOCKED"):
        align_fiscal_periods(same, prev)
