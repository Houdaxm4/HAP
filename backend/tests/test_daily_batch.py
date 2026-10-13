import sys
from pathlib import Path

from openpyxl import Workbook

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "scripts"))

import run_batch  # noqa: E402
from services.daily_sheet_service import DailySheet, custom_filter_status, find_latest_sheet  # noqa: E402


def _sheet(path: Path) -> Path:
    wb = Workbook()
    sp = wb.active
    sp.title = "S&P"
    sp.append([None, "No.", "Ticker", "Interest", "Current Filter Status w/o PE10 Percentile", "TBV/Price"])
    sp.append(["Woodward, Inc.Common Stock", 201, "WWD", "N", "Buy", 0.0559])
    sp.append(["Standex International Corp", 167, "SXI", "U", "Out", -0.0074])
    ftse = wb.create_sheet("FTSE")
    ftse.append([None, "No.", "Ticker", "Interest", "Current Filter Status w/o PE10 Percentile", "TBV/Price"])
    ftse.append(["Victrex plc", 191, "VCT", "N", "N/A", None])
    other = wb.create_sheet("Owned Statistics")
    other.append(["Ticker", "ROE"])
    wb.save(path)
    return path


def test_sheet_fields_for_the_email(tmp_path: Path):
    sheet = DailySheet(_sheet(tmp_path / "Summary&sector tables.xlsx"))
    f = sheet.fields("wwd")
    assert (f.classified_as, f.pe10_status, f.tbv_p_text, f.market_cap_text) == ("N", None, "5.59%", None)      # the status with PE10 is not in this sheet
    assert sheet.fields("SXI").classified_as == "U" and sheet.fields("SXI").tbv_p_text == "-0.74%"
    assert sheet.company_name("WWD") == "Woodward, Inc."
    v = sheet.fields("VCT")                       # second tab; N/A status is not a status
    assert v.classified_as == "N" and v.pe10_status is None
    missing = sheet.fields("ZZZ")
    assert missing.classified_as is None and not sheet.has("ZZZ")


def test_latest_sheet_is_the_newest_workbook_with_an_sp_tab(tmp_path: Path):
    old = _sheet(tmp_path / "Summary old.xlsx")
    Workbook().save(tmp_path / "WWD 2026 Q3 template.xlsx")           # no S&P tab
    import os
    os.utime(old, (1, 1))
    new = _sheet(tmp_path / "Summary new.xlsx")
    assert find_latest_sheet(tmp_path) == new


def _touch(folder: Path, *names: str) -> None:
    for n in names:
        Workbook().save(folder / n)


def test_files_are_matched_by_ticker_period_and_filter_name(tmp_path: Path):
    _touch(tmp_path, "WWD 2026 Q3 - Industrial Template v28.0.xlsx", "Custom_Run_Filter_2026-08-27-(20-58)-WWD.xlsx",
           "Custom_Run_Filter_2026-08-27-(20-58)-SXI.xlsx", "SXI 2026 Q4 - Industrial Template v28.0.xlsx")
    m = run_batch.match_files(tmp_path, "WWD", "new_company")
    assert not m.problems and m.workbook.name.startswith("WWD 2026 Q3") and m.crf.name.endswith("-WWD.xlsx") and m.previous is None
    assert not run_batch.match_files(tmp_path, "SXI", "new_company").problems


def test_update_uses_newer_period_as_new_and_older_as_previous(tmp_path: Path):
    _touch(tmp_path, "ETD 2025 FY.xlsx", "ETD 2026 Q1.xlsx", "ETD 2026 Q2.xlsx", "custom_run_filter_2026-10-06_ETD.xlsx")
    m = run_batch.match_files(tmp_path, "ETD", "quarterly_update")
    assert m.workbook.name == "ETD 2026 Q2.xlsx" and m.previous.name == "ETD 2026 Q1.xlsx"
    _touch(tmp_path, "ATR 2025 Q4.xlsx", "ATR 2025 FY.xlsx", "custom_run_filter_2026-10-06_ATR.xlsx")
    n = run_batch.match_files(tmp_path, "ATR", "annual_update")
    assert n.workbook.name == "ATR 2025 FY.xlsx" and n.previous.name == "ATR 2025 Q4.xlsx"       # FY sorts after Q4


def test_a_finished_result_is_accepted_as_the_previous_workbook(tmp_path: Path):
    # the real RPM folder: the new template, last period's finished result, the filter and the sheet
    _touch(tmp_path, "RPM 2027 Q1 - Industrial Template v28.2.xlsx", "2026 FY RPM FA.xlsx",
           "Custom_Run_Filter_2026-10-12-(16-47)-RPM.xlsx")
    m = run_batch.match_files(tmp_path, "RPM", "quarterly_update")
    assert not m.problems
    assert m.workbook.name.startswith("RPM 2027 Q1") and m.previous.name == "2026 FY RPM FA.xlsx"
    # a result for another ticker or the other naming style never leaks in
    _touch(tmp_path, "2026 Q2 TKR FA.xlsx")
    assert run_batch.match_files(tmp_path, "RPM", "quarterly_update").previous.name == "2026 FY RPM FA.xlsx"
    assert run_batch.match_files(tmp_path, "TKR", "quarterly_update").problems       # one TKR workbook only: an update needs two


def test_two_newest_workbooks_of_the_same_period_are_reported_not_used(tmp_path: Path):
    # the result of an earlier run of this very period must not become its own previous workbook
    _touch(tmp_path, "RPM 2027 Q1 - Industrial Template v28.2.xlsx", "2027 Q1 RPM FA.xlsx", "custom_run_filter_2026-10-12_RPM.xlsx")
    m = run_batch.match_files(tmp_path, "RPM", "quarterly_update")
    assert m.workbook is None and any("same period" in p for p in m.problems)


def test_missing_or_ambiguous_files_are_reported_not_guessed(tmp_path: Path):
    _touch(tmp_path, "TKR 2026 Q2 a.xlsx", "TKR 2026 Q1 b.xlsx")
    m = run_batch.match_files(tmp_path, "TKR", "new_company")
    assert m.workbook is None
    assert any("custom run filter" in p for p in m.problems) and any("exactly one" in p for p in m.problems)
    assert any("no workbook" in p for p in run_batch.match_files(tmp_path, "QQQ", "new_company").problems)
    assert run_batch.match_files(tmp_path, "TKR", "quarterly_update").workbook is not None      # two workbooks are enough for an update


def test_live_sheet_is_downloaded_from_the_published_link_and_falls_back_to_the_file(tmp_path: Path):
    import functools
    import http.server
    import threading

    from services.daily_sheet_service import LIVE_FILE, open_daily_sheet

    served = tmp_path / "served"
    served.mkdir()
    _sheet(served / "live.xlsx")
    (served / "notasheet").write_text("<html>sign in</html>", encoding="utf-8")
    handler = functools.partial(http.server.SimpleHTTPRequestHandler, directory=str(served))
    server = http.server.ThreadingHTTPServer(("127.0.0.1", 0), handler)
    threading.Thread(target=server.serve_forever, daemon=True).start()
    base = f"http://127.0.0.1:{server.server_address[1]}"
    try:
        work = tmp_path / "work"
        work.mkdir()
        _sheet(work / "Summary&sector tables.xlsx")
        sheet, source = open_daily_sheet(work, f"{base}/live.xlsx")
        assert "live Google Sheet" in source and sheet.has("WWD") and (work / LIVE_FILE).exists()
        # a sign-in page, a missing link or no link at all: the morning download is used
        (work / LIVE_FILE).unlink()
        for url in (f"{base}/notasheet", f"{base}/missing.xlsx", None):
            sheet, source = open_daily_sheet(work, url)
            assert "Summary&sector tables.xlsx" in source and sheet.has("WWD") and not (work / LIVE_FILE).exists()
        assert ("could not be reached" in open_daily_sheet(work, f"{base}/missing.xlsx")[1])
    finally:
        server.shutdown()


def test_status_with_pe10_comes_from_the_custom_run_filter_summary(tmp_path: Path):
    wb = Workbook()
    wb.active.title = "WWD"
    summary = wb.create_sheet("Summary")
    summary.append([None, "Ticker", "Price", "Sector", "Industry", "Current Status without PE10 Percentile", "Current Status with PE10 Percentile"])
    summary.append(["WOODWARD INC", "WWD", 346.34, "Industrial", "Controls", "Buy", "Out"])
    summary.append(["OTHER INC", "OTH", 10.0, "Industrial", "Controls", "Buy", "Buy"])
    path = tmp_path / "Custom_Run_Filter_2026-08-27-(20-58)-WWD.xlsx"
    wb.save(path)
    assert custom_filter_status(path, "wwd") == "Out"          # the "without" column says Buy; the email needs the "with" column
    assert custom_filter_status(path, "OTH") == "Buy"
    assert custom_filter_status(path, "ZZZ") is None and custom_filter_status(None, "WWD") is None
    assert custom_filter_status(tmp_path / "missing.xlsx", "WWD") is None
