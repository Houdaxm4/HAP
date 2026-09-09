"""Tests for completion / validation separation."""

from __future__ import annotations

from pathlib import Path

import pytest
from openpyxl import Workbook

from models.completion import CompletionDecision
from models.prefill_validation import PrefillValidationStatus
from services.completion_service import CompletionService
from services.excel_fill_service import ExcelFillService
from services.prefill_validation_service import PrefillValidationService, values_agree
from workbook_mapping.engine import IntentDecision, WriteIntent, WriteIntentReport


def _intent(
    *,
    intent_id: str,
    cell: str,
    value,
    decision: IntentDecision = IntentDecision.WRITE,
    sheet: str = "Income - GAAP",
    period: str = "FY2018",
) -> WriteIntent:
    return WriteIntent(
        intent_id=intent_id,
        mapping_id=intent_id,
        sheet=sheet,
        cell=cell,
        metric="Revenue",
        period=period,
        value=value,
        source="sec_companyfacts",
        write_policy_result="writable:Hybrid",
        decision=decision,
        reason="test",
        cfm_path="income_statement.revenue",
        confidence=0.9,
        existing_cell_classification="Writable Input",
    )


@pytest.fixture
def workbook_with_revenue(tmp_path: Path) -> Path:
    wb = Workbook()
    ws = wb.active
    ws.title = "Income - GAAP"
    ws["C3"] = "FY2025"
    ws["E9"] = 265595  # already correct
    ws["F9"] = None  # blank — needs fill when in-scope
    ws["G9"] = 999999  # wrong vs SEC
    ws["L9"] = None  # FY2025 blank
    ws["M9"] = "=E9*2"  # formula
    path = tmp_path / "upload.xlsx"
    wb.save(path)
    wb.close()
    return path


def test_populated_correct_cell_already_present_no_write(workbook_with_revenue: Path, tmp_path: Path):
    # In-scope for annual_update target year with populated value
    intents = WriteIntentReport(
        analysis_id="c1",
        ticker="AAPL",
        intents=[
            _intent(
                intent_id="is.revenue:L9",
                cell="L9",
                value=391035.0,
                period="FY2025",
            ),
        ],
        write_count=1,
    )
    # Prefill L9 so ALREADY_PRESENT
    from openpyxl import load_workbook

    wb = load_workbook(workbook_with_revenue)
    wb["Income - GAAP"]["L9"] = 391035
    wb.save(workbook_with_revenue)
    wb.close()

    completion, fill = CompletionService().plan(
        analysis_id="c1",
        ticker="AAPL",
        analysis_type="annual_update",
        source_workbook_path=workbook_with_revenue,
        intent_report=intents,
        target_fiscal_year="FY2025",
    )
    assert completion.already_present_count == 1
    assert completion.fill_count == 0
    assert completion.entries[0].decision == CompletionDecision.ALREADY_PRESENT
    assert fill.write_count == 0
    assert fill.intents[0].decision == IntentDecision.SKIP

    dest = tmp_path / "out.xlsx"
    _, cell_diff, written, _, _ = ExcelFillService().apply_write_intents(
        analysis_id="c1",
        ticker="AAPL",
        source_workbook_path=workbook_with_revenue,
        destination_workbook_path=dest,
        intent_report=fill,
    )
    assert written == 0
    assert cell_diff.changed_count == 0

    prefill = PrefillValidationService().validate(completion)
    assert prefill.validated_count == 1
    assert prefill.entries[0].status == PrefillValidationStatus.VALIDATED


def test_blank_required_cell_fill(workbook_with_revenue: Path, tmp_path: Path):
    intents = WriteIntentReport(
        analysis_id="c2",
        ticker="AAPL",
        intents=[
            _intent(
                intent_id="is.revenue:L9",
                cell="L9",
                value=391035.0,
                period="FY2025",
            ),
        ],
        write_count=1,
    )
    completion, fill = CompletionService().plan(
        analysis_id="c2",
        ticker="AAPL",
        analysis_type="annual_update",
        source_workbook_path=workbook_with_revenue,
        intent_report=intents,
        target_fiscal_year="FY2025",
    )
    assert completion.fill_count == 1
    assert completion.entries[0].decision == CompletionDecision.FILL
    assert fill.write_count == 1

    dest = tmp_path / "out.xlsx"
    _, cell_diff, written, _, _ = ExcelFillService().apply_write_intents(
        analysis_id="c2",
        ticker="AAPL",
        source_workbook_path=workbook_with_revenue,
        destination_workbook_path=dest,
        intent_report=fill,
    )
    assert written == 1
    assert cell_diff.changed_count == 1


def test_populated_incorrect_cell_discrepancy_not_overwritten(
    workbook_with_revenue: Path, tmp_path: Path
):
    intents = WriteIntentReport(
        analysis_id="c3",
        ticker="AAPL",
        intents=[
            _intent(
                intent_id="is.revenue:G9",
                cell="G9",
                value=274515.0,
                period="FY2020",
            ),
        ],
        write_count=1,
    )
    # FY2020 is prior year under annual_update FY2025 → OUT_OF_SCOPE,
    # but validation still reviews the populated value.
    completion, fill = CompletionService().plan(
        analysis_id="c3",
        ticker="AAPL",
        analysis_type="annual_update",
        source_workbook_path=workbook_with_revenue,
        intent_report=intents,
        target_fiscal_year="FY2025",
    )
    assert completion.entries[0].decision == CompletionDecision.OUT_OF_SCOPE
    assert fill.write_count == 0

    dest = tmp_path / "out.xlsx"
    ExcelFillService().apply_write_intents(
        analysis_id="c3",
        ticker="AAPL",
        source_workbook_path=workbook_with_revenue,
        destination_workbook_path=dest,
        intent_report=fill,
    )
    from openpyxl import load_workbook

    wb = load_workbook(dest)
    try:
        assert wb["Income - GAAP"]["G9"].value == 999999  # unchanged
    finally:
        wb.close()

    prefill = PrefillValidationService().validate(completion)
    assert prefill.discrepancy_count == 1
    assert prefill.entries[0].status == PrefillValidationStatus.DISCREPANCY
    assert prefill.entries[0].correction_applied is False
    assert prefill.entries[0].workbook_value == 999999
    assert prefill.entries[0].sec_value == 274515.0


def test_formula_blocked(workbook_with_revenue: Path):
    # Put formula on in-scope FY2025 cell
    from openpyxl import load_workbook

    wb = load_workbook(workbook_with_revenue)
    wb["Income - GAAP"]["L9"] = "=E9*2"
    wb.save(workbook_with_revenue)
    wb.close()

    intents = WriteIntentReport(
        analysis_id="c4",
        ticker="AAPL",
        intents=[
            _intent(
                intent_id="x:L9",
                cell="L9",
                value=1.0,
                period="FY2025",
            ),
        ],
        write_count=1,
    )
    completion, fill = CompletionService().plan(
        analysis_id="c4",
        ticker="AAPL",
        analysis_type="annual_update",
        source_workbook_path=workbook_with_revenue,
        intent_report=intents,
        target_fiscal_year="FY2025",
    )
    assert completion.blocked_count == 1
    assert fill.intents[0].decision == IntentDecision.BLOCK


def test_values_agree_tolerance():
    assert values_agree(265595, 265595.0)
    assert values_agree(1000.0, 1000.4)
    assert not values_agree(999999, 274515)


def test_new_company_historical_blank_not_filled(workbook_with_revenue: Path):
    intents = WriteIntentReport(
        analysis_id="c5",
        ticker="AAPL",
        intents=[
            _intent(intent_id="is.revenue:F9", cell="F9", value=260174.0, period="FY2019"),
        ],
        write_count=1,
    )
    completion, fill = CompletionService().plan(
        analysis_id="c5",
        ticker="AAPL",
        analysis_type="new_company",
        source_workbook_path=workbook_with_revenue,
        intent_report=intents,
    )
    assert completion.entries[0].decision == CompletionDecision.OUT_OF_SCOPE
    assert fill.write_count == 0
