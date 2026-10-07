"""Plain-language notes in one place per tab, full cumulative statements that tie, green error cells, ROIC formatting, valuation reviews."""

from __future__ import annotations

from datetime import datetime
from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook

from services.new_company_cumulative_service import NewCompanyCumulativeService
from services.tab_notes import NOTES_TITLE, add_notes, note, notes_column, read_notes
from services.valuation_inputs_review_service import ValuationInputsReviewService
from tests.test_wwd_feedback import LQ_CF, LQ_IS, _facts, _quarter_workbook


# ---------------------------------------------------------------- notes: one block, plain language
def test_note_answers_what_why_and_source_in_plain_sentences():
    text = note("Filled the year-to-date columns", "Bloomberg gives only the three-month quarter", "the company's 10-Q")
    assert text == "Filled the year-to-date columns because Bloomberg gives only the three-month quarter. Source: the company's 10-Q."
    assert note("Nothing changed") == "Nothing changed."
    assert note("Reviewed X", "It can distort Y", because=False) == "Reviewed X. It can distort Y."


def test_one_notes_block_per_tab_old_blocks_removed_and_no_repeats():
    wb = Workbook()
    ws = wb.active
    ws["A1"], ws["B1"] = "Revenue", 100
    ws["A20"] = "HAP ANALYSIS — NOTES"
    ws["A21"], ws["B21"] = "Presentation decision", "BLOOMBERG_PRESERVE"
    add_notes(ws, ["First note.", "Second note."])
    add_notes(ws, ["Second note.", "Third note."])
    notes = read_notes(ws)
    assert notes == ["First note.", "Second note.", "Third note."]            # merged, nothing repeated
    values = [str(c.value) for row in ws.iter_rows() for c in row if c.value]
    assert "HAP ANALYSIS — NOTES" not in values and "BLOOMBERG_PRESERVE" not in values
    assert sum(1 for v in values if v == NOTES_TITLE) == 1                       # a single block
    add_notes(ws, ["Revised second note."], replace_containing=("Second note",))
    assert read_notes(ws) == ["First note.", "Third note.", "Revised second note."]


def test_notes_use_column_d_when_column_b_is_hidden():
    wb = Workbook()
    ws = wb.active
    ws["A1"] = "Revenue"
    ws.column_dimensions["B"].hidden = True
    assert notes_column(ws) == 4
    add_notes(ws, ["A note."])
    assert all(ws.cell(r, 2).value is None for r in range(1, ws.max_row + 1))
    assert any(ws.cell(r, 4).value == "A note." for r in range(1, ws.max_row + 1))


# ---------------------------------------------------------------- cumulative statements tie out
def test_income_statement_lines_tie_the_check_rows():
    v = {"revenue": 3196.727, "cost_of_revenue": 2238.752, "gross_profit": 957.975, "operating_income": 521.049, "sga": 330.0, "rd": 137.678,
         "interest_expense": 40.0, "interest_income": 5.0, "pretax": 480.0, "tax": 65.6, "net_income": 414.407}
    lines = NewCompanyCumulativeService._income_lines(v, {"derived": [], "yahoo": [], "plug": []})
    g = lambda code: lines.get(code, (0.0, ""))[0]
    assert round(g("SALES_REV_TURN") - g("IS_COGS_TO_FE_AND_PP_AND_G") - g("GROSS_PROFIT"), 1) == 0                     # row 22
    assert round(g("IS_OPERATING_EXPN") - g("IS_SG&A_EXPENSE") - g("IS_OPERATING_EXPENSES_R&D") - g("OTHER_OPERATING_EXPENSES_RATIO"), 1) == 0   # row 31
    assert round(g("GROSS_PROFIT") - g("IS_OPERATING_EXPN") - g("IS_OPER_INC"), 1) == 0                                # row 33
    assert round(g("NONOP_INCOME_LOSS") - g("IS_NET_INTEREST_EXPENSE") - g("OTHER_NONOP_INCOME_LOSS"), 1) == 0          # row 43
    assert round(g("IS_OPER_INC") - g("NONOP_INCOME_LOSS") - g("PRETAX_INC"), 1) == 0                                  # row 45
    assert round(g("PRETAX_INC") - g("IS_INC_TAX_EXP") - g("IS_SH_PRO_EQY_MT_INV_NET_OF_TAX") - g("IS_INC_BEF_XO_ITEM"), 1) == 0   # row 53
    assert round(g("IS_INC_BEF_XO_ITEM") - g("NET_INCOME"), 1) == 0                                                    # row 61
    assert round(g("NET_INCOME") - g("EARN_FOR_COMMON"), 1) == 0                                                       # row 65


def test_cash_flow_lines_tie_the_check_rows():
    v = {"net_income": 414.407, "cfo": 351.937, "cfi": -286.795, "cff": 88.769, "da": 120.0, "sbc": 20.0, "deferred_tax": -5.0, "ar": -60.0,
         "inventory": -80.0, "ap": 10.0, "capex": -156.337, "ppe_sale": 3.0, "dividends": -26.0, "lt_proceeds": 300.0, "lt_repay": -150.0,
         "repurchase": -100.0, "issuance": 12.0, "net_change_cash": 147.42}
    lines = NewCompanyCumulativeService._cash_flow_lines(v)
    g = lambda code: lines.get(code, (0.0, ""))[0]
    assert round(g("NON_CASH_ITEMS_DETAILED") - g("CF_STOCK_BASED_COMPENSATION") - g("CF_DEF_INC_TAX") - g("OTHER_NON_CASH_ADJ_LESS_DETAILED"), 1) == 0     # row 19
    assert round(g("CF_NET_INC") + g("CF_DEPR_AMORT") + g("NON_CASH_ITEMS_DETAILED") + g("CF_CHNG_NON_CASH_WORK_CAP") - g("CF_CASH_FROM_OPER"), 1) == 0  # row 28
    assert round(g("CF_CASH_FROM_INV_ACT") - g("CHG_IN_FXD_&_INTANG_AST_DETAILED") - g("OTHER_INVESTING_ACT_DETAILED"), 1) == 0                       # row 50
    assert round(g("PROC_FR_REPURCH_EQTY_DETAILED") - g("CF_INCR_CAP_STOCK") - g("CF_DECR_CAP_STOCK"), 1) == 0                                        # row 61
    assert round(g("CF_DVD_PAID") + g("PROC_FR_REPAYMNTS_BOR_DETAILED") + g("PROC_FR_REPURCH_EQTY_DETAILED") + g("CF_OTHER_FINANCING_ACT_EXCL_FX")
                 - g("CFF_ACTIVITIES_DETAILED"), 1) == 0                                                                                            # row 65


def test_error_cells_are_green_when_zero_and_red_otherwise(tmp_path: Path):
    path = _quarter_workbook(tmp_path / "wb.xlsx")
    report = NewCompanyCumulativeService().apply(workbook_path=path, company_facts=_facts(), latest_quarter=3)
    assert report.error_cells == {LQ_IS: "L1", LQ_CF: "H1"}
    wb = load_workbook(path)
    for sheet, addr in ((LQ_IS, "L1"), (LQ_CF, "H1")):
        rules = [r for rng, rs in wb[sheet].conditional_formatting._cf_rules.items() if addr in str(rng.sqref) for r in rs]
        colors = {(r.operator, r.dxf.fill.bgColor.rgb[-6:]) for r in rules}
        assert ("equal", "C6EFCE") in colors and ("greaterThan", "FFC7CE") in colors


# ---------------------------------------------------------------- ROIC tab formatting
def test_projection_columns_are_coloured_and_ratios_are_percentages(tmp_path: Path):
    from tests.test_wwd_feedback import _ic_workbook
    from services.ic_projection_writer import IcProjectionWriter

    path = _ic_workbook(tmp_path / "wb.xlsx", {"revenue": 3000.0, "oi": 450.0})
    IcProjectionWriter().apply(workbook_path=path, latest_quarter=3)
    ws = load_workbook(path)["IC & NOPAT & ROIC "]
    assert ws["M3"].fill.fill_type == "solid" and ws["N4"].fill.fill_type == "solid" and ws["M3"].fill.fgColor.rgb != ws["N4"].fill.fgColor.rgb
    for addr in ("L23", "M23", "M24", "M25", "N4"):
        assert ws[addr].number_format == "0.0%", addr
    assert ws["M20"].number_format != "0.0%"                                      # amounts stay amounts
    notes = [str(c.value) for row in ws.iter_rows() for c in row if c.value]
    assert any(n.startswith("Projected ROIC (column M) and ROCE (column N) were calculated") for n in notes)
    assert notes.count(NOTES_TITLE) == 1


# ---------------------------------------------------------------- valuation inputs review
def _valuation_workbook(path: Path, *, oe: list[float], roe: list[float], g: float = 0.11, annual: float | None = None) -> Path:
    wb = Workbook()
    fm = wb.active
    fm.title = "Final Metrics"
    years = [f"FY{2016 + i}" for i in range(len(oe))]
    for i, fy in enumerate(years):
        fm.cell(1, 3 + i).value = fy
        fm.cell(4, 3 + i).value = roe[i]
        fm.cell(10, 3 + i).value = oe[i]
    fm["A4"], fm["A10"] = "ROE", "Owners Earnings"
    er = wb.create_sheet("Expected Returns & Buybacks")
    er["A11"], er["A14"], er["C5"], er["B8"], er["E2"], er["A2"], er["E14"] = g, sum(roe) / len(roe), 0.84, 40.0, 29.0, 335.0, 0.033
    ev = wb.create_sheet("Enterprise Value")
    first, last = oe[0], oe[-1]
    ev["C6"] = annual if annual is not None else (last / first) ** (1 / (len(oe) - 1)) - 1
    ev["B2"], ev["B17"], ev["B19"], ev["B20"] = 0.12, 2000.0, 20000.0, 400.0
    wb.save(path)
    return path


def test_owner_earnings_growth_very_high_negative_and_low_are_flagged(tmp_path: Path):
    oe_high = [74, 200, 260, 280, 290, 300, 310, 330, 380, 424]
    review = ValuationInputsReviewService().review(_valuation_workbook(tmp_path / "a.xlsx", oe=oe_high, roe=[0.13] * 10))
    flagged = [f for f in review.findings if f.topic == "owner_earnings"]
    assert flagged[0].verdict == "flag" and "very high" in flagged[0].text and "FY2016" in flagged[0].text
    assert "company value per share" in review.alternatives["owner_earnings"]
    review_neg = ValuationInputsReviewService().review(_valuation_workbook(tmp_path / "b.xlsx", oe=[300, 280, 260, 250, 240, 230, 220, 210, 205, 200], roe=[0.13] * 10))
    assert "negative" in review_neg.findings[-1].text
    review_low = ValuationInputsReviewService().review(_valuation_workbook(tmp_path / "c.xlsx", oe=[300, 301, 302, 303, 304, 305, 306, 307, 308, 309], roe=[0.13] * 10))
    assert "unrealistically low" in review_low.findings[-1].text
    review_ok = ValuationInputsReviewService().review(_valuation_workbook(tmp_path / "d.xlsx", oe=[200, 215, 230, 248, 265, 285, 305, 330, 355, 380], roe=[0.13] * 10))
    assert review_ok.findings[-1].verdict == "reasonable"
    assert all(len(n) < 600 for n in review.notes("owner_earnings"))
    assert review.notes("owner_earnings")[-1] == "No model value was changed. The flagged inputs are for your decision."


def test_expected_return_flags_distorted_roe_and_unrealistic_book_value_growth(tmp_path: Path):
    roe = [0.12, 0.13, 0.12, 0.14, 0.13, 0.70, 0.12, 0.13, 0.14, 0.13]       # one extraordinary year pulls the average up
    path = _valuation_workbook(tmp_path / "a.xlsx", oe=[100] * 10, roe=roe, g=0.30)
    review = ValuationInputsReviewService().review(path)
    by = {f.metric: f for f in review.findings if f.topic == "expected_return"}
    assert by["average_roe"].verdict == "flag" and "FY2021" in by["average_roe"].text
    assert by["book_value_growth"].verdict == "flag" and "aggressive" in by["book_value_growth"].text
    assert "typical return on equity" in review.alternatives["expected_return"]
    notes = review.notes("expected_return")
    assert notes[0].startswith("Reviewed the expected return model (book value growth and return on equity)") and "Source:" in notes[0]
    calm = ValuationInputsReviewService().review(_valuation_workbook(tmp_path / "b.xlsx", oe=[100] * 10, roe=[0.13] * 10, g=0.11))
    assert [f.verdict for f in calm.findings if f.topic == "expected_return"][0] == "reasonable"
    assert "look realistic" in calm.notes("expected_return")[0]


def test_template_error_cell_rules_are_not_duplicated(tmp_path: Path):
    from openpyxl.formatting.rule import CellIsRule
    from openpyxl.styles import PatternFill

    path = _quarter_workbook(tmp_path / "wb.xlsx")
    wb = load_workbook(path)
    green = PatternFill("solid", start_color="92D050", end_color="92D050")
    red = PatternFill("solid", start_color="FF0000", end_color="FF0000")
    for sheet, addr in ((LQ_IS, "L1"), (LQ_CF, "H1")):
        wb[sheet].conditional_formatting.add(addr, CellIsRule(operator="equal", formula=["0"], fill=green))
        wb[sheet].conditional_formatting.add(addr, CellIsRule(operator="notEqual", formula=["0"], fill=red))
    wb.save(path)
    NewCompanyCumulativeService().apply(workbook_path=path, company_facts=_facts(), latest_quarter=3)
    wb = load_workbook(path)
    for sheet, addr in ((LQ_IS, "L1"), (LQ_CF, "H1")):
        rules = [r for rng, rs in wb[sheet].conditional_formatting._cf_rules.items() if addr in str(rng.sqref) for r in rs]
        assert len(rules) == 2          # the template's own green (0) and red (not 0) rules are used as they are


# ---------------------------------------------------------------- corrections are applied, not only flagged
def _apply(path: Path):
    from services.valuation_correction_service import ValuationCorrectionService

    return ValuationCorrectionService().apply(path)


def test_a_clearly_wrong_owner_earnings_growth_replaces_the_model_value(tmp_path: Path):
    path = _valuation_workbook(tmp_path / "c.xlsx", oe=[74, 200, 260, 280, 290, 300, 310, 330, 380, 424], roe=[0.13] * 10)
    review = _apply(path)
    ev = load_workbook(path)["Enterprise Value"]
    assert 0.10 < ev["C6"].value < 0.15 and ev["C6"].number_format == "0.0%"             # was about 21%, now the 3-year-average rate
    applied = [a for a in review.applied if a.topic == "owner_earnings"][0]
    assert applied.kind == "replaced" and "was replaced with" in applied.what and "FY2016" in applied.why


def test_a_borderline_owner_earnings_growth_gets_a_parallel_calculation(tmp_path: Path):
    oe = [280, 300, 320, 345, 370, 400, 430, 460, 490, 520]      # about 7% a year; the start year is normal
    path = _valuation_workbook(tmp_path / "d.xlsx", oe=oe, roe=[0.13] * 10, annual=0.17)
    review = _apply(path)
    ev = load_workbook(path)["Enterprise Value"]
    applied = [a for a in review.applied if a.topic == "owner_earnings"]
    assert applied and applied[0].kind == "parallel" and ev["C6"].value == 0.17            # the original stays; a corrected copy sits below
    assert "HAP Alternative" in str(ev["A54"].value) or any("HAP" in str(c.value) for row in ev.iter_rows(min_row=50) for c in row if c.value)


def test_realistic_inputs_are_not_touched(tmp_path: Path):
    path = _valuation_workbook(tmp_path / "e.xlsx", oe=[200, 215, 230, 248, 265, 285, 305, 330, 355, 380], roe=[0.13] * 10, g=0.11)
    before = path.read_bytes()
    review = _apply(path)
    assert not review.applied and path.read_bytes() == before


def _er_workbook(path: Path, *, model_return: float, growth: float = 0.0491, eps0: float = 3.17, bv0: float = 14.5621, price: float = 93.34,
                 pe: float = 25.117) -> Path:
    """An Expected Returns tab as the template computes it (values in place of the cached results), with formulas in the two input cells."""
    wb = Workbook()
    er = wb.active
    er.title = "Expected Returns & Buybacks"
    er["A2"], er["E2"], er["B5"], er["B8"], er["C30"] = price, pe, growth, bv0, eps0
    er["A11"], er["A14"] = "=C5*A14", "='Final Metrics'!B4"
    er["D17"], er["E17"] = 4.78, 2.17                                           # first projected EPS and dividend: payout 45.4%
    er["E14"], er["F14"] = model_return - 0.01, model_return
    wb.create_sheet("Final Metrics")
    wb.save(path)
    return path


def test_an_unrealistic_expected_return_switches_to_the_eps_growth_model(tmp_path: Path):
    path = _er_workbook(tmp_path / "a.xlsx", model_return=0.14)
    review = _apply(path)
    wb = load_workbook(path)
    er = wb["Expected Returns & Buybacks"]
    assert er["A14"].value == pytest.approx(3.17 / 14.5621, abs=1e-4)             # current EPS / book value per share
    assert er["A11"].value == pytest.approx(0.0491)                                # the EPS growth rate
    assert er["A11"].number_format == "0.0%"
    applied = [a for a in review.applied if a.topic == "expected_return"]
    assert {a.cell.split("!")[-1] for a in applied} == {"A11", "A14"} and applied[0].original == "='Final Metrics'!B4"
    ledger = [str(c.value) for row in wb["HAP Adjustments"].iter_rows(min_row=5) for c in row if c.value]
    assert "Expected return input" in ledger and "='Final Metrics'!B4" in ledger and "=C5*A14" in ledger        # the original formulas are kept
    notes = review.notes("expected_return")
    assert len(notes) == 2 and "EPS growth of 4.9% a year" in notes[0] and "it was 14.0%" in notes[0] and "too high" in notes[0]
    # the same arithmetic as the template: EPS grows 4.91% a year from 3.17, valued at 25.1x, plus dividends at the model's payout
    eps = [3.17 * 1.0491 ** t for t in range(1, 11)]
    expected = ((eps[-1] * 25.117 + (2.17 / 4.78) * sum(eps)) / 93.34) ** 0.1 - 1
    assert f"the return is now {expected * 100:.1f}%" in notes[0]


def test_a_negative_expected_return_also_switches_models(tmp_path: Path):
    review = _apply(_er_workbook(tmp_path / "b.xlsx", model_return=-0.03))
    assert review.applied and "negative" in review.applied[0].why


def test_a_realistic_expected_return_is_left_alone(tmp_path: Path):
    path = _er_workbook(tmp_path / "c.xlsx", model_return=0.08)
    before = path.read_bytes()
    review = _apply(path)
    assert not review.applied and path.read_bytes() == before
    assert "within a realistic range of 0.0% to 12.0%" in review.notes("expected_return")[0]


def test_the_eps_model_is_not_used_when_eps_growth_is_extreme_or_eps_is_not_positive(tmp_path: Path):
    for name, kwargs in (("d", {"growth": 0.40}), ("e", {"eps0": -1.0}), ("f", {"bv0": 0.0})):
        path = _er_workbook(tmp_path / f"{name}.xlsx", model_return=0.20, **kwargs)
        before = path.read_bytes()
        review = _apply(path)
        assert not review.applied and path.read_bytes() == before and review.left_as_is["expected_return"].startswith("the EPS-growth model could not be used")
