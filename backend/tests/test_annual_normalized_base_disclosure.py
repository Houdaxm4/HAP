"""Normalized earnings-power analyst disclosure — communication only."""

from __future__ import annotations

from pathlib import Path

from docx import Document
from openpyxl import load_workbook

from models.annual_update import (
    DISCLOSE_DISTORTED_BASE,
    DISCLOSE_NORMALIZED_BASE,
    INSUFFICIENT_BASE_EVIDENCE,
    KEEP_REPORTED_BASE,
    MANAGEMENT_STATEMENT,
    NO_DISCLOSURE_REQUIRED,
    REMAINS_INSUFFICIENT,
    RESEARCHED,
    AnnualAnalystJudgmentReport,
    AnnualExpectedReturnReport,
    AnnualResearchReport,
    AnalyticalResearchEvidence,
    AnalyticalResearchReport,
    AnalyticalResearchSynthesis,
    GROWTH_CAPEX,
    JudgmentRecord,
    ResearchQuestion,
)
from services.annual_deliverables_service import AnnualDeliverablesService
from services.annual_judgment_service import AnnualJudgmentService
from services.annual_normalized_base_disclosure_service import (
    AnnualNormalizedBaseDisclosureService,
    format_millions,
)
from tests.test_annual_analytical_research import _keep_base
from tests.test_annual_growth_analysis import _geo, _valuation_workbook
from tests.test_annual_normalized_base import _analyze, _snapshot_ev


def _expansion_research() -> AnalyticalResearchReport:
    claim = (
        "We plan to invest approximately $90 million in equipment and infrastructure "
        "continuing into early fiscal 2027. We expect total capital expenditures of "
        "approximately $48 million in fiscal 2027, including project completion and "
        "ongoing maintenance spending."
    )
    return AnalyticalResearchReport(
        analysis_id="t",
        ticker="ZZ",
        fiscal_year=2026,
        analysis_as_of_date="2026-08-19",
        status=RESEARCHED,
        questions=[
            ResearchQuestion(
                question=(
                    "Is the CapEx increase a disclosed expansion program, and does that "
                    "establish a sustainable post-project reinvestment level?"
                ),
                concept="oe_base",
            )
        ],
        evidence=[
            AnalyticalResearchEvidence(
                company="ZZ",
                metric_or_issue="operating_earnings_base",
                research_question="capex program",
                source_type=MANAGEMENT_STATEMENT,
                source_title="FY2026 Form 10-K",
                source_locator="https://www.sec.gov/Archives/example/d10k.htm#item7",
                management_claim=claim,
                verified_financial_fact=(
                    "Purchases of property and equipment were $28.3 million, $50.7 million "
                    "and $88.1 million."
                ),
                available_as_of_analysis=True,
                analysis_as_of_date="2026-08-19",
                publication_date="2026-08-19",
            )
        ],
        synthesis=AnalyticalResearchSynthesis(
            issue="elevated_capex",
            question="sustainable post-project capex",
            hap_conclusion=GROWTH_CAPEX,
            decision_effect=REMAINS_INSUFFICIENT,
            confidence="MEDIUM",
            missing_evidence=[
                "Sustainable post-project reinvestment versus historical intensity is not established."
            ],
            hap_interpretation="Program identified; sustainable maintenance CapEx is not established.",
        ),
        writes_to_workbook=False,
        selected_normalized_base_written=False,
    )


def _jbss_like_base():
    return _analyze(
        ni=[62.0] * 10,
        da=[20.0] * 9 + [28.0],
        capex=[-25.0] * 7 + [-28.3, -50.7, -88.1],
        revenue=[1000.0] * 10,
        operating_income=[89.0] * 10,
    )


def _mild_keep_flag():
    return _analyze(
        ni=[62.0] * 10,
        da=[20.0] * 10,
        capex=[-20.0] * 9 + [-45.0],
        revenue=[1000.0] * 10,
        operating_income=[85.0] * 10,
    )


def _collapse_base():
    return _analyze(
        ni=[40.0] * 9 + [6.0],
        da=[12.0] * 10,
        capex=[-11.0] * 10,
        revenue=[400.0] * 10,
        operating_income=[48.0] * 9 + [7.0],
    )


def _forbidden(text: str) -> None:
    low = text.lower()
    for phrase in (
        "true oe is",
        "normalized oe is",
        "maintenance capex should",
        "capex will revert",
        "temporary capex",
        "the stock is undervalued",
        "the valuation is wrong",
        "fair value should",
        "not hap facts",
        "those characterizations are management statements",
    ):
        assert phrase not in low, phrase


def test_jbss_profile_produces_disclose_distorted_base():
    base = _jbss_like_base()
    assert base.reported_base_status == "materially_distorted"
    assert base.decision == INSUFFICIENT_BASE_EVIDENCE
    disc = AnnualNormalizedBaseDisclosureService().build(base, _expansion_research())
    assert disc.decision == DISCLOSE_DISTORTED_BASE
    assert disc.decision != DISCLOSE_NORMALIZED_BASE
    assert disc.reported_base == base.reported_current_base
    assert "28.3" in disc.display_text and "50.7" in disc.display_text and "88.1" in disc.display_text
    assert "90" in disc.research_summary
    assert "48" in disc.research_summary


def test_disclosure_reports_reported_base_not_candidate():
    base = _jbss_like_base()
    selected_before = base.selected_normalized_base
    disc = AnnualNormalizedBaseDisclosureService().build(base, _expansion_research())
    assert disc.reported_base == base.reported_current_base
    assert abs(disc.reported_base - 1.9) < 0.2
    assert selected_before is None
    assert base.selected_normalized_base is None
    candidates = [
        c.result
        for c in base.candidate_bases
        if c.result is not None and c.method != "REPORTED_CURRENT_BASE"
    ]
    for result in candidates:
        if abs(result - disc.reported_base) < 1:
            continue
        token = format_millions(result)
        assert f"normalized oe is ${token}" not in disc.display_text.lower()
        assert f"true oe is ${token}" not in disc.display_text.lower()


def test_disclosure_does_not_imply_candidate_equals_normalized_base():
    base = _jbss_like_base()
    disc = AnnualNormalizedBaseDisclosureService().build(base, _expansion_research())
    low = disc.display_text.lower()
    assert "does not substitute" in low or "has not replaced" in low
    assert disc.decision != DISCLOSE_NORMALIZED_BASE


def test_insufficient_evidence_language_preserves_uncertainty():
    base = _jbss_like_base()
    disc = AnnualNormalizedBaseDisclosureService().build(base, _expansion_research())
    low = disc.display_text.lower()
    assert "cannot determine" in low
    assert "sufficient confidence" in low
    _forbidden(disc.display_text)
    assert "will revert" not in low


def test_same_year_capex_guide_is_not_treated_as_subsequent():
    claim = (
        "Beginning in fiscal 2025 and continuing into early fiscal 2027, we will invest "
        "approximately $90.0 million in capital expenditures and related expenses. "
        "We expect total capital expenditures for equipment purchases and upgrades for "
        "fiscal 2025 to be approximately $95.0 million."
    )
    research = _expansion_research()
    research.evidence[0].management_claim = claim
    research.evidence[0].filing_type = "10-K"
    research.evidence[0].source_date = "2026-08-19"
    disc = AnnualNormalizedBaseDisclosureService().build(_jbss_like_base(), research)
    assert "90" in disc.research_summary
    assert "95" not in disc.research_summary
    assert "48" not in disc.research_summary


def test_research_evidence_can_inform_disclosure():
    base = _jbss_like_base()
    without = AnnualNormalizedBaseDisclosureService().build(base, None)
    with_r = AnnualNormalizedBaseDisclosureService().build(base, _expansion_research())
    assert with_r.research_summary
    assert "90" in with_r.display_text
    assert "48" in with_r.display_text
    assert with_r.provenance
    assert any("annual_analytical_research_report.json" in p for p in with_r.provenance)
    assert without.decision == DISCLOSE_DISTORTED_BASE
    assert "90" not in (without.research_summary or "")


def test_management_claims_remain_qualified():
    disc = AnnualNormalizedBaseDisclosureService().build(_jbss_like_base(), _expansion_research())
    low = disc.display_text.lower()
    assert "management's characterization" in low or "management characterization" in low
    assert "independently established conclusion" in low
    assert "not hap facts" not in low
    assert "those characterizations are management statements" not in low


def test_valuation_implication_warns_without_declaring_wrong():
    disc = AnnualNormalizedBaseDisclosureService().build(_jbss_like_base(), _expansion_research())
    low = disc.valuation_implication.lower()
    assert "interpreted cautiously" in low
    assert "original valuation" in low
    assert "undervalued" not in low
    assert "wrong" not in low
    assert "fair value" not in low


def test_original_ev_cells_untouched_when_disclosure_written(tmp_path: Path):
    path = tmp_path / "wb.xlsx"
    base = _jbss_like_base()
    _valuation_workbook(
        path,
        ni=[62.0] * 10,
        da=[20.0] * 9 + [28.0],
        capex=[-25.0] * 7 + [-28.3, -50.7, -88.1],
        oi=[89.0] * 10,
        revenue=[1000.0] * 10,
        b6_total=-0.95,
        c6_annual=-0.28,
    )
    before = _snapshot_ev(path)
    disc = AnnualNormalizedBaseDisclosureService().build(base, _expansion_research())
    base.disclosure = disc
    _, judge = AnnualJudgmentService().apply(
        analysis_id="a1",
        ticker="ZZ",
        workbook_path=path,
        context={
            "oe_base_analysis": base,
            "analytical_research": _expansion_research(),
            "book_value_growth": 0.03,
            "owner_earnings_growth": 0.04,
            "eps_growth": 0.04,
            "bv_classification": "REASONABLE",
            "oe_classification": "REASONABLE",
            "eps_classification": "REASONABLE",
        },
    )
    after = _snapshot_ev(path)
    assert after == before
    assert after["B9"] == "='Final Metrics'!$L$10*(1+$C$6)^B8"
    wb = load_workbook(path, data_only=False)
    try:
        ev = wb["Enterprise Value"]
        found = False
        for row in ev.iter_rows(max_row=40, max_col=20):
            for cell in row:
                if cell.value == "OE BASE — HAP ANALYTICAL OBSERVATION":
                    found = True
                    assert ev.cell(cell.row, cell.column + 1).value == disc.display_text
        assert found
    finally:
        wb.close()
    assert judge.oe_base_analysis is not None
    assert judge.oe_base_analysis.disclosure is not None
    assert judge.oe_base_analysis.disclosure.decision == DISCLOSE_DISTORTED_BASE


def test_disclosure_does_not_create_selected_normalized_base():
    base = _jbss_like_base()
    assert base.selected_normalized_base is None
    AnnualNormalizedBaseDisclosureService().build(base, _expansion_research())
    assert base.selected_normalized_base is None
    collapse = _collapse_base()
    before = collapse.selected_normalized_base
    AnnualNormalizedBaseDisclosureService().build(collapse, None)
    assert collapse.selected_normalized_base == before


def test_etd_profile_produces_no_disclosure_required():
    base = _keep_base()
    assert base.decision == KEEP_REPORTED_BASE
    disc = AnnualNormalizedBaseDisclosureService().build(base, None)
    assert disc.decision == NO_DISCLOSURE_REQUIRED
    assert disc.display_text == ""


def test_csco_profile_produces_no_disclosure_required():
    rec = _analyze(
        ni=_geo(36.0, 0.04, 10),
        da=[15.0 + i * 0.5 for i in range(10)],
        capex=[-12.0 - i * 0.4 for i in range(10)],
        revenue=_geo(800.0, 0.037, 10),
        operating_income=_geo(60.0, 0.04, 10),
    )
    disc = AnnualNormalizedBaseDisclosureService().build(rec, None)
    assert rec.reported_base_status in {"usable", "potentially_distorted"}
    assert rec.decision == KEEP_REPORTED_BASE
    assert disc.decision == NO_DISCLOSURE_REQUIRED


def test_potentially_distorted_keep_does_not_force_disclosure():
    rec = _mild_keep_flag()
    assert rec.decision == KEEP_REPORTED_BASE
    assert rec.reported_base_status == "potentially_distorted"
    disc = AnnualNormalizedBaseDisclosureService().build(rec, None)
    assert disc.decision == NO_DISCLOSURE_REQUIRED
    assert disc.display_text == ""


def test_disclosure_works_for_non_capex_distortion():
    rec = _collapse_base()
    assert "ONE_TIME_EARNINGS_COLLAPSE" in rec.distortion_type
    assert rec.reported_base_status == "materially_distorted"
    disc = AnnualNormalizedBaseDisclosureService().build(rec, None)
    assert disc.decision == DISCLOSE_DISTORTED_BASE
    assert disc.primary_driver == "one_time_earnings_level"
    assert "elevated capital" not in disc.issue.lower()
    assert "one-period" in disc.display_text.lower() or "earnings-level" in disc.display_text.lower()
    assert disc.reported_base == rec.reported_current_base
    _forbidden(disc.display_text)


def test_no_ticker_specific_disclosure_logic():
    src = (
        Path(__file__).resolve().parents[1]
        / "services"
        / "annual_normalized_base_disclosure_service.py"
    ).read_text(encoding="utf-8")
    lowered = src.lower()
    for token in ("jbss", "etd", "csco", "mzti"):
        assert token not in lowered
    for name in (
        "annual_judgment_service.py",
        "annual_deliverables_service.py",
        "annual_update_runner.py",
    ):
        text = (Path(__file__).resolve().parents[1] / "services" / name).read_text(encoding="utf-8")
        assert "if ticker ==" not in text
        assert "ticker == \"" not in text
        assert "ticker == '" not in text


def test_word_distinguishes_original_valuation_from_hap_observation(tmp_path: Path):
    path = tmp_path / "wb.xlsx"
    _valuation_workbook(path)
    base = _jbss_like_base()
    base.disclosure = AnnualNormalizedBaseDisclosureService().build(base, _expansion_research())
    judgment = AnnualAnalystJudgmentReport(
        analysis_id="a1",
        ticker="ZZ",
        owner_earnings_growth=JudgmentRecord(
            metric="owner_earnings_growth",
            original_value=0.08,
            selected_value=0.08,
            adjusted=False,
            rationale="Existing OE growth appears reasonable prospectively.",
        ),
        graham_eps_growth=JudgmentRecord(
            metric="graham_eps_growth",
            original_value=0.07,
            selected_value=0.07,
            adjusted=False,
            rationale="Historical EPS 10-year growth was not automatically used as the prospective Graham rate.",
        ),
        hap_analysis_cells=["Enterprise Value!G12"],
        oe_base_analysis=base,
    )
    er = AnnualExpectedReturnReport(
        analysis_id="a1",
        ticker="ZZ",
        original_growth_rate=0.06,
        original_expected_return=0.11,
        selected_growth_rate=0.06,
        final_expected_return=None,
        reasonableness="reasonable",
        rationale="Existing assumption appears reasonable; no adjustment is recommended.",
        selected_methodology="BOOK_VALUE_GROWTH",
    )
    research = AnnualResearchReport(analysis_id="a1", ticker="ZZ", fiscal_year=2026)
    perf = AnnualDeliverablesService().build_performance(
        analysis_id="a1", ticker="ZZ", fiscal_year=2026, workbook_path=path, research=research
    )
    report = AnnualDeliverablesService().produce(
        analysis_id="a1",
        ticker="ZZ",
        fiscal_year=2026,
        completed_workbook_path=path,
        output_dir=tmp_path,
        performance=perf,
        research=research,
        judgment=judgment,
        expected_return=er,
    )
    text = "\n".join(p.text for p in Document(tmp_path / report.word_filename).paragraphs)
    assert "ORIGINAL ANALYST VALUATION" in text
    assert "HAP ANALYTICAL OBSERVATION" in text
    assert "has not replaced" in text.lower()
    _forbidden(text)
