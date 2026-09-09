"""Tests for analyst intelligence layer: completeness, tax research, judgment context."""

from __future__ import annotations

from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook

from services.annual_analyst_intelligence_service import AnnualAnalystIntelligenceService
from services.annual_judgment_service import AnnualJudgmentService
from services.annual_tax_research_service import AnnualTaxResearchService


def _mini_workbook(path: Path, *, tax_blank: bool = True) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Inputs"
    ws["A7"] = "Line"
    for i, y in enumerate(range(2017, 2027)):
        ws.cell(7, 3 + i, f"FY{y}")
    ws["A10"] = "PE10"
    ws.cell(10, 12, 15.5)
    ws["A106"] = "Tax Table"
    ws["A107"] = "Federal Tax"
    ws["A108"] = "State Taxes"
    ws["A109"] = "Foreign Taxes"
    ws["A110"] = "R&D Tax Credits"
    ws["A111"] = "All Other Items"
    ws["A112"] = "Income Tax Expense"
    if not tax_blank:
        ws.cell(107, 12, 0.21)
        ws.cell(108, 12, 0.03)
        ws.cell(111, 12, -0.02)
    ws["B63"] = 100.0

    er = wb.create_sheet("Expected Returns & Buybacks")
    er["A11"] = 0.03
    er["B5"] = 0.06
    er["A12"] = "Growth assumption"
    er["B12"] = None

    fm = wb.create_sheet("Final Metrics")
    fm["L31"] = 0.06
    fm["C31"] = 0.05
    for row, val in ((38, 50), (44, 10), (46, -8)):
        ws.cell(row, 3, val * 0.8)
        ws.cell(row, 12, val)

    ev = wb.create_sheet("Enterprise Value")
    ev["A6"] = "Owners Earnings Growth Rate"
    ev["B6"] = 0.45
    ev["A41"] = "EPS 10 Year CAGR"
    ev["B41"] = 0.06

    inc = wb.create_sheet("Income - GAAP")
    inc["A7"] = "Line"
    for i, y in enumerate(range(2017, 2027)):
        inc.cell(7, 3 + i, f"FY{y}")

    wb.save(path)
    wb.close()


def test_completeness_gate_flags_missing_tax(tmp_path: Path):
    wb_path = tmp_path / "wb.xlsx"
    _mini_workbook(wb_path, tax_blank=True)
    report = AnnualAnalystIntelligenceService().assess_completeness(
        analysis_id="a1",
        ticker="ZZ",
        workbook_path=wb_path,
        new_fiscal_year="FY2026",
    )
    missing = [i for i in report.items if i.status == "MISSING_REQUIRED_INPUT"]
    assert any(i.concept.startswith("tax_") for i in missing)
    pe10 = next(i for i in report.items if i.concept == "pe10_new_fy")
    assert pe10.status == "POPULATED"


def test_tax_research_computes_etr_from_sec():
    company_facts = {
        "facts": {
            "us-gaap": {
                "IncomeTaxExpenseBenefit": {
                    "units": {
                        "USD": [
                            {
                                "val": 100.0,
                                "fy": 2026,
                                "fp": "FY",
                                "form": "10-K",
                                "filed": "2026-08-19",
                                "end": "2026-06-25",
                            }
                        ]
                    }
                },
                "IncomeLossFromContinuingOperationsBeforeIncomeTaxesExtraordinaryItemsNoncontrollingInterest": {
                    "units": {
                        "USD": [
                            {
                                "val": 400.0,
                                "fy": 2026,
                                "fp": "FY",
                                "form": "10-K",
                                "filed": "2026-08-19",
                                "end": "2026-06-25",
                            }
                        ]
                    }
                },
            }
        }
    }
    payload, questions = AnnualTaxResearchService().extract_tax_inputs(
        company_facts=company_facts,
        fiscal_year="FY2026",
        ticker="ZZ",
    )
    assert payload["reported_effective_rate"] == pytest.approx(0.25)
    assert questions[0].unsuccessful is False


def test_judgment_context_classifies_aggressive_oe(tmp_path: Path):
    wb_path = tmp_path / "wb.xlsx"
    _mini_workbook(wb_path)
    ctx = AnnualAnalystIntelligenceService().build_judgment_context(
        analysis_id="a1",
        ticker="ZZ",
        workbook_path=wb_path,
    )
    assert ctx["oe_classification"] in {"AGGRESSIVE", "VERY_AGGRESSIVE", "DISTORTED"}
    assert ctx.get("normalized_oe_growth") is not None


def test_judgment_writes_only_assumption_cells(tmp_path: Path):
    wb_path = tmp_path / "wb.xlsx"
    _mini_workbook(wb_path)
    svc = AnnualJudgmentService()
    _, judge = svc.apply(
        analysis_id="a1",
        ticker="ZZ",
        workbook_path=wb_path,
        context={
            "book_value_growth": 0.03,
            "owner_earnings_growth": 0.45,
            "eps_growth": 0.06,
            "oe_classification": "AGGRESSIVE",
            "normalized_oe_growth": 0.08,
            "assumption_cells": {"er_growth": "Expected Returns & Buybacks!B12"},
        },
    )
    wb = load_workbook(wb_path, data_only=True)
    assert wb["Expected Returns & Buybacks"]["B12"].value == pytest.approx(0.03)
    assert judge.owner_earnings_growth and judge.owner_earnings_growth.adjusted
    wb.close()
