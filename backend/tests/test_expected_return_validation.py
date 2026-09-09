"""Focused tests for Expected Return / EPS growth review."""

from __future__ import annotations

from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook

from models.expected_return_validation import (
    AdjustmentNecessity,
    ExpectedReturnDecision,
)
from services.expected_return_validation_service import (
    ExpectedReturnValidationService,
    _cagr,
    _yoy,
    interpret_attractiveness,
    reconstruct_expected_annual_return,
)


def test_normal_eps_series_cagr():
    assert _cagr(2.0, 4.0, 9) == pytest.approx((4 / 2) ** (1 / 9) - 1)


def test_split_adjusted_eps_not_discrepancy(tmp_path: Path, monkeypatch):
    path = _er_workbook(
        tmp_path / "split.xlsx",
        eps=[2.0775, 2.3, 3.0, 3.0, 3.28, 5.61, 6.11, 6.13, 6.08, 7.46],
        price=200.0,
        max_pe=25.0,
        terminal_price=300.0,
    )
    sec_map = {
        "FY2016": 8.31,
        "FY2017": 9.2,
        "FY2018": 12.0,
        "FY2019": 12.0,
        "FY2020": 3.28,
        "FY2021": 5.61,
        "FY2022": 6.11,
        "FY2023": 6.13,
        "FY2024": 6.08,
        "FY2025": 7.46,
    }
    monkeypatch.setattr(
        ExpectedReturnValidationService,
        "_sec_eps_map",
        lambda self, facts: sec_map,
    )
    report = ExpectedReturnValidationService().validate(
        analysis_id="er1",
        ticker="AAPL",
        workbook_path=path,
        company_facts={"ok": True},
    )
    early = [p for p in report.eps_history if p.period in {"FY2016", "FY2017", "FY2018"}]
    assert early
    assert all(p.decision == ExpectedReturnDecision.NOT_COMPARABLE for p in early)
    late = next(p for p in report.eps_history if p.period == "FY2025")
    assert late.decision == ExpectedReturnDecision.VALIDATED


def test_negative_near_zero_eps_denominator():
    assert _yoy(0.0, 1.0) is None
    assert _cagr(-1.0, 2.0, 5) is None
    assert _cagr(0.0, 2.0, 5) is None


def test_one_time_abnormal_year_with_filing_support(tmp_path: Path, monkeypatch):
    path = _er_workbook(
        tmp_path / "one_time.xlsx",
        eps=[1.0, 1.1, 1.2, 1.3, 1.4, 1.5, 0.2, 1.6, 1.7, 1.8],  # big drop year
        price=50.0,
        max_pe=15.0,
        terminal_price=80.0,
    )

    def fake_hint(self, period, company_facts):
        if period == "FY2022":
            return "RestructuringCharges=500000000"
        return None

    monkeypatch.setattr(
        ExpectedReturnValidationService,
        "_one_time_item_hint",
        fake_hint,
    )
    report = ExpectedReturnValidationService().validate(
        analysis_id="er2",
        ticker="TEST",
        workbook_path=path,
        company_facts={},
    )
    outs = [o for o in report.outliers if o.necessity == AdjustmentNecessity.OPTIONAL]
    assert outs
    assert outs[0].sec_evidence


def test_statistical_outlier_no_evidence_do_not_adjust(tmp_path: Path):
    path = _er_workbook(
        tmp_path / "stat.xlsx",
        eps=[1.0, 1.05, 1.1, 1.15, 1.2, 1.25, 2.5, 2.6, 2.7, 2.8],
        price=100.0,
        max_pe=20.0,
        terminal_price=150.0,
    )
    report = ExpectedReturnValidationService().validate(
        analysis_id="er3",
        ticker="TEST",
        workbook_path=path,
        company_facts=None,
    )
    assert report.outliers
    assert all(o.necessity == AdjustmentNecessity.NOT_JUSTIFIED for o in report.outliers)
    assert all(o.proposed_normalized_eps is None for o in report.outliers)


def test_growth_assumption_comparison(tmp_path: Path):
    eps = [2.0, 2.2, 2.4, 2.6, 2.8, 3.0, 3.3, 3.6, 4.0, 4.4]
    cagr = (4.4 / 2.0) ** (1 / 9) - 1
    path = _er_workbook(
        tmp_path / "growth.xlsx",
        eps=eps,
        price=120.0,
        max_pe=18.0,
        terminal_price=200.0,
        eps_growth=cagr,
    )
    report = ExpectedReturnValidationService().validate(
        analysis_id="er4",
        ticker="TEST",
        workbook_path=path,
    )
    assert report.growth.independent_eps_cagr == pytest.approx(cagr, rel=1e-4)
    assert report.growth.workbook_growth_assumption == pytest.approx(cagr, rel=1e-4)
    assert report.growth.growth_decision in {
        ExpectedReturnDecision.VALIDATED,
        ExpectedReturnDecision.WATCH,
    }


def test_expected_return_reconstruction():
    er = reconstruct_expected_annual_return(terminal_price=259.37, current_price=100.0, years=10)
    assert er == pytest.approx((259.37 / 100) ** 0.1 - 1)
    assert reconstruct_expected_annual_return(terminal_price=100, current_price=0) is None


def test_attractiveness_bands():
    assert interpret_attractiveness(0.15, treasury=0.04) == "attractive"
    assert interpret_attractiveness(0.08, treasury=0.04) == "acceptable"
    assert interpret_attractiveness(0.055, treasury=0.04) == "marginal"
    assert interpret_attractiveness(0.02, treasury=0.04) == "unattractive"


def test_no_workbook_writes(tmp_path: Path):
    path = _er_workbook(
        tmp_path / "nowrite.xlsx",
        eps=[1, 1.1, 1.2, 1.3, 1.4, 1.5, 1.6, 1.7, 1.8, 1.9],
        price=80.0,
        max_pe=16.0,
        terminal_price=120.0,
    )
    before = path.read_bytes()
    ExpectedReturnValidationService().validate(
        analysis_id="er5",
        ticker="TEST",
        workbook_path=path,
    )
    after = path.read_bytes()
    assert before == after
    # also mtime path via openpyxl round-trip safety
    wb = load_workbook(path)
    assert wb["Income - GAAP"]["C71"].value == 1
    wb.close()


def test_er_end_to_end_with_price(tmp_path: Path):
    path = _er_workbook(
        tmp_path / "er_ok.xlsx",
        eps=[2, 2.2, 2.4, 2.7, 3.0, 3.3, 3.6, 4.0, 4.4, 4.8],
        price=100.0,
        max_pe=20.0,
        terminal_price=200.0,
        workbook_er=(200 / 100) ** 0.1 - 1,
    )
    report = ExpectedReturnValidationService().validate(
        analysis_id="er6",
        ticker="TEST",
        workbook_path=path,
    )
    assert report.workbook_expected_return == pytest.approx((2.0) ** 0.1 - 1)
    assert report.independent_expected_return == pytest.approx((2.0) ** 0.1 - 1)
    assert report.expected_return_decision == ExpectedReturnDecision.VALIDATED
    review = ExpectedReturnValidationService().build_analyst_review(report)
    assert any(i.topic == "expected_return_output" for i in review.items)


def _er_workbook(
    path: Path,
    *,
    eps: list[float],
    price: float,
    max_pe: float,
    terminal_price: float,
    eps_growth: float | None = None,
    workbook_er: float | None = None,
) -> Path:
    assert len(eps) == 10
    wb = Workbook()
    wb.remove(wb.active)
    income = wb.create_sheet("Income - GAAP")
    bs = wb.create_sheet("Balance Sheet - Standardized")
    er = wb.create_sheet("Expected Returns & Buybacks")
    fm = wb.create_sheet("Final Metrics")
    inputs = wb.create_sheet("Inputs")

    for i, val in enumerate(eps):
        col = 3 + i
        year = 2016 + i
        income.cell(71, col).value = val
        bs.cell(8, col).value = __import__("datetime").datetime(year, 9, 30)

    if eps_growth is None:
        eps_growth = (eps[-1] / eps[0]) ** (1 / 9) - 1
    if workbook_er is None:
        workbook_er = (terminal_price / price) ** 0.1 - 1 if price > 0 else None

    er["A2"] = price
    er["B2"] = 0.045
    er["E2"] = max_pe
    er["B5"] = eps_growth
    er["B14"] = terminal_price / max_pe if max_pe else None
    er["C14"] = terminal_price
    er["E14"] = workbook_er
    er["A14"] = 0.2
    fm["L31"] = eps_growth
    fm["B51"] = price
    fm["B57"] = None
    inputs["B63"] = price
    inputs["B69"] = None

    wb.save(path)
    wb.close()
    return path
