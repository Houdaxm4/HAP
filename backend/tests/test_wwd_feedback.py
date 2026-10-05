"""WWD feedback: title rows stay blank, notes avoid hidden column B, cumulative sections are filled, projection columns are written."""

from __future__ import annotations

from datetime import datetime
from pathlib import Path

from openpyxl import Workbook, load_workbook

from services.hap_analysis_layout_service import HapAnalysisLayoutService
from services.ic_projection_writer import (
    LQ_BS,
    LQ_CF,
    LQ_IS,
    IcProjectionWriter,
    build_row_map,
    translate_to_last_quarter,
)
from services.new_company_cumulative_service import NewCompanyCumulativeService, sec_period_value
from services.new_company_statement_validation_service import NewCompanyStatementValidationService
from services.title_row_guard import clear_title_rows, is_title_row, real_row_for_label, title_rows

BSN = "Balance Sheet - Standardized"


def _statement(ws, rows: list[tuple[int, str, str | None]]) -> None:
    for row, label, code in rows:
        ws.cell(row, 1).value = label
        if code:
            ws.cell(row, 2).value = code


# ---------------------------------------------------------------- 1. title rows
def test_title_row_is_found_and_never_chosen():
    wb = Workbook()
    ws = wb.active
    ws.title = BSN
    _statement(ws, [(9, "Total Assets", None), (10, "  + Cash", "C&CE"), (61, "Total Assets", "BS_TOT_ASSET")])
    assert title_rows(ws) == {9}
    assert is_title_row(ws, 9) and not is_title_row(ws, 61)
    assert real_row_for_label(ws, "total assets") == 61
    assert NewCompanyStatementValidationService._find_row(ws, ("total assets",)) == 61


def test_clear_title_rows_removes_numbers_only_from_titles():
    wb = Workbook()
    ws = wb.active
    ws.title = "Cash Flow - Standardized"
    _statement(ws, [(9, "Cash from Operating Activities", None), (24, "Cash from Operating Activities", "CF_CASH_FROM_OPER")])
    ws["D9"], ws["E9"], ws["D24"] = 307.5, "=D24", 307.5
    ws["C9"] = datetime(2016, 9, 30)
    report = clear_title_rows(wb)
    assert report.cleared == ["Cash Flow - Standardized!D9"]
    assert ws["D9"].value is None and ws["E9"].value == "=D24" and ws["D24"].value == 307.5
    assert ws["C9"].value == datetime(2016, 9, 30)


# ---------------------------------------------------------------- 2. notes never in hidden column B
def test_notes_go_to_d_and_e_when_column_b_is_hidden():
    wb = Workbook()
    ws = wb.active
    ws.title = LQ_IS
    ws["A1"] = "Revenue"
    ws.column_dimensions["B"].hidden = True
    HapAnalysisLayoutService().write_notes_section(ws, [("Primary source", "workbook")])
    assert all(ws.cell(r, 2).value is None for r in range(1, ws.max_row + 1))
    assert any(ws.cell(r, 4).value == "Primary source" and ws.cell(r, 5).value == "workbook" for r in range(1, ws.max_row + 1))
    visible = Workbook().active
    visible["A1"] = "x"
    HapAnalysisLayoutService().write_notes_section(visible, [("Primary source", "workbook")])
    assert any(visible.cell(r, 2).value == "workbook" for r in range(1, visible.max_row + 1))


# ---------------------------------------------------------------- 3. cumulative sections
def _entry(val, start, end, form="10-Q"):
    return {"start": start, "end": end, "val": val, "form": form, "accn": "0001-26-000001", "filed": end}


def _facts() -> dict:
    nine_now, nine_prev = ("2025-10-01", "2026-06-30"), ("2024-10-01", "2025-06-30")
    gaap = {
        "RevenueFromContractWithCustomerExcludingAssessedTax": {"units": {"USD": [_entry(3_196_727_000, *nine_now), _entry(2_571_800_000, *nine_prev)]}},
        "CostOfGoodsAndServicesSold": {"units": {"USD": [_entry(2_238_752_000, *nine_now), _entry(1_892_908_000, *nine_prev)]}},
        "SellingGeneralAndAdministrativeExpense": {"units": {"USD": [_entry(300_000_000, *nine_now), _entry(250_000_000, *nine_prev)]}},
        "ResearchAndDevelopmentExpense": {"units": {"USD": [_entry(137_678_000, *nine_now), _entry(100_766_000, *nine_prev)]}},
        "NetIncomeLoss": {"units": {"USD": [_entry(414_407_000, *nine_now), _entry(304_488_000, *nine_prev)]}},
        "EarningsPerShareBasic": {"units": {"USD/shares": [_entry(6.95, *nine_now), _entry(5.12, *nine_prev)]}},
        "EarningsPerShareDiluted": {"units": {"USD/shares": [_entry(6.76, *nine_now), _entry(4.96, *nine_prev)]}},
        "NetCashProvidedByUsedInOperatingActivities": {"units": {"USD": [_entry(351_937_000, *nine_now), _entry(237_976_000, *nine_prev)]}},
        "NetCashProvidedByUsedInInvestingActivities": {"units": {"USD": [_entry(-286_795_000, *nine_now), _entry(-27_518_000, *nine_prev)]}},
        "NetCashProvidedByUsedInFinancingActivities": {"units": {"USD": [_entry(88_769_000, *nine_now), _entry(-26_126_000, *nine_prev)]}},
    }
    return {"facts": {"us-gaap": gaap}}


def _quarter_workbook(path: Path) -> Path:
    wb = Workbook()
    is_ws = wb.active
    is_ws.title = LQ_IS
    cf = wb.create_sheet(LQ_CF)
    for ws in (is_ws, cf):
        ws["C4"], ws["D4"] = datetime(2026, 6, 30), datetime(2025, 6, 30)
        ws.column_dimensions["B"].hidden = True   # as in the real Last Quarter tabs
    lines = [(11, "Revenue", "SALES_REV_TURN"), (12, "    + Sales & Services Rev", "IS_SALES_AND_SERVICES_REVENUES"),
             (16, "  - Cost of Revenue", "IS_COGS_TO_FE_AND_PP_AND_G"), (17, "    + Cost of Goods & Services", "IS_COG_AND_SERVICES_SOLD"),
             (21, "Gross Profit", "GROSS_PROFIT"), (32, "Operating Income (Loss)", "IS_OPER_INC"), (60, "Net Income, GAAP", "NET_INCOME"),
             (64, "Net Income Avail to Common, GAAP", "EARN_FOR_COMMON"), (67, "Basic EPS, GAAP", "IS_EPS"), (71, "Diluted EPS, GAAP", "IS_DILUTED_EPS")]
    _statement(is_ws, lines)
    is_ws["C11"] = 1109.71
    cf_lines = [(11, "Cash from Operating Activities", None), (13, "  + Net Income", "CF_NET_INC"), (27, "Cash from Operating Activities", "CF_CASH_FROM_OPER"),
                (49, "Cash from Investing Activities", "CF_CASH_FROM_INV_ACT"), (64, "Cash from Financing Activities", "CFF_ACTIVITIES_DETAILED"),
                (68, "Net Changes in Cash", "CF_NET_CHNG_CASH")]
    _statement(cf, cf_lines)
    cf["C27"] = 351.937   # already there: must be kept
    wb.save(path)
    return path


def test_sec_period_value_matches_end_date_and_length():
    facts = _facts()
    assert sec_period_value(facts, ("NetIncomeLoss",), datetime(2026, 6, 30).date(), "9m")[0] == 414.407
    assert sec_period_value(facts, ("NetIncomeLoss",), datetime(2026, 6, 30).date(), "6m") is None
    assert sec_period_value(facts, ("EarningsPerShareDiluted",), datetime(2025, 6, 30).date(), "9m", per_share=True)[0] == 4.96


def test_cumulative_income_statement_and_cash_flow_are_filled(tmp_path: Path):
    path = _quarter_workbook(tmp_path / "wb.xlsx")
    report = NewCompanyCumulativeService().apply(workbook_path=path, company_facts=_facts(), latest_quarter=3)
    wb = load_workbook(path)
    is_ws, cf = wb[LQ_IS], wb[LQ_CF]
    assert (is_ws["G11"].value, is_ws["H11"].value) == (3196.727, 2571.8)
    assert is_ws["G21"].value == round(3196.727 - 2238.752, 6)                       # gross profit derived
    assert round(is_ws["G32"].value, 3) == round(3196.727 - 2238.752 - 300.0 - 137.678, 3)   # operating income derived (no tag)
    assert (is_ws["G60"].value, is_ws["G64"].value, is_ws["G67"].value, is_ws["G71"].value) == (414.407, 414.407, 6.95, 6.76)
    assert (cf["D27"].value, cf["D49"].value, cf["D64"].value) == (237.976, -27.518, -26.126)
    assert (cf["C13"].value, cf["D13"].value) == (414.407, 304.488)
    assert cf["C27"].value == 351.937 and any("C27" in k for k in report.kept)       # existing value kept
    assert cf["C11"].value is None and cf["D11"].value is None                        # title row stays blank
    # margins: K/L built from the cumulative columns
    assert is_ws["K21"].value == "=G21/G11" and is_ws["L21"].value == "=H21/H11"
    assert is_ws["K32"].value == "=G32/G11" and is_ws["K64"].value == "=G64/G11"
    # notes in D/E (column B is where field codes live and is hidden in the real tab)
    notes = [str(is_ws.cell(r, 4).value) for r in range(1, is_ws.max_row + 1) if is_ws.cell(r, 4).value]
    assert "Notes" in notes and any("year-to-date columns of the income statement were filled" in n for n in notes)
    assert all(is_ws.cell(r, 2).value is None or str(is_ws.cell(r, 2).value).isupper() or "_" in str(is_ws.cell(r, 2).value) or "&" in str(is_ws.cell(r, 2).value) for r in range(60, is_ws.max_row + 1))
    ledger = wb["HAP Adjustments"]
    assert ledger["E5"].value == "Cumulative from 10-Q"


def test_full_year_means_nothing_to_fill(tmp_path: Path):
    path = _quarter_workbook(tmp_path / "wb.xlsx")
    assert not NewCompanyCumulativeService().apply(workbook_path=path, company_facts=_facts(), latest_quarter=4).changed


# ---------------------------------------------------------------- 4 to 7. projection columns
def test_row_map_matches_by_label_across_shifted_tabs():
    wb = Workbook()
    annual, quarter = wb.active, wb.create_sheet(LQ_BS)
    annual.title = BSN
    _statement(annual, [(10, "+ Cash & Cash Equivalents", None), (11, "check", None), (12, "+ Inventories", None), (13, "+ Prepaid Expenses", None)])
    _statement(quarter, [(11, "+ Cash & Cash Equivalents", None), (12, "check", None), (13, "+ Inventories", None), (14, "+ Prepaid expenses", None)])
    assert build_row_map(annual, quarter, {10, 12, 13}) == {10: 11, 12: 13, 13: 14}
    unmapped: list[int] = []
    out = translate_to_last_quarter(f"=IF('{BSN}'!L10=\"\",0,'{BSN}'!L10)+'{BSN}'!L99", {10: 11}, unmapped)
    assert "'Last Quarter BS Standardized'!C11" in out and unmapped == [99]


def _ic_workbook(path: Path, quarter_factor_cells: dict[str, float]) -> Path:
    wb = Workbook()
    ic = wb.active
    ic.title = "IC & NOPAT & ROIC "
    inputs = wb.create_sheet("Inputs")
    bs = wb.create_sheet(BSN)
    lq_bs = wb.create_sheet(LQ_BS)
    lq_is = wb.create_sheet(LQ_IS)
    lq_cf = wb.create_sheet(LQ_CF)
    wb.create_sheet("Final Metrics")["L7"] = 0.10
    _statement(bs, [(10, "+ Cash", None), (11, "+ Payables", None)])
    _statement(lq_bs, [(11, "+ Cash", None), (12, "+ Payables", None)])
    inputs["L81"] = f"=IF('{BSN}'!L10=\"\",0,'{BSN}'!L10)"
    inputs["L82"] = "=0"
    inputs["L84"] = f"=IF('{BSN}'!L11=\"\",0,'{BSN}'!L11)"
    inputs["L85"] = "=0"
    ic["L3"] = "=Inputs!L80"
    for row, value in {5: 20.0, 6: 500.0, 13: 400.0, 14: 0.0, 15: -1.5, 16: 140.0, 17: 130.0, 19: 100.0}.items():
        ic[f"L{row}"] = value
    lq_bs["C11"], lq_bs["C12"], lq_bs["C61"] = 1000.0, 300.0, 5000.0
    lq_is["G11"], lq_is["G32"] = quarter_factor_cells["revenue"], quarter_factor_cells["oi"]
    lq_cf["C27"] = 300.0
    wb.save(path)
    return path


def test_projection_columns_follow_the_house_method(tmp_path: Path):
    path = _ic_workbook(tmp_path / "wb.xlsx", {"revenue": 3000.0, "oi": 450.0})
    report = IcProjectionWriter().apply(workbook_path=path, latest_quarter=3)
    assert report.changed and not report.unmapped_rows
    ws = load_workbook(path)["IC & NOPAT & ROIC "]
    assert ws["M1"].value == "Projected ROIC" and ws["N1"].value == "ROCE"
    assert ws["M5"].value == "=L5" and ws["M6"].value == "=L6"                       # leases and R&D = previous year
    assert ws["M7"].value == "=M3-M4+M5+M6"
    assert ws["M11"].value == "='Last Quarter IS Standardized'!G11*4/3"              # Q3: x 4/3
    assert ws["M13"].value == "='Last Quarter IS Standardized'!G32*4/3"
    assert ws["M14"].value == "=L14" and ws["M15"].value == "=L15" and ws["M16"].value == "=L16" and ws["M17"].value == "=L17"
    assert ws["M18"].value == "=M13"
    assert ws["M19"].value == '=IF(OR(L13="",L13=0),"",M13/L13*L19)'                  # taxes scale with operating income
    assert ws["M20"].value == "=M13+M14-M15+M16-M17-M19"
    assert ws["M23"].value == "=M20/M7"
    assert ws["N4"].value == "='Last Quarter CF Standardized'!C27*4/3/'Last Quarter BS Standardized'!C61"
    assert "'Last Quarter BS Standardized'!C11" in ws["M3"].value and "'Last Quarter BS Standardized'!C12" in ws["M4"].value


def test_q2_uses_x2_and_full_year_writes_nothing(tmp_path: Path):
    path = _ic_workbook(tmp_path / "wb.xlsx", {"revenue": 2000.0, "oi": 300.0})
    IcProjectionWriter().apply(workbook_path=path, latest_quarter=2)
    ws = load_workbook(path)["IC & NOPAT & ROIC "]
    assert ws["M11"].value.endswith("G11*2") and ws["N4"].value.startswith("='Last Quarter CF Standardized'!C27*2/")
    again = IcProjectionWriter().apply(workbook_path=path, latest_quarter=2)
    assert not again.changed and any("already holds" in s for s in again.skipped)       # existing cells are not overwritten
    assert not IcProjectionWriter().apply(workbook_path=path, latest_quarter=4).changed


# ---------------------------------------------------------------- Yahoo fallback (quarters added up over the fiscal year)
class _FakeFetcher:
    def __init__(self, payload):
        self.payload = payload

    def get(self, url, accept=None):
        import json

        return 200, json.dumps(self.payload)


def _yahoo_payload(series: dict[str, dict[str, float]]) -> dict:
    result = []
    for kind, points in series.items():
        result.append({"meta": {"type": [kind]}, kind: [{"asOfDate": d, "reportedValue": {"raw": v * 1_000_000}} for d, v in points.items()]})
    return {"timeseries": {"result": result}}


def test_yahoo_sums_the_fiscal_quarters_when_revenue_agrees():
    from datetime import date

    from research.yahoo_fundamentals import YahooFallback

    quarters = {"2025-12-31": 1000.0, "2026-03-31": 1087.0, "2026-06-30": 1109.71}
    net = {"2025-12-31": 130.0, "2026-03-31": 137.0, "2026-06-30": 146.68}
    yahoo = YahooFallback(fetcher=_FakeFetcher(_yahoo_payload({"quarterlyTotalRevenue": quarters, "quarterlyNetIncome": net})))
    value, source = yahoo.ytd_value("WWD", "net_income", date(2026, 6, 30), 3, anchor_revenue=1109.71)
    assert round(value, 2) == 413.68 and "Yahoo" in source
    assert yahoo.ytd_value("WWD", "net_income", date(2026, 6, 30), 3, anchor_revenue=900.0) is None   # revenue disagrees: not used
    assert yahoo.ytd_value("WWD", "net_income", date(2026, 6, 30), 4, anchor_revenue=1109.71) is None  # quarter 4 is not a YTD
    gappy = {"2025-12-31": 1000.0, "2026-06-30": 1109.71}
    y2 = YahooFallback(fetcher=_FakeFetcher(_yahoo_payload({"quarterlyTotalRevenue": gappy, "quarterlyNetIncome": gappy})))
    assert y2.ytd_value("WWD", "net_income", date(2026, 6, 30), 3, anchor_revenue=1109.71) is None    # missing quarter: never guess


def test_cumulative_service_uses_yahoo_only_where_sec_has_nothing(tmp_path: Path):
    path = _quarter_workbook(tmp_path / "wb.xlsx")
    facts = _facts()
    del facts["facts"]["us-gaap"]["NetIncomeLoss"]
    quarters = {"2026-03-31": 1087.0, "2026-06-30": 1109.71, "2025-12-31": 1000.0}
    net = {"2025-12-31": 130.0, "2026-03-31": 137.0, "2026-06-30": 146.68}
    from research.yahoo_fundamentals import YahooFallback

    yahoo = YahooFallback(fetcher=_FakeFetcher(_yahoo_payload({"quarterlyTotalRevenue": quarters, "quarterlyNetIncome": net})))
    report = NewCompanyCumulativeService().apply(workbook_path=path, company_facts=facts, latest_quarter=3, ticker="WWD", yahoo=yahoo)
    ws = load_workbook(path)[LQ_IS]
    assert round(ws["G60"].value, 2) == 413.68                                    # from Yahoo (SEC lacked it)
    assert ws["G11"].value == 3196.727                                            # from SEC
    assert any(f.concept == "net_income" and f.method == "yahoo_quarterly_sum" for f in report.fills)
    assert ws["H60"].value is None                                                # prior year: Yahoo has no such history, nothing guessed


def test_zero_left_in_the_empty_year_to_date_column_is_replaced(tmp_path: Path):
    path = _quarter_workbook(tmp_path / "wb.xlsx")
    wb = load_workbook(path)
    wb[LQ_IS]["G25"] = 0          # the template leaves 0 in some year-to-date cells; that is not a reported figure
    wb[LQ_IS]["G25"].value = 0
    wb[LQ_IS].cell(25, 2).value = "IS_SG&A_EXPENSE"
    wb.save(path)
    NewCompanyCumulativeService().apply(workbook_path=path, company_facts=_facts(), latest_quarter=3)
    assert load_workbook(path)[LQ_IS]["G25"].value == 300.0


def test_quarterly_fill_never_writes_to_a_title_row(tmp_path: Path):
    from services.quarterly_presentation_service import QuarterlyPresentationService
    from models.quarterly_presentation import QuarterlyStatementKind

    wb = Workbook()
    ws = wb.active
    ws.title = LQ_CF
    ws["A11"] = "Cash from Operating Activities"                 # title row: no field code in column B
    ws["A27"], ws["B27"] = "Cash from Operating Activities", "CF_CASH_FROM_OPER"
    facts = {"facts": {"us-gaap": {"NetCashProvidedByUsedInOperatingActivities": {"units": {"USD": [
        {"start": "2025-10-01", "end": "2026-06-30", "val": 351_937_000, "form": "10-Q", "fp": "Q3", "fy": 2026, "accn": "0001", "filed": "2026-07-30"}]}}}}}
    from services.title_row_guard import title_rows

    assert title_rows(ws) == {11}
    QuarterlyPresentationService._fill_blanks_from_sec(wb, QuarterlyStatementKind.CASH_FLOW, LQ_CF, facts, 2026, "Q3", _NoDeps())
    assert ws["C11"].value is None


class _NoDeps:
    def feeds_metrics(self, *_args) -> bool:
        return False
