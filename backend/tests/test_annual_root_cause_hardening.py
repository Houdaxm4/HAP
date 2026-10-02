"""Targeted Annual Update root-cause regressions (circular refs, rollover, CRF, leases, R&D, HAP_ANALYSIS)."""

from __future__ import annotations

from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook
from openpyxl.comments import Comment

from models.annual_update import AnnualAnalystJudgmentReport, JudgmentRecord
from models.custom_run import CustomRunData, CustomRunPeriods
from models.write_actions import WriteActionClass
from services.annual_continuity_service import AnnualContinuityService
from services.annual_deliverables_service import AnnualDeliverablesService
from services.annual_inputs_service import AnnualInputsService
from services.annual_judgment_service import AnnualJudgmentService
from services.annual_leases_service import AnnualLeasesService
from services.annual_output_gate_service import AnnualOutputGateService
from services.annual_rd_service import AnnualRdService
from services.annual_statement_validation_service import AnnualStatementValidationService
from services.circular_reference_service import CircularReferenceService
from services.completion_scope import AnalysisTypeMode, WorkbookSection, sections_for_mode
from services.custom_run_service import CustomRunService, _collapse_to_fy_end, _collapse_to_fy_end_any
from services.excel_fill_service import ExcelFillService
from services.excel_recalc_service import ExcelRecalcReport, genuine_excel_com_recalc
from services.formula_utils import formula_would_self_reference
from services.hap_analysis_layout_service import HapAnalysisLayoutService, cell_is_occupied_or_formula
from services.statement_validation_service import StatementValidationService
from services.workbook_flag_service import STRUCTURAL_COMMENT, SUGGESTION_COMMENT
from workbook_mapping.engine import IntentDecision, WriteIntent, WriteIntentReport
from models.statement_validation import StatementValidationDecision


def _headers(ws, years: list[int], row: int = 7, start_col: int = 3) -> None:
    for i, y in enumerate(years):
        ws.cell(row, start_col + i, f"FY{y}")


def _intent(*, sheet: str, cell: str, value, decision=IntentDecision.WRITE) -> WriteIntent:
    return WriteIntent(
        intent_id=f"{sheet}:{cell}",
        mapping_id=f"{sheet}:{cell}",
        sheet=sheet,
        cell=cell,
        metric="test",
        period="FY2025",
        value=value,
        source="test",
        write_policy_result="writable",
        decision=decision,
        reason="test",
        cfm_path="inputs.pe10",
    )


def _inputs_pair(tmp: Path) -> tuple[Path, Path]:
    prev = tmp / "prev.xlsx"
    tmpl = tmp / "tmpl.xlsx"

    def book(path: Path, years: list[int], *, c121, d121) -> None:
        wb = Workbook()
        ws = wb.active
        ws.title = "Inputs"
        _headers(ws, years)
        ws["B115"] = "ok"
        ws["C121"] = c121
        ws["D121"] = d121
        ws["A121"] = "Compounded EPS Annualized Growth 10 Years"
        inc = wb.create_sheet("Income - GAAP")
        _headers(inc, years)
        inc["A11"] = "Revenue"
        for i, y in enumerate(years):
            inc.cell(11, 3 + i, 100 + i)
        wb.save(path)
        wb.close()

    book(prev, [2021, 2022, 2023, 2024], c121=10, d121='=IF(B115="","Error",C121-1)')
    book(tmpl, [2022, 2023, 2024, 2025], c121=None, d121=None)
    return prev, tmpl


# ---------------------------------------------------------------------------
# Circular references
# ---------------------------------------------------------------------------

def test_hap_cannot_introduce_circular_reference_via_continuity(tmp_path: Path):
    prev, tmpl = _inputs_pair(tmp_path)
    out = tmp_path / "out.xlsx"
    out.write_bytes(tmpl.read_bytes())
    AnnualContinuityService().apply(
        analysis_id="a1",
        ticker="ZZ",
        template_path=tmpl,
        previous_workbook_path=prev,
        workbook_path=out,
    )
    wb = load_workbook(out)
    c121 = wb["Inputs"]["C121"].value
    wb.close()
    if isinstance(c121, str) and c121.startswith("="):
        assert not formula_would_self_reference(c121, 3, 121, sheet="Inputs")
    report = CircularReferenceService().classify_against_source(
        analysis_id="a1", ticker="ZZ", current_path=out, source_path=prev, template_path=tmpl
    )
    assert report.hap_introduced == []
    assert report.status != "BLOCKING_STRUCTURAL_ERROR"


def test_proposed_formula_write_that_cycles_is_rejected(tmp_path: Path):
    src = tmp_path / "src.xlsx"
    dest = tmp_path / "dest.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "Inputs"
    ws["C121"] = None
    wb.save(src)
    wb.close()
    report = WriteIntentReport(
        analysis_id="a1",
        ticker="ZZ",
        intents=[_intent(sheet="Inputs", cell="C121", value="=C121-1")],
    )
    ExcelFillService().apply_write_intents(
        analysis_id="a1",
        ticker="ZZ",
        source_workbook_path=src,
        destination_workbook_path=dest,
        intent_report=report,
    )
    out = load_workbook(dest)
    assert out["Inputs"]["C121"].value in (None, "")
    out.close()
    cycles = CircularReferenceService().scan_workbook(dest)
    assert cycles == []


def test_preexisting_circular_reference_is_flagged_not_repaired(tmp_path: Path):
    prev = tmp_path / "prev.xlsx"
    tmpl = tmp_path / "tmpl.xlsx"
    out = tmp_path / "out.xlsx"
    for path, years, c121 in (
        (prev, [2021, 2022, 2023, 2024], 10),
        (tmpl, [2022, 2023, 2024, 2025], "=C121"),
    ):
        wb = Workbook()
        ws = wb.active
        ws.title = "Inputs"
        _headers(ws, years)
        ws["C121"] = c121
        wb.save(path)
        wb.close()
    out.write_bytes(tmpl.read_bytes())
    AnnualContinuityService().apply(
        analysis_id="a1",
        ticker="ZZ",
        template_path=tmpl,
        previous_workbook_path=prev,
        workbook_path=out,
    )
    wb = load_workbook(out)
    assert wb["Inputs"]["C121"].value == "=C121"
    wb.close()
    classified = CircularReferenceService().classify_against_source(
        analysis_id="a1", ticker="ZZ", current_path=out, source_path=prev, template_path=tmpl
    )
    assert classified.hap_introduced == []
    assert classified.template_native or classified.pre_existing
    from openpyxl import load_workbook as lw
    from services.workbook_flag_service import flag_structural

    flagged = tmp_path / "flagged.xlsx"
    flagged.write_bytes(out.read_bytes())
    fwb = lw(flagged)
    cycles = classified.pre_existing or classified.template_native
    flag_structural(fwb["Inputs"], "C121", cycle=cycles[0].cells)
    fwb.save(flagged)
    fwb.close()
    fwb = lw(flagged)
    assert STRUCTURAL_COMMENT in (fwb["Inputs"]["C121"].comment.text or "")
    assert fwb["Inputs"]["C121"].value == "=C121"
    fwb.close()


def test_c121_regression_uses_generic_circular_logic_not_a_c121_hack(tmp_path: Path):
    prev, tmpl = _inputs_pair(tmp_path)
    out = tmp_path / "out.xlsx"
    out.write_bytes(tmpl.read_bytes())
    AnnualContinuityService().apply(
        analysis_id="a1",
        ticker="ZZ",
        template_path=tmpl,
        previous_workbook_path=prev,
        workbook_path=out,
    )
    wb = load_workbook(out)
    # Helper stays at D121 (same address); it is not FY-relocated onto C121.
    assert wb["Inputs"]["D121"].value == '=IF(B115="","Error",C121-1)'
    assert wb["Inputs"]["C121"].value != '=IF(B115="","Error",C121-1)'
    wb.close()
    wb = Workbook()
    ws = wb.active
    ws.title = "Inputs"
    cycle = CircularReferenceService().write_would_create_cycle(
        wb, sheet="Inputs", cell="Z99", formula="=Z99+1"
    )
    assert cycle is not None
    assert "Z99" in str(cycle.cells)


# ---------------------------------------------------------------------------
# Rollover
# ---------------------------------------------------------------------------

def test_rollover_preserves_y2_y10_adds_y11_drops_y1(tmp_path: Path):
    prev = tmp_path / "prev.xlsx"
    tmpl = tmp_path / "tmpl.xlsx"
    out = tmp_path / "out.xlsx"
    prev_years = list(range(2016, 2026))
    new_years = list(range(2017, 2027))

    def book(path: Path, years: list[int], *, hist: bool) -> None:
        wb = Workbook()
        inc = wb.active
        inc.title = "Income - GAAP"
        _headers(inc, years)
        inc["A11"] = "Revenue"
        inc["A12"] = "Operating income"
        for i, y in enumerate(years):
            if hist or y < years[-1]:
                inc.cell(11, 3 + i, y)
            inc.cell(12, 3 + i, f"={chr(67 + i)}11*0.2")
        wb.save(path)
        wb.close()

    book(prev, prev_years, hist=True)
    book(tmpl, new_years, hist=False)
    out.write_bytes(tmpl.read_bytes())
    report = AnnualContinuityService().apply(
        analysis_id="a1",
        ticker="ZZ",
        template_path=tmpl,
        previous_workbook_path=prev,
        workbook_path=out,
    )
    wb = load_workbook(out)
    ws = wb["Income - GAAP"]
    assert ws.cell(7, 3).value == "FY2017"
    assert ws.cell(7, 12).value == "FY2026"
    assert ws.cell(11, 3).value == 2017  # old Y2 preserved
    assert ws.cell(11, 11).value == 2025  # old Y10 preserved
    assert ws.cell(11, 12).value in (None, "")  # Y11 not taken from previous 2025 window
    assert str(ws.cell(12, 3).value).startswith("=")
    wb.close()
    assert report.new_fiscal_year == "FY2026"


def test_analyst_override_survives_rollover(tmp_path: Path):
    prev = tmp_path / "prev.xlsx"
    tmpl = tmp_path / "tmpl.xlsx"
    out = tmp_path / "out.xlsx"

    def book(path: Path, years: list[int], rates: list) -> None:
        wb = Workbook()
        ls = wb.active
        ls.title = "Leases"
        _headers(ls, years)
        for i, r in enumerate(rates):
            ls.cell(18, 3 + i, r)
        wb.save(path)
        wb.close()

    book(prev, [2021, 2022, 2023, 2024], [0.04, 0.04, 0.04, 0.04])
    book(tmpl, [2022, 2023, 2024, 2025], [None, None, None, 0.12])
    out.write_bytes(tmpl.read_bytes())
    AnnualContinuityService().apply(
        analysis_id="a1",
        ticker="ZZ",
        template_path=tmpl,
        previous_workbook_path=prev,
        workbook_path=out,
    )
    wb = load_workbook(out)
    assert wb["Leases"]["C18"].value == pytest.approx(0.04)
    assert wb["Leases"]["E18"].value == pytest.approx(0.04)
    wb.close()


# ---------------------------------------------------------------------------
# Financial statements
# ---------------------------------------------------------------------------

def test_historical_fs_discrepancy_flagged_not_overwritten(tmp_path: Path):
    from tests.test_statement_analyst_review import _facts_revenue

    path = tmp_path / "wb.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "Income - GAAP"
    for i, col in enumerate(list("CDEFGHIJKL")):
        ws[f"{col}7"] = f"FY{2016 + i}"
    ws["A9"] = "Revenue"
    ws["E9"] = 215639
    ws["F9"] = 260174
    for name in ("Balance Sheet - Standardized", "Cash Flow - Standardized"):
        wb.create_sheet(name)
    wb.save(path)
    wb.close()
    original = 215639
    report = StatementValidationService().validate(
        analysis_id="v1",
        ticker="AAPL",
        workbook_path=path,
        company_facts=_facts_revenue(),
        annotate_workbook=True,
    )
    e2018 = next(e for e in report.entries if e.metric == "Revenue" and e.fiscal_period == "FY2018")
    assert e2018.decision == StatementValidationDecision.DISCREPANCY
    wb = load_workbook(path)
    assert wb["Income - GAAP"]["E9"].value == original
    comment = wb["Income - GAAP"]["E9"].comment
    assert comment and "FINANCIAL-STATEMENT DISCREPANCY" in comment.text
    assert "Value was NOT automatically changed" in comment.text
    wb.close()


def test_new_fy_discrepancy_not_rewritten(tmp_path: Path):
    path = tmp_path / "wb.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "Income - GAAP"
    _headers(ws, [2022, 2023, 2024, 2025])
    ws["A11"] = "Revenue"
    ws["F11"] = 10.0
    wb.create_sheet("Balance Sheet - Standardized")
    wb.create_sheet("Cash Flow - Standardized")
    wb.save(path)
    wb.close()
    AnnualStatementValidationService().validate(
        analysis_id="a1",
        ticker="ZZ",
        workbook_path=path,
        fiscal_year="FY2025",
        filing_overrides={"revenue": 999.0},
    )
    wb = load_workbook(path)
    assert wb["Income - GAAP"]["F11"].value == pytest.approx(10.0)
    assert wb["Income - GAAP"]["F11"].comment is not None
    wb.close()


def test_formula_cells_never_overwritten_by_fill(tmp_path: Path):
    src = tmp_path / "src.xlsx"
    dest = tmp_path / "dest.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "Income - GAAP"
    ws["C12"] = "=C11*0.2"
    wb.save(src)
    wb.close()
    ExcelFillService().apply_write_intents(
        analysis_id="a1",
        ticker="ZZ",
        source_workbook_path=src,
        destination_workbook_path=dest,
        intent_report=WriteIntentReport(
            analysis_id="a1",
            ticker="ZZ",
            intents=[_intent(sheet="Income - GAAP", cell="C12", value=99.0)],
        ),
    )
    out = load_workbook(dest)
    assert out["Income - GAAP"]["C12"].value == "=C11*0.2"
    out.close()


# ---------------------------------------------------------------------------
# CRF inputs
# ---------------------------------------------------------------------------

def test_five_crf_metrics_populate_newest_fy(tmp_path: Path, monkeypatch):
    path = tmp_path / "wb.xlsx"
    crf = tmp_path / "crf.xlsx"
    crf.write_bytes(b"pk")
    wb = Workbook()
    ws = wb.active
    ws.title = "Inputs"
    years = list(range(2017, 2027))
    _headers(ws, years)
    ws["A57"] = "PE10"
    ws["A58"] = "E10"
    ws["A59"] = "EPS 10-Year Growth"
    ws["A60"] = "EPS 10-Year Direction"
    ws["A61"] = "Revenue 10-Year Growth"
    ws["A65"] = "Current PE10"
    ws["B65"] = None
    wb.save(path)
    wb.close()
    token = "FY2026"
    fake = CustomRunData(
        source_filename="crf.xlsx",
        ticker="ZZ",
        ticker_sheet_name="ZZ",
        metadata={
            "inputs_annual_pe10": {token: 14.2},
            "inputs_annual_e10": {token: 5.1},
            "inputs_annual_eps_10y_growth": {token: 0.083},
            "inputs_annual_eps_10y_direction": {token: "Increasing"},
            "inputs_annual_revenue_10y_growth": {token: 0.061},
        },
        scalars={"Current PE10": 13.9},
        periods=CustomRunPeriods(fiscal_years=[token]),
    )
    monkeypatch.setattr(CustomRunService, "parse", lambda self, *a, **k: fake)
    monkeypatch.setattr(AnnualInputsService, "_scan_crf_scalar", staticmethod(lambda *a, **k: None))
    monkeypatch.setattr(AnnualInputsService, "_scan_crf_fy_series", staticmethod(lambda *a, **k: None))
    from models.quarterly_update import CurrentDataRefreshReport

    monkeypatch.setattr(
        "services.annual_inputs_service.CurrentDataRefreshService.apply",
        lambda self, **k: CurrentDataRefreshReport(analysis_id="a1", ticker="ZZ"),
    )
    report = AnnualInputsService().apply(
        analysis_id="a1",
        ticker="ZZ",
        workbook_path=path,
        custom_run_path=crf,
        new_fiscal_year=token,
    )
    wb = load_workbook(path)
    assert wb["Inputs"]["L57"].value == pytest.approx(14.2)
    assert wb["Inputs"]["L58"].value == pytest.approx(5.1)
    assert wb["Inputs"]["L59"].value == pytest.approx(0.083)
    assert wb["Inputs"]["L60"].value == "Increasing"
    assert wb["Inputs"]["L61"].value == pytest.approx(0.061)
    wb.close()
    assert report.pe10 and report.pe10.fiscal_year == token
    assert report.e10 and report.e10.source_value == pytest.approx(5.1)
    assert report.eps_10y_growth and report.eps_10y_growth.source_value == pytest.approx(0.083)
    assert report.eps_10y_direction and report.eps_10y_direction.source_value == "Increasing"
    assert report.revenue_10y_growth and report.revenue_10y_growth.fiscal_year == token
    assert "workbook-history substitute" in (report.eps_10y_growth.transformation or "")


def test_crf_fy_end_collapse_matches_pe10_rule():
    years = ["FY2024"] * 4 + ["FY2025"] * 4
    values = [1, 2, 3, 4, 10, 11, 12, 13]
    annual = _collapse_to_fy_end(years, values)
    assert annual["FY2024"] == 4
    assert annual["FY2025"] == 13
    direction = _collapse_to_fy_end_any(years, ["up"] * 4 + ["down"] * 3 + ["Increasing"])
    assert direction["FY2025"] == "Increasing"


# ---------------------------------------------------------------------------
# Leases
# ---------------------------------------------------------------------------

def test_reproducible_lease_methodology_is_carried(tmp_path: Path):
    path = tmp_path / "wb.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "Leases"
    _headers(ws, [2021, 2022, 2023, 2024, 2025])
    for i, r in enumerate([0.04, 0.041, 0.039, 0.04, 0.12]):
        ws.cell(18, 3 + i, r)
    wb.save(path)
    wb.close()
    report = AnnualLeasesService().apply(
        analysis_id="a1", ticker="ZZ", workbook_path=path, fiscal_year="FY2025"
    )
    wb = load_workbook(path)
    assert wb["Leases"]["G18"].value == pytest.approx(0.04)
    wb.close()
    assert report.methodology_carried is True
    assert report.normalized_to_prior is True


def test_ambiguous_lease_adjustment_is_not_guessed(tmp_path: Path):
    path = tmp_path / "wb.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "Leases"
    _headers(ws, [2021, 2022, 2023, 2024, 2025])
    for i, r in enumerate([0.04, 0.07, 0.02, 0.11, 0.12]):
        ws.cell(18, 3 + i, r)
    wb.save(path)
    wb.close()
    report = AnnualLeasesService().apply(
        analysis_id="a1", ticker="ZZ", workbook_path=path, fiscal_year="FY2025"
    )
    wb = load_workbook(path)
    assert wb["Leases"]["G18"].value == pytest.approx(0.12)
    comment = wb["Leases"]["G18"].comment
    if report.suggestion_only:
        assert comment and SUGGESTION_COMMENT in comment.text
    wb.close()
    assert report.methodology_carried is False
    assert report.normalized_to_prior is False


def test_lease_suggestion_does_not_overwrite(tmp_path: Path):
    path = tmp_path / "wb.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "Leases"
    _headers(ws, [2021, 2022, 2023, 2024, 2025])
    for i, r in enumerate([0.04, 0.041, 0.039, "=F17", 0.12]):
        ws.cell(18, 3 + i, r)
    wb.save(path)
    wb.close()
    report = AnnualLeasesService().apply(
        analysis_id="a1", ticker="ZZ", workbook_path=path, fiscal_year="FY2025"
    )
    wb = load_workbook(path)
    if report.suggestion_only:
        assert SUGGESTION_COMMENT in (wb["Leases"]["G18"].comment.text or "")
        assert wb["Leases"]["G18"].value == pytest.approx(0.12)
    wb.close()


# ---------------------------------------------------------------------------
# R&D
# ---------------------------------------------------------------------------

def test_rd_life_preserved_and_prewindow_history_copied(tmp_path: Path):
    prev = tmp_path / "prev.xlsx"
    out = tmp_path / "out.xlsx"
    years_prev = list(range(2016, 2026))
    years_new = list(range(2017, 2027))

    def book(path: Path, years: list[int], *, life: float, prewindow: list[float] | None) -> None:
        wb = Workbook()
        inp = wb.active
        inp.title = "Inputs"
        _headers(inp, years)
        inp["A103"] = "R&D Expense"
        for i, y in enumerate(years):
            if y != years[-1]:
                inp.cell(103, 3 + i, 8 + i)
        rd = wb.create_sheet("R&D")
        rd["B8"] = life
        rd["A2"] = "R&D expense"
        if prewindow:
            for col, val in enumerate(prewindow, start=2):
                rd.cell(1, col, f"=RIGHT(\"FY{2010 + col}\")")
                rd.cell(2, col, val)
        else:
            for col in range(2, 5):
                rd.cell(1, col, f"=RIGHT(\"pre{col}\")")
                rd.cell(2, col, '=IF(Inputs!C103="",0,Inputs!C103)')
        wb.save(path)
        wb.close()

    book(prev, years_prev, life=5, prewindow=[1.0, 1.0, 0.979])
    book(out, years_new, life=99, prewindow=None)
    report = AnnualRdService().apply(
        analysis_id="a1",
        ticker="ZZ",
        workbook_path=out,
        previous_workbook_path=prev,
        fiscal_year="FY2026",
        filing_rd=12.0,
    )
    wb = load_workbook(out)
    assert wb["R&D"]["B8"].value == 5
    assert wb["R&D"]["B2"].value in (1.0, 0.979, 1)
    assert not (isinstance(wb["R&D"]["B2"].value, str) and ",0," in wb["R&D"]["B2"].value)
    wb.close()
    assert report.useful_life_unchanged is True
    assert report.useful_life == 5


def test_missing_required_rd_history_is_flagged_not_zeroed(tmp_path: Path):
    prev = tmp_path / "prev.xlsx"
    out = tmp_path / "out.xlsx"
    for path in (prev, out):
        wb = Workbook()
        inp = wb.active
        inp.title = "Inputs"
        _headers(inp, list(range(2017, 2027)))
        rd = wb.create_sheet("R&D")
        rd["B8"] = 5
        rd["A2"] = "R&D expense"
        for col in range(2, 5):
            rd.cell(1, col, f"=RIGHT(\"pre{col}\")")
            rd.cell(2, col, None)
        wb.save(path)
        wb.close()
    report = AnnualRdService().apply(
        analysis_id="a1",
        ticker="ZZ",
        workbook_path=out,
        previous_workbook_path=prev,
        fiscal_year="FY2026",
    )
    assert report.lookback_complete is False or report.missing_required_history
    wb = load_workbook(out)
    assert wb["R&D"]["B2"].value not in (0, 0.0)
    wb.close()


# ---------------------------------------------------------------------------
# Expected Returns / EV / Graham — analysis, not population
# ---------------------------------------------------------------------------

def test_er_ev_not_source_fill_targets():
    sections = sections_for_mode(AnalysisTypeMode.ANNUAL_UPDATE)
    assert WorkbookSection.EXPECTED_RETURN not in sections
    assert WorkbookSection.VALUATION not in sections


def _valuation_book(path: Path, *, b5=0.06, d17=2.0, b6=0.08, e14=0.11) -> None:
    wb = Workbook()
    er = wb.active
    er.title = "Expected Returns & Buybacks"
    er["B5"] = b5
    er["E14"] = e14
    for r in range(17, 27):
        er[f"D{r}"] = d17 + (r - 17) * 0.1
    er["D17"] = f"=A14*(1+B5)"
    ev = wb.create_sheet("Enterprise Value")
    ev["B6"] = b6
    ev["C6"] = b6
    ev["B20"] = 50
    ev["B27"] = 0.2
    ev["B32"] = 40
    wb.save(path)
    wb.close()


def test_reasonable_assumptions_not_overwritten_no_unnecessary_hap_rate(tmp_path: Path):
    path = tmp_path / "wb.xlsx"
    _valuation_book(path, b5=0.06, b6=0.08)
    er, judge = AnnualJudgmentService().apply(
        analysis_id="a1",
        ticker="ZZ",
        workbook_path=path,
        context={
            "book_value_growth": 0.06,
            "expected_return": 0.11,
            "owner_earnings_growth": 0.08,
            "eps_growth": 0.07,
        },
    )
    wb = load_workbook(path)
    assert wb["Expected Returns & Buybacks"]["B5"].value == pytest.approx(0.06)
    assert str(wb["Expected Returns & Buybacks"]["D17"].value).startswith("=")
    assert wb["Enterprise Value"]["B6"].value == pytest.approx(0.08)
    assert wb["Enterprise Value"]["B20"].value == 50
    assert wb["Enterprise Value"]["B27"].value == pytest.approx(0.2)
    wb.close()
    assert er.selected_methodology == "BOOK_VALUE_GROWTH"
    assert judge.expected_return and judge.expected_return.adjusted is False
    assert judge.owner_earnings_growth and judge.owner_earnings_growth.adjusted is False
    assert "not automatically" in (judge.graham_eps_growth.rationale or "").lower() or (
        judge.graham_eps_growth and judge.graham_eps_growth.adjusted is False
    )


def test_unrealistic_assumptions_write_hap_analysis_beside_originals(tmp_path: Path):
    path = tmp_path / "wb.xlsx"
    _valuation_book(path, b5=0.80, b6=0.70, d17=1.0)
    orig_d17 = None
    wb = load_workbook(path)
    orig_d17 = wb["Expected Returns & Buybacks"]["D17"].value
    orig_e14 = wb["Expected Returns & Buybacks"]["E14"].value
    wb.close()
    er, judge = AnnualJudgmentService().apply(
        analysis_id="a1",
        ticker="ZZ",
        workbook_path=path,
        context={
            "book_value_growth": 0.80,
            "expected_return": 0.11,
            "eps_growth": 0.08,
            "prefer_eps": True,
            "owner_earnings_growth": 0.70,
            "eps_distorted": True,
            "normalized_eps_growth": 0.09,
            "normalized_oe_growth": 0.08,
            "evidence": ["depressed starting EPS; one-time items"],
        },
    )
    wb = load_workbook(path)
    assert wb["Expected Returns & Buybacks"]["B5"].value == pytest.approx(0.80)
    assert wb["Expected Returns & Buybacks"]["D17"].value == orig_d17
    assert wb["Expected Returns & Buybacks"]["E14"].value == orig_e14
    assert wb["Enterprise Value"]["B6"].value == pytest.approx(0.70)
    hap_text = " ".join(
        str(c.value)
        for row in wb["Expected Returns & Buybacks"].iter_rows()
        for c in row
        if c.value
    )
    assert "HAP ANALYSIS — NOTES" in hap_text
    assert "HAP" in hap_text
    wb.close()
    assert judge.hap_analysis_cells
    assert er.selected_methodology == "EPS_GROWTH"
    assert judge.owner_earnings_growth.adjusted
    assert judge.graham_eps_growth.adjusted
    assert "1." in (judge.expected_return.rationale or "")
    assert judge.expected_return.original_value == pytest.approx(0.80)


def test_hap_analysis_cannot_overwrite_occupied_or_formula_cells(tmp_path: Path):
    path = tmp_path / "wb.xlsx"
    _valuation_book(path)
    wb = load_workbook(path)
    ws = wb["Expected Returns & Buybacks"]
    ws["G1"] = "occupied"
    ws["H5"] = "=B5"
    wb.save(path)
    wb.close()
    AnnualJudgmentService().apply(
        analysis_id="a1",
        ticker="ZZ",
        workbook_path=path,
        context={"book_value_growth": 0.80, "prefer_eps": True, "eps_growth": 0.08, "owner_earnings_growth": 0.08},
    )
    wb = load_workbook(path)
    assert wb["Expected Returns & Buybacks"]["G1"].value == "occupied"
    assert wb["Expected Returns & Buybacks"]["H5"].value == "=B5"
    assert wb["Expected Returns & Buybacks"]["B5"].value == pytest.approx(0.06)
    wb.close()
    ws = Workbook().active
    ws["A1"] = "=B1"
    assert cell_is_occupied_or_formula(ws, "A1") is True
    layout = HapAnalysisLayoutService()
    col, _ = layout.allocate(ws, start_col=1)
    assert col > 1


def test_output_gate_blocks_hap_introduced_circulars():
    from services.circular_reference_service import CircularCycle, CircularReferenceReport

    circular = CircularReferenceReport(
        analysis_id="a1",
        ticker="ZZ",
        hap_introduced=[CircularCycle(cells=["Inputs!C121"], origin="hap_introduced")],
        status="BLOCKING_STRUCTURAL_ERROR",
    )
    gate = AnnualOutputGateService().evaluate(
        analysis_id="a1",
        ticker="ZZ",
        workbook_path=Path("missing.xlsx"),
        fiscal_year="FY2026",
        tax=None,
        inputs=None,
        rd=None,
        valuation=None,
        recalc=None,
        circular=circular,
    )
    assert "HAP_INTRODUCED_CIRCULAR_REFERENCE" in gate.blockers
    assert gate.gates.get("circular_references") == "fail"


def test_stale_cache_is_not_certified_output():
    fake = ExcelRecalcReport(
        analysis_id="a1",
        ticker="ZZ",
        status="ok",
        method="openpyxl_data_only",
        workbook_path="x.xlsx",
        com_invoked=False,
    )
    assert genuine_excel_com_recalc(fake) is False
    real = ExcelRecalcReport(
        analysis_id="a1",
        ticker="ZZ",
        status="ok",
        method="excel_com_calculate_full_rebuild",
        workbook_path="x.xlsx",
        com_invoked=True,
    )
    assert genuine_excel_com_recalc(real) is True


def test_word_distinguishes_original_from_hap(tmp_path: Path):
    path = tmp_path / "wb.xlsx"
    _valuation_book(path)
    judgment = AnnualAnalystJudgmentReport(
        analysis_id="a1",
        ticker="ZZ",
        owner_earnings_growth=JudgmentRecord(
            metric="owner_earnings_growth",
            original_value=0.08,
            selected_value=0.08,
            adjusted=False,
            rationale="Existing OE growth appears reasonable prospectively.",
        ),
        graham_eps_growth=JudgmentRecord(
            metric="graham_eps_growth",
            original_value=0.07,
            selected_value=0.07,
            adjusted=False,
            rationale="Historical EPS 10-year growth was not automatically used as the prospective Graham rate.",
        ),
        hap_analysis_cells=["Expected Returns & Buybacks!G2"],
    )
    from models.annual_update import AnnualExpectedReturnReport, AnnualResearchReport

    er = AnnualExpectedReturnReport(
        analysis_id="a1",
        ticker="ZZ",
        original_growth_rate=0.06,
        original_expected_return=0.11,
        selected_growth_rate=0.06,
        final_expected_return=None,
        reasonableness="reasonable",
        rationale="Existing assumption appears reasonable; no adjustment is recommended.",
        selected_methodology="BOOK_VALUE_GROWTH",
    )
    research = AnnualResearchReport(analysis_id="a1", ticker="ZZ", fiscal_year=2026)
    perf = AnnualDeliverablesService().build_performance(
        analysis_id="a1", ticker="ZZ", fiscal_year=2026, workbook_path=path, research=research
    )
    report = AnnualDeliverablesService().produce(
        analysis_id="a1",
        ticker="ZZ",
        fiscal_year=2026,
        completed_workbook_path=path,
        output_dir=tmp_path,
        performance=perf,
        research=research,
        judgment=judgment,
        expected_return=er,
    )
    from docx import Document

    text = "\n".join(p.text for p in Document(tmp_path / report.word_filename).paragraphs)
    assert "ORIGINAL WORKBOOK RESULT" in text
    assert "HAP-ADJUSTED ANALYSIS" in text
    assert "not automatically" in text.lower() or "not required" in text.lower()


def test_write_action_classes_exist():
    assert WriteActionClass.SOURCE_FILL.value == "SOURCE_FILL"
    assert WriteActionClass.HAP_ANALYSIS.value == "HAP_ANALYSIS"
    assert WriteActionClass.BLOCKING_STRUCTURAL_ERROR.value == "BLOCKING_STRUCTURAL_ERROR"
    assert WriteActionClass.DISCREPANCY.value == "DISCREPANCY"
