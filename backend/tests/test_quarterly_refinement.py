"""Refinement pass: model continuity, ignored sheets, Yahoo fallback, Word research."""

from __future__ import annotations

from copy import copy
from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook
from openpyxl.styles import PatternFill

from models.quarterly_update import CarryForwardDecision
from services.quarterly_model_continuity_service import QuarterlyModelContinuityService
from services.quarterly_research_service import QuarterlyResearchService
from workbook_mapping.sheet_policies import IGNORED_TEMPLATE_SHEETS


def _wb_with_ignored_and_persistent(tmp_path: Path) -> tuple[Path, Path, Path]:
    """template, previous, carried(starts as template copy)."""
    template = tmp_path / "template.xlsx"
    previous = tmp_path / "previous.xlsx"
    carried = tmp_path / "carried.xlsx"

    for path in (template, previous):
        wb = Workbook()
        wb.active.title = "Inputs"
        wb["Inputs"]["C57"] = 24.0 if path == previous else None
        wb["Inputs"]["B63"] = 100.0 if path == previous else None
        # Ignored sheet
        ig = wb.create_sheet("Last Quarter IS As Reported")
        ig["A1"] = "Template As Reported"
        ig["B2"] = 999
        if path == previous:
            ig["B2"] = 888  # must NOT propagate
        # Leases manual override
        ls = wb.create_sheet("Leases")
        ls["J18"] = 0.04 if path == previous else None
        ls["B18"] = "=B16+1" if path == previous else "=B16+1"
        # Formatting on previous
        if path == previous:
            wb["Inputs"]["C57"].fill = PatternFill("solid", fgColor="FFFF00")
        wb.save(path)
        wb.close()

    import shutil

    shutil.copy2(template, carried)
    return template, previous, carried


def test_ignored_sheets_match_template(tmp_path: Path):
    template, previous, carried = _wb_with_ignored_and_persistent(tmp_path)
    QuarterlyModelContinuityService().apply_persistent_continuity(
        analysis_id="mc1",
        ticker="AAPL",
        new_template_path=template,
        previous_workbook_path=previous,
        workbook_path=carried,
    )
    report = QuarterlyModelContinuityService().verify_final_deliverable(
        analysis_id="mc1f",
        ticker="AAPL",
        new_template_path=template,
        workbook_path=carried,
    )
    wb = load_workbook(carried)
    try:
        assert wb["Last Quarter IS As Reported"]["B2"].value == 999
        assert "Last Quarter IS As Reported" in report.ignored_sheets_verified
    finally:
        wb.close()


def test_model_continuity_carries_value_and_format(tmp_path: Path):
    template, previous, carried = _wb_with_ignored_and_persistent(tmp_path)
    QuarterlyModelContinuityService().apply_persistent_continuity(
        analysis_id="mc2",
        ticker="AAPL",
        new_template_path=template,
        previous_workbook_path=previous,
        workbook_path=carried,
    )
    wb = load_workbook(carried)
    try:
        assert wb["Inputs"]["C57"].value == 24.0
        assert wb["Inputs"]["B63"].value is None  # refresh region excluded
        assert wb["Leases"]["J18"].value == 0.04
        fill = wb["Inputs"]["C57"].fill.fgColor.rgb
        assert fill is not None
    finally:
        wb.close()


def test_restore_formula_from_previous(tmp_path: Path):
    template = tmp_path / "t.xlsx"
    previous = tmp_path / "p.xlsx"
    carried = tmp_path / "c.xlsx"
    for path, formula in ((template, None), (previous, "=Inputs!L87")):
        wb = Workbook()
        wb.active.title = "Tax"
        wb["Tax"]["L13"] = formula
        wb.save(path)
        wb.close()
    import shutil

    shutil.copy2(template, carried)
    apply_report = QuarterlyModelContinuityService().apply_persistent_continuity(
        analysis_id="mc3",
        ticker="AAPL",
        new_template_path=template,
        previous_workbook_path=previous,
        workbook_path=carried,
    )
    wb = load_workbook(carried)
    try:
        assert str(wb["Tax"]["L13"].value).startswith("=")
        actions = [e.final_action for e in apply_report.entries if e.cell == "L13"]
        assert CarryForwardDecision.RESTORE_FORMULA in actions
    finally:
        wb.close()


def test_research_includes_sec_manifest():
    manifest = {
        "selected_filings": [
            {
                "filing_type": "8-K",
                "filing_date": "2024-10-31",
                "accession_number": "0000320193-24-000123",
                "url": "https://www.sec.gov/example",
            }
        ]
    }
    report = QuarterlyResearchService().gather(
        analysis_id="r1",
        ticker="AAPL",
        fiscal_year=2024,
        fiscal_quarter=4,
        sec_manifest=manifest,
    )
    assert report.sources
    assert any(s.source_kind == "sec_8k" for s in report.sources)
    assert report.reported_facts


def test_restore_ignored_sheets_after_modification(tmp_path: Path):
    template = tmp_path / "template.xlsx"
    out = tmp_path / "out.xlsx"
    wb = Workbook()
    ig = wb.active
    ig.title = "DividendHelper"
    ig["K2"] = "KEEP"
    wb.save(template)
    wb.close()
    import shutil

    shutil.copy2(template, out)
    wb2 = load_workbook(out)
    wb2["DividendHelper"]["K2"] = "CHANGED"
    wb2.save(out)
    wb2.close()
    restored = QuarterlyModelContinuityService().restore_ignored_sheets(
        new_template_path=template, workbook_path=out
    )
    wb3 = load_workbook(out)
    try:
        assert "DividendHelper" in restored
        assert wb3["DividendHelper"]["K2"].value == "KEEP"
    finally:
        wb3.close()


def test_yahoo_millions_conversion():
    from services.yahoo_quarterly_statement_service import _to_millions

    assert _to_millions(94_930_000_000) == pytest.approx(94930.0)
