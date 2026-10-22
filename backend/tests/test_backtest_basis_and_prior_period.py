"""Back-test fixes: share basis of a point-in-time filter, stale prior-quarter column, 10-K of the period, new company by period."""

from __future__ import annotations

import datetime as dt
from pathlib import Path

from openpyxl import Workbook, load_workbook

from scripts.run_batch import match_files
from services.company_description_service import as_of_date, current_10k
from services.point_in_time_filter_service import _clean_factor, _scale_kind
from services.stale_prior_quarter_service import (
    LQ_BS,
    clear_stale_prior_quarter,
    prior_quarter_is_earlier,
)


def test_clean_factor_rounds_to_a_split_and_ignores_noise():
    assert _clean_factor(69.78 / 7.7536) == 9.0           # two 3-for-1 splits (TPL)
    assert _clean_factor(3.02) == 3.0
    assert _clean_factor(1.5) == 1.5
    assert _clean_factor(1.0) == 1.0 and _clean_factor(1.02) == 1.0
    assert _clean_factor(4.25) == 1.0                     # not close to a clean ratio: no conversion


def test_only_per_share_series_are_scaled():
    assert _scale_kind("Closing Stock Price (Q. Average)") == "per_share"
    assert _scale_kind("E10") == "per_share" and _scale_kind("Dividends per Share") == "per_share"
    assert _scale_kind("Shares Outstanding Diluted Average (MM)") == "shares"
    for ratio in ("PE10", "Tangible Book Value per Share/Price", "Dividend Yield TTM", "Price change in 10 years", "Revenue"):
        assert _scale_kind(ratio) is None


def test_as_of_date_for_a_year_and_a_quarter():
    assert as_of_date(2019, None) == "2019-12-31"
    assert as_of_date(2020, 3) == "2020-09-30"
    assert as_of_date(2021, 1) == "2021-03-31"
    assert as_of_date(None, None) is None


class _FakeSec:
    def list_recent_filings(self, cik, forms=None):
        return [
            {"report_date": "2025-12-31", "document_url": "u25"},
            {"report_date": "2020-12-31", "document_url": "u20"},
            {"report_date": "2019-12-31", "document_url": "u19"},
            {"report_date": "2018-12-31", "document_url": "u18"},
        ]


def test_10k_of_the_period_is_not_the_latest():
    sec = _FakeSec()
    assert current_10k(sec, "1", None)["document_url"] == "u25"
    assert current_10k(sec, "1", "2019-12-31")["document_url"] == "u19"
    assert current_10k(sec, "1", "2020-09-30")["document_url"] == "u19"      # the FY2020 10-K did not exist yet
    assert current_10k(sec, "1", "2010-12-31") is None


def _lq_bs(path: Path, prior_date) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = LQ_BS
    ws["C4"], ws["D4"] = dt.datetime(2019, 12, 31), prior_date
    ws["C5"], ws["D5"] = "2019 Q4", "2026 Q1"
    ws["B6"] = "2026 Q2"
    ws["D2"] = '="F"&B6'
    ws["D8"] = "=D4"
    ws["B11"], ws["C11"], ws["D11"] = "C&CE", 303.65, 247.57
    ws["B61"], ws["C61"], ws["D61"] = "BS_TOT_ASSET", 598.18, 1751.45
    ws["D14"] = "=ROUND(D11,1)"                      # a formula of the column itself stays
    ws["I10"] = "=D61/D106"
    ws["E11"] = "=IF(D11=0,\"n/a\",C11/D11-1)"
    wb.save(path)


def test_a_prior_column_dated_after_the_quarter_is_cleared(tmp_path):
    path = tmp_path / "w.xlsx"
    _lq_bs(path, dt.datetime(2026, 3, 31))
    assert prior_quarter_is_earlier(load_workbook(path)[LQ_BS]) is False
    report = clear_stale_prior_quarter(path)
    assert report["cleared"] >= 4 and report["wrapped_formulas"] == 1
    ws = load_workbook(path)[LQ_BS]
    assert ws["D11"].value is None and ws["D61"].value is None and ws["D4"].value is None
    assert ws["C11"].value == 303.65                      # the quarter itself is untouched
    assert ws["D14"].value == "=ROUND(D11,1)"
    assert ws["I10"].value == '=IFERROR(D61/D106,"n/a")'
    assert ws["D2"].value == "n/a"
    assert "HAP Adjustments" in load_workbook(path).sheetnames


def test_a_genuine_prior_quarter_is_kept(tmp_path):
    path = tmp_path / "w.xlsx"
    _lq_bs(path, dt.datetime(2019, 9, 30))
    assert prior_quarter_is_earlier(load_workbook(path)[LQ_BS]) is True
    assert clear_stale_prior_quarter(path)["cleared"] == 0
    assert load_workbook(path)[LQ_BS]["D61"].value == 1751.45


def test_new_company_can_name_its_period_when_the_folder_holds_two_workbooks(tmp_path):
    for name in ("TPL 2019 Q4 - Industrial Template v28.2.xlsx", "TPL 2020 Q3 - Industrial Template v28.2.xlsx",
                 "custom_run_filter_2026-10-16_TPL.xlsx"):
        (tmp_path / name).write_bytes(b"")
    both = match_files(tmp_path, "TPL", "new_company")
    assert both.problems and "exactly one" in both.problems[0]
    q4 = match_files(tmp_path, "TPL", "new_company", (2019, 4))
    q3 = match_files(tmp_path, "TPL", "new_company", (2020, 3))
    assert not q4.problems and q4.workbook.name.startswith("TPL 2019 Q4")
    assert not q3.problems and q3.workbook.name.startswith("TPL 2020 Q3")
    missing = match_files(tmp_path, "TPL", "new_company", (2021, 1))
    assert missing.problems and "requested period" in missing.problems[0]


def test_derived_tax_rate_fills_only_blank_years(tmp_path):
    from types import SimpleNamespace

    from services.new_company_tax_service import NewCompanyTaxService

    wb = Workbook()
    ws = wb.active
    ws.title = "Inputs"
    ws["E106"] = 1
    for col, (fy, pretax, net) in zip("CDE", (("FY 2017", 100.0, 80.0), ("FY 2018", 200.0, 158.0), ("FY 2019", 300.0, 240.0))):
        ws[f"{col}1"], ws[f"{col}37"], ws[f"{col}38"] = fy, pretax, net
    ws["D112"] = 0.1987                                   # a year the filing extractor already filled
    path = tmp_path / "w.xlsx"
    wb.save(path)
    done = NewCompanyTaxService._fill_derived_rates(path, [SimpleNamespace(fiscal_year=y) for y in ("FY2017", "FY2018", "FY2019")])
    assert done == ["FY2017", "FY2019"]
    out = load_workbook(path)["Inputs"]
    assert out["C112"].value == "=IFERROR(MAX(0,MIN(0.6,1-C38/C37)),0)" and out["C107"].value == "=C112" and out["C108"].value == 0
    assert out["D112"].value == 0.1987 and out["D107"].value is None
    assert "HAP Adjustments" in load_workbook(path).sheetnames


def test_derived_tax_in_usd_millions_mode_is_the_tax_charge(tmp_path):
    from types import SimpleNamespace

    from services.new_company_tax_service import NewCompanyTaxService

    wb = Workbook()
    ws = wb.active
    ws.title = "Inputs"
    ws["E106"] = 0
    ws["C1"], ws["C37"], ws["C38"] = "FY 2017", 100.0, 80.0
    path = tmp_path / "w.xlsx"
    wb.save(path)
    assert NewCompanyTaxService._fill_derived_rates(path, [SimpleNamespace(fiscal_year="FY2017")]) == ["FY2017"]
    out = load_workbook(path)["Inputs"]
    assert out["C112"].value == "=IFERROR(MAX(0,C37-C38),0)" and out["C107"].value == "=C112"


def test_derived_tax_accepts_statement_links(tmp_path):
    """The template links Inputs rows 37-38 to the statements (formulas), so those count as present."""
    from types import SimpleNamespace

    from services.new_company_tax_service import NewCompanyTaxService

    wb = Workbook()
    ws = wb.active
    ws.title = "Inputs"
    ws["E106"] = 1
    ws["C1"], ws["C37"], ws["C38"] = "FY 2012", "='Income - GAAP'!C42", "='Income - GAAP'!C58"
    path = tmp_path / "w.xlsx"
    wb.save(path)
    assert NewCompanyTaxService._fill_derived_rates(path, [SimpleNamespace(fiscal_year="FY2012")]) == ["FY2012"]
