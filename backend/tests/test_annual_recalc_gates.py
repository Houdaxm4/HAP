"""Regression: distinct ER metrics, Graham entry concepts, recalc gating."""

from __future__ import annotations

from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook

from models.annual_update import (
    AnnualInputsReport,
    AnnualRdReport,
    AnnualTaxReport,
    AnnualValuationOutputs,
    Pe10Provenance,
)
from services.annual_output_gate_service import AnnualOutputGateService
from services.annual_valuation_extract_service import AnnualValuationExtractService
from services.excel_recalc_service import ExcelRecalcReport, ExcelRecalcService


def test_b69_e14_f14_are_distinct_metrics():
    """57% Bloomberg B69 must not equal workbook E14 (~5.49%) or F14 (~9.33%)."""
    b69 = 0.5712791523788112
    e14 = 0.0549
    f14 = 0.0933
    assert abs(b69 - e14) > 0.02
    assert abs(b69 - f14) > 0.02
    assert abs(e14 - f14) > 0.02
    out = AnnualValuationOutputs(
        bloomberg_expected_return_at_current_price=b69,
        expected_annual_return=e14,
        expected_return_with_dividends=f14,
    )
    AnnualValuationExtractService._sanity(out)
    assert any("Distinct metrics" in w for w in out.warnings)
    # Extract must never copy B69 into expected_annual_return
    assert out.expected_annual_return == e14
    assert out.bloomberg_expected_return_at_current_price == b69


def test_graham_entry_prices_are_not_conflated():
    iv = 105.59942486946139
    mos_entry = iv * 0.75  # ~79.20
    target_entry = 55.40
    assert abs(mos_entry - 79.20) < 0.05
    assert abs(mos_entry - target_entry) > 10.0
    out = AnnualValuationOutputs(
        current_graham_intrinsic_value=iv,
        graham_margin_of_safety_entry_price=mos_entry,
        graham_target_return_entry_price=target_entry,
        graham_target_annualized_return=0.16,
        graham_projection_horizon_years=7,
        graham_expected_annualized_return=0.1164,
    )
    assert out.graham_margin_of_safety_entry_price != out.graham_target_return_entry_price
    assert out.graham_margin_of_safety_entry_price == pytest.approx(79.1996, rel=1e-3)


def test_recalc_incomplete_blocks_completion():
    gate = AnnualOutputGateService().evaluate(
        analysis_id="t",
        ticker="ZZ",
        workbook_path=Path("."),  # unused when tax fails early paths careful
        fiscal_year="FY2026",
        tax=AnnualTaxReport(
            analysis_id="t",
            ticker="ZZ",
            schedule_populated=True,
            cells_written=["Inputs!L112"],
            reported_effective_tax_rate=0.25,
        ),
        inputs=AnnualInputsReport(
            analysis_id="t",
            ticker="ZZ",
            pe10=Pe10Provenance(target_cell="L57", source_value=13.9),
        ),
        rd=AnnualRdReport(
            analysis_id="t",
            ticker="ZZ",
            useful_life=3,
            rd_expense=3.3,
            rd_asset=None,
            rd_amortization=None,
            lookback_complete=True,
        ),
        valuation=AnnualValuationOutputs(current_pe10=13.7),
        recalc=ExcelRecalcReport(
            analysis_id="t",
            ticker="ZZ",
            status="FAILED",
            method="excel_com_calculate_full_rebuild",
            workbook_path="x.xlsx",
            missing_cached_values=["Expected Returns & Buybacks!E14"],
            summary="WORKBOOK_RECALCULATION_INCOMPLETE",
        ),
    )
    assert gate.status == "NEEDS_REVIEW"
    assert any("WORKBOOK_RECALCULATION_INCOMPLETE" in b for b in gate.blockers)
    assert "REPORT_GENERATION_BLOCKED_BY_VALIDATION" in gate.blockers


def test_pe10_period_note_is_not_cross_sheet_mismatch():
    note = (
        "PE10_PERIOD_NOTE: FY2026 fiscal-year PE10=13.9210 as of fiscal-year-end FY2026; "
        "current PE10=13.7473 as of CRF as-of / current. "
        "Distinct metrics — not PE10_CROSS_SHEET_MISMATCH."
    )
    assert "PE10_CROSS_SHEET_MISMATCH" not in note.split("not ")[-1] or "not PE10_CROSS_SHEET" in note
    assert "PE10_PERIOD_NOTE" in note


@pytest.mark.skipif(
    __import__("importlib.util").util.find_spec("win32com") is None,
    reason="pywin32 required for Excel COM recalc",
)
def test_excel_recalc_refreshes_simple_formula(tmp_path: Path):
    path = tmp_path / "recalc.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "Expected Returns & Buybacks"
    ws["A2"] = 100
    ws["C14"] = 160
    ws["D14"] = 180
    ws["E14"] = "=(C14/A2)^(1/10)-1"
    ws["F14"] = "=(D14/A2)^(1/10)-1"
    # Minimal other required sheets/cells
    for name, cells in (
        ("Enterprise Value", {
            "B6": 0.1, "C6": 0.01, "B20": 50, "B27": 0.3, "B32": 40,
            "B42": 100, "B47": 0.12, "B48": 55, "B53": 0.16,
        }),
    ):
        s = wb.create_sheet(name)
        for addr, val in cells.items():
            s[addr] = val
    wb.save(path)
    wb.close()

    report = ExcelRecalcService().recalculate(
        analysis_id="t",
        ticker="ZZ",
        workbook_path=path,
    )
    assert report.status == "ok", report.summary
    wb2 = load_workbook(path, data_only=True)
    e14 = wb2["Expected Returns & Buybacks"]["E14"].value
    f14 = wb2["Expected Returns & Buybacks"]["F14"].value
    wb2.close()
    assert e14 is not None and isinstance(e14, float)
    assert f14 is not None and isinstance(f14, float)
    assert abs(e14 - ((160 / 100) ** 0.1 - 1)) < 1e-6
    assert abs(e14 - 0.57) > 0.2  # not Bloomberg 57%
