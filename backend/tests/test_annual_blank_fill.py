"""Annual update: blank new-FY cells are filled from the 10-K; needed-but-unavailable ones are flagged; supplied values are kept."""

from openpyxl import Workbook, load_workbook

from services.annual_statement_validation_service import AnnualStatementValidationService
from tests.ledger_util import entry, notes_text


def build(path, *, formula_in_revenue=False):
    wb = Workbook()
    ws = wb.active
    ws.title = "Income - GAAP"
    for i, year in enumerate((2024, 2025)):
        ws.cell(7, 3 + i, f"FY{year}")
    ws.cell(11, 1, "Revenue")
    ws.cell(12, 1, "Operating Income")
    ws.cell(13, 1, "Net Income")
    ws.cell(13, 4, 100.0)                      # supplied, differs materially from the filing
    if formula_in_revenue:
        ws.cell(11, 4, "=C11*1.1")
    bs = wb.create_sheet("Balance Sheet - Standardized")
    bs.cell(7, 4, "FY2025")
    bs.cell(9, 1, "Total Assets")              # blank, filing has it, nothing in the metrics uses it
    wb.create_sheet("Cash Flow - Standardized")
    metrics = wb.create_sheet("Final Metrics")
    metrics["A1"] = "='Income - GAAP'!D11+'Income - GAAP'!D12+'Income - GAAP'!D13"
    wb.save(path)


def validate(path):
    return AnnualStatementValidationService().validate(
        analysis_id="a", ticker="T", workbook_path=path, fiscal_year="FY2025",
        filing_overrides={"revenue": 500.0, "net_income": 999.0, "total_assets": 4000.0},
    )


def test_blank_filled_missing_needed_flagged_and_supplied_kept(tmp_path):
    path = tmp_path / "w.xlsx"
    build(path)
    report = validate(path)
    status = {i.concept: i.status for i in report.items}
    assert status["revenue"] == "FILLED_FROM_SEC" and status["total_assets"] == "FILLED_FROM_SEC"
    assert status["operating_income"] == "MISSING_IMPORTANT"   # blank, not in the filing, feeds Final Metrics
    assert status["net_income"] == "DISCREPANCY"
    assert (report.filled, report.missing_important, report.discrepancies) == (2, 1, 1)
    wb = load_workbook(path)
    income = wb["Income - GAAP"]
    assert income["D11"].value == 500.0 and income["D11"].comment is None and entry(wb, "Income - GAAP", "D11", "Filled from filing")
    assert income["D12"].value is None and income["D12"].comment is None and entry(wb, "Income - GAAP", "D12", "Missing figure")
    assert "Beige cells were filled from a filing" in notes_text(income) and "Red cells are missing a figure" in notes_text(income)
    assert income["D13"].value == 100.0  # a supplied value is never overwritten


def test_formula_cells_are_never_filled(tmp_path):
    path = tmp_path / "w.xlsx"
    build(path, formula_in_revenue=True)
    report = validate(path)
    assert next(i for i in report.items if i.concept == "revenue").status == "FORMULA_PRESERVED"
    assert load_workbook(path)["Income - GAAP"]["D11"].value == "=C11*1.1"


def test_gate_does_not_block_on_filled_values_but_warns_when_needed_data_is_missing():
    from types import SimpleNamespace

    from services.annual_output_gate_service import AnnualOutputGateService

    items = [SimpleNamespace(status="FILLED_FROM_SEC", concept="cfo", fiscal_year="FY2025"),
             SimpleNamespace(status="MISSING_IMPORTANT", concept="cfi", fiscal_year="FY2025")]
    blockers, warnings, gates = [], [], {}
    AnnualOutputGateService._apply_statement_gates(gates, blockers, warnings, SimpleNamespace(items=items, discrepancies=0), None)
    assert not blockers and gates["statements"] == "pass"
    assert any(w.startswith("DATA_UNAVAILABLE: cfi@FY2025") for w in warnings)


def test_flag_titles_use_readable_concept_names():
    from services.report_flags import _concept

    assert _concept("cfo") == "Cash from operations" and _concept("total_assets") == "Total Assets" and _concept(None) == ""
