"""A year with known repurchase dollars but no share count in any filing is a warning, not a blocker (the share cell stays blank)."""

from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook, load_workbook

from models.new_company import BuybackAbsenceClass, BuybackYearResult, NewCompanyBuybackReport
from services.annual_output_gate_service import AnnualOutputGateService
from services.new_company_buyback_service import NewCompanyBuybackService


def _facts() -> dict:
    entry = {"start": "2022-10-01", "end": "2023-09-30", "val": 50_000_000, "fy": 2023, "fp": "FY", "form": "10-K", "filed": "2023-11-01", "accn": "0001"}
    return {"facts": {"us-gaap": {"PaymentsForRepurchaseOfCommonStock": {"units": {"USD": [entry]}}}}}


def _workbook(path: Path) -> Path:
    wb = Workbook()
    ws = wb.active
    ws.title = "Income - GAAP"
    ws["A5"] = "Fiscal Year"
    ws.cell(5, 3).value = "FY2023"
    ws["A10"] = "Repurchase of common stock"
    ws["A11"] = "Shares repurchased"
    wb.save(path)
    return path


def _apply(tmp_path: Path, filings_text):
    return NewCompanyBuybackService().apply(
        analysis_id="a", ticker="XYZ", workbook_path=_workbook(tmp_path / "w.xlsx"), fiscal_years=["FY2023"],
        company_facts=_facts(), filings_text=filings_text, write_policy="new_company",
    )


def test_searched_filings_without_a_share_count_give_not_disclosed_and_a_warning(tmp_path: Path):
    report = _apply(tmp_path, {"FY2023": "<p>The Company repurchased shares under its program during the year.</p>"})
    year = report.years[0]
    assert year.dollars == 50.0 and year.shares is None and year.absence_class == BuybackAbsenceClass.NOT_DISCLOSED
    assert "BUYBACK_SHARES_NOT_DISCLOSED_IN_FILINGS: FY2023" in report.warnings
    assert not any("COVERAGE_INCOMPLETE" in w for w in report.warnings) and report.complete is True
    ws = load_workbook(tmp_path / "w.xlsx")["Income - GAAP"]
    assert ws["C10"].value == 50.0 and ws["C11"].value is None        # dollars written, the share count is never invented
    assert ws["C11"].fill.fgColor.rgb.endswith("FF6B6B")               # the blank is flagged red in the income tab
    assert "Missing figure" in [c.value for row in load_workbook(tmp_path / "w.xlsx")["HAP Adjustments"].iter_rows() for c in row]
    assert report.years[0].write_action == "shares_not_disclosed_flagged"


def test_without_any_filing_text_the_gap_is_still_unresolved(tmp_path: Path):
    report = _apply(tmp_path, {})
    year = report.years[0]
    assert year.shares is None and year.absence_class is None
    assert "BUYBACK_SHARES_COVERAGE_INCOMPLETE: FY2023" in report.warnings


def _gate(year: BuybackYearResult):
    report = NewCompanyBuybackReport(analysis_id="a", ticker="XYZ", years=[year], analysis=None, complete=True)  # type: ignore[arg-type]
    gates: dict[str, str] = {}
    blockers: list[str] = []
    warnings: list[str] = []
    AnnualOutputGateService._apply_buyback_gate(gates, blockers, warnings, "FY2023", report)
    return gates, blockers, warnings


def test_annual_gate_passes_with_a_warning_when_shares_are_not_disclosed():
    gates, blockers, warnings = _gate(BuybackYearResult(fiscal_year="FY2023", dollars=50.0, absence_class=BuybackAbsenceClass.NOT_DISCLOSED))
    assert gates["buybacks"] == "pass" and not blockers
    assert warnings == ["BUYBACK_SHARES_NOT_DISCLOSED_IN_FILINGS: FY2023"]


def test_annual_gate_still_blocks_an_unexplained_share_gap_and_missing_dollars():
    _g, blockers, _w = _gate(BuybackYearResult(fiscal_year="FY2023", dollars=50.0))
    assert blockers == ["BUYBACK_SHARES_COVERAGE_INCOMPLETE: FY2023"]
    _g, blockers, _w = _gate(BuybackYearResult(fiscal_year="FY2023", absence_class=BuybackAbsenceClass.NOT_DISCLOSED))
    assert blockers == ["BUYBACK_DOLLARS_COVERAGE_INCOMPLETE: FY2023"]
