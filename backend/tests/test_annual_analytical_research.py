"""Analytical research evidence layer — question-driven, primary-source, no SOURCE_FILL."""

from __future__ import annotations

from pathlib import Path

import pytest
from openpyxl import load_workbook

from models.annual_update import (
    INSUFFICIENT_BASE_EVIDENCE,
    KEEP_REPORTED_BASE,
    MANAGEMENT_STATEMENT,
    NO_EXTERNAL_RESEARCH_REQUIRED,
    NORMALIZATION_SUPPORTABLE_AMOUNT_NOT_SELECTED,
    REMAINS_INSUFFICIENT,
    RESEARCHED,
    SEC_DISCLOSURE,
    TEMPORARY_PROJECT_CAPEX,
    USE_NORMALIZED_BASE,
)
from models.write_actions import WriteActionClass
from services.annual_analytical_research_service import AnnualAnalyticalResearchService
from services.annual_judgment_service import AnnualJudgmentService
from services.annual_normalized_base_service import AnnualNormalizedBaseService
from services.sec_service import SecServiceError
from tests.test_annual_growth_analysis import _geo, _valuation_workbook
from tests.test_annual_normalized_base import _analyze, _snapshot_ev


class FakeSec:
    def __init__(self, documents: dict[str, str] | None = None, filings: list[dict] | None = None) -> None:
        self.documents = documents or {}
        self.filings = filings or []
        self.cache_dir = None
        self.fetched: list[str] = []

    def fetch_document_text(self, url: str, *, cik: str | None = None, cache_name: str | None = None) -> str:
        self.fetched.append(url)
        if url not in self.documents:
            raise SecServiceError(f"missing {url}")
        return self.documents[url]

    def list_recent_filings(self, cik: str, *, forms=None, items_contains=None):
        return list(self.filings)


def _capex_base():
    return _analyze(
        ni=[62.0] * 10,
        da=[20.0] * 10,
        capex=[-25.0] * 7 + [-28.0, -51.0, -88.0],
        revenue=[1000.0] * 10,
        operating_income=[85.0] * 10,
    )


def _keep_base():
    return _analyze(
        ni=_geo(36.0, 0.04, 10),
        da=[15.0 + i * 0.5 for i in range(10)],
        capex=[-12.0 - i * 0.4 for i in range(10)],
        revenue=_geo(800.0, 0.037, 10),
        operating_income=_geo(60.0, 0.04, 10),
    )


def _manifest(*docs: dict) -> dict:
    return {"ticker": "ZZ", "cik": "0000000001", "selected_filings": list(docs)}


def _filing(url: str, form: str = "10-K", filed: str = "2026-08-01", fy: int = 2026) -> dict:
    return {
        "accession_number": "0001-26-000001",
        "filing_type": form,
        "filing_date": filed,
        "report_date": f"{fy}-06-30",
        "primary_document": "doc.htm",
        "document_url": url,
        "fiscal_year": fy,
    }


_PROJECT_ONLY = """
<html><body>
Item 7. Management's Discussion and Analysis
Capital expenditures increased to $88.1 million from $50.7 million in the prior year.
The increase was primarily due to construction of a new production facility and related equipment.
We continue to invest in the facility.
Item 8. Financial Statements
</body></html>
"""

_PROJECT_COMPLETE_GUIDE = """
<html><body>
Item 7. Management's Discussion and Analysis
Capital expenditures were $88.1 million in fiscal 2026 compared with $50.7 million in fiscal 2025,
primarily related to the Gary facility expansion project. The project is expected to be completed
in fiscal 2027. We expect capital expenditures to decrease after completion of the project as
spending returns to more normal replacement levels.
Item 8. Financial Statements
</body></html>
"""

_STRUCTURAL = """
<html><body>
Item 7. Management's Discussion
Capital expenditures remain elevated. Management expects a structurally higher level of capital
spending as ongoing investment in automation continues indefinitely.
Item 8.
</body></html>
"""

_CONTRA = """
<html><body>
Item 7. Management's Discussion
The Lake Street expansion project is expected to be completed in fiscal 2027 and we expect
capital expenditures to decrease after completion. Capital expenditures nonetheless remain elevated
and the company expects ongoing investment at this higher level of capital intensity.
Item 8.
</body></html>
"""


def test_research_never_source_fills_and_does_not_import_fill_pipeline():
    src = Path(__file__).resolve().parents[1] / "services" / "annual_analytical_research_service.py"
    text = src.read_text(encoding="utf-8")
    assert "WriteActionClass" not in text
    assert "excel_fill" not in text.lower()
    assert "load_workbook" not in text
    svc = AnnualAnalyticalResearchService(sec_service=FakeSec())
    assert not hasattr(svc, "source_fill")
    assert WriteActionClass.SOURCE_FILL.value == "SOURCE_FILL"


def test_research_begins_from_specific_unresolved_question():
    base = _capex_base()
    status, questions = AnnualAnalyticalResearchService(sec_service=FakeSec()).generate_questions("ZZ", base)
    assert status == RESEARCHED
    assert len(questions) == 1
    q = questions[0].question.lower()
    assert "capital expenditures" in q
    assert "28" in q and "51" in q and "88" in q
    assert "temporary" in q and "structural" in q
    assert "generic" not in q
    assert questions[0].concept == "capex_temporary_vs_structural"


def test_primary_sources_outrank_secondary():
    svc = AnnualAnalyticalResearchService(sec_service=FakeSec())
    from models.annual_update import AnalyticalResearchEvidence, SECONDARY_SOURCE

    rows = [
        AnalyticalResearchEvidence(
            company="ZZ",
            metric_or_issue="capex",
            research_question="q",
            source_type=SECONDARY_SOURCE,
            source_title="blog",
            evidence_text_or_summary="blog says temporary",
            source_tier=3,
        ),
        AnalyticalResearchEvidence(
            company="ZZ",
            metric_or_issue="capex",
            research_question="q",
            source_type=SEC_DISCLOSURE,
            source_title="10-K",
            evidence_text_or_summary="Capital expenditures were $88 million for a facility expansion.",
            source_tier=1,
            filing_type="10-K",
        ),
    ]
    kept = svc._drop_secondary(rows)
    assert all(r.source_type != SECONDARY_SOURCE for r in kept)
    assert kept[0].source_tier == 1


def test_management_claims_remain_labeled_and_hap_inference_is_separate():
    url = "https://www.sec.gov/Archives/example/10k.htm"
    svc = AnnualAnalyticalResearchService(
        sec_service=FakeSec({url: _PROJECT_COMPLETE_GUIDE})
    )
    report = svc.investigate(
        analysis_id="t",
        ticker="ZZ",
        base=_capex_base(),
        sec_manifest=_manifest(_filing(url)),
        fiscal_year=2026,
        analysis_as_of_date="2026-08-19",
        today="2026-08-19",
    )
    claims = [e for e in report.evidence if e.management_claim]
    assert claims
    assert all(e.source_type in {MANAGEMENT_STATEMENT, SEC_DISCLOSURE} for e in report.evidence)
    assert all(e.hap_interpretation is None for e in report.evidence)
    assert report.synthesis is not None
    assert report.synthesis.hap_interpretation.startswith("HAP inference")


def test_repeated_claim_is_not_independent_corroboration():
    url_a = "https://www.sec.gov/a.htm"
    url_b = "https://www.sec.gov/b.htm"
    html = _PROJECT_COMPLETE_GUIDE
    svc = AnnualAnalyticalResearchService(sec_service=FakeSec({url_a: html, url_b: html}))
    report = svc.investigate(
        analysis_id="t",
        ticker="ZZ",
        base=_capex_base(),
        sec_manifest=_manifest(
            _filing(url_a, filed="2026-08-01", fy=2026),
            _filing(url_b, form="10-Q", filed="2026-08-02", fy=2026),
        ),
        fiscal_year=2026,
        analysis_as_of_date="2026-08-19",
        today="2026-08-19",
    )
    assert report.synthesis is not None
    assert report.synthesis.repeated_claim_count >= 1
    assert report.synthesis.independent_fact_count < len(report.evidence)


def test_capex_spike_alone_cannot_establish_temporary_capex():
    svc = AnnualAnalyticalResearchService(sec_service=FakeSec())
    report = svc.investigate(
        analysis_id="t",
        ticker="ZZ",
        base=_capex_base(),
        sec_manifest=_manifest(),
        fiscal_year=2026,
        analysis_as_of_date="2026-08-19",
        today="2026-08-19",
    )
    assert report.synthesis.decision_effect == REMAINS_INSUFFICIENT
    assert report.synthesis.hap_conclusion != TEMPORARY_PROJECT_CAPEX
    updated = svc.apply_to_diagnostic(_capex_base(), report)
    assert updated.decision == INSUFFICIENT_BASE_EVIDENCE
    assert updated.selected_normalized_base is None


def test_disclosed_project_alone_does_not_establish_future_normalization():
    url = "https://www.sec.gov/project-only.htm"
    svc = AnnualAnalyticalResearchService(sec_service=FakeSec({url: _PROJECT_ONLY}))
    report = svc.investigate(
        analysis_id="t",
        ticker="ZZ",
        base=_capex_base(),
        sec_manifest=_manifest(_filing(url)),
        fiscal_year=2026,
        analysis_as_of_date="2026-08-19",
        today="2026-08-19",
    )
    assert report.synthesis.decision_effect == REMAINS_INSUFFICIENT
    assert "completion" in " ".join(report.synthesis.missing_evidence).lower() or report.synthesis.hap_conclusion != TEMPORARY_PROJECT_CAPEX


def test_project_completion_and_guidance_can_support_temporary_classification():
    url = "https://www.sec.gov/full.htm"
    svc = AnnualAnalyticalResearchService(sec_service=FakeSec({url: _PROJECT_COMPLETE_GUIDE}))
    report = svc.investigate(
        analysis_id="t",
        ticker="ZZ",
        base=_capex_base(),
        sec_manifest=_manifest(_filing(url)),
        fiscal_year=2026,
        analysis_as_of_date="2026-08-19",
        today="2026-08-19",
    )
    assert report.synthesis.hap_conclusion == TEMPORARY_PROJECT_CAPEX
    assert report.synthesis.decision_effect == NORMALIZATION_SUPPORTABLE_AMOUNT_NOT_SELECTED
    updated = svc.apply_to_diagnostic(_capex_base(), report)
    assert updated.decision == USE_NORMALIZED_BASE
    assert updated.selected_normalized_base is None
    assert updated.writes_to_workbook is False


def test_contradictory_evidence_lowers_confidence():
    url = "https://www.sec.gov/contra.htm"
    svc = AnnualAnalyticalResearchService(sec_service=FakeSec({url: _CONTRA}))
    report = svc.investigate(
        analysis_id="t",
        ticker="ZZ",
        base=_capex_base(),
        sec_manifest=_manifest(_filing(url)),
        fiscal_year=2026,
        analysis_as_of_date="2026-08-19",
        today="2026-08-19",
    )
    assert report.synthesis.contradictory_evidence
    assert report.synthesis.confidence == "LOW"
    assert report.synthesis.decision_effect == REMAINS_INSUFFICIENT


def test_as_of_date_excludes_look_ahead_filings():
    past = "https://www.sec.gov/past.htm"
    future = "https://www.sec.gov/future.htm"
    svc = AnnualAnalyticalResearchService(
        sec_service=FakeSec({past: _PROJECT_ONLY, future: _PROJECT_COMPLETE_GUIDE})
    )
    report = svc.investigate(
        analysis_id="t",
        ticker="ZZ",
        base=_capex_base(),
        sec_manifest=_manifest(
            _filing(past, filed="2026-08-01"),
            _filing(future, filed="2026-12-01"),
        ),
        fiscal_year=2026,
        analysis_as_of_date="2026-08-19",
        today="2026-08-19",
    )
    assert all(
        (e.publication_date or "") <= "2026-08-19" or e.available_as_of_analysis is False
        for e in report.evidence
    )
    used = [e for e in report.evidence if e.available_as_of_analysis]
    assert used
    assert all("decrease after completion" not in (e.evidence_text_or_summary or "") for e in used)
    assert report.synthesis.decision_effect == REMAINS_INSUFFICIENT


def test_research_persisted_with_provenance():
    url = "https://www.sec.gov/full.htm"
    svc = AnnualAnalyticalResearchService(sec_service=FakeSec({url: _PROJECT_COMPLETE_GUIDE}))
    report = svc.investigate(
        analysis_id="persist-1",
        ticker="ZZ",
        base=_capex_base(),
        sec_manifest=_manifest(_filing(url, filed="2026-08-19")),
        fiscal_year=2026,
        analysis_as_of_date="2026-08-19",
        today="2026-08-20",
    )
    payload = report.model_dump()
    assert payload["analysis_id"] == "persist-1"
    assert payload["queries"]
    assert payload["questions"]
    assert payload["evidence"][0]["source_locator"]
    assert payload["evidence"][0]["publication_date"] == "2026-08-19"
    assert payload["evidence"][0]["retrieval_date"] == "2026-08-20"
    assert payload["evidence"][0]["document_identity"]
    assert payload["synthesis"]["hap_interpretation"]
    assert payload["writes_to_workbook"] is False


def test_no_external_research_required_when_workbook_sufficient():
    svc = AnnualAnalyticalResearchService(sec_service=FakeSec())
    status, questions = svc.generate_questions("ETD", _keep_base())
    assert status == NO_EXTERNAL_RESEARCH_REQUIRED
    assert questions == []
    report = svc.investigate(
        analysis_id="t",
        ticker="ETD",
        base=_keep_base(),
        sec_manifest=_manifest(),
        fiscal_year=2026,
    )
    assert report.status == NO_EXTERNAL_RESEARCH_REQUIRED
    assert report.evidence == []
    assert svc.sec.fetched == []


def test_research_cannot_write_selected_normalized_base():
    url = "https://www.sec.gov/full.htm"
    svc = AnnualAnalyticalResearchService(sec_service=FakeSec({url: _PROJECT_COMPLETE_GUIDE}))
    base = _capex_base()
    assert base.selected_normalized_base is None
    report = svc.investigate(
        analysis_id="t",
        ticker="ZZ",
        base=base,
        sec_manifest=_manifest(_filing(url)),
        fiscal_year=2026,
        analysis_as_of_date="2026-08-19",
        today="2026-08-19",
    )
    updated = svc.apply_to_diagnostic(base, report)
    assert updated.selected_normalized_base is None
    assert report.selected_normalized_base_written is False


def test_original_workbook_untouched_after_research_apply(tmp_path: Path):
    path = tmp_path / "wb.xlsx"
    _valuation_workbook(
        path,
        ni=[62.0] * 10,
        da=[20.0] * 10,
        capex=[-25.0] * 7 + [-28.0, -51.0, -88.0],
        oi=[85.0] * 10,
        revenue=[1000.0] * 10,
        b6_total=-0.95,
        c6_annual=-0.28,
    )
    before = _snapshot_ev(path)
    url = "https://www.sec.gov/full.htm"
    svc = AnnualAnalyticalResearchService(sec_service=FakeSec({url: _PROJECT_COMPLETE_GUIDE}))
    base = AnnualNormalizedBaseService().analyze_path(path)
    report = svc.investigate(
        analysis_id="t",
        ticker="ZZ",
        base=base,
        sec_manifest=_manifest(_filing(url)),
        fiscal_year=2026,
        analysis_as_of_date="2026-08-19",
        today="2026-08-19",
    )
    updated = svc.apply_to_diagnostic(base, report)
    _, judge = AnnualJudgmentService().apply(
        analysis_id="t",
        ticker="ZZ",
        workbook_path=path,
        context={
            "oe_base_analysis": updated,
            "book_value_growth": 0.03,
            "owner_earnings_growth": -0.28,
            "eps_growth": 0.04,
            "bv_classification": "REASONABLE",
            "oe_classification": "INSUFFICIENT_EVIDENCE",
            "eps_classification": "REASONABLE",
        },
    )
    after = _snapshot_ev(path)
    assert after == before
    assert after["B9"] == "='Final Metrics'!$L$10*(1+$C$6)^B8"
    wb = load_workbook(path, data_only=False)
    try:
        assert wb["Enterprise Value"]["C6"].value == pytest.approx(-0.28)
    finally:
        wb.close()
    assert updated.selected_normalized_base is None
    assert updated.writes_to_workbook is False


def test_no_ticker_hardcodes_in_research_service():
    src = (
        Path(__file__).resolve().parents[1] / "services" / "annual_analytical_research_service.py"
    ).read_text(encoding="utf-8")
    lowered = src.lower()
    for token in ("jbss", "etd", "csco", "mzti"):
        assert token not in lowered
