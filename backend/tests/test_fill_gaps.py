"""Debt and capex rows are filled from SEC with the right row and sign; check rows explain themselves."""

from openpyxl import Workbook, load_workbook

from services.new_company_statement_validation_service import NewCompanyStatementValidationService
from services.statement_row_rules import STATEMENT_ROW_RULES, sign_for


def build(path, *, debt=None, capex=None):
    wb = Workbook()
    ws = wb.active
    ws.title = "Income - GAAP"
    ws.cell(3, 3, "FY 2025")
    bs = wb.create_sheet("Balance Sheet - Standardized")
    bs.cell(3, 3, "FY 2025")
    bs.cell(86, 1, "+ LT Debt")                  # NOT the debt row: only 'LT Borrowings' agrees with SEC
    bs.cell(86, 3, 7.0)
    bs.cell(87, 1, "+ LT Borrowings")
    bs.cell(87, 3, debt)
    bs.cell(124, 1, "Cost of Debt")              # a rate, never an amount
    bs.cell(124, 3, 2.4)
    cf = wb.create_sheet("Cash Flow - Standardized")
    cf.cell(3, 3, "FY 2025")
    cf.cell(32, 1, "+ Acq of Fixed Prod Assets")
    cf.cell(32, 3, capex)
    wb.create_sheet("Final Metrics")["A1"] = "='Balance Sheet - Standardized'!C87+'Cash Flow - Standardized'!C32"
    wb.save(path)


def validate(path):
    return NewCompanyStatementValidationService().validate(
        analysis_id="a", ticker="T", workbook_path=path, fiscal_years=["FY2025"],
        filing_overrides={"FY2025": {"debt": 300.0, "capex": 25.0}},
    )


def test_rules_use_the_rows_that_agree_with_supplied_data():
    assert STATEMENT_ROW_RULES["debt"]["needles"] == ("lt borrowings",)
    assert sign_for("capex") == -1 and sign_for("debt") == 1 and sign_for("revenue") == 1


def test_blank_debt_and_capex_are_filled_into_the_right_rows_with_the_workbook_sign(tmp_path):
    path = tmp_path / "w.xlsx"
    build(path)
    report = validate(path)
    wb = load_workbook(path)
    bs, cf = wb["Balance Sheet - Standardized"], wb["Cash Flow - Standardized"]
    assert bs["C87"].value == 300.0 and "HAP FILLED" in bs["C87"].comment.text
    assert cf["C32"].value == -25.0  # Bloomberg shows the outflow as a negative number
    assert bs["C86"].value == 7.0 and bs["C124"].value == 2.4  # other rows untouched
    assert sorted(f.concept for f in report.filled_missing) == ["capex", "debt"]


def test_supplied_debt_and_capex_are_validated_not_overwritten(tmp_path):
    path = tmp_path / "w.xlsx"
    build(path, debt=301.0, capex=-25.2)
    report = validate(path)
    assert {i["concept"]: i["status"] for i in report.items if i["concept"] in ("debt", "capex")} == {"debt": "validated", "capex": "validated"}
    assert not report.filled_missing and load_workbook(path)["Cash Flow - Standardized"]["C32"].value == -25.2


def test_quarterly_check_row_gets_an_explanatory_note_when_its_total_is_filled():
    from services.quarterly_presentation_service import CHECK_ROW_NOTE

    assert "expected" in CHECK_ROW_NOTE and "not a data error" in CHECK_ROW_NOTE


def test_flags_note_the_quarterly_check_rows(tmp_path):
    import json

    from services.report_flags import collect_flags

    (tmp_path / "quarterly_presentation_report.json").write_text(json.dumps({
        "filled_from_sec": [{"cell": "Last Quarter CF Standardized!C27", "label": "Cash from Operating Activities", "value": 98.6, "source": "SEC"}],
        "statements": [], "flagged_missing": [],
    }))
    titles = [f["title"] for f in collect_flags(tmp_path)["flags"]["notes"]]
    assert "Quarterly cash-flow check rows show differences" in titles
