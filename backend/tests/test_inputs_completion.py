"""Tests for Inputs-tab tax / PE10 / current_data completion."""

from __future__ import annotations

from pathlib import Path
from unittest.mock import patch

import pytest
from openpyxl import Workbook

from canonical_model.company import CompanyFinancialModel
from canonical_model.inputs_bridge import InputsBridge
from canonical_model.primitives import FinancialPoint, FinancialSeries, LineItemProvenance
from models.completion import CompletionDecision
from services.completion_service import CompletionService
from services.sec_tax_inputs import populate_inputs_tax_from_sec
from workbook_mapping.engine import (
    IntentDecision,
    WriteIntent,
    WriteIntentReport,
    load_mapping_specification,
    map_model_to_write_intents,
)
from workbook_mapping.inputs_mappings import inputs_mappings
from workbook_mapping.mapping_schema import CoverageStats, MappingSpecification, PeriodColumnMap


def _mini_sheets() -> list[dict]:
    return [
        {"name": "Inputs", "write_policy": "Hybrid", "fill_priority": "P0", "cells": []},
        {"name": "Income - GAAP", "write_policy": "Hybrid", "fill_priority": "P0", "cells": []},
    ]


def _model_with_inputs() -> CompanyFinancialModel:
    bridge = InputsBridge(
        pe10=FinancialSeries(
            name="PE10",
            points=[
                FinancialPoint(period="FY2018", value=24.69, source="bloomberg_custom_run"),
                FinancialPoint(period="FY2025", value=51.91, source="bloomberg_custom_run"),
            ],
        ),
        e10=FinancialSeries(
            name="E10",
            points=[
                FinancialPoint(period="FY2018", value=1.71, source="bloomberg_custom_run"),
                FinancialPoint(period="FY2025", value=5.18, source="bloomberg_custom_run"),
            ],
        ),
        tax_expense=FinancialSeries(
            name="Income Tax Expense",
            points=[
                FinancialPoint(
                    period="FY2018",
                    value=13_372_000_000.0,
                    source="sec_edgar_10k",
                    provenance=LineItemProvenance(
                        concept="IncomeTaxExpenseBenefit",
                        xbrl_tag="IncomeTaxExpenseBenefit",
                        filing_type="10-K",
                        accession_number="0000320193-18-000145",
                        source_document="SEC 10-K FY2018",
                    ),
                ),
                FinancialPoint(
                    period="FY2025",
                    value=20_719_000_000.0,
                    source="sec_edgar_10k",
                    provenance=LineItemProvenance(
                        concept="IncomeTaxExpenseBenefit",
                        xbrl_tag="IncomeTaxExpenseBenefit",
                        filing_type="10-K",
                        accession_number="0000320193-25-000001",
                        source_document="SEC 10-K FY2025",
                    ),
                ),
            ],
        ),
        tax_federal=FinancialSeries(
            name="Federal Tax",
            points=[
                FinancialPoint(period="FY2025", value=10_000_000_000.0, source="sec_edgar_10k"),
            ],
        ),
        tax_unit_flag=0,
        current_price=190.0,
        current_price_source="market_internet:test",
        current_pe10=49.6,
        current_e10=5.58,
    )
    return CompanyFinancialModel(
        analysis_id="inputs-1",
        ticker="AAPL",
        analysis_type="new_company",
        inputs=bridge,
        metadata={"custom_run_source": "test_crf.xlsx"},
    )


def _spec() -> MappingSpecification:
    return MappingSpecification(
        schema_version="1.0.0",
        milestone="inputs",
        template_family="Industrial Template",
        manifest_schema_version="1.0.0",
        manifest_source="test",
        baseline_ticker="AAPL",
        generated_at="2026-01-01T00:00:00Z",
        period_column_maps=[
            PeriodColumnMap(
                sheet="Inputs",
                header_row=1,
                date_row=2,
                columns={
                    "C": "FY2016",
                    "D": "FY2017",
                    "E": "FY2018",
                    "F": "FY2019",
                    "G": "FY2020",
                    "H": "FY2021",
                    "I": "FY2022",
                    "J": "FY2023",
                    "K": "FY2024",
                    "L": "FY2025",
                },
            ),
            PeriodColumnMap(
                sheet="Income - GAAP",
                header_row=7,
                date_row=8,
                columns={"E": "FY2018", "L": "FY2025"},
            ),
        ],
        mappings=inputs_mappings(),
        writable_dispositions=[],
        unmapped_cfm_metrics=[],
        coverage=CoverageStats(
            total_cfm_metrics=1,
            mapped_cfm_metrics=1,
            unmapped_cfm_metrics=0,
            cfm_coverage_pct=100.0,
            total_writable_workbook_cells=0,
            mapped_writable_cells=0,
            intentionally_empty_writable_cells=0,
            unsupported_writable_cells=0,
            future_feature_writable_cells=0,
            unexplained_writable_cells=0,
            writable_coverage_explained_pct=100.0,
        ),
    )


@pytest.fixture
def inputs_workbook(tmp_path: Path) -> Path:
    wb = Workbook()
    ws = wb.active
    ws.title = "Inputs"
    for col in list("CDEFGHIJKL"):
        ws[f"{col}57"] = None
        ws[f"{col}112"] = None
    ws["B63"] = None
    ws["B65"] = None
    ws["E106"] = 1
    income = wb.create_sheet("Income - GAAP")
    income["C3"] = "FY2025"
    income["E9"] = 265595
    path = tmp_path / "inputs.xlsx"
    wb.save(path)
    wb.close()
    return path


def test_inputs_mappings_present_in_loaded_spec():
    spec = load_mapping_specification()
    ids = {m.mapping_id for m in spec.mappings}
    assert "inputs.pe10" in ids
    assert "inputs.tax_expense" in ids
    assert "inputs.current_price" in ids
    assert any(p.sheet == "Inputs" for p in spec.period_column_maps)


def test_tax_mapping_sec_to_inputs_cell():
    model = _model_with_inputs()
    report = map_model_to_write_intents(
        model,
        _spec(),
        manifest_index={"sheets": {s["name"]: s for s in _mini_sheets()}, "cells": {}, "workbook_protection": {}},
    )
    tax = [i for i in report.intents if i.mapping_id == "inputs.tax_expense" and i.period == "FY2018"]
    assert len(tax) == 1
    assert tax[0].decision == IntentDecision.WRITE
    assert tax[0].cell == "E112"
    assert tax[0].value == pytest.approx(13372.0)
    assert tax[0].source == "sec_edgar_10k"
    assert tax[0].xbrl_tag == "IncomeTaxExpenseBenefit"
    assert tax[0].accession_number


def test_pe10_mapping_crf_to_inputs_cell():
    model = _model_with_inputs()
    report = map_model_to_write_intents(
        model,
        _spec(),
        manifest_index={"sheets": {s["name"]: s for s in _mini_sheets()}, "cells": {}, "workbook_protection": {}},
    )
    pe = [i for i in report.intents if i.mapping_id == "inputs.pe10" and i.period == "FY2018"]
    assert len(pe) == 1
    assert pe[0].decision == IntentDecision.WRITE
    assert pe[0].cell == "E57"
    assert pe[0].value == pytest.approx(24.69)
    assert "bloomberg" in (pe[0].source or "")


def test_current_data_mapping():
    model = _model_with_inputs()
    report = map_model_to_write_intents(
        model,
        _spec(),
        manifest_index={"sheets": {s["name"]: s for s in _mini_sheets()}, "cells": {}, "workbook_protection": {}},
    )
    price = next(i for i in report.intents if i.mapping_id == "inputs.current_price")
    pe = next(i for i in report.intents if i.mapping_id == "inputs.current_pe10")
    assert price.decision == IntentDecision.WRITE
    assert price.cell == "B63"
    assert price.value == pytest.approx(190.0)
    assert "market" in (price.source or "")
    assert pe.decision == IntentDecision.WRITE
    assert pe.value == pytest.approx(49.6)


def test_missing_source_skips_write():
    model = _model_with_inputs()
    model.inputs.current_price = None
    report = map_model_to_write_intents(
        model,
        _spec(),
        manifest_index={"sheets": {s["name"]: s for s in _mini_sheets()}, "cells": {}, "workbook_protection": {}},
    )
    price = next(i for i in report.intents if i.mapping_id == "inputs.current_price")
    assert price.decision == IntentDecision.SKIP


def test_new_company_scope_inputs_vs_annual(inputs_workbook: Path):
    model = _model_with_inputs()
    intent_report = map_model_to_write_intents(
        model,
        _spec(),
        manifest_index={"sheets": {s["name"]: s for s in _mini_sheets()}, "cells": {}, "workbook_protection": {}},
    )
    annual = WriteIntent(
        intent_id="is.revenue:E9",
        mapping_id="is.revenue",
        sheet="Income - GAAP",
        cell="E9",
        metric="Revenue",
        period="FY2018",
        value=265595.0,
        source="sec",
        write_policy_result="writable:Hybrid",
        decision=IntentDecision.WRITE,
        reason="test",
        cfm_path="income_statement.revenue",
    )
    merged = WriteIntentReport(
        analysis_id="x",
        ticker="AAPL",
        intents=list(intent_report.intents) + [annual],
        write_count=intent_report.write_count + 1,
    )
    completion, fill = CompletionService().plan(
        analysis_id="x",
        ticker="AAPL",
        analysis_type="new_company",
        source_workbook_path=inputs_workbook,
        intent_report=merged,
    )
    assert any(
        e.workbook_section == "tax" and e.decision == CompletionDecision.FILL
        for e in completion.entries
    )
    assert any(
        e.workbook_section == "pe10" and e.decision == CompletionDecision.FILL
        for e in completion.entries
    )
    assert any(
        e.workbook_section == "current_data" and e.decision == CompletionDecision.FILL
        for e in completion.entries
    )
    annual_e = next(e for e in completion.entries if e.intent_id == "is.revenue:E9")
    assert annual_e.decision == CompletionDecision.OUT_OF_SCOPE
    assert fill.write_count >= 1


def test_annual_update_target_year_only(inputs_workbook: Path):
    model = _model_with_inputs()
    intent_report = map_model_to_write_intents(
        model,
        _spec(),
        manifest_index={"sheets": {s["name"]: s for s in _mini_sheets()}, "cells": {}, "workbook_protection": {}},
    )
    completion, _ = CompletionService().plan(
        analysis_id="au",
        ticker="AAPL",
        analysis_type="annual_update",
        source_workbook_path=inputs_workbook,
        intent_report=intent_report,
        target_fiscal_year="FY2025",
    )
    pe_entries = [e for e in completion.entries if e.mapping_id == "inputs.pe10"]
    fy2018 = next(e for e in pe_entries if e.period == "FY2018")
    fy2025 = next(e for e in pe_entries if e.period == "FY2025")
    assert fy2018.decision == CompletionDecision.OUT_OF_SCOPE
    assert fy2025.decision == CompletionDecision.FILL
    tax2018 = next(
        e
        for e in completion.entries
        if e.mapping_id == "inputs.tax_expense" and e.period == "FY2018"
    )
    assert tax2018.decision == CompletionDecision.OUT_OF_SCOPE


def test_populated_inputs_already_present(inputs_workbook: Path):
    from openpyxl import load_workbook

    wb = load_workbook(inputs_workbook)
    wb["Inputs"]["L57"] = 51.91
    wb.save(inputs_workbook)
    wb.close()

    model = _model_with_inputs()
    intent_report = map_model_to_write_intents(
        model,
        _spec(),
        manifest_index={"sheets": {s["name"]: s for s in _mini_sheets()}, "cells": {}, "workbook_protection": {}},
    )
    completion, fill = CompletionService().plan(
        analysis_id="ap",
        ticker="AAPL",
        analysis_type="annual_update",
        source_workbook_path=inputs_workbook,
        intent_report=intent_report,
        target_fiscal_year="FY2025",
    )
    pe = next(e for e in completion.entries if e.mapping_id == "inputs.pe10" and e.period == "FY2025")
    assert pe.decision == CompletionDecision.ALREADY_PRESENT
    assert all(
        i.decision != IntentDecision.WRITE or i.cell != "L57" for i in fill.intents
    )


def test_sec_tax_populate_preserves_provenance():
    facts = {
        "facts": {
            "us-gaap": {
                "IncomeTaxExpenseBenefit": {
                    "label": "Income Tax",
                    "units": {
                        "USD": [
                            {
                                "val": 13372000000,
                                "fy": 2018,
                                "fp": "FY",
                                "form": "10-K",
                                "filed": "2018-11-05",
                                "accn": "0000320193-18-000145",
                                "frame": "CY2018",
                                "end": "2018-09-29",
                                "start": "2017-10-01",
                            }
                        ]
                    },
                }
            }
        }
    }
    bridge = InputsBridge()
    populate_inputs_tax_from_sec(bridge, facts, years=10)
    pt = bridge.tax_expense.point_for("FY2018")
    assert pt is not None
    assert pt.value == pytest.approx(13372000000)
    assert pt.provenance is not None
    assert pt.provenance.xbrl_tag == "IncomeTaxExpenseBenefit"
    assert pt.provenance.accession_number == "0000320193-18-000145"


def test_market_price_not_substituted_from_crf():
    with patch(
        "services.market_price_service.MarketPriceService.get_price",
        return_value=(None, None),
    ):
        from canonical_model.builder import CompanyFinancialModelBuilder

        model = CompanyFinancialModelBuilder().build(
            analysis_id="m",
            ticker="AAPL",
            custom_run=None,
            company_facts=None,
            workbook_cells=[],
        )
        assert model.inputs.current_price is None
