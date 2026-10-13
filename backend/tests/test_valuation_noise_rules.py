"""The valuation rules for noisy data: inputs first, then the other model, then red flags; two growth models for owner's earnings;
and an EPS growth rate for the Graham entry price that extraordinary years cannot skew."""

from __future__ import annotations

from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook

from services.valuation_correction_service import ValuationCorrectionService
from services.valuation_history import eps_trend, extraordinary_eps_years, loglinear, off_years, review_eps_growth

YEARS = [f"FY{2017 + i}" for i in range(10)]


def _series(values):
    return dict(zip(YEARS, values))


# ------------------------------------------------------------------------------------------------ the statistics
def test_off_years_are_the_few_years_far_from_the_typical_one():
    roe = _series([0.126, 0.207, 0.189, 0.241, 0.288, 0.248, 0.223, 0.234, 0.239, 0.200])
    retention = _series([0.136, 0.497, 0.305, 0.392, 0.617, 0.585, 0.553, 0.609, 0.631, 0.591])
    assert off_years(roe) == ["FY2017"]
    assert off_years(retention) == ["FY2017", "FY2019"]
    assert off_years(_series([0.2] * 10)) == []


def test_eps_trend_is_steady_for_growth_and_not_for_a_cycle():
    assert eps_trend(_series([1.36, 2.5, 2.01, 2.34, 3.87, 3.79, 3.72, 4.56, 5.35, 5.17])).reasonable
    cyc = eps_trend(_series([2.0, 3.5, 1.2, 4.0, 0.8, 3.9, 1.1, 4.2, 1.0, 4.4]))
    assert not cyc.reasonable and "steady trend" in cyc.reason
    assert not eps_trend(_series([1, 2, 3, -1, 5, 6, 7, 8, 9, 10])).reasonable          # a loss year is not a steady trend


def test_an_extraordinary_end_year_skews_the_endpoint_rate_but_a_middle_spike_does_not():
    steady = [1.0 * 1.1 ** i for i in range(10)]
    end_spike = steady[:-1] + [steady[-1] * 2.2]
    r = review_eps_growth(_series(end_spike), (end_spike[-1] / end_spike[0]) ** (1 / 9) - 1)
    assert r.skewed and r.extraordinary == ["FY2026"] and r.robust == pytest.approx(0.10, abs=1e-6) and r.endpoint > 0.19
    mid = list(steady)
    mid[5] *= 3
    r2 = review_eps_growth(_series(mid), (mid[-1] / mid[0]) ** (1 / 9) - 1)
    assert not r2.skewed                                                                   # the endpoint rate does not depend on a middle year
    assert extraordinary_eps_years(_series(mid)) == ["FY2022"]
    assert loglinear(_series(steady)).growth == pytest.approx(0.10)


# ------------------------------------------------------------------------------------------------ the workbook
def _book(path: Path, *, roe, retention, eps, price=60.0, pe=40.0, bv0=10.0, b41=None, oe=None, c6=None) -> Path:
    wb = Workbook()
    fm = wb.active
    fm.title = "Final Metrics"
    inputs = wb.create_sheet("Inputs")
    for i, fy in enumerate(YEARS):
        fm.cell(1, 3 + i).value = fy
        inputs.cell(1, 3 + i).value = fy
        fm.cell(4, 3 + i).value = roe[i]
        fm.cell(29, 3 + i).value = eps[i]
        inputs.cell(119, 3 + i).value = retention[i]
        if oe:
            fm.cell(10, 3 + i).value = oe[i]
    fm["A4"], fm["A29"], fm["A10"] = "ROE", "EPS", "Owners Earnings"
    inputs["A119"] = "Change in Retained Earnings per Share/EPS"
    avg_roe, avg_ret = sum(roe) / len(roe), sum(retention) / len(retention)
    er = wb.create_sheet("Expected Returns & Buybacks")
    er["A2"], er["E2"], er["B8"], er["C30"] = price, pe, bv0, eps[-1]
    er["B5"] = (eps[-1] / eps[0]) ** (1 / 9) - 1
    er["A14"], er["C5"] = avg_roe, avg_ret
    er["A11"] = "=C5*A14"
    model = ValuationCorrectionService._model_return(price, pe, bv0, avg_ret * avg_roe, avg_roe, 0.0)
    er["E14"] = er["F14"] = model
    ev = wb.create_sheet("Enterprise Value")
    ev["B41"] = b41 if b41 is not None else er["B5"].value
    ev["B2"], ev["B17"], ev["B19"], ev["B20"] = 0.12, 2000.0, 20000.0, 400.0
    ev["C6"] = c6 if c6 is not None else 0.10
    wb.save(path)
    return path


STEADY_EPS = [1.0 * 1.10 ** i for i in range(10)]
CYCLICAL_EPS = [2.0, 3.5, 1.2, 4.0, 0.8, 3.9, 1.1, 4.2, 1.0, 4.4]


def test_a_few_off_years_in_the_inputs_are_removed_and_the_model_is_kept(tmp_path: Path):
    roe = [0.10] * 9 + [3.00]                               # one extraordinary year pulls the average ROE up
    path = _book(tmp_path / "a.xlsx", roe=roe, retention=[0.5] * 10, eps=STEADY_EPS)
    before = load_workbook(path)["Expected Returns & Buybacks"]
    assert before["F14"].value > 0.20                        # the model gives an unrealistic return
    review = ValuationCorrectionService().apply(path)
    er = load_workbook(path)["Expected Returns & Buybacks"]
    assert er["A14"].value == pytest.approx(0.10) and er["C5"].value == pytest.approx(0.5)       # recalculated without FY2026; C5 had no off year, A11 stays a formula
    assert er["A11"].value == "=C5*A14"                      # the model is kept
    applied = [a for a in review.applied if a.topic == "expected_return"]
    assert [a.cell.split("!")[-1] for a in applied] == ["A14"] and "without FY2026" in applied[0].what
    ledger = [str(c.value) for row in load_workbook(path)["HAP Adjustments"].iter_rows(min_row=5) for c in row if c.value]
    assert "Expected return input" in ledger and "Applied" in ledger


def test_when_the_whole_history_is_off_but_eps_is_steady_the_other_model_is_used(tmp_path: Path):
    roe = [0.05, 0.45] * 5                                   # a cyclical business: no few years to remove
    path = _book(tmp_path / "b.xlsx", roe=roe, retention=[0.1, 0.9] * 5, eps=STEADY_EPS, price=40.0, pe=40.0)
    review = ValuationCorrectionService().apply(path)
    er = load_workbook(path)["Expected Returns & Buybacks"]
    assert er["A11"].value == pytest.approx(0.10, abs=1e-4)                         # EPS growth of the steady series
    assert er["A14"].value == pytest.approx(STEADY_EPS[-1] / 10.0, abs=1e-3)       # current EPS / book value per share
    applied = [a for a in review.applied if a.topic == "expected_return"]
    assert {a.cell.split("!")[-1] for a in applied} == {"A11", "A14"}
    notes = review.notes("expected_return")
    assert any("steady trend" in n for n in notes) and all(len(n) <= 400 for n in notes)        # why this model, in a note of its own


def test_when_eps_is_cyclical_too_the_model_is_kept_and_red_flags_are_written(tmp_path: Path):
    roe = [0.05, 0.45] * 5
    path = _book(tmp_path / "c.xlsx", roe=roe, retention=[0.1, 0.9] * 5, eps=CYCLICAL_EPS, price=40.0, pe=40.0)
    original = load_workbook(path)["Expected Returns & Buybacks"]
    a14, a11, c5 = original["A14"].value, original["A11"].value, original["C5"].value
    review = ValuationCorrectionService().apply(path)
    wb = load_workbook(path)
    er = wb["Expected Returns & Buybacks"]
    assert (er["A14"].value, er["A11"].value, er["C5"].value) == (a14, a11, c5)         # nothing in the model changed
    assert not [a for a in review.applied if a.topic == "expected_return"]
    flags = [f.text for f in review.flagged if f.topic == "expected_return"]
    assert any("Return on equity ranged" in t for t in flags) and any("steady trend" in t for t in flags)
    rows = [[str(c.value) for c in row] for row in wb["HAP Adjustments"].iter_rows(min_row=5) if row[0].value]
    assert any("Red flag" in r and "Expected return input" in r for r in rows)
    assert "No model value was changed" in " ".join(review.notes("expected_return"))


def test_the_email_marks_an_unrealistic_return_that_was_kept(tmp_path: Path):
    from services.email_draft_service import _er_adjusted, _er_red_flag

    path = _book(tmp_path / "d.xlsx", roe=[0.05, 0.45] * 5, retention=[0.1, 0.9] * 5, eps=CYCLICAL_EPS, price=40.0, pe=40.0)
    ValuationCorrectionService().apply(path)
    wb = load_workbook(path)
    assert _er_red_flag(wb) and not _er_adjusted(wb)


# ------------------------------------------------------------------------------------------------ owner's earnings
def test_both_owner_earnings_models_are_compared_in_the_notes_and_the_better_one_is_used(tmp_path: Path):
    oe = [74, 200, 260, 280, 290, 300, 310, 330, 380, 424]
    path = _book(tmp_path / "e.xlsx", roe=[0.13] * 10, retention=[0.5] * 10, eps=STEADY_EPS, oe=oe, c6=0.21)
    review = ValuationCorrectionService().apply(path)
    text = " ".join(review.notes("owner_earnings"))
    assert "Final Metrics M10" in text and "Final Metrics N10" in text and "recommended" in text
    assert [a for a in review.applied if a.topic == "owner_earnings"][0].kind == "replaced"
    irregular = [300, 120, 380, 90, 420, 150, 460, 100, 500, 130]
    path2 = _book(tmp_path / "f.xlsx", roe=[0.13] * 10, retention=[0.5] * 10, eps=STEADY_EPS, oe=irregular, c6=0.30)
    review2 = ValuationCorrectionService().apply(path2)
    assert "The average of the first three and last three years is recommended" in " ".join(review2.notes("owner_earnings"))


# ------------------------------------------------------------------------------------------------ Graham
def test_a_skewed_eps_growth_rate_is_corrected_for_the_graham_entry_price(tmp_path: Path):
    eps = STEADY_EPS[:-1] + [STEADY_EPS[-1] * 2.2]           # an extraordinary last year
    endpoint = (eps[-1] / eps[0]) ** (1 / 9) - 1
    path = _book(tmp_path / "g.xlsx", roe=[0.13] * 10, retention=[0.5] * 10, eps=eps, b41=endpoint)
    review = ValuationCorrectionService().apply(path)
    ev = load_workbook(path)["Enterprise Value"]
    assert ev["B41"].value == pytest.approx(0.10, abs=1e-3) and ev["B41"].value < endpoint - 0.05
    applied = [a for a in review.applied if a.topic == "graham"]
    assert applied and "FY2026" in applied[0].why and "was replaced" in applied[0].what
    assert any("EPS growth rate used for the Graham entry price" in n for n in review.notes("graham"))


def test_a_clean_eps_growth_rate_is_left_alone_for_graham(tmp_path: Path):
    path = _book(tmp_path / "h.xlsx", roe=[0.13] * 10, retention=[0.5] * 10, eps=STEADY_EPS, b41=0.1)
    review = ValuationCorrectionService().apply(path)
    assert load_workbook(path)["Enterprise Value"]["B41"].value == 0.1
    assert not [a for a in review.applied if a.topic == "graham"] and review.notes("graham") == []


def test_a_return_up_to_twenty_percent_is_realistic_and_above_it_is_not(tmp_path: Path):
    from tests.test_notes_and_reviews import _er_workbook

    keep = _er_workbook(tmp_path / "k.xlsx", model_return=0.19)
    before = keep.read_bytes()
    review = ValuationCorrectionService().apply(keep)
    assert not review.applied and keep.read_bytes() == before and "0.0% to 20.0%" in review.notes("expected_return")[0]
    assert ValuationCorrectionService().apply(_er_workbook(tmp_path / "s.xlsx", model_return=0.21)).applied          # the old-behaviour fallback switches above 20%
