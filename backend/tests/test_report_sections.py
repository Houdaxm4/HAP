"""Word report: flags come first; fundamentals and valuation sections follow the rules-based verdicts."""

import json
from types import SimpleNamespace

from docx import Document

from services.report_flags import collect_flags, write_flags_section
from services.report_opinion import fundamentals_assessment, valuation_assessment, write_assessment_sections


def write(path, name, payload):
    (path / name).write_text(json.dumps(payload), encoding="utf-8")


def artifacts(tmp_path, *, bq=75.0, authorized=True):
    write(tmp_path, "new_company_output_gate_report.json", {
        "status": "ok" if authorized else "NEEDS_REVIEW", "report_authorized": authorized,
        "blockers": [] if authorized else ["WORKBOOK_RECALCULATION_INCOMPLETE", "NEW_COMPANY_REPORT_NOT_AUTHORIZED"],
        "warnings": ["PE10_PERIOD_NOTE: dates not supplied", "RD_USEFUL_LIFE_AGENT_SELECTED", "DATA_UNAVAILABLE: x"],
    })
    write(tmp_path, "new_company_statement_validation_report.json", {
        "filled_missing": [{"concept": "total_assets", "fiscal_year": "FY2016", "new_value": 1727.9, "cell": "BS!C9", "source": "SEC 10-K"}],
        "flagged_missing": ["FY2018:capex:CF!E20"],
        "corrections": [{"concept": "cash", "fiscal_year": "FY2018", "old_value": 475.0, "new_value": 488.7, "cell": "BS!E11", "source": "SEC"}],
    })
    write(tmp_path, "rd_useful_life_decision.json", {"selected_useful_life": 3, "confidence": 0.4, "rationale": "thin evidence"})
    write(tmp_path, "final_recommendation_report.json", {
        "business_quality_score": bq, "business_quality_classification": "HIGH_QUALITY_BUSINESS",
        "current_price": 100.0, "intrinsic_value": 150.0, "margin_of_safety": 0.33, "valuation_status": "CHEAP",
        "reasons_against": ["Thin evidence."],
    })


def test_flags_are_categorised_and_written_first(tmp_path):
    artifacts(tmp_path, authorized=False)
    collected = collect_flags(tmp_path)
    counts = collected["counts"]
    assert counts["filled"] == 1 and counts["judgments"] == 1
    assert counts["attention"] >= 3  # unavailable-but-needed, material difference, blocker
    assert not any("NOT_AUTHORIZED" in f["detail"] for f in collected["flags"]["attention"])
    assert [f["title"] for f in collected["flags"]["notes"]] == ["PE10 period note"]  # agent decisions are not repeated as notes

    doc = Document()
    doc.add_heading("Title", level=0)
    write_flags_section(doc, collected)
    texts = [p.text for p in doc.paragraphs]
    assert texts[1] == "Flags"
    assert texts[2].startswith("Status: NOT AUTHORIZED")
    assert "Needs your attention (3)" in texts or any(t.startswith("Needs your attention") for t in texts)


def test_clean_run_says_so(tmp_path):
    doc = Document()
    write_flags_section(doc, collect_flags(tmp_path))
    texts = [p.text for p in doc.paragraphs]
    assert "No flags: nothing needed attention." in texts


def test_strong_fundamentals_get_the_valuation_discussion(tmp_path):
    artifacts(tmp_path, bq=75.0)
    valuation = SimpleNamespace(roic_wacc=0.12, roce=0.2, current_price=100.0, company_value_per_share=140.0,
                                enterprise_mos=0.3, current_pe10=15.0, expected_annual_return=0.11)
    assert fundamentals_assessment(tmp_path, valuation)["strong"]
    assert valuation_assessment(tmp_path, valuation)["verdict"] == "cheap"
    doc = Document()
    write_assessment_sections(doc, tmp_path, valuation)
    text = "\n".join(p.text for p in doc.paragraphs)
    assert "STRONG." in text and "the company looks cheap" in text


def test_weak_fundamentals_skip_the_cheapness_verdict(tmp_path):
    artifacts(tmp_path, bq=65.0)
    doc = Document()
    write_assessment_sections(doc, tmp_path, SimpleNamespace(roic_wacc=0.2, roce=0.2, enterprise_mos=0.5))
    text = "\n".join(p.text for p in doc.paragraphs)
    assert "NOT RATED STRONG" in text and "not treated as an opportunity" in text and "looks cheap" not in text


def test_opinion_slot_is_labelled(tmp_path):
    artifacts(tmp_path, bq=75.0)
    doc = Document()
    write_assessment_sections(doc, tmp_path, SimpleNamespace(roic_wacc=0.1, enterprise_mos=0.4), opinion={"fundamentals": "Solid moat.", "valuation": "Cheap on cash flow."})
    text = "\n".join(p.text for p in doc.paragraphs)
    assert text.count("not a HAP rules-based result") == 2 and "Solid moat." in text
