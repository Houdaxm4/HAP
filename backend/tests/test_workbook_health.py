import json

from openpyxl import Workbook

from services.workbook_health import buyback_consistency, health_report, scan_errors


def make_workbook(path, buybacks):
    wb = Workbook()
    ws = wb.active
    ws.title = "Income - GAAP"
    for i, year in enumerate(range(2022, 2025)):
        ws.cell(3, 3 + i, f"FY {year}")
        ws.cell(88, 3 + i, buybacks[i])
    ws.cell(88, 1, "$ Paid for Shares (Millions)")
    ratios = wb.create_sheet("All Ratios")
    ratios["B2"] = "#DIV/0!"
    q = wb.create_sheet("Last Quarter IS Standardized")
    q["A1"] = "#NAME?"
    wb.save(path)


def test_scan_reports_core_errors_and_ignores_quarter_sheets_for_annual(tmp_path):
    path = tmp_path / "completed_workbook.xlsx"
    make_workbook(path, [10, 20, 30])
    annual = scan_errors(path)
    assert any(f.startswith("All Ratios") for f in annual["findings"])
    assert not any("Last Quarter" in f for f in annual["findings"])
    assert any("Last Quarter" in f for f in scan_errors(path, quarterly=True)["findings"])


def test_buyback_mismatch_flagged_and_matching_passes(tmp_path):
    path = tmp_path / "completed_workbook.xlsx"
    make_workbook(path, [10, 0, 30])
    report = {"years": [{"fiscal_year": "FY2022", "dollars": 10.2}, {"fiscal_year": "FY2023", "dollars": 25.0},
                        {"fiscal_year": "FY2024", "dollars": 30.0}]}
    result = buyback_consistency(path, report)
    assert result["checked"] and [m["fiscal_year"] for m in result["mismatches"]] == ["FY2023"]
    (tmp_path / "new_company_buyback_report.json").write_text(json.dumps(report))
    full = health_report(tmp_path)
    assert full["available"] and any("Buybacks" in f for f in full["findings"])


def test_health_report_without_workbook(tmp_path):
    assert health_report(tmp_path) == {"available": False, "findings": [], "notes": []}
