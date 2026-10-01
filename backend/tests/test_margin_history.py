from openpyxl import Workbook

from services.margin_history import noncomparable_years, note_text, read_history, table_rows


def build(path, costs):
    wb = Workbook()
    ws = wb.active
    ws.title = "Income - GAAP"
    keys = {9: "SALES_REV_TURN", 14: "IS_COGS_TO_FE_AND_PP_AND_G", 19: "GROSS_PROFIT", 30: "IS_OPER_INC", 58: "NET_INCOME"}
    for r, k in keys.items():
        ws.cell(r, 2, k)
    for i, cost in enumerate(costs):
        col = 3 + i
        ws.cell(3, col, f"FY {2020 + i}")
        ws.cell(9, col, 100.0 * (i + 1))
        ws.cell(14, col, cost)
        ws.cell(19, col, 100.0 * (i + 1) - (cost or 0))
        ws.cell(30, col, 20.0 * (i + 1))
        ws.cell(58, col, 10.0 * (i + 1))
    wb.save(path)


def test_blank_cost_year_is_flagged_and_others_are_not(tmp_path):
    path = tmp_path / "w.xlsx"
    build(path, [20, 40, None, 80, 100])
    history = read_history(path)
    assert noncomparable_years(history) == ["FY2022"]
    rows = table_rows(history)
    assert rows[2][3] == "not comparable" and rows[1][3] == "80.0%" and rows[0][2] == "n/a" and rows[1][2] == "100.0%"
    assert "FY2022" in note_text(history)


def test_clean_history_has_no_note(tmp_path):
    path = tmp_path / "w.xlsx"
    build(path, [20, 40, 60, 80, 100])
    assert note_text(read_history(path)) is None
