"""Regression: read-only CRF iter_rows yields EmptyCell without .coordinate."""

from __future__ import annotations

from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook

from services.annual_inputs_service import AnnualInputsService
from services.annual_period_service import align_fiscal_periods, detect_year_columns

JBSS_ANALYSIS_ID = "51b5e28e-6f16-400f-b2ed-f62dee714bd3"
JBSS_UPLOADS = (
    Path(__file__).resolve().parents[1] / "storage" / "uploads" / JBSS_ANALYSIS_ID
)


def _mini_inputs_workbook(path: Path, *, years: list[int]) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Inputs"
    ws["A7"] = "Line"
    for i, y in enumerate(years):
        ws.cell(7, 3 + i, f"FY{y}")
    ws["A10"] = "PE10"
    wb.save(path)
    wb.close()


def _sparse_crf(path: Path) -> None:
    """Mimic JBSS CRF: PE10 label with blank value column before Current PE10."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Custom Run Filter"
    ws["A133"] = "PE10"
    ws["A159"] = "Current PE10"
    ws["B159"] = 13.74729693930744
    wb.save(path)
    wb.close()


def test_pe10_fill_survives_readonly_empty_cells(tmp_path: Path) -> None:
    workbook = tmp_path / "workbook.xlsx"
    crf = tmp_path / "crf.xlsx"
    _mini_inputs_workbook(workbook, years=list(range(2017, 2027)))
    _sparse_crf(crf)

    # Seed CRF with parseable structure is heavy; service must not raise EmptyCell.
    # With a minimal CRF lacking annual metadata, fiscal PE10 may be missing — but
    # Current PE10 must never be written into the historical FY column.
    report = AnnualInputsService().apply(
        analysis_id="t-empty-cell",
        ticker="JBSS",
        workbook_path=workbook,
        custom_run_path=crf,
        new_fiscal_year="FY2026",
    )

    wb = load_workbook(workbook, data_only=True)
    cols = detect_year_columns(wb["Inputs"], wb)
    assert cols["FY2026"] == 12
    fy_val = wb["Inputs"].cell(10, cols["FY2026"]).value
    # Must not equal Current PE10 pasted into historical column.
    if fy_val is not None:
        assert fy_val != pytest.approx(13.74729693930744) or (
            report.pe10 is not None and report.pe10.field_role == "pe10_fiscal_year"
        )
    wb.close()


@pytest.mark.skipif(
    not (JBSS_UPLOADS / "prefilled_workbook.xlsx").exists()
    or not (JBSS_UPLOADS / "previous_workbook.xlsx").exists(),
    reason="JBSS stored UI inputs not available",
)
def test_jbss_fiscal_year_alignment_and_inputs_pe10() -> None:
    tmpl = JBSS_UPLOADS / "prefilled_workbook.xlsx"
    prev = JBSS_UPLOADS / "previous_workbook.xlsx"
    crf = next(JBSS_UPLOADS.glob("Custom_Run_Filter*.xlsx"))

    alignment = align_fiscal_periods(tmpl, prev)
    assert alignment.new_fiscal_year == "FY2026"
    assert alignment.dropped_years == ("FY2016",)

    wb_t = load_workbook(tmpl, data_only=False)
    wb_p = load_workbook(prev, data_only=False)
    try:
        tmpl_years = sorted(
            k for k in detect_year_columns(wb_t["Balance Sheet - Standardized"], wb_t) if k.startswith("FY")
        )
        prev_years = sorted(
            k for k in detect_year_columns(wb_p["Balance Sheet - Standardized"], wb_p) if k.startswith("FY")
        )
    finally:
        wb_t.close()
        wb_p.close()

    assert tmpl_years == [f"FY{y}" for y in range(2017, 2027)]
    assert prev_years == [f"FY{y}" for y in range(2016, 2026)]
    assert "COL" not in "".join(tmpl_years + prev_years)

    out = JBSS_UPLOADS.parent.parent / "outputs" / JBSS_ANALYSIS_ID / "empty_cell_regression.xlsx"
    out.parent.mkdir(parents=True, exist_ok=True)
    out.write_bytes(tmpl.read_bytes())

    report = AnnualInputsService().apply(
        analysis_id=JBSS_ANALYSIS_ID,
        ticker="JBSS",
        workbook_path=out,
        custom_run_path=crf,
        new_fiscal_year=alignment.new_fiscal_year,
    )
    assert report.pe10 is not None
    assert isinstance(report.pe10.source_value, (int, float))
