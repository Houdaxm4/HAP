"""Point-in-time custom run filter: scalars come from the quarter's history column and nothing later is kept."""

from __future__ import annotations

import datetime as dt
import sys
from pathlib import Path

from openpyxl import Workbook, load_workbook

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "scripts"))
import run_batch  # noqa: E402

from services.market_price_service import MarketPriceService  # noqa: E402
from services.point_in_time_filter_service import PointInTimeError, available_quarters, build_point_in_time  # noqa: E402

QUARTERS = ["2024 Q4", "2025 Q1", "2025 Q2"]
SERIES = {
    "Closing Stock Price (Q. Average)": [175.0, 150.0, 180.0],
    "E10": [5.0, 5.5, 5.6],
    "PE10": [35.0, 27.27, 32.14],
    "Lowest of 20th Percentile PE10 or 7PE10": [30.0, 30.0, 30.5],
    "Lowest of 45th Percentile PE10 or 7PE10": [34.0, 34.0, 34.0],
    "Enterprise Value (not-diluted)": [1000.0, 900.0, 1100.0],
    "Market Value of Common": [1100.0, 1000.0, 1200.0],
    "PE10 Quantile Satisfied (Interpolation)": ["No", "Yes", "No"],
    "Weighted Average Cost of Capital": [12.0, 12.5, 11.5],
}


def _crf(path: Path) -> Path:
    wb = Workbook()
    ws = wb.active
    ws.title = "ABC"
    ws["A1"], ws["B1"] = None, "ABC US Equity"
    ws["A15"] = "date"
    ws["A16"] = "Fiscal Quarter"
    for i, q in enumerate(QUARTERS):
        ws.cell(15, 2 + i).value = dt.datetime(2024 + (i + 3) // 4, ((i + 3) % 4) * 3 + 1, 28)
        ws.cell(16, 2 + i).value = q
    for r, (lab, vals) in enumerate(SERIES.items(), start=18):
        ws.cell(r, 1).value = lab
        for i, v in enumerate(vals):
            ws.cell(r, 2 + i).value = v
    start = 18 + len(SERIES) + 2
    for r, lab in enumerate(["Current E10", "Current PE10", "Max Current Price to Buy", "Current Price (Live Price)", "Current Graham Instrinsic Value"], start=start):
        ws.cell(r, 1).value, ws.cell(r, 2).value = lab, 999.0
    ws.cell(start + 8, 2).value = "PE10"
    ws.cell(start + 9, 1).value, ws.cell(start + 9, 2).value = 0, 12.0
    wb.create_sheet("Summary")
    wb.save(path)
    return path


def test_scalars_come_from_the_quarter_and_later_history_is_removed(tmp_path: Path):
    crf = _crf(tmp_path / "Custom_Run_Filter_2026-01-01-(10-00)-ABC.xlsx")
    assert available_quarters(crf) == QUARTERS
    info = build_point_in_time(crf, "2025 Q1", tmp_path / "pit.xlsx")
    assert info["price_quarter_average"] == 150.0 and info["e10"] == 5.5
    assert abs(info["max_price_to_buy"] - 30.0 * 5.5) < 1e-9 and abs(info["first_exit_price"] - 34.0 * 5.5) < 1e-9
    ws = load_workbook(tmp_path / "pit.xlsx")["ABC"]
    assert ws.cell(16, 4).value is None and ws.cell(18, 4).value is None          # 2025 Q2 column is gone
    values = {ws.cell(r, 1).value: ws.cell(r, 2).value for r in range(1, ws.max_row + 1) if ws.cell(r, 1).value}
    assert values["Current E10"] == 5.5 and values["Current Price (Live Price)"] == 150.0
    assert values["Max Current Price to Buy"] == 165.0
    assert values["Current Graham Instrinsic Value"] is None                       # no history row: left empty, nothing from a later date
    assert "POINT-IN-TIME" in ws.cell(13, 1).value


def test_missing_quarter_is_reported(tmp_path: Path):
    crf = _crf(tmp_path / "c.xlsx")
    try:
        build_point_in_time(crf, "2019 Q1", tmp_path / "x.xlsx")
    except PointInTimeError as exc:
        assert "2019 Q1" in str(exc)
    else:
        raise AssertionError("expected PointInTimeError")


def test_registered_price_replaces_the_live_quote():
    MarketPriceService.set_point_in_time("abc", 150.0, "pit")
    try:
        assert MarketPriceService().get_price("ABC") == (150.0, "pit")
    finally:
        MarketPriceService.clear_point_in_time("ABC")


def test_custom_filter_file_names_with_a_suffix_are_recognised():
    parse = run_batch._parse_crf_name
    assert parse(Path("Custom_Run_Filter_2025-03-26-(14-25)-MPWR - net_income_updated (1).xlsx"))[0] == "MPWR"
    assert parse(Path("Custom_Run_Filter_2026-02-05-(14-44)-EXP - Copy.xlsx"))[0] == "EXP"
    assert parse(Path("Custom_Run_Filter_2025-12-16-(16-48)-AMAT (1).xlsx"))[0] == "AMAT"
