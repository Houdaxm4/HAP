"""Notes placement, quarterly dependencies, buybacks, and deliverable names."""

from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook, load_workbook

from models.new_company import BuybackAbsenceClass, BuybackYearResult
from models.quarterly_presentation import QuarterlyStatementKind
from services.deliverable_naming import excel_deliverable_name
from services.hap_analysis_layout_service import HapAnalysisLayoutService, discover_occupied_end_row
from services.new_company_buyback_service import NewCompanyBuybackService
from services.quarterly_dependency_service import QuarterlyDependencyService
from services.sec_10q_statement_service import derive_standalone_from_ytd
from tests.ledger_util import entry, notes_text


def test_notes_sit_below_hidden_and_formula_rows_without_shifting(tmp_path: Path):
    path = tmp_path / "notes.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "Leases"
    ws["A10"] = "Estimated Long-Term Rate"
    ws["C10"] = "=B10"
    ws["B18"] = 0.045
    ws.row_dimensions[18].hidden = True
    ws["A40"] = "Existing analyst note"
    formula = ws["C10"].value
    wb.save(path)
    wb.close()

    loaded = load_workbook(path)
    sheet = loaded["Leases"]
    end = discover_occupied_end_row(sheet)
    assert end == 40
    written = HapAnalysisLayoutService().write_notes_section(
        sheet,
        [
            ("Lease discount rate (selected)", "4.50%"),
            ("ASC 842 WtdAvg vs model long-term rate", "Template row 10 remains Interest Expense / Total Debt."),
        ],
    )
    loaded.save(path)
    loaded.close()

    out = load_workbook(path)
    try:
        leases = out["Leases"]
        assert leases["C10"].value == formula
        assert leases["B18"].value == 0.045
        assert leases["A40"].value == "Existing analyst note"
        assert leases["A41"].value in (None, "")
        assert leases["A42"].value == "HAP ANALYSIS — NOTES"
        assert leases["A43"].value == "Lease discount rate (selected)"
        assert "Leases!B43" in written
        assert leases["A10"].value == "Estimated Long-Term Rate"
    finally:
        out.close()


def test_dependency_remap_follows_semantic_field_not_row(tmp_path: Path):
    path = tmp_path / "dep.xlsx"
    wb = Workbook()
    lq = wb.active
    lq.title = "Last Quarter IS Standardized"
    lq["A11"] = "Revenue"
    lq["C11"] = 10
    lq["A12"] = "Net Income"
    lq["C12"] = 2
    inputs = wb.create_sheet("Inputs")
    inputs["B5"] = "='Last Quarter IS Standardized'!$C$11"
    inputs["B6"] = "=C5"
    other = wb.create_sheet("Final Metrics")
    other["A1"] = "=C11"
    wb.save(path)
    wb.close()

    book = load_workbook(path)
    service = QuarterlyDependencyService()
    snapshot = service.snapshot(book)
    sheet = book["Last Quarter IS Standardized"]
    sheet["A11"] = "Net Income"
    sheet["C11"] = 2
    sheet["A12"] = "Revenue"
    sheet["C12"] = 10
    diff = service.reconnect(book, snapshot)
    book.save(path)
    book.close()

    out = load_workbook(path)
    try:
        assert out["Inputs"]["B5"].value == "='Last Quarter IS Standardized'!$C$12"
        assert out["Inputs"]["B6"].value == "=C5"
        assert out["Final Metrics"]["A1"].value == "=C11"
        assert diff["changed_formulas"]
        assert diff["changed_formulas"][0]["field"] == "revenue"
        assert not diff["unresolved_dependencies"]
    finally:
        out.close()


def test_component_row_is_not_retargeted_to_the_parent_total(tmp_path: Path):
    path = tmp_path / "oi.xlsx"
    wb = Workbook()
    lq = wb.active
    lq.title = "Last Quarter IS Standardized"
    lq["A23"] = "  + Other Operating Income"
    lq["C23"] = 0
    lq["A32"] = "Operating Income (Loss)"
    lq["C32"] = 18.51
    bs = wb.create_sheet("Last Quarter BS Standardized")
    bs["H15"] = "=IF('Last Quarter IS Standardized'!C32=\"\",0,'Last Quarter IS Standardized'!C32)"
    wb.save(path)
    wb.close()

    book = load_workbook(path)
    service = QuarterlyDependencyService()
    snapshot = service.snapshot(book)
    diff = service.reconnect(book, snapshot)
    assert book["Last Quarter BS Standardized"]["H15"].value == (
        "=IF('Last Quarter IS Standardized'!C32=\"\",0,'Last Quarter IS Standardized'!C32)"
    )
    assert not diff["changed_formulas"]
    assert not diff["unresolved_dependencies"]
    book.close()


def test_removed_quarter_field_is_an_unresolved_dependency(tmp_path: Path):
    wb = Workbook()
    lq = wb.active
    lq.title = "Last Quarter IS Standardized"
    lq["A11"] = "Revenue"
    lq["C11"] = 10
    inputs = wb.create_sheet("Inputs")
    inputs["B5"] = "='Last Quarter IS Standardized'!C11"
    service = QuarterlyDependencyService()
    snapshot = service.snapshot(wb)
    lq["A11"] = None
    lq["C11"] = None
    diff = service.reconnect(wb, snapshot)
    assert inputs["B5"].value == "='Last Quarter IS Standardized'!C11"
    assert diff["unresolved_dependencies"]
    assert diff["unresolved_dependencies"][0]["field"] == "revenue"


def test_q2_cash_flow_derivation_uses_same_year_ytd():
    derived = derive_standalone_from_ytd(
        {"val": 80_000_000, "_unit": "USD", "fp": "Q2"},
        {"val": 30_000_000, "_unit": "USD", "fp": "Q1"},
        tag="NetCashProvidedByUsedInOperatingActivities",
    )
    assert derived is not None
    assert derived[0] == 50
    assert "Q2" in derived[1] and "Q1" in derived[1]


def test_buyback_rows_87_88_fill_blanks_and_do_not_invent_zero(tmp_path: Path):
    path = tmp_path / "bb.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "Income Statement"
    for index, year in enumerate(range(2016, 2026)):
        ws.cell(7, 3 + index, f"FY{year}")
    ws["A87"] = "10-k Total Shares Buybacks (Millions)"
    ws["A88"] = "$ Paid for Shares (Millions)"
    ws["C88"] = "=C87"
    ws["D87"] = 9
    wb.save(path)
    wb.close()
    years = [
        BuybackYearResult(
            fiscal_year="FY2016",
            dollars=5,
            shares=1,
            dollars_source="sec_xbrl:PaymentsForRepurchaseOfCommonStock",
            shares_source="sec_xbrl:StockRepurchasedDuringPeriodShares",
        ),
        BuybackYearResult(
            fiscal_year="FY2017",
            dollars=None,
            shares=None,
            absence_class=BuybackAbsenceClass.NOT_DISCLOSED,
        ),
        BuybackYearResult(
            fiscal_year="FY2018",
            dollars=5,
            shares=2,
            dollars_source="sec_xbrl:PaymentsForRepurchaseOfCommonStock",
            shares_source="sec_xbrl:StockRepurchasedDuringPeriodShares",
        ),
    ]
    NewCompanyBuybackService()._write_buyback_schedule(
        path,
        years,
        write_policy="new_company",
        new_fiscal_year=None,
    )
    out = load_workbook(path)
    try:
        sheet = out["Income Statement"]
        assert sheet["C88"].value == "=C87"
        assert sheet["C87"].value == 1
        assert sheet["D87"].value == 9
        assert sheet["D88"].value in (None, "")
        assert sheet["E87"].value == 2
        assert sheet["E88"].value == 5
        notes = [
            sheet.cell(row, 1).value
            for row in range(1, (sheet.max_row or 1) + 1)
            if sheet.cell(row, 1).value == "Notes"
        ]
        assert notes
    finally:
        out.close()


def test_annual_buyback_preserves_history_and_fills_new_year(tmp_path: Path):
    path = tmp_path / "annual_bb.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "Income Statement"
    ws["C7"] = "FY2024"
    ws["D7"] = "FY2025"
    ws["A87"] = "10-k Total Shares Buybacks (Millions)"
    ws["A88"] = "$ Paid for Shares (Millions)"
    ws["C88"] = 4
    ws["D88"] = None
    wb.save(path)
    wb.close()
    years = [
        BuybackYearResult(
            fiscal_year="FY2024",
            dollars=99,
            shares=3,
            dollars_source="sec_xbrl:PaymentsForRepurchaseOfCommonStock",
            shares_source="sec",
        ),
        BuybackYearResult(
            fiscal_year="FY2025",
            dollars=8,
            shares=2,
            dollars_source="sec_xbrl:PaymentsForRepurchaseOfCommonStock",
            shares_source="sec",
        ),
    ]
    _written, discrepancies = NewCompanyBuybackService()._write_buyback_schedule(
        path,
        years,
        write_policy="annual_update",
        new_fiscal_year="FY2025",
    )
    out = load_workbook(path)
    try:
        sheet = out["Income Statement"]
        assert sheet["C88"].value == 4
        assert sheet["D88"].value == 8
        assert sheet["D87"].value == 2
        assert discrepancies
        assert discrepancies[0]["fiscal_year"] == "FY2024"
        assert sheet["C88"].comment is None and entry(out, "Income Statement", "C88", "Differs from filing")
    finally:
        out.close()


def test_missing_buyback_evidence_does_not_write_zero(tmp_path: Path):
    path = tmp_path / "missing_bb.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "Income - GAAP"
    ws["L7"] = "FY2026"
    ws["A87"] = "10-k Total Shares Buybacks (Millions)"
    ws["A88"] = "$ Paid for Shares (Millions)"
    ws["L88"] = 0
    wb.save(path)
    wb.close()
    year = BuybackYearResult(fiscal_year="FY2026", dollars=None, shares=None)
    NewCompanyBuybackService()._write_buyback_schedule(
        path,
        [year],
        write_policy="annual_update",
        new_fiscal_year="FY2026",
    )
    out = load_workbook(path)
    try:
        assert out["Income - GAAP"]["L87"].value in (None, "")
        assert out["Income - GAAP"]["L88"].value == 0
        assert year.write_action == "not_written_evidence_missing"
        assert year.dollars_cell == "Income - GAAP!L88"
        assert year.workbook_dollars == 0
    finally:
        out.close()


def test_deliverable_names_follow_incorporated_period():
    assert (
        excel_deliverable_name(
            fiscal_year=2026, ticker="lnn", analysis_type="New Company", fiscal_quarter=2
        )
        == "2026 Q2 LNN New Company.xlsx"
    )
    assert (
        excel_deliverable_name(
            fiscal_year=2026, ticker="idcc", analysis_type="Quarterly Update", fiscal_quarter=2
        )
        == "2026 Q2 IDCC Quarterly Update.xlsx"
    )
    assert (
        excel_deliverable_name(fiscal_year=2025, ticker="jbss", analysis_type="Annual Update")
        == "2025 Fiscal Year JBSS Annual Update.xlsx"
    )
    assert (
        excel_deliverable_name(fiscal_year=2025, ticker="lnn", analysis_type="New Company")
        == "2025 Fiscal Year LNN New Company.xlsx"
    )
    assert QuarterlyStatementKind.CASH_FLOW.value
