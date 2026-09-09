"""Focused tests for optimized workbook parsing (performance + semantic safety)."""

from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook

from services.quarterly_carry_forward_service import QuarterlyCarryForwardService
from services.workbook_service import (
    PARSER_SCHEMA_VERSION,
    WorkbookService,
    workbook_sha256,
)


def _mini_wb(path: Path) -> Path:
    wb = Workbook()
    ws = wb.active
    ws.title = "Income Statement"
    ws["A1"] = "Revenue"
    ws["B5"] = 1000
    ws["B10"] = "=B5*2"
    # Inflate declared dimensions with a far-right style-only touch if possible
    # (openpyxl may still expand max_column when writing far cells)
    ws["ZZ1"] = None
    hid = wb.create_sheet("HiddenStuff")
    hid.sheet_state = "hidden"
    hid["A1"] = 1
    wb.create_sheet("Inputs")
    wb["Inputs"]["C107"] = 0.21
    wb["Inputs"]["C80"] = "=C81+C82"
    wb.create_sheet("R&D")
    wb["R&D"]["B8"] = 3
    wb["R&D"]["E2"] = '=IF(Inputs!C103="",0,Inputs!C103)'
    wb.create_sheet("Leases")
    wb["Leases"]["B16"] = "=SUM(B14:B15)"
    wb.save(path)
    wb.close()
    return path


def _formula_index(structure) -> dict[tuple[str, str], str]:
    out: dict[tuple[str, str], str] = {}
    for sheet in structure.worksheets:
        for cell in sheet.cells:
            if cell.is_formula:
                out[(sheet.name, cell.address)] = cell.formula or str(cell.value)
    return out


def _value_index(structure) -> dict[tuple[str, str], object]:
    out: dict[tuple[str, str], object] = {}
    for sheet in structure.worksheets:
        for cell in sheet.cells:
            if not cell.is_formula:
                out[(sheet.name, cell.address)] = cell.value
    return out


def test_manifest_semantic_equivalence_two_parses(tmp_path: Path):
    path = _mini_wb(tmp_path / "a.xlsx")
    svc = WorkbookService(cache_dir=tmp_path / "cache", use_cache=False)
    a = svc.parse_structure(path, "a.xlsx")
    b = svc.parse_structure(path, "a.xlsx")
    assert a.worksheet_names == b.worksheet_names
    assert a.formula_count == b.formula_count == 4  # B10, C80, E2, B16
    assert _formula_index(a) == _formula_index(b)
    assert _value_index(a) == _value_index(b)
    # blanks are not enumerated
    assert all(c.data_type != "blank" for s in a.worksheets for c in s.cells)


def test_formulas_still_detected(tmp_path: Path):
    path = _mini_wb(tmp_path / "b.xlsx")
    svc = WorkbookService(cache_dir=tmp_path / "cache", use_cache=False)
    structure = svc.parse_structure(path, "b.xlsx")
    assert svc.cell_contains_formula(structure, "Income Statement", "B10") is True
    assert svc.cell_contains_formula(structure, "Income Statement", "B5") is False
    assert svc.cell_contains_formula(structure, "Inputs", "C80") is True


def test_cache_hit_identical_workbook(tmp_path: Path):
    path = _mini_wb(tmp_path / "c.xlsx")
    cache = tmp_path / "cache"
    svc = WorkbookService(cache_dir=cache, use_cache=True)
    first = svc.parse_structure(path, "c.xlsx")
    assert svc.cache_hit_count == 0
    assert svc.last_parse_timings.get("cache") == "miss"
    second = svc.parse_structure(path, "c.xlsx")
    assert svc.cache_hit_count == 1
    assert svc.last_parse_timings.get("cache") == "memory"
    assert _formula_index(first) == _formula_index(second)
    # Disk cache file exists
    digest = workbook_sha256(path)
    assert (cache / f"{digest}_{PARSER_SCHEMA_VERSION}.json").exists()


def test_cache_invalidation_when_workbook_changes(tmp_path: Path):
    path = _mini_wb(tmp_path / "d.xlsx")
    cache = tmp_path / "cache"
    svc = WorkbookService(cache_dir=cache, use_cache=True)
    before = svc.parse_structure(path, "d.xlsx")
    hits_before = svc.cache_hit_count

    from openpyxl import load_workbook

    wb = load_workbook(path)
    wb["Income Statement"]["B5"] = 9999
    wb.save(path)
    wb.close()

    after = svc.parse_structure(path, "d.xlsx")
    # New content → cache miss (hit count unchanged on this call)
    assert svc.cache_hit_count == hits_before
    assert svc.last_parse_timings.get("cache") == "miss"
    assert _value_index(after)[("Income Statement", "B5")] == 9999
    assert _value_index(before)[("Income Statement", "B5")] == 1000


def test_no_repeated_uncached_parse_within_service_when_cached(tmp_path: Path):
    path = _mini_wb(tmp_path / "e.xlsx")
    svc = WorkbookService(cache_dir=tmp_path / "cache", use_cache=True)
    svc.parse_structure(path, "e.xlsx")
    svc.parse_structure(path, "e.xlsx")
    svc.parse_structure(path, "e.xlsx")
    assert svc.parse_call_count == 3
    assert svc.cache_hit_count == 2


def test_meaningful_used_range_skips_inflated_columns(tmp_path: Path):
    path = tmp_path / "inflated.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "Income - GAAP"
    ws["A1"] = "Revenue"
    ws["N10"] = 42
    # Force a very high declared column via a far cell then clear value
    ws.cell(1, 500).value = "tmp"
    ws.cell(1, 500).value = None
    wb.save(path)
    wb.close()

    svc = WorkbookService(cache_dir=tmp_path / "cache", use_cache=False)
    structure = svc.parse_structure(path, "inflated.xlsx")
    sheet = next(s for s in structure.worksheets if s.name == "Income - GAAP")
    # Must find real values; must not enumerate millions of blanks
    assert any(c.address == "N10" and c.value == 42 for c in sheet.cells)
    assert all(c.data_type != "blank" for c in sheet.cells)
    assert len(sheet.cells) < 100


def test_hidden_sheets_and_named_ranges_preserved(tmp_path: Path):
    path = _mini_wb(tmp_path / "f.xlsx")
    from openpyxl import load_workbook

    wb = load_workbook(path)
    wb.defined_names.add(
        __import__("openpyxl.workbook.defined_name", fromlist=["DefinedName"]).DefinedName(
            "RevCell", attr_text="'Income Statement'!$B$5"
        )
    )
    wb.save(path)
    wb.close()

    svc = WorkbookService(cache_dir=tmp_path / "cache", use_cache=False)
    structure = svc.parse_structure(path, "f.xlsx")
    assert "HiddenStuff" in structure.hidden_sheets
    assert "Income Statement" in structure.visible_sheets
    assert any(nr.name == "RevCell" for nr in structure.named_ranges)


def test_quarter_carry_forward_still_works_with_optimized_parse(tmp_path: Path):
    prev = _mini_wb(tmp_path / "prev.xlsx")
    new = _mini_wb(tmp_path / "new.xlsx")
    from openpyxl import load_workbook

    wb = load_workbook(new)
    wb["Inputs"]["C107"] = None
    wb.save(new)
    wb.close()

    dest = tmp_path / "out.xlsx"
    report = QuarterlyCarryForwardService().apply(
        analysis_id="parse-opt",
        ticker="AAPL",
        new_template_path=new,
        previous_workbook_path=prev,
        destination_path=dest,
    )
    assert report.carry_forward_count >= 1
    out = load_workbook(dest)
    assert out["Inputs"]["C107"].value == 0.21
    assert isinstance(out["Inputs"]["C80"].value, str) and out["Inputs"]["C80"].value.startswith("=")
    out.close()
