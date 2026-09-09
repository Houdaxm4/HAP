"""Certification hardening: YTD provenance, Word research, final continuity."""

from __future__ import annotations

from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook

from models.quarterly_presentation import QuarterlyStatementKind
from models.quarterly_update import CarryForwardDecision
from services.quarterly_model_continuity_service import QuarterlyModelContinuityService
from services.quarterly_research_service import QuarterlyResearchService
from services.yahoo_quarterly_statement_service import (
    DERIVED_QUARTER_SUM,
    DIRECT_REPORTED,
    SECONDARY_REPORTED,
    SOURCE_MISSING,
    YahooQuarterSnapshot,
    YahooQuarterlyBundle,
    YahooQuarterlyStatementService,
)


@pytest.fixture
def mock_sec_is_ytd(monkeypatch):
    """Patch SEC YTD map to return known revenue YTD."""

    def _fake_map(kind, company_facts, fiscal_year, fiscal_period):
        if kind == QuarterlyStatementKind.INCOME:
            return {"totalRevenue": 364357.0, "netIncome": 101464.0}
        if kind == QuarterlyStatementKind.CASH_FLOW:
            return {"totalCashFromOperatingActivities": 116996.0}
        return {}

    monkeypatch.setattr(
        YahooQuarterlyStatementService,
        "_sec_ytd_map",
        staticmethod(_fake_map),
    )


def _is_bundle(*quarter_values: float) -> YahooQuarterlyBundle:
    snaps = []
    for i, v in enumerate(quarter_values):
        snaps.append(
            YahooQuarterSnapshot(
                end_date=f"2024-{9-i*3:02d}-30",
                fields={"totalRevenue": v, "netIncome": v * 0.2},
            )
        )
    return YahooQuarterlyBundle(income_quarters=snaps)


def _cf_bundle(ytd_cfo: float) -> YahooQuarterlyBundle:
    return YahooQuarterlyBundle(
        cashflow_quarters=[
            YahooQuarterSnapshot(
                end_date="2024-09-30",
                fields={"totalCashFromOperatingActivities": ytd_cfo},
            )
        ]
    )


def test_sec_ytd_beats_derived_quarter_sum(mock_sec_is_ytd):
    bundle = _is_bundle(109417.0, 85777.0, 90753.0)  # Q3, Q2, Q1 standalone
    svc = YahooQuarterlyStatementService()
    result = svc.values_for_periods(
        bundle,
        QuarterlyStatementKind.INCOME,
        fiscal_quarter=3,
        company_facts={"facts": {}},
        fiscal_year=2024,
        fiscal_period="Q3",
    )
    prov = result.ytd_provenance["totalRevenue"]
    assert prov.provenance == SECONDARY_REPORTED
    assert prov.value == pytest.approx(364357.0)
    assert prov.components is None


def test_derived_quarter_sum_when_no_sec(mock_sec_is_ytd, monkeypatch):
    monkeypatch.setattr(
        YahooQuarterlyStatementService,
        "_sec_ytd_map",
        staticmethod(lambda *a, **k: {}),
    )
    bundle = _is_bundle(100.0, 90.0, 80.0)
    svc = YahooQuarterlyStatementService()
    result = svc.values_for_periods(
        bundle, QuarterlyStatementKind.INCOME, fiscal_quarter=3
    )
    prov = result.ytd_provenance["totalRevenue"]
    assert prov.provenance == DERIVED_QUARTER_SUM
    assert prov.components == [100.0, 90.0, 80.0]
    assert prov.value == pytest.approx(270.0)


def test_cf_ytd_direct_reported_not_quarter_sum():
    bundle = _cf_bundle(116996.0)
    svc = YahooQuarterlyStatementService()
    result = svc.values_for_periods(
        bundle, QuarterlyStatementKind.CASH_FLOW, fiscal_quarter=3
    )
    prov = result.ytd_provenance["totalCashFromOperatingActivities"]
    assert prov.provenance == DIRECT_REPORTED
    assert prov.period_identity == "Q3_9M_YTD"
    assert prov.value == pytest.approx(116996.0)


def test_q2_period_identity(mock_sec_is_ytd):
    bundle = _is_bundle(200.0, 180.0)
    svc = YahooQuarterlyStatementService()
    result = svc.values_for_periods(
        bundle,
        QuarterlyStatementKind.INCOME,
        fiscal_quarter=2,
        company_facts={"facts": {}},
        fiscal_year=2024,
        fiscal_period="Q2",
    )
    assert result.ytd_provenance["totalRevenue"].period_identity == "Q2_6M_YTD"


def test_research_prioritizes_sec_and_flags_missing_call():
    manifest = {
        "selected_filings": [
            {
                "filing_type": "8-K",
                "filing_date": "2024-10-31",
                "document_url": "https://sec.gov/x",
            }
        ]
    }
    report = QuarterlyResearchService().gather(
        analysis_id="r1",
        ticker="ZZZZ",
        fiscal_year=2024,
        fiscal_quarter=3,
        sec_manifest=manifest,
    )
    assert any(s.source_kind == "sec_8k" for s in report.sources)
    assert report.sources[0].url == "https://sec.gov/x"
    assert report.earnings_call_status == "EARNINGS_CALL_SOURCE_UNAVAILABLE"
    assert report.management_explanations == [] or report.earnings_call_status != "AVAILABLE"


def test_research_rejects_template_webcast_as_transcript():
    """IR webcast placeholders must not count as earnings-call transcript."""
    report = QuarterlyResearchService().gather(
        analysis_id="r3",
        ticker="AAPL",
        fiscal_year=2026,
        fiscal_quarter=3,
        sec_manifest={"selected_filings": [{"filing_type": "10-Q", "filing_date": "2026-07-31"}]},
    )
    assert report.official_release_used is True
    for m in report.management_explanations:
        assert "{{" not in m
    if report.earnings_call_status == "AVAILABLE":
        assert any(
            len(m) >= 200 or "operator:" in m.lower()
            for m in report.management_explanations
        )


def test_research_yahoo_news_ticker_filtered(monkeypatch):
    payload = {
        "news": [
            {"title": "What To Expect From Viking's (VIK) Q2 Earnings", "link": "https://x"},
            {"title": "AAPL Q3 earnings beat estimates", "link": "https://y", "summary": "Apple revenue"},
        ]
    }

    monkeypatch.setattr(QuarterlyResearchService, "_get_json", staticmethod(lambda url: payload))
    items = QuarterlyResearchService()._yahoo_news("AAPL", 2026, 3)
    titles = [i.title for i in items]
    assert any("AAPL" in t for t in titles)
    assert not any("Viking" in t for t in titles)


def test_research_earnings_call_not_from_generic_news(monkeypatch):
    def _fake_news(self, ticker, fy, q):
        return []

    def _fake_call(self, ticker, fy, q, ir_url):
        from models.quarterly_update import ResearchSourceEntry

        return ResearchSourceEntry(
            source_kind="earnings_call_secondary",
            title="AAPL Q3 transcript",
            url="https://seekingalpha.com/article/123",
            snippet="Operator: Good afternoon. Tim Cook discussed services growth and iPhone demand.",
            reliability="secondary",
        )

    monkeypatch.setattr(QuarterlyResearchService, "_yahoo_news", _fake_news)
    monkeypatch.setattr(QuarterlyResearchService, "_find_earnings_call", _fake_call)
    report = QuarterlyResearchService().gather(
        analysis_id="r2", ticker="AAPL", fiscal_year=2024, fiscal_quarter=3
    )
    assert report.earnings_call_status == "AVAILABLE_SECONDARY_TRANSCRIPT"
    assert report.management_explanations
    assert "Operator:" in report.management_explanations[0] or "Tim Cook" in report.management_explanations[0]


def test_final_continuity_array_formula_match(tmp_path: Path):
    """ArrayFormula cells must compare by text, not object identity."""
    from openpyxl.worksheet.formula import ArrayFormula

    template = tmp_path / "template_af.xlsx"
    out = tmp_path / "out_af.xlsx"
    formula = ArrayFormula(
        ref="K2",
        text="=VALUE(LEFT(INDEX('Balance Sheet - Standardized'!C$5:L$5, ROW()-1), 4))",
    )
    wb = Workbook()
    wb.active.title = "DividendHelper"
    wb["DividendHelper"]["K2"] = formula
    wb.create_sheet("Last Quarter IS As Reported")
    wb.save(template)
    wb.close()
    import shutil

    shutil.copy2(template, out)
    wb2 = load_workbook(out)
    wb2["DividendHelper"]["K2"] = ArrayFormula(
        ref="K2",
        text="=VALUE(LEFT(INDEX('Balance Sheet - Standardized'!C$5:L$5, ROW()-1), 4))",
    )
    wb2.save(out)
    wb2.close()
    report = QuarterlyModelContinuityService().verify_final_deliverable(
        analysis_id="fc2",
        ticker="AAPL",
        new_template_path=template,
        workbook_path=out,
    )
    assert report.status == "ok"
    assert "DividendHelper" in report.ignored_sheets_verified


def test_final_continuity_after_restore(tmp_path: Path):
    template = tmp_path / "template.xlsx"
    out = tmp_path / "out.xlsx"
    wb = Workbook()
    wb.active.title = "DividendHelper"
    wb["DividendHelper"]["K2"] = "TEMPLATE"
    ig = wb.create_sheet("Last Quarter IS As Reported")
    ig["A1"] = "As Reported"
    wb.save(template)
    wb.close()
    import shutil

    shutil.copy2(template, out)
    wb2 = load_workbook(out)
    wb2["DividendHelper"]["K2"] = "MODIFIED_BY_FILL"
    wb2.save(out)
    wb2.close()
    svc = QuarterlyModelContinuityService()
    svc.restore_ignored_sheets(new_template_path=template, workbook_path=out)
    report = svc.verify_final_deliverable(
        analysis_id="fc1",
        ticker="AAPL",
        new_template_path=template,
        workbook_path=out,
    )
    assert report.phase == "final"
    assert report.status == "ok"
    assert report.ignored_sheet_mismatches == []
    assert "DividendHelper" in report.ignored_sheets_verified
