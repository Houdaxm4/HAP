"""Company-specific valuation growth analysis — no generic 8% selector."""

from __future__ import annotations

from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook

from models.annual_update import ValuationAssumptionAnalysis
from services.annual_growth_analysis_service import AnnualGrowthAnalysisService
from services.annual_judgment_service import AnnualJudgmentService
from services.annual_analyst_intelligence_service import AnnualAnalystIntelligenceService


def _geo(start: float, rate: float, n: int) -> list[float]:
    return [start * ((1 + rate) ** i) for i in range(n)]


def _valuation_workbook(
    path: Path,
    *,
    years: int = 10,
    start_year: int = 2017,
    eps: list[float] | None = None,
    revenue: list[float] | None = None,
    oi: list[float] | None = None,
    ni: list[float] | None = None,
    da: list[float] | None = None,
    capex: list[float] | None = None,
    a11: float = 0.03,
    c5: float = 0.16,
    a14: float = 0.1875,
    b6_total: float | None = None,
    c6_annual: float | None = None,
    c31_garbage: float = -2.21,
    shares: list[float] | None = None,
) -> None:
    n = years
    eps = eps or _geo(3.0, 0.04, n)
    revenue = revenue or _geo(800.0, 0.037, n)
    oi = oi or _geo(60.0, 0.04, n)
    ni = ni or _geo(36.0, 0.04, n)
    da = da or [15.0 + i * 0.5 for i in range(n)]
    capex = capex or [-12.0 - i * 0.4 for i in range(n)]
    shares = shares or [11.4 + i * 0.03 for i in range(n)]
    if b6_total is None:
        oe0 = ni[0] + da[0] + capex[0]
        oe1 = ni[-1] + da[-1] + capex[-1]
        b6_total = oe1 / oe0 - 1.0 if oe0 else 0.0
    if c6_annual is None:
        c6_annual = (1.0 + b6_total) ** (1.0 / (n - 1)) - 1.0 if b6_total is not None else 0.0

    wb = Workbook()
    inp = wb.active
    inp.title = "Inputs"
    fy = [f"FY {start_year + i}" for i in range(n)]
    for i, label in enumerate(fy):
        inp.cell(1, 3 + i, label)
        inp.cell(7, 3 + i, label)
    inp["A38"] = "Consolidated Net Income"
    inp["A44"] = "Depreciation & Amortization"
    inp["A46"] = "Capital Expenditures"
    for i in range(n):
        inp.cell(38, 3 + i, ni[i])
        inp.cell(44, 3 + i, da[i])
        inp.cell(46, 3 + i, capex[i])
    inp["B63"] = 80.0

    inc = wb.create_sheet("Income - GAAP")
    inc["A9"] = "Revenue"
    inc["A30"] = "Operating Income (Loss)"
    inc["A70"] = "Diluted Weighted Avg Shares"
    inc["A71"] = "Diluted EPS, GAAP"
    inc["A105"] = "EPS FY 2008"
    inc["C105"] = -0.56
    for i, label in enumerate(fy):
        inc.cell(7, 3 + i, label)
        inc.cell(9, 3 + i, revenue[i])
        inc.cell(30, 3 + i, oi[i])
        inc.cell(70, 3 + i, shares[i])
        inc.cell(71, 3 + i, eps[i])

    fm = wb.create_sheet("Final Metrics")
    fm["A4"] = "ROE"
    fm["A31"] = "EPS 10 Year CAGR"
    fm["A40"] = "Operating Margin"
    for i, label in enumerate(fy):
        fm.cell(1, 3 + i, label)
        fm.cell(4, 3 + i, a14)
        fm.cell(40, 3 + i, oi[i] / revenue[i] if revenue[i] else None)
    fm["C29"] = eps[0]
    fm["L29"] = eps[-1]
    fm["C31"] = c31_garbage
    fm["L31"] = (eps[-1] / eps[0]) ** (1 / (n - 1)) - 1 if eps[0] > 0 else None
    fm["C10"] = ni[0] + da[0] + capex[0]
    fm["L10"] = ni[-1] + da[-1] + capex[-1]
    fm["B4"] = a14

    er = wb.create_sheet("Expected Returns & Buybacks")
    er["C5"] = c5
    er["A14"] = a14
    er["A11"] = a11
    er["B5"] = fm["L31"].value
    er["B8"] = 30.0
    er["A2"] = 70.0
    er["E2"] = 18.0
    er["C14"] = "=B14*E2"
    er["B11"] = "=B8*((1+A11)^10)"
    er["B14"] = "=B11*A14"
    er["E14"] = "=(C14/A2)^(1/10)-1"
    for r in range(17, 27):
        yr = r - 16
        er[f"B{r}"] = yr
        er[f"C{r}"] = f"=$B$8*((1+$A$11)^B{r})"
        er[f"D{r}"] = f"=C{r}*$A$14"
        er[f"E{r}"] = f"=D{r}*0.4"
    er["C17"] = "=$B$8*((1+$A$11)^B17)"

    ev = wb.create_sheet("Enterprise Value")
    ev["B6"] = b6_total
    ev["C6"] = c6_annual
    ev["B8"] = 1
    ev["B9"] = "='Final Metrics'!$L$10*(1+$C$6)^B8"
    ev["B20"] = 50
    ev["B27"] = 0.2
    ev["B32"] = 40
    ev["B40"] = eps[-1]
    ev["B41"] = fm["L31"].value
    ev["B42"] = "=(B51+(B41*100*B52))*B40"
    ev["B51"] = 8.5
    ev["B52"] = 2.0

    wb.save(path)
    wb.close()


def _load_analyses(path: Path) -> dict[str, ValuationAssumptionAnalysis]:
    wb = load_workbook(path, data_only=True)
    wf = load_workbook(path, data_only=False)
    try:
        return AnnualGrowthAnalysisService().analyze_workbooks(wb, wf)
    finally:
        wb.close()
        wf.close()


def test_ev_uses_c6_forecast_driver_not_b6_total(tmp_path: Path):
    path = tmp_path / "wb.xlsx"
    _valuation_workbook(path, b6_total=-0.95, c6_annual=-0.28)
    oe = _load_analyses(path)["oe"]
    assert oe.actual_formula_driver == "Enterprise Value!C6"
    assert oe.existing_assumption == pytest.approx(-0.28)
    assert oe.historical_observations["oe_total_change_b6"] == pytest.approx(-0.95)
    assert oe.existing_assumption_grain == "annualized_growth_rate"
    assert oe.historical_observations["oe_total_change_grain"] == "total_historical_change"


def test_graham_does_not_read_c31_as_5y_or_revenue(tmp_path: Path):
    path = tmp_path / "wb.xlsx"
    eps = _geo(3.17, 0.04, 10)
    # last 6 points nearly flat so true 5y is ~0 not -221%
    eps = [3.17, 2.84, 3.43, 4.69, 5.17, 5.33, 5.40, 5.15, 5.03, 5.26]
    _valuation_workbook(path, eps=eps, c31_garbage=-2.212)
    gr = _load_analyses(path)["graham"]
    assert gr.historical_observations["fm_c31_ignored"] == pytest.approx(-2.212)
    assert gr.historical_observations["eps_5y_cagr"] != pytest.approx(-2.212)
    assert gr.historical_observations["revenue_cagr"] != pytest.approx(-2.212)
    # 5y from FY2021 5.17 to FY2026 5.26
    assert gr.historical_observations["eps_5y_cagr"] == pytest.approx((5.26 / 5.17) ** (1 / 5) - 1, rel=1e-3)
    assert gr.historical_observations["eps_10y_cagr"] == pytest.approx((5.26 / 3.17) ** (1 / 9) - 1, rel=1e-3)


def test_true_revenue_cagr_from_revenue_series(tmp_path: Path):
    path = tmp_path / "wb.xlsx"
    rev = _geo(846.6, 0.037, 10)
    _valuation_workbook(path, revenue=rev, c31_garbage=-2.21)
    analyses = _load_analyses(path)
    expected = (rev[-1] / rev[0]) ** (1 / 9) - 1
    assert analyses["graham"].historical_observations["revenue_cagr"] == pytest.approx(expected, rel=1e-3)
    assert analyses["er"].historical_observations["revenue_cagr"] == pytest.approx(expected, rel=1e-3)
    assert analyses["oe"].historical_observations["revenue_cagr"] == pytest.approx(expected, rel=1e-3)


def test_eight_percent_not_candidate_or_fallback(tmp_path: Path):
    path = tmp_path / "wb.xlsx"
    # CapEx spike collapses OE; NI/OI stay ~flat. Must not invent 8%.
    ni = [60.0] * 10
    da = [20.0] * 10
    capex = [-15.0] * 8 + [-50.0, -88.0]
    oi = [85.0] * 10
    eps = [5.2] * 10
    rev = _geo(900.0, 0.03, 10)
    _valuation_workbook(path, ni=ni, da=da, capex=capex, oi=oi, eps=eps, revenue=rev, a11=0.027, b6_total=-0.95, c6_annual=-0.28)
    analyses = _load_analyses(path)
    for rec in analyses.values():
        assert rec.selected_prospective_rate != pytest.approx(0.08)
        joined = " ".join(rec.evidence) + rec.rationale + (rec.selection_method or "")
        assert "0.08 fallback" not in joined
        assert rec.selection_method != "generic_band"
    assert analyses["oe"].decision == "INSUFFICIENT_EVIDENCE"
    assert analyses["oe"].selected_prospective_rate is None
    assert any("capex" in d for d in analyses["oe"].distortions_identified)


def test_anomaly_thresholds_do_not_select_rate(tmp_path: Path):
    path = tmp_path / "wb.xlsx"
    _valuation_workbook(path, a11=0.40, eps=_geo(1.0, 0.06, 10), revenue=_geo(100.0, 0.055, 10), oi=_geo(10.0, 0.06, 10))
    er = _load_analyses(path)["er"]
    assert er.decision == "ADJUST"
    assert er.selected_prospective_rate != pytest.approx(0.08)
    assert er.selected_prospective_rate is not None
    assert 0.04 < er.selected_prospective_rate < 0.08
    assert er.selection_method == "median_of_corroborating_cluster"
    assert "generic" not in (er.selection_method or "")


def test_er_uses_a11_projection_driver_and_can_keep(tmp_path: Path):
    path = tmp_path / "wb.xlsx"
    eps = [3.17, 2.84, 3.43, 4.69, 5.17, 5.33, 5.40, 5.15, 5.03, 5.26]
    rev = _geo(847.0, 0.037, 10)
    oi = _geo(60.5, 0.044, 10)
    _valuation_workbook(path, eps=eps, revenue=rev, oi=oi, a11=0.0267, c5=0.143, a14=0.1866)
    er = _load_analyses(path)["er"]
    assert er.actual_formula_driver == "Expected Returns & Buybacks!A11"
    assert er.existing_assumption == pytest.approx(0.0267, rel=1e-3)
    assert er.existing_assumption_grain == "sustainable_g_retention_x_roe"
    assert er.decision == "KEEP_EXISTING"
    assert er.decision_basis
    assert "not sufficiently material" in (er.decision_basis or "").lower() or "consistent with" in (er.decision_basis or "").lower()
    assert "reasonable range" not in er.rationale.lower() or "0–18" not in er.rationale
    assert er.prospective_range_low is not None and er.prospective_range_high is not None
    assert er.prospective_range_low - 1e-9 <= er.existing_assumption <= er.prospective_range_high + 1e-9


def test_er_adjusts_when_evidence_cluster_disagrees(tmp_path: Path):
    path = tmp_path / "wb.xlsx"
    _valuation_workbook(
        path,
        a11=0.40,
        eps=_geo(2.0, 0.06, 10),
        revenue=_geo(200.0, 0.058, 10),
        oi=_geo(20.0, 0.062, 10),
    )
    er = _load_analyses(path)["er"]
    assert er.decision == "ADJUST"
    assert er.selected_prospective_rate == pytest.approx(0.06, abs=0.005)
    assert er.selection_method == "median_of_corroborating_cluster"
    assert er.provenance


def test_ev_decomposes_capex_and_can_be_insufficient(tmp_path: Path):
    path = tmp_path / "wb.xlsx"
    ni = [36, 33, 39, 54, 60, 62, 63, 60, 59, 62]
    da = [16, 15, 17, 18, 18, 18, 21, 25, 27, 28]
    capex = [-11, -13, -15, -15, -25, -18, -21, -28, -51, -88]
    _valuation_workbook(
        path,
        ni=ni,
        da=da,
        capex=capex,
        oi=[60, 56, 59, 79, 85, 87, 90, 85, 85, 89],
        eps=[3.17, 2.84, 3.43, 4.69, 5.17, 5.33, 5.4, 5.15, 5.03, 5.26],
        b6_total=-0.95,
        c6_annual=-0.283,
        a11=0.027,
    )
    oe = _load_analyses(path)["oe"]
    decomp = oe.historical_observations["oe_decomposition"]
    assert decomp[-1]["capex"] == pytest.approx(-88)
    assert decomp[-1]["ni"] == pytest.approx(62)
    assert any("capex" in d or "reinvestment" in d for d in oe.distortions_identified)
    assert "DISTORTED_BASE" in oe.distortions_identified
    assert oe.decision == "INSUFFICIENT_EVIDENCE"
    assert oe.selected_prospective_rate is None


def test_graham_independent_of_oe_and_can_keep(tmp_path: Path):
    path = tmp_path / "wb.xlsx"
    eps = [3.17, 2.84, 3.43, 4.69, 5.17, 5.33, 5.40, 5.15, 5.03, 5.26]
    _valuation_workbook(
        path,
        eps=eps,
        revenue=_geo(847.0, 0.037, 10),
        oi=_geo(60.5, 0.044, 10),
        ni=[36, 33, 39, 54, 60, 62, 63, 60, 59, 62],
        capex=[-11, -13, -15, -15, -25, -18, -21, -28, -51, -88],
        da=[16, 15, 17, 18, 18, 18, 21, 25, 27, 28],
        a11=0.027,
        b6_total=-0.95,
        c6_annual=-0.283,
    )
    analyses = _load_analyses(path)
    assert analyses["graham"].metric == "graham_eps_growth"
    assert analyses["oe"].decision != analyses["graham"].decision or analyses["er"].decision != analyses["oe"].decision
    assert analyses["graham"].historical_observations["eps_10y_cagr"] == pytest.approx((5.26 / 3.17) ** (1 / 9) - 1, rel=1e-3)
    # Three records are independently populated.
    rates = {
        analyses["er"].decision,
        analyses["oe"].decision,
        analyses["graham"].decision,
    }
    assert len(rates) >= 2


def test_range_selection_is_not_automatic_midpoint(tmp_path: Path):
    path = tmp_path / "wb.xlsx"
    # EPS/revenue cluster near 2.1%; operating income compounds at 10%.
    # Range is wide (~2–10%); selected rate must be the 2% cluster, not the midpoint.
    eps = _geo(2.0, 0.021, 10)
    rev = _geo(100.0, 0.022, 10)
    oi = _geo(10.0, 0.10, 10)
    _valuation_workbook(path, a11=0.20, eps=eps, revenue=rev, oi=oi)
    er = _load_analyses(path)["er"]
    assert er.decision == "ADJUST"
    selected = er.selected_prospective_rate
    assert selected is not None
    lo, hi = er.prospective_range_low, er.prospective_range_high
    assert lo is not None and hi is not None
    mid = (lo + hi) / 2
    assert selected == pytest.approx(0.021, abs=0.005)
    assert abs(selected - mid) > 0.02
    assert er.selection_method == "median_of_corroborating_cluster"


def test_existing_inside_range_preserved(tmp_path: Path):
    path = tmp_path / "wb.xlsx"
    _valuation_workbook(path, a11=0.04, eps=_geo(2.0, 0.05, 10), revenue=_geo(100.0, 0.03, 10), oi=_geo(10.0, 0.045, 10))
    er = _load_analyses(path)["er"]
    assert er.decision == "KEEP_EXISTING"
    assert er.selected_prospective_rate == pytest.approx(0.04)


def test_distorted_endpoint_cagr_not_projected(tmp_path: Path):
    path = tmp_path / "wb.xlsx"
    ni = [60.0] * 10
    capex = [-12.0] * 8 + [-50.0, -90.0]
    _valuation_workbook(path, ni=ni, capex=capex, da=[18.0] * 10, b6_total=-0.9, c6_annual=-0.22, a11=0.03)
    oe = _load_analyses(path)["oe"]
    usable_note = oe.selection_method or ""
    assert "oe_full" not in (oe.selection_method or "")
    assert oe.decision == "INSUFFICIENT_EVIDENCE"
    assert oe.selected_prospective_rate is None
    assert usable_note.startswith("insufficient")


def test_adjust_has_range_provenance_and_method(tmp_path: Path):
    path = tmp_path / "wb.xlsx"
    _valuation_workbook(path, a11=0.35, eps=_geo(1.0, 0.07, 10), revenue=_geo(50.0, 0.065, 10), oi=_geo(8.0, 0.068, 10))
    er = _load_analyses(path)["er"]
    assert er.decision == "ADJUST"
    assert er.prospective_range_low is not None
    assert er.prospective_range_high is not None
    assert er.selection_method
    assert er.provenance
    assert er.selected_prospective_rate is not None
    assert "evidence" in er.rationale.lower() or "cluster" in (er.selection_method or "")


def test_keep_has_evidence_rationale(tmp_path: Path):
    path = tmp_path / "wb.xlsx"
    _valuation_workbook(path, a11=0.035, eps=_geo(3.0, 0.05, 10), revenue=_geo(800.0, 0.03, 10), oi=_geo(60.0, 0.04, 10))
    er = _load_analyses(path)["er"]
    assert er.decision == "KEEP_EXISTING"
    assert "2.67%" not in er.rationale  # not a hardcoded JBSS sentence
    assert _pct_in_rationale(er.rationale, 0.035)
    assert er.prospective_range_low is not None


def test_original_cells_untouched_and_hap_adjacent_only(tmp_path: Path):
    path = tmp_path / "wb.xlsx"
    _valuation_workbook(path, a11=0.40, eps=_geo(1.0, 0.06, 10), revenue=_geo(100.0, 0.06, 10), oi=_geo(10.0, 0.06, 10))
    ctx = AnnualAnalystIntelligenceService().build_judgment_context(
        analysis_id="a1", ticker="ZZ", workbook_path=path
    )
    orig = load_workbook(path)
    b5 = orig["Expected Returns & Buybacks"]["B5"].value
    d17 = orig["Expected Returns & Buybacks"]["D17"].value
    b6 = orig["Enterprise Value"]["B6"].value
    c6 = orig["Enterprise Value"]["C6"].value
    orig.close()
    _, judge = AnnualJudgmentService().apply(
        analysis_id="a1", ticker="ZZ", workbook_path=path, context=ctx
    )
    wb = load_workbook(path)
    assert wb["Expected Returns & Buybacks"]["B5"].value == b5
    assert wb["Expected Returns & Buybacks"]["D17"].value == d17
    assert wb["Enterprise Value"]["B6"].value == b6
    assert wb["Enterprise Value"]["C6"].value == c6
    assert judge.hap_analysis_cells
    assert "Expected Returns & Buybacks!B5" in judge.original_cells_preserved
    wb.close()


def test_independent_metrics_need_not_share_a_rate(tmp_path: Path):
    path = tmp_path / "wb.xlsx"
    _valuation_workbook(
        path,
        a11=0.027,
        eps=_geo(3.0, 0.04, 10),
        revenue=_geo(800.0, 0.037, 10),
        oi=_geo(60.0, 0.04, 10),
        ni=[60] * 10,
        capex=[-12] * 8 + [-55, -90],
        b6_total=-0.92,
        c6_annual=-0.25,
    )
    analyses = _load_analyses(path)
    selected = {
        "er": analyses["er"].selected_prospective_rate if analyses["er"].decision == "ADJUST" else analyses["er"].existing_assumption,
        "oe": analyses["oe"].selected_prospective_rate,
        "graham": analyses["graham"].selected_prospective_rate if analyses["graham"].decision == "ADJUST" else analyses["graham"].existing_assumption,
    }
    assert analyses["oe"].decision == "INSUFFICIENT_EVIDENCE"
    assert selected["oe"] is None
    assert selected["er"] != selected["oe"]


def _pct_in_rationale(text: str, rate: float) -> bool:
    return f"{rate:.1%}" in text or f"{rate*100:.1f}" in text


def _hap_cells(ws, prefix="HAP"):
    found = {}
    for row in ws.iter_rows(min_row=1, max_row=30, max_col=24):
        for cell in row:
            lab = cell.value
            if isinstance(lab, str) and prefix in lab:
                neighbor = ws.cell(cell.row, cell.column + 1).value
                found[lab] = neighbor
    return found


def test_hap_growth_never_written_into_roe_semantic_field(tmp_path: Path):
    path = tmp_path / "wb.xlsx"
    _valuation_workbook(path, a11=0.40, eps=_geo(1.0, 0.06, 10), revenue=_geo(100.0, 0.06, 10), oi=_geo(10.0, 0.06, 10))
    ctx = AnnualAnalystIntelligenceService().build_judgment_context(
        analysis_id="a1", ticker="ZZ", workbook_path=path
    )
    _, judge = AnnualJudgmentService().apply(analysis_id="a1", ticker="ZZ", workbook_path=path, context=ctx)
    assert judge.er_analysis and judge.er_analysis.decision == "ADJUST"
    hap_rate = judge.er_analysis.selected_prospective_rate
    wb = load_workbook(path)
    er = wb["Expected Returns & Buybacks"]
    labels = _hap_cells(er)
    roe_val = next(v for k, v in labels.items() if "HAP ROE" in k)
    g_val = next(v for k, v in labels.items() if "prospective growth" in k or "HAP prospective growth" in k)
    wb.close()
    assert hap_rate is not None
    if isinstance(roe_val, (int, float)):
        assert roe_val != pytest.approx(hap_rate)
    else:
        assert isinstance(roe_val, str)
        assert "A14" in roe_val.replace("$", "")
        assert f"{hap_rate}" not in str(roe_val)
    assert g_val == pytest.approx(hap_rate)


def test_er_parallel_preserves_semantic_roles_and_e14(tmp_path: Path):
    path = tmp_path / "wb.xlsx"
    _valuation_workbook(path, a11=0.40, c5=0.16, a14=0.1875, eps=_geo(1.0, 0.06, 10), revenue=_geo(100.0, 0.06, 10), oi=_geo(10.0, 0.06, 10))
    ctx = AnnualAnalystIntelligenceService().build_judgment_context(
        analysis_id="a1", ticker="ZZ", workbook_path=path
    )
    er_rep, judge = AnnualJudgmentService().apply(analysis_id="a1", ticker="ZZ", workbook_path=path, context=ctx)
    assert judge.er_analysis.decision == "ADJUST"
    assert er_rep.semantic_substitution
    assert "ROE" in (er_rep.semantic_substitution or "")
    assert er_rep.hap_expected_return_cell
    assert (er_rep.hap_mechanics or {}).get("methodology") == "bv_x_roe"
    wb = load_workbook(path)
    er = wb["Expected Returns & Buybacks"]
    labels = _hap_cells(er)
    wb.close()
    assert any("implied retention" in k.lower() for k in labels)
    assert any("HAP ROE" in k for k in labels)
    assert any("Expected Return" in k for k in labels)
    assert any("BV path" in str(v) or "BV path" in k for k, v in labels.items()) or any("HAP BV" in k for k in labels)


def test_er_cannot_certify_adjust_without_coherent_parallel_model(tmp_path: Path):
    path = tmp_path / "wb.xlsx"
    _valuation_workbook(path, a11=0.40, eps=_geo(1.0, 0.06, 10), revenue=_geo(100.0, 0.06, 10), oi=_geo(10.0, 0.06, 10))
    wb = load_workbook(path)
    er = wb["Expected Returns & Buybacks"]
    er["A2"] = None
    er["E2"] = None
    er["B8"] = None
    wb.save(path)
    wb.close()
    ctx = AnnualAnalystIntelligenceService().build_judgment_context(
        analysis_id="a1", ticker="ZZ", workbook_path=path
    )
    # Analysis may still want ADJUST; judgment must not certify it without a model.
    er_rep, judge = AnnualJudgmentService().apply(analysis_id="a1", ticker="ZZ", workbook_path=path, context=ctx)
    assert judge.er_analysis.decision == "INSUFFICIENT_EVIDENCE"
    assert judge.expected_return.decision == "INSUFFICIENT_EVIDENCE"
    assert "INSUFFICIENT" in (judge.er_analysis.decision_basis or judge.er_analysis.rationale)


def test_slightly_outside_range_does_not_automatically_adjust(tmp_path: Path):
    """Long-term EPS ~6.43% vs independent max ~5.90% is not an automatic ADJUST."""
    path = tmp_path / "wb.xlsx"
    start, full, y5 = 1.90, 0.0643, 0.059
    end = start * ((1 + full) ** 9)
    mid5 = end / ((1 + y5) ** 5)
    r_early = (mid5 / start) ** (1 / 4)
    r_late = (end / mid5) ** (1 / 5)
    eps = [start * (r_early ** i) if i <= 4 else mid5 * (r_late ** (i - 4)) for i in range(10)]
    _valuation_workbook(
        path,
        eps=eps,
        revenue=_geo(100.0, 0.031, 10),
        oi=_geo(20.0, 0.028, 10),
        a11=0.03,
    )
    wb = load_workbook(path)
    wb["Enterprise Value"]["B41"] = (eps[-1] / eps[0]) ** (1 / 9) - 1
    wb.save(path)
    wb.close()
    gr = _load_analyses(path)["graham"]
    assert gr.existing_assumption == pytest.approx((eps[-1] / eps[0]) ** (1 / 9) - 1, rel=1e-3)
    assert gr.prospective_range_high is not None
    assert gr.existing_assumption > gr.prospective_range_high - 1e-12
    assert gr.decision == "KEEP_EXISTING"
    assert "CSCO" not in (gr.decision_basis or "")
    assert "ETD" not in (gr.rationale or "")


def test_material_conflict_can_adjust_even_when_gap_is_small(tmp_path: Path):
    """Opposite-sign operating/recent evidence can ADJUST even <1pp outside the range."""
    path = tmp_path / "wb.xlsx"
    # Existing slightly positive; all independent windows slightly negative and clustered.
    eps = [2.0]
    for _ in range(9):
        eps.append(eps[-1] * 0.992)  # ~ -0.8% 10y; 3y/5y also negative
    _valuation_workbook(
        path,
        eps=eps,
        revenue=_geo(100.0, -0.03, 10),
        oi=_geo(20.0, -0.028, 10),
        a11=0.03,
    )
    wb = load_workbook(path)
    ev = wb["Enterprise Value"]
    ev["B41"] = 0.006  # +0.6%, about 0.8pp above a -0.2% upper independent bound
    wb.save(path)
    wb.close()
    gr = _load_analyses(path)["graham"]
    assert gr.existing_assumption == pytest.approx(0.006)
    assert gr.decision == "ADJUST"
    assert "CSCO" not in gr.rationale and "ETD" not in gr.rationale


def test_keep_outside_numerical_range_when_historical_assumption_remains_evidence(tmp_path: Path):
    path = tmp_path / "wb.xlsx"
    eps = [2.0 * (1.06 ** i) for i in range(10)]
    _valuation_workbook(path, eps=eps, revenue=_geo(80.0, 0.03, 10), oi=_geo(12.0, 0.028, 10), a11=0.03)
    full = (eps[-1] / eps[0]) ** (1 / 9) - 1
    wb = load_workbook(path)
    wb["Enterprise Value"]["B41"] = full
    wb.save(path)
    wb.close()
    gr = _load_analyses(path)["graham"]
    assert _same_or_keep(gr)
    assert gr.historical_observations["eps_10y_cagr"] == pytest.approx(full, rel=1e-3)
    # Independent range excludes eps_full, so existing may sit outside it.
    assert gr.decision == "KEEP_EXISTING"


def _same_or_keep(gr):
    from services.annual_growth_analysis_service import _same_observation
    return _same_observation(gr.existing_assumption, gr.historical_observations.get("eps_10y_cagr"))


def test_adjust_inside_broad_numerical_range_when_contradicted(tmp_path: Path):
    path = tmp_path / "wb.xlsx"
    # Noisy positive 10y revenue widens the independent range, but recent EPS and
    # operating windows are persistently negative — existing 5% sits inside the
    # broad min-max and is still economically contradicted.
    eps = [3.0, 3.2, 3.5, 3.8, 4.0, 3.2, 2.5, 2.0, 1.6, 1.3]
    rev = [40.0, 50.0, 60.0, 80.0, 120.0, 200.0, 180.0, 150.0, 110.0, 90.0]
    oi = [5.0, 6.0, 8.0, 12.0, 18.0, 22.0, 18.0, 14.0, 10.0, 8.0]
    _valuation_workbook(path, eps=eps, revenue=rev, oi=oi, a11=0.03)
    wb = load_workbook(path)
    wb["Enterprise Value"]["B41"] = 0.05
    wb.save(path)
    wb.close()
    gr = _load_analyses(path)["graham"]
    lo, hi = gr.prospective_range_low, gr.prospective_range_high
    assert lo is not None and hi is not None
    assert lo <= 0.05 <= hi
    assert gr.decision == "ADJUST"


def test_no_fixed_materiality_tolerance_heuristic():
    src = Path(__file__).resolve().parents[1] / "services" / "annual_growth_analysis_service.py"
    text = src.read_text(encoding="utf-8")
    decide = text.split("def _decide(")[1].split("def _decision_sentence")[0]
    assert "0.01" not in decide or "single_observation" in decide
    assert "within 1" not in decide.lower()
    assert "0.10" not in decide
    assert "10%" not in decide
    assert "if existing" in decide.lower() or "existing is not None" in decide
    assert "lo <= existing <= hi" not in decide or "keep_support" in decide


def test_etd_csco_decisions_are_not_hard_coded():
    src = Path(__file__).resolve().parents[1] / "services" / "annual_growth_analysis_service.py"
    text = src.read_text(encoding="utf-8")
    assert "CSCO" not in text
    assert "ETD" not in text
    assert "JBSS" not in text
    assert "MZTI" not in text
    src2 = Path(__file__).resolve().parents[1] / "services" / "annual_judgment_service.py"
    text2 = src2.read_text(encoding="utf-8")
    assert "CSCO" not in text2 and "ETD" not in text2


def test_originals_untouched_and_hap_adjacent_after_er_clone(tmp_path: Path):
    path = tmp_path / "wb.xlsx"
    _valuation_workbook(path, a11=0.40, eps=_geo(1.0, 0.06, 10), revenue=_geo(100.0, 0.06, 10), oi=_geo(10.0, 0.06, 10))
    orig = load_workbook(path)
    snapshot = {
        "B5": orig["Expected Returns & Buybacks"]["B5"].value,
        "A11": orig["Expected Returns & Buybacks"]["A11"].value,
        "A14": orig["Expected Returns & Buybacks"]["A14"].value,
        "C5": orig["Expected Returns & Buybacks"]["C5"].value,
        "D17": orig["Expected Returns & Buybacks"]["D17"].value,
        "E14": orig["Expected Returns & Buybacks"]["E14"].value,
        "C17": orig["Expected Returns & Buybacks"]["C17"].value,
    }
    orig.close()
    ctx = AnnualAnalystIntelligenceService().build_judgment_context(
        analysis_id="a1", ticker="ZZ", workbook_path=path
    )
    _, judge = AnnualJudgmentService().apply(analysis_id="a1", ticker="ZZ", workbook_path=path, context=ctx)
    wb = load_workbook(path)
    er = wb["Expected Returns & Buybacks"]
    for addr, val in snapshot.items():
        assert er[addr].value == val
    assert judge.hap_analysis_cells
    assert all("!A11" not in c or c.endswith("A11") is False or "HAP" in str(er[c.split("!")[-1]].comment.text if False else "") for c in judge.hap_analysis_cells)
    wb.close()
