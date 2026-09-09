"""Focused tests for valuation / 25% MOS / entry-price review."""

from __future__ import annotations

from pathlib import Path

import pytest
from openpyxl import Workbook

from models.valuation_validation import ValuationAttractiveness, ValuationDecision
from services.valuation_validation_service import (
    ValuationValidationService,
    classify_attractiveness,
    entry_price_at_mos,
    graham_intrinsic,
    margin_of_safety,
    sensitivity_table,
)


def test_exact_valuation_reconstruction():
    assert graham_intrinsic(multiple=39.025, eps=7.46) == pytest.approx(291.1265, rel=1e-4)
    from services.valuation_validation_service import graham_multiple_from_cagr

    assert graham_multiple_from_cagr(eps_cagr=0.15262661592678528) == pytest.approx(
        39.02532318535705, rel=1e-6
    )


def test_missing_current_price_source_missing(tmp_path: Path):
    path = _val_wb(tmp_path / "noprice.xlsx", price=None, iv=200.0, multiple=20.0, eps=10.0)
    report = ValuationValidationService().validate(
        analysis_id="v1", ticker="T", workbook_path=path
    )
    assert report.current_price_decision == ValuationDecision.SOURCE_MISSING
    assert report.attractiveness == ValuationAttractiveness.INDETERMINATE
    assert report.decision == ValuationDecision.SOURCE_MISSING


def test_live_current_price_override_used(tmp_path: Path):
    path = _val_wb(tmp_path / "live.xlsx", price=None, iv=200.0, multiple=20.0, eps=10.0)
    report = ValuationValidationService().validate(
        analysis_id="v2",
        ticker="T",
        workbook_path=path,
        current_price_override=150.0,
        current_price_source="market_internet:yahoo_chart",
    )
    assert report.current_price == 150.0
    assert report.current_price_source == "market_internet:yahoo_chart"
    assert report.margin_of_safety == pytest.approx(0.25)


def test_positive_and_negative_valuation_gap():
    assert margin_of_safety(intrinsic=200, price=150) == pytest.approx(0.25)
    assert margin_of_safety(intrinsic=200, price=250) == pytest.approx(-0.25)


def test_mos_below_25_wait_or_near_entry():
    assert classify_attractiveness(0.10) == ValuationAttractiveness.WAIT
    assert classify_attractiveness(0.22) == ValuationAttractiveness.NEAR_ENTRY


def test_mos_at_least_25_attractive():
    assert classify_attractiveness(0.25) == ValuationAttractiveness.ATTRACTIVE
    assert classify_attractiveness(0.40) == ValuationAttractiveness.ATTRACTIVE


def test_premium_to_intrinsic_expensive():
    assert classify_attractiveness(-0.05) == ValuationAttractiveness.EXPENSIVE


def test_entry_price_calculation():
    assert entry_price_at_mos(intrinsic=200.0) == pytest.approx(150.0)


def test_valuation_sensitivity():
    sens = sensitivity_table(base_multiple=20.0, eps=10.0, price=160.0)
    assert len(sens["cases"]) >= 5
    assert "robust_to_adverse" in sens


def test_no_workbook_writes(tmp_path: Path):
    path = _val_wb(tmp_path / "nw.xlsx", price=100.0, iv=200.0, multiple=20.0, eps=10.0)
    before = path.read_bytes()
    ValuationValidationService().validate(analysis_id="v3", ticker="T", workbook_path=path)
    assert path.read_bytes() == before


def test_end_to_end_attractive(tmp_path: Path):
    # IV 200, price 140 → MOS 30% ≥ 25%
    path = _val_wb(tmp_path / "ok.xlsx", price=140.0, iv=200.0, multiple=20.0, eps=10.0)
    report = ValuationValidationService().validate(
        analysis_id="v4", ticker="T", workbook_path=path
    )
    assert report.intrinsic_decision == ValuationDecision.VALIDATED
    assert report.independent_intrinsic_value == pytest.approx(200.0)
    assert report.margin_of_safety == pytest.approx(0.30)
    assert report.required_entry_price == pytest.approx(150.0)
    assert report.attractiveness == ValuationAttractiveness.ATTRACTIVE
    review = ValuationValidationService().build_analyst_review(report)
    assert review.attractiveness == ValuationAttractiveness.ATTRACTIVE


def test_end_to_end_expensive(tmp_path: Path):
    path = _val_wb(tmp_path / "exp.xlsx", price=250.0, iv=200.0, multiple=20.0, eps=10.0)
    report = ValuationValidationService().validate(
        analysis_id="v5", ticker="T", workbook_path=path
    )
    assert report.attractiveness == ValuationAttractiveness.EXPENSIVE
    assert report.meets_mos_threshold is False
    assert report.pct_decline_to_entry == pytest.approx((250 - 150) / 250)


def _val_wb(
    path: Path,
    *,
    price: float | None,
    iv: float,
    multiple: float,
    eps: float,
) -> Path:
    wb = Workbook()
    wb.remove(wb.active)
    fm = wb.create_sheet("Final Metrics")
    inp = wb.create_sheet("Inputs")
    income = wb.create_sheet("Income - GAAP")
    ev = wb.create_sheet("Enterprise Value")
    fm["L32"] = multiple
    fm["L29"] = eps
    fm["L33"] = iv
    fm["B33"] = iv
    fm["B32"] = multiple
    fm["B29"] = eps
    fm["B51"] = price
    fm["L7"] = 0.09
    # 10y EPS ending at eps for optional reconstruction
    for i in range(10):
        income.cell(71, 3 + i).value = eps * (0.7 + 0.03 * i)
    income["L71"] = eps
    ev["B51"] = 8.5
    ev["B52"] = 2
    if price is not None:
        inp["B63"] = price
    else:
        inp["B63"] = None
    inp["B66"] = 18.0
    inp["B67"] = iv * 0.75
    inp["B68"] = iv * 1.1
    wb.save(path)
    wb.close()
    return path
