"""The projection reads ROIC (not the ROIC - WACC spread) from the Final Metrics tab."""

from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook

from services.new_company_projection_service import NewCompanyProjectionService, final_metrics_rows


def _final_metrics(path: Path, labels: dict[int, str], col: int = 3) -> Path:
    wb = Workbook()
    ws = wb.active
    ws.title = "Final Metrics"
    ws.cell(1, col).value = "FY2025"
    for row, label in labels.items():
        ws.cell(row, 1).value = label
    wb.save(path)
    return path


def test_rows_found_by_label_not_by_position(tmp_path: Path):
    template = _final_metrics(tmp_path / "a.xlsx", {3: "ROA", 4: "ROE", 5: "ROCE", 6: "ROIC Including Goodwill", 7: "WACC  Bloomberg", 8: "ROIC in - WACC"})
    from openpyxl import load_workbook

    assert final_metrics_rows(load_workbook(template)["Final Metrics"]) == (5, 6)
    shifted = _final_metrics(tmp_path / "b.xlsx", {4: "ROCE", 5: "ROIC Including Goodwill", 6: "WACC", 7: "ROIC in - WACC"})
    assert final_metrics_rows(load_workbook(shifted)["Final Metrics"]) == (4, 5)
    assert final_metrics_rows(None) == (5, 6)


def test_latest_roic_is_the_roic_row_not_the_spread(tmp_path: Path):
    path = _final_metrics(tmp_path / "wb.xlsx", {5: "ROCE", 6: "ROIC Including Goodwill", 7: "WACC  Bloomberg", 8: "ROIC in - WACC"})
    from openpyxl import load_workbook

    wb = load_workbook(path)
    ws = wb["Final Metrics"]
    ws["C5"], ws["C6"], ws["C7"], ws["C8"] = 0.102, 0.078, 0.104, -0.026
    wb.save(path)
    result = NewCompanyProjectionService._annual_returns(path, ["FY2025"])
    assert result["latest_roic"] == 0.078          # not -0.026, which is ROIC minus WACC
    assert result["latest_roce"] == 0.102
