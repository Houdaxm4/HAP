"""ROIC adjustments: Inputs formulas for operating assets/liabilities, Income tab cells for one-time items, all logged."""

from __future__ import annotations

from datetime import datetime
from pathlib import Path

from openpyxl import Workbook, load_workbook

from services.roic_adjustment_service import (
    RoicAdjustmentService,
    add_term,
    facts_for_year,
    remove_term,
)

BSN = "Balance Sheet - Standardized"


def _term(col: str, row: int) -> str:
    return f"IF('{BSN}'!{col}{row}=\"\",0,'{BSN}'!{col}{row})"


def _facts(tag: str, value: float, end: str = "2024-12-31", start: str = "2024-01-01") -> dict:
    return {"facts": {"us-gaap": {tag: {"units": {"USD": [
        {"start": start, "end": end, "val": value, "form": "10-K", "filed": "2025-02-20"},
        {"start": "2023-01-01", "end": "2023-12-31", "val": 1.0, "form": "10-K", "filed": "2025-02-20"},
    ]}}}}}


def _workbook(path: Path, *, other_opex=15.0, sga=40.0, opex_total=120.0, held_for_sale_in_formula=False, unbilled=5.0, accrued=2.0) -> Path:
    wb = Workbook()
    inputs = wb.active
    inputs.title = "Inputs"
    bs = wb.create_sheet(BSN)
    inc = wb.create_sheet("Income - GAAP")
    for ws in (inputs, bs, inc):
        ws["A1"] = "Year"
        ws["C1"], ws["D1"] = "FY2023", "FY2024"
    inc["A7"] = "In Millions of USD except Per Share"
    inc["C8"], inc["D8"] = datetime(2023, 12, 31), datetime(2024, 12, 31)
    for col, oi in (("C", 90.0), ("D", 100.0)):
        inc[f"{col}30"] = oi
        inc[f"{col}22"] = opex_total
        inc[f"{col}23"] = sga
        inc[f"{col}28"] = other_opex
    for col in ("C", "D"):
        bs[f"{col}11"] = 100.0
        bs[f"{col}14"] = 50.0
        bs[f"{col}19"] = unbilled   # unbilled revenues
        bs[f"{col}29"] = 25.0   # assets held for sale
        bs[f"{col}65"] = 30.0
        bs[f"{col}66"] = accrued   # accrued taxes
        extra = "+" + _term(col, 29) if held_for_sale_in_formula else ""
        inputs[f"{col}81"] = "=" + _term(col, 11) + "+" + _term(col, 14) + extra
        inputs[f"{col}82"] = "=0"
        inputs[f"{col}84"] = "=" + _term(col, 65)
        inputs[f"{col}85"] = "=0"
        inputs[f"{col}80"] = 500.0
        inputs[f"{col}83"] = 200.0
    wb.save(path)
    return path


def test_term_helpers_round_trip():
    f = "=" + _term("D", 11) + "+" + _term("D", 14)
    assert remove_term(f, "D", 14) == "=" + _term("D", 11)
    assert remove_term(f, "D", 11) == "=" + _term("D", 14)
    assert add_term(f, "D", 19).endswith("+" + _term("D", 19))


def test_facts_pick_the_fiscal_year_by_period_end():
    facts = _facts("RestructuringCharges", 12_000_000)
    from datetime import date
    assert facts_for_year(facts, "RestructuringCharges", date(2024, 12, 31)) == 12_000_000
    assert facts_for_year(facts, "RestructuringCharges", date(2022, 12, 31)) is None
    assert facts_for_year(facts, "MissingTag", date(2024, 12, 31)) is None


def test_one_time_charge_is_removed_from_component_and_total(tmp_path: Path):
    path = _workbook(tmp_path / "wb.xlsx")
    report = RoicAdjustmentService().apply(workbook_path=path, company_facts=_facts("RestructuringCharges", 12_000_000), newest_only_fy="FY2024")
    ws = load_workbook(path)["Income - GAAP"]
    assert ws["D28"].value == 3.0 and ws["D22"].value == 108.0
    assert ws["C28"].value == 15.0 and ws["C22"].value == 120.0  # earlier year untouched in annual mode
    assert "ADJ-" in ws["D28"].comment.text
    ledger = load_workbook(path)["HAP Adjustments"]
    cats = {ledger.cell(r, 5).value for r in range(5, 8) if ledger.cell(r, 1).value}
    assert cats == {"One-time operating income"}
    assert any(a.amount == 12.0 for a in report.adjustments)


def test_immaterial_item_is_ignored(tmp_path: Path):
    path = _workbook(tmp_path / "wb.xlsx")
    report = RoicAdjustmentService().apply(workbook_path=path, company_facts=_facts("RestructuringCharges", 1_000_000), newest_only_fy="FY2024")
    assert not [a for a in report.adjustments if a.category == "One-time operating income"]


def test_one_time_gain_is_added_back_only_when_visible_in_opex(tmp_path: Path):
    path = _workbook(tmp_path / "wb.xlsx", other_opex=-8.0)
    RoicAdjustmentService().apply(workbook_path=path, company_facts=_facts("GainLossOnSaleOfPropertyPlantEquipment", 8_000_000), newest_only_fy="FY2024")
    ws = load_workbook(path)["Income - GAAP"]
    assert ws["D28"].value == 0.0 and ws["D22"].value == 128.0
    path2 = _workbook(tmp_path / "wb2.xlsx", other_opex=5.0)
    report = RoicAdjustmentService().apply(workbook_path=path2, company_facts=_facts("GainLossOnSaleOfPropertyPlantEquipment", 8_000_000), newest_only_fy="FY2024")
    assert not report.changed and any("not visible" in s for s in report.skipped)


def test_operating_lines_are_added_and_non_operating_removed(tmp_path: Path):
    path = _workbook(tmp_path / "wb.xlsx", held_for_sale_in_formula=True, unbilled=60.0, accrued=20.0)
    report = RoicAdjustmentService().apply(workbook_path=path, newest_only_fy="FY2024")
    ws = load_workbook(path)["Inputs"]
    assert "29" not in ws["D81"].value                      # held-for-sale taken out of operating assets
    assert "D19" in ws["D81"].value                          # unbilled revenues (60 of 500 = 12%) added
    assert "D66" in ws["D84"].value                          # accrued taxes (20 of 200 = 10%) added
    assert "29" in ws["C81"].value                           # earlier year untouched
    cats = {a.category for a in report.adjustments}
    assert cats == {"Operating assets", "Operating liabilities"}


def test_new_company_adjusts_every_year(tmp_path: Path):
    path = _workbook(tmp_path / "wb.xlsx")
    facts = _facts("RestructuringCharges", 12_000_000)
    facts["facts"]["us-gaap"]["RestructuringCharges"]["units"]["USD"].append(
        {"start": "2023-01-01", "end": "2023-12-31", "val": 11_000_000, "form": "10-K", "filed": "2025-02-21"}
    )
    RoicAdjustmentService().apply(workbook_path=path, company_facts=facts)
    ws = load_workbook(path)["Income - GAAP"]
    assert ws["D22"].value == 108.0 and ws["C22"].value == 109.0


def test_overlapping_impairment_tags_are_counted_once(tmp_path: Path):
    from services.roic_adjustment_service import family_amount

    assert family_amount("impairment", {"AssetImpairmentCharges": 13.2, "ImpairmentOfIntangibleAssetsExcludingGoodwill": 13.2}) == (13.2, ["AssetImpairmentCharges"])
    amount, used = family_amount("impairment", {"GoodwillImpairmentLoss": 5.0, "ImpairmentOfLongLivedAssetsHeldForUse": 2.0, "ImpairmentOfIntangibleAssetsExcludingGoodwill": 3.0})
    assert amount == 8.0 and "GoodwillImpairmentLoss" in used
    path = _workbook(tmp_path / "wb.xlsx")
    facts = _facts("AssetImpairmentCharges", 12_000_000)
    facts["facts"]["us-gaap"]["ImpairmentOfIntangibleAssetsExcludingGoodwill"] = facts["facts"]["us-gaap"]["AssetImpairmentCharges"]
    RoicAdjustmentService().apply(workbook_path=path, company_facts=facts, newest_only_fy="FY2024")
    assert load_workbook(path)["Income - GAAP"]["D22"].value == 108.0  # 120 - 12, not 120 - 24
