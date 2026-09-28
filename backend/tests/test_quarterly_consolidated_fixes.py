"""Quarterly Update consolidated fixes: retrieval, carry-forward, ROIC format, valuation parity."""

from __future__ import annotations

from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook

from models.quarterly_presentation import PresentationDecision, QuarterlyStatementKind
from models.quarterly_update import CarryForwardDecision
from services.annual_analyst_intelligence_service import AnnualAnalystIntelligenceService
from services.annual_judgment_service import AnnualJudgmentService
from services.circular_reference_service import CircularReferenceService
from services.quarterly_carry_forward_service import QuarterlyCarryForwardService
from services.quarterly_presentation_service import QuarterlyPresentationService
from services.quarterly_projection_service import QuarterlyProjectionService
from services.sec_10q_statement_service import extract_sec_10q_statement
from test_annual_growth_analysis import _valuation_workbook
from test_quarterly_certification import _mini_proj_wb


def _usd_entry(*, val, fy, fp, form, start=None, end=None, filed="2024-11-01", accn="0001-24-000001"):
    row = {
        "val": val,
        "fy": fy,
        "fp": fp,
        "form": form,
        "filed": filed,
        "accn": accn,
        "end": end,
    }
    if start:
        row["start"] = start
    return row


def _facts(tags: dict[str, list[dict]]) -> dict:
    us_gaap = {}
    for tag, entries in tags.items():
        us_gaap[tag] = {"units": {"USD": entries}}
    return {"facts": {"us-gaap": us_gaap}}


def _geo(start: float, rate: float, n: int) -> list[float]:
    return [start * ((1 + rate) ** i) for i in range(n)]


def _blank_stmt_wb(path: Path, *, is_blank=True, cf_blank=True) -> Path:
    wb = Workbook()
    wb.active.title = "Income - GAAP"
    for name in (
        "Last Quarter IS Standardized",
        "Last Quarter BS Standardized",
        "Last Quarter CF Standardized",
    ):
        wb.create_sheet(name)
    if not is_blank:
        ws = wb["Last Quarter IS Standardized"]
        ws["A11"] = "Revenue"
        ws["C11"] = 10
    if not cf_blank:
        ws = wb["Last Quarter CF Standardized"]
        ws["A11"] = "Cash from Operating Activities"
        ws["C11"] = 10
    bs = wb["Last Quarter BS Standardized"]
    bs["A11"] = "Cash"
    bs["C11"] = 100
    bs["A12"] = "Total current assets"
    bs["C12"] = 200
    bs["A13"] = "Total assets"
    bs["C13"] = 500
    bs["A14"] = "Total liabilities"
    bs["C14"] = 200
    bs["A15"] = "Total shareholders equity"
    bs["C15"] = 300
    for r in range(20, 32):
        bs.cell(r, 1, f"BS Line {r}")
        bs.cell(r, 3, 1.0)
    wb.save(path)
    wb.close()
    return path


def _q2_cf_facts() -> dict:
    return _facts({
        "NetCashProvidedByUsedInOperatingActivities": [
            _usd_entry(val=40_000_000, fy=2024, fp="Q1", form="10-Q", start="2024-01-01", end="2024-03-31"),
            _usd_entry(val=100_000_000, fy=2024, fp="Q2", form="10-Q", start="2024-01-01", end="2024-06-30"),
        ],
        "RevenueFromContractWithCustomerExcludingAssessedTax": [
            _usd_entry(val=50_000_000, fy=2024, fp="Q2", form="10-Q", start="2024-04-01", end="2024-06-30"),
        ],
        "Assets": [
            _usd_entry(val=365_000_000_000, fy=2024, fp="Q2", form="10-Q", end="2024-06-30"),
        ],
    })


def _q4_cf_facts() -> dict:
    return _facts({
        "NetCashProvidedByUsedInOperatingActivities": [
            _usd_entry(val=150_000_000, fy=2024, fp="Q3", form="10-Q", start="2024-01-01", end="2024-09-30"),
            _usd_entry(val=200_000_000, fy=2024, fp="FY", form="10-K", start="2024-01-01", end="2024-12-31"),
        ],
    })


def test_1_missing_bloomberg_is_sec_retrieval(tmp_path: Path):
    src = _blank_stmt_wb(tmp_path / "src.xlsx", is_blank=True, cf_blank=False)
    dest = tmp_path / "out.xlsx"
    facts = _facts({
        "RevenueFromContractWithCustomerExcludingAssessedTax": [
            _usd_entry(val=94_930_000_000, fy=2024, fp="Q2", form="10-Q", start="2024-04-01", end="2024-06-30"),
        ],
        "OperatingIncomeLoss": [
            _usd_entry(val=29_000_000_000, fy=2024, fp="Q2", form="10-Q", start="2024-04-01", end="2024-06-30"),
        ],
        "NetIncomeLoss": [
            _usd_entry(val=14_000_000_000, fy=2024, fp="Q2", form="10-Q", start="2024-04-01", end="2024-06-30"),
        ],
        "Assets": [
            _usd_entry(val=1_000_000_000, fy=2024, fp="Q2", form="10-Q", end="2024-06-30"),
        ],
        "NetCashProvidedByUsedInOperatingActivities": [
            _usd_entry(val=10_000_000, fy=2024, fp="Q2", form="10-Q", start="2024-01-01", end="2024-06-30"),
        ],
    })
    report = QuarterlyPresentationService().plan_and_apply(
        analysis_id="t1",
        ticker="AAPL",
        source_workbook_path=src,
        destination_workbook_path=dest,
        company_facts=facts,
    )
    is_stmt = next(s for s in report.statements if s.statement == QuarterlyStatementKind.INCOME)
    assert is_stmt.decision == PresentationDecision.STATEMENT_INCOMPLETE
    assert "STATEMENT_INCOMPLETE" in is_stmt.reason
    assert report.input_blockers
    wb = load_workbook(dest)
    try:
        ws = wb["Last Quarter IS Standardized"]
        assert ws["C11"].value != pytest.approx(94930.0)
    finally:
        wb.close()


def test_2_missing_cf_standalone_q2_derivation():
    items = extract_sec_10q_statement(
        _q2_cf_facts(),
        QuarterlyStatementKind.CASH_FLOW,
        fiscal_year=2024,
        fiscal_period="Q2",
        include_unresolved=False,
    )
    cfo = next(i for i in items if "operating" in i.label.lower())
    assert cfo.extraction_method == "derived_ytd_subtract"
    assert cfo.duration_kind == "standalone_quarter"
    assert cfo.value == pytest.approx(60.0)
    assert cfo.derivation


def test_3_missing_is_and_cf_rebuild_tabs(tmp_path: Path):
    src = _blank_stmt_wb(tmp_path / "src.xlsx", is_blank=True, cf_blank=True)
    dest = tmp_path / "out.xlsx"
    facts = _q2_cf_facts()
    facts["facts"]["us-gaap"]["OperatingIncomeLoss"] = {
        "units": {"USD": [
            _usd_entry(val=12_000_000, fy=2024, fp="Q2", form="10-Q", start="2024-04-01", end="2024-06-30"),
        ]}
    }
    report = QuarterlyPresentationService().plan_and_apply(
        analysis_id="t3",
        ticker="AAPL",
        source_workbook_path=src,
        destination_workbook_path=dest,
        company_facts=facts,
    )
    kinds = {s.statement: s for s in report.statements}
    assert kinds[QuarterlyStatementKind.INCOME].decision == PresentationDecision.STATEMENT_INCOMPLETE
    assert kinds[QuarterlyStatementKind.CASH_FLOW].decision == PresentationDecision.STATEMENT_INCOMPLETE
    assert report.input_blockers
    wb = load_workbook(dest)
    try:
        assert wb["Last Quarter IS Standardized"]["A11"].value != "Revenue"
        assert wb["Last Quarter CF Standardized"]["C11"].value != pytest.approx(60.0)
    finally:
        wb.close()


def test_4_absent_from_both_sources_unresolved_not_zero(tmp_path: Path):
    src = _blank_stmt_wb(tmp_path / "src.xlsx")
    dest = tmp_path / "out.xlsx"
    facts = _facts({
        "Assets": [_usd_entry(val=1_000_000, fy=2024, fp="Q2", form="10-Q", end="2024-06-30")],
    })
    report = QuarterlyPresentationService().plan_and_apply(
        analysis_id="t4",
        ticker="AAPL",
        source_workbook_path=src,
        destination_workbook_path=dest,
        company_facts=facts,
    )
    is_stmt = next(s for s in report.statements if s.statement == QuarterlyStatementKind.INCOME)
    assert is_stmt.unresolved_facts
    wb = load_workbook(dest)
    try:
        ws = wb["Last Quarter IS Standardized"]
        for r in range(11, 25):
            lab = ws.cell(r, 1).value
            val = ws.cell(r, 3).value
            if lab and "Revenue" in str(lab):
                assert val is None
    finally:
        wb.close()


def test_5_sec_conflict_does_not_overwrite_blank_statement(tmp_path: Path):
    src = _blank_stmt_wb(tmp_path / "src.xlsx")
    dest = tmp_path / "out.xlsx"
    facts = _facts({
        "RevenueFromContractWithCustomerExcludingAssessedTax": [
            _usd_entry(val=80_000_000, fy=2024, fp="Q2", form="10-Q", start="2024-04-01", end="2024-06-30"),
        ],
        "Assets": [_usd_entry(val=1_000_000, fy=2024, fp="Q2", form="10-Q", end="2024-06-30")],
    })
    report = QuarterlyPresentationService().plan_and_apply(
        analysis_id="t5",
        ticker="AAPL",
        source_workbook_path=src,
        destination_workbook_path=dest,
        company_facts=facts,
    )
    is_stmt = next(s for s in report.statements if s.statement == QuarterlyStatementKind.INCOME)
    assert is_stmt.decision == PresentationDecision.STATEMENT_INCOMPLETE
    assert is_stmt.data_source_primary == "workbook"
    wb = load_workbook(dest)
    try:
        assert wb["Last Quarter IS Standardized"]["C11"].value != pytest.approx(80.0)
    finally:
        wb.close()


def test_6_q4_cf_derivation_from_10k_minus_9m():
    items = extract_sec_10q_statement(
        _q4_cf_facts(),
        QuarterlyStatementKind.CASH_FLOW,
        fiscal_year=2024,
        fiscal_period="Q4",
        include_unresolved=False,
    )
    cfo = next(i for i in items if "operating" in i.label.lower())
    assert cfo.extraction_method == "derived_ytd_subtract"
    assert cfo.value == pytest.approx(50.0)


def test_6b_balance_sheet_never_derived_by_subtraction():
    facts = _facts({
        "Assets": [
            _usd_entry(val=100_000_000, fy=2024, fp="Q1", form="10-Q", end="2024-03-31"),
            _usd_entry(val=180_000_000, fy=2024, fp="Q2", form="10-Q", end="2024-06-30"),
        ],
    })
    items = extract_sec_10q_statement(
        facts, QuarterlyStatementKind.BALANCE_SHEET, fiscal_year=2024, fiscal_period="Q2"
    )
    assets = next(i for i in items if i.label == "Total assets")
    assert assets.extraction_method == "reported_instant"
    assert assets.value == pytest.approx(180.0)
    assert assets.duration_kind == "instant"


def _bs_pair(tmp_path: Path) -> tuple[Path, Path]:
    prev = tmp_path / "prev.xlsx"
    dest_tmpl = tmp_path / "new.xlsx"

    def _build(
        path: Path,
        *,
        ticker: str,
        i_values: dict[str, object] | None,
        h_shift: int,
        period: str,
        outputs: dict[str, object] | None = None,
    ):
        wb = Workbook()
        wb.remove(wb.active)
        inp = wb.create_sheet("Inputs")
        bs = wb.create_sheet("Last Quarter BS Standardized")
        inc = wb.create_sheet("Last Quarter IS Standardized")
        inc["A1"] = f"{ticker} Quarterly Income"
        inc["C5"] = period
        mapping = [
            (26 - h_shift, "B63"),
            (27 - h_shift, "B67"),
            (28 - h_shift, "B73"),
            (29 - h_shift, "B74"),
            (30 - h_shift, "B75"),
            (31 - h_shift, "B70"),
            (32 - h_shift, "B71"),
        ]
        for row, bref in mapping:
            if row < 1:
                continue
            bs.cell(row, 8).value = f"=Inputs!{bref}"
        if i_values:
            for addr, val in i_values.items():
                cell = bs[addr]
                cell.value = val
                if isinstance(val, float) and abs(val) < 2:
                    cell.number_format = "0.00%"
                elif isinstance(val, (int, float)) and val > 10:
                    cell.number_format = '"$"#,##0.00'
        bs["I22"] = "=H22"
        for addr, val in (outputs or {}).items():
            cell = inp[addr]
            cell.value = val
            if isinstance(val, float) and abs(val) < 2:
                cell.number_format = "0.00%"
            elif isinstance(val, (int, float)) and val > 10:
                cell.number_format = '"$"#,##0.00'
        wb.save(path)
        wb.close()

    _build(
        prev,
        ticker="AAPL",
        period="2026 Q1",
        i_values={
            "I26": 1.0,
            "I27": 2.0,
            "I28": 3.0,
            "I29": 4.0,
            "I30": 5.0,
            "I31": 6.0,
        },
        outputs={
            "B67": 95.01,
            "B73": 1.0377,
            "B74": "P",
            "B75": 0.4952,
            "B70": 0.0218,
            "B71": 0.1709,
        },
        h_shift=1,
    )
    _build(dest_tmpl, ticker="AAPL", period="2026 Q2", i_values=None, h_shift=0)
    return prev, dest_tmpl


def test_7_i27_i32_carry_by_inputs_mapping(tmp_path: Path):
    prev, tmpl = _bs_pair(tmp_path)
    dest = tmp_path / "carried.xlsx"
    report = QuarterlyCarryForwardService().apply(
        analysis_id="c7",
        ticker="AAPL",
        new_template_path=tmpl,
        previous_workbook_path=prev,
        destination_path=dest,
    )
    carried = {
        e.cell: e
        for e in report.entries
        if e.sheet == "Last Quarter BS Standardized" and e.cell.startswith("I")
    }
    assert carried["I27"].final_action == CarryForwardDecision.CARRY_FORWARD
    assert carried["I27"].final_value == pytest.approx(95.01)
    assert carried["I28"].final_value == pytest.approx(1.0377)
    assert carried["I29"].final_value == "P"
    assert carried["I30"].final_value == pytest.approx(0.4952)
    assert carried["I31"].final_value == pytest.approx(0.0218)
    assert carried["I32"].final_value == pytest.approx(0.1709)
    wb = load_workbook(dest)
    try:
        bs = wb["Last Quarter BS Standardized"]
        assert bs["I27"].value == pytest.approx(95.01)
        assert "%" in str(bs["I28"].number_format)
        assert bs["I22"].value == "=H22"
        assert bs["I27"].value != 2.0
    finally:
        wb.close()


def test_7b_wrong_quarter_and_uncached_formula_are_review(tmp_path: Path):
    prev, tmpl = _bs_pair(tmp_path)
    pwb = load_workbook(prev)
    pwb["Last Quarter IS Standardized"]["C5"] = "2025 Q4"
    pwb.save(prev)
    pwb.close()
    report = QuarterlyCarryForwardService().apply(
        analysis_id="c7b",
        ticker="AAPL",
        new_template_path=tmpl,
        previous_workbook_path=prev,
        destination_path=tmp_path / "wrong.xlsx",
    )
    i_entries = [
        e for e in report.entries
        if e.sheet == "Last Quarter BS Standardized" and e.cell == "I27"
    ]
    assert i_entries[0].final_action == CarryForwardDecision.REVIEW_REQUIRED
    assert "immediately before" in i_entries[0].reason

    (tmp_path / "cache").mkdir()
    prev2, tmpl2 = _bs_pair(tmp_path / "cache")
    pwb = load_workbook(prev2)
    pwb["Inputs"]["B67"] = "=1+1"
    pwb.save(prev2)
    pwb.close()
    report2 = QuarterlyCarryForwardService().apply(
        analysis_id="c7c",
        ticker="AAPL",
        new_template_path=tmpl2,
        previous_workbook_path=prev2,
        destination_path=tmp_path / "cache" / "out.xlsx",
    )
    missing = next(
        e for e in report2.entries
        if e.sheet == "Last Quarter BS Standardized" and e.cell == "I27"
    )
    assert missing.final_action == CarryForwardDecision.REVIEW_REQUIRED
    assert "Cached calculated value" in missing.reason


def test_8_missing_previous_workbook_no_substitution(tmp_path: Path):
    tmpl = tmp_path / "new.xlsx"
    wb = Workbook()
    wb.active.title = "Inputs"
    wb.save(tmpl)
    wb.close()
    report = QuarterlyCarryForwardService().apply(
        analysis_id="c8",
        ticker="AAPL",
        new_template_path=tmpl,
        previous_workbook_path=None,
        destination_path=tmp_path / "out.xlsx",
    )
    assert report.status == "MISSING_PREVIOUS_WORKBOOK"


def test_8b_mismatched_company_no_silent_i_column(tmp_path: Path):
    prev, tmpl = _bs_pair(tmp_path)
    pwb = load_workbook(prev)
    pwb["Last Quarter IS Standardized"]["A1"] = "MSFT Quarterly Income"
    pwb.save(prev)
    pwb.close()
    dest = tmp_path / "carried.xlsx"
    report = QuarterlyCarryForwardService().apply(
        analysis_id="c8b",
        ticker="AAPL",
        new_template_path=tmpl,
        previous_workbook_path=prev,
        destination_path=dest,
    )
    i_entries = [
        e for e in report.entries
        if e.sheet == "Last Quarter BS Standardized" and e.cell in {f"I{r}" for r in range(27, 33)}
    ]
    assert i_entries
    assert all(e.final_action == CarryForwardDecision.REVIEW_REQUIRED for e in i_entries)
    wb = load_workbook(dest)
    try:
        bs = wb["Last Quarter BS Standardized"]
        for r in range(27, 33):
            assert bs.cell(r, 9).value in (None, "")
    finally:
        wb.close()


def test_9_roic_roce_percent_and_yellow_fill(tmp_path: Path):
    path = _mini_proj_wb(tmp_path / "q2.xlsx", q=2)
    report = QuarterlyProjectionService().apply(
        analysis_id="p9", ticker="AAPL", workbook_path=path, fiscal_quarter=2
    )
    assert report.yellow_fill_applied is True
    assert "M23" in report.percent_format_cells
    assert "M24" in report.percent_format_cells
    assert "M25" in report.percent_format_cells
    assert "N4" in report.percent_format_cells
    wb = load_workbook(path)
    try:
        ic = wb["IC & NOPAT & ROIC "]
        yellow = "FFFF00"
        for row in range(1, 26):
            for col in (13, 14):
                fill = ic.cell(row, col).fill
                fg = fill.fgColor.rgb if fill.fgColor is not None else ""
                assert yellow in str(fg).upper()
        assert ic["M23"].number_format == "0.00%"
        assert ic["M24"].number_format == "0.00%"
        assert ic["M25"].number_format == "0.00%"
        assert ic["N4"].number_format == "0.00%"
        assert ic["M24"].value == pytest.approx(0.10)
        assert ic["M3"].number_format != "0.00%"
        assert ic["M23"].value == "=M20/M7"
    finally:
        wb.close()


def test_10_valuation_parity_keep_adjust_insufficient(tmp_path: Path):
    keep = tmp_path / "keep.xlsx"
    adj = tmp_path / "adj.xlsx"
    insuff = tmp_path / "ins.xlsx"
    _valuation_workbook(
        keep,
        a11=0.04,
        eps=_geo(2.0, 0.05, 10),
        revenue=_geo(100.0, 0.03, 10),
        oi=_geo(10.0, 0.045, 10),
    )
    _valuation_workbook(
        adj,
        a11=0.20,
        eps=_geo(2.0, 0.021, 10),
        revenue=_geo(100.0, 0.022, 10),
        oi=_geo(10.0, 0.10, 10),
    )
    ni = [36, 33, 39, 54, 60, 62, 63, 60, 59, 62]
    da = [16, 15, 17, 18, 18, 18, 21, 25, 27, 28]
    capex = [-11, -13, -15, -15, -25, -18, -21, -28, -51, -88]
    _valuation_workbook(
        insuff,
        ni=ni,
        da=da,
        capex=capex,
        oi=[60, 56, 59, 79, 85, 87, 90, 85, 85, 89],
        eps=[3.17, 2.84, 3.43, 4.69, 5.17, 5.33, 5.4, 5.15, 5.03, 5.26],
        b6_total=-0.95,
        c6_annual=-0.283,
        a11=0.027,
    )

    intel = AnnualAnalystIntelligenceService()
    judge_svc = AnnualJudgmentService()

    def _run(path: Path):
        ctx = intel.build_judgment_context(analysis_id="q", ticker="X", workbook_path=path)
        return judge_svc.apply(analysis_id="q", ticker="X", workbook_path=path, context=ctx)

    _keep_er, keep_j = _run(keep)
    _adj_er, adj_j = _run(adj)
    _ins_er, ins_j = _run(insuff)
    assert keep_j.er_analysis.decision in {"KEEP_EXISTING", "KEEP", "ADJUST"}
    assert adj_j.er_analysis.decision == "ADJUST"
    assert ins_j.oe_analysis.decision == "INSUFFICIENT_EVIDENCE"
    assert ins_j.er_analysis.decision != ins_j.oe_analysis.decision or ins_j.graham_analysis.decision != ins_j.oe_analysis.decision
    wb = load_workbook(adj)
    try:
        assert wb["Expected Returns & Buybacks"]["A11"].value == pytest.approx(0.20)
    finally:
        wb.close()


def test_11_formulas_and_historical_tabs_preserved(tmp_path: Path):
    prev, tmpl = _bs_pair(tmp_path)
    wb = load_workbook(tmpl)
    hist = wb.create_sheet("Income - GAAP")
    hist["C10"] = "=C9+C8"
    hist["C11"] = 999
    wb["Last Quarter BS Standardized"]["I22"] = "=H22"
    wb.save(tmpl)
    wb.close()
    dest = tmp_path / "carried.xlsx"
    QuarterlyCarryForwardService().apply(
        analysis_id="c11",
        ticker="AAPL",
        new_template_path=tmpl,
        previous_workbook_path=prev,
        destination_path=dest,
    )
    out = load_workbook(dest)
    try:
        assert out["Income - GAAP"]["C10"].value == "=C9+C8"
        assert out["Income - GAAP"]["C11"].value == 999
        assert out["Last Quarter BS Standardized"]["I22"].value == "=H22"
    finally:
        out.close()


def test_12_no_hap_introduced_circular_on_projection(tmp_path: Path):
    path = _mini_proj_wb(tmp_path / "circ.xlsx", q=2)
    QuarterlyProjectionService().apply(
        analysis_id="c12", ticker="AAPL", workbook_path=path, fiscal_quarter=2
    )
    cycles = CircularReferenceService().scan_workbook(path, analysis_id="c12", ticker="AAPL")
    assert cycles == [] or all(c.origin != "hap_introduced" for c in cycles)
    wb = load_workbook(path)
    try:
        ic = wb["IC & NOPAT & ROIC "]
        assert ic["M23"].value == "=M20/M7"
        assert ic["M7"].value == "=M3-M4+M5+M6"
    finally:
        wb.close()


def test_empty_sec_does_not_fill_statement(tmp_path: Path):
    src = _blank_stmt_wb(tmp_path / "src.xlsx")
    dest = tmp_path / "out.xlsx"
    report = QuarterlyPresentationService().plan_and_apply(
        analysis_id="yf",
        ticker="AAPL",
        source_workbook_path=src,
        destination_workbook_path=dest,
        company_facts=None,
    )
    is_stmt = next(s for s in report.statements if s.statement == QuarterlyStatementKind.INCOME)
    assert is_stmt.decision == PresentationDecision.STATEMENT_INCOMPLETE
    assert is_stmt.data_source_primary == "workbook"
    assert report.input_blockers
    wb = load_workbook(dest)
    try:
        assert wb["Last Quarter IS Standardized"]["A11"].value != "Total Revenues"
        assert wb["Last Quarter IS Standardized"]["C11"].value != pytest.approx(12.0)
    finally:
        wb.close()
