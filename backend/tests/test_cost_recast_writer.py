from openpyxl import Workbook, load_workbook

from services.cost_recast_writer import CostRecastWriter
from tests.ledger_util import entry, notes_text

STATEMENT_22 = ("CONSOLIDATED STATEMENTS OF INCOME (in thousands) OPERATING EXPENSES: Research and portfolio development 185,202 200,484 204,360 "
                "Licensing 71,419 64,625 50,464 General and administrative 47,377 61,217 48,999 Restructuring activities 3,280 27,877 - "
                "Total Operating expenses 307,278 354,203 303,823 Income from operations 150,516")
KEYS = {9: "SALES_REV_TURN", 14: "IS_COGS_TO_FE_AND_PP_AND_G", 15: "IS_COG_AND_SERVICES_SOLD", 19: "GROSS_PROFIT",
        22: "IS_OPERATING_EXPN", 23: "IS_SG&A_EXPENSE", 26: "IS_OPERATING_EXPENSES_R&D", 28: "OTHER_OPERATING_EXPENSES_RATIO", 30: "IS_OPER_INC"}


def build(path, formula_in_gp=False):
    wb = Workbook()
    ws = wb.active
    ws.title = "Income - GAAP"
    for r, k in KEYS.items():
        ws.cell(r, 2, k)
    # FY2021 in column C (old provider definition), FY2022 in column D (already latest definition)
    for col, fy in ((3, 2021), (4, 2022)):
        ws.cell(3, col, f"FY {fy}")
    data = {
        3: dict(rev=425.409, cost=175.741, gp=249.668, opex=178.462, sga=61.217, rd=89.368, other=27.877, opinc=71.206),
        4: dict(rev=457.794, cost=71.419, gp=386.375, opex=235.859, sga=47.377, rd=185.202, other=3.280, opinc=150.516),
    }
    for col, v in data.items():
        for row, key in ((9, "rev"), (14, "cost"), (15, "cost"), (19, "gp"), (22, "opex"), (23, "sga"), (26, "rd"), (28, "other"), (30, "opinc")):
            ws.cell(row, col, v[key])
    if formula_in_gp:
        ws.cell(19, 3, "=C9-C14")
    wb.save(path)


def write_filing(tmp_path):
    folder = tmp_path / "filings"
    folder.mkdir()
    (folder / "10k_buybacks_2022.htm").write_text(f"<html><body><p>{STATEMENT_22}</p></body></html>", encoding="utf-8")
    return folder


def test_recast_keeps_totals_and_flags_cells(tmp_path):
    path = tmp_path / "w.xlsx"
    build(path)
    report = CostRecastWriter().apply(workbook_path=path, filings_dir=write_filing(tmp_path))
    assert report["applied"] and report["label"] == "licensing"
    assert [y["fiscal_year"] for y in report["years"]] == ["FY2021"] and report["years"][0]["status"] == "written"
    ws = load_workbook(path)["Income - GAAP"]
    cost, gp, opex, rd, other = (ws.cell(r, 3).value for r in (14, 19, 22, 26, 28))
    assert (cost, rd, other) == (64.625, 200.484, 27.877)
    assert round(425.409 - cost - opex, 3) == 71.206 and round(gp, 3) == round(425.409 - cost, 3)
    assert abs(ws.cell(23, 3).value + rd + other - opex) < 1e-6
    recast = entry(ws.parent, ws.title, "C14", "Recast to latest definition")
    assert recast and recast["original"] == "175.741"          # the original is kept in the HAP Adjustments tab
    assert ws.cell(14, 4).value == 71.419 and ws.cell(14, 4).comment is None  # already on the latest definition: untouched


def test_skips_when_a_target_is_a_formula_or_totals_do_not_reconcile(tmp_path):
    path = tmp_path / "w.xlsx"
    build(path, formula_in_gp=True)
    report = CostRecastWriter().apply(workbook_path=path, filings_dir=write_filing(tmp_path))
    assert not report["applied"] and "formula" in report["years"][0]["reason"]
    assert load_workbook(path)["Income - GAAP"].cell(14, 3).value == 175.741

    path2 = tmp_path / "w2.xlsx"
    build(path2)
    ws_wb = load_workbook(path2)
    ws_wb["Income - GAAP"].cell(22, 3, 150.0)  # opex no longer reconciles with the filing
    ws_wb.save(path2)
    report2 = CostRecastWriter().apply(workbook_path=path2, filings_dir=write_filing(tmp_path / "x") if False else tmp_path / "filings")
    assert not report2["applied"] and "reconcile" in report2["years"][0]["reason"]


def test_no_filings_means_no_change(tmp_path):
    path = tmp_path / "w.xlsx"
    build(path)
    assert CostRecastWriter().apply(workbook_path=path, filings_dir=None)["applied"] is False
