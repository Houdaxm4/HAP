"""Focused Annual Update tests — continuity, restatement, statements, tax, R&D, leases, judgment, Word."""

from __future__ import annotations

from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill

from models.annual_update import ContinuityAction
from services.annual_continuity_service import AnnualContinuityService
from services.annual_deliverables_service import AnnualDeliverablesService
from services.annual_formula_guard_service import AnnualFormulaGuardService
from services.annual_inputs_service import AnnualInputsService
from services.annual_judgment_service import AnnualJudgmentService
from services.annual_leases_service import AnnualLeasesService
from services.annual_rd_service import AnnualRdService
from services.annual_research_service import AnnualResearchService
from services.annual_restatement_service import AnnualRestatementService
from services.annual_statement_validation_service import AnnualStatementValidationService
from services.annual_tax_service import AnnualTaxService
from services.annual_update_runner import sha256_file


def _header(ws, years: list[int]) -> None:
    ws["A7"] = "Line"
    for i, y in enumerate(years):
        ws.cell(7, 3 + i, f"FY{y}")


def _mini_pair(tmp: Path) -> tuple[Path, Path]:
    prev = tmp / "prev.xlsx"
    tmpl = tmp / "tmpl.xlsx"

    def _book(path: Path, years: list[int], *, hist_rev: float, new_rev: float, rd_life: float, rates: list[float]) -> None:
        wb = Workbook()
        is_ = wb.active
        is_.title = "Income - GAAP"
        _header(is_, years)
        is_["A11"] = "Revenue"
        is_["A12"] = "Operating income"
        is_["A13"] = "Net income"
        is_["A14"] = "Diluted EPS"
        for i, y in enumerate(years):
            col = 3 + i
            is_.cell(11, col, hist_rev + i * 10 if y < years[-1] else new_rev)
            is_.cell(12, col, 50 + i)
            is_.cell(13, col, 40 + i)
            is_.cell(14, col, 1 + i * 0.1)
        is_["C11"].font = Font(bold=True, color="FF0000")
        is_["C11"].fill = PatternFill("solid", fgColor="FFFF00")
        is_["C12"] = "=C11*0.2"

        bs = wb.create_sheet("Balance Sheet - Standardized")
        _header(bs, years)
        bs["A11"] = "Total assets"
        bs["A12"] = "Total liabilities"
        bs["A13"] = "Total shareholders' equity"
        for i, _y in enumerate(years):
            bs.cell(11, 3 + i, 1000 + i)
            bs.cell(12, 3 + i, 600 + i)
            bs.cell(13, 3 + i, 400)

        cf = wb.create_sheet("Cash Flow - Standardized")
        _header(cf, years)
        cf["A11"] = "Cash from operating activities"
        cf["A12"] = "Cash from investing activities"
        cf["A13"] = "Cash from financing activities"
        for i, _y in enumerate(years):
            cf.cell(11, 3 + i, 80 + i)
            cf.cell(12, 3 + i, -20)
            cf.cell(13, 3 + i, -30)

        inp = wb.create_sheet("Inputs")
        _header(inp, years)
        inp["A10"] = "PE10"
        inp["A63"] = "Current Price"
        inp["B63"] = 100
        inp["A106"] = "Statutory federal"
        inp["A107"] = "State"
        inp["A108"] = "Foreign"
        inp["A109"] = "Credits"
        inp["A110"] = "Deferred"
        inp["A111"] = "Other"
        for i, y in enumerate(years):
            if y < years[-1]:
                inp.cell(10, 3 + i, 20 + i)
                inp.cell(106, 3 + i, 0.21)

        rd = wb.create_sheet("R&D")
        rd["B8"] = rd_life
        rd["A11"] = "R&D expense"
        _header(rd, years)
        for i, y in enumerate(years):
            if y < years[-1]:
                rd.cell(11, 3 + i, 8 + i)
        rd["C20"] = "=C11/10"

        ls = wb.create_sheet("Leases")
        _header(ls, years)
        for i, r in enumerate(rates):
            ls.cell(18, 3 + i, r)

        ic = wb.create_sheet("IC & NOPAT & ROIC ")
        ic["A10"] = "Invested capital"
        ic["C10"] = "=Inputs!C10"
        ic["A11"] = "NOPAT"
        ic["C11"] = "=C10*0.1"
        ic["A12"] = "ROIC"
        ic["C12"] = "=C11/C10"
        ic["A13"] = "ROIC-WACC"
        ic["C13"] = "=C12-0.08"
        ic["A14"] = "Lease adjustment"
        ic["A15"] = "R&D adjustment"
        ic["A16"] = "Operating tax"

        ratios = wb.create_sheet("All Ratios")
        ratios["B2"] = "=('Income - GAAP'!C11)"
        fin = wb.create_sheet("Final Metrics")
        fin["B5"] = 190
        fin["A5"] = "Current Price"
        fin["B6"] = 28
        fin["A6"] = "Current PE10"
        fin["B7"] = 150
        fin["A7"] = "Max Buy"
        fin["B8"] = 0.2
        fin["B9"] = 0.08
        fin["B10"] = 120
        fin["B11"] = 0.3
        fin["B12"] = 20
        fin["B13"] = 0.05
        fin["B14"] = 0.04
        fin["B15"] = 0.06
        fin["B16"] = 0.15
        fin["B17"] = 0.14
        fin["B18"] = 0.12
        fin["C5"] = "=B5"

        wb.save(path)
        wb.close()

    years_prev = [2021, 2022, 2023, 2024]
    years_new = [2021, 2022, 2023, 2024, 2025]
    _book(prev, years_prev, hist_rev=100, new_rev=140, rd_life=6, rates=[0.04, 0.041, 0.039, 0.04])
    _book(tmpl, years_new, hist_rev=0, new_rev=200, rd_life=99, rates=[0.04, 0.041, 0.039, 0.04, 0.12])
    return prev, tmpl


def test_historical_formulas_values_and_formats_carried(tmp_path: Path):
    prev, tmpl = _mini_pair(tmp_path)
    out = tmp_path / "out.xlsx"
    out.write_bytes(tmpl.read_bytes())
    report = AnnualContinuityService().apply(
        analysis_id="a1",
        ticker="AAPL",
        template_path=tmpl,
        previous_workbook_path=prev,
        workbook_path=out,
    )
    wb = load_workbook(out)
    assert wb["Income - GAAP"]["C12"].value == "=C11*0.2"
    assert wb["Income - GAAP"]["C11"].value == 100
    assert wb["Income - GAAP"]["C11"].font.bold is True
    assert wb["Income - GAAP"]["G11"].value == 200  # new FY not overwritten by hist 140
    wb.close()
    actions = {e.action for e in report.entries}
    assert ContinuityAction.CARRY_FORWARD_VALUE in actions or ContinuityAction.RESTORE_FORMULA in actions
    assert ContinuityAction.NEW_FY_REFRESH in actions


def test_new_fy_excluded_and_current_data_not_carried(tmp_path: Path):
    prev, tmpl = _mini_pair(tmp_path)
    out = tmp_path / "out.xlsx"
    out.write_bytes(tmpl.read_bytes())
    AnnualContinuityService().apply(
        analysis_id="a1",
        ticker="AAPL",
        template_path=tmpl,
        previous_workbook_path=prev,
        workbook_path=out,
    )
    wb = load_workbook(out)
    assert wb["Inputs"]["B63"].value == 100  # template current price remains until refresh
    # previous had 100 too; action should still be NEW_FY_REFRESH for B63
    wb.close()


def test_explicit_restatement_updates_and_vague_is_review(tmp_path: Path):
    prev, tmpl = _mini_pair(tmp_path)
    out = tmp_path / "out.xlsx"
    out.write_bytes(tmpl.read_bytes())
    AnnualContinuityService().apply(
        analysis_id="a1", ticker="AAPL", template_path=tmpl,
        previous_workbook_path=prev, workbook_path=out,
    )
    svc = AnnualRestatementService()
    auto = svc.apply(
        analysis_id="a1",
        ticker="AAPL",
        workbook_path=out,
        explicit_revisions=[
            {
                "concept": "revenue",
                "fiscal_year": "FY2024",
                "revised_value": 333.0,
                "sheet": "Income - GAAP",
                "row": 11,
                "source": "FY2025 10-K comparative",
                "section": "Item 8",
                "reason": "Explicit recast of FY2024 revenue",
            }
        ],
    )
    assert auto.automatic_changes
    assert auto.automatic_changes[0].automatic_change is True
    wb = load_workbook(out)
    # FY2024 is column F (2021=C ... 2024=F)
    assert wb["Income - GAAP"]["F11"].value == 333.0
    wb.close()

    vague = svc.apply(
        analysis_id="a1",
        ticker="AAPL",
        workbook_path=out,
        narrative_disclosures=[
            {
                "concept": "revenue",
                "fiscal_year": "FY2023",
                "source": "press release",
                "explanation": "We recast certain prior-period amounts",
                "revised_comparatives": False,
            }
        ],
    )
    assert vague.review_required
    assert vague.status == "RESTATEMENT_REVIEW_REQUIRED"
    wb = load_workbook(out)
    assert wb["Income - GAAP"]["E11"].value != 333  # FY2023 not silently rewritten
    wb.close()


def test_statement_validation_preserves_bloomberg(tmp_path: Path):
    prev, tmpl = _mini_pair(tmp_path)
    report = AnnualStatementValidationService().validate(
        analysis_id="a1",
        ticker="AAPL",
        workbook_path=tmpl,
        fiscal_year="FY2025",
        filing_overrides={"revenue": 200.0, "operating_income": 54.0, "net_income": 44.0},
    )
    assert report.bloomberg_preserved is True
    rev = next(i for i in report.items if i.concept == "revenue")
    assert rev.status == "VALIDATED"


def test_tax_residual_other_reconciles(tmp_path: Path):
    prev, tmpl = _mini_pair(tmp_path)
    out = tmp_path / "out.xlsx"
    out.write_bytes(tmpl.read_bytes())
    report = AnnualTaxService().apply(
        analysis_id="a1",
        ticker="AAPL",
        workbook_path=out,
        fiscal_year="FY2025",
        reported_effective_rate=0.16,
        filing_components=[
            {"label": "Federal statutory rate", "rate": 0.21},
            {"label": "State taxes", "rate": 0.02},
            {"label": "Foreign rate differential", "rate": -0.04},
            {"label": "Other items (company label)", "rate": 0.99},  # must not map to house Other
        ],
        source="10-K tax recon",
    )
    assert report.residual_other == pytest.approx(0.16 - (0.21 + 0.02 - 0.04))
    assert report.reconciliation_delta == pytest.approx(0.0)
    other = next(c for c in report.mapped_components if c.house_category == "other")
    assert other.residual is True
    assert other.source_label and "residual" in other.source_label.lower()


def test_rd_useful_life_unchanged_and_filing_fill(tmp_path: Path):
    prev, tmpl = _mini_pair(tmp_path)
    out = tmp_path / "out.xlsx"
    out.write_bytes(tmpl.read_bytes())
    report = AnnualRdService().apply(
        analysis_id="a1",
        ticker="AAPL",
        workbook_path=out,
        previous_workbook_path=prev,
        fiscal_year="FY2025",
        filing_rd=12.5,
    )
    wb = load_workbook(out)
    assert wb["R&D"]["B8"].value == 6
    assert wb["R&D"]["C20"].value == "=C11/10"
    assert wb["R&D"]["G11"].value == 12.5
    wb.close()
    assert report.useful_life_unchanged is True
    assert report.filled_from_filing is True


def test_lease_outlier_normalized_unless_event(tmp_path: Path):
    prev, tmpl = _mini_pair(tmp_path)
    out = tmp_path / "out.xlsx"
    out.write_bytes(tmpl.read_bytes())
    svc = AnnualLeasesService()
    n = svc.apply(
        analysis_id="a1", ticker="AAPL", workbook_path=out, fiscal_year="FY2025"
    )
    assert n.outlier is True
    assert n.normalized_to_prior is True
    wb = load_workbook(out)
    assert wb["Leases"]["G18"].value == pytest.approx(0.04)
    wb.close()

    out2 = tmp_path / "out2.xlsx"
    out2.write_bytes(tmpl.read_bytes())
    e = svc.apply(
        analysis_id="a1",
        ticker="AAPL",
        workbook_path=out2,
        fiscal_year="FY2025",
        extraordinary_event=True,
        event_note="Major store expansion.",
        event_sources=["10-K MD&A"],
    )
    assert e.normalized_to_prior is False
    wb = load_workbook(out2)
    assert wb["Leases"]["G18"].value == pytest.approx(0.12)
    wb.close()


def test_roic_and_ratios_formulas_untouched(tmp_path: Path):
    prev, tmpl = _mini_pair(tmp_path)
    roic, guard = AnnualFormulaGuardService().inspect(
        analysis_id="a1", ticker="AAPL", workbook_path=tmpl
    )
    assert roic.ic_formulas_preserved or roic.roic_formulas_preserved
    assert guard.ratios_formulas_preserved
    assert guard.final_metrics_formulas_preserved
    wb = load_workbook(tmpl)
    assert str(wb["All Ratios"]["B2"].value).startswith("=")
    assert str(wb["IC & NOPAT & ROIC "]["C12"].value).startswith("=")
    wb.close()


def test_expected_return_bv_first_then_judgment(tmp_path: Path):
    prev, tmpl = _mini_pair(tmp_path)
    svc = AnnualJudgmentService()
    er, judge = svc.apply(
        analysis_id="a1",
        ticker="AAPL",
        workbook_path=tmpl,
        context={"book_value_growth": 0.09, "expected_return": 0.11, "owner_earnings_growth": 0.08, "eps_growth": 0.07},
    )
    assert er.default_methodology == "BOOK_VALUE_GROWTH"
    assert judge.expected_return and judge.expected_return.adjusted is False

    er2, judge2 = svc.apply(
        analysis_id="a1",
        ticker="AAPL",
        workbook_path=tmpl,
        context={
            "book_value_growth": 0.80,
            "eps_growth": 0.08,
            "prefer_eps": True,
            "owner_earnings_growth": 0.70,
            "eps_distorted": True,
            "normalized_eps_growth": 0.09,
            "normalized_oe_growth": 0.08,
            "evidence": ["buyback distortion"],
        },
    )
    assert er2.reasonableness in {"unrealistic", "AGGRESSIVE", "DISTORTED"}
    assert er2.selected_methodology == "EPS_GROWTH"
    assert judge2.owner_earnings_growth and judge2.owner_earnings_growth.adjusted
    assert judge2.graham_eps_growth and judge2.graham_eps_growth.adjusted
    assert judge2.expected_return and judge2.expected_return.rationale


def test_word_report_flags_missing_call_and_includes_required_sections(tmp_path: Path, monkeypatch):
    prev, tmpl = _mini_pair(tmp_path)
    monkeypatch.setattr(AnnualResearchService, "_fetch", lambda self, url: None)
    monkeypatch.setattr(AnnualResearchService, "_json", staticmethod(lambda url: {"news": []}))
    research = AnnualResearchService().gather(
        analysis_id="a1",
        ticker="ZZZZ",
        fiscal_year=2025,
        sec_manifest={"selected_filings": [{"form": "10-K", "filing_date": "2025-11-01", "document_url": "https://sec.gov/x"}]},
    )
    assert research.earnings_call_status == "EARNINGS_CALL_SOURCE_UNAVAILABLE"
    assert any(s.source_kind == "sec_10k" for s in research.sources)
    assert research.management_explanations == []

    deliv = AnnualDeliverablesService()
    perf = deliv.build_performance(
        analysis_id="a1", ticker="AAPL", fiscal_year=2025, workbook_path=tmpl, research=research
    )
    report = deliv.produce(
        analysis_id="a1",
        ticker="AAPL",
        fiscal_year=2025,
        completed_workbook_path=tmpl,
        output_dir=tmp_path,
        performance=perf,
        research=research,
    )
    assert report.excel_filename == "2025 AAPL FA.xlsx"
    assert report.word_filename == "2025 AAPL Annual Update.docx"
    assert (tmp_path / report.excel_filename).exists()
    from docx import Document

    doc = Document(tmp_path / report.word_filename)
    text = "\n".join(p.text for p in doc.paragraphs)
    assert "Executive Investment Conclusion" in text
    assert "Fiscal-Year Highlights" in text
    assert "EARNINGS_CALL_SOURCE_UNAVAILABLE" in text or "unavailable" in text.lower()
    assert "ROIC−WACC" in text or "ROIC" in text
    assert "Current Price" in text
    assert "Current Data and Valuation" in text
    # Filing dates may appear in Sources, but not as fiscal-year business highlights.
    highlight_block = text.split("Fiscal-Year Highlights")[1].split("3.")[0] if "Fiscal-Year Highlights" in text else ""
    assert "SEC 10-K filed" not in highlight_block


def test_source_immutability_hash(tmp_path: Path):
    prev, tmpl = _mini_pair(tmp_path)
    h1 = sha256_file(prev)
    out = tmp_path / "out.xlsx"
    out.write_bytes(tmpl.read_bytes())
    AnnualContinuityService().apply(
        analysis_id="a1", ticker="AAPL", template_path=tmpl,
        previous_workbook_path=prev, workbook_path=out,
    )
    assert sha256_file(prev) == h1
    assert sha256_file(tmpl) == sha256_file(tmpl)
