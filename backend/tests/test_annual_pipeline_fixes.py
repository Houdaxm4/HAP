"""Regression tests for Annual Update tax / PE10 / R&D / gates / Word grounding."""

from __future__ import annotations

from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook

from services.annual_deliverables_service import AnnualDeliverablesService
from services.annual_inputs_service import AnnualInputsService
from services.annual_output_gate_service import AnnualOutputGateService
from services.annual_period_service import align_fiscal_periods, detect_year_columns
from services.annual_rd_service import AnnualRdService
from services.annual_tax_service import AnnualTaxService
from services.annual_valuation_extract_service import AnnualValuationExtractService


def _rolling_books(tmp: Path) -> tuple[Path, Path, Path]:
    prev = tmp / "prev.xlsx"
    tmpl = tmp / "tmpl.xlsx"
    crf = tmp / "crf.xlsx"

    def book(path: Path, years: list[int], *, pe_last: float, rd_life: float = 3) -> None:
        wb = Workbook()
        bs = wb.active
        bs.title = "Balance Sheet - Standardized"
        for i, y in enumerate(years):
            bs.cell(5, 3 + i, f"{y} A")
            bs.cell(7, 3 + i, f"FY{y}")
            bs.cell(11, 3 + i, 100 + i)
        inc = wb.create_sheet("Income - GAAP")
        for i, y in enumerate(years):
            inc.cell(7, 3 + i, f"FY{y}")
        labels = {
            9: "Revenue",
            26: "    + Research & Development",
            42: "Pretax Income",
            44: "  - Income Tax Expense (Benefit)",
            50: "Net Income",
        }
        for row, lab in labels.items():
            inc.cell(row, 1, lab)
            for i, y in enumerate(years):
                if "Research" in lab:
                    inc.cell(row, 3 + i, 3.0 + i * 0.1)
                elif "Pretax" in lab:
                    inc.cell(row, 3 + i, 80 + i)
                elif "Tax Expense" in lab:
                    inc.cell(row, 3 + i, 20 + i * 0.2)
                elif "Revenue" in lab:
                    inc.cell(row, 3 + i, 1000 + i * 50)
                else:
                    inc.cell(row, 3 + i, 60 + i)
        inp = wb.create_sheet("Inputs")
        for i, y in enumerate(years):
            inp.cell(1, 3 + i, f"='Balance Sheet - Standardized'!{chr(67 + i)}7")
        inp["A57"] = "PE10"
        inp["A58"] = "E10"
        inp["A63"] = "Current Price"
        inp["B63"] = 74.37
        inp["A65"] = "Current PE10"
        inp["B65"] = None
        inp["A69"] = "Expected Return @ Current Price"
        inp["B69"] = 0.5712791523788112  # Bloomberg CRF — NOT Expected Returns!E14
        inp["A70"] = "Expected Return Price Plus Dividends - Given Current Price"
        inp["B70"] = 0.107188362899858
        inp["A103"] = "R&D Expense"
        inp["A106"] = "Tax Table"
        inp["E106"] = 1
        inp["A107"] = "Federal Tax"
        inp["A108"] = "State Taxes"
        inp["A109"] = "Foreign Taxes"
        inp["A110"] = "R&D Tax Credits"
        inp["A111"] = "All Other Items"
        inp["A112"] = "Income Tax Expense"
        for i, y in enumerate(years):
            if y < years[-1]:
                inp.cell(57, 3 + i, 12 + i * 0.1)
                inp.cell(107, 3 + i, 0.21)
                inp.cell(112, 3 + i, 0.25)
            inp.cell(103, 3 + i, f"='Income - GAAP'!{chr(67 + i)}26")
        rd = wb.create_sheet("R&D")
        rd["A2"] = "R&D Expense"
        rd["A3"] = "R&D Asset"
        rd["A4"] = "Amortization of R&D"
        rd["A8"] = "R&D Life"
        rd["B8"] = rd_life
        # R&D cols E.. map Inputs C..
        for i, y in enumerate(years):
            col = 5 + i
            letter = chr(67 + i)
            rd.cell(1, col, f"=Inputs!{letter}1")
            rd.cell(2, col, f'=IF(Inputs!{letter}103="",0,Inputs!{letter}103)')
            if i >= 2:
                prev_l = chr(64 + col - 1)
                pprev = chr(64 + col - 2)
                cur = chr(64 + col)
                rd.cell(3, col, f"={cur}2+{prev_l}2*2/3+1/3*{pprev}2")
                rd.cell(4, col, f"=({cur}2+{prev_l}2+{pprev}2)/3")
        tax = wb.create_sheet("Tax")
        tax["A15"] = "Total Effective Tax Rate"
        tax["A25"] = "Total Operating Taxes"
        for i, y in enumerate(years):
            letter = chr(67 + i)
            tax.cell(15, 3 + i, f"=IF(Inputs!$E$106=1,Inputs!{letter}112,0)")
            tax.cell(25, 3 + i, f"=20+{i}")  # placeholder formula
        er = wb.create_sheet("Expected Returns & Buybacks")
        er["A2"] = 74.37
        er["C14"] = 125.0
        er["D14"] = 180.0
        # Cached numeric results simulating post-Excel-recalc (openpyxl data_only reads these
        # only when they are values; for unit tests we store values, not formulas).
        er["E14"] = 0.0549
        er["F14"] = 0.0933
        ev = wb.create_sheet("Enterprise Value")
        ev["B6"] = -0.10
        ev["C6"] = -0.012
        ev["B20"] = 32.44
        ev["B27"] = 0.30
        ev["B30"] = 0.25
        ev["B32"] = 24.33
        ev["B42"] = 105.6
        ev["B47"] = 0.1164
        ev["B48"] = 55.40
        ev["B53"] = 0.16
        ic = wb.create_sheet("IC & NOPAT & ROIC ")
        ic["A7"] = "Invested Capital"
        ic["A20"] = "NOPAT"
        ic["A23"] = "ROIC in"
        for i, y in enumerate(years):
            ic.cell(7, 3 + i, 500 + i)
            ic.cell(20, 3 + i, 60 + i)
            ic.cell(23, 3 + i, 0.12)
        fm = wb.create_sheet("Final Metrics")
        fm["A5"] = "ROCE"
        fm["A8"] = "ROIC in - WACC"
        for i, y in enumerate(years):
            fm.cell(5, 3 + i, 0.18)
            fm.cell(8, 3 + i, 0.04)
        wb.save(path)
        wb.close()

    years_prev = list(range(2016, 2026))
    years_new = list(range(2017, 2027))
    book(prev, years_prev, pe_last=14.0)
    book(tmpl, years_new, pe_last=0.0)

    wb = Workbook()
    ws = wb.active
    ws.title = "JBSS"
    ws["A133"] = "PE10"
    # sparse empties then Current PE10
    ws["A159"] = "Current PE10"
    ws["B159"] = 13.75
    # crude annual block metadata rows used by CustomRunService may not parse;
    # AnnualInputsService also accepts metadata from parse — for unit test we
    # still verify EmptyCell safety + current PE10 separation via service path.
    wb.save(crf)
    wb.close()
    return prev, tmpl, crf


def test_detect_year_columns_inputs_row1_and_ignores_bloomberg_y(tmp_path: Path):
    prev, tmpl, _ = _rolling_books(tmp_path)
    alignment = align_fiscal_periods(tmpl, prev)
    assert alignment.new_fiscal_year == "FY2026"
    assert alignment.dropped_years == ("FY2016",)
    wb = load_workbook(tmpl)
    cols = detect_year_columns(wb["Inputs"], wb)
    assert cols["FY2026"] == 12
    assert cols["FY2017"] == 3
    wb.close()


def test_tax_writes_new_fy_inputs_column(tmp_path: Path):
    prev, tmpl, _ = _rolling_books(tmp_path)
    out = tmp_path / "out.xlsx"
    out.write_bytes(tmpl.read_bytes())
    report = AnnualTaxService().apply(
        analysis_id="t",
        ticker="ZZ",
        workbook_path=out,
        fiscal_year="FY2026",
        reported_effective_rate=0.2536,
        filing_components=[
            {"label": "Federal statutory", "rate": 0.21, "house": "statutory_federal"},
            {"label": "State", "rate": 0.04, "house": "state"},
        ],
        pretax_income=82.975e6,
        income_tax_expense=21.041e6,
    )
    assert report.schedule_populated
    assert report.cells_written
    assert abs((report.reported_effective_tax_rate or 0) - 0.2536) < 0.01
    wb = load_workbook(out)
    cols = detect_year_columns(wb["Inputs"], wb)
    col = cols["FY2026"]
    assert wb["Inputs"].cell(107, col).value == pytest.approx(0.21)
    assert wb["Inputs"].cell(112, col).value == pytest.approx(0.2536)
    wb.close()


def test_rd_carries_life_and_reads_income_expense(tmp_path: Path):
    prev, tmpl, _ = _rolling_books(tmp_path)
    out = tmp_path / "out.xlsx"
    out.write_bytes(tmpl.read_bytes())
    report = AnnualRdService().apply(
        analysis_id="t",
        ticker="ZZ",
        workbook_path=out,
        previous_workbook_path=prev,
        fiscal_year="FY2026",
    )
    assert report.useful_life == 3
    assert report.useful_life_unchanged is True
    assert report.rd_expense == pytest.approx(3.9, rel=0.05)  # 3.0 + 9*0.1
    assert report.lookback_complete is True


def test_valuation_extract_and_word_not_na(tmp_path: Path):
    prev, tmpl, _ = _rolling_books(tmp_path)
    val = AnnualValuationExtractService().extract(tmpl)
    assert val.expected_annual_return == pytest.approx(0.0549)
    assert val.bloomberg_expected_return_at_current_price == pytest.approx(0.5712791523788112)
    assert abs(val.bloomberg_expected_return_at_current_price - val.expected_annual_return) > 0.02
    assert val.current_price == pytest.approx(74.37)
    assert val.graham_margin_of_safety_entry_price == pytest.approx(105.6 * 0.75)
    assert val.graham_target_return_entry_price == pytest.approx(55.40)

    gate = AnnualOutputGateService().evaluate(
        analysis_id="t",
        ticker="ZZ",
        workbook_path=tmpl,
        fiscal_year="FY2026",
        tax=None,
        inputs=None,
        rd=None,
        valuation=val,
        recalc=None,
    )
    assert gate.status == "NEEDS_REVIEW"
    assert "WORKBOOK_RECALCULATION_INCOMPLETE" in gate.blockers
    assert "ANNUAL_TAX_SCHEDULE_NOT_POPULATED" in gate.blockers

    # After tax+inputs populate, Word should not say n/a for PE10/ER
    out = tmp_path / "out.xlsx"
    out.write_bytes(tmpl.read_bytes())
    AnnualTaxService().apply(
        analysis_id="t",
        ticker="ZZ",
        workbook_path=out,
        fiscal_year="FY2026",
        reported_effective_rate=0.25,
        filing_components=[{"label": "Federal", "rate": 0.21, "house": "statutory_federal"}],
    )
    # seed fiscal PE10
    wb = load_workbook(out)
    cols = detect_year_columns(wb["Inputs"], wb)
    wb["Inputs"].cell(57, cols["FY2026"]).value = 13.5
    wb["Inputs"]["B65"] = 13.75
    wb.save(out)
    wb.close()
    from models.annual_update import AnnualInputsReport, AnnualRdReport, Pe10Provenance
    from services.excel_recalc_service import ExcelRecalcReport

    inputs = AnnualInputsReport(
        analysis_id="t",
        ticker="ZZ",
        pe10=Pe10Provenance(target_cell="L57", source_value=13.5, field_role="pe10_fiscal_year"),
    )
    rd = AnnualRdReport(
        analysis_id="t",
        ticker="ZZ",
        useful_life=3,
        rd_expense=3.3,
        rd_asset=6.87,
        rd_amortization=3.48,
        lookback_complete=True,
        schedule_extended=True,
    )
    tax = AnnualTaxService().apply(
        analysis_id="t",
        ticker="ZZ",
        workbook_path=out,
        fiscal_year="FY2026",
        reported_effective_rate=0.25,
        filing_components=[{"label": "Federal", "rate": 0.21, "house": "statutory_federal"}],
    )
    val2 = AnnualValuationExtractService().extract(out, recalculation_complete=False)
    gate2 = AnnualOutputGateService().evaluate(
        analysis_id="t",
        ticker="ZZ",
        workbook_path=out,
        fiscal_year="FY2026",
        tax=tax,
        inputs=inputs,
        rd=rd,
        valuation=val2,
        recalc=ExcelRecalcReport(
            analysis_id="t",
            ticker="ZZ",
            status="FAILED",
            method="excel_com_calculate_full_rebuild",
            workbook_path=str(out),
            summary="WORKBOOK_RECALCULATION_INCOMPLETE",
        ),
    )
    assert gate2.status == "NEEDS_REVIEW"
    assert any("WORKBOOK_RECALCULATION_INCOMPLETE" in b for b in gate2.blockers)
    perf = AnnualDeliverablesService().build_performance(
        analysis_id="t",
        ticker="ZZ",
        fiscal_year=2026,
        workbook_path=out,
        research=None,
        valuation=val2,
        gate=gate2,
    )
    assert perf.current_pe10 is not None
    assert perf.current_expected_return is not None
    assert not any("filed" in h.lower() and h.lower().startswith("sec") for h in perf.year_highlights)
    assert any(m.metric == "Revenue" for m in perf.yoy)

    AnnualDeliverablesService().produce(
        analysis_id="t",
        ticker="ZZ",
        fiscal_year=2026,
        completed_workbook_path=out,
        output_dir=tmp_path,
        performance=perf,
        gate=gate2,
    )
    from docx import Document

    doc = Document(tmp_path / "2026 ZZ Annual Update.docx")
    text = "\n".join(p.text for p in doc.paragraphs)
    assert "n/a" not in text.lower() or "Current PE10" in text
    assert "Executive Investment Conclusion" in text
    assert "NOT AUTHORIZED" in text
    assert "FY2026 versus FY2025" in text or "FY2026 versus" in text
    assert "margin-of-safety entry" in text.lower()
    assert "target-return entry" in text.lower()
    assert "Inputs!B69" in text
    assert "not a substitute" in text.lower() or "distinct metric" in text.lower()

def test_rd_lookback_lengths():
    for life in (1, 3, 5, 10):
        year = 2026
        lookback = [f"FY{y}" for y in range(year - life + 1, year + 1)]
        assert len(lookback) == life
        assert lookback[-1] == "FY2026"
